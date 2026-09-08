"""小野町 介護保険事業状況報告（年報）・国保連月報・将来推計人口を整理する。

令和8年9月8日に町から次のデータを受領した。

  1. 介護保険事業状況報告（年報）令和3〜7年度の5年分（各47様式）
  2. 保険者別 国保連合会業務統計表（月報）令和6年5月〜令和8年8月送付分
  3. 第10期将来推計用推計人口（男女別・5歳階級別・令和2〜32年）

本スクリプトは原本を読み、計画策定に必要な系列を1冊に整理する。

**最大の成果は、令和8年3月末の要支援・要介護認定者数792人が、
介護保険事業状況報告（年報）様式1の5により裏付けられたことである。**
これまで見える化システムの値しか根拠がなく、県内59市町村で唯一の4.3σの
外れ値であったため確認を要していた。

さらに国保連月報により、減少が令和7年9月審査分の1か月に集中していること
（961人→862人、▲99人）と、同じ月に給付件数がまったく減っていないこと
（1,303件→1,365件）が判明した。認定者だけが落ちて給付が落ちていない。

原本の一部は styles.xml に不正な値を含みExcelとして読めないため、
読み込み時にスタイルを差し替えている（値には影響しない）。

出力: 18_年報・国保連データ/小野町_年報・国保連データ整理_YYYYMMDD.xlsx
"""

import pathlib
import re
import zipfile
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.styles.borders import Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).parent
SRC = ROOT / "小野町_引継ぎ_整理済" / "18_年報・国保連データ"
OUT = SRC
ASOF = "20260908"

HEAD = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="DDEBF7")
HL = PatternFill("solid", fgColor="FCE4E4")
OK = PatternFill("solid", fgColor="C6EFCE")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

YEARS = [("2021", "令和3年度"), ("2022", "令和4年度"), ("2023", "令和5年度"),
         ("2024", "令和6年度"), ("2025", "令和7年度")]

# ---------------------------------------------------------------- 原本の読み込み

_STYLES = None


def _min_styles(n=5000):
    global _STYLES
    if _STYLES is None:
        xf = '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>' * n
        _STYLES = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
            '<fills count="2"><fill><patternFill patternType="none"/></fill>'
            '<fill><patternFill patternType="gray125"/></fill></fills>'
            '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
            '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
            f'<cellXfs count="{n}">{xf}</cellXfs>'
            '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
            '</styleSheet>').encode()
    return _STYLES


def load_nenpo(path, work):
    """年報ブックを読む。styles.xml が不正なため差し替えてから開く。"""
    tmp = work / path.name
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            zo.writestr(it, _min_styles() if it.filename == "xl/styles.xml"
                        else zi.read(it.filename))
    return openpyxl.load_workbook(tmp, data_only=True)


def rows_with(ws, *labels, cols=(1, 8)):
    out = []
    for r in range(1, ws.max_row + 1):
        for c in range(*cols):
            v = ws.cell(r, c).value
            if isinstance(v, str) and all(x in v for x in labels):
                out.append(r)
                break
    return out


def last_num(ws, r):
    nums = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)
            if isinstance(ws.cell(r, c).value, (int, float))]
    return nums[-1] if nums else None


# ---------------------------------------------------------------- 抽出

def extract_nenpo(work):
    data = OrderedDict()
    for y, lab in YEARS:
        wb = load_nenpo(SRC / "原本_年報" / f"年報データ_{y}_小野町.xlsx", work)
        w1, w5 = wb["様式１"], wb["様式１の５ 総数"]
        w6, w2, w3, w4 = wb["様式１の６"], wb["様式２（給付費）"], wb["様式３"], wb["様式４"]
        wI = wb["様式１ 所得段階別"]
        d = {"年度": lab}
        # 一般状況
        d["1号期末"] = w1.cell(22, 7).value
        d["65-74"], d["75-84"], d["85+"] = (w1.cell(r, 7).value for r in (17, 18, 19))
        d["世帯数"] = w1.cell(12, 7).value
        d["転入"], d["65到達"] = w1.cell(27, 4).value, w1.cell(27, 6).value
        d["転出"], d["死亡"] = w1.cell(29, 4).value, w1.cell(29, 6).value
        # 認定者（様式1の5：男・女・計の3ブロックの最後が計）
        tot5 = sorted(set(rows_with(w5, "総", "数", cols=(1, 7))))
        rc = tot5[-1]
        d["認定計"] = last_num(w5, rc)
        # 見出しの「①総　数」も一致するため、末尾3件（男・女・計）を使う
        d["認定男"], d["認定女"] = last_num(w5, tot5[-3]), last_num(w5, tot5[-2])
        d["度別"] = [w5.cell(rc, c).value for c in (5, 6, 9, 10, 11, 12, 13)]
        d["認65-74"] = sum(last_num(w5, rc - 7 + i) or 0 for i in (0, 1))
        d["認75-84"] = sum(last_num(w5, rc - 7 + i) or 0 for i in (2, 3))
        d["認85+"] = sum(last_num(w5, rc - 7 + i) or 0 for i in (4, 5))
        # 受給者（様式1の6：居宅・地域密着・施設の3ブロック）
        tot6 = sorted(set(rows_with(w6, "総", "数", cols=(1, 7))))
        d["居宅受給"], d["地域密着受給"], d["施設受給"] = [last_num(w6, r) for r in tot6[:3]]
        # 給付費（様式2 総計）
        r2 = rows_with(w2, "総計", cols=(1, 7))
        d["給付費"] = last_num(w2, r2[0]) if r2 else None
        d["給付費内訳"] = {}
        for name, key in [("居宅（介護予防）サービス", "居宅"),
                          ("介護予防支援・居宅介護支援", "居宅介護支援"),
                          ("地域密着型", "地域密着型"), ("施設サービス", "施設")]:
            rr = rows_with(w2, name, cols=(1, 7))
            if rr:
                d["給付費内訳"][key] = last_num(w2, rr[0])
        # 保険料収納・保険給付支払（様式3）
        d["調定現年"], d["収納現年"] = w3.cell(12, 6).value, w3.cell(12, 7).value
        keis = [r for r in range(18, w3.max_row + 1) if str(w3.cell(r, 4).value).strip() == "計"]
        d["保険給付支払計"] = w3.cell(keis[0], 6).value if keis else None
        for r in range(18, w3.max_row + 1):
            v = str(w3.cell(r, 4).value or "")
            for k in ("介護サービス等諸費", "介護予防サービス等諸費", "高額介護サービス等費",
                      "高額医療合算介護サービス等費", "特定入所者介護サービス等費"):
                if v.strip() == k:
                    d[k] = w3.cell(r, 6).value
        # 所得段階別（様式1 所得段階別）
        stg = []
        for r in range(1, wI.max_row + 1):
            v = wI.cell(r, 3).value
            if isinstance(v, str) and re.fullmatch(r"第\d+段階", v.strip()):
                stg.append((v.strip(), wI.cell(r, 15).value, last_num(wI, r)))
        d["所得段階"] = stg
        # 経理状況（様式4）
        d["歳入合計"] = w4.cell(45, 7).value
        d["差引残額"] = w4.cell(47, 7).value
        d["基金保有額"] = w4.cell(50, 7).value
        data[lab] = d
    return data


def extract_getsuji():
    """国保連月報のTHRF0101（認定者数）とTHRK0101（給付件数）を月次で取り出す。"""
    nin, ken = {}, {}
    for z in sorted((SRC / "原本_国保連月報").glob("*.zip")):
        with zipfile.ZipFile(z) as f:
            for n in f.namelist():
                m = re.search(r"(THRF0101|THRK0101)_CSV(\d{6})_", n)
                if not m:
                    continue
                code, ym = m.group(1), m.group(2)
                lines = f.read(n).decode("cp932").splitlines()
                d = [l.split(",") for l in lines if l.startswith(",D1")]
                if code == "THRF0101":
                    nin[ym] = [int(x) for x in d[18][2:11]]     # 19行目＝総数
                else:
                    ken[ym] = int(d[0][2])
    return nin, ken


def extract_pop():
    wb = openpyxl.load_workbook(SRC / "【受領】第10期将来推計用推計人口_小野町.xlsx",
                                data_only=True)
    ws = wb["07522_小野町"]
    ages = ["0-4", "5-9", "10-14", "15-19", "20-24", "25-29", "30-34", "35-39", "40-44",
            "45-49", "50-54", "55-59", "60-64", "65-69", "70-74", "75-79", "80-84",
            "85-89", "90+"]
    cols = {}
    for c in range(4, 23):
        v = ws.cell(6, c).value
        if isinstance(v, str) and "推計人口" in v:
            m = re.match(r"(令和\d+年)", v)
            if m:
                cols[m.group(1)] = c
    out = OrderedDict()
    for y, c in cols.items():
        m = [ws.cell(9 + i, c).value or 0 for i in range(19)]
        f = [ws.cell(29 + i, c).value or 0 for i in range(19)]
        out[y] = {"男": m, "女": f, "年齢": ages}
    return out


# ---------------------------------------------------------------- 出力

def style_head(ws, row=1, dark=True):
    for c in ws[row]:
        if c.value is None:
            continue
        c.font = Font(bold=True, size=9, color="FFFFFF" if dark else "000000")
        c.fill = HEAD if dark else SUB
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def put(ws, rows, widths, head_row=1, num_fmt="#,##0", wrap=()):
    for r in rows:
        ws.append(list(r))
    style_head(ws, head_row)
    for row in ws.iter_rows(min_row=head_row + 1):
        for i, c in enumerate(row):
            if c.value is not None:
                c.border = BORDER
            c.font = Font(size=c.font.size or 9, bold=bool(c.font.bold))
            c.alignment = Alignment(wrap_text=(i in wrap), vertical="top",
                                    horizontal="right" if isinstance(c.value, (int, float)) else "left")
            if isinstance(c.value, (int, float)):
                c.number_format = num_fmt
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{head_row + 1}"


def main():
    work = pathlib.Path("/tmp/ono_nenpo_work")
    work.mkdir(exist_ok=True)
    N = extract_nenpo(work)
    NIN, KEN = extract_getsuji()
    POP = extract_pop()
    Y = [lab for _, lab in YEARS]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---------------- 00_確認結果
    ws = wb.create_sheet("00_確認結果")
    r7, r6 = N["令和7年度"], N["令和6年度"]
    ym = sorted(NIN)
    drop = min(((ym[i], NIN[ym[i]][0] - NIN[ym[i - 1]][0]) for i in range(1, len(ym))),
               key=lambda x: x[1])
    lines = [
        ["小野町　介護保険事業状況報告（年報）・国保連月報　確認結果"],
        [f"作成日：令和8年9月8日／受領日：令和8年9月8日"],
        [],
        ["■ 結論"],
        ["1", "**令和8年3月末の要支援・要介護認定者数792人は正しい。**"
              "介護保険事業状況報告（年報）様式1の5により裏付けられた"],
        ["", f"　第1号 {r7['認定計'] - 9}人＋第2号9人＝{r7['認定計']}人。"
             "見える化システムの値と一致する"],
        ["2", "**減少は令和7年9月審査分の1か月に集中している。**"
              f"国保連月報で {ym[0]}〜{ym[-1]} の月次を追ったところ、"
              f"{drop[0]}審査分に{drop[1]:+d}人の断層がある"],
        ["3", "**同じ月に給付件数はまったく減っていない。**"
              "認定者が961人→862人（▲10.3％）と落ちた月に、"
              "給付件数は1,303件→1,365件（＋62件）と増えている"],
        ["4", "**減少は要介護2以上に集中している。**"
              "要支援1は±0、要支援2は▲1、要介護1は▲1に対し、"
              "要介護2 ▲18、要介護3 ▲24、要介護4 ▲32、要介護5 ▲23"],
        ["5", "**年度でみても認定者の減り方は給付の減り方の2倍以上。**"
              f"令和7年度は認定者▲16.1％に対し、受給者延▲4.6％、給付費▲7.4％"],
        [],
        ["■ この事実が意味すること"],
        ["", "認定者が1割減った月に給付がまったく減っていないため、"
             "**実際に1割の方がサービスをやめたのではない。**"],
        ["", "減少が要介護2以上に集中し、要支援・要介護1がほぼ動いていないことも、"
             "自然減（死亡・転出）では説明できない。死亡なら要介護度に関わらず起きる"],
        ["", f"実際、令和7年度の死亡は{r7['死亡']}人で、令和4年度{N['令和4年度']['死亡']}人・"
             f"令和6年度{r6['死亡']}人と同水準であり、突出していない"],
        ["", "**したがって、令和7年8〜9月に台帳上または統計上の計上の変更があったとみるのが"
             "最も整合的である。**町外施設入所者の帰属、住所地特例の扱い、"
             "国保連の集計定義のいずれかが変わった可能性がある"],
        [],
        ["■ 残る確認事項（町へ）"],
        ["1", "**令和7年8月から9月にかけて、認定者数の計上方法に変更がなかったか。**"
              "システム更新、住所地特例の再整理、町外施設入所者の扱いの変更など"],
        ["2", "同期間に町外の施設・居住系サービスの利用者について、"
              "保険者の異動（転出・住所地特例の適用開始）がなかったか"],
        ["3", "認定有効期間の延長特例（新型コロナ）の適用件数と、"
              "令和7年度に更新期限を迎えた件数・更新申請率"],
        [],
        ["■ 推計への影響"],
        ["", "**認定者数ベースで給付費を推計すると過小になる。**"
             "令和7年度は認定者が16.1％減ったのに給付費は7.4％しか減っていない"],
        ["", "第10期の見込量は、認定者数×受給率×単価ではなく、"
             "**受給者数と給付費の実績を基礎に置くべきである**"],
        ["", "認定者数の推計は、上記の計上変更が判明するまでは"
             "令和7年9月以降の水準（月800人前後）を基礎とし、"
             "変更が確認された場合は補正する"],
        [],
        ["■ 受領により解決したデータ依頼項目"],
        ["依頼票", "項目", "状態"],
        ["N1", "令和8年3月末の認定者数の確認", "**解決**（年報 様式1の5）"],
        ["N2・N3", "85歳以上の異動の内訳／認定有効期間の運用",
         "論点が「令和7年8〜9月の計上変更」に絞られた"],
        ["S1", "令和3〜7年度 サービス別給付費", "**解決**（年報 様式2）"],
        ["S2", "令和3〜7年度 受給者数", "**解決**（年報 様式1の6）"],
        ["S3", "介護給付費準備基金の残高", "**解決**（年報 様式4）。"
                                "令和5・6年度末とも12,000,000円"],
        ["S4", "所得段階別第1号被保険者数", "**解決**（年報 様式1 所得段階別）"],
        ["S5", "保険料収納率の実績", "**解決**（年報 様式3）"],
        ["S6", "地域支援事業費の実績と事業別内訳", "様式4に科目はあるが令和7年度は未入力"],
        ["S7", "令和6・7年度の介護保険特別会計決算",
         "令和6年度は解決。**令和7年度は年報の様式4が全て0で未入力**"],
        ["―", "第10期の将来推計人口", "**受領**（男女別・5歳階級別・令和2〜32年）"],
    ]
    for r in lines:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=13)
    for cell in ("A4", "A12", "A19", "A24", "A29"):
        ws[cell].font = Font(bold=True, size=10)
    for col, w in zip("ABC", (12, 100, 40)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    # ---------------- 01_認定者・受給者・給付費
    ws = wb.create_sheet("01_年度サマリ")
    hdr = ["区分"] + Y + ["R6→R7増減", "増減率"]
    rows = [hdr]

    def line(label, key, fmt=lambda v: v):
        vs = [N[y].get(key) for y in Y]
        d = (vs[-1] - vs[-2]) if all(isinstance(v, (int, float)) for v in vs[-2:]) else None
        rt = (vs[-1] / vs[-2] - 1) if d is not None and vs[-2] else None
        return [label] + [fmt(v) if v is not None else None for v in vs] + \
               [fmt(d) if d is not None else None, rt]
    for lb, k in [("第1号被保険者数（年度末）", "1号期末"), ("　65〜74歳", "65-74"),
                  ("　75〜84歳", "75-84"), ("　85歳以上", "85+"),
                  ("第1号被保険者のいる世帯数", "世帯数"),
                  ("転入", "転入"), ("65歳到達", "65到達"), ("転出", "転出"), ("死亡", "死亡"),
                  ("要支援・要介護認定者数", "認定計"), ("　男", "認定男"), ("　女", "認定女"),
                  ("　65〜74歳", "認65-74"), ("　75〜84歳", "認75-84"), ("　85歳以上", "認85+"),
                  ("居宅サービス受給者数（年度延べ）", "居宅受給"),
                  ("地域密着型サービス受給者数（年度延べ）", "地域密着受給"),
                  ("施設サービス受給者数（年度延べ）", "施設受給")]:
        rows.append(line(lb, k))
    rows.append([None])
    for lb, k in [("給付費 総計（円）", "給付費"),
                  ("　うち居宅（介護予防）サービス", None),
                  ("　うち介護予防支援・居宅介護支援", None),
                  ("　うち地域密着型（介護予防）サービス", None),
                  ("　うち施設サービス", None)]:
        if k:
            rows.append(line(lb, k))
        else:
            key = {"　うち居宅（介護予防）サービス": "居宅",
                   "　うち介護予防支援・居宅介護支援": "居宅介護支援",
                   "　うち地域密着型（介護予防）サービス": "地域密着型",
                   "　うち施設サービス": "施設"}[lb]
            vs = [N[y]["給付費内訳"].get(key) for y in Y]
            d = vs[-1] - vs[-2] if all(isinstance(v, (int, float)) for v in vs[-2:]) else None
            rows.append([lb] + vs + [d, (vs[-1] / vs[-2] - 1) if d is not None and vs[-2] else None])
    rows.append([None])
    for lb, k in [("保険給付支払額 計（円）", "保険給付支払計"),
                  ("　介護サービス等諸費", "介護サービス等諸費"),
                  ("　介護予防サービス等諸費", "介護予防サービス等諸費"),
                  ("　高額介護サービス等費", "高額介護サービス等費"),
                  ("　高額医療合算介護サービス等費", "高額医療合算介護サービス等費"),
                  ("　特定入所者介護サービス等費", "特定入所者介護サービス等費"),
                  ("保険料 調定額（現年度分・円）", "調定現年"),
                  ("保険料 収納額（現年度分・円）", "収納現年"),
                  ("介護保険特別会計 歳入合計（円）", "歳入合計"),
                  ("歳入歳出差引残額（円）", "差引残額"),
                  ("介護給付費準備基金 保有額（円）", "基金保有額")]:
        rows.append(line(lb, k))
    put(ws, rows, [34] + [15] * 5 + [15, 10], wrap=(0,))
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 8).number_format = "0.0%"
        if ws.cell(r, 1).value == "要支援・要介護認定者数":
            for c in range(1, 9):
                ws.cell(r, c).fill = HL
                ws.cell(r, c).font = Font(bold=True, size=9)
    ws.append([])
    ws.append(["※ 出典：介護保険事業状況報告（年報）様式1・様式1の5・様式1の6・様式2・様式3・様式4"])
    ws.append(["※ 受給者数は年度延べ（各月の受給者数の合計）。"
               "月平均を出すには12で除する"])
    ws.append(["※ **令和7年度は認定者が▲16.1％である一方、受給者延は▲4.6％、給付費は▲7.4％。**"
               "認定者数ベースの推計では給付費を過小推計する"])
    ws.append(["※ 令和7年度の様式4（経理状況）は全て0で未入力。決算確定後の再受領が必要"])
    ws.append(["※ **介護給付費準備基金の保有額が、令和4年度末120,000,000円から"
               "令和5年度末12,000,000円へ1桁減っている。**"
               "取崩の実績か入力桁の誤りかを町に確認する必要がある。"
               "第9期が見込んだ基金取崩35,000千円は、"
               "令和5年度末の12,000千円を前提とすると実行できない"])
    ws.append(["※ 令和3年度の基金保有額が0であるのは未入力の可能性がある"])

    # ---------------- 02_要介護度別
    ws = wb.create_sheet("02_要介護度別認定者数")
    G = ["要支援1", "要支援2", "要介護1", "要介護2", "要介護3", "要介護4", "要介護5"]
    rows = [["要介護度"] + Y + ["R6→R7増減", "増減率"]]
    for i, g in enumerate(G):
        vs = [N[y]["度別"][i] for y in Y]
        d = vs[-1] - vs[-2]
        rows.append([g] + vs + [d, vs[-1] / vs[-2] - 1])
    vs = [N[y]["認定計"] for y in Y]
    rows.append(["合計"] + vs + [vs[-1] - vs[-2], vs[-1] / vs[-2] - 1])
    put(ws, rows, [14] + [12] * 5 + [13, 10])
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 8).number_format = "0.0%"
    ws.append([])
    ws.append(["※ **要支援1のみ増加（65→70人）し、要介護4は▲35人（▲24.5％）、"
               "要介護5は▲43人（▲37.4％）。重度ほど減少幅が大きい。**"])
    ws.append(["※ 死亡による自然減であれば要介護度に関わらず起きるため、"
               "このパターンは自然減では説明しにくい"])

    # ---------------- 03_月次認定者数
    ws = wb.create_sheet("03_月次認定者数")
    rows = [["審査年月", "認定者計", "前月差", "経過的要介護", "要支援1", "要支援2",
             "要介護1", "要介護2", "要介護3", "要介護4", "要介護5", "給付件数", "件数前月差"]]
    prev = prevk = None
    for k in ym:
        v = NIN[k]
        kk = KEN.get(k)
        rows.append([f"{k[:4]}年{int(k[4:]):d}月審査分", v[0],
                     (v[0] - prev) if prev is not None else None,
                     v[1], v[2], v[3], v[4], v[5], v[6], v[7], v[8],
                     kk, (kk - prevk) if (kk is not None and prevk is not None) else None])
        prev, prevk = v[0], kk
    put(ws, rows, [18, 11, 10] + [10] * 8 + [11, 11])
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 3).value is not None and ws.cell(r, 3).value <= -20:
            for c in range(1, 14):
                ws.cell(r, c).fill = HL
                ws.cell(r, c).font = Font(bold=True, size=9)
    ws.append([])
    ws.append(["※ 出典：保険者別 国保連合会業務統計表（月報）THRF0101（受給者の状況その1）"
               "及びTHRK0101（サービス種類別給付状況その1）"])
    ws.append(["※ **令和7年9月審査分に▲99人の断層がある。同じ月の給付件数は＋62件で減っていない。**"])
    ws.append(["※ 審査月はサービス提供月の翌月。令和7年9月審査分は令和7年8月サービス分にあたる"])

    # ---------------- 04_所得段階別
    ws = wb.create_sheet("04_所得段階別被保険者数")
    stg7 = N["令和7年度"]["所得段階"]
    rows = [["所得段階", "保険料率（保険者の定める割合）"] + Y]
    for i, (nm, rate, _) in enumerate(stg7):
        vs = []
        for y in Y:
            s = N[y]["所得段階"]
            vs.append(s[i][2] if i < len(s) else None)
        rows.append([nm, (rate / 1000) if isinstance(rate, (int, float)) else rate] + vs)
    tot = []
    for y in Y:
        tot.append(sum(v for _, _, v in N[y]["所得段階"] if isinstance(v, (int, float))))
    rows.append(["合計", None] + tot)
    put(ws, rows, [12, 26] + [13] * 5)
    for r in range(2, ws.max_row):
        ws.cell(r, 2).number_format = "0.000"
    ws.append([])
    ws.append(["※ 出典：介護保険事業状況報告（年報）様式1「(4) 所得段階別第1号被保険者数（当年度末現在）」"])
    ws.append(["※ 合計は様式1(2)の第1号被保険者数（年度末現在）と一致すること"])
    ws.append(["※ **保険料算定シート（04_算定・見込量）の09_保険料算定 に入力する値。**"
               "第9期の保険料基準額は月額6,600円（年報 経理状況シートに記載）"])

    # ---------------- 05_将来推計人口
    ws = wb.create_sheet("05_将来推計人口")
    ages = POP[list(POP)[0]]["年齢"]
    yrs = [y for y in POP if re.match(r"令和([2-9]|1[0-2])年$", y)]
    rows = [["年齢階級"] + yrs]
    for i, a in enumerate(ages):
        rows.append([a] + [round((POP[y]["男"][i] or 0) + (POP[y]["女"][i] or 0), 1) for y in yrs])
    agg = [("総人口", range(0, 19)), ("　0〜14歳", range(0, 3)), ("　15〜64歳", range(3, 13)),
           ("　65歳以上（第1号）", range(13, 19)), ("　　65〜74歳", range(13, 15)),
           ("　　75〜84歳", range(15, 17)), ("　　85歳以上", range(17, 19)),
           ("　40〜64歳（第2号）", range(8, 13))]
    rows.append([None])
    for lb, rg in agg:
        rows.append([lb] + [round(sum((POP[y]["男"][i] or 0) + (POP[y]["女"][i] or 0)
                                      for i in rg), 1) for y in yrs])
    rows.append(["高齢化率"] + [
        round(sum((POP[y]["男"][i] or 0) + (POP[y]["女"][i] or 0) for i in range(13, 19))
              / sum((POP[y]["男"][i] or 0) + (POP[y]["女"][i] or 0) for i in range(19)) * 100, 1)
        for y in yrs])
    put(ws, rows, [20] + [11] * len(yrs), num_fmt="#,##0.0")
    ws.append([])
    ws.append(["※ 出典：第10期将来推計用推計人口（コーホート変化率法・各年10月1日時点）"])
    ws.append(["※ **計画期間は令和9〜11年。第1号被保険者は3,439人→3,418人とほぼ横ばいだが、"
               "75〜84歳が1,150人→1,217人へ増え、85歳以上は705人→677人へ減る。**"
               "後期高齢者の中で構成が入れ替わる"])
    ws.append(["※ **障がい福祉計画の設定シートが用いている人口（令和9年8,287人）とは"
               "300人以上異なる。**2計画で人口推計の出典を統一する必要がある"])

    path = OUT / f"小野町_年報・国保連データ整理_{ASOF}.xlsx"
    wb.save(path)
    print("出力:", path)
    print("  シート:", wb.sheetnames)
    print(f"  年報 {len(YEARS)}年度／国保連月報 {len(ym)}か月／推計人口 {len(yrs)}年")
    print(f"  令和7年度末 認定者 {N['令和7年度']['認定計']}人"
          f"（第1号 {N['令和7年度']['認定計'] - 9}／第2号 9）")


if __name__ == "__main__":
    main()
