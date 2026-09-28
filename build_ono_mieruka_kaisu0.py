"""小野町 見える化システム 利用回（日）数が0のままで給付費が潰れている件。

令和8年9月28日のご指示
  「『利用回（日）数の施策反映へ』では自然体推計へすべて戻すボタンがなく、
    すべて0のままです。」

**原因が確定しました**
  利用者数は正常に出ているのに、給付費だけが激減しています。
    訪問介護       令和8年度45,553千円 → 令和9年度  2,421千円（5.3％）
    通所介護       令和8年度151,267千円 → 令和9年度 17,190千円（11.4％）
    短期入所生活介護 令和8年度 90,099千円 → 令和9年度  6,193千円（6.9％）
    **福祉用具貸与  令和8年度 33,050千円 → 令和9年度 35,387千円（107％）**
  **福祉用具貸与だけが正常です。福祉用具貸与は利用回（日）数を持たない
  サービスです。**これが決定的な証拠で、
  給付費＝利用者数×利用回（日）数×単価 の「利用回（日）数」が0のため
  給付費が潰れています。

**伸びの選択肢（3画面とも4つだけ。「0」も「全国」もない）**
  施設・居住系サービス利用率の変化（S212000）①〜④
  在宅サービス利用率の変化（S212200）      ①〜④
  在宅サービス1人1月あたり利用回（日）数の変化（S212300）①〜④
  認定率の変化（S211900）だけ⑤全国がある。

**もう1つの発見**
  「在宅サービス利用回(日)数の実績見込み値の編集」画面（令和8年度）に
  **訪問看護 要支援2 ＝ −14.0回**という負の値がある。
  介護予防訪問看護の給付費が負になっていた原因である。

出力
  11_見える化出力依頼/小野町_見える化_利用回数が0のままの件_YYYYMMDD.docx
  11_見える化出力依頼/小野町_見える化_利用回数が0のままの件_YYYYMMDD.xlsx
"""

import pathlib
import warnings

warnings.filterwarnings("ignore")

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = pathlib.Path(__file__).parent
OUT = ROOT / "小野町_引継ぎ_整理済" / "11_見える化出力依頼"
ASOF = "20260928"
ASOF_JP = "令和8年9月28日"
JP_MIN = "游明朝"
JP_GO = "游ゴシック"

HEAD = PatternFill("solid", fgColor="1F3864")
NG = PatternFill("solid", fgColor="FCE4E4")
WARN = PatternFill("solid", fgColor="FFE0B2")
OK = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

OUR_GETSU = 6048
          # 令和8年9月28日訂正。調整交付金見込額の基数に総合事業費を
          # 含めていなかった（6,095円→6,048円）。切上げ6,100円は不変。

# 給付費と利用者数の対比（ワークシート 2_サービス別給付費）
SHOKO = [
    ("訪問介護", "回", 45552.7, 2421, 67.0, 65.0),
    ("訪問看護", "回", 20605.1, 3595, 46.0, 50.0),
    ("通所介護", "回", 151267.2, 17190, 169.0, 168.0),
    ("短期入所生活介護", "日", 90099.3, 6193, 60.0, 61.0),
    ("**福祉用具貸与**", "**なし**", 33050.4, 35387, 204.0, 210.0),
]

# 伸びの選択肢（3画面で共通）
SENTAKUSHI = [
    ("①", "当該市町村における令和3年度→令和8年度の伸び", "5年",
     "**施設・居住系の利用率は合計すべて0.0で、実質「変化なし」と同じ。**"
     "在宅の利用率も訪問介護▲0.3％と穏当。**戻す場合の第一候補**"),
    ("②", "当該市町村における令和7年度→令和8年度の伸び", "1年",
     "避ける。令和8年度は国保連月報の1か月分"),
    ("③", "当該市町村における令和6年度→令和8年度の伸び", "2年",
     "避ける。同じく令和8年度を含む"),
    ("④", "当該市町村における令和6年度→令和7年度の伸び", "1年",
     "**いま設定されている。**令和7年8〜9月の認定者数の断層が入る。"
     "在宅サービスの利用率が下がり続ける向きに効いている"),
]

# 直し方
AN = [
    ("A案", "「利用者数の施策反映へ」の画面で"
     "「自然体推計に全て戻す」を押す",
     "**まずこれをお試しください。**"
     "利用回（日）数の画面にボタンがないのは、"
     "同じ「在宅サービス利用者数等」の画面群の"
     "利用者数側のボタンが両方に効く作りだからかもしれません",
     "1分"),
    ("B案", "利用回（日）数の画面で、"
     "**令和7年度の行の値を令和9〜11年度にそのまま写す**",
     "**確実です。**画面に令和7年度の行がグレーで表示されています。"
     "サービスを1つずつ切り替えて、実績のある8サービスだけ入れれば足ります。"
     "**これは「変化なし」と同じ結果になり、当方の算定とも整合します**",
     "1時間"),
    ("C案", "厚生労働省のヘルプデスクに問い合わせる",
     "画面右上の「操作に困った時は」から。"
     "**「在宅サービス利用回（日）数の施策反映が0のまま戻せない」**"
     "とお伝えください",
     "―"),
]

# B案で入れるサービス
B_SVC = [
    ("訪問介護", "回", "**要介護1〜5の5区分。**"
     "令和7年度は9.2／12.1／19.4／23.7／39.1"),
    ("訪問入浴介護", "回", "要支援1〜要介護5の7区分"),
    ("訪問看護", "回", "**要支援2の令和8年度が▲14.0回。**"
     "令和7年度の値を使ってください"),
    ("通所介護", "回", "要介護1〜5の5区分"),
    ("通所リハビリテーション", "回", "要介護1〜5の5区分"),
    ("短期入所生活介護", "日", "要介護1〜5の5区分"),
    ("短期入所療養介護（老健）", "日", "実績のある区分のみ"),
    ("認知症対応型通所介護", "回", "実績のある区分のみ"),
]

# 令和8年度の実績見込み値で直すもの（画面から読み取り）
R8_NAOSU = [
    ("訪問看護", "要支援2", -14.0, "**負の値。0.0に直す**", "最優先"),
    ("訪問介護", "要介護5", 61.3, "令和7年度は39.1回", "高"),
    ("短期入所生活介護", "要介護1", 18.6, "令和7年度の水準へ", "高"),
    ("短期入所生活介護", "要介護4", 22.9, "同上", "高"),
    ("通所リハビリテーション", "要介護5", 13.3, "同上", "中"),
]

MEYASU = [
    ("在宅サービス（令和9年度）", "209,069千円", "**400,000〜520,000千円**"),
    ("居住系サービス（令和9年度）", "144,729千円", "140,000〜155,000千円（妥当）"),
    ("施設サービス（令和9年度）", "387,688千円", "380,000〜440,000千円（妥当）"),
    ("総給付費（令和9年度）", "741,486千円", "**950,000〜1,070,000千円**"),
    ("保険料基準額（月額）", "4,829円", f"**6,000円台**（当方{OUR_GETSU:,}円）"),
]


# ---------------------------------------------------------------- 体裁

def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = JP_MIN
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), JP_MIN)
    return doc


def head(doc, text, level):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = JP_GO
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_GO)


def _emph(p, text):
    for i, part in enumerate(str(text).split("**")):
        if not part:
            continue
        r = p.add_run(part)
        r.font.name = JP_MIN
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_MIN)
        r.bold = bool(i % 2)


def body(doc, *paras):
    for t in paras:
        _emph(doc.add_paragraph(), t)


def note(doc, text):
    p = doc.add_paragraph()
    _emph(p, text)
    for r in p.runs:
        r.font.size = Pt(9)


def table(doc, header_row, rows, right_from=99):
    t = doc.add_table(rows=1, cols=len(header_row))
    t.style = "Table Grid"
    for i, h in enumerate(header_row):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = JP_GO
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_GO)
        r.bold = True
        r.font.size = Pt(9)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            _emph(p, v)
            for r in p.runs:
                r.font.size = Pt(9)
            if i >= right_from:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()


def style_head(ws, row=1):
    for c in ws[row]:
        if c.value is None:
            continue
        c.fill = HEAD
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        c.border = BORDER


def body_style(ws, first=2, wrap=()):
    for r in ws.iter_rows(min_row=first):
        for c in r:
            c.font = Font(size=9)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=c.column in wrap)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


# ---------------------------------------------------------------- Word

def build_docx():
    doc = new_doc()
    head(doc, "利用回（日）数が0のままで給付費が潰れている件", 0)
    body(doc, f"小野町　第10期介護保険事業計画　／　{ASOF_JP}")

    head(doc, "0　結論", 1)
    table(doc, ["#", "内容"], [
        ("1", "**ご指摘のとおりです。利用回（日）数が0のままのため、"
              "給付費が潰れています。**"
              "利用者数は正常に出ているのに、給付費だけが"
              "訪問介護で5.3％、通所介護で11.4％まで落ちています（第1章）"),
        ("2", "**決定的な証拠は福祉用具貸与です。**"
              "福祉用具貸与だけが令和9年度に令和8年度の107％と正常に出ています。"
              "**福祉用具貸与は利用回（日）数を持たないサービスです**"),
        ("3", "**まず「利用者数の施策反映へ」の画面で"
              "「自然体推計に全て戻す」を押してみてください。**"
              "利用回（日）数の画面にボタンがないのは、"
              "利用者数側のボタンが両方に効く作りだからかもしれません（第2章）"),
        ("4", "**戻らない場合は、利用回（日）数の画面で"
              "令和7年度の行の値を令和9〜11年度にそのまま写してください。**"
              "**これは「変化なし」と同じ結果になり、当方の算定とも整合します**"),
        ("5", "**あわせて、令和8年度の実績見込み値に負の数があります。**"
              "「在宅サービス利用回(日)数の実績見込み値の編集」画面で、"
              "**訪問看護の要支援2が▲14.0回**です。"
              "介護予防訪問看護の給付費が負になっていた原因です（第4章）"),
    ])

    # ---- 1 証拠
    head(doc, "1　利用者数は正常、給付費だけが潰れています", 1)
    table(doc, ["サービス", "回（日）数の有無", "給付費\n令和8年度",
                "給付費\n令和9年度", "比", "利用者数\n令和8年度",
                "利用者数\n令和9年度"],
          [(nm, tani, f"{k8:,.0f}千円", f"**{k9:,.0f}千円**",
            f"**{k9 / k8:.1%}**", f"{n8:,.0f}人", f"{n9:,.0f}人")
           for nm, tani, k8, k9, n8, n9 in SHOKO], right_from=2)
    body(doc,
         "**福祉用具貸与だけが107％と正常です。**"
         "福祉用具貸与は利用回（日）数を持たず、利用者数だけで給付費が決まります。"
         "**ほかのサービスは 給付費＝利用者数×利用回（日）数×単価 で計算されるため、"
         "利用回（日）数が0だと給付費が潰れます。**",
         "**利用者数はいずれも正常に出ています。**"
         "訪問介護は令和8年度67.0人に対し令和9年度65.0人、"
         "通所介護は169.0人に対し168.0人です。"
         "**つまり利用率の設定は効いており、"
         "利用回（日）数だけが入っていません。**")

    # ---- 2 直し方
    head(doc, "2　直し方", 1)
    table(doc, ["案", "すること", "内容", "目安"],
          [(a, b, c, d) for a, b, c, d in AN], right_from=3)

    head(doc, "(1)　B案で入れる値", 2)
    body(doc,
         "**画面の令和7年度の行がグレーで表示されています。"
         "その値を令和9・10・11年度にそのまま写してください。**"
         "サービスを1つずつ切り替えます。"
         "実績のある8サービスだけ入れれば足ります。")
    table(doc, ["サービス", "単位", "留意点"],
          [(a, b, c) for a, b, c in B_SVC])
    note(doc, "※ 令和9・10・11年度はいずれも令和7年度と同じ値にします。"
              "**これが「1人1月あたり利用回（日）数の変化なし」の意味です。**"
              "※ 訪問リハビリテーション・地域密着型通所介護・"
              "短期入所療養介護（病院等）（介護医療院）は"
              "小野町に利用がないため0のままで構いません。")

    # ---- 3 伸びの選択肢
    head(doc, "3　伸びの選択肢について", 1)
    body(doc,
         "**3つの画面とも選択肢は4つだけで、「0（ゼロ）」も「全国」も"
         "ありませんでした。**"
         "認定率の画面（S211900）だけに⑤全国があります。")
    table(doc, ["#", "選択肢", "窓", "評価"],
          [(a, b, c, d) for a, b, c, d in SENTAKUSHI])
    body(doc,
         "**いま④（令和6年度→令和7年度）が設定されています。**"
         "利用者数は妥当に出ているので、当面はこのままで構いません。",
         "**A案で自然体推計に戻す場合は、①に変えることをお勧めします。**"
         "施設・居住系サービスの利用率は①なら合計がすべて0.0で、"
         "実質「変化なし」と同じです。"
         "在宅サービスの利用率も①なら訪問介護▲0.3％と穏当です。",
         "**B案で値を直接入れる場合は、伸びの設定は効かなくなるため"
         "④のままで構いません。**")

    # ---- 4 令和8年度の負の値
    head(doc, "4　令和8年度の実績見込み値に負の数があります", 1)
    body(doc,
         "**「在宅サービス利用回(日)数の実績見込み値の編集」画面（令和8年度）に、"
         "訪問看護の要支援2が▲14.0回と入っています。**"
         "利用回数が負になることは現実にはあり得ません。"
         "**これが、将来推計0925で介護予防訪問看護の給付費が"
         "▲3,345千円と負になっていた原因です。**")
    table(doc, ["サービス", "区分", "令和8年度の値", "直す内容", "優先度"],
          [(a, b, f"**{c:.1f}**", d, e) for a, b, c, d, e in R8_NAOSU],
          right_from=2)
    note(doc, "※ 同じ画面の右上に「計算値に戻す」ボタンがあります。"
              "**押すと国保連月報の値（＝いまの値）に戻りますので、"
              "直したあとは押さないでください。**"
              "※ 令和8年度は国保連月報が1か月分しかありません。"
              "負の値は、過誤調整（返戻）が1か月に集中したためと考えられます。")

    # ---- 5 確認の目安
    head(doc, "5　直ったかどうかの見分け方", 1)
    table(doc, ["項目", "いま", "あるべき姿"],
          [(a, b, c) for a, b, c in MEYASU], right_from=1)
    note(doc, "※ 当方の算定は総給付費3年計2,869,936千円"
              f"（年956,645千円）、保険料月額{OUR_GETSU:,}円です。"
              "**システムは令和8年度の利用者数を基準にするため、"
              "当方より高めに出ます。**一致しなくても差し支えありません。")

    head(doc, "(1)　お願いする事項", 2)
    table(doc, ["#", "内容"], [
        ("1", "**「利用者数の施策反映へ」の画面で"
              "「自然体推計に全て戻す」を押した結果をお知らせください。**"
              "利用回（日）数も一緒に戻れば、それで解決です"),
        ("2", "**戻らない場合は、B案（手入力）に進んでください。**"
              "令和7年度の行の値を写すだけです"),
        ("3", "**訪問看護の要支援2の▲14.0回は、"
              "どちらの案を採る場合でも必ず直してください**"),
        ("4", "**直したあと、ワークシートと総括表を出力してお送りください。**"
              "当方で突き合わせます"),
    ])

    p = OUT / f"小野町_見える化_利用回数が0のままの件_{ASOF}.docx"
    doc.save(p)
    return p


# ---------------------------------------------------------------- Excel

def build_xlsx():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("00_要点")
    ws.append(["#", "内容", "判定"])
    style_head(ws)
    for i, (a, ok) in enumerate([
        ("利用回（日）数が0のままで給付費が潰れている（ご指摘のとおり）", False),
        ("利用者数は正常（訪問介護 令和8年度67.0人→令和9年度65.0人）", True),
        ("福祉用具貸与だけ正常（107％）。回（日）数を持たないサービスだから",
         True),
        ("A案：「利用者数の施策反映へ」の「自然体推計に全て戻す」を押す", None),
        ("B案：利用回（日）数の画面で令和7年度の行を令和9〜11年度に写す", None),
        ("伸びの選択肢は3画面とも4つだけ。「0」も「全国」もない", None),
        ("令和8年度の実績見込み値に負の数（訪問看護 要支援2 ▲14.0回）", False),
    ], start=1):
        ws.append([i, a, "○" if ok else ("×" if ok is False else "―")])
        ws.cell(ws.max_row, 3).fill = \
            OK if ok else (NG if ok is False else WARN)
    body_style(ws, wrap=(2,))
    widths(ws, [4, 76, 6])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("01_証拠")
    ws.append(["サービス", "回（日）数", "給付費 令和8年度", "給付費 令和9年度",
               "比", "利用者数 令和8年度", "利用者数 令和9年度"])
    style_head(ws)
    for nm, tani, k8, k9, n8, n9 in SHOKO:
        ws.append([nm.replace("**", ""), tani.replace("**", ""), k8, k9,
                   k9 / k8, n8, n9])
        ws.cell(ws.max_row, 5).fill = OK if k9 / k8 > 0.8 else NG
    for r in range(2, ws.max_row + 1):
        for c in (3, 4):
            ws.cell(r, c).number_format = "#,##0"
        ws.cell(r, 5).number_format = "0.0%"
        for c in (6, 7):
            ws.cell(r, c).number_format = "#,##0.0"
    body_style(ws)
    widths(ws, [24, 12, 18, 18, 10, 18, 18])
    ws.freeze_panes = "A2"
    ws.append([])
    ws.append(["※ 出典：ワークシート「2_サービス別給付費」。給付費は千円、"
               "利用者数は1月あたり。"])
    ws.append(["※ 福祉用具貸与は利用回（日）数を持たないため、"
               "利用者数だけで給付費が決まる。ここだけ正常なのが決定的な証拠。"])

    ws = wb.create_sheet("02_直し方")
    ws.append(["案", "すること", "内容", "目安"])
    style_head(ws)
    for a, b, c, d in AN:
        ws.append([a, b.replace("**", ""), c.replace("**", ""), d])
    body_style(ws, wrap=(2, 3))
    widths(ws, [6, 44, 74, 8])
    ws.freeze_panes = "A2"
    ws.append([])
    ws.append(["【B案で入れるサービス】"])
    ws.append(["サービス", "単位", "留意点"])
    style_head(ws, ws.max_row)
    hr = ws.max_row
    for a, b, c in B_SVC:
        ws.append([a, b, c.replace("**", "")])
    body_style(ws, hr + 1, wrap=(3,))

    ws = wb.create_sheet("03_令和8年度で直す値")
    ws.append(["サービス", "区分", "令和8年度の値", "直す内容", "優先度"])
    style_head(ws)
    for a, b, c, d, e in R8_NAOSU:
        ws.append([a, b, c, d.replace("**", ""), e])
        ws.cell(ws.max_row, 3).fill = NG
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 3).number_format = "0.0"
    body_style(ws, wrap=(4,))
    widths(ws, [24, 12, 16, 40, 10])
    ws.freeze_panes = "A2"
    ws.append([])
    ws.append(["※ 画面：在宅サービス利用回(日)数の実績見込み値の編集（令和8年度）。"])
    ws.append(["※ 右上の「計算値に戻す」を押すと国保連月報の値に戻るため、"
               "直したあとは押さない。"])

    ws = wb.create_sheet("04_伸びの選択肢")
    ws.append(["#", "選択肢", "窓", "評価"])
    style_head(ws)
    for a, b, c, d in SENTAKUSHI:
        ws.append([a, b, c, d.replace("**", "")])
        ws.cell(ws.max_row, 2).fill = OK if a == "①" else \
            (WARN if a == "④" else NG)
    body_style(ws, wrap=(2, 4))
    widths(ws, [4, 42, 6, 80])
    ws.freeze_panes = "A2"
    ws.append([])
    ws.append(["※ 施設・居住系サービス利用率の変化（S212000）、"
               "在宅サービス利用率の変化（S212200）、"
               "在宅サービス1人1月あたり利用回（日）数の変化（S212300）"
               "の3画面とも同じ4つ。"])
    ws.append(["※ 認定率の変化（S211900）だけ⑤全国における"
               "令和6年度→令和7年度の伸びがある。"])

    ws = wb.create_sheet("05_確認の目安")
    ws.append(["項目", "いま", "あるべき姿"])
    style_head(ws)
    for a, b, c in MEYASU:
        ws.append([a, b, c.replace("**", "")])
    body_style(ws, wrap=(3,))
    widths(ws, [30, 20, 44])
    ws.freeze_panes = "A2"

    p = OUT / f"小野町_見える化_利用回数が0のままの件_{ASOF}.xlsx"
    wb.save(p)
    return p


# ---------------------------------------------------------------- main

def selfcheck():
    bad = []
    fu = [x for x in SHOKO if "福祉用具" in x[0]][0]
    if fu[3] / fu[2] < 1.0:
        bad.append("福祉用具貸与が正常でない")
    for nm, _t, k8, k9, _n8, _n9 in SHOKO:
        if "福祉用具" not in nm and k9 / k8 > 0.5:
            bad.append(f"{nm} の潰れが小さい")
    if len(SENTAKUSHI) != 4:
        bad.append(f"選択肢 {len(SENTAKUSHI)}件")
    return bad


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    bad = selfcheck()
    p1 = build_docx()
    p2 = build_xlsx()
    print("出力:", p1)
    print("     ", p2)
    print("  給付費の潰れ（令和9年度÷令和8年度）")
    for nm, tani, k8, k9, n8, n9 in SHOKO:
        print(f"    {nm.replace('**',''):<16}{k9 / k8:>7.1%}"
              f"　利用者数 {n8:.0f}→{n9:.0f}人（回数{tani.replace('**','')}）")
    print(f"  自己点検 {'OK' if not bad else 'NG: ' + ' / '.join(bad)}")


if __name__ == "__main__":
    main()
