"""小野町 推計手法（基本推計／包括推計）の妥当性を国のデータで検証する。

令和8年9月25日に、見える化システムの「推計パターン毎の乖離状況」ボタンから
出力されたブック2件を受領した。**厚生労働省が全国の保険者に標準の推計方法を
当てはめ、推計値÷実績値を算定したものである。**

  推計パターン毎の乖離状況_認定者数.xlsx            令和6・7年度　1,573保険者
  推計パターン毎の乖離状況_サービス別利用者数.xlsx    令和6年度　　　1,573保険者

これにより、見込量算定の再検証（令和8年9月25日）で唯一検証できていなかった
「推計手法の選択の妥当性」が確かめられる。

**結論は2つある。**
  ① 基本推計と包括推計の優劣は年度で逆転し、小野町ではどちらとも言えない。
     ただし2年の平均では基本推計がわずかに優り、現行の設定を変える理由はない。
  ② **標準の推計方法そのものが小野町では当たらない。**
     全国1,570保険者中1,267位・1,015位で中位より下であり、
     サービス別では誤差が2桁・3桁に達するものが並ぶ。
     **当方が実績を基点として総額を月報で裏づける置き方を採っていることの
     裏づけになる。**

出力
  04_算定・見込量/小野町_第10期_推計手法の妥当性検証_YYYYMMDD.docx
  04_算定・見込量/小野町_第10期_推計手法の妥当性検証_YYYYMMDD.xlsx
"""

import pathlib
import statistics
import warnings

warnings.filterwarnings("ignore")

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = pathlib.Path(__file__).parent
BASE = ROOT / "小野町_引継ぎ_整理済"
SRC = BASE / "11_見える化出力依頼" / "原本_見える化出力_20260925"
OUT = BASE / "04_算定・見込量"
ASOF = "20260925"
ASOF_JP = "令和8年9月25日"
JP_MIN = "游明朝"
JP_GO = "游ゴシック"

ONO = "07522"          # 小野町の保険者番号
KEN = "福島県"

HEAD = PatternFill("solid", fgColor="1F3864")
NG = PatternFill("solid", fgColor="FCE4E4")
WARN = PatternFill("solid", fgColor="FFE0B2")
OK = PatternFill("solid", fgColor="E2EFDA")
CALC = PatternFill("solid", fgColor="EAF1FB")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NINTEI = SRC / "推計パターン毎の乖離状況_認定者数.xlsx"
SERVICE = SRC / "推計パターン毎の乖離状況_サービス別利用者数.xlsx"


# ---------------------------------------------------------------- 読み取り

def nintei():
    """認定者数の推計誤差。年度ごとに小野町の値と全国・県内の位置を返す。"""
    wb = openpyxl.load_workbook(NINTEI, data_only=True)
    out = {}
    for nm in wb.sheetnames:
        ws = wb[nm]
        rows = []
        for r in range(15, ws.max_row + 1):
            no = ws.cell(r, 2).value
            if not ws.cell(r, 4).value:
                continue
            rows.append({"番号": str(no), "県": str(ws.cell(r, 3).value),
                         "保険者": str(ws.cell(r, 4).value),
                         "第1号": ws.cell(r, 5).value,
                         "認定者": ws.cell(r, 6).value,
                         "基本推計": ws.cell(r, 7).value,
                         "包括推計": ws.cell(r, 8).value})
        o = next(x for x in rows if x["番号"] == ONO)
        d = {"保険者数": len(rows), "第1号": o["第1号"], "認定者": o["認定者"]}
        for k in ("基本推計", "包括推計"):
            v = o[k]
            err = abs(v - 1)
            zen = [abs(x[k] - 1) for x in rows
                   if isinstance(x[k], (int, float))]
            ken = [abs(x[k] - 1) for x in rows
                   if x["県"] == KEN and isinstance(x[k], (int, float))]
            d[k] = {"値": v, "誤差": err,
                    "全国順位": sum(1 for e in zen if e < err) + 1,
                    "全国数": len(zen), "全国中央値": statistics.median(zen),
                    "県内順位": sum(1 for e in ken if e < err) + 1,
                    "県内数": len(ken), "県中央値": statistics.median(ken)}
        out[nm] = d
    return out


def service():
    """サービス別の利用者数の推計誤差（令和6年度）。"""
    wb = openpyxl.load_workbook(SERVICE, data_only=True)
    ws = wb["令和6年度"]
    svc, cur = {}, None
    for c in range(5, ws.max_column + 1):
        v = ws.cell(6, c).value
        if v:
            cur = str(v).strip()
        k = ws.cell(7, c).value
        if cur and k:
            svc.setdefault(cur, {})[str(k).strip()] = c
    rows = {}
    for r in range(8, ws.max_row + 1):
        no = ws.cell(r, 2).value
        if no:
            rows[str(no)] = r
    ro = rows[ONO]
    out = []
    for nm, cc in svc.items():
        a = ws.cell(ro, cc["基本推計"]).value if "基本推計" in cc else None
        b = ws.cell(ro, cc["包括推計"]).value if "包括推計" in cc else None
        if not (isinstance(a, (int, float)) and isinstance(b, (int, float))):
            out.append({"サービス": nm, "基本推計": None, "包括推計": None,
                        "判定": "―", "全国順位": None, "全国数": None})
            continue
        ea, eb = abs(a - 1), abs(b - 1)
        zen = [abs(ws.cell(rr, cc["基本推計"]).value - 1) for rr in rows.values()
               if isinstance(ws.cell(rr, cc["基本推計"]).value, (int, float))]
        out.append({
            "サービス": nm, "基本推計": a, "包括推計": b,
            "判定": "基本" if ea < eb else ("包括" if eb < ea else "同"),
            "全国順位": sum(1 for e in zen if e < ea) + 1, "全国数": len(zen),
            "全国中央値": statistics.median(zen) if zen else None})
    return out


# ---------------------------------------------------------------- 体裁

def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = JP_MIN
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), JP_MIN)
    return doc


def head(doc, text, level):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = JP_GO
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_GO)


def _emph(p, text):
    for i, part in enumerate(str(text).split("**")):
        if not part:
            continue
        r = p.add_run(part)
        r.font.name = JP_MIN
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_MIN)
        r.bold = bool(i % 2)


def body(doc, *paras):
    for t in paras:
        _emph(doc.add_paragraph(), t)


def note(doc, text):
    p = doc.add_paragraph()
    _emph(p, text)
    for r in p.runs:
        r.font.size = Pt(9)


def table(doc, header_row, rows, right_from=99):
    t = doc.add_table(rows=1, cols=len(header_row))
    t.style = "Table Grid"
    for i, h in enumerate(header_row):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = JP_GO
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_GO)
        r.bold = True
        r.font.size = Pt(9)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            _emph(p, v)
            for r in p.runs:
                r.font.size = Pt(9)
            if i >= right_from:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()


def pct(v):
    return f"{(v - 1) * 100:+.2f}％" if isinstance(v, (int, float)) else "―"


# ---------------------------------------------------------------- Word

def build_docx(N, S):
    yrs = list(N)
    heikin = {k: statistics.mean(N[y][k]["誤差"] for y in yrs)
              for k in ("基本推計", "包括推計")}
    yushi = min(heikin, key=heikin.get)
    ooki = sorted([s for s in S if s["基本推計"] is not None],
                  key=lambda x: -abs(x["基本推計"] - 1))

    doc = new_doc()
    head(doc, "第10期 推計手法の妥当性検証", 0)
    body(doc, f"小野町　第10期介護保険事業計画　／　{ASOF_JP}")

    head(doc, "0　結論", 1)
    table(doc, ["#", "確かめたこと", "結果"], [
        ("1", "基本推計と包括推計のどちらが当たるか",
         "**年度で逆転しており、小野町ではどちらとも言えない。**"
         f"{yrs[0]}は包括推計が、{yrs[1]}は基本推計が当たる。"
         f"2年の平均では基本推計{heikin['基本推計'] * 100:.2f}％・"
         f"包括推計{heikin['包括推計'] * 100:.2f}％で"
         f"**{yushi}がわずかに優る。現行の設定（基本推計）を変える理由はない**"),
        ("2", "標準の推計方法は小野町で当たるか",
         f"**当たらない。**認定者数の誤差は全国"
         f"{N[yrs[0]]['基本推計']['全国順位']:,}位／"
         f"{N[yrs[0]]['基本推計']['全国数']:,}保険者、"
         f"{N[yrs[1]]['基本推計']['全国順位']:,}位／"
         f"{N[yrs[1]]['基本推計']['全国数']:,}保険者で、"
         f"**いずれも中位より下**である"),
        ("3", "サービス別ではどうか",
         f"**さらに外れる。**誤差が10％を超えるサービスが"
         f"{sum(1 for s in ooki if abs(s['基本推計'] - 1) > 0.10)}件あり、"
         f"最大は{ooki[0]['サービス']}の{pct(ooki[0]['基本推計'])}である"),
        ("4", "当方の算定への含意",
         "**当方は標準の推計方法（認定者数×利用率を積み上げる3段連鎖）ではなく、"
         "令和7年度の実績を基点として総額を国保連月報で裏づける置き方を"
         "採っている。**"
         "標準の推計方法がこれだけ外れる小野町では、"
         "実績を基点とする置き方のほうが安全である。"
         "**本検証は当方の算定の方針を支持する**"),
    ])

    head(doc, "1　この資料の位置づけ", 1)
    body(doc,
         "**見える化システムの「推計方法の設定」画面にある"
         "「推計パターン毎の乖離状況」ボタンから出力されたものです。**"
         "厚生労働省が全国1,573保険者に標準の推計方法を当てはめ、"
         "推計値÷実績値を算定したものです。",
         "**これまで、推計手法の妥当性を確かめるには"
         "厚生労働省の保険者別配布シートが要ると申し上げてきましたが、"
         "システム内に同じものがありました。**"
         "見込量算定の再検証（本日付）で唯一検証できていなかった項目が"
         "これで埋まります。")
    table(doc, ["区分", "内容"], [
        ("推計値", "令和5年9月の第1号被保険者数・認定者数を起点として、"
                 "各年度9月の値を推計したもの"),
        ("実績値", "厚生労働省 介護保険事業状況報告"),
        ("基本推計", "性別・年齢5歳階級別・要介護度別に推計"),
        ("包括推計", "性別・年齢前期・後期別に推計"),
    ])

    head(doc, "2　認定者数の推計誤差", 1)
    table(doc, ["年度", "第1号被保険者数", "認定者数",
                "基本推計", "包括推計"],
          [(y, f"{N[y]['第1号']:,}人", f"{N[y]['認定者']:,}人",
            f"**{pct(N[y]['基本推計']['値'])}**",
            f"**{pct(N[y]['包括推計']['値'])}**") for y in yrs],
          right_from=1)
    body(doc, "**全国・県内での位置は次のとおりです。**")
    table(doc, ["年度", "手法", "誤差", "全国順位", "全国の中央値",
                "県内順位", "県内の中央値"],
          [(y, k, f"{N[y][k]['誤差'] * 100:.2f}％",
            f"{N[y][k]['全国順位']:,}／{N[y][k]['全国数']:,}",
            f"{N[y][k]['全国中央値'] * 100:.2f}％",
            f"{N[y][k]['県内順位']}／{N[y][k]['県内数']}",
            f"{N[y][k]['県中央値'] * 100:.2f}％")
           for y in yrs for k in ("基本推計", "包括推計")], right_from=2)
    body(doc,
         f"**{yrs[0]}は包括推計"
         f"（{N[yrs[0]]['包括推計']['誤差'] * 100:.2f}％）が基本推計"
         f"（{N[yrs[0]]['基本推計']['誤差'] * 100:.2f}％）を上回り、"
         f"{yrs[1]}は逆に基本推計"
         f"（{N[yrs[1]]['基本推計']['誤差'] * 100:.2f}％）が包括推計"
         f"（{N[yrs[1]]['包括推計']['誤差'] * 100:.2f}％）を上回ります。**"
         "他町村の案件では両年度とも基本推計が明確に優っていましたが、"
         "**小野町では年度で逆転しており、どちらが当たるとも言えません。**",
         f"2年の平均では基本推計{heikin['基本推計'] * 100:.2f}％・"
         f"包括推計{heikin['包括推計'] * 100:.2f}％で"
         f"{yushi}がわずかに優ります。"
         "**現行の設定は基本推計であり、これを変える理由はありません。**",
         "**より重要なのは、小野町がいずれの年度も全国の中位より下にあることです。**"
         f"全国の|誤差|中央値は{yrs[0]}"
         f"{N[yrs[0]]['基本推計']['全国中央値'] * 100:.2f}％・"
         f"{yrs[1]}{N[yrs[1]]['基本推計']['全国中央値'] * 100:.2f}％であるのに対し、"
         f"小野町は{N[yrs[0]]['基本推計']['誤差'] * 100:.2f}％・"
         f"{N[yrs[1]]['基本推計']['誤差'] * 100:.2f}％です。"
         "**標準の推計方法は、小野町ではあまり当たりません。**")

    head(doc, "3　サービス別の利用者数の推計誤差", 1)
    body(doc,
         "**認定者数よりさらに外れます。**"
         f"誤差が大きい順に並べます（{yrs[0]}・基本推計）。")
    table(doc, ["サービス", "基本推計", "包括推計", "当たるのは", "全国順位"],
          [(s["サービス"],
            f"**{pct(s['基本推計'])}**" if abs(s["基本推計"] - 1) > 0.10
            else pct(s["基本推計"]),
            pct(s["包括推計"]), s["判定"],
            f"{s['全国順位']:,}／{s['全国数']:,}")
           for s in ooki], right_from=1)
    n10 = sum(1 for s in ooki if abs(s["基本推計"] - 1) > 0.10)
    kih = sum(1 for s in S if s["判定"] == "基本")
    hou = sum(1 for s in S if s["判定"] == "包括")
    body(doc,
         f"**誤差が10％を超えるサービスが{n10}件あります。**"
         f"最大は{ooki[0]['サービス']}の{pct(ooki[0]['基本推計'])}、"
         f"次が{ooki[1]['サービス']}の{pct(ooki[1]['基本推計'])}です。"
         "利用者数の少ないサービスは1人の増減で比が大きく動くため、"
         "小規模保険者では避けられない面があります。",
         f"**基本推計が当たるのが{kih}件、包括推計が当たるのが{hou}件で拮抗しています。**"
         "サービス別に見ても、どちらかが一貫して優るということはありません。")
    note(doc, "※ 「―」は実績が0であるか推計値が得られないサービス。"
              "※ 全国順位は|誤差|の小さい順。"
              "**※ 介護医療院の▲100.00％は、標準の推計方法が0人と推計する一方で"
              "実績があることを示す。**"
              "第9期計画も介護医療院を0人と見込んでいたが、"
              "令和7年度の実績は月1.2人・4,310千円である"
              "（第9期計画の見込みと実績の対比）。"
              "**当方は令和7年度の実績を基点としているため、この取りこぼしは生じない。**")

    head(doc, "4　当方の算定への含意", 1)
    body(doc,
         "**当方の算定は、標準の推計方法とは別の経路を採っています。**")
    table(doc, ["区分", "標準の推計方法", "当方の算定"], [
        ("組み立て方",
         "認定者数 × 利用率 × 1人1月あたり回（日）数 × 単価を積み上げる",
         "**令和7年度のサービス別実績（年報 様式2）を基点とし、"
         "区分別の係数を乗じる**"),
        ("令和8年度の置き方",
         "月報（システムの設定では令和8年5月の1か月）",
         "**国保連月報3か月（令和8年4〜6月提供分）の前年同期比0.9598**"),
        ("総額の裏づけ", "なし（積み上げの結果）",
         "**月報の実績で総額を裏づける**"),
    ])
    body(doc,
         "**標準の推計方法がこれだけ外れる小野町では、"
         "実績を基点として総額を裏づける置き方のほうが安全です。**"
         "積み上げの各段（認定率・利用率・回数・単価）に誤差が乗ると、"
         "サービス別に見たとおりの振れが総額に及びます。",
         "**ただし、当方の置き方にも限界があります。**"
         "サービスの構成と1人当たりの水準を令和7年度のまま動かさないため、"
         "構成の変化を見込めません。"
         "**要介護度別の構成変化を織り込む方式（他町村で用いているもの）は"
         "バックテストで検討しましたが、令和7年9月の断層のため採れませんでした。**"
         "断層の計上方法が確認できれば、再度検討する余地があります。")

    head(doc, "5　協議会にお諮りする事項", 1)
    table(doc, ["No.", "事項", "受託者の案"], [
        ("1", "推計手法を基本推計のままとするか",
         f"**そのままとする。**2年の平均では{yushi}がわずかに優り、"
         "変える理由がない"),
        ("2", "標準の推計方法によらず実績を基点とすることの是非",
         "**実績を基点とする。**標準の推計方法は小野町では全国中位より下であり、"
         "サービス別では2桁・3桁の誤差が並ぶ"),
        ("3", "令和8年度の置き方",
         "**国保連月報3か月の実績による。**"
         "システムの設定は令和8年5月の1か月であり、"
         "1か月の振れが3年間に効く"),
    ])

    head(doc, "6　町にご確認いただきたいこと", 1)
    table(doc, ["#", "内容", "優先度"], [
        ("1", "**令和7年度の乖離状況（サービス別）。**"
              "受領したサービス別のブックは令和6年度のみである", "B"),
        ("2", "令和7年8〜9月の認定者数の計上方法（断層の原因）", "A"),
        ("3", "令和8年度の残りの月の国保連月報", "A"),
    ])
    note(doc, "※ 2が確認できれば、要介護度別の構成変化を織り込む方式を"
              "再度検討できる（第4章）。")

    OUT.mkdir(parents=True, exist_ok=True)
    p = OUT / f"小野町_第10期_推計手法の妥当性検証_{ASOF}.docx"
    doc.save(p)
    print("出力:", p)


# ---------------------------------------------------------------- Excel

def style_head(ws, row):
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row, c)
        cell.fill = HEAD
        cell.font = Font(bold=True, color="FFFFFF", size=9, name=JP_GO)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = BORDER


def xbody(ws, first, wrap=()):
    for r in range(first + 1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.font = Font(size=9, name=JP_MIN)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=c in wrap)


def widths(ws, w):
    for i, x in enumerate(w, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = x


def notes(ws, lines):
    ws.append([])
    for t in lines:
        ws.append([t])
        ws.cell(ws.max_row, 1).font = Font(size=9, italic=True, name=JP_MIN)


def build_xlsx(N, S):
    yrs = list(N)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("00_認定者数の誤差")
    ws.append([f"小野町 推計手法の妥当性検証　認定者数の推計誤差　（{ASOF_JP}）"])
    ws.cell(1, 1).font = Font(bold=True, size=13, name=JP_GO)
    ws.append(["出所：見える化システム「推計パターン毎の乖離状況」"
               "（厚生労働省が全国1,573保険者に標準の推計方法を当てはめたもの）"])
    ws.append([])
    ws.append(["年度", "第1号被保険者数", "認定者数", "手法", "推計値÷実績値",
               "誤差", "全国順位", "全国の中央値", "県内順位", "県内の中央値"])
    hr = ws.max_row
    for y in yrs:
        for k in ("基本推計", "包括推計"):
            d = N[y][k]
            ws.append([y, N[y]["第1号"], N[y]["認定者"], k, d["値"], d["誤差"],
                       f"{d['全国順位']:,}／{d['全国数']:,}", d["全国中央値"],
                       f"{d['県内順位']}／{d['県内数']}", d["県中央値"]])
    style_head(ws, hr)
    for r in range(hr + 1, ws.max_row + 1):
        ws.cell(r, 5).number_format = "0.0000"
        for c in (6, 8, 10):
            ws.cell(r, c).number_format = "0.00%"
        v = ws.cell(r, 6).value
        med = ws.cell(r, 8).value
        ws.cell(r, 6).fill = NG if v > med * 1.5 else WARN if v > med else OK
    xbody(ws, hr)
    widths(ws, [12, 16, 12, 12, 14, 10, 16, 14, 12, 14])
    notes(ws, [
        "**※ 赤＝全国中央値の1.5倍超／橙＝中央値超／緑＝中央値以下。**",
        "**※ 令和6年度は包括推計、令和7年度は基本推計が当たり、逆転している。**"
        "2年の平均では基本推計がわずかに優る。",
        "**※ いずれの年度も全国の中位より下である。**"
        "標準の推計方法は小野町ではあまり当たらない。",
    ])

    ws = wb.create_sheet("01_サービス別の誤差")
    ws.append([f"サービス別の利用者数の推計誤差（{yrs[0]}）"])
    ws.cell(1, 1).font = Font(bold=True, size=12, name=JP_GO)
    ws.append([])
    ws.append(["サービス", "基本推計", "包括推計", "当たるのは",
               "全国順位（基本推計）", "全国の中央値"])
    hr = ws.max_row
    for s in sorted(S, key=lambda x: -(abs(x["基本推計"] - 1)
                                       if x["基本推計"] is not None else -1)):
        ws.append([s["サービス"], s["基本推計"], s["包括推計"], s["判定"],
                   f"{s['全国順位']:,}／{s['全国数']:,}" if s["全国順位"] else "―",
                   s.get("全国中央値")])
    style_head(ws, hr)
    for r in range(hr + 1, ws.max_row + 1):
        for c in (2, 3):
            ws.cell(r, c).number_format = "0.0000"
        ws.cell(r, 6).number_format = "0.00%"
        v = ws.cell(r, 2).value
        if isinstance(v, (int, float)):
            e = abs(v - 1)
            ws.cell(r, 2).fill = NG if e > 0.25 else WARN if e > 0.10 else OK
    xbody(ws, hr, wrap=(1,))
    widths(ws, [36, 12, 12, 12, 20, 14])
    notes(ws, [
        "**※ 赤＝誤差25％超／橙＝10％超／緑＝10％以下。**",
        "※ 「―」は実績が0であるか推計値が得られないサービス。",
        "**※ 標準の推計方法がこれだけ外れるため、"
        "当方は実績を基点として総額を国保連月報で裏づける置き方を採っている。**",
    ])

    p = OUT / f"小野町_第10期_推計手法の妥当性検証_{ASOF}.xlsx"
    wb.save(p)
    print("出力:", p)


def main():
    N, S = nintei(), service()
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx(N, S)
    build_xlsx(N, S)
    yrs = list(N)
    for y in yrs:
        print(f"  {y} 基本推計{(N[y]['基本推計']['値'] - 1) * 100:+.2f}％"
              f"（全国{N[y]['基本推計']['全国順位']:,}/{N[y]['基本推計']['全国数']:,}）"
              f"／包括推計{(N[y]['包括推計']['値'] - 1) * 100:+.2f}％")
    n10 = sum(1 for s in S if s["基本推計"] is not None
              and abs(s["基本推計"] - 1) > 0.10)
    print(f"  サービス別で誤差10％超 {n10}件")


if __name__ == "__main__":
    main()
