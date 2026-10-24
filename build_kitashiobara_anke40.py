# -*- coding: utf-8 -*-
"""
北塩原村　アンケートの割合を有効回答40件の集計に差し替える

経緯
  令和8年10月21日に報告書用集計表が改訂され、障がい者調査の有効回答数が
  38件から40件になった。正本 第2章7 等の割合は他メンバー版の本文にあり、
  当方は差し替えを他メンバーにお願いしていた（Ｓ-39）。
  令和8年10月24日の他ＡＩの横断レビュー（RT-01・RT-10）で、
  正本・概要版・協議会資料・報告書の間で38件と40件が混在していることを
  「委員会提出前に是正が必要」と指摘され、同日、修正のご指示を受けた。
  38件の値は確定した集計表と合わない事実の誤りであるため、
  他メンバー版の本文であっても当方で差し替える（書き換えの取決めの範囲）。

何をするか
  正本のうち他メンバー版の本文にある段落と表のセルについて、
  割合が出る順に「どの指標か」を割り当てる（SASHIKAE・SASHIKAE_HYOU）。
  ・割り当てた指標の38件の値が、いまの文の値とちょうど一致することを確かめる
    （割り当ての誤りを止める。検査①）
  ・40件の値は kitashiobara_anke から読む（書き写さない）
  ・None を割り当てたものは障がい別・障がい児調査・制度の値であり動かさない
    （障がい別が動かないことは kitashiobara_anke の側で確かめている）
  当方が足した記述（anke_mikomi・ishoku_okurijo・uchiawase_ronten・
  anke_kaitei 等）は、それぞれの生成器が kitashiobara_anke から読む形に改めた。

出力
  output/北塩原村_アンケート40件への差し替え.xlsx
  正本への反映（AN40_KAKIKAE・AN40_HYOU）は build_kitashiobara_honpon.py が読み込む。
"""

import os
import re
import sys
import warnings

import docx
from openpyxl import Workbook

import kitashiobara_anke as A
from kitashiobara_common import (
    OUT_DIR, SRC_DIR, add_sheet, ensure_out_dir, style_header_row, write_row,
)

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

OUT_FILE = f"{OUT_DIR}/北塩原村_アンケート40件への差し替え.xlsx"
MEMBER = (f"{SRC_DIR}/他メンバー素案_20261008/"
          "北塩原村_第8期障がい福祉計画_第4期障がい児福祉計画_"
          "素案_他メンバー版.docx")
HONPON = f"{OUT_DIR}/北塩原村_計画素案_正本_移植後.docx"

PAT = re.compile(r"[0-9]+\.[0-9]％")

# ===========================================================================
# 1. 段落（他メンバー版の本文）
#    (番号, 箇所, 段落の書き出し, [割合が出る順の指標], [(言い回しの置換)])
# ===========================================================================
SASHIKAE = [
    ("AN40-1", "第2章7（5）調査結果の概要（冒頭）", "障がい者調査は有効回答率が",
     ["有効回答率_障がい者", "65歳以上", "身体障がい（主たる）"], []),
    ("AN40-2", "第2章7 日常生活・将来の暮らし", "日常生活で困っていることは",
     ["困りごと_特にない", "困りごと_将来の生活", "困りごと_将来の生活_65歳未満",
      "暮らし_今の自宅", "暮らし_家族と一緒", "在宅を希望"], []),
    ("AN40-3", "第2章7 外出・移動", "外出時に困ることは",
     ["外出_冬季", "外出_冬季_裏磐梯", "移動_自分で運転", "移動_家族の送迎"], []),
    ("AN40-4", "第2章7 相談", "困ったときの相談相手は",
     ["相談相手_家族", "相談相手_相談支援専門員", "相談相手_村役場",
      "相談への期待_情報提供", "相談への期待_サービス利用",
      "相談への期待_将来の暮らし"], []),
    ("AN40-5", "第2章7 福祉サービス・地域生活支援拠点", "今後３年間で利用したいサービスは",
     ["利用したい_施設入所支援", "利用したい_グループホーム", "利用したい_短期入所",
      "知らない_重度訪問介護", "拠点_緊急時の受入れ", "拠点_緊急時の相談"], []),
    ("AN40-6", "第2章7 日中活動・就労", "日中は「自宅で過ごすことが多い」",
     ["日中_自宅", "働き方_短時間", "就労支援_職場の理解",
      "就労支援_体調に合わせた勤務"], []),
    ("AN40-7", "第2章7 権利擁護", "差別や偏見、不利益を感じたことがある方は",
     ["差別を感じた", "配慮_わかりやすい説明", "配慮_手続きの支援",
      "成年後見_知らない"], []),
    ("AN40-8", "第2章7 災害", "避難行動要支援者名簿・個別避難計画を「知らない」が",
     ["個別避難計画_知らない", "災害不安_避難先での生活", "災害不安_薬・医療機器",
      "災害不安_冬季の避難"], []),
    #   40件では「相談支援体制」と「災害時支援」が同率でなくなる
    ("AN40-9", "第2章7 村が力を入れるべきこと", "「相談支援体制」「災害時支援」（各",
     ["施策_相談支援体制", "施策_在宅生活支援"],
     [("「相談支援体制」「災害時支援」（各{施策_相談支援体制}）が最も高く、",
       "「相談支援体制」（{施策_相談支援体制}）が最も高く、"
       "「災害時支援」（{施策_災害時支援}）、")]),
    ("AN40-10", "第3章 第4次計画の指標の参考値", "なお、今回のアンケート調査では、①の指標",
     ["差別を感じた", "満足（利用者）"], []),
    ("AN40-11", "第4章2（2）退院者の見込み", "北塩原村においては、第７期計画策定時に",
     ["在宅を希望", None], []),
    ("AN40-12", "第4章2（3）一般就労への移行等", "一方で、今回のアンケート調査では、今後の働き方",
     ["働き方_短時間", "働き方_一般企業", "就労支援_職場の理解",
      "就労支援_体調に合わせた勤務", None, None], []),
    ("AN40-13", "第4章2（3）就労選択支援", "令和７年10月に開始された就労選択支援については、令和９年４月",
     ["知らない_就労選択支援"], []),
    ("AN40-14", "第5章1（2）短期入所の見込量", "短期入所（福祉型）は、今回のアンケート調査で",
     ["拠点_緊急時の受入れ"], []),
    ("AN40-15", "第5章3 移動支援事業の見込量", "理解促進研修・啓発事業については、第４次",
     ["外出_冬季", None, "外出_付き添い"], []),
    ("AN40-16", "第6章 成年後見制度の周知", "今回のアンケート調査では、成年後見制度や権利擁護",
     ["成年後見_知らない", "成年後見_名前だけ", "成年後見_十分に知らない"], []),
    #   割合を含まない段落。表2-6の下の注記の回答者数（38人→40人）。
    #   割合の走査では拾えないため、令和8年10月24日の正本の走査で見つけた
    ("AN40-17", "第2章7 表2-6 アンケート調査結果の概要（注記）",
     "（障がい者調査、回答者数", [],
     [("回答者数38人", "回答者数{n}人")]),
]

# ===========================================================================
# 2. 表のセル（他メンバー版の表）
#    (番号, 箇所, 表頭の1列目, 行の1列目, 列, [指標], [(言い回しの置換)])
#    行の1列目が None のものは表頭の行そのもの
# ===========================================================================
SASHIKAE_HYOU = [
    ("AN40-T1", "第2章7 表2-x 調査の概要", "調査区分", "障がい者調査（18歳以上）", 2,
     [], [("38人", "{n}人")]),
    ("AN40-T2", "第2章7 表2-x 調査の概要", "調査区分", "障がい者調査（18歳以上）", 3,
     ["有効回答率_障がい者"], []),
    ("AN40-T3", "第2章7 表2-x 調査の概要", "調査区分", "合計", 2,
     [], [("41人", "{n_goukei}人")]),
    ("AN40-T4", "第2章7 表2-x 調査の概要", "調査区分", "合計", 3,
     ["有効回答率_合計"], []),
    ("AN40-T5", "第2章7 主な調査結果", "分野|主な調査結果", "日常生活", 2, ["困りごと_将来の生活"], []),
    ("AN40-T6", "第2章7 主な調査結果", "分野|主な調査結果", "将来の暮らし", 2, ["在宅を希望"], []),
    ("AN40-T7", "第2章7 主な調査結果", "分野|主な調査結果", "外出・移動", 2, ["外出_冬季"], []),
    ("AN40-T8", "第2章7 主な調査結果", "分野|主な調査結果", "相談@1", 2, ["相談相手_家族"], []),
    ("AN40-T9", "第2章7 主な調査結果", "分野|主な調査結果", "相談@2", 2, ["相談への期待_情報提供"], []),
    ("AN40-T10", "第2章7 主な調査結果", "分野|主な調査結果", "意思決定支援", 2, ["意思決定支援_必要"], []),
    ("AN40-T11", "第2章7 主な調査結果", "分野|主な調査結果", "福祉サービス", 2,
     ["利用したい_施設入所支援"], []),
    ("AN40-T12", "第2章7 主な調査結果", "分野|主な調査結果", "地域生活支援拠点等", 2,
     ["拠点_緊急時の受入れ"], []),
    ("AN40-T13", "第2章7 主な調査結果", "分野|主な調査結果", "就労", 2, ["就労支援_職場の理解"], []),
    ("AN40-T14", "第2章7 主な調査結果", "分野|主な調査結果", "権利擁護@1", 2, ["差別を感じた"], []),
    ("AN40-T15", "第2章7 主な調査結果", "分野|主な調査結果", "権利擁護@2", 2, ["成年後見_知らない"], []),
    ("AN40-T16", "第2章7 主な調査結果", "分野|主な調査結果", "災害", 2, ["個別避難計画_知らない"], []),
    ("AN40-T17", "第2章7 主な調査結果", "分野|主な調査結果", "村の施策", 1, [],
     [("力を入れるべきこと「相談支援体制」「災害時支援」",
       "力を入れるべきこと「相談支援体制」（次点は「災害時支援」{施策_災害時支援}）")]),
    ("AN40-T18", "第2章7 主な調査結果", "分野|主な調査結果", "村の施策", 2, ["施策_相談支援体制"],
     [("各{施策_相談支援体制}", "{施策_相談支援体制}")]),
    ("AN40-T19", "第2章7 障がい別の比較", "項目|全体", None, 1, [],
     [("（n=38）", "（n={n}）")]),
    ("AN40-T20", "第2章7 障がい別の比較", "項目|全体", "65歳以上の割合（問2）", 1, ["65歳以上"], []),
    ("AN40-T21", "第2章7 障がい別の比較", "項目|全体", "将来も在宅での生活を希望（問11）", 1,
     ["在宅を希望"], []),
    ("AN40-T22", "第2章7 障がい別の比較", "項目|全体",
     "外出時に「冬季・積雪時に外出しにくい」（問14）", 1, ["外出_冬季"], []),
    ("AN40-T23", "第2章7 障がい別の比較", "項目|全体",
     "困ったときの相談相手「家族・親族」（問22）", 1, ["相談相手_家族"], []),
    ("AN40-T24", "第2章7 障がい別の比較", "項目|全体",
     "働くために必要な支援「体調に合わせた勤務」（問33）", 1,
     ["就労支援_体調に合わせた勤務"], []),
    ("AN40-T25", "第2章7 障がい別の比較", "項目|全体",
     "差別や偏見、不利益を感じたことがある（問34）", 1, ["差別を感じた"], []),
    ("AN40-T26", "第2章7 障がい別の比較", "項目|全体", "災害時の不安「避難先での生活」（問38）", 1,
     ["災害不安_避難先での生活"], []),
    ("AN40-T27", "第2章7 障がい別の比較", "項目|全体",
     "避難行動要支援者名簿・個別避難計画を「知らない」（問39）", 1,
     ["個別避難計画_知らない"], []),
    ("AN40-T28", "第2章7 課題の整理", "主な課題", "相談支援体制の充実と\nわかりやすい情報提供", 1,
     ["相談相手_家族", "相談への期待_情報提供"], []),
    ("AN40-T29", "第2章7 課題の整理", "主な課題", "移動・外出支援\n（冬季への対応を含む）", 1,
     ["外出_冬季", None], []),
    ("AN40-T30", "第2章7 課題の整理", "主な課題", "地域生活の継続と\n緊急時の備え", 1,
     ["在宅を希望", "拠点_緊急時の受入れ"], []),
    ("AN40-T31", "第2章7 課題の整理", "主な課題", "災害時の支援体制", 1,
     ["個別避難計画_知らない", "災害不安_避難先での生活", None], []),
    ("AN40-T32", "第2章7 課題の整理", "主な課題", "権利擁護・差別解消", 1,
     ["差別を感じた", "成年後見_知らない"], []),
    ("AN40-T33", "第2章7 課題の整理", "主な課題", "就労・日中活動", 1,
     ["就労支援_職場の理解", None], []),
]


def _vals(kyu):
    d = {k: A.pc(k, kyu) for k in A.M}
    d["n"] = A.n(kyu)
    d["n_goukei"] = A.kazu("有効回答率_合計", kyu)[0]
    return d


OLD, NEW = _vals(True), _vals(False)


def _okikae(text, names, rep, src_vals, dst_vals):
    """割合を出る順に指標の値へ置き換える。src_vals の値と一致しなければ止める。"""
    ms = list(PAT.finditer(text))
    if len(ms) != len(names):
        raise ValueError(f"割合の数が合わない（{len(ms)}≠{len(names)}）: {text[:40]}")
    out, pos = [], 0
    for m, nm in zip(ms, names):
        out.append(text[pos:m.start()])
        if nm is None:
            out.append(m.group(0))
        else:
            if m.group(0) != src_vals[nm]:
                raise ValueError(f"38件の値と一致しない: {nm} {m.group(0)}≠{src_vals[nm]}"
                                 f"／{text[:40]}")
            out.append(dst_vals[nm])
        pos = m.end()
    out.append(text[pos:])
    s = "".join(out)
    for a, b in rep:
        a2, b2 = a.format(**dst_vals), b.format(**dst_vals)
        if s.count(a2) != 1:
            raise ValueError(f"言い回しの置換先が一意でない: {a2[:30]}")
        s = s.replace(a2, b2)
    return s


def _member_paras():
    d = docx.Document(MEMBER)
    return [p.text for p in d.paragraphs]


def _member_tables():
    d = docx.Document(MEMBER)
    out = []
    for t in d.tables:
        rows = [[c.text for c in r.cells] for r in t.rows]
        out.append(rows)
    return out


def _find_row(tables, atama, row_key):
    """表頭の1列目と行の1列目で行を一つに絞る。'相談@2' は同じ見出しの2つ目。"""
    key, nth = (row_key.split("@") + ["1"])[:2] if row_key else (None, "1")
    hits = []
    a1, a2 = (atama.split("|") + [""])[:2]
    for t in tables:
        if t[0][0].strip() != a1:
            continue
        if a2 and not (len(t[0]) > 1 and t[0][1].startswith(a2)):
            continue
        if key is None:
            hits.append(t[0])
            continue
        rs = [r for r in t[1:] if r[0] == key]
        if len(rs) >= int(nth):
            hits.append(rs[int(nth) - 1])
    if len(hits) != 1:
        raise LookupError(f"行が一意でない（{len(hits)}）: {atama} {row_key}")
    return hits[0]


def _build():
    paras = _member_paras()
    kaki = []
    for no, saki, kakidashi, names, rep in SASHIKAE:
        hits = [p for p in paras if p.startswith(kakidashi)]
        if len(hits) != 1:
            raise LookupError(f"段落が一意でない（{len(hits)}）: {no}")
        mae = hits[0]
        ato = _okikae(mae, names, rep, OLD, NEW)
        kaki.append(dict(no=no, saki=saki, mae=mae, ato=ato,
                         konkyo=A.SHIN, ryo="本文1段落の割合の差し替え"))
    tables = _member_tables()
    hyou = []
    for no, saki, atama, row_key, col, names, rep in SASHIKAE_HYOU:
        row = _find_row(tables, atama, row_key)
        mae = row[col]
        ato = _okikae(mae, names, rep, OLD, NEW)
        if ato == mae:
            raise ValueError(f"差し替えても変わらない: {no}")
        #   mae・ato は1セルの組で持つ（対応履歴の照合が表の行で行うため）
        hyou.append(dict(no=no, saki=saki, atama=atama,
                         row_key=row_key, col=col, mae=(mae,), ato=(ato,),
                         konkyo=A.SHIN, ryo="表の1セルの割合の差し替え"))
    return kaki, hyou


AN40_KAKIKAE, AN40_HYOU = _build()


def shin_ichi(ichi):
    """他メンバー版の段落の一部（挿入位置の引用）を、差し替え後の文にする。

    当方の追記は、他メンバー版の段落を挿入位置として引用している
    （juten12 の W-2・W-9、shihyou_sokutei の ＳＳ-1 など）。追記は差し替えの
    前に入れるため引用は38件の値のままでよいが、出来上がった正本で
    挿入位置を確かめるときは、差し替え後の文で照らす必要がある。
    引用の範囲にある割合だけを、その段落の割り当てに従って置き換える。
    """
    paras = _member_paras_cache()
    for _no, _saki, kakidashi, names, rep in SASHIKAE:
        hits = [p for p in paras if p.startswith(kakidashi)]
        if len(hits) != 1 or ichi not in hits[0]:
            continue
        para = hits[0]
        idx = para.index(ichi)
        out = ichi
        pairs = list(zip(PAT.finditer(para), names))
        for m, nm in reversed(pairs):
            if nm is None:
                continue
            if m.start() >= idx and m.end() <= idx + len(ichi):
                a, b = m.start() - idx, m.end() - idx
                out = out[:a] + NEW[nm] + out[b:]
        for a0, b0 in rep:
            a2, b2 = a0.format(**NEW), b0.format(**NEW)
            if a2 in out:
                out = out.replace(a2, b2)
        return out
    return ichi


_MP = []


def _member_paras_cache():
    if not _MP:
        _MP.extend(_member_paras())
    return _MP


# ===========================================================================
# 3. 自己点検
# ===========================================================================
def kensa(honpon_text):
    ng = []
    #   ① 各指標の38件の値が、集計表（38件版）から出したものと同じ
    #     （_okikae が一致しなければ止まるため、ここまで来れば合っている）
    #   ② 障がい別の集計（手帳による区分）が新旧で動いていない
    from openpyxl import load_workbook
    def betsu(p):
        wb = load_workbook(p, read_only=True, data_only=True)
        #   数値を含む行だけを比べる（注記の文言は版で変わる。
        #   令和8年10月21日版は「青字」を「斜体」に改めている）
        v = [tuple(r) for r in wb["集計_障がい別"].iter_rows(values_only=True)
             if any(isinstance(x, (int, float)) for x in r)]
        wb.close()
        return v
    bk, bs = betsu(A.KYU), betsu(A.SHIN)
    if len(bk) < 100:
        ng.append(f"② 障がい別の集計の数値の行が少なすぎる: {len(bk)}行")
    if bk != bs:
        ng.append("② 障がい別の集計が新旧で違う（None を割り当てた値を見直す）")
    #   ③ 正本に38件の表記が残っていない
    if honpon_text is not None:
        for k in ("（n=38）", "38人 | 22.1％", "回答者数38人"):
            if k in honpon_text:
                ng.append(f"③ 正本に38件の表記が残っている: {k}")
        if "有効回答率が22.1％" in honpon_text or "有効回答率は22.1％" in honpon_text:
            ng.append("③ 正本に有効回答率22.1％が残っている")
        for r in AN40_KAKIKAE:
            if "".join(r["ato"].split()) not in honpon_text:
                ng.append(f"③ 正本に差し替えた文が無い: {r['no']}")
    return ng


def _honpon_text():
    if not os.path.exists(HONPON):
        return None
    d = docx.Document(HONPON)
    out = ["".join(p.text.split()) for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            out.append(" | ".join("".join(c.text.split()) for c in r.cells))
    return "\n".join(out)


def verify():
    ng = kensa(_honpon_text())
    #   否定試験71：割り当てを1つずらすと①（_okikae）が止めること
    try:
        _okikae(SASHIKAE_TEXT_TEST, ["外出_冬季"], [], OLD, NEW)
        ng.append("否定試験71: 違う指標を割り当てても止まらない")
    except ValueError:
        pass
    if ng:
        print("自己点検 不合格:")
        for e in ng:
            print("   -", e)
        raise SystemExit(1)
    return (f"段落{len(AN40_KAKIKAE)}・セル{len(AN40_HYOU)}の割合を38件の値と照合して"
            "40件へ差し替え／障がい別の集計が新旧で不変／正本に38件の表記なし／否定試験1件")


#   否定試験用：65歳以上（65.8％）の段落に「外出_冬季」（34.2％）を割り当てる
SASHIKAE_TEXT_TEST = "回答者の65.8％が65歳以上"


def main():
    ensure_out_dir()
    msg = verify()
    wb = Workbook()
    wb.remove(wb.active)
    ws = add_sheet(wb, "00_指標の新旧", "アンケートの指標の新旧（38件→40件）",
                   "kitashiobara_anke の定義による。件数÷分母。障がい別・障がい児は動かない。",
                   [34, 14, 14, 14, 14])
    r = 5
    style_header_row(ws, r, ["指標", "38件（件数／分母）", "38件", "40件（件数／分母）", "40件"])
    r += 1
    for k in A.M:
        a, b = A.kazu(k, True), A.kazu(k)
        write_row(ws, r, [k, f"{a[0]}／{a[1]}", OLD[k], f"{b[0]}／{b[1]}", NEW[k]],
                  aligns=["left", "center", "center", "center", "center"])
        r += 1
    ws = add_sheet(wb, "01_段落の差し替え", f"正本の段落の差し替え（{len(AN40_KAKIKAE)}段落）",
                   "他メンバー版の本文。割合が出る順に指標を割り当て、38件の値と一致することを確かめた。",
                   [10, 26, 70, 70])
    r = 5
    style_header_row(ws, r, ["番号", "箇所", "直す前", "直した後"])
    r += 1
    for x in AN40_KAKIKAE:
        write_row(ws, r, [x["no"], x["saki"], x["mae"], x["ato"]])
        ws.row_dimensions[r].height = 120
        r += 1
    ws = add_sheet(wb, "02_表の差し替え", f"正本の表のセルの差し替え（{len(AN40_HYOU)}セル）",
                   "他メンバー版の表。", [10, 26, 30, 50, 50])
    r = 5
    style_header_row(ws, r, ["番号", "箇所", "行", "直す前", "直した後"])
    r += 1
    for x in AN40_HYOU:
        write_row(ws, r, [x["no"], x["saki"], x["row_key"] or "（表頭）", x["mae"][0],
                          x["ato"][0]])
        ws.row_dimensions[r].height = 40
        r += 1
    wb.save(OUT_FILE)
    print(f"  自己検証: {msg}")
    print(f"作成: {OUT_FILE}")


if __name__ == "__main__":
    sys.exit(main())
