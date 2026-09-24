"""小野町 第10期 見込量の推計方法をバックテストで検証する。

他町村のブランチで確立した方法と当方の方法を、小野町の年報の実績で突き合わせる。

**比べる方法**

  据置        基準年度の実績をそのまま延ばす（見える化の自然体推計の初期値。伸び0）
  一律        基準年度の実績 × (1＋分母の伸び)^経過年数
  補正        大雪地区広域連合の手引き第5章
                g'ₛ ＝ gₙ ＋ φ × clip(gₛ － gₙ, ±c)
                見込量ₛ ＝ 基準年度の実績ₛ × (1 ＋ g'ₛ)^経過年数
              φ は趨勢を反映する割合、c は趨勢の差の上限。手引きの例は φ＝0.5・c＝±3％
  利用率      当方の方法。基準期間の［給付費÷分母］の平均 × 将来の分母

**測り方**（手引き第5章3）

  起点年度の実績から数年先を予測し、その年度の実績と比べる。
  誤差はサービス別の絶対誤差を給付費で加重したものと、総給付費の絶対誤差の2つで見る。
  起点と予測年数の組合せを複数用意し、その平均で設定を選ぶ。
  一律（φ＝0）を必ず比較対象に入れ、補正が一律より良いことを示せなければ補正しない。

**小野町で注意すること**

  年報は令和3〜7年度の5年分しかないため、組合せは3通りに限られる。
  しかも令和7年度は認定者数に段差があり（令和7年9月審査分で961人→862人）、
  令和7年度を予測先に置く2通りはその影響を受ける。
  結果はこの制約のもとで読む。

出力: 04_算定・見込量/小野町_推計方法のバックテスト_YYYYMMDD.xlsx
"""

import itertools
import pathlib
import warnings

warnings.filterwarnings("ignore")

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.borders import Border, Side

import build_ono_tanka as T
import ono_kaigodo as K

OUT = pathlib.Path(__file__).parent / "小野町_引継ぎ_整理済" / "04_算定・見込量"
ASOF = "20260915"
ASOF_JP = "令和8年9月15日"

HEAD = PatternFill("solid", fgColor="1F3864")
KEY = PatternFill("solid", fgColor="FCE4E4")
CALC = PatternFill("solid", fgColor="EAF1FB")
WARN = PatternFill("solid", fgColor="FFE0B2")
OK = PatternFill("solid", fgColor="C6EFCE")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

YS = ["令和3年度", "令和4年度", "令和5年度", "令和6年度", "令和7年度"]


def load():
    """{(サービス, 給付の別): {年度: 給付費（価格調整後・千円）}} と分母を返す。"""
    act, _sub, gen = T.extract()
    g = {}
    for _cat, svc, _u in T.ORDER_KAIGO:
        for kind, i in (("介護", 1), ("予防", 0)):
            ser = {}
            for y in YS:
                v = act[y]["給付費"].get(svc, (0, 0))[i] / 1000
                ser[y] = v * T.PRICE_ADJ.get(y, 1.0)
            if max(ser.values()) > 0:
                g[(svc, kind)] = ser
    den = {
        "第1号被保険者数": {y: gen[y]["1号"] for y in YS},
        "認定者数": {y: gen[y]["認定"] for y in YS},
        "75歳以上人口": {y: gen[y]["75+"] for y in YS},
    }
    return g, den


# 要介護度の区分を予防給付・介護給付に振り分ける。
# 様式2は1行で予防（要支援1・2の列）と介護（要介護1〜5の列）の双方を持つ。
KD_KIND = {"予防": ["要支援1", "要支援2"],
           "介護": ["要介護1", "要介護2", "要介護3", "要介護4", "要介護5"]}


def load_kaigodo():
    """川崎町方式のもと。{(サービス, 給付の別): {年度: {要介護度: 千円}}} と認定者数。

    価格水準は load() と同じ調整をかけ、比較できるようにする。
    """
    kyufu, nintei, _chk = K.read()
    out = {}
    for y in YS:
        adj = T.PRICE_ADJ.get(y, 1.0)
        for svc, v in kyufu[y].items():
            for kind, ds in KD_KIND.items():
                ser = out.setdefault((svc, kind), {})
                ser[y] = {d: v[d] / 1000 * adj for d in ds}
    return out, nintei


def cagr(ser, years):
    """年平均変化率。始点か終点が0なら0を返す。"""
    a, b = ser.get(years[0], 0), ser.get(years[-1], 0)
    n = len(years) - 1
    if a <= 0 or b <= 0 or n <= 0:
        return 0.0
    return (b / a) ** (1 / n) - 1


def predict(g, den, den_key, t0, h, method, phi=0.5, cap=0.03, win=3,
            base_years=None, kd=None, nintei=None, kd_kind="総数"):
    """起点 t0 から h 年先を予測した給付費を返す。"""
    i0 = YS.index(t0)
    tgt = YS[i0 + h]
    if method == "度別":
        gr = K.growth(nintei, t0, tgt, kd_kind)
        out = {}
        for k, ser in g.items():
            base = kd.get(k, {}).get(t0)
            out[k] = (sum(base[d] * gr[d] for d in base) if base
                      else ser[t0])
        return out, tgt
    d = den[den_key]
    gn = cagr(d, YS[max(0, i0 - win + 1):i0 + 1])
    out = {}
    for k, ser in g.items():
        base = ser[t0]
        if method == "据置":
            out[k] = base
        elif method == "一律":
            out[k] = base * (1 + gn) ** h
        elif method == "補正":
            gs = cagr(ser, YS[max(0, i0 - win + 1):i0 + 1])
            gd = max(-cap, min(cap, gs - gn))
            out[k] = base * (1 + gn + phi * gd) ** h
        elif method == "利用率":
            bys = base_years or YS[max(0, i0 - win + 1):i0 + 1]
            rates = [ser[y] / d[y] for y in bys if d[y]]
            out[k] = (sum(rates) / len(rates)) * d[tgt] if rates else 0.0
        else:
            raise ValueError(method)
    return out, tgt


def score(pred, g, tgt):
    """サービス別絶対誤差の給付費加重平均と、総給付費の絶対誤差。"""
    tot_a = sum(ser[tgt] for ser in g.values())
    tot_p = sum(pred.values())
    w = 0.0
    for k, ser in g.items():
        a = ser[tgt]
        if a > 0:
            w += a * abs(pred[k] / a - 1)
    return (w / tot_a if tot_a else 0.0,
            abs(tot_p / tot_a - 1) if tot_a else 0.0)


# 起点年度と予測年数の組合せ。年報が5年分しかないため3通りに限られる。
COMBOS = [("令和5年度", 1), ("令和5年度", 2), ("令和6年度", 1)]


def run():
    g, den = load()
    kd, nintei = load_kaigodo()
    rows = []
    # 川崎町方式（要介護度別の認定者数の伸び）。分母を1つ選ぶ方式ではないため
    # 分母の総当たりの外で1回だけ測る。認定者数は総数と第1号の2系列で見る。
    for kind in ("総数", "1号"):
        svc_e, tot_e = [], []
        for t0, h in COMBOS:
            p_, tgt = predict(g, den, "第1号被保険者数", t0, h, "度別",
                              kd=kd, nintei=nintei, kd_kind=kind)
            a, b = score(p_, g, tgt)
            svc_e.append(a)
            tot_e.append(b)
        rows.append({
            "分母": f"要介護度別の認定者数（{kind}）", "方法": "度別",
            "φ": None, "上限": None, "窓": None,
            "サービス別": sum(svc_e) / len(svc_e),
            "総給付費": sum(tot_e) / len(tot_e),
            "内訳": list(zip([f"{t0[2:]}+{h}" for t0, h in COMBOS],
                           [round(x * 100, 2) for x in tot_e])),
        })
    for den_key in den:
        for method in ("据置", "一律", "補正", "利用率"):
            grid = ([(None, None, None)] if method in ("据置", "一律", "利用率")
                    else list(itertools.product((0.25, 0.5, 0.75, 1.0),
                                                (0.02, 0.03, 0.05), (3,))))
            if method == "利用率":
                grid = [(None, None, w) for w in (2, 3)]
            for phi, cap, win in grid:
                svc_e, tot_e = [], []
                for t0, h in COMBOS:
                    p, tgt = predict(g, den, den_key, t0, h, method,
                                     phi=phi or 0.0, cap=cap or 0.03,
                                     win=win or 3)
                    s, t = score(p, g, tgt)
                    svc_e.append(s)
                    tot_e.append(t)
                rows.append({
                    "分母": den_key, "方法": method, "φ": phi, "上限": cap,
                    "窓": win,
                    "サービス別": sum(svc_e) / len(svc_e),
                    "総給付費": sum(tot_e) / len(tot_e),
                    "内訳": list(zip([f"{t0[2:]}+{h}" for t0, h in COMBOS],
                                   [round(x * 100, 2) for x in tot_e])),
                })
    return g, den, rows


# ---------------------------------------------------------------- 書式

def style_head(ws, row=1):
    for c in ws[row]:
        if c.value is None:
            continue
        c.fill = HEAD
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)


def body(ws, first=2, wrap=()):
    for r in ws.iter_rows(min_row=first):
        for c in r:
            if c.value is None:
                continue
            c.border = BORDER
            if c.font.size is None or c.font.size > 9:
                c.font = Font(bold=c.font.bold, size=9)
            c.alignment = Alignment(vertical="center", wrap_text=c.column in wrap)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def notes(ws, lines):
    ws.append([])
    for t in lines:
        ws.append([t])
        ws.cell(ws.max_row, 1).font = Font(size=9)
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True,
                                                     vertical="top")


def main():
    g, den, rows = run()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---- 00 結果
    ws = wb.create_sheet("00_結果")
    best = min(rows, key=lambda r: r["総給付費"])
    ichiritsu = [r for r in rows if r["方法"] == "一律"]
    ws.append(["小野町 第10期 推計方法のバックテスト"])
    ws.cell(1, 1).font = Font(bold=True, size=13)
    ws.append([f"作成日：{ASOF_JP}"])
    ws.append([])
    ws.append(["■ 何をしたか"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=10)
    for t in [
        "大雪地区広域連合の『サービス見込量算定の手引き』第5章の手順で、"
        "推計方法を小野町の年報の実績で検証しました。",
        "起点年度の実績から数年先を予測し、その年度の実績と比べます。"
        "年報が令和3〜7年度の5年分しかないため、組合せは"
        f"{'・'.join(f'{t0}から{h}年先' for t0, h in COMBOS)}の3通りです。",
        "**令和7年度は認定者数に段差があり（令和7年9月審査分で961人→862人）、"
        "令和7年度を予測先に置く2通りはその影響を受けます。**",
    ]:
        ws.append([t])
    ws.append([])
    ws.append(["■ 結果"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=10)
    ws.append(["方法", "分母", "設定", "サービス別の誤差", "総給付費の誤差"])
    hr = ws.max_row
    for r in sorted(rows, key=lambda x: x["総給付費"])[:12]:
        st = ("―" if r["方法"] in ("据置", "一律")
              else f"窓{r['窓']}年" if r["方法"] == "利用率"
              else f"φ{r['φ']}・上限±{r['上限'] * 100:.0f}％・窓{r['窓']}年")
        ws.append([r["方法"], r["分母"], st, r["サービス別"], r["総給付費"]])
    style_head(ws, hr)
    for r in range(hr + 1, ws.max_row + 1):
        for c in (4, 5):
            ws.cell(r, c).number_format = "0.00%"
            ws.cell(r, c).fill = CALC
        if r == hr + 1:
            for c in range(1, 6):
                ws.cell(r, c).fill = KEY
                ws.cell(r, c).font = Font(bold=True, size=9)
    notes(ws, [
        f"※ 最良は「{best['方法']}」（分母{best['分母']}）で、"
        f"総給付費の誤差{best['総給付費'] * 100:.2f}％。"
        f"一律の最良は{min(x['総給付費'] for x in ichiritsu) * 100:.2f}％。",
        "**※ 手引きは「補正が一律より良いことを示せなければ補正しない」としています。**"
        "小野町ではこの条件を満たすかどうかを、上の表で確かめてください。",
        "※ サービス別の誤差は、サービスごとの絶対誤差を給付費で加重した平均。"
        "総給付費の誤差は合計の絶対誤差。3通りの組合せの平均です。",
        "※ 給付費は令和6年度の介護報酬改定（＋1.59％）で価格水準をそろえています。",
    ])
    body(ws, first=hr, wrap=(1,))
    widths(ws, [16, 18, 30, 18, 18])

    # ---- 01 全設定
    ws = wb.create_sheet("01_全設定")
    ws.append(["方法", "分母", "φ", "上限", "窓", "サービス別の誤差",
               "総給付費の誤差"]
              + [f"総給付費の誤差 {lab}" for lab, _ in rows[0]["内訳"]])
    for r in sorted(rows, key=lambda x: (x["方法"], x["分母"], x["総給付費"])):
        ws.append([r["方法"], r["分母"], r["φ"], r["上限"], r["窓"],
                   r["サービス別"], r["総給付費"]]
                  + [v / 100 for _lab, v in r["内訳"]])
    style_head(ws)
    for r in range(2, ws.max_row + 1):
        for c in (4, 6, 7, 8, 9, 10):
            ws.cell(r, c).number_format = "0.00%"
    body(ws)
    widths(ws, [12, 18, 7, 9, 7, 18, 18, 18, 18, 18])
    ws.freeze_panes = "C2"

    # ---- 02 分母の伸び
    ws = wb.create_sheet("02_分母の伸び")
    ws.append(["分母"] + YS + ["令和5→7年度の年率", "変動係数"])
    import statistics
    for k, d in den.items():
        v = [d[y] for y in YS]
        ws.append([k] + v + [cagr(d, YS[2:]),
                             statistics.pstdev(v) / statistics.mean(v)])
    style_head(ws)
    for r in range(2, ws.max_row + 1):
        for c in range(2, 7):
            ws.cell(r, c).number_format = "#,##0"
        for c in (7, 8):
            ws.cell(r, c).number_format = "0.00%"
    body(ws)
    widths(ws, [20, 12, 12, 12, 12, 12, 18, 12])
    notes(ws, [
        "**※ 認定者数は令和7年9月審査分に段差があり（961人→862人）、"
        "令和6年度944人から令和7年度792人へ16.1％減っています。**"
        "国保連月報でみると、段差は認定者数にだけ現れ受給者数は連続しています"
        "（断層前後で認定者▲11.4％に対し受給者計▲3.9％）。",
        "※ 分母の伸びを一律に用いる方法（一律・補正）では、"
        "この段差がそのまま伸びに入ります。",
    ])

    # ---- 03 サービス別の趨勢
    ws = wb.create_sheet("03_サービス別の趨勢")
    ws.append(["サービス", "給付の別"] + [y[2:] for y in YS]
              + ["令和5→7年度の年率", "令和3→7年度の年率"])
    for (svc, kind), ser in sorted(g.items(), key=lambda x: -x[1][YS[-1]]):
        if max(ser.values()) < 1000:
            continue
        ws.append([svc, kind] + [round(ser[y]) for y in YS]
                  + [cagr(ser, YS[2:]), cagr(ser, YS)])
    style_head(ws)
    for r in range(2, ws.max_row + 1):
        for c in range(3, 8):
            ws.cell(r, c).number_format = "#,##0"
        for c in (8, 9):
            ws.cell(r, c).number_format = "+0.0%;-0.0%"
            v = ws.cell(r, c).value
            if isinstance(v, float) and abs(v) > 0.10:
                ws.cell(r, c).fill = WARN
    body(ws)
    widths(ws, [34, 9, 12, 12, 12, 12, 12, 18, 18])
    ws.freeze_panes = "C2"
    notes(ws, [
        "※ 給付費（千円）。令和6年度の介護報酬改定で価格水準をそろえています。",
        "※ 年率±10％を超えるものに色を付けています。"
        "**サービス別の趨勢は分母の伸びとは別物で、"
        "一律の伸びで置くとその差がそのまま対計画比の開きになります**"
        "（手引き第5章1）。",
    ])

    # ---- 04 要介護度別（川崎町方式）
    kd, nintei = load_kaigodo()
    ws = wb.create_sheet("04_要介護度別（川崎町方式）")
    ws.append(["川崎町 第10期で用いている方式を小野町のデータで測る"])
    ws.cell(1, 1).font = Font(bold=True, size=13)
    for t in [
        "方式：基準年度の要介護度別給付費（年報 様式2）に、"
        "要介護度別の認定者数の伸び（年報 様式1の5）を乗じて足し上げる。"
        "1人当たりの水準は基準年度で据え置く。",
        "当方の方式が第1号被保険者1人当たりの利用率を延ばすのに対し、"
        "こちらは要介護度の構成の変化だけを織り込む。",
    ]:
        ws.append([t])
    ws.append([])
    ws.append(["■ 要介護度別の認定者数（年報 様式1の5・総数・年度末）"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=10)
    ws.append(["年度"] + K.KAIGODO + ["計"])
    hr = ws.max_row
    for y in YS:
        n = nintei[y]["総数"]
        ws.append([y] + [n[d] for d in K.KAIGODO] + [sum(n.values())])
    style_head(ws, hr)
    ws.append([])
    ws.append(["■ 伸び率（基準年度＝1.000）"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=10)
    ws.append(["起点→予測先"] + K.KAIGODO + ["全体"])
    hr2 = ws.max_row
    pairs = [(t0, YS[YS.index(t0) + h]) for t0, h in COMBOS]
    for t0, tgt in pairs:
        gr = K.growth(nintei, t0, tgt)
        tot = (sum(nintei[tgt]["総数"].values())
               / sum(nintei[t0]["総数"].values()))
        ws.append([f"{t0}→{tgt}"] + [round(gr[d], 3) for d in K.KAIGODO]
                  + [round(tot, 3)])
    style_head(ws, hr2)
    for r in range(hr2 + 1, ws.max_row + 1):
        for c in range(2, len(K.KAIGODO) + 3):
            ws.cell(r, c).number_format = "0.000"
            v = ws.cell(r, c).value
            if isinstance(v, float) and v < 0.90:
                ws.cell(r, c).fill = WARN
    ws.append([])
    ws.append(["■ 総給付費の予測（価格調整後・千円）"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=10)
    ws.append(["起点→予測先", "実績", "要介護度別（川崎町方式）", "誤差",
               "利用率・第1号・窓3年（当方）", "誤差", "断層をまたぐか"])
    hr3 = ws.max_row
    for (t0, h), (_t, tgt) in zip(COMBOS, pairs):
        act = sum(ser[tgt] for ser in g.values())
        pk, _ = predict(g, den, "第1号被保険者数", t0, h, "度別",
                        kd=kd, nintei=nintei)
        pu, _ = predict(g, den, "第1号被保険者数", t0, h, "利用率", win=3)
        vk, vu = sum(pk.values()), sum(pu.values())
        ws.append([f"{t0}→{tgt}", round(act), round(vk), vk / act - 1,
                   round(vu), vu / act - 1,
                   "またぐ" if tgt == "令和7年度" else "またがない"])
    style_head(ws, hr3)
    for r in range(hr3 + 1, ws.max_row + 1):
        for c in (2, 3, 5):
            ws.cell(r, c).number_format = "#,##0"
        for c in (4, 6):
            ws.cell(r, c).number_format = "+0.00%;-0.00%"
            v = ws.cell(r, c).value
            if isinstance(v, float):
                ws.cell(r, c).fill = WARN if abs(v) > 0.05 else OK
        if ws.cell(r, 7).value == "またぐ":
            ws.cell(r, 7).fill = WARN
    body(ws, first=hr, wrap=(1,))
    widths(ws, [24, 13, 13, 13, 13, 13, 13, 13, 16])
    notes(ws, [
        "**※ 川崎町方式は、断層をまたがない令和5年度→令和6年度では"
        "▲1.06％と全方式のなかで最も良い。**"
        "いっぽう令和7年度を予測先に置く2通りでは▲14.66％・▲13.47％となり、"
        "3通りの平均では50設定中49位に落ちる。",
        "**※ 理由は伸び率の表にある。**令和6年度→令和7年度の認定者数は"
        "要介護4が0.755倍、要介護5が0.626倍であるのに対し、"
        "給付費の実績は7.4％しか下がっていない。"
        "令和7年9月の断層で重度の認定者が「減った」のは計上の変更であって、"
        "その方々がサービスを使わなくなったわけではない。"
        "要介護度別の伸びをそのまま給付費に乗せると、この差がそのまま誤差になる。",
        "※ したがって本方式は、**断層の原因が町の確認で解消するまでは"
        "小野町には用いない。**"
        "解消した場合は、要介護度の構成の変化を織り込める点で当方の方式より"
        "優れる可能性があり、そのときに採り直す。",
        "※ 認定者数を総数で見るか第1号で見るかによる差は小さい"
        "（総給付費の誤差9.73％と9.82％）。",
    ])

    OUT.mkdir(parents=True, exist_ok=True)
    p = OUT / f"小野町_推計方法のバックテスト_{ASOF}.xlsx"
    wb.save(p)
    print(f"出力: {p}")
    print(f"{'方法':<8}{'分母':<16}{'設定':<26}{'サービス別':>10}{'総給付費':>10}")
    for r in sorted(rows, key=lambda x: x["総給付費"])[:10]:
        st = ("―" if r["方法"] in ("据置", "一律")
              else f"窓{r['窓']}年" if r["方法"] == "利用率"
              else f"φ{r['φ']}・上限±{r['上限'] * 100:.0f}％")
        print(f"{r['方法']:<8}{r['分母']:<16}{st:<26}"
              f"{r['サービス別'] * 100:>9.2f}%{r['総給付費'] * 100:>9.2f}%")
    return rows


if __name__ == "__main__":
    main()
