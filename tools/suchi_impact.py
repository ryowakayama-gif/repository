# -*- coding: utf-8 -*-
"""数値を直したときに、文章と他の成果品への及び方を確かめる。

図表は数値を直せば追随するが、**その数値を述べた文章は追随しない。**
また、同じ数値が別の成果品にも書かれていることがある。
本ツールは次の3つを出す。

    1 変わった数値       ある版と今の版を比べ、消えた数値・現れた数値を出す
    2 文章への影響       消えた数値を含む文がどこに残っているかを出し、
                         大小・向きを述べている文を「要判断」とする
    3 修正漏れ           消えた数値が他の成果品・ソースに残っていないかを出す

使い方
    python3 tools/suchi_impact.py                   直前のコミットとの差を見る
    python3 tools/suchi_impact.py --rev HEAD~3      指定した版との差を見る
    python3 tools/suchi_impact.py --value 6740 6436 旧→新を直接に指定する
    python3 tools/suchi_impact.py --map 6436        その数値がどこにあるかだけ出す
    python3 tools/suchi_impact.py --csv <path>      結果をCSVにも書き出す
    python3 tools/suchi_impact.py --min 4           4桁以上の数値だけを見る

桁の少ない整数は別の量として偶然に一致しやすい。
騒がしいときは --min を上げる（既定は3。小数・桁区切り付きは桁数によらず見る）。

判定
    不適合  旧の値が他の成果品に残っている（修正漏れ）
    要判断  旧の値を含む文に、大小・向き・順位を述べる語がある
    未実施  読めなかったファイル（理由を残す。適合に丸めない）

**「消えた数値」は直した値とは限らない。** 表の区分を入れ替えた、
年度を1つ進めたといった理由でも消える。**必ず人が中身を見て決める。**

不適合が1件でもあると終了コード1で終わる。
"""

import csv
import os
import re
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# 数値らしい並び。桁区切り付き・小数・3桁以上の整数に限る（1桁2桁は拾わない）
NUM = re.compile(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+\.\d+|\d{3,}")
# 大小・向き・順位を述べる語。これを含む文は数値を直したら読み直す
MUKI = ("上回", "下回", "増え", "減り", "減少", "増加", "最も", "最大", "最小",
        "倍", "ポイント", "超え", "満た", "以上", "以下", "未満", "高い", "低い",
        "多い", "少ない", "伸び", "落ち", "位", "順")
BUN = re.compile(r"[^。\n]+。|[^。\n]+$")

FUTEKI, YOHAN, MIJISSHI = "不適合", "要判断", "未実施"
MIJISSHI_LOG = []


def canon(s):
    """「6,740」「6740」「4.00」「4.0」を同じものとして扱うための正規化。

    桁区切りを落とし、小数の末尾の0を落とす。
    表示の桁数の違いを「数値が変わった」と読まないためである。
    表示そのものを比べたいときは元の文字列を用いる。
    """
    s = s.replace(",", "")
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return s or "0"


# ------------------------------------------------------------ 読み取り
def read_docx(path, blob=None):
    from docx import Document
    from docx.oxml.ns import qn
    import io
    d = Document(io.BytesIO(blob) if blob is not None else path)
    out = []
    for p in d.element.body.iter(qn("w:p")):
        t = "".join(n.text or "" for n in p.iter(qn("w:t")))
        if t.strip():
            out.append(("本文", t))
    for i, s in enumerate(d.sections, start=1):
        for nm, part in (("ヘッダー", s.header), ("フッター", s.footer)):
            try:
                for p in part._element.iter(qn("w:p")):
                    t = "".join(n.text or "" for n in p.iter(qn("w:t")))
                    if t.strip():
                        out.append(("第%d節の%s" % (i, nm), t))
            except Exception:
                pass
    return out


def read_xlsx(path, blob=None):
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(blob) if blob is not None else path,
                       read_only=True, data_only=False)
    out = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None or v == "":
                    continue
                if isinstance(v, str) and v.startswith("="):
                    continue
                out.append(("%s!%s" % (sn, c.coordinate), str(v)))
    wb.close()
    return out


def read_py(path, blob=None):
    s = blob.decode("utf-8", "replace") if blob is not None else \
        open(path, encoding="utf-8", errors="replace").read()
    return [("%d行" % (i + 1), ln) for i, ln in enumerate(s.split("\n")) if ln.strip()]


def read_any(path, blob=None):
    try:
        if path.endswith(".docx"):
            return read_docx(path, blob)
        if path.endswith(".xlsx"):
            return read_xlsx(path, blob)
        if path.endswith((".py", ".md")):
            return read_py(path, blob)
    except Exception as e:
        MIJISSHI_LOG.append((path, str(e)[:70]))
    return []


def targets():
    """点検の対象。成果品（docx・xlsx）と、数値を持つソース。"""
    out = []
    od = os.path.join(ROOT, "output")
    if os.path.isdir(od):
        for f in sorted(os.listdir(od)):
            if f.endswith((".docx", ".xlsx")) and not f.startswith("~$"):
                out.append(os.path.join("output", f))
    for f in sorted(os.listdir(ROOT)):
        if f.endswith(".py") and (f.startswith("data_") or f.startswith("build_")):
            out.append(f)
    out.append("CLAUDE.md")
    return [p for p in out if os.path.exists(os.path.join(ROOT, p))]


# ------------------------------------------------------------ 索引
def build_index(paths):
    """数値 → [(ファイル, 位置, 文字列)] の索引を作る。"""
    idx = {}
    for rel in paths:
        for place, text in read_any(os.path.join(ROOT, rel)):
            for m in set(NUM.findall(text)):
                idx.setdefault(canon(m), []).append((rel, place, text))
    return idx


def nums_of(paths, rev=None):
    """指定の版（rev が None なら今）の数値の集合をファイルごとに返す。"""
    out = {}
    for rel in paths:
        blob = None
        if rev is not None:
            try:
                blob = subprocess.check_output(
                    ["git", "show", "%s:%s" % (rev, rel)], cwd=ROOT,
                    stderr=subprocess.DEVNULL)
            except subprocess.CalledProcessError:
                continue
        s = set()
        for _, text in read_any(os.path.join(ROOT, rel), blob):
            for m in NUM.findall(text):
                s.add(canon(m))
        out[rel] = s
    return out


def sentences(text, val):
    """その数値を含む文だけを返す。"""
    return [b.strip() for b in BUN.findall(text)
            if val in canon(b) or val in canon(b.replace(" ", ""))]


# ------------------------------------------------------------ 出力
def show_map(idx, val, limit=40):
    hits = idx.get(canon(val), [])
    print("　%s … %d箇所" % (val, len(hits)))
    for rel, place, text in hits[:limit]:
        t = text if len(text) < 90 else text[:88] + "…"
        print("    %s ／ %s　%s" % (os.path.basename(rel), place, t))
    if len(hits) > limit:
        print("    … ほか %d箇所" % (len(hits) - limit))
    return hits


def main(argv):
    global NUM
    if "--min" in argv:
        i = argv.index("--min")
        n = int(argv[i + 1])
        NUM = re.compile(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+\.\d+|\d{%d,}" % n)
        argv = argv[:i] + argv[i + 2:]
    csv_out = None
    if "--csv" in argv:
        i = argv.index("--csv")
        csv_out = argv[i + 1]
        argv = argv[:i] + argv[i + 2:]
    paths = targets()
    rows = []

    if "--map" in argv:
        val = argv[argv.index("--map") + 1]
        print("■ 同じ数値がどこにあるか")
        idx = build_index(paths)
        hits = show_map(idx, val)
        for rel, place, text in hits:
            rows.append(["所在", val, "", os.path.basename(rel), place, text[:200]])
        if csv_out:
            _csv(csv_out, rows)
        return 0

    if "--value" in argv:
        i = argv.index("--value")
        kie, arawa = {canon(argv[i + 1])}, {canon(argv[i + 2])}
        print("■ 指定された旧→新　%s → %s" % (argv[i + 1], argv[i + 2]))
        per_file = None
    else:
        rev = "HEAD"
        if "--rev" in argv:
            rev = argv[argv.index("--rev") + 1]
        print("■ %s と今の版を比べる" % rev)
        old, new = nums_of(paths, rev), nums_of(paths, None)
        per_file = {}
        kie, arawa = set(), set()
        for rel in paths:
            o, n = old.get(rel, set()), new.get(rel, set())
            if not o and not n:
                continue
            k, a = o - n, n - o
            if k or a:
                per_file[rel] = (k, a)
                kie |= k
                arawa |= a
        # 別のファイルに同じ数値が現れているものは「移した」のであって
        # 「直した」のではない。修正漏れの候補から外す。
        utsushi = kie & arawa
        kie -= arawa
        print("　変わったファイル %d件／消えた数値 %d件／現れた数値 %d件"
              "（うち別のファイルへ移したとみられるもの %d件は除く）"
              % (len(per_file), len(kie), len(arawa), len(utsushi)))
        for rel, (k, a) in sorted(per_file.items()):
            print("　　%s　消えた %d／現れた %d"
                  % (os.path.basename(rel), len(k), len(a)))

    if not kie:
        print("消えた数値がない。文章への影響・修正漏れの点検は行わない。")
        return 0

    idx = build_index(paths)

    # 残っている箇所が多い数値は、別の量として偶然に同じ数が出ているとみて
    # 要判断にする。少ない箇所に残っているものほど直し漏れである見込みが高い。
    OOI = 30
    print("\n■ 消えた数値が、まだ残っているところ（修正漏れの候補）")
    print("　箇所の少ないものから示す。%d箇所を超えるものは"
          "別の量として偶然に同じ数が出ているとみて要判断とする。" % OOI)
    n_f = 0
    nokori = sorted(((v, idx.get(v, [])) for v in kie if idx.get(v)),
                    key=lambda x: len(x[1]))
    for v, hits in nokori[:60]:
        han = FUTEKI if len(hits) <= OOI else YOHAN
        if han == FUTEKI:
            n_f += 1
        print("　[%s] %s … %d箇所" % (han, v, len(hits)))
        for rel, place, text in hits[:6]:
            t = text if len(text) < 90 else text[:88] + "…"
            print("    %s ／ %s　%s" % (os.path.basename(rel), place, t))
            rows.append([han, v, "", os.path.basename(rel), place, text[:200]])
        if len(hits) > 6:
            print("    … ほか %d箇所" % (len(hits) - 6))

    print("\n■ その数値を述べた文のうち、大小・向き・順位に触れているもの（要判断）")
    n_y = 0
    for v in sorted(kie, key=lambda x: -len(x))[:60]:
        for rel, place, text in idx.get(v, []):
            for b in sentences(text, v):
                if any(w in b for w in MUKI):
                    n_y += 1
                    print("　%s ／ %s　%s" % (os.path.basename(rel), place,
                                             b if len(b) < 120 else b[:118] + "…"))
                    rows.append([YOHAN, v, "", os.path.basename(rel), place, b[:200]])
                    break

    for path, err in MIJISSHI_LOG:
        rows.append([MIJISSHI, "", "", os.path.basename(path), "", err])
    if MIJISSHI_LOG:
        print("\n■ 読めなかったもの（未実施）")
        for path, err in MIJISSHI_LOG:
            print("　%s　%s" % (os.path.basename(path), err))

    print("\n修正漏れの候補 %d件／要判断 %d件／未実施 %d件"
          % (n_f, n_y, len(MIJISSHI_LOG)))
    print("※ 消えた数値は直した値とは限らない（区分の入替え・年度の進行でも消える）。"
          "中身は人が見て決める。")
    if csv_out:
        _csv(csv_out, rows)
    return 1 if n_f else 0


def _csv(path, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as fp:
        w = csv.writer(fp)
        w.writerow(["判定", "数値", "備考", "ファイル", "位置", "内容"])
        w.writerows(rows)
    print("書き出し", path)


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
