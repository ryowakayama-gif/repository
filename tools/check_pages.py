# -*- coding: utf-8 -*-
"""docx の紙面を実際に数えて、体裁の崩れを見つける。

使い方
    python3 tools/check_pages.py <pdf> <docx>

PDF は docx を LibreOffice で変換したもの。

    soffice --headless --convert-to pdf --outdir <dir> <docx>

数えるもの
  - 総ページ数
  - 中身の少ないページ（表紙・最終ページを除き400字未満）
  - 紙面をまたぐ表（表の先頭の行と最後の行が別のページにあるもの）

表の位置は、表題（「表N　…」）があればそれで、なければ
文書の並び順に沿って先頭の行・最後の行の文字を探して定める。
"""

import re
import sys

import pymupdf
from docx import Document


# PDF の文字の取り出しでは、丸数字や矢印が読み順から外れて置かれることが
# ある（「表6　調査①の要介護度分布」が「表6①調査の要介護度分布」になる）。
# 突き合わせの前に、その種の文字を両方から落としてそろえる。
# かっこ付きの数字（⑴⑵）も同じく読み順から外れる
# （「市町村内の一部地域において⑴」が「市町村内の一部地域に⑴おいて」になる）。
# ローマ数字（Ⅰ〜Ⅻ）は落とさない。「支援Ⅰ」「支援Ⅲ」のように
# これだけが違う行があり、落とすと別の行に当たってしまう。
_DROP = re.compile(r"[\s　①-⑳⓵-⓾⑴-⒇⒈-⒛→←↑↓⇒]")


def norm(s):
    return _DROP.sub("", s or "")


def main(pdf, docx):
    pages = [norm(p.get_text()) for p in pymupdf.open(pdf)]

    def pageof(s, frm=1):
        s = norm(s)
        if not s:
            return None
        for i in range(max(frm, 1) - 1, len(pages)):
            if s in pages[i]:
                return i + 1
        return None

    doc = Document(docx)
    capmap = {}
    for p in doc.paragraphs:
        m = re.match(r"表(\d+)[　 ]", p.text)
        if m:
            capmap[int(m.group(1))] = p.text

    split, unknown, cur = [], [], 1
    for n, t in enumerate(doc.tables, 1):
        # 表の先頭を探す手掛かり。表題があればそれを使う。
        # 1行1列の枠（囲み）も表である。先頭の行だけで探す
        # （従前は2行以上の表しか手掛かりを作らず、枠が必ず
        #  「位置を定められなかった表」として挙がっていた）。
        r0 = "".join(c.text for c in t.rows[0].cells)
        r1 = ("".join(c.text for c in t.rows[1].cells)
              if len(t.rows) > 1 else "")
        # 見出しの行の文字は紙面で折り返されると取り出しの順が変わるため、
        # 行をつないだものだけでは見つからないことがある
        # （表題を「表N　」の形で持たない文書で実際に起きた）。
        # 2行目（表の中身の最初の行）・見出しの行を別の手掛かりとしても探す。
        # 2行目を先に試すのは、見出しの行は他の表と同じ文字であることがあり、
        # 直前の表のページに当たってしまうためである。
        heads = [h for h in (capmap.get(n), r0 + r1, r1, r0) if h]
        head = heads[0]
        lastfull = "".join(c.text for c in t.rows[-1].cells)
        # 表題と表の並び順が一致しない文書では、直前の表のページから探すと
        # 見つからない。見つからなければ先頭のページから探し直す。
        pa = None
        for _h in heads:
            pa = pageof(_h, cur) or pageof(_h, 1)
            if pa is None:
                for _k in (40, 28, 18):
                    pa = pageof(_h[:_k], 1)
                    if pa:
                        break
            if pa:
                break
        # PDF の文字の取り出しでは矢印などが読み順から外れることがあるため、
        # 長さを段々短くして探す（60字で見つからないことが実際にあった）。
        last, pb = lastfull[:60], None
        for _k in (60, 40, 28, 18):
            last = lastfull[:_k]
            pb = pageof(last, pa or cur)
            if pb:
                break
        if pa:
            cur = pa
        if pa and pb and pa != pb:
            split.append((n, pa, pb, (capmap.get(n) or last)[:28]))
        elif pa is None or pb is None:
            # 位置を定められなかった表。黙って飛ばすと「紙面をまたぐ表0件」が
            # 数えた結果でなくなる（CLAUDE.md §4）。
            unknown.append((n, (capmap.get(n) or last)[:28],
                            "先頭" if pa is None else "最終行"))

    short = [(i, len(t)) for i, t in enumerate(pages, 1)
             if len(t) < 400 and 1 < i < len(pages)]
    print("総ページ %d" % len(pages))
    print("中身の少ないページ（表紙・最終ページを除く）:", short)
    print("紙面をまたぐ表 %d件（位置を定められなかった表 %d件）"
          % (len(split), len(unknown)))
    for x in split:
        print("   表%d 先頭p%d→最終行p%d  %s" % x)
    for n, lab, which in unknown:
        print("   未確定 表%d %s（%sを紙面で見つけられない）" % (n, lab, which))
    return len(split) + len(unknown)


if __name__ == "__main__":
    sys.exit(1 if main(sys.argv[1], sys.argv[2]) else 0)
