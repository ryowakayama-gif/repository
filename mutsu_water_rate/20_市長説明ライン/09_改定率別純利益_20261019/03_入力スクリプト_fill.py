import openpyxl, copy, warnings
from openpyxl.styles import PatternFill
warnings.filterwarnings('ignore')
SPW="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
exec(open(f"{SPW}/work/recalc.py").read().split("# --- 検証")[0])
import decimal
def rnd(v): return int(decimal.Decimal(v/1000).quantize(0, rounding=decimal.ROUND_HALF_UP))

wb=openpyxl.load_workbook(f"{SPW}/upload3/req.xlsx")
ws=wb["Sheet"]
NOFILL=PatternFill(fill_type=None)
ref=ws["E3"]                                   # 既記入行（現行料金）の書式を基準にする
I=range(4,10); COLS=["E","F","G","H","I","J"]

for rate,row in [(0.10,5),(0.20,7),(0.30,9)]:
    d=run(rate, True)                          # True＝二重計上を是正
    for col,i in zip(COLS,I):
        c=ws[f"{col}{row}"]
        c.value=rnd(d["r28"][i])
        c.number_format=ref.number_format       # '0;"△ "0'
        c.font=copy.copy(ref.font)              # Yu Gothic 11
        c.alignment=copy.copy(ref.alignment)    # 右寄せ
        c.fill=NOFILL                           # 入力指示の黄色を解除（E3:J3と同じ無地）

# 出典・前提の注記（既存の※行に合わせる）
note=ws["B12"]
note.value=("※当期純利益は投資財政計画の純損益。改定率はR11以降の給水収益に乗じて算定（資産維持費は純損益には算入されない）")
note.font=copy.copy(ws["B11"].font)

out=f"{SPW}/out/改定率別純利益シミュレーション.xlsx"
wb.save(out); print("保存:",out)
