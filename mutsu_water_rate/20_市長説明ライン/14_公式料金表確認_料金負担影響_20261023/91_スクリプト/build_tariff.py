# -*- coding: utf-8 -*-
import openpyxl, warnings
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
warnings.filterwarnings('ignore')
WD="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
F="Yu Gothic"
HDRF=Font(name=F,sz=10,b=True,color="FFFFFF"); HDRB=PatternFill("solid",fgColor="1C3C60")
TTL=Font(name=F,sz=13,b=True,color="1C3C60"); BF=Font(name=F,sz=10,b=True); NF=Font(name=F,sz=10)
NOTE=Font(name=F,sz=9,color="555555"); LBLB=PatternFill("solid",fgColor="E1EBF7")
TH=Border(*[Side(style="thin",color="BFCDDD")]*4)
BASE={"13mm・20mm":1826,"25mm":3289}; JURYO=284; FREE=10
RATES=[("現行料金",0.0),("改定率10％",0.10),("改定率20％",0.20),("改定率30％",0.30)]

wb=openpyxl.Workbook(); ws=wb.active; ws.title="料金負担影響"
ws["A1"]="料金改定による使用者負担の影響（一般用・税込）"; ws["A1"].font=TTL
ws["A2"]=("出典：むつ市上下水道局 水道料金表（令和8年7月使用分〜）。基本料金は基本水量10㎥込み、"
          "従量料金は10㎥超過分に284円/㎥。改定率は基本料金・従量料金に一律適用（端数切捨て）。")
ws["A2"].font=NOTE
r=4
for dia,b in BASE.items():
    c=ws.cell(row=r,column=1,value=f"口径 {dia}"); c.font=TTL; r+=1
    for j,v in enumerate(["区分","基本料金","従量単価/㎥","月10㎥","増加額","月20㎥","増加額","月30㎥","増加額"]):
        c=ws.cell(row=r,column=1+j,value=v); c.font=HDRF; c.fill=HDRB
        c.alignment=Alignment(horizontal="center"); c.border=TH
    hdr=r; r+=1
    cur={}
    for nm,rate in RATES:
        bb=int(b*(1+rate)); jj=int(JURYO*(1+rate))
        bills={v: bb+max(0,v-FREE)*jj for v in (10,20,30)}
        if rate==0: cur=dict(bills)
        c=ws.cell(row=r,column=1,value=nm); c.font=BF; c.fill=LBLB; c.border=TH
        for col,val in [(2,bb),(3,jj)]:
            cc=ws.cell(row=r,column=col,value=val); cc.font=NF; cc.number_format='#,##0'
            cc.alignment=Alignment(horizontal="right"); cc.border=TH
        col=4
        for v in (10,20,30):
            cc=ws.cell(row=r,column=col,value=bills[v]); cc.font=NF; cc.number_format='#,##0'
            cc.alignment=Alignment(horizontal="right"); cc.border=TH
            inc=bills[v]-cur[v]
            cd=ws.cell(row=r,column=col+1,value=(inc if inc else None))
            cd.font=NF; cd.number_format='+#,##0;-#,##0;"—"'
            cd.alignment=Alignment(horizontal="right"); cd.border=TH
            col+=2
        r+=1
    r+=1
ws.cell(row=r,column=1,value="参考：大口利用者（口径100mm・月1,000㎥）").font=TTL; r+=1
for j,v in enumerate(["区分","基本料金","従量単価/㎥","月1,000㎥","増加額"]):
    c=ws.cell(row=r,column=1+j,value=v); c.font=HDRF; c.fill=HDRB
    c.alignment=Alignment(horizontal="center"); c.border=TH
r+=1; base0=None
for nm,rate in RATES:
    bb=int(73150*(1+rate)); jj=int(JURYO*(1+rate)); bill=bb+1000*jj
    if rate==0: base0=bill
    c=ws.cell(row=r,column=1,value=nm); c.font=BF; c.fill=LBLB; c.border=TH
    for col,val,fmt in [(2,bb,'#,##0'),(3,jj,'#,##0'),(4,bill,'#,##0'),
                        (5,(bill-base0) if bill!=base0 else None,'+#,##0;-#,##0;"—"')]:
        cc=ws.cell(row=r,column=col,value=val); cc.font=NF; cc.number_format=fmt
        cc.alignment=Alignment(horizontal="right"); cc.border=TH
    r+=1
r+=1
for t in ["※口径40mm以上は料金表に基本水量の記載がないため、全量を従量計算した仮置き。市に確認が必要。",
          "※令和8年7〜9月使用分は物価高騰対策として口径13〜25mmの基本料金が1か月1,650円（税込）減額されている（従量料金は対象外）。",
          "※上表は減額前の料金表に基づく。下水道使用料は含まない。"]:
    ws.cell(row=r,column=1,value=t).font=NOTE; r+=1
ws.column_dimensions["A"].width=16
for col in "BCDEFGHI": ws.column_dimensions[col].width=13
out=f"{WD}/w8/料金負担影響試算.xlsx"; wb.save(out); print("保存:",out)
