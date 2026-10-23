# -*- coding: utf-8 -*-
import warnings; warnings.filterwarnings('ignore')
from pptx import Presentation
from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
WD="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
FONT="BIZ UDPゴシック"
NAVY=RGBColor(0x1C,0x3C,0x60); INK=RGBColor(0x19,0x1C,0x23); WHITE=RGBColor(0xFF,0xFF,0xFF)
LBL=RGBColor(0xE1,0xEB,0xF7); RA=RGBColor(0xFF,0xFF,0xFF); RB=RGBColor(0xF4,0xF7,0xFB)
RED=RGBColor(0xBE,0x41,0x41)
def ea(run):
    rPr=run._r.get_or_add_rPr()
    for t in ("a:latin","a:ea","a:cs"):
        e=rPr.find(qn(t))
        if e is None: e=rPr.makeelement(qn(t),{}); rPr.append(e)
        e.set("typeface",FONT)
def txt(s,l,t,w,h,v,size,bold=False,color=INK):
    tb=s.shapes.add_textbox(Cm(l),Cm(t),Cm(w),Cm(h)); tf=tb.text_frame; tf.word_wrap=True
    for i,line in enumerate(v.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        r=p.add_run(); r.text=line
        r.font.size=Pt(size); r.font.bold=bold; r.font.name=FONT; r.font.color.rgb=color; ea(r)
    return tb
def cell(c,v,*,bold=False,color=INK,bg=None,al=PP_ALIGN.CENTER,size=9):
    tf=c.text_frame; tf.clear(); p=tf.paragraphs[0]; p.alignment=al
    r=p.add_run(); r.text=v; r.font.size=Pt(size); r.font.bold=bold
    r.font.name=FONT; r.font.color.rgb=color; ea(r)
    if bg is not None: c.fill.solid(); c.fill.fore_color.rgb=bg

BASE=[("13mm・20mm",1826),("25mm",3289)]; J=284; FREE=10
RATES=[("現行料金",0.0),("10％改定",0.10),("20％改定",0.20),("30％改定",0.30)]
for tag in ["A_現行基準","B_是正基準"]:
    prs=Presentation(f"{WD}/w8/04_図表編_{tag}.pptx")
    s=prs.slides.add_slide(prs.slide_layouts[6])
    for sh in list(s.shapes): sh._element.getparent().remove(sh._element)
    txt(s,1.14,0.64,30.7,1.14,"料金改定による使用者負担の影響",22,True,NAVY)
    txt(s,1.22,1.98,30.5,0.64,"一般家庭（月20㎥）では10％改定で月462円、30％改定で月1,397円の増加",12)
    txt(s,32.0,18.03,1.4,0.46,"4",10)
    rows=1+len(BASE)*len(RATES); cols=8
    gf=s.shapes.add_table(rows,cols,Cm(1.6),Cm(3.3),Cm(30.6),Cm(1.0*rows))
    tb=gf.table
    for i,w in enumerate([4.2,3.6,3.6,3.6,3.6,4.0,4.0,4.0]): tb.columns[i].width=Cm(w)
    for j,h in enumerate(["口径","区分","基本料金","従量単価/㎥","月10㎥","月20㎥","月30㎥","月20㎥の増加額"]):
        cell(tb.cell(0,j),h,bold=True,color=WHITE,bg=NAVY)
    r=1
    for dia,b in BASE:
        cur={}
        for k,(nm,rate) in enumerate(RATES):
            bb=int(b*(1+rate)); jj=int(J*(1+rate))
            bills={v: bb+max(0,v-FREE)*jj for v in (10,20,30)}
            if rate==0: cur=dict(bills)
            bg = RA if k%2==0 else RB
            cell(tb.cell(r,0),dia if k==0 else "",bold=True,bg=LBL,al=PP_ALIGN.LEFT)
            cell(tb.cell(r,1),nm,bold=True,bg=LBL,al=PP_ALIGN.LEFT)
            cell(tb.cell(r,2),f"{bb:,}円",bg=bg)
            cell(tb.cell(r,3),f"{jj:,}円",bg=bg)
            for c2,v in zip((4,5,6),(10,20,30)): cell(tb.cell(r,c2),f"{bills[v]:,}円",bg=bg)
            inc=bills[20]-cur[20]
            cell(tb.cell(r,7),("—" if inc==0 else f"+{inc:,}円"),bg=bg,bold=(inc!=0),color=(INK if inc==0 else RED))
            r+=1
    txt(s,1.65,3.4+1.0*rows+0.4,30.5,2.2,
        "・基本料金には基本水量10㎥が含まれ、従量料金は10㎥超過分に課される。\n"
        "・改定率は基本料金・従量料金に一律適用。月10㎥以下の世帯は基本料金の増加のみとなる。\n"
        "・参考：口径100mm・月1,000㎥の大口利用者では、10％改定で月35,315円、30％改定で月106,945円の増加。",11)
    txt(s,1.65,16.3,30.5,1.5,
        "出典：むつ市上下水道局 水道料金表（令和8年7月使用分〜）。税込。下水道使用料は含まない。"
        "令和8年7〜9月使用分は物価高騰対策として口径13〜25mmの基本料金が1か月1,650円減額されているが、上表は減額前の料金表による。"
        "口径40mm以上は料金表に基本水量の記載がないため、大口の例は全量を従量計算した仮置き。",10,False,RGBColor(0x55,0x55,0x55))
    out=f"{WD}/w8/04_図表編_{tag}_v2.pptx"
    prs.save(out); print(f"{tag}: {len(prs.slides._sldIdLst)}枚（料金負担影響を追加）")
