# -*- coding: utf-8 -*-
import json, warnings
warnings.filterwarnings('ignore')
from pptx import Presentation
SP="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
V=json.load(open(f"{SP}/work/vals.json"))["B"]
def fmt(n): return f"△{abs(n):,}" if n<0 else f"{n:,}"

def set_para(p, text):
    """run0の書式を保ったまま段落テキストを差し替える"""
    assert p.runs, "runがない段落"
    p.runs[0].text=text
    for r in list(p.runs)[1:]: r._r.getparent().remove(r._r)

def set_cell(cell, text): set_para(cell.text_frame.paragraphs[0], text)

prs=Presentation(f"{SP}/pat/base15.pptx")
log=[]

# ---------------- P9（index 8）----------------
s=prs.slides[8]; sh=list(s.shapes)
tb=sh[3].table
for j,v in enumerate(V["p9_hoten"][3:]):          # R11..R16 → 列4..9
    set_cell(tb.cell(8,j+4), fmt(v)); log.append(f"P9 表r8c{j+4}={fmt(v)}")
set_para(sh[4].text_frame.paragraphs[0],
    "R9〜R11は補填財源残高がマイナスとなり、一般会計支援、追加起債、投資平準化、料金改定の組合せが必要となる。")
log.append("P9 本文p0 書換（R9・R10→R9〜R11）")
set_para(sh[5].text_frame.paragraphs[0],
    "単位：百万円。補填財源残高はR7末183百万円を起点とした簡便推計。"
    "「改定なし」は現行料金を維持した場合、「改定あり」はR11以降にスライド10の資産維持費0％ケース"
    "（必要改定率7.3％）を反映した場合。営業収益の二重計上を是正済み。")
log.append("P9 脚注 書換（是正の明記）")

# ---------------- P10（index 9）----------------
s=prs.slides[9]; sh=list(s.shapes)
tb=sh[4].table
for i,p in enumerate(V["p10"]):
    set_cell(tb.cell(i+1,2), f"{p['keijo']:.1f}%")      # R11経常収支比率
    set_cell(tb.cell(i+1,5), fmt(p['genkin']))          # R16現預金
    log.append(f"P10 表r{i+1}: 経常={p['keijo']:.1f}% 現預金={fmt(p['genkin'])}")
set_para(sh[1].text_frame.paragraphs[0],
    "経常収支比率・料金回収率・現預金（営業費用・営業収益 修正後）")
log.append("P10 サブタイトル 書換")
set_para(sh[5].text_frame.paragraphs[3],
    "・現行料金のままでは補填財源がR9以降マイナスで推移（R9末△26・R10末△83百万円、現預金もR10末△24百万円）。"
    "7.3％改定でもR11末は△40百万円で、回復はR12以降となる。")
log.append("P10 本文p3 書換（R11改定で回復→R12以降）")
set_para(sh[6].text_frame.paragraphs[0],
    "※R7は決算確定値（累計4条不足に反映）。現預金は百万円。企業債充当率100%（建設改良費−国庫−工事負担金、"
    "利率3%・40年元利均等）、過年度建設仮勘定純額12百万円のR9借入、建設仮勘定の減価償却・戻入を本体へ統合済み。"
    "補填財源はR7末183百万円（0708修正版補てん財源経過表）起点の簡便推計。"
    "営業収益は「うち料金改定増収」の二重計上を是正済み。")
log.append("P10 脚注 書換（是正の明記）")

# ---------------- P11（index 10）----------------
s=prs.slides[10]; sh=list(s.shapes)
tb=sh[4].table
for i,key in enumerate(["0%","1%","2%","3%"]):
    for j,v in enumerate(V["p11"][key]["vals"]):
        set_cell(tb.cell(i+1,j+2), fmt(v))
    log.append(f"P11 表r{i+1} 維持費{key}: "+" ".join(fmt(x) for x in V['p11'][key]['vals']))
t=V["p11"]
set_para(sh[5].text_frame.paragraphs[0],
    f"・資産維持費0％（7.3％改定）はR13まで黒字だがR14以降は再び赤字となり、R11-R16累計も"
    f"{fmt(t['0%']['total'])}百万円。1％以上は全期間黒字を維持。"
    f"R11-R16累計は0%={fmt(t['0%']['total'])}、1%={fmt(t['1%']['total'])}、"
    f"2%={fmt(t['2%']['total'])}、3%={fmt(t['3%']['total'])}百万円。")
log.append("P11 本文p0 書換（0%は全期間黒字でない旨＋累計更新）")
set_para(sh[5].text_frame.paragraphs[2],
    "・0％では10年間の黒字を維持できないため、1〜2％を現実的検討帯、3％を上限として、段階改定・5年周期で調整する。")
log.append("P11 本文p2 書換")
set_para(sh[6].text_frame.paragraphs[0],
    "単位：百万円。改定後3条純利益は各シナリオSIMの純損益（r28）。4条収支不足は資本的収入−資本的支出（r54）で"
    "全ケース共通。R7-R16累計4条不足は△4,887百万円で改定率に依存しない。"
    "営業収益の二重計上を是正済み。")
log.append("P11 脚注 書換（是正の明記）")

out=f"{SP}/pat/B_市長説明資料_是正基準.pptx"
prs.save(out)
print("保存:",out); print(f"変更点 {len(log)}件")
for l in log: print("  -",l)
