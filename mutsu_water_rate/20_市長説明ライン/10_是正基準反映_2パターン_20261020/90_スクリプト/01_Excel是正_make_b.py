import openpyxl, warnings
warnings.filterwarnings('ignore')
SP="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
SRC="/home/user/repository/mutsu_water_rate/20_市長説明ライン/01_final_deliverables/02_mutsu_toushizaisei_plan_v7_final.xlsx"
wb=openpyxl.load_workbook(SRC)
changed=[]
for sn in [s for s in wb.sheetnames if s.startswith("SIM")]:
    ws=wb[sn]
    # r9（営業収益）の集計範囲から r11（うち料金改定増収）を除外
    assert (ws.cell(row=11,column=2).value or "").__contains__("うち") or "うち" in str(ws.cell(row=11,column=1).value), sn
    for col in range(2,12):
        c=ws.cell(row=9,column=col)
        L=c.column_letter
        before=c.value
        assert before==f"=SUM({L}10:{L}12)", (sn,L,before)
        c.value=f"=SUM({L}10,{L}12)"
        changed.append(f"{sn}!{L}9")
    # 備考を追記（2行目の注記末尾）
    note=ws.cell(row=2,column=1)
    if isinstance(note.value,str) and "営業収益" not in note.value:
        note.value=note.value+"／営業収益は給水収益＋その他営業収益（『うち料金改定増収』は内数のため集計から除外）"
print("修正セル数:",len(changed),"（12シート×10列）")
print("例:",changed[:3],"...",changed[-3:])
wb.save(f"{SP}/pat/B_是正基準.xlsx")
print("保存: B_是正基準.xlsx")
