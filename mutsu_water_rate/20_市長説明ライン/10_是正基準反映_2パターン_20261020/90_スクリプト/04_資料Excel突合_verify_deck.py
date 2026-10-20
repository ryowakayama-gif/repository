import json, openpyxl, warnings, decimal
warnings.filterwarnings('ignore')
from pptx import Presentation
SP="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
def rnd(v): return int(decimal.Decimal(v/1000).quantize(0,rounding=decimal.ROUND_HALF_UP))
def fmt(n): return f"△{abs(n):,}" if n<0 else f"{n:,}"
def num(s): 
    s=s.strip().replace(",","").replace("%","")
    return -float(s.replace("△","")) if s.startswith("△") else float(s)

for tag,ppt,xls in [("A","pat/base15.pptx","pat/A_現行基準.xlsx"),
                    ("B","pat/B_市長説明資料_是正基準.pptx","pat/B_是正基準.xlsx")]:
    wb=openpyxl.load_workbook(f"{SP}/{xls}",data_only=True)
    prs=Presentation(f"{SP}/{ppt}")
    sim=wb["SIM_R8予算置換_0%"]; cmp_=wb["03_SIM比較_R8予算反映"]
    bad=[]
    # P9 r8 補填財源（改定あり）R8..R16
    tb=[s for s in prs.slides[8].shapes if getattr(s,'has_table',False)][0].table
    for j,c in enumerate(range(3,12)):
        want=rnd(sim.cell(row=72,column=c).value); got=num(tb.cell(8,j+1).text)
        if abs(want-got)>0.5: bad.append(f"P9 R{8+j}: 資料{got} vs Excel{want}")
    # P9 r7 補填財源（改定なし）は両パターン共通（改定率0）
    # P10
    tb=[s for s in prs.slides[9].shapes if getattr(s,'has_table',False)][0].table
    for i,r in enumerate(range(4,8)):
        for col,xc,lab in [(2,4,"R11経常"),(3,5,"R11料金"),(4,6,"維持費込")]:
            want=cmp_.cell(row=r,column=xc).value; got=num(tb.cell(i+1,col).text)
            if abs(want-got)>0.06: bad.append(f"P10 r{i+1} {lab}: 資料{got} vs Excel{want:.2f}")
        want=rnd(cmp_.cell(row=r,column=7).value); got=num(tb.cell(i+1,5).text)
        if abs(want-got)>0.5: bad.append(f"P10 r{i+1} R16現預金: 資料{got} vs Excel{want}")
    # P11
    tb=[s for s in prs.slides[10].shapes if getattr(s,'has_table',False)][0].table
    for i,s_ in enumerate(["0%","1%","2%","3%"]):
        w=wb[f"SIM_R8予算置換_{s_}"]
        for j,c in enumerate(range(6,12)):
            want=rnd(w.cell(row=28,column=c).value); got=num(tb.cell(i+1,j+2).text)
            if abs(want-got)>0.5: bad.append(f"P11 維持費{s_} R{11+j}: 資料{got} vs Excel{want}")
    # P7（両パターン共通・現行料金ベース）
    tb=[s for s in prs.slides[6].shapes if getattr(s,'has_table',False)][0].table
    o=wb["02_R8予算_年次計画置換"]
    for ri,orow in [(1,7),(2,5),(3,17)]:
        for j,c in enumerate(range(3,12)):
            want=rnd(o.cell(row=orow,column=c).value); got=num(tb.cell(ri,j+1).text)
            if abs(want-got)>0.5: bad.append(f"P7 r{ri} R{8+j}: 資料{got} vs Excel{want}")
    print(f"■ パターン{tag}: 資料とExcelの突合 → {'全一致' if not bad else '★不一致'}")
    for b in bad: print("    ",b)
