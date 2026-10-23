# -*- coding: utf-8 -*-
import openpyxl, warnings, decimal
warnings.filterwarnings('ignore')
WD="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
REPO="/home/user/repository/mutsu_water_rate/20_市長説明ライン/10_是正基準反映_2パターン_20261020"
exec(open(f"{WD}/w6/recalc.py").read().split("# --- 検証")[0])
def rnd(v): return int(decimal.Decimal(v/1000).quantize(0,rounding=decimal.ROUND_HALF_UP))

def scen(xlsx):
    wb=openpyxl.load_workbook(xlsx,data_only=True)
    keijo=[wb["03_SIM比較_R8予算反映"].cell(row=r,column=4).value for r in range(4,8)]
    jun={i:[rnd(wb[f"SIM_R8予算置換_{s}"].cell(row=28,column=c).value) for c in range(6,12)]
         for i,s in enumerate(["0%","1%","2%","3%"])}
    return keijo,jun

for tag,fix,xlsx in [("A_現行基準",False,f"{REPO}/A_現行基準/02_投資財政計画_A_現行基準.xlsx"),
                     ("B_是正基準",True, f"{REPO}/B_是正基準/02_投資財政計画_B_是正基準.xlsx")]:
    keijo,jun=scen(xlsx)
    wb=openpyxl.load_workbook(f"{WD}/upload4/fmt.xlsx")   # 市提供フォーマット原本
    changed=[]

    # Sheet（可視）：改定率10/20/30％の当期純利益 ─ 値のみ
    ws=wb["Sheet"]
    for rate,row in [(0.10,5),(0.20,7),(0.30,9)]:
        d=run(rate,fix)
        for j,col in enumerate("EFGHIJ"):
            ws[f"{col}{row}"].value=rnd(d["r28"][4+j])
            changed.append(f"Sheet!{col}{row}")

    # Sheet2（非表示）：資料P.10/P.11の元数値 ─ 値のみ
    s2=wb["Sheet2"]
    for i in range(4):
        s2.cell(row=3+i,column=4).value=round(keijo[i]/100,3)      # D列 R11経常収支比率
        changed.append(f"Sheet2!D{3+i}")
        for j in range(6):
            s2.cell(row=3+i,column=9+j).value=jun[i][j]            # I〜N列 改定後純利益
            changed.append(f"Sheet2!{'IJKLMN'[j]}{3+i}")

    out=f"{WD}/w6/パワポ用資料_{tag}.xlsx"
    wb.save(out)
    print(f"保存: {out}  （値を更新したセル {len(changed)}個）")
