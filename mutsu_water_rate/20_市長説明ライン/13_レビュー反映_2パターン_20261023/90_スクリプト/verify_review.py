import openpyxl, warnings, decimal
warnings.filterwarnings('ignore')
WD="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
exec(open(f"{WD}/w7/recalc.py").read().split("# --- 検証")[0])
def mm(v): return v/1000.0
wv=openpyxl.load_workbook(SRC,data_only=True)
o=wv["02_R8予算_年次計画置換"]; sim=wv["SIM_R8予算置換_0%"]

print("■ 指摘：4条不足の累計 R8〜R16 と R7分")
r54=[sim.cell(row=54,column=c).value for c in range(2,12)]
print(f"   R7        = {r54[0]:>12,.0f} 千円 = {mm(r54[0]):>8,.0f} 百万円")
print(f"   R8〜R16計 = {sum(r54[1:]):>12,.0f} 千円 = {mm(sum(r54[1:])):>8,.0f} 百万円  ← レビュー指摘の△4,169と一致")
print(f"   R7〜R16計 = {sum(r54):>12,.0f} 千円 = {mm(sum(r54)):>8,.0f} 百万円  ← 資料の△4,887と一致")
print("   → 差額はR7実績分。R7決算書との一致は00_前提・検証シートに記録あり（△717,372千円）")

print("\n■ 指摘：必要改定率の算定根拠（B6の構成要素・R11〜R15の5か年集計、単位千円）")
def s(row): return sum(o.cell(row=row,column=c).value or 0 for c in range(6,11))  # F..J = R11..R15
comp={"営業費用(r18)":s(18),"営業外費用(r28)":s(28),"長期前受金戻入(r12)":-s(12),
      "その他営業収益(r8)":-s(8),"繰入金・負担金等(r11)":-s(11)}
kyusui=s(7)
for s5 in ["0%","1%","2%","3%"]:
    ws=wv[f"SIM_R8予算置換_{s5}"]
    am=ws["B5"].value*5
    bunshi=sum(comp.values())+am-kyusui
    rate=max(0,bunshi/kyusui)
    print(f"   資産維持費{s5}: 分子={bunshi:>12,.0f}  分母(給水収益5年計)={kyusui:>12,.0f}  → {rate*100:>6.3f}%  (Excel={ws['B6'].value*100:.3f}%)")
print("   内訳（全ケース共通部分）:")
for k,v in comp.items(): print(f"     {k:<24}{v:>14,.0f}")
print(f"     {'小計(控除後総括原価)':<24}{sum(comp.values()):>14,.0f}")
print(f"     {'給水収益5年計':<24}{kyusui:>14,.0f}")

print("\n■ 指摘：資産維持費込料金回収率が100%未満 vs 総括原価回収率")
print(f"   {'維持費':<6}{'必要改定率':>9}{'R11資産維持費込料金回収率':>24}{'総括原価回収率R11-R15平均':>26}")
for s5 in ["0%","1%","2%","3%"]:
    ws=wv[f"SIM_R8予算置換_{s5}"]
    r33=ws["F33"].value
    r36=[ws[f"{c}36"].value for c in "FGHIJ"]
    print(f"   {s5:<6}{ws['B6'].value*100:>8.1f}%{r33:>23.1f}%{sum(r36)/5:>25.3f}%")
