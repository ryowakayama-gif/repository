import openpyxl, warnings
warnings.filterwarnings('ignore')
WD="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
exec(open(f"{WD}/w7/recalc.py").read().split("# --- 検証")[0])
wv=openpyxl.load_workbook(SRC,data_only=True)
B=openpyxl.load_workbook("/home/user/repository/mutsu_water_rate/20_市長説明ライン/10_是正基準反映_2パターン_20261020/B_是正基準/02_投資財政計画_B_是正基準.xlsx",data_only=True)
sim=B["SIM_R8予算置換_0%"]; Y=["R7","R8","R9","R10","R11","R12","R13","R14","R15","R16"]
def row(r): return [sim.cell(row=r,column=c).value for c in range(2,12)]
def mm(v): return round(v/1000)
r28,r21,r22,r14=row(28),row(21),row(22),row(14)
r54,r57,r59,r72=row(54),row(57),row(59),row(72)
print("■ 補填財源残高の分解（是正基準・資産維持費0%／単位：百万円）")
hdr=f"{'':<26}"+"".join(f"{y:>8}" for y in Y)
print(hdr)
for nm,v in [("当年度発生① 純損益",r28),("② 減価償却費",r21),("③ 資産減耗・除却費",r22),
             ("④ 長期前受金戻入(控除)",[-x for x in r14])]:
    print(f"  {nm:<25}"+"".join(f"{mm(x):>8,}" for x in v))
print(f"  {'当年度発生計 A=①+②+③+④':<25}"+"".join(f"{mm(x):>8,}" for x in r57))
print(f"  {'4条収支不足額 B':<25}"+"".join(f"{mm(x):>8,}" for x in r54))
print(f"  {'当年度増減 A+B':<25}"+"".join(f"{mm(a+b):>8,}" for a,b in zip(r57,r54)))
print("  " + "-"*(25+8*10))
print(f"  {'補填財源 期首':<25}"+"".join(f"{mm(r72[i]-(r57[i]+r54[i])) if i>0 else 0:>8,}" for i in range(10)))
print(f"  {'補填財源 期末':<25}"+"".join(f"{mm(x):>8,}" for x in r72))
print(f"  {'現預金 期末':<25}"+"".join(f"{mm(x):>8,}" for x in r59))
print(f"  {'差（補填財源−現預金）':<25}"+"".join(f"{mm(r72[i]-r59[i]):>8,}" for i in range(10)))
print("  → 両者の差は全年度△59百万円で一定（起点の差：R7末 補填財源183 − 現預金243）")

print("\n■ R7→R8繰越工事182,069千円の感度（是正基準・資産維持費0%／単位：百万円）")
print("   繰越工事は資本的支出の増。財源充当154,100千円の扱いで影響幅が変わる")
for nm,shortfall in [("ケースa：財源充当154,100千円を全額確保（純増28百万円）",-27969),
                     ("ケースb：財源充当なし（純増182百万円）",-182069)]:
    print(f"   {nm}")
    h=[r72[i]+(shortfall/1000*1000 if i>=1 else 0) for i in range(10)]
    g=[r59[i]+(shortfall/1000*1000 if i>=1 else 0) for i in range(10)]
    print(f"     {'補填財源 期末':<16}"+"".join(f"{mm(x):>8,}" for x in h))
    print(f"     {'現預金 期末':<16}"+"".join(f"{mm(x):>8,}" for x in g))
print("   ※ R8以降一律シフトの簡便評価。企業債充当に回る場合は元利償還が後年度に加算されるため別途精査要")
