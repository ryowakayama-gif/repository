import openpyxl, warnings, decimal
warnings.filterwarnings('ignore')
from pptx import Presentation
WD="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
def rnd(v): return int(decimal.Decimal(v/1000).quantize(0,rounding=decimal.ROUND_HALF_UP))
def num(s):
    s=s.strip().replace(",","").replace("%","").replace("△ ","△")
    return -float(s.replace("△","")) if s.startswith("△") else float(s)
for tag in ["A_現行基準","B_是正基準"]:
    xl=openpyxl.load_workbook(f"{WD}/w7/out/投資財政計画_{tag}.xlsx",data_only=True)
    rt=openpyxl.load_workbook(f"{WD}/w7/out/改定率別純利益_{tag}.xlsx",data_only=True)
    pp=Presentation(f"{WD}/w7/out/市長説明資料_{tag}.pptx")
    cg=Presentation(f"{WD}/w7/out/図表編_{tag}.pptx")
    sim=xl["SIM_R8予算置換_0%"]; cmp_=xl["03_SIM比較_R8予算反映"]
    bad=[]
    # 1) 資料P9 ↔ Excel r72
    tb=[s for s in pp.slides[8].shapes if getattr(s,'has_table',False)][0].table
    for j,c in enumerate(range(3,12)):
        w=rnd(sim.cell(row=72,column=c).value); g=num(tb.cell(8,j+1).text)
        if abs(w-g)>0.5: bad.append(f"資料P9 R{8+j}: {g} vs Excel {w}")
    # 2) 資料P10 ↔ Excel 03_SIM比較
    tb=[s for s in pp.slides[9].shapes if getattr(s,'has_table',False)][0].table
    for i in range(4):
        for col,xc in [(2,4),(3,5),(4,6)]:
            w=cmp_.cell(row=4+i,column=xc).value; g=num(tb.cell(i+1,col).text)
            if abs(w-g)>0.06: bad.append(f"資料P10 r{i+1}c{col}: {g} vs Excel {w:.2f}")
        w=rnd(cmp_.cell(row=4+i,column=7).value); g=num(tb.cell(i+1,5).text)
        if abs(w-g)>0.5: bad.append(f"資料P10 r{i+1} R16現預金: {g} vs Excel {w}")
    # 3) 資料P11 ↔ Excel r28 ↔ 改定率別Excel Sheet2
    tb=[s for s in pp.slides[10].shapes if getattr(s,'has_table',False)][0].table
    s2=rt["Sheet2"]
    for i,s in enumerate(["0%","1%","2%","3%"]):
        w=xl[f"SIM_R8予算置換_{s}"]
        for j,c in enumerate(range(6,12)):
            v=rnd(w.cell(row=28,column=c).value)
            g=num(tb.cell(i+1,j+2).text)
            h=s2.cell(row=3+i,column=9+j).value
            if abs(v-g)>0.5: bad.append(f"資料P11 {s} R{11+j}: {g} vs Excel {v}")
            if v!=h: bad.append(f"改定率別Sheet2 {s} R{11+j}: {h} vs Excel {v}")
    # 4) 改定率別Excel Sheet2 D列 ↔ Excel
    for i in range(4):
        w=cmp_.cell(row=4+i,column=4).value; h=s2.cell(row=3+i,column=4).value*100
        if abs(w-h)>0.06: bad.append(f"改定率別Sheet2 D{3+i}: {h:.1f} vs Excel {w:.2f}")
    # 5) 図表編 ↔ Excel
    c1=[s for s in cg.slides[0].shapes if s.has_chart][0].chart
    ho=[rnd(sim.cell(row=72,column=c).value) for c in range(3,12)]
    gv=[v for v in list(c1.series)[1].values]
    if [round(x) for x in gv]!=ho: bad.append(f"図表編S1 改定あり: {gv} vs {ho}")
    # 6) 06シート 検証列
    s6=xl["06_必要改定率算定根拠"]
    hr=next(r for r in range(1,40) if s6.cell(row=r,column=1).value=="資産維持費率")
    for r in range(hr+1,hr+5):
        if s6.cell(row=r,column=8).value!="一致": bad.append(f"06シート r{r} 検証={s6.cell(row=r,column=8).value}")
    print(f"■ {tag}: 資料／Excel／改定率別Excel／図表編 の4点突合 → {'全一致' if not bad else '★不一致'}")
    for b in bad: print("    ",b)
