# -*- coding: utf-8 -*-
import openpyxl, warnings, decimal, shutil
warnings.filterwarnings('ignore')
WD="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
exec(open(f"{WD}/w7/recalc.py").read().split("# --- 検証")[0])
def rnd(v): return int(decimal.Decimal(v/1000).quantize(0,rounding=decimal.ROUND_HALF_UP))

def scen(xlsx):
    wb=openpyxl.load_workbook(xlsx,data_only=True)
    keijo=[wb["03_SIM比較_R8予算反映"].cell(row=r,column=4).value for r in range(4,8)]
    jun={i:[rnd(wb[f"SIM_R8予算置換_{s}"].cell(row=28,column=c).value) for c in range(6,12)]
         for i,s in enumerate(["0%","1%","2%","3%"])}
    return keijo,jun

def fill(tmpl,out,fix,xlsx):
    keijo,jun=scen(xlsx)
    wb=openpyxl.load_workbook(tmpl)
    ws=wb["Sheet"]
    for rate,row in [(0.10,5),(0.20,7),(0.30,9)]:
        d=run(rate,fix)
        for j,col in enumerate("EFGHIJ"):
            ws[f"{col}{row}"].value=rnd(d["r28"][4+j])
    s2=wb["Sheet2"]
    for i in range(4):
        s2.cell(row=3+i,column=4).value=round(keijo[i]/100,3)
        for j in range(6):
            s2.cell(row=3+i,column=9+j).value=jun[i][j]
    wb.save(out)

REPO="/home/user/repository/mutsu_water_rate/20_市長説明ライン"
TMPL=f"{WD}/upload4/fmt.xlsx"
for tag,fix in [("A_現行基準",False),("B_是正基準",True)]:
    fill(TMPL,f"{WD}/w7/out/改定率別純利益_{tag}.xlsx",fix,f"{WD}/w7/out/投資財政計画_{tag}.xlsx")
    print(f"生成: 改定率別純利益_{tag}.xlsx")

# レビュー指摘①の原因ファイル（10/20納品のB版）をその場で是正
stale=f"{REPO}/10_是正基準反映_2パターン_20261020/B_是正基準/03_改定率別純利益_B_是正基準.xlsx"
fill(stale,stale,True,f"{WD}/w7/out/投資財政計画_B_是正基準.xlsx")
print("是正: 10_是正基準反映_2パターン_20261020/B_是正基準/03_改定率別純利益_B_是正基準.xlsx")
