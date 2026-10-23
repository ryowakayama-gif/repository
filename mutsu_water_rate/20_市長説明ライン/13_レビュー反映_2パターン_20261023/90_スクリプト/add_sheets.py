# -*- coding: utf-8 -*-
import openpyxl, warnings, shutil, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
warnings.filterwarnings('ignore')
WD="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
REPO="/home/user/repository/mutsu_water_rate/20_市長説明ライン/10_是正基準反映_2パターン_20261020"
O2="'02_R8予算_年次計画置換'"
SIM0="'SIM_R8予算置換_0%'"
SCEN=["0%","1%","2%","3%"]
HDRF=Font(name="Yu Gothic",sz=10,b=True,color="FFFFFF")
HDRB=PatternFill("solid",fgColor="1C3C60")
TTL=Font(name="Yu Gothic",sz=12,b=True,color="1C3C60")
BF=Font(name="Yu Gothic",sz=10,b=True)
NF=Font(name="Yu Gothic",sz=10)
NOTE=Font(name="Yu Gothic",sz=9,color="555555")
THIN=Border(*[Side(style="thin",color="BFCDDD")]*4)
def hdr(ws,row,vals,start=1):
    for j,v in enumerate(vals):
        c=ws.cell(row=row,column=start+j,value=v); c.font=HDRF; c.fill=HDRB
        c.alignment=Alignment(horizontal="center"); c.border=THIN
def lab(ws,row,text,col=1,bold=False):
    c=ws.cell(row=row,column=col,value=text); c.font=BF if bold else NF; c.border=THIN; return c
def num(ws,row,col,formula,fmt='#,##0;"△ "#,##0'):
    c=ws.cell(row=row,column=col,value=formula); c.font=NF; c.number_format=fmt
    c.alignment=Alignment(horizontal="right"); c.border=THIN; return c

def build(path):
    wb=openpyxl.load_workbook(path)
    for nm in ["06_必要改定率算定根拠","07_補填財源分解","08_繰越工事感度"]:
        if nm in wb.sheetnames: del wb[nm]

    # ===== 06 必要改定率算定根拠 =====
    ws=wb.create_sheet("06_必要改定率算定根拠")
    ws["A1"]="必要改定率の算定根拠（料金算定期間 R11〜R15 の5か年）"; ws["A1"].font=TTL
    ws["A2"]="改定率＝MAX(0,(Σ営業費用＋Σ営業外費用−Σ長期前受金戻入−Σその他営業収益−Σ繰入金・負担金等＋資産維持費×5−Σ給水収益)÷Σ給水収益)"
    ws["A2"].font=NOTE
    YR=["令和11年度","令和12年度","令和13年度","令和14年度","令和15年度"]
    hdr(ws,4,["項目"]+YR+["5か年計"])
    items=[("営業費用",18,1),("営業外費用",28,1),("長期前受金戻入（控除）",12,-1),
           ("その他営業収益（控除）",8,-1),("繰入金・負担金等（控除）",11,-1)]
    r=5
    for nm,orow,sign in items:
        lab(ws,r,nm)
        for j,col in enumerate("FGHIJ"):
            num(ws,r,2+j,f"={'' if sign>0 else '-'}{O2}!{col}{orow}")
        num(ws,r,7,f"=SUM(B{r}:F{r})"); r+=1
    lab(ws,r,"小計：控除後の原価（資産維持費を除く）",bold=True)
    for j in range(5): num(ws,r,2+j,f"=SUM({chr(66+j)}5:{chr(66+j)}{r-1})")
    num(ws,r,7,f"=SUM(B{r}:F{r})"); SUB=r; r+=1
    lab(ws,r,"給水収益（現行料金）",bold=True)
    for j,col in enumerate("FGHIJ"): num(ws,r,2+j,f"={O2}!{col}7")
    num(ws,r,7,f"=SUM(B{r}:F{r})"); KYU=r; r+=2

    ws.cell(row=r,column=1,value="シナリオ別の必要改定率").font=TTL; r+=1
    hdr(ws,r,["資産維持費率","年額資産維持費","資産維持費×5年","分子（不足額）","分母（給水収益5年計）","必要改定率","SIMシートのB6","検証"])
    r+=1
    for i,s in enumerate(SCEN):
        sh=f"'SIM_R8予算置換_{s}'"
        lab(ws,r,s,bold=True).alignment=Alignment(horizontal="center")
        num(ws,r,2,f"={sh}!$B$5")
        num(ws,r,3,f"={sh}!$B$5*5")
        num(ws,r,4,f"=$G${SUB}+C{r}-$G${KYU}")
        num(ws,r,5,f"=$G${KYU}")
        num(ws,r,6,f"=MAX(0,D{r}/E{r})",'0.000%')
        num(ws,r,7,f"={sh}!$B$6",'0.000%')
        c=ws.cell(row=r,column=8,value=f'=IF(ABS(F{r}-G{r})<0.0000001,"一致","★差異")')
        c.font=NF; c.alignment=Alignment(horizontal="center"); c.border=THIN
        r+=1
    r+=1
    for t in ["※単位：千円。資産維持費は損益計算書の費用ではなく、総括原価にのみ算入される（当期純利益には影響しない）。",
              "※「必要改定率」は算定期間R11〜R15の5か年で総括原価を回収する水準。単年度の料金回収率が100％に達することを意味しない。",
              "※料金回収率（給水収益÷給水原価対象費用）は分母からその他営業収益・繰入金等を控除しないため、総括原価回収率より低く出る。"]:
        ws.cell(row=r,column=1,value=t).font=NOTE; r+=1
    ws.column_dimensions["A"].width=34
    for col in "BCDEFGH": ws.column_dimensions[col].width=17

    # ===== 07 補填財源分解 =====
    ws=wb.create_sheet("07_補填財源分解")
    ws["A1"]="補填財源残高の分解（資産維持費0％ケース）"; ws["A1"].font=TTL
    ws["A2"]="年度末残高＝前年度末残高＋当年度発生補填財源＋資本的収支不足額。現預金は起点のみ異なる同一の増減構造。"
    ws["A2"].font=NOTE
    YR2=[f"令和{y}年度" for y in range(7,17)]
    hdr(ws,4,["項目"]+YR2)
    rows=[("① 当期純損益",28,1),("② 減価償却費",21,1),("③ 資産減耗・除却費",22,1),("④ 長期前受金戻入（控除）",14,-1)]
    r=5
    for nm,srow,sign in rows:
        lab(ws,r,nm)
        for j,col in enumerate("BCDEFGHIJK"):
            num(ws,r,2+j,f"={'' if sign>0 else '-'}{SIM0}!{col}{srow}")
        r+=1
    lab(ws,r,"当年度発生補填財源 A（①+②+③+④）",bold=True)
    for j in range(10): num(ws,r,2+j,f"=SUM({chr(66+j)}5:{chr(66+j)}{r-1})")
    A=r; r+=1
    lab(ws,r,"資本的収支不足額 B",bold=True)
    for j,col in enumerate("BCDEFGHIJK"): num(ws,r,2+j,f"={SIM0}!{col}54")
    Bq=r; r+=1
    lab(ws,r,"当年度増減 A＋B",bold=True)
    for j in range(10): num(ws,r,2+j,f"={chr(66+j)}{A}+{chr(66+j)}{Bq}")
    r+=2
    lab(ws,r,"補填財源 年度末残高",bold=True)
    for j,col in enumerate("BCDEFGHIJK"): num(ws,r,2+j,f"={SIM0}!{col}72")
    H=r; r+=1
    lab(ws,r,"現預金 年度末残高",bold=True)
    for j,col in enumerate("BCDEFGHIJK"): num(ws,r,2+j,f"={SIM0}!{col}59")
    G=r; r+=1
    lab(ws,r,"差（補填財源−現預金）")
    for j in range(10): num(ws,r,2+j,f"={chr(66+j)}{H}-{chr(66+j)}{G}")
    r+=2
    for t in ["※単位：千円。補填財源はR7末183,463千円（0708修正版 補てん財源経過表）、現預金はR7末242,588千円（R7決算）を起点とする簡便推計。",
              "※両者の差は起点の差（△59,125千円）で一定。補填財源は予算上の財源区分、現預金は実際の支払能力であり定義が異なる。",
              "※残高がマイナスとなることは、追加財源（一般会計繰入・追加起債）または投資の平準化・料金改定が必要であることを示す。"
              "それ自体で資金ショートや企業債償還不能が確定するものではない。",
              "※消費税資本的収支調整額・利益剰余金処分の時期差は未反映。"]:
        ws.cell(row=r,column=1,value=t).font=NOTE; r+=1
    ws.column_dimensions["A"].width=34
    for col in "BCDEFGHIJK": ws.column_dimensions[col].width=14

    # ===== 08 繰越工事感度 =====
    ws=wb.create_sheet("08_繰越工事感度")
    ws["A1"]="R7→R8 繰越工事182,069千円の感度分析（資産維持費0％ケース）"; ws["A1"].font=TTL
    ws["A2"]="本SIMのR8建設改良費（年次計画888,542千円）に繰越工事は未反映。実行時の影響幅を財源充当の前提別に示す。"
    ws["A2"].font=NOTE
    lab(ws,4,"繰越工事額",bold=True);      num(ws,4,2,"=182069")
    lab(ws,5,"うち財源充当額",bold=True);  num(ws,5,2,"=154100")
    lab(ws,6,"純増（財源未充当分）",bold=True); num(ws,6,2,"=B4-B5")
    hdr(ws,8,["ケース"]+[f"令和{y}年度" for y in range(7,17)])
    r=9
    for nm,shift in [("現行推計（繰越未反映）","0"),("ケースa：財源充当を全額確保（純増 $B$6）","-$B$6"),
                     ("ケースb：財源充当なし（全額 $B$4）","-$B$4")]:
        lab(ws,r,"補填財源 "+nm,bold=(r==9))
        for j,col in enumerate("BCDEFGHIJK"):
            adj = "0" if j==0 or shift=="0" else shift
            num(ws,r,2+j,f"={SIM0}!{col}72+{adj}")
        r+=1
    for nm,shift in [("現行推計（繰越未反映）","0"),("ケースa：財源充当を全額確保（純増 $B$6）","-$B$6"),
                     ("ケースb：財源充当なし（全額 $B$4）","-$B$4")]:
        lab(ws,r,"現預金 "+nm)
        for j,col in enumerate("BCDEFGHIJK"):
            adj = "0" if j==0 or shift=="0" else shift
            num(ws,r,2+j,f"={SIM0}!{col}59+{adj}")
        r+=1
    r+=1
    for t in ["※単位：千円。R8以降へ一律シフトさせた簡便評価。",
              "※繰越額が企業債に充当される場合は後年度に元金償還と支払利息が加算されるため、本表より影響は大きくなる。",
              "※繰越額の全額が新規の追加支出となるとは限らない。既計上額との重複整理と、市の執行見込の確定が必要。"]:
        ws.cell(row=r,column=1,value=t).font=NOTE; r+=1
    ws.column_dimensions["A"].width=42
    for col in "BCDEFGHIJK": ws.column_dimensions[col].width=14
    wb.save(path)
    return [s for s in wb.sheetnames]

for tag,src in [("A_現行基準",f"{REPO}/A_現行基準/02_投資財政計画_A_現行基準.xlsx"),
                ("B_是正基準",f"{REPO}/B_是正基準/02_投資財政計画_B_是正基準.xlsx")]:
    dst=f"{WD}/w7/out/投資財政計画_{tag}.xlsx"
    shutil.copy(src,dst)
    names=build(dst)
    print(f"{tag}: {len(names)}シート  追加3シート → {names[-3:]}")
