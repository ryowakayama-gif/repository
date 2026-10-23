# -*- coding: utf-8 -*-
import warnings, shutil, openpyxl
warnings.filterwarnings('ignore')
from pptx import Presentation
from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
WD="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
REPO="/home/user/repository/mutsu_water_rate/20_市長説明ライン/10_是正基準反映_2パターン_20261020"
FONT="BIZ UDPゴシック"

def set_para(p,text):
    assert p.runs, "run無し"
    p.runs[0].text=text
    for r in list(p.runs)[1:]: r._r.getparent().remove(r._r)

def find_para(slide,needle):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                if needle in p.text: return p
    return None

for tag in ["A_現行基準","B_是正基準"]:
    src=f"{REPO}/{tag}/01_市長説明資料_{tag}.pptx"
    dst=f"{WD}/w7/out/市長説明資料_{tag}.pptx"
    shutil.copy(src,dst)
    prs=Presentation(dst)
    log=[]

    # ① P2：日程を時点中立化（８月は既に経過）
    p=find_para(prs.slides[1],"到達点")
    set_para(p,"到達点：理事者説明で、料金改定検討を進める必要性と次年度以降の工程を合意する。")
    log.append("P2 到達点の月表記を削除（時点中立化）")

    # ② P8：R14以降の国庫補助金ゼロ前提を明記
    p=find_para(prs.slides[7],"単位：百万円。建設改良費は将来年次計画")
    set_para(p,"単位：百万円。建設改良費は将来年次計画、企業債は充当率100％方針による置換後。"
               "R14〜R16は将来年次計画上、国庫補助金・工事負担金の計上がないため企業債が建設改良費と同額となる"
               "（補助制度の適用可否は未確定のため、確定後に再計算が必要）。")
    log.append("P8 脚注にR14〜R16の国庫補助金ゼロ前提を追記")

    # ③ P9：断定表現を資金不足リスクへ緩和
    p=find_para(prs.slides[8],"現金が一時的に減る")
    set_para(p,"市長説明では「現金が一時的に減る」ではなく、"
               "「現行の投資・財源条件のままでは補填財源が不足し、予定工事の予算化や企業債元金償還に支障が生じるリスクがある」"
               "流れで説明する。")
    log.append("P9 本文の断定表現をリスク表現へ緩和")

    # ④ P9 脚注：簡便推計の限界と分解シート参照
    p=find_para(prs.slides[8],"起点とした簡便推計")
    base=p.runs[0].text
    set_para(p,base+"残高のマイナスは追加財源・投資平準化・料金改定のいずれかが必要であることを示すもので、"
                    "資金ショートや企業債償還不能が確定したことを意味しない。"
                    "内訳はExcel「07_補填財源分解」シート参照。")
    log.append("P9 脚注に簡便推計の限界と分解シート参照を追記")

    # ⑤ P10：必要改定率の算定根拠シート参照
    p=find_para(prs.slides[9],"必要改定率は、二重計上除外に加え")
    set_para(p,"・必要改定率は、二重計上除外に加え、企業債充当率100%の支払利息と建設仮勘定供用開始"
               "（減価償却・戻入）を総括原価へ統合した水準。算定根拠はExcel「06_必要改定率算定根拠」シート参照。")
    log.append("P10 本文に算定根拠シート参照を追記")

    # ⑥ P10：資産維持費3％の位置づけを議事録どおりに
    p=find_para(prs.slides[9],"3％は感度分析の上限")
    set_para(p,"・市長説明では、0％は収支均衡の比較下限、1〜2％を現実的な検討帯として示し、"
               "本来的には3％を目標として段階改定・5年後見直しを選択肢とする（議事録2回目③）。")
    log.append("P10 本文の3％の位置づけを議事録準拠に修正")

    # ⑦ P11：同様
    p=find_para(prs.slides[10],"3％を上限として、段階改定")
    cur=p.text
    new=("・0％では10年間の黒字を維持できないため、" if "0％では" in cur else "・")\
        +"1〜2％を現実的な検討帯とし、本来的には3％を目標として段階改定・5年周期で調整する。"
    set_para(p,new)
    log.append("P11 本文の3％の位置づけを議事録準拠に修正")

    prs.save(dst)
    print(f"【{tag}】{len(prs.slides._sldIdLst)}ページ／修正{len(log)}件")
    for l in log: print("   -",l)
