# -*- coding: utf-8 -*-
import openpyxl, copy, warnings, decimal
from openpyxl.styles import PatternFill
warnings.filterwarnings('ignore')
WD="/tmp/claude-0/-home-user-repository/05ae8dfb-f2f7-5c38-805c-45df053cf245/scratchpad"
REPO="/home/user/repository/mutsu_water_rate/20_市長説明ライン/10_是正基準反映_2パターン_20261020"
exec(open(f"{WD}/w4/recalc.py").read().split("# --- 検証")[0])
def rnd(v): return int(decimal.Decimal(v/1000).quantize(0,rounding=decimal.ROUND_HALF_UP))
NOFILL=PatternFill(fill_type=None)
NEGFMT='0;"△ "0'

# 資産維持費シナリオの値は再計算済みExcelから取得（権威ソース）
def scen(xlsx):
    wb=openpyxl.load_workbook(xlsx,data_only=True)
    cmp_=wb["03_SIM比較_R8予算反映"]
    keijo=[cmp_.cell(row=r,column=4).value for r in range(4,8)]          # R11経常収支比率(%)
    jun={}
    for i,s in enumerate(["0%","1%","2%","3%"]):
        w=wb[f"SIM_R8予算置換_{s}"]
        jun[i]=[rnd(w.cell(row=28,column=c).value) for c in range(6,12)] # R11..R16
    return keijo,jun

for tag,fix,xlsx in [("A_現行基準",False,f"{REPO}/A_現行基準/02_投資財政計画_A_現行基準.xlsx"),
                     ("B_是正基準",True, f"{REPO}/B_是正基準/02_投資財政計画_B_是正基準.xlsx")]:
    keijo,jun=scen(xlsx)
    wb=openpyxl.load_workbook(f"{WD}/upload4/fmt.xlsx")

    # ---- Sheet（可視）：改定率10/20/30％の当期純利益 ----
    ws=wb["Sheet"]; ref=ws["E3"]
    for rate,row in [(0.10,5),(0.20,7),(0.30,9)]:
        d=run(rate,fix)
        for j,col in enumerate("EFGHIJ"):
            c=ws[f"{col}{row}"]
            c.value=rnd(d["r28"][4+j])
            c.number_format=ref.number_format
            c.font=copy.copy(ref.font); c.alignment=copy.copy(ref.alignment)
            c.fill=NOFILL
    ws["B12"]=("※当期純利益は投資財政計画の純損益。改定率はR11以降の給水収益に乗じて算定"
               f"（資産維持費は純損益には算入されない）／{'二重計上を是正した基準' if fix else '現行モデル基準'}")
    ws["B12"].font=copy.copy(ws["B11"].font)

    # ---- Sheet2（非表示・資料P10/P11の元数値）----
    s2=wb["Sheet2"]
    for i in range(4):
        s2.cell(row=3+i,column=4).value=round(keijo[i]/100,3)      # D列 R11経常収支比率（分数格納）
        for j in range(6):
            c=s2.cell(row=3+i,column=9+j)                          # I..N 改定後純利益
            c.value=jun[i][j]
            if any(v<0 for v in jun[i]): c.number_format=NEGFMT     # 負値は本ファイルの△規約へ
    s2["O1"]=f"※{'二重計上を是正した基準' if fix else '現行モデル基準'}（D列R11経常収支比率・I〜N列改定後純利益）"
    s2["O1"].font=copy.copy(s2["B2"].font)

    out=f"{WD}/w4/改定率別純利益_{tag}.xlsx"
    wb.save(out); print("保存:",out)
