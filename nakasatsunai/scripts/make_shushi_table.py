# -*- coding: utf-8 -*-
"""投資・財政計画（収支計画）の表を、シートの印刷範囲からPNGとして生成する。

本文p.31に貼られていたEMFは「下水道現況予測 P2-15% (2)」の印刷範囲（$A$25:$Y$127）を
貼り込んだもの。採用案を協議会案へ変更したため作り直す。あわせて、貼込幅9.56inが
本文幅6.54inを超えてはみ出していた問題も是正する。
"""
import sys, os, re, warnings
warnings.filterwarnings('ignore')
import openpyxl
from PIL import Image, ImageDraw, ImageFont

#  使い方: make_shushi_table.py SRC SHEET OUT [開始行] [終了行] [見出し] [最大高(in)]
SRC   = sys.argv[1]
SHEET = sys.argv[2]
OUT   = sys.argv[3]
R0    = int(sys.argv[4]) if len(sys.argv) > 4 else 25
R1    = int(sys.argv[5]) if len(sys.argv) > 5 else 127
TITLE = sys.argv[6] if len(sys.argv) > 6 else '■投資・財政計画（収支計画）'
HMAX  = float(sys.argv[7]) if len(sys.argv) > 7 else 9.00
os.makedirs(os.path.dirname(OUT) or '.', exist_ok=True)

FONT_PATH = '/usr/share/fonts/opentype/ipafont-gothic/ipagp.ttf'
DPI = 300
W_IN = 6.54
HEADBG, BANDBG, ALTBG = '#7E9BB2', '#DCE5EB', '#F4F7F9'
LINE, GRID, TXT, HEADTX = '#7E9BB2', '#C9D3DA', '#333333', '#FFFFFF'
NEG = '#C0392B'
F = lambda s: ImageFont.truetype(FONT_PATH, max(6, int(s / 72 * DPI)))

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb[SHEET]
VCOLS = list('PQRSTUVWXY')               # 令和7〜16年度
LABC  = list('BCDEFG')                   # ラベル（インデント段階）
ANNC  = list('HIJ')                      # (A)・(C)-(D) などの注記

def fmt(v, nf):
    if v is None or v == '':
        return ''
    if isinstance(v, str):
        return v
    if '%' in nf:
        d = 2 if '0.00%' in nf else 1
        return f'{v * 100:.{d}f}%'
    #  書式から小数桁数を読む。General は整数ならそのまま、小数なら1桁。
    m = re.search(r'0\.(0+)', nf)
    if m:
        d = len(m.group(1))
    elif nf in ('General', '') and abs(v - round(v)) > 1e-9:
        d = 1
    else:
        d = 0
    if v < 0:
        return f'△ {abs(v):,.{d}f}'
    return f'{v:,.{d}f}'

#  縦方向にマージされたB〜G列のセルは「大区分」。項目名から切り離して帯行にする。
GROUPS, GROUPCELL = {}, set()
for m in ws.merged_cells.ranges:
    #  表示する行範囲と重なる縦マージのみを大区分として扱う
    if m.max_col <= 7 and m.max_row > m.min_row and m.min_row <= R1 and m.max_row >= R0:
        head = max(m.min_row, R0)       # 範囲の途中から始まる場合は先頭行に帯を置く
        GROUPS[(m.min_col, head)] = str(ws.cell(m.min_row, m.min_col).value)
        GROUPCELL |= {(m.min_col, r) for r in range(m.min_row, m.max_row + 1)}

rows = []
for r in range(R0, R1 + 1):
    if ws.row_dimensions[r].hidden:
        continue
    for ci in range(2, 8):                       # B=2 … G=7
        if (ci, r) in GROUPS:
            rows.append({'r': r, 'lab': GROUPS[(ci, r)], 'ind': ci - 2,
                         'ann': '', 'vals': [''] * 10, 'kind': 'band'})
    lab, ind = '', 0
    for i, c in enumerate(LABC):
        if (i + 2, r) in GROUPCELL:              # 大区分は項目名に含めない
            continue
        v = ws[f'{c}{r}'].value
        if v not in (None, ''):
            if not lab:
                ind = i
            lab = (lab + ' ' + str(v)).strip()
    ann = ' '.join(str(ws[f'{c}{r}'].value) for c in ANNC
                   if ws[f'{c}{r}'].value not in (None, ''))
    nf = ws[f'P{r}'].number_format
    vals = [fmt(ws[f'{c}{r}'].value, nf) for c in VCOLS]
    #  西暦の行は桁区切りを外す（1900〜2200の整数が並ぶ行を西暦とみなす）
    raw = [ws[f'{c}{r}'].value for c in VCOLS]
    nums = [x for x in raw if isinstance(x, (int, float))]
    if nums and all(isinstance(x, (int, float)) and float(x).is_integer()
                    and 1900 <= x <= 2200 for x in nums):
        vals = [str(int(x)) if isinstance(x, (int, float)) else (str(x) if x else '')
                for x in raw]
    if not lab and not ann and not any(vals):
        continue
    kind = 'head' if r in (25, 26, 66, 67, 99, 100, 109, 110) else 'data'
    lab = re.sub(r'\s{2,}', ' ', lab).replace('㎥', 'm³').replace('㎡', 'm²')
    if re.fullmatch(r'[０-９0-9]+条', lab) and not any(v for v in vals):
        continue        # 「3条」「4条」だけの区切り行は表に出さない
    rows.append({'r': r, 'lab': lab, 'ind': ind,
                 'ann': ann, 'vals': vals, 'kind': kind})

#  行の高さは、ページ内（本文高さ 9.0in 以内）に収まるよう自動で決める
TITLE_H, NOTE_H, PAD = 0.20, 0.30, 0.05
rh = min(0.090, (HMAX - TITLE_H - NOTE_H - PAD * 2) / len(rows))
fs = rh / 0.090 * 5.6
fonts = {'head': F(fs * 0.95), 'body': F(fs), 'band': F(fs), 'note': F(fs * 0.85),
         'title': F(7.5)}

W = int(W_IN * DPI)
pad = int(PAD * DPI)
H = int((TITLE_H + len(rows) * rh + NOTE_H + PAD * 2) * DPI)
img = Image.new('RGB', (W, H), 'white')
d = ImageDraw.Draw(img)

LABW, ANNW = 1.52, 0.42
xs = [pad]
unit = (W - 2 * pad) / (LABW + ANNW + 10)
xs.append(pad + LABW * unit)                     # 項目名／注記 の境界
xs.append(xs[-1] + ANNW * unit)                  # 注記／数値 の境界
for i in range(10):
    xs.append(xs[-1] + unit)

def fit(text, maxw, base_pt):
    """ラベルが枠に収まるまでフォントを小さくする。"""
    pt = base_pt
    while pt > base_pt * 0.48:
        f = F(pt)
        if d.textlength(text, font=f) <= maxw:
            return f
        pt -= 0.25
    return F(pt)

def put(box, text, font, fill, align):
    if not text:
        return
    x0, y0, x1, y1 = box
    l, t, r2, b = d.textbbox((0, 0), text, font=font)
    w, h = r2 - l, b - t
    x = {'c': x0 + (x1 - x0 - w) / 2, 'l': x0 + 0.02 * DPI, 'r': x1 - 0.03 * DPI - w}[align] - l
    d.text((x, y0 + (y1 - y0 - h) / 2 - t), text, font=font, fill=fill)

y = pad
put((xs[0], y, xs[-1], y + TITLE_H * DPI), TITLE, fonts['title'], TXT, 'l')
y += TITLE_H * DPI
top = y
for row in rows:
    h = rh * DPI
    bg = {'head': HEADBG, 'band': BANDBG, 'data': 'white'}[row['kind']]
    if row['kind'] == 'data' and row['r'] % 2 == 0:
        bg = ALTBG
    d.rectangle([xs[0], y, xs[-1], y + h], fill=bg)
    fg = HEADTX if row['kind'] == 'head' else TXT
    fnt = fonts[row['kind'] if row['kind'] in ('head', 'band') else 'body']
    ind = row['ind'] * 0.050 * DPI
    if row['kind'] == 'head':
        put((xs[0], y, xs[2], y + h), row['lab'] or row['ann'], fnt, fg, 'c')
    else:
        lf = fit(row['lab'], xs[1] - xs[0] - ind - 0.05 * DPI, fs)
        put((xs[0] + ind, y, xs[1], y + h), row['lab'], lf, fg, 'l')
        put((xs[1], y, xs[2], y + h), row['ann'], fonts['note'], fg, 'r')
    for i, v in enumerate(row['vals']):
        col = NEG if v.startswith('△') else fg
        put((xs[i + 2], y, xs[i + 3], y + h), v, fnt,
            col, 'c' if row['kind'] == 'head' else 'r')
    for i in range(2, 13):
        d.line([xs[i], y, xs[i], y + h], fill=GRID, width=1)
    d.line([xs[0], y + h, xs[-1], y + h], fill=GRID, width=1)
    y += h
d.rectangle([xs[0], top, xs[-1], y], outline=LINE, width=2)
put((xs[0], y, xs[-1], y + NOTE_H * DPI),
    f'※シミュレーションExcel「{SHEET}」の{R0}〜{R1}行から作成。',
    fonts['note'], '#555555', 'l')
img.save(OUT, dpi=(DPI, DPI))
print(f'{OUT}: {img.width}x{img.height}px → {img.width/DPI:.2f}×{img.height/DPI:.2f}in  '
      f'（{len(rows)}行／行高 {rh:.3f}in／本文 {fs:.1f}pt）')
