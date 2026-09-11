# -*- coding: utf-8 -*-
"""大雪地区広域連合 第10期介護保険事業計画
見える化システム 将来推計の設定資料の受領点検（協議資料）.

令和8年9月11日のご指示
  「見える化システム上、施設・居住系サービス利用者数の施策反映から
    次の施策反映へのボタン押下が出来ない状況です。
    現状のデータ状況踏まえて添付ファイルと内容の確認をお願い致します」

受領した4ファイルの点検と、施設・居住系サービスの画面から先へ進めない件の
切り分けを行う。受託者の点検であり確定値ではない。

━━ 判明したこと ━━

1 4ファイルはいずれも「将来推計」の「推計方法の設定」画面で
  伸びを選ぶための参考シートである（シートの操作方法にその旨の記載がある）。
  4件とも認定率・在宅サービスに関するものであり、
  施設・居住系サービス利用者数の伸びを設定するためのシートは含まれていない。

2 当広域連合（保険者番号01832）では、
  推計誤差が算定されていないサービスが5件ある（実績がないため）。
  うち施設・居住系に属するのは地域密着型特定施設入居者生活介護の1件である。

3 認定者数の推計は、基本推計が2か年とも実績をわずかに下回り、
  包括推計は2か年とも実績を上回る。
  令和6年度 基本▲0.72％／包括＋0.55％、令和7年度 基本▲0.50％／包括＋1.91％。
  令和6年度だけを見れば包括推計の方が1.00に近いが、
  2か年の平均では基本推計0.61％・包括推計1.23％で基本推計が小さく、
  振れも小さい。

4 推計誤差が大きいサービスは、当方が対計画比の補正試算で
  「補正では解消しない外れ」とした6件と重なる。
  独立した資料で同じサービスが外れることが確かめられた。

5 在宅サービス利用率は令和8年度まで出力されている。
  率は月数に中立であるが、令和8年度は4月から7月までの4か月であり、
  当方の季節性の検証（在宅は冬季が夏季より2.48％低い）に照らすと
  夏季に偏った率である。伸びの基礎にそのまま用いると在宅を高めに見込む。

6 在宅サービス利用率のサービス一覧に「登録施設介護支援」が既に置かれている。
  令和8年法律第51号により新設される区分であり、確認事項No.115に関わる。

シート構成
  00_受領資料の概要
  01_推計誤差
  02_認定率と在宅サービス利用率の実績
  03_当方の分析との突合
  04_施設・居住系で進めない件の切り分け
  05_自己点検
  06_確認事項

出力
  output/第10期計画_見える化_将来推計の設定資料の受領点検.xlsx

自己点検で1件でも不適合があると終了コード1で終わる。
"""

import io
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import data_mieru_suikei as MS

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ODIR = "/home/user/repository/output"
OUT = os.path.join(ODIR, "第10期計画_見える化_将来推計の設定資料の受領点検.xlsx")
KIJUNBI = "令和8年9月11日"

FONT = "游ゴシック"
NAVY, HEAD = "1F3864", "4472C4"
IN_Y, OK_G, NG_O, MID_B, GRAY = ("FFF2CC", "E2EFDA", "FCE4D6",
                                 "DEEBF7", "F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
wb.remove(wb.active)

CHECKS = []


def chk(no, naiyo, shiki, kekka, ok):
    CHECKS.append((no, naiyo, shiki, kekka, "適合" if ok else "不適合"))
    return ok


# ============================================================ 体裁
def sheet(name, title, subtitle, widths, freeze="A5"):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=FONT, size=9)
    ws["A2"].fill = PatternFill("solid", fgColor=GRAY)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    n = max(len(widths), 6)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 66
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    return ws


def header(ws, row, cols, height=32):
    for i, hh in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=hh)
        c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEAD)
        c.alignment = Alignment(wrap_text=True, horizontal="center",
                                vertical="center")
        c.border = BORDER
    ws.row_dimensions[row].height = height
    return row + 1


def body(ws, row, vals, fills=None, height=24, align=None, bold=False,
         numfmt=None):
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(name=FONT, size=9, bold=bold)
        c.border = BORDER
        ha = (align or {}).get(i, "left" if isinstance(v, str) else "right")
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal=ha)
        if numfmt and numfmt.get(i) and isinstance(v, float):
            c.number_format = numfmt[i]
        if fills and fills.get(i):
            c.fill = PatternFill("solid", fgColor=fills[i])
    ws.row_dimensions[row].height = height
    return row + 1


def lead(ws, row, text, span=8):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=10.5, bold=True, color=NAVY)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    ws.row_dimensions[row].height = 20
    return row + 1


def note(ws, row, text, span=8, height=None):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=8.5)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    ws.row_dimensions[row].height = height or (14 * (1 + text.count("\n")))
    return row + 1


# ============================================================ 00
JURYO = [
    ("①", "要介護認定率の伸びのパターンの比較",
     "グラフ・データ・リストの3シート。"
     "性別・年齢階級別・要介護度別の認定率（実績値）を令和3年度から令和8年度まで",
     "認定率の伸びの選択",
     "「将来推計」の「推計方法の設定」画面から、認定率の伸びを選択してください"),
    ("②", "在宅サービス利用率の伸びのパターンの比較",
     "グラフ・データ・リストの3シート。"
     "サービス別・要介護度別の在宅サービス利用率（実績値）を"
     "令和3年度から令和8年度まで。22サービス",
     "在宅サービス利用率の伸びの選択",
     "「将来推計」の「推計方法の設定」画面から、"
     "在宅サービス利用率の伸びを選択してください"),
    ("③", "要介護認定者数の推計誤差",
     "全国の保険者について、令和6年度・令和7年度の"
     "第1号被保険者数・要介護認定者数と、基本推計・包括推計の誤差",
     "基本推計と包括推計のいずれを採るかの判断",
     "推計誤差＝当年度の認定者数（推計値）÷当年度の認定者数（実績値）"),
    ("④", "サービス利用者数の推計誤差",
     "全国の保険者について、令和6年度のサービス別の"
     "基本推計・包括推計の誤差。25サービス",
     "サービス別に推計がどれだけ外れるかの確認",
     "誤差＝令和6年度の利用者数（推計値）÷令和6年度の利用者数（実績値）"),
]

ws = sheet("00_受領資料の概要",
           "見える化システム 将来推計の設定資料の受領点検",
           "令和8年9月11日に受領した4ファイルの点検。"
           "4件はいずれも「将来推計」の「推計方法の設定」画面で"
           "伸びを選ぶための参考シートである。"
           "受託者の点検であり確定値ではない。基準日 " + KIJUNBI + "。",
           [4, 30, 42, 24, 40])

r = 4
r = lead(ws, r, "1　受領したファイル", span=5)
r = header(ws, r, ["", "資料名", "収録内容", "用途", "シートの記載"])
for no, nm, naiyo, yoto, kisai in JURYO:
    r = body(ws, r, [no, nm, naiyo, yoto, kisai], height=64)

r = note(ws, r + 1,
         "注1）①②はいずれも「本シートの操作方法」の3番目に、"
         "「将来推計」の「推計方法の設定」画面から伸びを選択するよう"
         "記されています。伸びを決めるための資料です。\n"
         "注2）③④の推計誤差は、前年9月の認定率・利用率を"
         "当年9月の第1号被保険者数に乗じて求めた推計値を、"
         "当年9月の実績値で除したものです。"
         "1.00を上回れば推計が実績を上回り、下回れば推計が実績を下回ります。\n"
         "注3）基本推計は性別・年齢5歳階級別・要介護度別、"
         "包括推計は性別・年齢前期後期別に推計する方法です。", span=5)

r += 1
r = lead(ws, r, "2　4件に含まれていないもの", span=5)
r = header(ws, r, ["", "含まれていないもの", "影響", "", ""])
for a, b in [
    ("施設・居住系サービス利用者数の伸びを設定するための参考シート",
     "今回のお問合せはこの画面のものである。"
     "①②に相当するシート（施設・居住系サービス利用者数の伸びのパターンの比較）が"
     "4件の中にない。別に出力できるかを確認いただきたい"),
    ("地域支援事業の量の見込みに関する資料",
     "確認事項No.112（地域支援事業の量の見込み）に関わる。"
     "本件とは別に整理している"),
    ("令和7年度のサービス利用者数の推計誤差",
     "④は令和6年度のみである。"
     "③（認定者数）は令和6年度・令和7年度の2か年がある"),
]:
    r = body(ws, r, ["", a, b, "", ""], height=44)


# ============================================================ 01
ws = sheet("01_推計誤差",
           "当広域連合の推計誤差（保険者番号01832）",
           "推計誤差は、推計値を実績値で除したものである。"
           "1.00に近いほど推計が実績に合っている。"
           "1.00を上回れば推計が実績を上回り（過大）、"
           "下回れば推計が実績を下回る（過小）。",
           [4, 34, 12, 12, 12, 36])

r = 4
r = lead(ws, r, "1　要介護認定者数の推計誤差", span=6)
r = header(ws, r, ["", "年度", "第1号被保険者数", "要介護認定者数",
                   "基本推計", "包括推計"])
for y, nm in (("R6", "令和6年度"), ("R7", "令和7年度")):
    h, n, kihon, hokatsu = MS.NINTEI_GOSA[y]
    r = body(ws, r, ["", nm, h, n, round(kihon, 4), round(hokatsu, 4)],
             fills={5: OK_G, 6: (IN_Y if abs(hokatsu - 1) >
                                 abs(kihon - 1) else None)},
             numfmt={5: "0.0000", 6: "0.0000"}, height=24)

# 年ごとに見ると令和6年度は包括推計の方が1.00に近い。
# 2か年の平均の絶対誤差で比べる。
_avg_kihon = sum(abs(MS.NINTEI_GOSA[y][2] - 1) for y in ("R6", "R7")) / 2
_avg_hokatsu = sum(abs(MS.NINTEI_GOSA[y][3] - 1) for y in ("R6", "R7")) / 2
_kihon_win = _avg_kihon < _avg_hokatsu
_kihon_muki = all(MS.NINTEI_GOSA[y][2] < 1 for y in ("R6", "R7"))
_hokatsu_muki = all(MS.NINTEI_GOSA[y][3] > 1 for y in ("R6", "R7"))
r = note(ws, r + 1,
         "注1）基本推計は2か年とも実績をわずかに下回り（▲0.72％・▲0.50％）、"
         "包括推計は2か年とも実績を上回ります（＋0.55％・＋1.91％）。"
         "令和6年度だけを見れば包括推計の方が1.00に近いのですが、"
         "2か年の平均の絶対誤差は基本推計0.61％・包括推計1.23％で、"
         "基本推計の方が小さく、振れも小さくなっています。\n"
         "注2）第1号被保険者数・要介護認定者数は各年9月末時点"
         "（介護保険事業状況報告 月報）です。"
         "当方が計画素案で用いている認定者数1,984人は令和8年3月末時点であり、"
         "時点が異なります。\n"
         "注3）要介護認定者数には第2号被保険者の認定者を含みます。"
         "認定率のシート（02シート）の第1号被保険者の認定率とは"
         "分子の範囲が異なります。", span=6)

r += 1
r = lead(ws, r, "2　サービス利用者数の推計誤差（令和6年度）", span=6)
r = header(ws, r, ["", "サービス", "基本推計", "包括推計", "外れ", "内容"])

SHISETSU = set(MS.SHISETSU_KYOJU)
HAZURE_T = 0.10
hazure, jisseki_nashi = [], []
for nm, v in MS.RIYOSHA_GOSA.items():
    kubun = "施設・居住系" if nm in SHISETSU else "在宅・地域密着型"
    if v is None:
        jisseki_nashi.append(nm)
        r = body(ws, r, ["", nm, "―", "―", "実績なし",
                         "当該年度に利用実績がないため推計誤差が算定されていない"
                         "（" + kubun + "）"],
                 fills={5: NG_O}, height=26)
        continue
    a, b = v
    d = abs(a - 1)
    if d >= HAZURE_T:
        hazure.append(nm)
    r = body(ws, r, ["", nm, round(a, 4), round(b, 4),
                     "★" if d >= HAZURE_T else "",
                     ("推計が実績を{:.1f}％".format(abs(a - 1) * 100)
                      + ("上回る" if a > 1 else "下回る"))
                     if d >= HAZURE_T else ""],
             fills={3: (IN_Y if d >= HAZURE_T else None)},
             numfmt={3: "0.0000", 4: "0.0000"}, height=24)

r = note(ws, r + 1,
         "注4）★は基本推計の誤差が±10％以上のもので{}件です。"
         "「実績なし」は{}件で、うち施設・居住系に属するのは"
         "地域密着型特定施設入居者生活介護の1件です。\n"
         "注5）基本推計と包括推計で外れの向きが異なるのは訪問看護だけで、"
         "いずれも誤差1％未満です（基本1.0004・包括0.9967）。"
         "したがって、いずれを採っても外れるサービスは変わりません。\n"
         "注6）推計誤差は令和6年度のものです。"
         "第10期の推計にそのまま当てはまるものではありませんが、"
         "どのサービスで推計が外れやすいかを示します."
         .format(len(hazure), len(jisseki_nashi)), span=6)


# ============================================================ 02
ws = sheet("02_認定率と在宅サービス利用率の実績",
           "認定率と在宅サービス利用率の実績（令和3年度〜令和8年度）",
           "伸びを選ぶための基礎となる実績値。"
           "令和8年度は令和8年4月から7月までの4か月による値である。"
           "率は月数に中立であるが、季節の偏りは残る。",
           [4, 34, 11, 11, 11, 11, 30])

r = 4
r = lead(ws, r, "1　要介護認定率（計・第1号被保険者・全体）", span=7)
r = header(ws, r, ["", "区分", "令和3年度", "令和6年度", "令和7年度",
                   "令和8年度", "内容"])
NR = MS.NINTEIRITSU
r = body(ws, r, ["", "認定率", NR["R3"], NR["R6"], NR["R7"], NR["R8"],
                 "令和3年度から令和8年度で＋{:.2f}ポイント".format(
                     (NR["R8"] - NR["R3"]) * 100)],
         numfmt={3: "0.00%", 4: "0.00%", 5: "0.00%", 6: "0.00%"}, height=24)
r = note(ws, r + 1,
         "注1）令和7年度と令和8年度は同じ20.84％です。伸びが止まっています。\n"
         "注2）当方が計画素案で用いている認定率21.8％（令和7年度）とは"
         "分子・分母の取り方が異なります。"
         "本シートは第1号被保険者の認定率（第2号を含まない）であり、"
         "当方の値は第2号を含む認定者数を第1号被保険者数で除したものです。"
         "どちらを用いるかを揃える必要があります。", span=7)

r += 1
r = lead(ws, r, "2　在宅サービス利用率（全体）", span=7)
r = header(ws, r, ["", "サービス", "令和3年度", "令和6年度", "令和7年度",
                   "令和8年度", "令和8年度÷令和7年度"])
ZR = MS.ZAITAKU_RITSU
for nm, (a, b, c, d) in ZR.items():
    hi = (d / c) if c else None
    f = {}
    if hi is not None and (hi >= 1.10 or hi <= 0.90):
        f[7] = IN_Y
    if a == b == c == d == 0.0:
        f[2] = GRAY
    r = body(ws, r, ["", nm, a, b, c, d,
                     round(hi, 3) if hi is not None else "―"],
             fills=f, numfmt={3: "0.00%", 4: "0.00%", 5: "0.00%",
                              6: "0.00%", 7: "0.000"}, height=22)

_zero = [k for k, v in ZR.items() if all(x == 0.0 for x in v)]
_natsu = [k for k, v in ZR.items() if v[2] and (v[3] / v[2]) >= 1.10]
r = note(ws, r + 1,
         "注3）全年度0.00％のサービスが{}件あります。"
         "区域内に事業所がなく実績もないもののほか、"
         "令和8年法律第51号により新設される「登録施設介護支援」を含みます。\n"
         "注4）令和8年度÷令和7年度が1.10以上のものが{}件あります。"
         "令和8年度は4月から7月までの4か月であり、"
         "当方の季節性の検証では在宅サービスは冬季が夏季より2.48％低いため、"
         "夏季に偏った率になっています。"
         "伸びの基礎にそのまま用いると在宅を高めに見込みます。\n"
         "注5）定期巡回・随時対応型訪問介護看護は全年度0.00％ですが、"
         "給付実績はあります（令和7年度・令和8年度）。"
         "小口のサービスは率では捉えられません。"
         .format(len(_zero), len(_natsu)), span=7)


# ============================================================ 03
OUR_HAZURE = {
    "居宅療養管理指導": "居宅療養管理指導",
    "地域密着型通所介護": "地域密着通所介護",
    "短期入所生活介護": "短期入所生活介護",
    "短期入所療養介護": "短期入所療養介護（介護老人保健施設）",
    "訪問看護": "訪問看護",
    "訪問リハビリテーション": "訪問リハビリテーション",
}
_ours = set(OUR_HAZURE.values())
_both = sorted(_ours & set(hazure))
_only_ours = sorted(_ours - set(hazure))
_only_gosa = sorted(set(hazure) - _ours)

ws = sheet("03_当方の分析との突合",
           "当方の分析との突合",
           "受領した推計誤差と、当方が対計画比の補正試算で示した"
           "「補正では解消しない外れ6件」を突き合わせる。"
           "出所の異なる資料で同じサービスが外れるかどうかを見る。",
           [4, 32, 20, 20, 44])

r = 4
r = lead(ws, r, "1　外れるサービスの突合", span=5)
r = header(ws, r, ["", "サービス", "見える化の推計誤差",
                   "当方の対計画比の補正", "内容"])
for nm in sorted(_ours | set(hazure)):
    g = MS.RIYOSHA_GOSA.get(nm)
    in_g = nm in hazure
    in_o = nm in _ours
    r = body(ws, r, ["", nm,
                     ("誤差{:.3f}".format(g[0]) if g else "―")
                     + ("　★" if in_g else ""),
                     "外れ6件に含む" if in_o else "―",
                     "両方で外れる" if (in_g and in_o) else
                     ("当方のみ" if in_o else "推計誤差のみ")],
             fills={5: (OK_G if (in_g and in_o) else IN_Y)}, height=24)
r = body(ws, r, ["", "計", "★{}件".format(len(hazure)),
                 "6件", "一致{}件・当方のみ{}件・推計誤差のみ{}件".format(
                     len(_both), len(_only_ours), len(_only_gosa))],
         bold=True, height=22)
r = note(ws, r + 1,
         "注1）出所も方法も異なる2つの資料で{}件が一致しました。"
         "当方の外れの指摘は、見える化システムの推計誤差からも裏づけられます。\n"
         "注2）一致しないものもあります。"
         "訪問看護は当方では外れますが推計誤差は1.000です。"
         "訪問入浴介護と介護医療院は推計誤差では外れますが"
         "当方の6件には入っていません。"
         "当方の対計画比は計画値との比、推計誤差は前年からの推計との比であり、"
         "測っているものが異なります。".format(len(_both)), span=5)

r += 1
r = lead(ws, r, "2　当方の整理に加えるべきこと", span=5)
r = header(ws, r, ["", "事項", "受領資料から分かること", "当方の現在の整理",
                   "どうするか"])
for a, b, c, d in [
    ("推計方式は基本推計",
     "基本推計は2か年とも実績をわずかに下回り（▲0.72％・▲0.50％）、"
     "包括推計は2か年とも上回る（＋0.55％・＋1.91％）。"
     "2か年の平均の絶対誤差は基本0.61％・包括1.23％",
     "採用パターンP3は年齢階級別（65〜74・75〜84・85歳以上）の人口に"
     "階級別認定率を乗じている",
     "当方の方法は基本推計に近い。見える化で推計方式を選ぶ場面では"
     "基本推計を選ぶ"),
    ("令和8年度の率は夏季に偏る",
     "在宅サービス利用率は令和8年度まで出力されるが、4月から7月までの4か月である",
     "基準年度は令和7年度としている（数量は令和7年度のみ）",
     "率であっても令和8年度は季節の偏りを含む。"
     "伸びの基礎に令和8年度を用いない。"
     "当方の採用パターンP3は令和7年度基準であり、この点は変わらない"),
    ("登録施設介護支援の区分が既にある",
     "在宅サービス利用率のサービス一覧に「登録施設介護支援」が置かれている"
     "（全年度0.00％）",
     "確認事項No.115で「令和9年4月に新設される区分を見込量に立てるか」を"
     "お諮りしている",
     "見える化システムは既に区分を持っている。"
     "区分を立てない場合でも、システム上は0を入力することになる"),
    ("夜間対応型訪問介護の区分がなお残る",
     "サービス一覧に夜間対応型訪問介護が置かれている（全年度0.00％）",
     "令和8年法律第51号により定期巡回・随時対応型訪問介護看護へ統合される"
     "（計画素案 第1章第9節2）",
     "統合の時期とシステムの区分の扱いが一致するかを確認する"),
    ("認定率の定義が当方と異なる",
     "第1号被保険者の認定率は令和7年度20.84％",
     "計画素案は認定率21.8％（令和7年度）としている",
     "分子に第2号被保険者の認定者を含むかどうかが異なる。"
     "計画に掲げる認定率の定義を1つに揃える"),
]:
    r = body(ws, r, ["", a, b, c, d], height=64)


# ============================================================ 04
ws = sheet("04_施設・居住系で進めない件の切り分け",
           "施設・居住系サービス利用者数の画面から先へ進めない件の切り分け",
           "当方は見える化システムに接続できないため、"
           "画面の状態を確認できない。"
           "受領資料から分かることを根拠に、確かめる順序を示す。"
           "原因を特定したものではない。",
           [4, 30, 40, 34, 14])

r = 4
r = header(ws, r, ["", "確かめること", "根拠（受領資料から分かること）",
                   "確かめ方", "見込み"])
KIRIWAKE = [
    ("実績のないサービスの欄が空でないか",
     "当広域連合では地域密着型特定施設入居者生活介護の推計誤差が"
     "算定されていない（実績がないため）。"
     "施設・居住系7区分のうちこの1件だけが「―」である",
     "施設・居住系の画面で、地域密着型特定施設入居者生活介護の欄が"
     "空欄になっていないかを見る。空欄であれば0を入力する",
     "高い"),
    ("介護療養型医療施設の欄が残っていないか",
     "受領した推計誤差の一覧に介護療養型医療施設はない"
     "（令和6年3月末で廃止済み）。"
     "画面に区分が残っている場合は入力を求められる可能性がある",
     "画面に当該区分があれば0を入力する",
     "中"),
    ("前の段階の設定が確定しているか",
     "受領した4件はいずれも「推計方法の設定」画面で"
     "認定率の伸び・在宅サービス利用率の伸びを選ぶための資料である。"
     "この2つの設定を求められている状況であることを示す",
     "認定率の伸びと在宅サービス利用率の伸びが"
     "選択済み・保存済みかを確認する。"
     "未選択のまま先の画面に入っていると戻される場合がある",
     "高い"),
    ("定員との整合の検査に掛かっていないか",
     "施設・居住系は必要利用定員総数と結びつく。"
     "当区域は認知症対応型共同生活介護の定員が未確定である"
     "（99人＋［要確認］。確認事項No.88）",
     "定員の入力欄があるか、"
     "入力した利用者数が定員を超えていないかを見る",
     "中"),
    ("入力値の形式が誤っていないか",
     "―",
     "全角数字・小数点・空白・負の値・カンマの混入を見る。"
     "特に他の表から貼り付けた場合に起きやすい",
     "中"),
    ("画面側の問題でないか",
     "―",
     "別のブラウザで試す。"
     "セッションが切れていないか（一定時間放置していないか）を見る。"
     "未入力の欄に自動で移動しない場合は、"
     "画面を上から順に空欄がないか目視する",
     "低い"),
    ("システム側の不具合でないか",
     "―",
     "上記で解消しない場合は、"
     "見える化システムのヘルプデスクへ照会する。"
     "保険者番号01832・該当画面名・操作手順・"
     "入力済みの値を添えると早い",
     "―"),
]
for i, (a, b, c, d) in enumerate(KIRIWAKE, start=1):
    r = body(ws, r, [i, a, b, c, d],
             fills={5: (IN_Y if d == "高い" else None)}, height=70)

r = note(ws, r + 1,
         "注1）当方は mieruka.mhlw.go.jp に接続できないため、"
         "画面の挙動を確かめられません。"
         "上記は受領資料から言えることにとどまります。\n"
         "注2）1と3の見込みを「高い」としたのは、"
         "当区域に実績のないサービスが施設・居住系に1件あること、"
         "及び今回お送りいただいた4件が"
         "いずれも前段の設定（認定率・在宅サービス利用率の伸び）の"
         "資料であることによります。\n"
         "注3）画面の状態が分かる資料（画面の写し、"
         "エラーの表示、入力済みの値）をいただければ、"
         "さらに絞り込めます。", span=5)

r += 1
r = lead(ws, r, "参考　施設・居住系7区分の状況", span=5)
r = header(ws, r, ["", "サービス", "推計誤差（基本推計）", "定員", "内容"])
TEIIN = {
    "介護老人福祉施設": ("160人", "北海道 特別養護老人ホーム名簿"),
    "介護老人保健施設": ("［要確認］", "区域内の施設の定員を未受領"),
    "介護医療院": ("［要確認］", "同上"),
    "特定施設入居者生活介護": ("156人", "北海道の指定事業所一覧。"
                                "混合型か介護専用型かは未確認（確認事項No.113）"),
    "認知症対応型共同生活介護": ("99人＋［要確認］",
                                "くるみの郷（東川町）の定員が未確定"
                                "（確認事項No.88）"),
    "地域密着型特定施設入居者生活介護": ("―", "区域内に事業所がない"),
    "地域密着型介護老人福祉施設入所者生活介護": ("62人",
                                                "北海道 特別養護老人ホーム名簿"),
}
for nm in MS.SHISETSU_KYOJU:
    g = MS.RIYOSHA_GOSA.get(nm)
    t, src = TEIIN[nm]
    r = body(ws, r, ["", nm,
                     round(g[0], 4) if g else "―（実績なし）",
                     t, src],
             fills={3: (NG_O if g is None else None),
                    4: (IN_Y if "［要確認］" in t else None)},
             numfmt={3: "0.0000"}, height=32)


# ============================================================ 05
ws = sheet("05_自己点検",
           "自己点検（エラーチェック）",
           "受領資料の収録値と、本表の記述の整合を点検する。",
           [4, 40, 26, 34, 10, 10])

chk(1, "受領したファイルの件数が4件であること",
    "資料一覧", "{}件".format(len(JURYO)), len(JURYO) == 4)

chk(2, "サービス利用者数の推計誤差が25サービス収録されていること",
    "④の収録数", "{}サービス".format(len(MS.RIYOSHA_GOSA)),
    len(MS.RIYOSHA_GOSA) == 25)

chk(3, "施設・居住系7区分がすべて推計誤差の一覧にあること",
    "7区分",
    "{}区分が一覧にある".format(
        sum(1 for k in MS.SHISETSU_KYOJU if k in MS.RIYOSHA_GOSA)),
    all(k in MS.RIYOSHA_GOSA for k in MS.SHISETSU_KYOJU))

_nashi_shisetsu = [k for k in MS.SHISETSU_KYOJU
                   if MS.RIYOSHA_GOSA.get(k) is None]
chk(4, "施設・居住系のうち実績がないものが1件であること",
    "実績なし全体{}件".format(len(jisseki_nashi)),
    "施設・居住系は{}件（{}）".format(
        len(_nashi_shisetsu), "・".join(_nashi_shisetsu)),
    len(_nashi_shisetsu) == 1)

chk(5, "認定者数の推計は2か年の平均で基本推計が実績に近いこと",
    "平均|誤差－1| 基本{:.4f}／包括{:.4f}".format(_avg_kihon, _avg_hokatsu),
    "R6 基本{:+.2%}／包括{:+.2%}、R7 基本{:+.2%}／包括{:+.2%}".format(
        MS.NINTEI_GOSA["R6"][2] - 1, MS.NINTEI_GOSA["R6"][3] - 1,
        MS.NINTEI_GOSA["R7"][2] - 1, MS.NINTEI_GOSA["R7"][3] - 1),
    _kihon_win)

chk(14, "基本推計は2か年とも過小、包括推計は2か年とも過大であること",
    "誤差の符号",
    "基本は{}、包括は{}".format(
        "2か年とも1.00未満" if _kihon_muki else "符号が混在",
        "2か年とも1.00超" if _hokatsu_muki else "符号が混在"),
    _kihon_muki and _hokatsu_muki)

# 向きが異なるものがあっても、誤差が1.00に極めて近ければ実質の差はない。
_muki = [nm for nm, v in MS.RIYOSHA_GOSA.items()
         if v and ((v[0] - 1) * (v[1] - 1) < 0)]
_muki_big = [nm for nm in _muki
             if max(abs(MS.RIYOSHA_GOSA[nm][0] - 1),
                    abs(MS.RIYOSHA_GOSA[nm][1] - 1)) >= 0.01]
chk(6, "基本推計と包括推計で外れの向きが異なるものが、"
       "誤差1％未満のものに限られること",
    "(基本－1)×(包括－1)＜0 のもの {}件".format(len(_muki)),
    ("向きが異なるのは{}のみで、いずれも誤差1％未満".format("・".join(_muki))
     if _muki and not _muki_big else
     ("全サービスで同じ" if not _muki else
      "誤差1％以上で向きが異なる " + "・".join(_muki_big))),
    not _muki_big)

chk(7, "推計誤差の外れ（±10％以上）と当方の外れ6件が一部一致すること",
    "推計誤差★{}件／当方6件".format(len(hazure)),
    "一致{}件・当方のみ{}件・推計誤差のみ{}件".format(
        len(_both), len(_only_ours), len(_only_gosa)),
    len(_both) >= 3 and len(_only_ours) + len(_only_gosa) > 0)

chk(8, "当方の外れ6件がすべて推計誤差の一覧にあること",
    "6件",
    "{}件が一覧にある".format(sum(1 for k in _ours if k in MS.RIYOSHA_GOSA)),
    all(k in MS.RIYOSHA_GOSA for k in _ours))

chk(9, "認定率が令和3年度から令和8年度まで6か年収録されていること",
    "①の収録年度", "{}か年".format(len(MS.NINTEIRITSU)),
    len(MS.NINTEIRITSU) == 6)

chk(10, "在宅サービス利用率が22サービス収録されていること",
    "②の収録数", "{}サービス".format(len(MS.ZAITAKU_RITSU)),
    len(MS.ZAITAKU_RITSU) == 22)

chk(11, "在宅サービス利用率に令和8年法律第51号で新設される区分があること",
    "登録施設介護支援",
    "一覧にある（全年度0.00％）" if "登録施設介護支援" in MS.ZAITAKU_RITSU
    else "一覧にない",
    "登録施設介護支援" in MS.ZAITAKU_RITSU)

chk(12, "令和8年度の在宅サービス利用率が令和7年度を上回るものが多いこと"
        "（夏季に偏ることの傍証）",
    "令和8年度÷令和7年度",
    "1.00超{}件・1.00以下{}件".format(
        sum(1 for v in MS.ZAITAKU_RITSU.values() if v[2] and v[3] / v[2] > 1),
        sum(1 for v in MS.ZAITAKU_RITSU.values()
            if v[2] and v[3] / v[2] <= 1)),
    sum(1 for v in MS.ZAITAKU_RITSU.values() if v[2] and v[3] / v[2] > 1) >
    sum(1 for v in MS.ZAITAKU_RITSU.values() if v[2] and v[3] / v[2] <= 1))

chk(13, "切り分けの各項目に確かめ方が示されていること",
    "{}件".format(len(KIRIWAKE)),
    "確かめ方のないもの{}件".format(
        sum(1 for x in KIRIWAKE if not x[2].strip())),
    all(x[2].strip() for x in KIRIWAKE))

r = 4
r = header(ws, r, ["No", "点検の内容", "式・対象", "結果", "判定", ""])
for no, naiyo, shiki, kekka, han in CHECKS:
    r = body(ws, r, [no, naiyo, shiki, kekka, han, ""],
             fills={5: OK_G if han == "適合" else NG_O}, height=32)
NG = sum(1 for c in CHECKS if c[4] == "不適合")
r = body(ws, r, ["", "計", "{}件".format(len(CHECKS)),
                 "適合{}件・不適合{}件".format(len(CHECKS) - NG, NG),
                 "適合" if NG == 0 else "不適合", ""],
         fills={5: OK_G if NG == 0 else NG_O}, bold=True)
r = note(ws, r + 1,
         "注1）収録値は受領したファイルから読み取った当広域連合の行のみです。"
         "全国の他の保険者の値は収録していません。\n"
         "注2）1件でも不適合があると、本表を作るスクリプトは"
         "終了コード1で終わります。", span=6)


# ============================================================ 06
ws = sheet("06_確認事項",
           "確認事項",
           "本表により新たに起票した確認事項を示す。"
           "業務工程管理表 03_確認事項一覧に同じ番号で登載する。",
           [7, 30, 46, 20, 12, 12])

r = 4
r = header(ws, r, ["No.", "確認事項", "内容", "止めている成果物",
                   "確認先", "回答期限"])
KAKUNIN = [
    ("No.123", "施設・居住系の画面から進めない件の確認",
     "当方は見える化システムに接続できないため画面の状態を確認できない。"
     "受領資料から、次の2点を先に確かめていただきたい。"
     "①施設・居住系7区分のうち地域密着型特定施設入居者生活介護は"
     "区域内に事業所がなく、推計誤差も算定されていない。"
     "当該欄が空欄になっていないか（空欄であれば0を入力する）。"
     "②今回お送りいただいた4件はいずれも"
     "「推計方法の設定」画面で認定率の伸び・在宅サービス利用率の伸びを"
     "選ぶための資料である。この2つの設定が保存済みか。"
     "③あわせて、画面の写し・エラーの表示・入力済みの値を"
     "ご提供いただければ、さらに絞り込める。"
     "④上記で解消しない場合は見える化システムのヘルプデスクへの照会をお願いしたい。",
     "将来推計 第2段階・第3段階\nサービス見込量 第1次概算",
     "発注者", "R8.9"),
    ("No.124", "施設・居住系サービス利用者数の伸びの参考シートの有無",
     "認定率（①）と在宅サービス利用率（②）については"
     "伸びのパターンを比較する参考シートが出力されているが、"
     "施設・居住系サービス利用者数について同じ形のシートが"
     "受領資料に含まれていない。"
     "見える化システムから出力できるものであればご提供いただきたい。"
     "施設・居住系は定員に縛られ、在宅とは伸びの置き方が異なるため、"
     "実績の推移を確かめたうえで伸びを選ぶ必要がある。",
     "将来推計 第2段階\n計画素案 第6章第2節・第4節",
     "発注者", "R8.9"),
    ("No.125", "計画に掲げる認定率の定義を揃えること",
     "受領した認定率のシートでは、第1号被保険者の認定率は"
     "令和7年度20.84％である。"
     "一方、計画素案は認定率21.8％（令和7年度）としている。"
     "差は分子に第2号被保険者の認定者を含むかどうかによると考えられる。"
     "①計画に掲げる認定率を、第1号被保険者のみとするか、"
     "第2号被保険者を含むかを決めていただきたい。"
     "②見える化システムで推計する場合は同システムの定義に従うことになるため、"
     "計画本文の値と食い違わないよう揃える必要がある。"
     "③時点も異なる（受領資料は各年9月末、当方は年度末）。",
     "計画素案 第2章第1節3・第6章第1節\n将来推計 第1段階・第2段階",
     "発注者", "R8.10"),
    ("No.126", "推計方式は基本推計でよいか",
     "受領した認定者数の推計誤差では、"
     "令和6年度は基本推計▲0.72％・包括推計＋0.55％、"
     "令和7年度は基本推計▲0.50％・包括推計＋1.91％である。"
     "令和6年度だけを見れば包括推計の方が1.00に近いが、"
     "基本推計は2か年とも実績をわずかに下回る側で振れが小さく"
     "（2か年平均0.61％）、"
     "包括推計は2か年とも上回る側で振れが大きい（同1.23％）。"
     "当方の採用パターンP3は年齢階級別の人口に階級別認定率を乗じる方法であり、"
     "基本推計に近い。"
     "見える化システムで推計方式を選ぶ場面では基本推計を選ぶ整理でよいか。",
     "将来推計 第2段階\nサービス見込量 第1次概算",
     "発注者", "R8.9"),
]
for no, ken, naiyo, tome, saki, kigen in KAKUNIN:
    r = body(ws, r, [no, ken, naiyo, tome, saki, kigen], height=140)
r = note(ws, r + 1,
         "注1）本表により新たに起票した確認事項は{}件（No.123〜No.{}）です。\n"
         "注2）確認事項は業務工程管理表 03_確認事項一覧で一元管理します。"
         "計画素案の本文には注記しません。"
         .format(len(KAKUNIN), 122 + len(KAKUNIN)), span=6)


# ============================================================ 出力
os.makedirs(ODIR, exist_ok=True)
wb.save(OUT)

print("出力：" + OUT)
print("受領資料：{}件".format(len(JURYO)))
print("サービス利用者数の推計誤差：{}サービス（実績なし{}件・外れ★{}件）"
      .format(len(MS.RIYOSHA_GOSA), len(jisseki_nashi), len(hazure)))
print("　うち施設・居住系で実績なし：{}".format("・".join(_nashi_shisetsu)))
print("当方の外れ6件との突合：一致{}件・当方のみ{}件・推計誤差のみ{}件"
      .format(len(_both), len(_only_ours), len(_only_gosa)))
print("認定者数の推計：平均|誤差| 基本{:.2%}／包括{:.2%}（{}）"
      .format(_avg_kihon, _avg_hokatsu,
              "基本推計が小さい" if _kihon_win else "包括推計が小さい"))
print("新たに起票した確認事項：{}件（No.123〜No.{}）"
      .format(len(KAKUNIN), 122 + len(KAKUNIN)))
print("自己点検：{}件（適合{}件・不適合{}件）"
      .format(len(CHECKS), len(CHECKS) - NG, NG))

if NG:
    sys.exit(1)
