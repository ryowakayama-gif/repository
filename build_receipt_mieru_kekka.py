# -*- coding: utf-8 -*-
"""大雪地区広域連合 第10期介護保険事業計画
見える化システムの推計結果（推計パターン 9028（１））の受領点検と修正箇所.

令和8年9月30日のご依頼
  「添付ファイル確認の上、修正必要な箇所をご教示ください」

受領したもの
  ① 見える化システムの画面の写し（34点。認定者数・施設居住系・在宅の施策反映、
     訪問型通所型サービス、地域支援事業費、所得段階別、保険料額の算定）
  ② 第10期介護保険事業（支援）計画策定に向けたワークシート（6シート）

受領ファイルはいずれもファイルの属性に作成者名があるためリポジトリに格納しない。
集計値のみを `data_mieru_kekka.py` に収めた。個票は含まれない。

━━ 点検の骨格 ━━

推計結果の値そのものを直すのではなく、**どの設定・どの入力欄を直すと
値が変わるか**を、当方の算定（サービス見込量の算定）と対照して示す。

当方の算定との主な差
  総給付費（第10期3か年）　見える化 9,685,143千円／当方 9,031,900千円
  保険料基準額（月額）　　　見える化と当方の対照は本文の表による
  （固定値を書かない。算定を改めると追随する）
  第9期の再現　　　　　　　見える化 6,427.89円／当方 6,427.92円（一致）

シート構成
  00_この点検について
  01_修正が必要な箇所
  02_設定の対照
  03_令和8年度の基準値の点検
  04_所得段階別の入力
  05_保険料の対照
  06_サービス別の対照
  07_人口の設定の入力値
  08_弾力化の入力値（基準所得金額と割合）
  09_再出力の点検（人口の設定を案Cに改めた後）
  10_ワーニングチェック
  11_自己点検

出力
  output/第10期計画_見える化_推計結果の受領点検.xlsx

自己点検で1件でも不適合があると終了コード1で終わる。
"""

import io
import os
import runpy
import sys
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import data_mieru_danryoku as GD
import mieru_anc as AN
import data_mieru_jinko as J
import data_mieru_kekka as K
import repo_paths as RP

if isinstance(sys.stdout, io.TextIOWrapper):
    sys.stdout.reconfigure(encoding="utf-8")

OUT = os.path.join(RP.OUTPUT, "第10期計画_見える化_推計結果の受領点検.xlsx")
KIJUNBI = "令和8年9月30日"

FONT = "游ゴシック"
NAVY, HEAD = "1F3864", "4472C4"
IN_Y, OK_G, NG_O, MID_B, GRAY = ("FFF2CC", "E2EFDA", "FCE4D6",
                                 "DEEBF7", "F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
wb.remove(wb.active)
CHECKS = []


def chk(no, naiyo, shiki, kekka, ok):
    CHECKS.append((no, naiyo, shiki, kekka, "適合" if ok else "不適合"))
    return ok


def _plain(v):
    """xlsx は Markdown を解釈しないため、書き出しの時点で ** を落とす。"""
    return v.replace("**", "") if isinstance(v, str) else v


# ============================================================ 体裁
def sheet(name, title, subtitle, widths, freeze="A5"):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A2"] = _plain(subtitle)
    ws["A2"].font = Font(name=FONT, size=9)
    ws["A2"].fill = PatternFill("solid", fgColor=GRAY)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    n = max(len(widths), 6)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 72
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def header(ws, row, cols, height=28):
    for i, hh in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=_plain(hh))
        c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEAD)
        c.alignment = Alignment(wrap_text=True, horizontal="center",
                                vertical="center")
        c.border = BORDER
    ws.row_dimensions[row].height = height
    return row + 1


def body(ws, row, vals, fills=None, height=20, align=None, bold=False,
         fmt=None):
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=row, column=i, value=_plain(v))
        c.font = Font(name=FONT, size=9, bold=bold)
        c.border = BORDER
        ha = (align or {}).get(i, "left" if isinstance(v, str) else "right")
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal=ha)
        if fmt and fmt.get(i) and isinstance(v, (int, float)):
            c.number_format = fmt[i]
        if fills and fills.get(i):
            c.fill = PatternFill("solid", fgColor=fills[i])
    ws.row_dimensions[row].height = height
    return row + 1


def lead(ws, row, text, span=10):
    c = ws.cell(row=row, column=1, value=_plain(text))
    c.font = Font(name=FONT, size=10.5, bold=True, color=NAVY)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    ws.row_dimensions[row].height = 20
    return row + 1


def note(ws, row, text, span=10, height=None):
    c = ws.cell(row=row, column=1, value=_plain(text))
    c.font = Font(name=FONT, size=8.5)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    ws.row_dimensions[row].height = height or (14 * (1 + text.count("\n")))
    return row + 1


# ============================================================ 当方の算定
_S = runpy.run_path(os.path.join(RP.ROOT, "build_mikomiryo_santei.py"))
G_DO = _S["G_DO"]
KYUFU3 = _S["KYUFU3"]
KYUFU_Y = _S["KYUFU_Y"]
Y3 = _S["Y3"]
HIHO3_TO = _S["HIHO3"]

TOU_SOGAKU3 = sum(KYUFU3)
TOU_GETSU = G_DO["月額"]
TOU_A = G_DO["A"]
TOU_B = G_DO["B"]
TOU_J = G_DO["J"]
TOU_HOSEI3 = G_DO["③"]

MI_SOGAKU3 = sum(K.SOGAKU["合計"][y] for y in ("R9", "R10", "R11")) * 1000
MI_A = K.SHUNO["標準給付費見込額"]["計"]
MI_B = K.SHUNO["地域支援事業費"]["計"]
MI_J = K.SHUNO["保険料収納必要額"]["計"]
MI_GETSU = K.HOKENRYO["5_保険料推計"]["第10期"]
MI_GETSU1 = K.HOKENRYO["1_推計値サマリ"]["第10期"]

# 標準13段階の割合（画面。当連合の条例の乗率とは違う）
GAMEN13_JOR = [0.4550, 0.6850, 0.6900, 0.9000, 1.0000, 1.2000, 1.3000,
               1.5000, 1.7000, 1.9000, 2.1000, 2.3000, 2.4000]
DAN_Y = ("R9", "R10", "R11")
HOSEI13 = sum(
    sum(K.DANKAI["第%d段階" % (i + 1)][y] * GAMEN13_JOR[i] for i in range(13))
    for y in DAN_Y)
KEISU_TO = TOU_HOSEI3 / HIHO3_TO          # 条例16段階の係数
HOSEI16 = K.DANKAI["計"]["計"] * KEISU_TO
GETSU16 = MI_J / 0.99 / HOSEI16 / 12
DANRYOKU_SA = GETSU16 - MI_GETSU

# 地域支援事業費を入れた場合の効き
WARI_CHOSEI = (K.SHUNO["調整交付金見込額"]["計"]
               / K.SHUNO["標準給付費見込額"]["計"])
DJ = TOU_B * (0.23 + 0.05 - WARI_CHOSEI)
CHIIKI_SA = DJ / 0.99 / HOSEI13 / 12


# ============================================================ 00
ws = sheet("00_この点検について",
           "見える化システムの推計結果の受領点検と修正箇所",
           "%s受領。推計パターン %s（保険者番号 %s）。"
           "画面の写し34点とワークシート6シートを当方の算定と突き合わせ、"
           "直す必要のある箇所を優先順に整理したものです。"
           "受領ファイルそのものは保管せず、集計値のみを用いています。"
           % (K.META["受領日"], K.META["推計パターン名"],
              K.META["保険者番号"]),
           [4, 34, 26, 26, 30])
r = 4
r = lead(ws, r, "1　要点", span=5)
r = header(ws, r, ["#", "要点", "見える化", "当方の算定", "扱い"])
YOTEN = [
    ("保険料基準額（月額）が同じ出力の中で2通りある",
     "%s円（1シート）／%s円（5シート）"
     % ("{:,.0f}".format(MI_GETSU1), "{:,.0f}".format(MI_GETSU)),
     "{:,.0f}円".format(TOU_GETSU),
     "5シートが所得段階別の入力を反映した値。1シートは反映前の値とみられる"),
    ("地域支援事業費が全ての欄で0",
     "0円", "{:,}円".format(TOU_B),
     "入力を要する。入れると月額が約{:,.0f}円上がる".format(CHIIKI_SA)),
    ("所得段階が標準13段階のまま（弾力化していない）",
     "標準13段階", "条例16段階",
     "弾力化の画面に入力を要する。月額 約{:+,.0f}円".format(DANRYOKU_SA)),
    ("令和8年度の実績見込み値を編集していない",
     "月報からの計算値のまま", "年報 令和7年度の月平均に差し替える案あり",
     "令和8年度は年度の実績ではない。18区分が令和7年度と±10％超ずれる"),
    ("人口の設定が社人研のまま",
     "第1号 第10期3か年 {:,}人".format(K.DANKAI["計"]["R9"]
                                        + K.DANKAI["計"]["R10"]
                                        + K.DANKAI["計"]["R11"]),
     "案C {:,.0f}人".format(HIHO3_TO),
     "確認事項No.127。給付費の側も人口の設定で動く"),
    ("所得段階別の人数が第10期と令和12年度以降で別の人口による",
     "第10期＝画面の計／令和12年度以降＝案Cの計",
     "同じ系列でそろえる",
     "加入割合の合計が令和12年度以降で1.0を超える（最大1.0149）"),
    ("1人1月あたり給付費の実績値に令和8年度を選んでいる",
     "令和8年度", "令和7年度",
     "確認事項No.137。ご判断を要する"),
    ("介護サービス利用者数の自然体推計手法が包括推計",
     "要介護度を包括して推計する", "要介護度別",
     "確認事項No.126。当方の検証では要介護度別（基本推計）の方が誤差が小さい"),
    ("施策反映の記入欄が空欄",
     "「出力後に入力」のまま", "―",
     "計画に載せる前に記入を要する"),
]
for i, (a, b, c, d) in enumerate(YOTEN, start=1):
    r = body(ws, r, [i, a, b, c, d], height=34,
             fills={1: MID_B})
r = note(ws, r,
         "注1）本表は受領した推計結果を当方の算定と突き合わせたものであり、"
         "どちらが正しいというものではありません。"
         "設定と入力を直したうえで、なお残る差を協議の材料とします。\n"
         "注2）第9期の保険料基準額は見える化が%s円、当方の再現が%s円で"
         "一致しており、算定式そのものは同じものです。\n"
         "注3）当方の算定は必要資料が未受領であることによる暫定値です。"
         % ("{:,.2f}".format(K.HOKENRYO["第9期"]), "6,427.92"),
         span=5, height=60)

# ============================================================ 01
ws = sheet("01_修正が必要な箇所",
           "修正が必要な箇所（優先順）",
           "画面のどこを、どう直すかを示します。"
           "「効き」は第10期の保険料基準額（月額）への影響の見当です。",
           [4, 8, 26, 30, 34, 18])
r = 4
r = header(ws, r, ["#", "優先", "画面・シート", "現在", "直す内容", "効き"])
SHUSEI = [
    (1, "A", "保険料額の算定 → 保険料収納必要額の算出に必要な数値",
     "地域支援事業費が全ての年度で0",
     "総合事業費・包括的支援事業（センター運営）・任意事業費・"
     "包括的支援事業（社会保障充実分）を入れる。"
     "当方の算定では第10期3か年で{:,}円（令和6年度決算の据え置き）".format(TOU_B),
     "月額 約{:+,.0f}円".format(CHIIKI_SA)),
    (2, "A", "保険料額の算定 → 所得段階別第1号被保険者数②（弾力化）",
     "全欄が0（弾力化していない）",
     "当連合の所得段階は第9期の条例により16段階である。"
     "②の画面に16段階の人数と基準額に対する割合を入れる",
     "月額 約{:+,.0f}円".format(DANRYOKU_SA)),
    (3, "A", "実績及び推計方法の設定 → 令和8年度の実績見込み値",
     "編集していない（月報からの計算値のまま）",
     "令和8年度は年度の実績ではない。"
     "認定者数・施設居住系の利用者数・在宅の利用者数・"
     "在宅の利用回（日）数を年報（令和7年度）の月平均に差し替える",
     "総給付費の水準が変わる"),
    (4, "A", "所得段階別第1号被保険者数①（標準段階区分）",
     "第10期は画面の計に合わせ、令和12年度以降は案Cの計になっている",
     "人口の設定を案Cに改めたうえで全8年度を案Cでそろえるか、"
     "社人研のままとして全8年度を画面の計に合わせるか、どちらかにそろえる",
     "令和12年度以降で約1.0〜1.5％"),
    (5, "B", "実績及び推計方法の設定 → 人口の設定",
     "社人研推計のまま",
     "総人口を案C（地方創生総合戦略＋住民基本台帳の実績趨勢）に改める。"
     "確認事項No.127",
     "給付費・保険料の双方"),
    (6, "B", "実績及び推計方法の設定 → 1人1月あたり給付費の実績値",
     "施設・居住系／在宅とも令和8年度",
     "令和7年度に改めるかご判断をいただく。確認事項No.137。"
     "令和8年度は月報（4〜7月）からの計算値である",
     "総給付費の水準"),
    (7, "B", "実績及び推計方法の設定 → 自然体推計手法",
     "介護サービス利用者数が「要介護度を包括して推計する」",
     "「要介護度別」に改める。"
     "当方の検証では2か年平均の絶対誤差が基本推計0.61％・包括推計1.23％で"
     "要介護度別の方が小さい。確認事項No.126",
     "サービス別の構成"),
    (8, "B", "保険料額の算定 → 準備基金",
     "残高・取崩額とも0",
     "令和7年度末の基金残高と取崩方針を入れる。未受領のため当方も0で置いている",
     "1億円の取崩しで月額 約▲310円"),
    (9, "B", "保険料額の算定 → 保険者機能強化推進交付金の交付見込額",
     "0",
     "当方は地域支援事業費から保険者機能強化推進事業費・"
     "保険者努力支援事業費を除く方法を採っている。"
     "この欄と併用すると二重控除になる。確認事項No.86",
     "併用すると過小になる"),
    (10, "C", "4_施策反映の解説",
     "「施策反映の全体方針」が「出力後に入力」のまま。"
     "設定内容・施策反映内容の欄も空欄",
     "計画の基本目標に沿った記述を入れる。"
     "実績見込み値を編集した場合はその理由も記入する",
     "―"),
    (11, "C", "1_推計値サマリ ７．介護保険料基準額",
     "%s円で、5_保険料推計の%s円と違う"
     % ("{:,.2f}".format(MI_GETSU1), "{:,.2f}".format(MI_GETSU)),
     "「保険料額の更新」を行ったうえで出力し直し、"
     "どちらが確定値かを画面で確かめる",
     "―"),
]
for x in SHUSEI:
    r = body(ws, r, list(x), height=52,
             fills={2: (NG_O if x[1] == "A" else
                        (IN_Y if x[1] == "B" else GRAY))})
r = note(ws, r,
         "注1）優先Aは第2次概算（10月30日）までに直すことを要するもの、"
         "優先Bはご判断又は資料の受領を要するもの、"
         "優先Cは計画に載せる前に整えるものです。\n"
         "注2）「効き」は他の条件を動かさずにその1件だけを直した場合の見当です。"
         "複数を直すと単純な足し算にはなりません。\n"
         "注3）#1の額は当方の算定における地域支援事業費であり、"
         "令和6年度決算を据え置いたものです（総合事業の令和7年度実績が未受領）。",
         span=6, height=60)

# ============================================================ 02
ws = sheet("02_設定の対照",
           "設定の対照（画面の設定と当方の推奨）",
           "4_施策反映の解説に記録されている設定を、当方の推奨と対照します。",
           [4, 34, 30, 30, 10])
r = 4
r = header(ws, r, ["#", "設定項目", "画面の設定", "当方の推奨", "判定"])
SUISHO = {
    "認定者数の自然体推計手法":
        ("性別／年齢5歳階級別／要介護度別", "一致"),
    "介護サービス利用者数の自然体推計手法":
        ("要介護度別（基本推計）。確認事項No.126", "要改"),
    "推計に用いた認定率の伸び":
        ("令和6年度→令和7年度の伸び（令和8年度を含まない唯一の選択肢）",
         "一致"),
    "推計に用いた利用率の伸び":
        ("令和6年度→令和7年度の伸び", "一致"),
    "推計に用いた1人1月あたりの給付費の実績値":
        ("令和7年度（令和8年度は月報4〜7月からの計算値）。確認事項No.137",
         "要改"),
    "推計に用いた1人1月あたり利用回（日）数の伸び":
        ("令和6年度→令和7年度の伸び", "一致"),
}
_hantei = []
for i, (k, v) in enumerate(K.SETTEI, start=1):
    rec = SUISHO.get(k)
    sui, hn = rec if rec else ("―", "―")
    _hantei.append(hn)
    r = body(ws, r, [i, k, v.replace("　", " "), sui, hn], height=32,
             fills={5: (OK_G if hn == "一致" else
                        (NG_O if hn == "要改" else GRAY))})
r = note(ws, r,
         "注1）「一致」は画面の設定が当方の推奨と同じものです。"
         "認定率・利用率・利用回（日）数の伸びはいずれも"
         "令和6年度→令和7年度が選ばれており、令和8年度の影響を避ける"
         "唯一の選択肢です。\n"
         "注2）1人1月あたり給付費の実績値は施設・居住系と在宅で別に選べます。"
         "いずれも令和8年度になっています。",
         span=5, height=46)

# ============================================================ 03
ws = sheet("03_令和8年度の基準値の点検",
           "令和8年度の実績見込み値の点検",
           "令和8年度の列は年度の実績ではなく、月報（令和8年4〜7月）からの"
           "計算値です。これを基準年として延ばしているため、"
           "令和8年度に跳ねた区分・落ちた区分がそのまま将来に残ります。",
           [4, 34, 14, 14, 12, 14, 34])
r = 4
r = lead(ws, r, "1　令和8年度が令和7年度と±10％を超えてずれている区分", span=7)
r = header(ws, r, ["#", "サービス", "R7 給付費\n（千円）",
                   "R8 給付費\n（千円）", "R8／R7",
                   "R11 給付費\n（千円）", "見方"])
ZURE = []
for (ku, nm), d in K.SVC.items():
    g = d.get("給付費")
    if not g:
        continue
    r7, r8 = g.get("R7"), g.get("R8")
    if not isinstance(r7, (int, float)) or not isinstance(r8, (int, float)):
        continue
    if r7 == 0:
        continue
    rate = r8 / r7
    if abs(rate - 1) > 0.10:
        ZURE.append((nm, r7, r8, rate, g.get("R11")))
ZURE.sort(key=lambda x: -abs(x[3] - 1))
for i, (nm, r7, r8, rate, r11) in enumerate(ZURE, start=1):
    mikata = ("令和8年度が0。令和9年度以降も0のまま残る" if r8 == 0 else
              ("令和8年度が高い。将来がその水準で延びる" if rate > 1 else
               "令和8年度が低い。将来がその水準で延びる"))
    r = body(ws, r, [i, nm, round(r7, 1), round(r8, 1), round(rate, 3),
                     (round(r11, 1) if isinstance(r11, (int, float))
                      else "―"), mikata],
             height=20, fills={5: NG_O},
             fmt={3: "#,##0.0", 4: "#,##0.0", 5: "0.000", 6: "#,##0.0"})
r = note(ws, r,
         "注1）給付費が令和8年度に0となっている区分は、"
         "償還払い（住宅改修費・特定福祉用具購入費）と"
         "令和7年度の利用が小口であった区分です。"
         "月報には現物給付しか計上されないため0になります。\n"
         "注2）令和9年度以降が0のまま残るため、"
         "法定の記載事項である見込量が立ちません。",
         span=7, height=46)

r += 1
r = lead(ws, r, "2　令和8年度の人数の欄が整数であること", span=7)
_int8 = _tot8 = _dec7 = _tot7 = 0
for (ku, nm), d in K.SVC.items():
    p = d.get("人数")
    if not p:
        continue
    v8, v7 = p.get("R8"), p.get("R7")
    if isinstance(v8, (int, float)) and v8:
        _tot8 += 1
        _int8 += abs(v8 - round(v8)) < 1e-9
    if isinstance(v7, (int, float)) and v7:
        _tot7 += 1
        _dec7 += abs(v7 - round(v7)) > 1e-9
r = header(ws, r, ["", "年度", "非ゼロの区分", "整数の区分", "小数の区分",
                   "", "意味"])
r = body(ws, r, ["", "令和7年度", _tot7, _tot7 - _dec7, _dec7, "",
                 "年報・月報の年間延べを12で除した月平均であるため小数になる"],
         height=20)
r = body(ws, r, ["", "令和8年度", _tot8, _int8, _tot8 - _int8, "",
                 "全て整数。実績見込み値が編集されていないことを示す"],
         height=20, fills={4: NG_O})
r = note(ws, r,
         "注）人数の欄は整数で扱われます（確認事項No.135）。"
         "令和8年度の人数が全て整数であることは、"
         "システムが月報から計算した値がそのまま入っていることを示します。",
         span=7, height=32)

# ============================================================ 04
ws = sheet("04_所得段階別の入力",
           "所得段階別第1号被保険者数の入力の点検",
           "第10期（令和9〜11年度）と令和12年度以降とで、"
           "計に用いている人口の系列が違っています。",
           [4, 16, 12, 12, 12, 12, 12, 12, 12, 12])
r = 4
YR = ["R9", "R10", "R11", "R12", "R17", "R22", "R27", "R32"]
YRL = ["令和9年度", "令和10年度", "令和11年度", "令和12年度",
       "令和17年度", "令和22年度", "令和27年度", "令和32年度"]
r = lead(ws, r, "1　入力されている計と、画面が表示する第1号被保険者数", span=10)
r = header(ws, r, ["", "区分"] + YRL)
r = body(ws, r, ["", "入力されている計"] + [K.DANKAI["計"][y] for y in YR],
         height=20, fmt={i: "#,##0" for i in range(3, 11)}, fills={2: MID_B})
r = body(ws, r, ["", "画面の第1号被保険者数"] + [K.HIHO[y] for y in YR],
         height=20, fmt={i: "#,##0" for i in range(3, 11)})
r = body(ws, r, ["", "加入割合の合計"]
         + [round(K.WARIAI_KEI[y], 4) for y in YR],
         height=20, fmt={i: "0.0000" for i in range(3, 11)},
         fills={i: (NG_O if abs(K.WARIAI_KEI[YR[i - 3]] - 1) > 1e-6 else OK_G)
                for i in range(3, 11)})
r = note(ws, r,
         "注1）加入割合の合計は1.0000でなければなりません。"
         "令和12年度以降で1.0を超えているのは、"
         "入力した計が画面の第1号被保険者数より多いためです。\n"
         "注2）第10期（令和9〜11年度）は画面の計に合わせてあり、"
         "加入割合の合計は1.0000です。\n"
         "注3）給付費の側は画面の第1号被保険者数（社人研）で推計されています。"
         "保険料の分母だけを別の人口で置くと、"
         "令和12年度以降の保険料が1.0〜1.5％低く出ます。",
         span=10, height=60)

r += 1
r = lead(ws, r, "2　どちらかにそろえる", span=10)
r = header(ws, r, ["#", "案", "内容", "第10期の計", "令和12年度以降の計",
                   "", "", "", "", ""])
r = body(ws, r, [1, "人口の設定を案Cに改める（推奨）",
                 "確認事項No.127。給付費の側も案Cになり、"
                 "所得段階別の計と一致する",
                 "現在のまま直す（案Cの計に改める）", "現在のままでよい",
                 "", "", "", "", ""], height=34, fills={2: OK_G})
r = body(ws, r, [2, "社人研のままとする",
                 "所得段階別の計を全8年度とも画面の第1号被保険者数に合わせる",
                 "現在のままでよい", "画面の計に改める",
                 "", "", "", "", ""], height=34)
r = note(ws, r,
         "注）いずれの案でも、加入割合の合計が全年度で1.0000になることを"
         "画面で確かめてください。",
         span=10, height=20)

r += 1
r = lead(ws, r, "3　段階の数（弾力化）", span=10)
r = header(ws, r, ["", "区分", "段階", "第10期3か年の補正後被保険者数",
                   "算定上の月額", "", "", "", "", ""])
r = body(ws, r, ["", "現在（標準段階区分）", 13, round(HOSEI13, 1),
                 round(MI_GETSU, 2), "", "", "", "", ""],
         height=20, fmt={4: "#,##0.0", 5: "#,##0.00"}, fills={3: NG_O})
r = body(ws, r, ["", "弾力化（条例の16段階）", 16, round(HOSEI16, 1),
                 round(GETSU16, 2), "", "", "", "", ""],
         height=20, fmt={4: "#,##0.0", 5: "#,##0.00"}, fills={3: OK_G})
r = body(ws, r, ["", "当方の算定", 16, round(TOU_HOSEI3, 1),
                 round(TOU_GETSU, 2), "", "", "", "", ""],
         height=20, fmt={4: "#,##0.0", 5: "#,##0.00"}, fills={2: MID_B})
r = note(ws, r,
         "注1）当連合の所得段階は第9期の条例により16段階です"
         "（第13段階を多段階化して第14〜16段階を置いています）。"
         "画面の標準は13段階で、9つの段階で割合が違います。\n"
         "注2）弾力化の欄は第30段階まであり、現在は全欄が0です。\n"
         "注3）当方の算定と補正後被保険者数が違うのは、"
         "人口の系列（案Cと社人研）が違うためです。乗率は同じです。\n"
         "注4）第10期の段階数・乗率は政令改正により変わり得ます"
         "（確認事項No.33）。",
         span=10, height=60)

# ============================================================ 05
ws = sheet("05_保険料の対照",
           "保険料の対照（見える化と当方の算定）",
           "同じ算定式であることは第9期の再現で確かめられます。"
           "第10期の差は、入力と設定の違いによるものです。",
           [4, 34, 22, 22, 18, 34])
r = 4
r = lead(ws, r, "1　第10期3か年（円）", span=6)
r = header(ws, r, ["#", "項目", "見える化", "当方の算定", "差", "見方"])
TAI = [
    ("総給付費", MI_SOGAKU3, TOU_SOGAKU3,
     "令和8年度を基準年としているため見える化が高い"),
    ("標準給付費見込額", MI_A, TOU_A,
     "上乗せ（特定入所者・高額・高額医療合算・審査支払手数料）は"
     "見える化5.35％、当方5.909％（令和6年度決算による割増率）"),
    ("地域支援事業費", MI_B, TOU_B,
     "見える化は未入力。入力を要する"),
    ("第1号被保険者負担分相当額", K.SHUNO["第1号被保険者負担分相当額"]["計"],
     G_DO["C"], "いずれも23％。負担割合は一致する"),
    ("調整交付金相当額", K.SHUNO["調整交付金相当額"]["計"], G_DO["D"],
     "いずれも5％"),
    ("調整交付金見込額", K.SHUNO["調整交付金見込額"]["計"], G_DO["E"],
     "見える化は補正係数から年度別に算定（7.61→7.07％）。"
     "当方は7.3755％で固定"),
    ("保険料収納必要額", MI_J, TOU_J,
     "地域支援事業費が入っていない分と、給付費が高い分が相殺している"),
]
for i, (nm, a, b, mi) in enumerate(TAI, start=1):
    r = body(ws, r, [i, nm, round(a), round(b), round(a - b), mi],
             height=32, fmt={3: "#,##0", 4: "#,##0", 5: "#,##0"})
r = body(ws, r, ["", "算定上の月額（円）", round(MI_GETSU, 2),
                 round(TOU_GETSU, 2), round(MI_GETSU - TOU_GETSU, 2),
                 "5_保険料推計による"],
         height=22, bold=True, fmt={3: "#,##0.00", 4: "#,##0.00",
                                    5: "#,##0.00"},
         fills={i: MID_B for i in range(1, 7)})
r = note(ws, r,
         "注1）第9期の保険料基準額は見える化%s円、当方の再現%s円で"
         "一致しています。算定式そのものは同じものです。\n"
         "注2）見える化の月額は地域支援事業費が0のままの値です。"
         "当方の地域支援事業費を入れると約%s円上がり、約%s円になります。\n"
         "注3）さらに弾力化（条例16段階）にすると約%s円になります。"
         % ("{:,.2f}".format(K.HOKENRYO["第9期"]), "6,427.92",
            "{:,.0f}".format(CHIIKI_SA),
            "{:,.0f}".format(MI_GETSU + CHIIKI_SA),
            "{:,.0f}".format(GETSU16 + DJ / 0.99 / HOSEI16 / 12)),
         span=6, height=52)

r += 1
r = lead(ws, r, "2　同じ出力の中の2つの保険料基準額（円）", span=6)
r = header(ws, r, ["", "出所"] + K.HOKENRYO["年度"] + [""])
r = body(ws, r, ["", "1_推計値サマリ ７"]
         + [round(K.HOKENRYO["1_推計値サマリ"][y], 1)
            for y in K.HOKENRYO["年度"]] + [""],
         height=20, fmt={i: "#,##0.0" for i in range(3, 9)}, fills={2: NG_O})
r = body(ws, r, ["", "5_保険料推計 ２"]
         + [round(K.HOKENRYO["5_保険料推計"][y], 1)
            for y in K.HOKENRYO["年度"]] + [""],
         height=20, fmt={i: "#,##0.0" for i in range(3, 9)}, fills={2: OK_G})
r = note(ws, r,
         "注）5シートの値は、入力されている所得段階別被保険者数と"
         "標準13段階の割合から当方が検算して一致します"
         "（%s円÷0.99÷%s人÷12＝%s円）。"
         "1シートの値は所得段階別の登録前のものとみられます。"
         "画面の上部に表示されている暫定値も1シートと同じです。"
         % ("{:,.0f}".format(MI_J), "{:,.2f}".format(HOSEI13),
            "{:,.4f}".format(MI_GETSU)),
         span=6, height=34)

# ============================================================ 06
ws = sheet("06_サービス別の対照",
           "サービス別の対照（令和11年度の利用者数）",
           "見える化は介護と予防を別の行に持ちます。"
           "当方の算定は総括表の区分（介護＋予防）であるため、"
           "対照では見える化の側を合算しています。",
           [4, 34, 14, 14, 12, 34])
r = 4
r = header(ws, r, ["#", "サービス", "見える化\nR11（人／月）",
                   "当方\nR11（人／月）", "差", "見方"])
# 当方の令和11年度の利用者数（サービス名→人／月）。総括表の区分による。
_TOU_R11 = {_S["short"](l): _S["sogaku"](l, Y3[2]) for l in _S["SVC"]}
PAIR = [
    ("訪問介護", ["訪問介護"]),
    ("訪問入浴介護", ["訪問入浴介護", "介護予防訪問入浴介護"]),
    ("訪問看護", ["訪問看護", "介護予防訪問看護"]),
    ("訪問リハビリテーション", ["訪問リハビリテーション",
                                "介護予防訪問リハビリテーション"]),
    ("居宅療養管理指導", ["居宅療養管理指導", "介護予防居宅療養管理指導"]),
    ("通所介護", ["通所介護"]),
    ("通所リハビリテーション", ["通所リハビリテーション",
                                "介護予防通所リハビリテーション"]),
    ("短期入所生活介護", ["短期入所生活介護", "介護予防短期入所生活介護"]),
    ("短期入所療養介護（老健）", ["短期入所療養介護（老健）",
                                  "介護予防短期入所療養介護（老健）"]),
    ("福祉用具貸与", ["福祉用具貸与", "介護予防福祉用具貸与"]),
    ("特定福祉用具販売", ["特定福祉用具購入費",
                          "特定介護予防福祉用具購入費"]),
    ("住宅改修", ["住宅改修費", "介護予防住宅改修"]),
    ("定期巡回・随時対応型訪問介護看護",
     ["定期巡回・随時対応型訪問介護看護"]),
    ("地域密着型通所介護", ["地域密着型通所介護"]),
    ("認知症対応型通所介護", ["認知症対応型通所介護",
                              "介護予防認知症対応型通所介護"]),
    ("小規模多機能型居宅介護", ["小規模多機能型居宅介護",
                                "介護予防小規模多機能型居宅介護"]),
    ("認知症対応型共同生活介護", ["認知症対応型共同生活介護",
                                  "介護予防認知症対応型共同生活介護"]),
    ("特定施設入居者生活介護", ["特定施設入居者生活介護",
                                "介護予防特定施設入居者生活介護"]),
    ("地域密着型介護老人福祉施設入所者生活介護",
     ["地域密着型介護老人福祉施設入所者生活介護"]),
    ("介護老人福祉施設", ["介護老人福祉施設"]),
    ("介護老人保健施設", ["介護老人保健施設"]),
    ("介護医療院", ["介護医療院"]),
]
_by = {}
for (ku, nm), d in K.SVC.items():
    p = d.get("人数")
    if p and isinstance(p.get("R11"), (int, float)):
        _by[nm] = p["R11"]
_ng = []
for i, (tou, mis) in enumerate(PAIR, start=1):
    mv = sum(_by.get(m, 0) for m in mis)
    tv = _TOU_R11.get(tou)
    if tv is None:
        continue
    sa = mv - tv
    mikata = ""
    if abs(sa) >= max(5.0, abs(tv) * 0.12):
        mikata = "令和8年度の基準値の違いが効いている"
        _ng.append(tou)
    r = body(ws, r, [i, tou, round(mv, 1), round(tv, 1), round(sa, 1),
                     mikata], height=20,
             fmt={3: "#,##0.0", 4: "#,##0.0", 5: "#,##0.0"},
             fills={5: (NG_O if mikata else None)})
r = note(ws, r,
         "注1）差が大きい区分は、令和8年度の実績見込み値が"
         "令和7年度と大きく違うものです（03シート）。\n"
         "注2）住宅改修・定期巡回は見える化が0です。"
         "令和8年度の欄が0のまま将来に残っています。\n"
         "注3）当方の算定は要介護度別に利用率を固定して延ばしたものであり、"
         "見える化の包括推計とは延ばし方も違います。",
         span=6, height=46)

# ============================================================ 07 人口の設定
# 案C（総人口＝地方創生総合戦略、年齢階級別＝住民基本台帳の実績趨勢）を
# 「独自データを登録する」の欄へ入れるための値。
ANC = AN.ANC
JY = J.YEARS
pop_tot, pop_juki, CL3 = AN.pop_tot, AN.pop_juki, AN.CL3
_r0, _saidai = AN.r0, AN.saidai
JITSU, PAIR3 = AN.JITSU, AN.PAIR3


def _i1(y):
    return sum(ANC[y]["男"].values()) + sum(ANC[y]["女"].values())


def _kouki(d):
    return sum(d["男"][a] + d["女"][a] for a in J.AGE[2:])


ws = sheet("07_人口の設定の入力値",
           "「総人口と被保険者数の設定」に入れる値（案C）",
           "「独自データを登録する」を選んだ場合の入力値です。"
           "**令和6年度・令和7年度は実績であるため画面の値をそのまま入れ、"
           "令和8年度以降を案C（総人口＝地方創生総合戦略、"
           "年齢階級別＝住民基本台帳の実績趨勢）に改めます。**"
           "白い欄（6階級と第2号被保険者）が入力欄で、"
           "第1号被保険者と総数は自動計算されます。",
           [4, 18, 11, 11, 11, 11, 11, 11, 11, 11, 11, 11, 11])
r = 4
_AL = {j: "right" for j in range(3, 14)}
_FMT = {j: "#,##0" for j in range(3, 14)}
_HD = ["", "区分"] + JY

r = lead(ws, r, "1　総人口（人）", span=13)
r = header(ws, r, _HD)
r = body(ws, r, ["", "案C（入力する値）"] + [ANC[y]["総人口"] for y in JY],
         height=20, fills={j: IN_Y for j in range(3, 14)}, fmt=_FMT,
         align=_AL, bold=True)
r = body(ws, r, ["", "画面の現在値"] + [J.SOJINKO[y] for y in JY],
         height=20, fmt=_FMT, align=_AL)
r = body(ws, r, ["", "差（案C−画面）"]
         + [ANC[y]["総人口"] - J.SOJINKO[y] for y in JY],
         height=20, fills={j: GRAY for j in range(1, 14)},
         fmt={j: "+#,##0;-#,##0;0" for j in range(3, 14)}, align=_AL)

r += 1
r = lead(ws, r, "2　被保険者数（1）男（人）", span=13)
r = header(ws, r, _HD)
for a in J.AGE:
    r = body(ws, r, ["", J.AGEL[a]] + [ANC[y]["男"][a] for y in JY],
             height=18, fills={j: IN_Y for j in range(3, 14)}, fmt=_FMT,
             align=_AL)
r = body(ws, r, ["", "第2号被保険者"] + [J.dan("第2号", y) for y in JY],
         height=18, fills={j: IN_Y for j in range(3, 14)}, fmt=_FMT,
         align=_AL)
r = body(ws, r, ["", "第1号被保険者（自動）"]
         + [sum(ANC[y]["男"].values()) for y in JY],
         height=18, fills={j: GRAY for j in range(1, 14)}, fmt=_FMT,
         align=_AL)

r += 1
r = lead(ws, r, "3　被保険者数（2）女（人）", span=13)
r = header(ws, r, _HD)
for a in J.AGE:
    r = body(ws, r, ["", J.AGEL[a]] + [ANC[y]["女"][a] for y in JY],
             height=18, fills={j: IN_Y for j in range(3, 14)}, fmt=_FMT,
             align=_AL)
r = body(ws, r, ["", "第2号被保険者"] + [J.jo("第2号", y) for y in JY],
         height=18, fills={j: IN_Y for j in range(3, 14)}, fmt=_FMT,
         align=_AL)
r = body(ws, r, ["", "第1号被保険者（自動）"]
         + [sum(ANC[y]["女"].values()) for y in JY],
         height=18, fills={j: GRAY for j in range(1, 14)}, fmt=_FMT,
         align=_AL)

r += 1
r = lead(ws, r, "4　入力後に画面が表示する値（自動計算）と現在値の対照", span=13)
r = header(ws, r, _HD)
r = body(ws, r, ["", "第1号被保険者 案C"] + [_i1(y) for y in JY],
         height=20, fills={j: OK_G for j in range(3, 14)}, fmt=_FMT,
         align=_AL, bold=True)
r = body(ws, r, ["", "第1号被保険者 画面"] + [J.ichigo(y) for y in JY],
         height=20, fmt=_FMT, align=_AL)
r = body(ws, r, ["", "差（案C−画面）"]
         + [_i1(y) - J.ichigo(y) for y in JY],
         height=20, fills={j: GRAY for j in range(1, 14)},
         fmt={j: "+#,##0;-#,##0;0" for j in range(3, 14)}, align=_AL)
r = body(ws, r, ["", "後期（75歳〜）の割合 案C"]
         + [round(_kouki(ANC[y]) / _i1(y), 4) for y in JY],
         height=20, fmt={j: "0.0000" for j in range(3, 14)}, align=_AL)
r = body(ws, r, ["", "同 画面"]
         + [round(sum(J.kei(a, y) for a in J.AGE[2:]) / J.ichigo(y), 4)
            for y in JY],
         height=20, fmt={j: "0.0000" for j in range(3, 14)}, align=_AL)
r = note(ws, r,
         "注1）**令和6年度・令和7年度は実績であるため画面の値をそのまま入れます。**"
         "令和8年度以降が案Cによる値です。\n"
         "注2）**令和8年度の第1号被保険者数は案Cでも9,117人で画面と同じ**であり、"
         "動くのは年齢階級別の内訳だけです。\n"
         "注3）当方の案Cは65歳以上を65〜74歳・75〜84歳・85歳以上の3区分で推計"
         "しています。**5歳階級への分け方と男女の分け方は、"
         "画面が現に表示している補正データの同じ年の構成比によっています。**"
         "住民基本台帳の5歳階級の趨勢をそのまま延ばすと、"
         "65〜69歳が年▲4.2％・70〜74歳が年＋1.0％という"
         "特定の世代が移っていく動きをそのまま繰り返すことになり、"
         "令和11年度の65〜69歳の割合が0.41（画面は0.51）まで下がります。"
         "3区分より細かい構成は画面の側によるのが妥当と考えます。\n"
         "注4）**第2号被保険者（40歳から64歳）は画面の値をそのままとしています。**"
         "案Cは総人口と65歳以上の趨勢を定めるもので、"
         "40歳から64歳の内訳を持っていません。"
         "第1号被保険者の保険料の算定には用いられません。\n"
         "注5）**後期高齢者の割合が画面より高くなります**"
         "（令和11年度 0.6340対0.6174）。"
         "調整交付金の後期高齢者加入割合補正係数が上がるため、"
         "調整交付金見込額が増え保険料を下げる向きに働きます。"
         "一方で認定者数が増えるため給付費は上がる向きに働きます。"
         "**どちらが上回るかは入力後の算定によります。**",
         span=13, height=140)

r += 1
r = lead(ws, r, "5　併せて入れ直すもの", span=13)
r = header(ws, r, ["#", "画面", "入れ直す内容", "", "", "", "", "", "",
                   "", "", "", ""])
for _i, (_a, _b) in enumerate([
        ("所得段階別第1号被保険者数①（標準段階区分）",
         "計が案Cの第1号被保険者数（9,147／9,179／9,213／9,211／9,250／"
         "9,491／9,395／9,127人）になる表に入れ直す。"
         "これにより加入割合の合計が全年度で1.0000になる"),
        ("所得段階別第1号被保険者数②（弾力化）",
         "当連合の条例による16段階で入れる。"
         "①と②の計は同じ第1号被保険者数にそろえる"),
        ("実績及び推計方法の設定（認定率・利用率の伸び）",
         "「令和6年度→令和7年度の伸び」のままでよい（変更は要らない）"),
        ("保険料額の算定（保険料収納必要額）",
         "地域支援事業費を入れる。人口の設定とは別の作業である")],
        start=1):
    r = body(ws, r, [_i, _a, _b] + [""] * 10, height=34)
r = note(ws, r,
         "注）人口の設定を改めると、認定者数・サービス見込量・給付費・"
         "保険料のすべてが計算し直されます。"
         "所得段階別の表を入れ直さないと、"
         "加入割合の合計が1.0にならない状態が残ります。",
         span=13, height=34)


# ============================================================ 08 弾力化の入力値
def _stage():
    """当連合の16段階の定義をソースから読む。

    `build_premium_bracket_review.py` の `STAGE` を `ast` で読み、
    段階・対象者要件・割合（公費軽減前）と、対象者要件の
    「合計所得金額○○万円以上」から基準所得金額を求める。
    固定値を書き写すと、段階の定義を改めたときに本表とずれる。
    """
    import ast
    import re
    src = io.open(os.path.join(RP.ROOT, "build_premium_bracket_review.py"),
                  encoding="utf-8").read()
    m = re.search(r"^STAGE = (\[.*?\n\])\s*$", src, re.S | re.M)
    if not m:
        raise RuntimeError("STAGE が見つからない")
    out = []
    for row in ast.literal_eval(m.group(1)):
        g = re.search(r"合計所得金額([0-9,]+)万円以上", row[1])
        out.append((row[0], row[1].replace("\n", ""), row[2],
                    int(g.group(1).replace(",", "")) * 10000 if g else None))
    return out


STAGE16 = _stage()
_DAN_YM = dict(zip(GD.YEARS, ("2027", "2028", "2029", "2030", "2035",
                              "2040", "2045", "2050")))
_DAN_TO = _load_dankai = runpy.run_path(
    os.path.join(RP.ROOT, "build_shotoku_dankai.py"))["DANKAI"]

ws = sheet("08_弾力化の入力値",
           "「②保険料基準額に対する割合の弾力化」に入れる基準所得金額と割合",
           "画面には既に所得段階別第1号被保険者数（16段階）が入っており、"
           "**未入力は基準所得金額と基準額に対する割合の2つ**です。"
           "本シートはその2つに入れる値です。"
           "基準所得金額は第7段階から第16段階の10件、"
           "割合は16段階すべてで、いずれも年度によらず同じ値です。",
           [4, 42, 11, 15, 13, 13, 30])
r = 4
_ALS = {3: "center", 4: "right", 5: "right", 6: "center"}
r = lead(ws, r, "1　基準所得金額と基準額に対する割合（全年度共通）", span=7)
r = header(ws, r, ["段階", "対象者要件", "割合", "基準所得金額（円）",
                   "画面の現在値", "標準13段階の割合", "備考"])
for _no, _yoken, _jor, _kin in STAGE16:
    _std = GAMEN13_JOR[_no - 1] if _no <= 13 else None
    _bik = []
    if _no <= 3:
        _bik.append("公費軽減の対象（割合は軽減前）")
    if _std is None:
        _bik.append("多段階化した分。標準13段階にない")
    elif abs(_std - _jor) > 1e-9:
        _bik.append("標準と違う")
    if _kin is None:
        _bik.append("基準所得金額の欄なし")
    r = body(ws, r, [_no, _yoken, _jor,
                     _kin if _kin else "―",
                     "0" if _kin else "―",
                     _std if _std is not None else "―",
                     "／".join(_bik) or ""],
             fills={3: IN_Y, 4: (IN_Y if _kin else GRAY)},
             fmt={4: "#,##0", 3: "0.000", 6: "0.0000"},
             align=_ALS, height=30)
r = note(ws, r,
         "注1）**割合は公費軽減前の値です。**"
         "第1段階から第3段階は公費軽減の対象ですが、"
         "補正後被保険者数はこの割合で算定します。"
         "見える化システムが「①標準段階区分・割合」の画面に表示する"
         "0.4550・0.6850・0.6900…も公費軽減前の値であり、基準がそろっています。"
         "公費軽減後の値（第1段階0.285など）は入れません。\n"
         "注2）**基準所得金額は第7段階から第16段階の10件です。**"
         "第1段階から第6段階は本人の合計所得金額の下限による区分ではないため"
         "（生活保護受給・世帯の課税状況・合計所得金額＋課税年金収入額による）"
         "画面も当該欄を受け付けません。"
         "第7段階から第13段階の額は画面が標準として表示している額と一致し、"
         "第14段階から第16段階が多段階化した分です。\n"
         "注3）**割合は9つの段階で標準13段階と違い、"
         "第14段階から第16段階は標準にありません。**"
         "これが弾力化を要する理由です。\n"
         "注4）第10期の段階数・割合・基準所得金額は政令改正により"
         "変わり得ます（確認事項No.33）。"
         "令和8年3月の全国課長会議資料は、保険料段階の基準額を"
         "80.9万円から82.65万円に改める政令改正が済んでおり、"
         "**条例改正の手続を要する**としています。"
         "本表は第9期の条例による暫定の値です。",
         span=7, height=140)

r += 1
r = lead(ws, r, "2　画面に入力済みの第1号被保険者数と当方の値の対照（人）",
         span=7)
r = header(ws, r, ["", "年度", "画面の計", "当方の計（案C）", "差", "段階別",
                   "備考"])
for _i, _y in enumerate(GD.YEARS, start=1):
    _to = _DAN_TO[_DAN_YM[_y]]
    _sa = GD.KEI[_y] - sum(_to)
    _ng = [j + 1 for j in range(GD.NDAN) if GD.HIHO[_y][j] != _to[j]]
    r = body(ws, r, [_i, GD.YEARSL[_i - 1], GD.KEI[_y], sum(_to), _sa,
                     "一致" if not _ng else "違い 第%s段階"
                     % "・".join(str(x) for x in _ng),
                     "画面の計＝段階別の和" if sum(GD.HIHO[_y]) == GD.KEI[_y]
                     else "画面の計と段階別の和が合わない"],
             fills={6: (OK_G if not _ng else NG_O)},
             fmt={3: "#,##0", 4: "#,##0", 5: "+#,##0;-#,##0;0"},
             align={3: "right", 4: "right", 5: "right", 6: "center"},
             height=20)
r = note(ws, r,
         "注1）画面に入っている16段階の人数は、"
         "当方が人口の設定（案C）からお示しした値と全年度・全段階で"
         "一致しています。"
         "**人数の入れ直しは要りません。**"
         "第10期と令和12年度以降が同じ人口の系列でそろっており、"
         "01シートに掲げた「第10期と令和12年度以降で別の人口による」"
         "（加入割合の合計が1.0を超える）は、この画面では解消しています。\n"
         "注2）**画面の上部に表示されている保険料額は暫定値であり、"
         "従前の写しから動いています**"
         "（第10期 %s円から%s円、令和12年度以降 %s円から%s円）。"
         "地域支援事業費が0のままであることなど、"
         "01シートの修正が済むまでは当方の算定と突き合わせられません。\n"
         "注3）第17段階から第30段階は空欄のままとします"
         "（当連合の所得段階は16段階です）。"
         % ("{:,}".format(GD.HOKENRYO_ZENKAI["第10期"]),
            "{:,}".format(GD.HOKENRYO_GAMEN["第10期"]),
            "{:,}".format(GD.HOKENRYO_ZENKAI["令和12年度以降"]),
            "{:,}".format(GD.HOKENRYO_GAMEN["令和12年度以降"])),
         span=7, height=104)


# ============================================================ 09 再出力の点検
K2 = __import__("data_mieru_kekka2")
GW = __import__("data_mieru_warning")

# ワーニングの要因（サービスごと）。件数を数える側より先に置く。
_YOIN = {
    "訪問介護": ("当区域の実態", "当方の年報令和7年度でも要介護計55.4回／月で"
                 "全国の上限50回を上回ります。訪問介護の利用回数が多いことは"
                 "当区域の実態であり、入力の誤りではありません。"
                 "理由を記入して進めます"),
    "訪問入浴介護": ("令和8年度の実績見込み値",
                     "当方の年報令和7年度は要介護計6.3回／月で閾値の範囲内です。"
                     "令和8年度（月報4〜7月）の値が使われていることによります"),
    "短期入所生活介護": ("令和8年度の実績見込み値",
                         "当方の年報令和7年度は要支援7.9日／月で閾値の範囲内です。"
                         "同上"),
    "地域密着型通所介護": ("令和8年度の実績見込み値",
                           "当方の年報令和7年度は要介護計7.9回／月であり、"
                           "画面の34.3〜36.2回／月は4倍を超えます。"
                           "月に31回を超える値であり、そのままでは用いられません"),
    "認知症対応型通所介護": ("令和8年度の実績見込み値",
                             "利用回（日）数が0のまま令和9年度以降も0で"
                             "残っています。当方の年報令和7年度は"
                             "要介護計7.9回／月です"),
    "特定施設入居者生活介護": ("1人1月あたり給付費の年度",
                               "当方の年報令和7年度は要介護5で238,772円／月です。"
                               "画面の350,399円／月は令和8年度（月報4〜7月）に"
                               "よるものとみられます（確認事項No.137）"),
}

_W_R8 = sum(1 for _w in GW.MEISAI
            if _YOIN.get(_w[1], ("", ""))[0] == "令和8年度の実績見込み値")
_W_TANKA = sum(1 for _w in GW.MEISAI
               if _YOIN.get(_w[1], ("", ""))[0] == "1人1月あたり給付費の年度")
_W_JITTAI = sum(1 for _w in GW.MEISAI
                if _YOIN.get(_w[1], ("", ""))[0] == "当区域の実態")

_Y3 = ("R9", "R10", "R11")
_STD13 = [0.4550, 0.6850, 0.6900, 0.9000, 1.0000, 1.2000, 1.3000,
          1.5000, 1.7000, 1.9000, 2.1000, 2.3000, 2.4000]


def _to13(v):
    return list(v[:12]) + [sum(v[12:])]


def _quiet(name):
    buf, old = io.StringIO(), sys.stdout
    sys.stdout = buf
    try:
        return runpy.run_path(os.path.join(RP.ROOT, name))
    finally:
        sys.stdout = old


_JOR16 = _quiet("build_shotoku_dankai.py")["JORITSU"]
_H16 = sum(sum(a * b for a, b in zip(GD.HIHO[y], _JOR16)) for y in _Y3)
_H13 = sum(sum(a * b for a, b in zip(_to13(GD.HIHO[y]), _STD13))
           for y in _Y3)
_J2 = K2.SHUNO["保険料収納必要額"]["第10期計"]
_G13 = _J2 / 0.99 / _H13 / 12
_G16 = _J2 / 0.99 / _H16 / 12
_WARI = K2.KOSEI["調整交付金見込交付割合"]
_DJ = sum(TOU_B / 3 * (0.23 + 0.05 - _WARI[y]) for y in _Y3)
_G13B = (_J2 + _DJ) / 0.99 / _H13 / 12
_G16B = (_J2 + _DJ) / 0.99 / _H16 / 12

ws = sheet("09_再出力の点検",
           "人口の設定を案Cに改めた後の再出力の点検",
           "%s受領。推計パターン %s。"
           "人口の設定を案Cに改めた後の出力を、改める前の出力及び"
           "当方の算定と突き合わせたものです。"
           % (K2.META["出力日"], K2.META["推計パターン名"]),
           [4, 30, 22, 22, 22, 36])
r = 4
r = lead(ws, r, "1　人口の設定の変更が及んでいる箇所・及んでいない箇所", span=6)
r = header(ws, r, ["#", "項目", "改める前", "改めた後", "差", "見方"])
_ROWS9 = [
    ("第1号被保険者数（第10期3か年）",
     "{:,}人".format(K.DANKAI["計"]["R9"] + K.DANKAI["計"]["R10"]
                     + K.DANKAI["計"]["R11"]),
     "{:,}人".format(sum(K2.HIHO[y] for y in _Y3)),
     "+{:,}人".format(sum(K2.HIHO[y] for y in _Y3)
                      - (K.DANKAI["計"]["R9"] + K.DANKAI["計"]["R10"]
                         + K.DANKAI["計"]["R11"])),
     "案Cに変わった。当方の算定と一致する", True),
    ("後期高齢者加入割合補正係数（令和11年度）",
     "{:.4f}".format(K.KOSEI["後期高齢者加入割合補正係数"]["R11"]),
     "{:.4f}".format(K2.KOSEI["後期高齢者加入割合補正係数"]["R11"]),
     "{:+.4f}".format(K2.KOSEI["後期高齢者加入割合補正係数"]["R11"]
                      - K.KOSEI["後期高齢者加入割合補正係数"]["R11"]),
     "後期高齢者の割合が上がったため下がった", True),
    ("調整交付金見込額（第10期3か年）",
     "{:,}円".format(int(K.SHUNO["調整交付金見込額"]["計"])),
     "{:,}円".format(int(K2.SHUNO["調整交付金見込額"]["計"])),
     "{:+,}円".format(int(K2.SHUNO["調整交付金見込額"]["計"]
                         - K.SHUNO["調整交付金見込額"]["計"])),
     "増えた。保険料を下げる向きに働く", True),
    ("保険料収納必要額（第10期3か年）",
     "{:,}円".format(int(K.SHUNO["保険料収納必要額"]["計"])),
     "{:,}円".format(int(_J2)),
     "{:+,}円".format(int(_J2 - K.SHUNO["保険料収納必要額"]["計"])),
     "減った額は調整交付金の増加額と一致する", True),
    ("要介護（支援）認定者数（令和11年度・第1号）",
     "{:,}人".format(K.NINTEI["第1号"]["R11"]),
     "{:,}人".format(K2.NINTEI_ICHI["R11"]),
     "±0人",
     "人口が変わったのに1人も動いていない", False),
    ("総給付費（第10期3か年）",
     "{:,}千円".format(int(sum(K.SOGAKU["合計"][y] for y in _Y3))),
     "{:,}千円".format(int(sum(K2.KYUFU[y] for y in _Y3))),
     "±0千円", "同上", False),
    ("標準給付費見込額（第10期3か年）",
     "{:,}円".format(int(K.SHUNO["標準給付費見込額"]["計"])),
     "{:,}円".format(int(K2.SHUNO["標準給付費見込額"]["計"])),
     "±0円", "同上", False),
]
for _i, (_a, _b, _c, _d, _e, _ok) in enumerate(_ROWS9, start=1):
    r = body(ws, r, [_i, _a, _b, _c, _d, _e],
             fills={5: (OK_G if _ok else NG_O)},
             align={3: "right", 4: "right", 5: "right"}, height=24)
r = note(ws, r,
         "注1）**人口の設定の変更は、第1号被保険者数と調整交付金にしか"
         "及んでいません。**"
         "認定者数・サービス見込量・総給付費・標準給付費見込額は"
         "改める前の出力と1円も変わっていません。"
         "認定者数は8年度すべてで同じ値です。\n"
         "注2）認定者数は性別・年齢5歳階級別・要介護度別の認定率に"
         "将来人口を乗じて推計するものであり、"
         "人口が変われば認定者数も変わります。"
         "**「実績及び推計方法の設定」を保存したうえで、"
         "施策反映（認定者数）→（施設・居住系）→（在宅）→ 地域支援事業 →"
         "保険料額の算定 の順に各画面を開いて登録し直すことを要します。**"
         "登録し直せたかどうかは、"
         "**令和11年度の認定者数（第1号）が%s人から変わること**で"
         "確かめられます。\n"
         "注3）保険料収納必要額の減少額%s円は、"
         "調整交付金見込額の増加額と円単位で一致します。"
         "給付費の側が動いていないことの裏づけです。"
         % ("{:,}".format(K2.NINTEI_ICHI["R11"]),
            "{:,}".format(int(K.SHUNO["保険料収納必要額"]["計"] - _J2))),
         span=6, height=110)

r += 1
r = lead(ws, r, "2　修正が必要な箇所11件の現況", span=6)
r = header(ws, r, ["#", "修正が必要な箇所", "前回", "今回", "", "見方"])
_GEN = [
    ("保険料基準額が同じ出力の中で2通りある",
     "1シート%s円／5シート%s円"
     % ("{:,.0f}".format(MI_GETSU1), "{:,.0f}".format(MI_GETSU)),
     "1シートも5シートも{:,.0f}円".format(K2.HOKENRYO["第10期"]),
     "解消", "所得段階別が登録されたため一致した"),
    ("所得段階別の計が第10期と令和12年度以降で別の人口による",
     "加入割合の合計が最大1.0149", "全8年度とも1.0000",
     "解消", "所得段階別加入割合補正係数も全8年度0.9856にそろった"),
    ("人口の設定が社人研のまま", "社人研", "案C",
     "解消", "ただし給付費の側へ及んでいない（本シート1）"),
    ("地域支援事業費が全ての欄で0", "0円", "0円",
     "未", "当方の第10期3か年{:,}円。月額 約＋{:,.0f}円"
     .format(TOU_B, _G13B - _G13)),
    ("所得段階を弾力化していない", "標準13段階", "標準13段階",
     "未", "条例は16段階。月額 約{}{:,.0f}円"
     .format("＋" if _G16 >= _G13 else "▲", abs(_G16 - _G13))),
    ("令和8年度の実績見込み値を編集していない",
     "月報からの計算値のまま", "月報からの計算値のまま",
     "未", "ワーニング{}件のうち{}件がこれによる（10シート）"
     .format(len(GW.MEISAI), _W_R8)),
    ("1人1月あたり給付費の実績値が令和8年度",
     "令和8年度", "令和8年度",
     "未", "確認事項No.137。ワーニング{}件がこれによる".format(_W_TANKA)),
    ("介護サービス利用者数の自然体推計手法が包括推計",
     "包括推計", "包括推計", "未", "確認事項No.126。当方は要介護度別"),
    ("準備基金が0", "0円", "0円", "未", "令和7年度末の基金残高の提供を要する"),
    ("保険者機能強化推進交付金の欄が0", "0円", "0円",
     "未", "確認事項No.86。Bを入れるときに二重控除しない"),
    ("施策反映の記入欄が空欄", "「出力後に入力」のまま",
     "「出力後に入力」のまま", "未", "計画に載せる前に記入を要する"),
]
for _i, (_a, _b, _c, _st, _e) in enumerate(_GEN, start=1):
    r = body(ws, r, [_i, _a, _b, _c, _st, _e],
             fills={5: (OK_G if _st == "解消" else NG_O)},
             align={5: "center"}, height=26)
r = note(ws, r,
         "注）11件のうち**3件が解消し、8件が残って**います。"
         "残る8件のうち月額を動かすのは、"
         "地域支援事業費・弾力化・令和8年度の実績見込み値・"
         "1人1月あたり給付費の年度・準備基金の5件です。",
         span=6, height=34)

r += 1
r = lead(ws, r, "3　保険料基準額（月額）の見通し", span=6)
r = header(ws, r, ["#", "置き方", "補正後被保険者数", "月額", "現況との差",
                   "備考"])
_MIT = [
    ("現況（標準13段階・地域支援事業費0）", _H13, _G13,
     "出力の値%.2f円と%.2f円の差" % (K2.HOKENRYO["第10期"],
                                     _G13 - K2.HOKENRYO["第10期"])),
    ("弾力化のみ（条例16段階）", _H16, _G16, "条例の割合による"),
    ("地域支援事業費のみ", _H13, _G13B, "当方のB {:,}円".format(TOU_B)),
    ("弾力化＋地域支援事業費", _H16, _G16B, "この2件を入れた場合の見通し"),
]
for _i, (_a, _h, _g, _e) in enumerate(_MIT, start=1):
    r = body(ws, r, [_i, _a, round(_h, 1), round(_g, 2),
                     "{:+,.0f}円".format(_g - _G13), _e],
             fills={4: (OK_G if _i == 4 else None)},
             fmt={3: "#,##0.0", 4: "#,##0.00"},
             align={3: "right", 4: "right", 5: "right"}, height=22)
r = body(ws, r, ["", "（参考）当方の算定", TOU_HOSEI3, TOU_GETSU,
                 "{:+,.0f}円".format(TOU_GETSU - _G13),
                 "令和7年度基準・要介護度別・地域支援事業費を含む"],
         fills={i: GRAY for i in range(1, 7)},
         fmt={3: "#,##0.0", 4: "#,##0.00"},
         align={3: "right", 4: "right", 5: "right"}, height=22)
r = note(ws, r,
         "注1）本シートの月額は保険料収納必要額%s円から計算したものです。"
         "地域支援事業費を入れた行は、当方のB（%s円）を3か年に均等に置き、"
         "年度別の調整交付金見込交付割合で第1号負担分を求めたものです。\n"
         "注2）**現況と当方の算定の差は、"
         "地域支援事業費・弾力化を入れてもなお残ります。**"
         "残る差は、基準年度が令和8年度であること（確認事項No.137）、"
         "包括推計であること（同No.126）、"
         "給付費の側が人口の変更を受けていないこと（本シート1）によります。\n"
         "注3）**いずれも暫定値です。**"
         % ("{:,}".format(int(_J2)), "{:,}".format(TOU_B)),
         span=6, height=72)


# ============================================================ 10 ワーニング
def _ours(svc, do):
    """当方の年報令和7年度の値（要支援計／要介護計）。"""
    m = {"訪問介護": "在宅サービス 訪問介護",
         "訪問入浴介護": "在宅サービス 訪問入浴介護",
         "短期入所生活介護": "在宅サービス 短期入所生活介護",
         "地域密着型通所介護": "在宅サービス 地域密着型通所介護",
         "認知症対応型通所介護": "在宅サービス 認知症対応型通所介護"}
    if svc in m:
        v = _S["kaisu_tanka"](m[svc])
        x = v[0] if do.startswith("要支援") else v[1]
        return ("%.1f" % x) if x else "―"
    if svc == "特定施設入居者生活介護":
        t = _S["tanka_do"]("居住系サービス 特定施設入居者生活介護")[0]
        i = ["要支援1", "要支援2", "要介護1", "要介護2", "要介護3",
             "要介護4", "要介護5"].index(do) if do != "合計" else 6
        return "{:,.0f}".format(t[i]) if t[i] else "―"
    return "―"



ws = sheet("10_ワーニングチェック",
           "将来推計ワーニングチェック結果の点検",
           "%s受領。施策反映の3画面（認定者数・施設居住系・在宅）は"
           "いずれも該当無しであり、"
           "ワーニングは参考として示される%d件に限られます。"
           "それぞれについて、当方の年報令和7年度の値と突き合わせ、"
           "手当てを要するかどうかを判定しました。"
           % (GW.META["出力日"], len(GW.MEISAI)),
           [4, 26, 13, 11, 11, 16, 13, 46])
r = 4
r = lead(ws, r, "1　サービス別の要約", span=8)
r = header(ws, r, ["#", "サービス", "件数", "区分", "", "要因", "", "見方"])
_SVC_N = {}
for _w in GW.MEISAI:
    _SVC_N.setdefault(_w[1], []).append(_w)
for _i, (_sv, _ws_) in enumerate(sorted(_SVC_N.items(),
                                        key=lambda x: -len(x[1])), start=1):
    _y = _YOIN.get(_sv, ("―", "―"))
    r = body(ws, r, [_i, _sv, len(_ws_), _ws_[0][0], "", _y[0], "", _y[1]],
             fills={6: (OK_G if _y[0] == "当区域の実態" else NG_O)},
             align={3: "center"}, height=40)
r = body(ws, r, ["", "計", len(GW.MEISAI), "", "", "", "", ""],
         fills={i: MID_B for i in range(1, 9)}, bold=True,
         align={3: "center"}, height=20)
r = note(ws, r,
         "注1）**ワーニングは推計を止めるものではありません。**"
         "全国の分布から定めた閾値との対比であり、"
         "外れていても理由が説明できれば進められます。"
         "施策反映の3画面（推計そのもの）は該当無しです。\n"
         + ("注2）**%d件のうち%d件（訪問介護）は当区域の実態です。**"
            % (len(GW.MEISAI), _W_JITTAI))
         + "当方の年報令和7年度でも要介護計55.4回／月で全国の上限50回を"
         "上回ります。手当ては要しません。\n"
         + ("注3）**残る%d件は令和8年度の実績見込み値と"
            % (len(GW.MEISAI) - _W_JITTAI))
         + "1人1月あたり給付費の年度によるものです**"
         "（確認事項No.128・No.137）。"
         "令和7年度の値に差し替えれば解消する見込みです。"
         "とくに地域密着型通所介護は月に31回を超える値になっており、"
         "そのままでは用いられません。",
         span=8, height=90)

r += 1
r = lead(ws, r, "2　ワーニングの明細（%d件）" % len(GW.MEISAI), span=8)
r = header(ws, r, ["#", "サービス", "区分", "年度", "要介護度",
                   "画面の値", "当方（年報R7）", "内容"])
for _i, (_ku, _sv, _yr, _do, _v, _lo, _hi, _msg) in enumerate(
        GW.MEISAI, start=1):
    r = body(ws, r, [_i, _sv, _ku.split("（")[0], _yr, _do, _v,
                     _ours(_sv, _do), _msg],
             align={6: "right", 7: "right"}, height=18)
r = note(ws, r,
         "注1）「当方（年報R7）」は要支援計・要介護計の値であり、"
         "要介護度別ではありません。"
         "当方は1人1月あたり回（日）数を要支援・要介護の2区分で"
         "固定しているためです。\n"
         "注2）特定施設入居者生活介護は1人1月あたり給付費（円）、"
         "そのほかは1人1月あたり利用回（日）数です。\n"
         "注3）認知症対応型通所介護は画面の値が0であり、"
         "下限を下回るものとして挙がっています。"
         "区域内に事業所はありませんが区域外の事業所で給付があり、"
         "当方は令和11年度に月1.05人と算定しています。",
         span=8, height=58)


# ============================================================ 11 自己点検
ws = sheet("11_自己点検", "自己点検",
           "本表の内的整合を機械で確かめた記録です。"
           "1件でも不適合があると出力そのものを止めます。",
           [6, 44, 34, 34, 10])

chk(1, "第9期の保険料基準額が当方の再現と1円以内で一致すること",
    "|見える化 − 当方の再現6,427.92| < 1",
    "見える化 %.4f円／当方 6,427.92円" % K.HOKENRYO["第9期"],
    abs(K.HOKENRYO["第9期"] - 6427.92) < 1)

_j = K.SHUNO["第1号被保険者負担分相当額"]["計"] \
    + K.SHUNO["調整交付金相当額"]["計"] - K.SHUNO["調整交付金見込額"]["計"]
chk(2, "保険料収納必要額が内訳から再現できること",
    "第1号負担分＋調整交付金相当額−調整交付金見込額",
    "再現 %.2f円／記載 %.2f円" % (_j, MI_J), abs(_j - MI_J) < 1)

_g = MI_J / 0.99 / HOSEI13 / 12
chk(3, "5シートの月額が所得段階別の入力と標準13段階の割合から再現できること",
    "J÷0.99÷補正後被保険者数÷12",
    "再現 %.4f円／記載 %.4f円" % (_g, MI_GETSU), abs(_g - MI_GETSU) < 0.01)

chk(4, "第1号被保険者負担分相当額が標準給付費見込額の23％であること",
    "（標準給付費見込額＋地域支援事業費）×0.23",
    "%.2f円／0.23倍は %.2f円"
    % (K.SHUNO["第1号被保険者負担分相当額"]["計"], (MI_A + MI_B) * 0.23),
    abs(K.SHUNO["第1号被保険者負担分相当額"]["計"] - (MI_A + MI_B) * 0.23) < 1)

_u = sum(v["計"] for v in K.HYOJUN_UCHI.values())
chk(5, "標準給付費見込額が内訳の和と一致すること",
    "総給付費＋特定入所者＋高額＋高額医療合算＋審査支払手数料",
    "内訳の和 %.0f円／記載 %.0f円" % (_u, MI_A), abs(_u - MI_A) < 1)

chk(6, "地域支援事業費が0であること（未入力であることの確認）",
    "5_保険料推計 地域支援事業費",
    "%.0f円" % MI_B, MI_B == 0)

_w = [y for y in YR if abs(K.WARIAI_KEI[y] - 1) > 1e-6]
chk(7, "加入割合の合計が1.0にならない年度を拾えていること",
    "|加入割合の合計 − 1| > 0.000001 の年度",
    "該当 %s" % ("・".join(_w) if _w else "なし"),
    _w == ["R12", "R17", "R22", "R27", "R32"])

_d3 = [y for y in ("R9", "R10", "R11") if K.DANKAI["計"][y] != K.HIHO[y]]
chk(8, "第10期の所得段階別の計が画面の第1号被保険者数と一致すること",
    "所得段階別の計 ＝ 画面の第1号被保険者数",
    "合わない年度 %s" % ("・".join(_d3) if _d3 else "なし"), not _d3)

_d8 = [y for y in ("R12", "R17", "R22", "R27", "R32")
       if K.DANKAI["計"][y] == K.HIHO[y]]
chk(9, "令和12年度以降の所得段階別の計が画面の第1号被保険者数と違うこと",
    "所得段階別の計 ≠ 画面の第1号被保険者数",
    "一致してしまう年度 %s" % ("・".join(_d8) if _d8 else "なし"), not _d8)

_ds = [y for y in ("R12", "R17", "R22", "R27", "R32")
       if K.DANKAI["計"][y] != sum(K.DANKAI["第%d段階" % (i + 1)][y]
                                   for i in range(13))]
chk(10, "所得段階別の段階ごとの和が計と一致すること",
    "Σ第1〜13段階 ＝ 計",
    "合わない年度 %s" % ("・".join(_ds) if _ds else "なし"), not _ds)

chk(11, "1シートと5シートの保険料基準額が違うこと（2通りあることの確認）",
    "1_推計値サマリ 第10期 ≠ 5_保険料推計 第10期",
    "%.4f円／%.4f円" % (MI_GETSU1, MI_GETSU),
    abs(MI_GETSU1 - MI_GETSU) > 1)

chk(12, "令和8年度の人数の欄が全て整数であること",
    "非ゼロの区分について小数部が0",
    "非ゼロ%d件／整数%d件" % (_tot8, _int8), _tot8 == _int8 and _tot8 > 0)

chk(13, "令和7年度の人数の欄に小数があること（月平均であることの確認）",
    "非ゼロの区分について小数部が0でないものがある",
    "非ゼロ%d件／小数%d件" % (_tot7, _dec7), _dec7 > 0)

chk(14, "令和8年度が令和7年度と±10％超ずれる区分を拾えていること",
    "|R8／R7 − 1| > 0.10",
    "%d件" % len(ZURE), len(ZURE) >= 10)

_zero = [nm for nm, r7, r8, rate, r11 in ZURE if r8 == 0]
chk(15, "令和8年度が0で令和11年度も0のままの区分があること",
    "R8＝0 かつ R11＝0",
    "%d件（%s）" % (len(_zero), "・".join(_zero)), len(_zero) >= 3)

chk(16, "設定の対照が受領した設定を残らず判定していること",
    "判定が「―」のものがないこと。要改の件数",
    "%d件／判定なし %d件／要改 %d件"
    % (len(K.SETTEI), sum(1 for h in _hantei if h == "―"),
       sum(1 for h in _hantei if h == "要改")),
    len(K.SETTEI) == 8 and "―" not in _hantei
    and sum(1 for h in _hantei if h == "要改") == 3)

# 固定値で突き合わせると算定を改めるたびに不適合になる（CLAUDE.md §4）。
# 算定が妥当な範囲にあることだけを見る。
chk(17, "当方の算定の月額が算定スクリプトの値と一致すること",
    "build による算定値をそのまま用いている（固定値で突き合わせない）",
    "%.2f円（保険料基準額 %d円）" % (TOU_GETSU, round(TOU_GETSU, -2)),
    6000 < TOU_GETSU < 7000)

chk(18, "弾力化にした場合の月額が標準13段階より低くなること",
    "条例16段階の乗率は標準より補正後被保険者数が大きくなる",
    "13段階 %.2f円／16段階 %.2f円" % (MI_GETSU, GETSU16),
    GETSU16 < MI_GETSU)

_pi = [v for v in K.META.values()]
chk(19.1, "画面の被保険者数の読み取りが受領した出力と一致すること",
    "6階級の和＝1_推計値サマリの第1号被保険者数",
    "第10期3か年 %s／受領 %s"
    % ([J.ichigo(y) for y in ("R9", "R10", "R11")],
       [K.HIHO[y] for y in ("R9", "R10", "R11")]),
    all(J.ichigo(y) == K.HIHO[y] for y in JY))

chk(19.2, "画面の第1号＋第2号が受領した出力の総数と一致すること",
    "第1号＋第2号（男女計）＝1_推計値サマリの総数",
    "R6 %d人" % (J.ichigo("R6") + J.dan("第2号", "R6")
                 + J.jo("第2号", "R6")),
    J.ichigo("R6") + J.dan("第2号", "R6") + J.jo("第2号", "R6") == 18386)

_ai = [y for y in JY if y not in JITSU
       and _i1(y) != _r0(sum(pop_juki(c, J.SEIREKI[y]) for c in CL3))]
chk(19.3, "案Cの入力値の和が当方の推計と一致すること（整数化で合計を保つ）",
    "男6階級＋女6階級の和 ＝ 案Cの65歳以上",
    "合わない年度 %s" % ("・".join(_ai) if _ai else "なし"), not _ai)

_r67 = [y for y in JITSU
        if ANC[y]["男"] != {a: J.dan(a, y) for a in J.AGE}
        or ANC[y]["総人口"] != J.SOJINKO[y]]
chk(19.4, "令和6年度・令和7年度は画面の値をそのまま用いていること",
    "実績年の入力値＝画面の値",
    "変えている年度 %s" % ("・".join(_r67) if _r67 else "なし"), not _r67)

chk(19.5, "令和11年度の第1号被保険者数が当方の算定と一致すること",
    "案Cの令和11年度 ＝ 9,213人",
    "%d人" % _i1("R11"), _i1("R11") == 9213)

chk(19.6, "後期高齢者の割合が画面より高くなること（案Cの趨勢による）",
    "75歳以上÷65歳以上を第10期3か年で比べる",
    "案C %s／画面 %s"
    % ([round(_kouki(ANC[y]) / _i1(y), 4) for y in ("R9", "R10", "R11")],
       [round(sum(J.kei(a, y) for a in J.AGE[2:]) / J.ichigo(y), 4)
        for y in ("R9", "R10", "R11")]),
    all(_kouki(ANC[y]) / _i1(y)
        > sum(J.kei(a, y) for a in J.AGE[2:]) / J.ichigo(y)
        for y in ("R9", "R10", "R11")))

_dk_ng = [(_y, _j + 1) for _y in GD.YEARS for _j in range(GD.NDAN)
          if GD.HIHO[_y][_j] != _DAN_TO[_DAN_YM[_y]][_j]]
chk(19.7, "画面に入力済みの16段階の人数が当方の値と一致すること",
    "弾力化②の画面の写し ＝ 所得段階別の将来推計（案C）",
    "全%d年度×%d段階＝%d件／違い %s"
    % (len(GD.YEARS), GD.NDAN, len(GD.YEARS) * GD.NDAN,
       _dk_ng or "なし"), not _dk_ng)

_kei_ng = [_y for _y in GD.YEARS if sum(GD.HIHO[_y]) != GD.KEI[_y]]
chk(19.8, "画面の段階別の和が画面の「計」の行と一致すること",
    "読み取りの検算（Σ第1〜16段階 ＝ 計）",
    "全%d年度／合わない %s" % (len(GD.YEARS), _kei_ng or "なし"),
    not _kei_ng)

_ks = [x[3] for x in STAGE16]
_ks_none = [x[0] for x in STAGE16 if x[3] is None]
_ks_val = [x for x in _ks if x is not None]
chk(19.9, "基準所得金額が第7段階から第16段階の10件で昇順であること",
    "第1〜6段階は欄なし／第7段階以降は金額かつ昇順",
    "欄なし 第%s段階／金額%d件（%s〜%s円）"
    % ("・".join(str(x) for x in _ks_none), len(_ks_val),
       "{:,}".format(_ks_val[0]), "{:,}".format(_ks_val[-1])),
    _ks_none == [1, 2, 3, 4, 5, 6] and len(_ks_val) == 10
    and all(a < b for a, b in zip(_ks_val, _ks_val[1:])))

_jor_sa = [x[0] for x in STAGE16
           if x[0] <= 13 and abs(x[2] - GAMEN13_JOR[x[0] - 1]) > 1e-9]
chk(19.95, "条例の割合が標準13段階と違う段階を拾えていること",
    "第1〜13段階で 条例の割合 ≠ 標準の割合 となる段階を数える",
    "違う %d段階（第%s段階）／標準にない 第14〜16段階"
    % (len(_jor_sa), "・".join(str(x) for x in _jor_sa)),
    len(_jor_sa) == 9)

_h_ng = [y for y in GD.YEARS if K2.HIHO[y] != sum(GD.HIHO[y])]
chk(21, "再出力の第1号被保険者数が当方の案Cと一致すること",
    "総括表 1．被保険者数 ＝ 所得段階別の将来推計の計",
    "全%d年度／合わない %s" % (len(GD.YEARS), _h_ng or "なし"), not _h_ng)

_fix = [("認定者数（第1号）", K.NINTEI["第1号"], K2.NINTEI_ICHI),
        ("総給付費（千円）", K.SOGAKU["合計"], K2.KYUFU)]
_fix_ng = [(nm, y) for nm, a, b in _fix for y in GD.YEARS if a[y] != b[y]]
chk(22, "人口の設定を改めても認定者数・総給付費が変わっていないこと",
    "改める前の出力と改めた後の出力を8年度で突き合わせる",
    "認定者数・総給付費とも全%d年度で同じ／違い %s"
    % (len(GD.YEARS), _fix_ng or "なし"), not _fix_ng)

_sa_j = K.SHUNO["保険料収納必要額"]["計"] - _J2
_sa_c = K2.SHUNO["調整交付金見込額"]["計"] - K.SHUNO["調整交付金見込額"]["計"]
chk(23, "保険料収納必要額の減少額が調整交付金見込額の増加額と一致すること",
    "給付費の側が動いていないことの裏づけ",
    "収納必要額 ▲%s円／調整交付金 +%s円"
    % ("{:,.0f}".format(_sa_j), "{:,.0f}".format(_sa_c)),
    abs(_sa_j - _sa_c) < 1)

chk(24, "標準13段階で再現した月額が出力の月額と1円以内で一致すること",
    "保険料収納必要額÷0.99÷補正後被保険者数÷12",
    "再現 %.2f円／出力 %.2f円" % (_G13, K2.HOKENRYO["第10期"]),
    abs(_G13 - K2.HOKENRYO["第10期"]) < 1)

_N_R6 = _quiet("build_shotoku_dankai.py")["N_R6"]
_k16 = [x / sum(_N_R6) for x in _N_R6]
_k13 = _k16[:12] + [sum(_k16[12:])]
_d13 = {y: _saidai(sum(GD.HIHO[y]), _k13) for y in GD.YEARS}
_reg_ng = [(y, i + 1) for y in GD.YEARS for i in range(13)
           if K2.DANKAI13["第%d段階" % (i + 1)][y] != _d13[y][i]]
chk(25, "画面に登録されている所得段階別が当方の標準13段階の値と一致すること",
    "総括表 5．所得段階別被保険者数 ＝ 08シート2（13段階へ束ねる案）",
    "全%d年度×13段階＝%d件／違い %s"
    % (len(GD.YEARS), len(GD.YEARS) * 13, _reg_ng or "なし"), not _reg_ng)

_g3 = [x for x in GW.GAIYO[:3] if x[1] != "該当無し"]
chk(26, "施策反映の3画面にワーニングがないこと",
    "認定者数・施設居住系・在宅の3画面の件数",
    "該当のあるもの %s" % (_g3 or "なし"), not _g3)

_wn = sum(x[1] for x in GW.GAIYO if isinstance(x[1], int))
chk(27, "ワーニングの明細の件数が概要の件数と一致すること",
    "概要の件数の和 ＝ 明細の行数",
    "概要 %d件／明細 %d件" % (_wn, len(GW.MEISAI)), _wn == len(GW.MEISAI))

_ast = []
for _w in wb.worksheets:
    for _row in _w.iter_rows():
        for _c in _row:
            if isinstance(_c.value, str) and "**" in _c.value:
                _ast.append("%s!%s" % (_w.title, _c.coordinate))
chk(19, "強調の指定（**）がセルに残っていないこと",
    "全シートの全セルを走査（xlsx は Markdown を解釈しない）",
    "%d件" % len(_ast), not _ast)

import re as _re
_PI = [_re.compile(r"\d{2,4}-\d{2,4}-\d{3,4}"),
       _re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}")]
_pi = []
for _w in wb.worksheets:
    for _row in _w.iter_rows():
        for _c in _row:
            if isinstance(_c.value, str):
                for _p in _PI:
                    if _p.search(_c.value):
                        _pi.append("%s!%s" % (_w.title, _c.coordinate))
chk(20, "個人を特定する値を収めていないこと",
    "電話番号・メールアドレスの形を全セルで走査",
    "%d件" % len(_pi), not _pi)

r = header(ws, 4, ["No.", "点検した内容", "式・条件", "結果", "判定"])
for c in CHECKS:
    body(ws, r, list(c), fills={5: (OK_G if c[4] == "適合" else NG_O)},
         height=30)
    r += 1
_NG = sum(1 for c in CHECKS if c[4] != "適合")
body(ws, r, ["", "自己点検の結果", "",
             "適合%d件・不適合%d件" % (len(CHECKS) - _NG, _NG),
             "適合" if _NG == 0 else "不適合"],
     fills={5: (OK_G if _NG == 0 else NG_O)}, height=24, bold=True)
r += 1
note(ws, r,
     "注）本シートが確かめているのは受領した推計結果の内的整合と、"
     "本表がそれを正しく読み取れていることです。"
     "推計そのものの当否を確かめるものではありません。",
     span=5, height=32)


# ============================================================ 出力
for ws_ in wb.worksheets:
    ws_.oddFooter.left.text = "見える化システムの推計結果の受領点検（%s）" \
        % KIJUNBI
    ws_.oddFooter.right.text = "&P / &N"

wb.save(OUT)
print("書き出しました:", OUT)
for ws_ in wb.worksheets:
    print("  -", ws_.title, ws_.max_row, "rows")
print("見える化 %s円／当方 %s円（第10期の算定上の月額）"
      % ("{:,.2f}".format(MI_GETSU), "{:,.2f}".format(TOU_GETSU)))
print("修正が必要な箇所 %d件（優先A %d件）"
      % (len(SHUSEI), sum(1 for x in SHUSEI if x[1] == "A")))
print("令和8年度が令和7年度と±10％超ずれる区分 %d件（うち0になるもの %d件）"
      % (len(ZURE), len(_zero)))
print("自己点検 %d件：適合%d件・不適合%d件"
      % (len(CHECKS), len(CHECKS) - _NG, _NG))
if _NG:
    print("不適合があります。")
    sys.exit(1)
print("すべての点検に適合しました。")
