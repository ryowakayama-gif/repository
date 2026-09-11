# -*- coding: utf-8 -*-
"""リポジトリ内のパス.

ビルドスクリプトの出力先を、スクリプト自身の置かれている場所から求める。

従前は "/home/user/repository/output" を絶対パスで固定していたため、
同じ計算機の別の場所（別ブランチの作業ツリーなど）でスクリプトを実行しても、
出力は常に /home/user/repository/output へ書き込まれていた。
別の案件と並行して作業する場合に、意図しない場所へ生成物が落ちる。

ROOT はこのファイルの置かれているディレクトリであり、
スクリプトはいずれも同じディレクトリにあるため、
どこから実行してもそのスクリプトのリポジトリの output へ書き込む。
"""

import os

ROOT = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(ROOT, "output")
SOURCE = os.path.join(ROOT, "source")


DRAFT = os.path.join(OUTPUT, "第10期介護保険事業計画_協議用素案_令和8年8月.docx")


def draft_size():
    """計画素案の段落数・表数・図数を実物から数える。

    件数を固定値で書くと、素案を改訂したときに管理表の記載とずれる。
    （令和8年9月11日の点検で、管理表の667段落114表34図・620段落108表34図が
      実物と合っていないことが分かった。）

    図は本文に埋め込んだ画像の数を数える。
    素案が生成されていない場合は (0, 0, 0) を返す。
    """
    if not os.path.exists(DRAFT):
        return (0, 0, 0)
    from docx import Document
    d = Document(DRAFT)
    zu = len(d.inline_shapes)
    return (len(d.paragraphs), len(d.tables), zu)


def draft_label():
    """「806段落130表36図」の形の文字列を返す。"""
    p, t, z = draft_size()
    return "%d段落%d表%d図" % (p, t, z)
