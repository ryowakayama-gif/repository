# -*- coding: utf-8 -*-
"""
成果物の再現性点検

目的
  金ヶ崎町の案件で「本文を組み立てる生成器がスクラップパッドにしか無く、
  環境が作り直された時点で消えた」という事象が起きた。北塩原村で同じことが
  起きないことを、人の記憶ではなく機械で確かめるための点検である。

  要点は一つで、成果物そのものではなく「成果物を作り直せる素」が
  すべてリポジトリに追跡されているかを見る。追跡されていれば環境が
  消えても作り直せる。追跡されていなければ、作り直す道が無い。

点検項目（既定）
  C-0  必要なライブラリが入っている
                                   欠けると成果物が静かに欠落する
  C-1  未追跡ファイルが無い        作業ツリーにだけ在るものを許さない
  C-2  Python が全て追跡済み        生成器が1本でも未追跡なら作り直せない
  C-3  成果物と生成器が対応する    生成器の無い成果物・成果物の無い生成器を検出
  C-4  リポジトリ内の絶対パスを直書きしていない
                                   別の場所に展開すると出力先が狂う
  C-5  外部資料に入手先の記録がある
                                   容量の都合で格納しない資料は入手先が要る

点検項目（--regenerate を付けたとき）
  C-6  全生成器を実行し、HEAD と内容署名が一致する

  docx・xlsx は ZIP であるため、作り直すと生成時刻の違いでバイト列は必ず
  変わる。よって一致はバイトではなく内容（段落・表・セルの値）で判定する。

使い方
  python3 check_reproducibility.py              静的点検のみ（数秒）
  python3 check_reproducibility.py --regenerate 作り直して内容を照合（数分）

  いずれも、ひとつでも不合格があれば終了コード 1 を返す。
"""

import hashlib
import os
import subprocess
import sys
import zipfile

REPO_ROOT = os.path.dirname(os.path.abspath(__file__))

# ============================================================
# 生成器と成果物の対応（依存の順に並べる＝作り直しの手順書でもある）
# ============================================================
BUILDERS = [
    # (生成器, 成果物, 備考)
    ("build_excel.py", [
        "output/00_全計画マスター管理表.xlsx",
        "output/01_共通_基本コラム部品.xlsx",
        "output/02_高齢者介護保険事業計画.xlsx",
        "output/03_障がい福祉計画.xlsx",
        "output/04_こども計画.xlsx",
    ], "全計画の管理表。最初に作る"),
    ("build_component_images.py", [
        "output/images_basic/BC-01_ポイント.png",
        "output/images_basic/BC-02_コラム.png",
        "output/images_basic/BC-03_事例紹介.png",
        "output/images_basic/BC-04_解説.png",
        "output/images_basic/BC-05_データの見方.png",
        "output/images_basic/BC-06_注意・留意点.png",
    ], "build_excel.py の成果物に画像を貼るため、その後に実行する"),
    ("build_kitashiobara_village_data.py", ["output/北塩原村_村提供実績データ.xlsx"], ""),
    ("build_kitashiobara_projection.py", ["output/北塩原村_将来推計.xlsx"], ""),
    ("build_kitashiobara_evaluation.py", ["output/北塩原村_現行計画評価.xlsx"], ""),
    ("build_kitashiobara_service_estimate.py", ["output/北塩原村_サービス見込量.xlsx"], ""),
    ("build_kitashiobara_katsudou.py", ["output/北塩原村_活動指標.xlsx"], ""),
    ("build_kitashiobara_chiiki.py", ["output/北塩原村_地域生活支援事業.xlsx"], ""),
    ("build_kitashiobara_shigen.py", ["output/北塩原村_圏域サービス資源.xlsx"], ""),
    ("build_kitashiobara_finance.py", ["output/北塩原村_財源構成案.xlsx"], ""),
    ("build_kitashiobara_cross.py", ["output/北塩原村_クロス集計設計.xlsx"], ""),
    ("build_kitashiobara_progress.py", ["output/北塩原村_業務進捗管理.xlsx"], ""),
    ("build_kitashiobara_mikomiryo.py", ["output/北塩原村_サービス見込量算定.xlsx"],
     "将来推計の計算を runpy で取り込むため、推計の後に実行する"),
    ("build_kitashiobara_survey_check.py",
     ["output/北塩原村_アンケート結果報告書_点検結果.xlsx"], ""),
    ("build_kitashiobara_kpi_hikaku.py",
     ["output/北塩原村_目標KPI_前回比較可能性.xlsx"], ""),
    ("build_kitashiobara_kosshi_rev.py", ["output/北塩原村_計画素案.docx"],
     "計画素案の本文。原本 source/北塩原村_骨子案_原本_20260731.docx から組む"),
    ("build_kitashiobara_mokuji.py", ["output/北塩原村_計画素案.docx"],
     "計画素案に目次を入れる（同じファイルを上書きするため素案の後）"),
    ("build_kitashiobara_review.py",
     ["output/北塩原村_骨子案レビュー_基本指針網羅性.xlsx"],
     "計画素案を読んで網羅性を見るため、素案の後に実行する"),
    ("build_kitashiobara_kaigi.py", [
        "output/会議資料/01_協議会資料1_計画策定の概要.docx",
        "output/会議資料/02_協議会資料2_国の基本指針の改正点.docx",
        "output/会議資料/03_協議会資料8_提供体制の課題.docx",
        "output/会議資料/04_庁議資料2_関連計画との整合.docx",
        "output/会議資料/05_庁議資料4_庁内各課への依頼事項.docx",
    ], ""),
    ("build_kitashiobara_mece.py",
     ["output/北塩原村_見込量算定_MECEチェック.xlsx"],
     "計画素案と地域生活支援事業の生成器を読んで指摘の対応状況を判定するため、"
     "素案・目次の後に実行する"),
    ("build_kitashiobara_houkokusho_review.py",
     ["output/北塩原村_アンケート分析報告書_レビュー.xlsx"],
     "計画素案と見込量算定MECEチェックの生成器を読んで当方の修正の反映を"
     "確かめるため、両者の後に実行する"),
    ("build_kitashiobara_houkokusho_kousei.py",
     ["output/北塩原村_アンケート報告書_構成案.xlsx"],
     "計画素案を読んで人材確保の記述の有無を確かめるため、素案の後に実行する"),
    ("build_kitashiobara_kenkeikaku.py",
     ["output/北塩原村_福島県計画との整合.xlsx"],
     "計画素案を読んで県計画の確認結果の反映を確かめるため、素案の後に実行する"),
    ("build_kitashiobara_shuusei_rireki.py",
     ["output/北塩原村_計画素案_修正履歴.docx"],
     "原本・計画素案・素案の生成器を読んで差分を出すため、素案の後に実行する"),
    ("build_kitashiobara_kaigo_seigo.py",
     ["output/北塩原村_障がい計画と介護保険事業計画の整合.xlsx"],
     "計画素案と介護保険事業計画の素案を読んで切り分けを見るため、"
     "素案の後に実行する"),
    ("build_kitashiobara_review_taiou.py",
     ["output/北塩原村_レビュー対応と素案修正方針.xlsx"],
     "村資料の実績と第1次概算を突き合わせるため、見込量算定の後に実行する"),
    ("build_kitashiobara_member_soan_check.py",
     ["output/北塩原村_他メンバー素案_網羅性点検.xlsx"],
     "他メンバーの素案と告示本文・村資料を読んで網羅性を見るため、"
     "当方の整理が揃った後に実行する"),
    ("build_kitashiobara_review2_taiou.py",
     ["output/北塩原村_レビュー対応_第2回.xlsx"],
     "第2回レビューの反映を計画素案の現物で照合するため、素案の後に実行する"),
    ("build_kitashiobara_uchiawase.py",
     ["output/北塩原村_打合せ記録_20261009.xlsx"],
     "加筆版と当方の網羅性点検の差分を取るため、網羅性点検の後に実行する"),
    ("build_kitashiobara_ronten.py",
     ["output/北塩原村_論点整理_20261009.xlsx"],
     "論点メモ・当方の素案・正本の3つを突き合わせるため、素案の後に実行する"),
    ("build_kitashiobara_ishoku_okurijo.py",
     ["output/北塩原村_申し送り_他メンバー版への移植_20261009.docx"],
     "挿入位置の引用が他メンバー版に実在することを照合するため、"
     "他メンバー版の格納後に実行する"),
    ("build_kitashiobara_graph.py",
     ["output/図表/01_給付費の推移.png",
      "output/図表/02_サービス別の給付費.png",
      "output/図表/03_高齢化率の推移.png",
      "output/図表/04_計画の体系.png",
      "output/図表/05_関係機関との連携.png",
      "output/図表/06_地域生活支援拠点の5機能.png",
      "output/図表/07_緊急時の4場面.png",
      "output/図表/08_65歳到達時の判定の流れ.png",
      "output/図表/09_ライフコースと制度の移行.png",
      "output/図表/10_両計画の関係.png",
      "output/図表/11_ライフコース_概要版.png"],
     "申し送りの表と正本の表から数値を読んで図を描くため、"
     "申し送りの後・正本の前に実行する"),
    ("build_kitashiobara_honpon.py",
     ["output/北塩原村_計画素案_正本_移植後.docx"],
     "申し送りの生成器から文章を読み込んで正本に入れるため、申し送りの後に実行する"),
    ("build_kitashiobara_murashiryo.py",
     ["output/北塩原村_村資料点検_20261010.xlsx"],
     "村資料にあって正本にない事実を正本の現物で照合するため、"
     "移植後の正本を作った後に実行する"),
    ("build_kitashiobara_zuhyo.py", ["output/北塩原村_図表一覧.xlsx"],
     "正本の図表を現物から数え上げるため、移植後の正本を作った後に実行する"),
    ("build_kitashiobara_hp_shisaku.py",
     ["output/北塩原村_村HP掲載施策の棚卸し.xlsx"],
     "村ホームページの施策が正本に入ったかを現物で照合するため、"
     "移植後の正本を作った後に実行する"),
    ("build_kitashiobara_hp_3bunya.py",
     ["output/北塩原村_村HP3分野の棚卸し.xlsx"],
     "高齢者・児童・社会福祉の施策が正本に入ったかを現物で照合するため、"
     "移植後の正本を作った後に実行する"),
    ("build_kitashiobara_zenkai_hyouka.py",
     ["output/北塩原村_前回計画の評価整理.xlsx"],
     "前回計画の評価が正本の表と一致するかを現物で照合するため、"
     "移植後の正本を作った後に実行する"),
    ("build_kitashiobara_anke_mikomi.py",
     ["output/北塩原村_アンケートと見込量の整合.xlsx"],
     "正本の調査結果の記述と割合が合うかを現物で照合するため、"
     "移植後の正本を作った後に実行する"),
    ("build_kitashiobara_mikomi_redteam.py",
     ["output/北塩原村_見込量算定のRedTeam再レビュー.xlsx"],
     "直した定義が入れ替わったかを正本と告示の現物で照合するため、"
     "移植後の正本を作った後に実行する"),
    ("build_kitashiobara_mece_houshu.py",
     ["output/北塩原村_MECE_利用の流れと報酬改定.xlsx"],
     "直した表現が入れ替わったか、足した記述が入ったかを現物で照合するため、"
     "移植後の正本を作った後に実行する"),
    ("build_kitashiobara_gaiyou.py", ["output/北塩原村_計画素案_概要版.docx"],
     "令和8年10月に加えた3節（方策・年齢到達・他計画との連携）は正本に"
     "しかないため、その現物と突き合わせる。移植後の正本と図の後に実行する"),
    ("build_kitashiobara_taihi.py", ["output/北塩原村_目次構成_前回対比.xlsx"],
     "計画素案と概要版を読むため、両方の後に実行する"),
    ("build_kitashiobara_zuhyo_bangou.py",
     ["output/北塩原村_図表番号一覧.xlsx"],
     "付けた番号・表題が正本の図表と1対1に対応するかを現物で照合するため、"
     "移植後の正本を作った後に実行する"),
    ("build_kitashiobara_redteam_kyukyu.py",
     ["output/北塩原村_RedTeam_救急防災冬季.xlsx"],
     "直した表現が入れ替わったか、足した記述が入ったかを現物で照合するため、"
     "移植後の正本を作った後に実行する"),
    ("build_kitashiobara_juten12.py",
     ["output/北塩原村_12重点施策_評価設計.xlsx"],
     "12重点施策の接続を正本の現物で照合するため、"
     "移植後の正本を作った後に実行する"),
    ("build_kitashiobara_kouiki_redteam.py",
     ["output/北塩原村_広域連携_RedTeam再レビュー.xlsx"],
     "直した表現が入れ替わったか、足した記述と表が入ったかを現物で照合する"
     "ため、移植後の正本を作った後に実行する"),
    ("build_kitashiobara_shoukai_bun.py",
     ["output/北塩原村_村への照会_6通.docx"],
     "各点検・レビューの照会を読み込んで組むため、それらの後に実行する"),
    ("build_shinchoku_hyo.py", ["output/北塩原村_業務進捗管理表.xlsx"],
     "build_kitashiobara_progress.py のデータを取り込むため、その後に実行する"),
]

# 外部資料を入力とする生成器（容量の都合でリポジトリに格納しないもの）。
# 入手先の記録が生成器本体にあることを C-5 で確かめる。成果物は追跡済みなので、
# 原典が手元に無くても計画本文の執筆は止まらない。
EXTERNAL_INPUT = {
    "build_fukushima_seishin_extract.py": {
        "outputs": ["output/福島県_精神医療指標抽出.xlsx"],
        "why": "NDB集計が1ファイル20MB前後あるため格納しない",
        "needs": ["http"],
    },
}

# 生成器ではない共有モジュール
MODULES = ["kitashiobara_common.py", "check_reproducibility.py"]

# 生成に必要な外部ライブラリ。import 名と requirements.txt での名を対にする。
# PIL が欠けると build_component_images.py だけが落ち、00_・01_ の画像シートが
# 欠けた成果物になる。気づきにくいため、実行前に必ず確かめる。
REQUIRED_LIBS = [
    ("openpyxl", "openpyxl"),
    ("docx", "python-docx"),
    ("PIL", "Pillow"),
    ("pypdf", "pypdf"),
    ("pdfplumber", "pdfplumber"),
]
REQUIREMENTS = "requirements.txt"


# ============================================================
# git
# ============================================================
def git(*args):
    return subprocess.run(["git", "-C", REPO_ROOT, *args],
                          capture_output=True, text=True, check=True).stdout


def tracked_files():
    """追跡済みファイルの一覧。git は非ASCIIのパスを \\xxx に直すため -z で読む。"""
    out = subprocess.run(["git", "-C", REPO_ROOT, "ls-files", "-z"],
                         capture_output=True, check=True).stdout
    return set(p.decode("utf-8") for p in out.split(b"\0") if p)


def untracked_files():
    out = subprocess.run(["git", "-C", REPO_ROOT, "status", "--porcelain", "-z",
                          "--untracked-files=all"],
                         capture_output=True, check=True).stdout
    paths = []
    for rec in out.split(b"\0"):
        if rec.startswith(b"?? "):
            paths.append(rec[3:].decode("utf-8"))
    return paths


def head_blob(path):
    r = subprocess.run(["git", "-C", REPO_ROOT, "show", f"HEAD:{path}"],
                       capture_output=True)
    return r.stdout if r.returncode == 0 else None


# ============================================================
# 内容署名（ZIP の生成時刻に左右されない形で中身を要約する）
# ============================================================
def sig_docx(data):
    """(段落数, 表数, 本文テキストのSHA-256先頭12桁)"""
    import io
    import re
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    paras = re.findall(r"<w:p[ >].*?</w:p>|<w:p/>", xml, re.S)
    tables = len(re.findall(r"<w:tbl>", xml))
    texts = re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml, re.S)
    h = hashlib.sha256("".join(texts).encode("utf-8")).hexdigest()[:12]
    return (len(paras), tables, h)


def sig_xlsx(data):
    """(シート数, 値のあるセル数, 値のSHA-256先頭12桁)"""
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    h = hashlib.sha256()
    cells = 0
    for ws in wb.worksheets:
        h.update(ws.title.encode("utf-8"))
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    cells += 1
                    h.update(f"{c.coordinate}={c.value}".encode("utf-8"))
    n = len(wb.worksheets)
    wb.close()
    return (n, cells, h.hexdigest()[:12])


def sig_bytes(data):
    return ("-", len(data), hashlib.sha256(data).hexdigest()[:12])


def signature(path, data):
    if path.endswith(".docx"):
        return sig_docx(data)
    if path.endswith(".xlsx"):
        return sig_xlsx(data)
    return sig_bytes(data)


# ============================================================
# 点検
# ============================================================
class Report:
    def __init__(self):
        self.fails = []

    def check(self, tag, title, bad, detail=""):
        mark = "OK  " if not bad else "NG  "
        print(f"{mark}{tag}  {title}")
        if detail:
            print(f"      {detail}")
        for b in bad:
            print(f"      - {b}")
        if bad:
            self.fails.append(tag)
        return not bad


def c0_libs(rep, tracked):
    import importlib.util
    bad = []
    if REQUIREMENTS not in tracked:
        bad.append(f"{REQUIREMENTS} が追跡されていない（環境を作り直せない）")
    req = ""
    if os.path.exists(os.path.join(REPO_ROOT, REQUIREMENTS)):
        req = open(os.path.join(REPO_ROOT, REQUIREMENTS), encoding="utf-8").read()
    for mod, pkg in REQUIRED_LIBS:
        if importlib.util.find_spec(mod) is None:
            bad.append(f"{pkg}（import {mod}）が入っていない："
                       f"pip install -r {REQUIREMENTS} を実行する")
        if pkg.lower() not in req.lower():
            bad.append(f"{pkg} が {REQUIREMENTS} に書かれていない")
    return rep.check("C-0", "必要なライブラリが入っている", bad,
                     "／".join(p for _, p in REQUIRED_LIBS))


def c1_untracked(rep):
    bad = untracked_files()
    return rep.check("C-1", "未追跡ファイルが無い", bad,
                     f"未追跡 {len(bad)}件" if bad else "未追跡 0件")


def c2_python_tracked(rep, tracked):
    here = sorted(f for f in os.listdir(REPO_ROOT) if f.endswith(".py"))
    bad = [f for f in here if f not in tracked]
    return rep.check("C-2", "Python が全て追跡済み", bad,
                     f"{len(here)}本を点検（生成器{len(BUILDERS) + len(EXTERNAL_INPUT)}本"
                     f"＋共有{len(here) - len(BUILDERS) - len(EXTERNAL_INPUT)}本）")


def c3_pairs(rep, tracked):
    declared = {}
    for script, outs, _ in BUILDERS:
        for o in outs:
            declared.setdefault(o, []).append(script)
    for script, spec in EXTERNAL_INPUT.items():
        for o in spec["outputs"]:
            declared.setdefault(o, []).append(script)

    bad = []
    for script in list(dict.fromkeys([s for s, _, _ in BUILDERS])) + list(EXTERNAL_INPUT):
        if script not in tracked:
            bad.append(f"生成器が未追跡: {script}")
        if not os.path.exists(os.path.join(REPO_ROOT, script)):
            bad.append(f"生成器が存在しない: {script}")
    for out, scripts in sorted(declared.items()):
        if out not in tracked:
            bad.append(f"成果物が未追跡: {out}")
        elif not os.path.exists(os.path.join(REPO_ROOT, out)):
            bad.append(f"成果物が存在しない: {out}")

    # 追跡されている成果物に生成器が無いもの（手作業で置いた＝作り直せないもの）
    for p in sorted(tracked):
        if not p.startswith("output/"):
            continue
        if p not in declared:
            bad.append(f"生成器の対応が無い成果物: {p}")

    # 生成器なのに BUILDERS/EXTERNAL_INPUT に載っていないもの
    listed = set(s for s, _, _ in BUILDERS) | set(EXTERNAL_INPUT)
    for f in sorted(os.listdir(REPO_ROOT)):
        if f.startswith("build_") and f.endswith(".py") and f not in listed:
            bad.append(f"対応表に未記載の生成器: {f}")

    # 複数の生成器が同じ成果物に書くもの（後の生成器が前の出力を加工する）。
    # 禁じるのではなく、順序が効くことを毎回表に出す。
    # 前の生成器だけを単独で走らせると、後の加工が消えたまま残る。
    kasanari = []
    for out, scripts in sorted(declared.items()):
        if len(scripts) > 1:
            kasanari.append(f"{out} ← {' → '.join(scripts)}")

    nokori = f"生成器 {len(listed)}本 ／ 成果物 {len(declared)}点"
    if kasanari:
        nokori += ("　※順序が効く成果物（前の生成器だけを単独で走らせない）: "
                   + "／".join(kasanari))
    return rep.check("C-3", "成果物と生成器が対応する", bad, nokori)


def c4_abs_paths(rep):
    """リポジトリ内を指す絶対パスの直書きを禁じる（別の場所に展開すると狂う）。"""
    bad = []
    for f in sorted(os.listdir(REPO_ROOT)):
        if not f.endswith(".py"):
            continue
        with open(os.path.join(REPO_ROOT, f), encoding="utf-8") as fh:
            for i, line in enumerate(fh, 1):
                if f'"{REPO_ROOT}' in line or f"'{REPO_ROOT}" in line:
                    bad.append(f"{f}:{i} {line.strip()[:70]}")
    return rep.check("C-4", "リポジトリ内の絶対パスを直書きしていない", bad,
                     f"基準は __file__ から求める（現在 {REPO_ROOT}）")


def c5_external(rep, tracked):
    bad = []
    for script, spec in EXTERNAL_INPUT.items():
        src = open(os.path.join(REPO_ROOT, script), encoding="utf-8").read()
        for need in spec["needs"]:
            if need not in src:
                bad.append(f"{script} に入手先（{need}）の記録が無い")
        for o in spec["outputs"]:
            if o not in tracked:
                bad.append(f"{script} の成果物が未追跡: {o}（原典も手元も無くなる）")
    return rep.check("C-5", "外部資料に入手先の記録がある", bad,
                     "／".join(f"{k}：{v['why']}" for k, v in EXTERNAL_INPUT.items()))


def c6_regenerate(rep, tracked):
    """全生成器を実行し、HEAD と内容署名を照合する。"""
    print()
    print("  作り直しを実行します（数分かかります）")
    ran, failed = [], []
    for script, _, _ in BUILDERS:
        r = subprocess.run([sys.executable, script], cwd=REPO_ROOT,
                           capture_output=True, text=True)
        if r.returncode != 0:
            failed.append(f"{script} が終了コード {r.returncode}："
                          f"{(r.stderr or r.stdout).strip().splitlines()[-1][:80]}")
        else:
            ran.append(script)
        print(f"    {'OK ' if r.returncode == 0 else 'NG '} {script}")
    rep.check("C-6a", "全生成器が正常終了する", failed,
              f"実行 {len(ran)}/{len(BUILDERS)}本")

    targets = sorted({o for _, outs, _ in BUILDERS for o in outs} & tracked)
    bad, same = [], 0
    for p in targets:
        blob = head_blob(p)
        if blob is None:
            bad.append(f"HEAD に無い: {p}")
            continue
        with open(os.path.join(REPO_ROOT, p), "rb") as fh:
            now = fh.read()
        s_head, s_now = signature(p, blob), signature(p, now)
        if s_head == s_now:
            same += 1
        else:
            bad.append(f"{p}  HEAD={s_head}  再生成={s_now}")
    return rep.check("C-6b", "再生成した内容が HEAD と一致する", bad,
                     f"一致 {same}/{len(targets)}点"
                     "（バイト差はZIP内の生成時刻によるもので内容には影響しない）")


def main():
    regen = "--regenerate" in sys.argv[1:]
    tracked = tracked_files()

    print("=" * 72)
    print("  北塩原村　成果物の再現性点検")
    print(f"  リポジトリ: {REPO_ROOT}")
    print(f"  HEAD: {git('rev-parse', '--short', 'HEAD').strip()}"
          f"  追跡 {len(tracked)}件")
    print("=" * 72)

    rep = Report()
    c0_libs(rep, tracked)
    c1_untracked(rep)
    c2_python_tracked(rep, tracked)
    c3_pairs(rep, tracked)
    c4_abs_paths(rep)
    c5_external(rep, tracked)
    if regen:
        c6_regenerate(rep, tracked)
    else:
        print("--  C-6  再生成による内容照合は省略（--regenerate を付けると実行）")

    print("=" * 72)
    if rep.fails:
        print(f"  不合格 {len(rep.fails)}件: {' '.join(rep.fails)}")
        print("=" * 72)
        return 1
    print("  すべて合格。環境が作り直されても、追跡済みのファイルだけで"
          "成果物を復元できます。")
    print("=" * 72)
    return 0


if __name__ == "__main__":
    sys.exit(main())
