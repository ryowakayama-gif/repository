# -*- coding: utf-8 -*-
"""大雪地区広域連合 第10期介護保険事業計画　3町合同意見交換会 資料.

令和8年10月1日のご依頼
  「意見交換会の資料を作成して下さい」

業務仕様書５の令和8年10月の工程（地域課題整理、将来推計、
第10期計画骨子案作成、構成町との協議・調整）のうち、
3町合同の意見交換会（令和8年10月上旬）で用いる資料である。
町ごとの個別協議（10月中旬〜下旬）は
「10月 地域課題の整理と3町協議の確認事項」の03〜05シートによる。

━━ 本資料の考え方 ━━

1 ご決定をお願いする事項には、すべて「決まらない場合の当方の扱い」を
  置いている。したがって第2次概算（10月30日）は決定を待たずに組める。
  置いているのは受託者の仮置きであり、本意見交換会で内容を諮る。

2 保険料の月額を動かす前提は、国の告示・公布によるものと
  広域連合のご判断によるものに分けて示す。

3 地域の課題は、根拠となる数値・出所・計画本文の反映先と対にして示す。

━━ 数値の出どころ ━━

見込量・給付費・保険料はサービス見込量の算定（第1次概算）による。
確認事項とその扱い、資料のご提供のお願いは業務工程管理表による。
地域の課題は「10月 地域課題の整理と3町協議の確認事項」による。

構成
  第1節　意見交換会の次第
  第2節　第10期計画の策定の状況
  第3節　サービス見込量・給付費・保険料の試算
  第4節　保険料の月額を動かす前提
  第5節　地域の課題
  第6節　ご決定をお願いする事項（3町共通）
  第7節　認知症施策推進計画の扱い
  第8節　町ごとの個別協議でお伺いする事項
  第9節　資料のご提供をお願いしたいもの
  第10節　10月以降の進め方

出力
  output/第10期計画_3町合同意見交換会資料.docx

自己点検で1件でも不適合があると終了コード1で終わる。
"""

import docx_fix
import ast
import contextlib
import io
import os
import re
import runpy
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import data_kitei as DK                                   # noqa: E402
import repo_paths as RP                                   # noqa: E402

OUT = RP.ROOT + "/output/第10期計画_3町合同意見交換会資料.docx"
KADAI_XLSX = (RP.ROOT + "/output/"
              "第10期計画_10月 地域課題の整理と3町協議の確認事項.xlsx")
KIJUNBI = "令和8年10月1日"
TOWNS = ["東川町", "美瑛町", "東神楽町"]
FONT = "游ゴシック"
NAVY = RGBColor(0x1F, 0x38, 0x64)

CHECKS = []


def chk(no, naiyo, kekka, ok):
    CHECKS.append((no, naiyo, kekka, "適合" if ok else "不適合"))
    return ok


# ==================================================== 算定と台帳を読む
def _load(name):
    with contextlib.redirect_stdout(io.StringIO()):
        return runpy.run_path(RP.ROOT + "/" + name)


_M = _load("build_mikomiryo_santei.py")          # サービス見込量の算定
_NASHI_SRC = open(RP.ROOT + "/build_mikomi_juryo_nashi.py",
                  encoding="utf-8").read()

SVC = _M["SVC"]
Y3, Y3L = _M["Y3"], _M["Y3L"]
YLONG_HON = _M["YLONG_HON"]
KYUFU3, KYUFU_Y = _M["KYUFU3"], _M["KYUFU_Y"]
G3 = _M["G_DO"]
G_ITTEI, G_GENKO = _M["G_ITTEI"], _M["G_GENKO"]
SUEOKI = _M["SUEOKI"]
NIN, HIHO = _M["NIN"], _M["HIHO"]
KAISU_MAP = _M["KAISU_MAP"]
YLONG, YLONGL = _M["YLONG"], _M["YLONGL"]
# 計画本文に掲げる中長期の年度（令和17年度・令和22年度）の表示名
YL_HON_L = [YLONGL[YLONG.index(y)] for y in YLONG_HON]


def nin(y):
    """認定者数（要介護度別の和）。"""
    return sum(NIN[y])
SOGO_RYO, KAYOI_RYO = _M["SOGO_RYO"], _M["KAYOI_RYO"]
GETSU = G3["月額"]
SOU3 = sum(KYUFU3)
KIJUN = int(round(GETSU / 100.0)) * 100          # 百円未満四捨五入


def _literal(src, name):
    for node in ast.parse(src).body:
        if (isinstance(node, ast.Assign)
                and getattr(node.targets[0], "id", None) == name):
            return ast.literal_eval(node.value)
    raise RuntimeError("%s を読めない" % name)


# 据え置きの区分（Z＝受託者の側で確定したもの／A＝確定できるもの／
# B＝発注者・3町の資料待ち／C＝国の告示・公布待ち）。
KUBUN = _literal(_NASHI_SRC, "KUBUN")
NZ = sum(1 for v in KUBUN.values() if v == "Z")
NA = sum(1 for v in KUBUN.values() if v == "A")
NB = sum(1 for v in KUBUN.values() if v == "B")
NC = sum(1 for v in KUBUN.values() if v == "C")

_PC_SRC = open(RP.ROOT + "/build_process_control.py", encoding="utf-8").read()
CHECK = _literal(_PC_SRC, "CHECK")
LACK = _literal(_PC_SRC, "LACK")
BY = {c[0]: c for c in CHECK}
# 列：No./業務内容/工程/確認事項/確認をお願いする内容/止めている成果物/
#     確認先/状態/回答期限/回答・対応
SUMI = ("完了", "了承済", "了承済（保管せず廃棄）", "代替により解消", "解決")
KITEI = DK.all_kitei()

MIKETSU = [c for c in CHECK if c[7] not in SUMI]
MACHI = [c for c in MIKETSU if "町" in c[6]]
KYOTSU = [c for c in MACHI if not any(t in c[6] for t in TOWNS)]
KOBETSU = {t: [c for c in MACHI if t in c[6]] for t in TOWNS}


def tome(c):
    """止めている成果物（概算・保険料・計画素案）。"""
    s = str(c[5])
    out = []
    if "概算" in s:
        out.append("概算")
    if "保険料" in s or "将来推計" in s:
        out.append("保険料")
    if "計画素案" in s:
        out.append("計画素案")
    return out


def _kigen(c):
    """回答期限（R8.9 のような表記）を並べ替えられる値にする。

    文字のまま並べると「R8.10」が「R8.9」より前になる。
    """
    m = re.match(r"R(\d+)\.(\d+)", str(c[8]))
    return (int(m.group(1)), int(m.group(2)), c[0]) if m else (99, 99, c[0])


YUSEN = sorted((c for c in KYOTSU
                if "概算" in tome(c) or "保険料" in tome(c)), key=_kigen)
SONOTA = sorted((c for c in KYOTSU if c not in YUSEN), key=_kigen)

# 令和8年10月1日に、ご決定を待たずに仮置きのとおり計画素案へ反映した事項。
# 反映した内容が計画素案に現に入っていることは自己点検13から18で確かめる。
SUSUMETA = [
    (88, "第6章第4節4",
     "3町（3圏域）を基盤整備単位とし、必要利用定員総数は［要協議］のまま"
     "残した"),
    (113, "第6章第4節4",
     "混合型特定施設入居者生活介護に係る必要利用定員総数は"
     "定めないこととした"),
    (131, "第6章第2節・第3節2",
     "利用回（日）数と、介護予防・日常生活支援総合事業及び"
     "一般介護予防事業の量の見込みを掲げた"),
    (112, "第6章第3節2",
     "令和6年度の利用者実人数を基礎として量の見込みを掲げ、"
     "令和7年度の実施状況は単位が異なるため別に示した"),
    (106, "第6章第6節2",
     "サービス別の趨勢を反映する補正は行わず、"
     "補正した場合の月額への効きを併記した"),
    (105, "第6章第2節",
     "施策による見込量の増減を織り込まなかった（自然体）"),
]

# ==================================================== 地域の課題を読む
_WB = load_workbook(KADAI_XLSX, data_only=True)
_KWS = _WB["01_地域課題の整理"]
KADAI = []
for row in _KWS.iter_rows(min_row=5, values_only=True):
    if row[0] is None or not str(row[0]).strip().isdigit():
        continue
    KADAI.append([("" if v is None else str(v)) for v in row[:6]])

# ==================================================== 体裁
doc = Document()
_st = doc.styles["Normal"]
_st.font.name = FONT
_st.font.size = Pt(10.5)
_st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
_st.paragraph_format.space_after = Pt(4)
_st.paragraph_format.line_spacing = 1.15
_sec = doc.sections[0]
_sec.page_width, _sec.page_height = Cm(21.0), Cm(29.7)
_sec.top_margin = _sec.bottom_margin = Cm(1.8)
_sec.left_margin = _sec.right_margin = Cm(1.9)
TEXTW = 21.0 - 1.9 * 2


def _shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    # shd は vAlign 等より前に置く（OOXML の要素の順序。CLAUDE.md §4）
    after = None
    for tag in ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                "w:vAlign", "w:hideMark"):
        after = tcPr.find(qn(tag))
        if after is not None:
            break
    if after is None:
        tcPr.append(shd)
    else:
        after.addprevious(shd)


def _cellmargin(t):
    mar = OxmlElement("w:tblCellMar")
    for tag, v in (("top", 0.06), ("left", 0.12), ("bottom", 0.06),
                   ("right", 0.12)):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:w"), str(int(v * 567)))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    # tblCellMar は tblLook より前に置く（CLAUDE.md §4）
    tblPr = t._tbl.tblPr
    after = None
    for tag in ("w:tblLook", "w:tblCaption", "w:tblDescription"):
        after = tblPr.find(qn(tag))
        if after is not None:
            break
    if after is None:
        tblPr.append(mar)
    else:
        after.addprevious(mar)


def _runs(p, text, size, bold=False, color=None):
    """** で囲んだ部分を太字にする（docx は Markdown を解釈しない）。"""
    for i, part in enumerate(str(text).split("**")):
        if not part:
            continue
        r = p.add_run(part)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold or (i % 2 == 1)
        if color is not None:
            r.font.color.rgb = color
        r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if not p.runs:
        r = p.add_run("")
        r.font.name = FONT
        r.font.size = Pt(size)
        r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return p


def P(text="", size=10.5, bold=False, align=None, after=4):
    p = doc.add_paragraph()
    _runs(p, text, size, bold)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    return p


# 節の先頭で改ページする節の番号。全ての節に入れると、
# 直前の節の本文があふれた1〜2行だけのページができる。
# 実際の紙面（tools/check_pages.py）で確かめて決める。
BRK = set(os.environ.get("BRK", "2,3,5,6,7,9,10").split(",")) - {""}
_SECNO = [0]


def H1(text, brk=None):
    _SECNO[0] += 1
    if brk is None:
        brk = str(_SECNO[0]) in BRK
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = brk
    p.paragraph_format.space_before = Pt(0 if brk else 14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def H2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(11)
    r.font.bold = True
    r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def BUL(text, size=10.5, mark="・"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.6)
    p.paragraph_format.space_after = Pt(3)
    _runs(p, mark + text, size)
    return p


def NOTE(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    _runs(p, "※ " + text, 9)
    return p


TBLNO = [0]


def CAP(text):
    TBLNO[0] += 1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("表%d　%s" % (TBLNO[0], text))
    r.font.name = FONT
    r.font.size = Pt(10)
    r.font.bold = True
    r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


# 表を紙面で分割しない上限（行数）。これを超える表は分割を許す
# （1ページに収まらないため、押し出すと空白が大きくなる）。
KEEP_MAX = 14


def TBL(head, rows, widths, size=9, center=None, first_bold=False,
        bold=None, keep=None):
    center = center or set()
    bold = bold or set()
    t = doc.add_table(rows=1 + len(rows), cols=len(head))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    s = sum(widths)
    if s > TEXTW:
        widths = [w * TEXTW / s for w in widths]
    _cellmargin(t)
    tr = t.rows[0]
    tr._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for j, v in enumerate(head):
        c = tr.cells[j]
        c.width = Cm(widths[j])
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _shade(c, "1F3864")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(v)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for i, row in enumerate(rows):
        tr = t.rows[i + 1]
        tr._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, v in enumerate(row):
            c = tr.cells[j]
            c.width = Cm(widths[j])
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if i % 2 == 1:
                _shade(c, "F2F5FA")
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if j in center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(p, "" if v is None else v, size,
                  bold=(first_bold and j == 0) or j in bold)
    if keep if keep is not None else (1 + len(rows) <= KEEP_MAX):
        for tr in t.rows[:-1]:
            for c in tr.cells:
                for pp in c.paragraphs:
                    pp.paragraph_format.keep_with_next = True
    return t


def _n(v, nd=0):
    return format(round(v, nd), ",.%df" % nd)


def _yen(v):
    return format(int(round(v)), ",")


def _sa(v):
    return ("＋%s円" % _n(v)) if v >= 0 else ("▲%s円" % _n(-v))


# ============================================================ 表紙
P("大雪地区広域連合　第10期介護保険事業計画", size=12,
  align=WD_ALIGN_PARAGRAPH.CENTER)
_p = P("構成3町合同　意見交換会　資料", size=18, bold=True,
       align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
for _r in _p.runs:
    _r.font.color.rgb = NAVY
P("令和8年10月　　大雪地区広域連合・東川町・美瑛町・東神楽町",
  size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
P("")
P("第10期介護保険事業計画（令和9年度から令和11年度まで）の策定に当たり、"
  "構成3町と広域連合で共通の認識を持っていただきたい事項と、"
  "ご決定をお願いする事項をまとめたものです。基準日 %s。" % KIJUNBI)
P("ご決定をお願いする事項には、**すべて「決まらない場合の当方の扱い」を"
  "置いています。** これは協議を進めるための仮置きであり、"
  "本意見交換会でその内容を諮るものです。"
  "ご決定がない場合もサービス見込量 第2次概算（令和8年10月30日）は"
  "組むことができますが、計画として確定する前にご判断を要します。")

# ============================================================ 第1節
H1("第1節　意見交換会の次第")

CAP("次第")
TBL(["区分", "内容"],
    [["日時", "令和8年10月　〔　　　　　　　　　　〕"],
     ["場所", "〔　　　　　　　　　　　　　　　〕"],
     ["出席", "大雪地区広域連合、東川町、美瑛町、東神楽町、受託者"],
     ["1", "第10期計画の策定の状況（第2節）"],
     ["2", "サービス見込量・給付費・保険料の試算（第3節・第4節）"],
     ["3", "地域の課題（第5節）"],
     ["4", "ご決定をお願いする事項　%d件（第6節）" % len(KYOTSU)],
     ["5", "認知症施策推進計画の扱い（第7節）"],
     ["6", "資料のご提供のお願いと今後の進め方（第9節・第10節）"]],
    [2.6, 14.6], center={0}, first_bold=True)
NOTE("町ごとの個別協議（令和8年10月中旬から下旬）でお伺いする事項は"
     "第8節に掲げています。本意見交換会では扱いません。")

# ============================================================ 第2節
H1("第2節　第10期計画の策定の状況")

H2("1　算定が済んでいるもの")
P("次の事項は、現に確認できている資料により算定を終えています。"
  "資料の追加のご提供を待たずに第2次概算に用いることができます。", size=10)
CAP("算定を終えている事項")
TBL(["事項", "内容", "計画本文"],
    [["人口・第1号被保険者数", "令和9年度から令和32年度まで（総合戦略及び"
      "住民基本台帳の実績趨勢による案C）", "第2章第1節・第6章第1節"],
     ["認定者数", "性別・年齢階級別・要介護度別。%s%s人、%s%s人"
      % (Y3L[-1], _n(nin(Y3[-1])), YL_HON_L[-1], _n(nin(YLONG_HON[-1]))),
      "第6章第1節"],
     ["サービス見込量", "%d区分の1月当たり利用者数（要介護度別）" % len(SVC),
      "第6章第2節"],
     ["利用回（日）数", "%d区分" % len(KAISU_MAP), "第6章第2節"],
     ["地域支援事業の量", "総合事業%d項目・一般介護予防事業%d項目"
      % (len(SOGO_RYO), len(KAYOI_RYO)), "第6章第3節"],
     ["給付費", "要介護度別の単価による。3か年計%s円" % _yen(SOU3),
      "第6章第3節"],
     ["保険料", "算定上の月額基準額%s円（記号A〜J）" % _n(GETSU),
      "第6章第6節"],
     ["中長期の推計", "／".join(YL_HON_L)
      + "の見込量・給付費・保険料", "第6章第1節・第3節"]],
    [3.6, 9.4, 4.2], first_bold=True)

H2("2　ご決定をお待ちしているもの")
CAP("確認事項の件数")
TBL(["区分", "件数", "内容"],
    [["確認事項（全体）", "%d件" % len(CHECK),
      "業務工程管理表 03_確認事項一覧の番号です（欠番・重複はありません）"],
     ["うち決着していないもの", "%d件" % len(MIKETSU), "―"],
     ["うち構成3町に関わるもの", "%d件" % len(MACHI),
      "確認先に町を含むものです"],
     ["　3町共通（本意見交換会）", "%d件" % len(KYOTSU),
      "広域連合として決めていただく事項です（第6節）"],
     ["　町ごと（個別協議）", "%d件" % sum(len(v) for v in KOBETSU.values()),
      "／".join("%s%d件" % (t, len(KOBETSU[t])) for t in TOWNS)
      + "（第8節）"]],
    [5.0, 2.0, 10.2], center={1}, first_bold=True, bold={1})

H2("3　決定がない場合の進め方")
P("算定の前提のうち、資料の追加のご提供や決定を待っているものを"
  "%d件整理しています。出どころにより扱いが分かれます。" % len(SUEOKI),
  size=10)
CAP("算定の前提の扱い")
TBL(["区分", "件数", "扱い"],
    [["受託者の側で確定したもの", "%d件" % NZ,
      "資料の受領を待たずに確定しました（第3節の算定に反映しています）"],
     ["受託者の側で確定できるもの", "%d件" % NA,
      "残っていません"],
     ["発注者・構成3町の資料をお願いしているもの", "%d件" % NB,
      "届かない場合は現に確認できている値のまま確定します。算定は止まりません"],
     ["国の告示・公布を待つもの", "%d件" % NC,
      "告示・公布が出た時点で必ず置き換えます。"
      "報酬改定率、第1号被保険者負担割合、所得段階の政令改正です"]],
    [6.2, 1.8, 9.2], center={1}, first_bold=True, bold={1})
P("したがって、**発注者及び構成3町から追加の資料のご提供を受けられない場合でも、"
  "国の告示が出れば計画は確定します。** "
  "一方で、国の告示によるものは催促のできないものであり、"
  "いずれも保険料を引き上げる向きに働きます（第4節）。", size=10)

def _oki(s):
    """算定の前提の置き方を本資料の文に直す。

    算定の側の表の中で用いている言い回し（自己点検の番号、シートの番号、
    「本表」）は本資料では意味を持たないため落とす。
    確定した日を述べる文も、本文の注で述べるためはじめに置かない。
    """
    v = str(s).replace("**", "")
    if v.startswith("令和8年10月1日") and "。" in v:
        v = v.split("。", 1)[1]
    v = re.sub(r"（点検[0-9０-９・点検]*）", "", v)
    v = re.sub(r"[0-9０-９]+シートに", "サービス見込量の算定に", v)
    v = v.replace("本表は", "見込量は")
    return v.strip()


CAP("受託者の側で確定した算定の前提")
TBL(["事項", "確定した内容", "保険料への効き"],
    [[s[0], _oki(s[3]), str(s[5])]
     for i, s in enumerate(SUEOKI, 1) if KUBUN.get(i) == "Z"],
    [3.0, 10.0, 4.2], first_bold=True)
NOTE("いずれも令和8年10月1日に確定し、第3節の算定に反映しています。")

# ============================================================ 第3節
H1("第3節　サービス見込量・給付費・保険料の試算")
P("サービス見込量 第1次概算（令和8年9月26日）の結果です。"
  "基準年度は令和7年度、単位は保険者（広域連合）です。"
  "介護保険法第3条及び地方自治法第284条により、"
  "サービス見込量・給付費・保険料・施策は保険者を単位として算定するもので、"
  "町別には算定しません。実績は町別に掲げています。")

H2("1　第1号被保険者数と認定者数")
CAP("第1号被保険者数と認定者数")
TBL(["区分"] + Y3L + ["令和7年度（実績）"],
    [["第1号被保険者数（人）"] + [_n(HIHO[y], 1) for y in Y3]
     + [_n(HIHO["2025"], 1)],
     ["認定者数（人）"] + [_n(nin(y), 1) for y in Y3]
     + [_n(nin("2025"), 1)]],
    [5.0, 3.0, 3.0, 3.0, 3.2], center={1, 2, 3, 4}, first_bold=True)
SRC_P = ("資料：住民基本台帳及び地方創生総合戦略（案C）、"
         "厚生労働省「介護保険事業状況報告」年報（令和7年度）")
P(SRC_P, size=8.5)

H2("2　給付費")
CAP("介護給付費・介護予防給付費の見込み")
TBL(["区分"] + Y3L + ["3か年計"],
    [["総給付費（円）"] + [_yen(KYUFU_Y[y]) for y in Y3] + [_yen(SOU3)]],
    [4.4, 3.4, 3.4, 3.4, 2.6], center={1, 2, 3, 4}, first_bold=True)
P("資料：サービス見込量の算定（第1次概算）", size=8.5)
NOTE("基準年度（令和7年度）の給付費は、"
     "地域包括ケア「見える化」システムの総括表と円単位で一致します。")

H2("3　保険料")
CAP("第10期保険料算定表")
_HR = [("標準給付費見込額", "A", "A"),
       ("地域支援事業費", "B", "B"),
       ("第1号被保険者負担分相当額", "C＝①×23％", "C"),
       ("調整交付金相当額", "D＝②×5％", "D"),
       ("調整交付金見込額", "E＝②×7.3755％", "E"),
       ("保険料収納必要額", "J＝C＋D－E＋F＋G±H－I", "J")]
TBL(["項目", "記号・算式", "第10期計（円）"],
    [[nm, kg, _yen(G3[k])] for nm, kg, k in _HR]
    + [["財政安定化基金拠出金・償還金等", "F・G・H", "0"],
       ["財政調整基金取崩額", "I", "0［要協議］"],
       ["予定保険料収納率", "収納率", "99.0％"],
       ["補正後被保険者数", "所得段階等補正", _n(G3["③"], 1) + "人"],
       ["算定上の月額基準額", "J÷収納率÷補正人数÷12",
        _n(GETSU, 0) + "円"],
       ["採用する月額基準額", "条例・計画案", "［要協議］"]],
    [6.0, 5.4, 5.8], first_bold=True, bold={2})
P("資料：サービス見込量の算定（第1次概算）、将来推計 第3段階", size=8.5)
P("算定上の月額基準額%s円は、百円未満を四捨五入すると**%s円**となり、"
  "第9期の基準額6,400円と同額になります。"
  "第9期の算定上の月額基準額は6,428円でした。"
  % (_n(GETSU), _n(KIJUN)), size=10)
NOTE("暫定の算定です。第10期の第1号被保険者負担割合、介護報酬改定率、"
     "令和7年度末の基金残高、所得段階の政令改正及びサービス見込量の採用値が"
     "確定するまでは、条例で定める基準額の算定には用いられません。"
     "財政安定化基金拠出金・償還金等（F・G・H）は当広域連合ではいずれも0です。")

# ============================================================ 第4節
H1("第4節　保険料の月額を動かす前提", brk=False)
P("算定上の月額基準額を動かす前提を、国の告示・公布によるものと"
  "広域連合のご判断によるものに分けて掲げます。"
  "**いずれも当方の算定には織り込んでいません。**", size=10)

_DE = {"C": "国の告示・公布", "B": "広域連合のご判断・資料"}
# 月額への効きが記されている算定の前提（確定したものを除く）。
# 効きの額は算定の側で求めたものを引く。
_LEV = [[_DE[KUBUN[i]], s[0], _oki(s[3]), str(s[5])]
        for kb in ("C", "B")
        for i, s in enumerate(SUEOKI, 1)
        if KUBUN.get(i) == kb and "月額" in str(s[5])]
CAP("月額基準額を動かす前提　%d件" % len(_LEV))
TBL(["出どころ", "事項", "当方の置き方", "月額への効き"],
    _LEV, [2.8, 3.4, 6.2, 4.8], bold={3})
P("資料：サービス見込量の算定（第1次概算）、対計画比の偏りの補正試算", size=8.5)
P("介護報酬改定率は1％当たり月額%sです。"
  "令和9年度は改定の年に当たり、加えて令和9年度改定を待たず"
  "令和8年度にも改定が行われることとされています。"
  "基準年度である令和7年度の単価には令和8年度改定が反映されていないため、"
  "**単価を令和7年度で固定することは2回分の改定を織り込んでいません。**"
  % _sa(_M["gaku_kaitei"](0.01)["月額"] - GETSU), size=10)
P("**上向きに働くものの多くは国の告示・公布によるもので、"
  "資料のご提供とは関わりなく生じます。"
  "下向きに働くものはいずれも織り込んでいません。** "
  "したがって現在の算定は、保険料が低く出る側に偏っています。", size=10)
NOTE("各欄は、当方の算定から当該の前提のみを動かした場合の値です。"
     "複数を同時に動かした場合の値ではありません。"
     "改定率を乗じることと単価の趨勢を延ばすことを重ねると二重になります。")

# ============================================================ 第5節
H1("第5節　地域の課題")
P("第10期計画に載せる地域の課題です。"
  "根拠となる数値と計画本文の反映先を対にして掲げています。"
  "「町の別」が「町で分かれる」「○○町」のものは、"
  "町ごとの個別協議でも扱います。")

CAP("地域の課題　%d件" % len(KADAI))
TBL(["No.", "地域の課題", "根拠となる事実（数値）", "町の別"],
    [[k[0], k[1], k[2], k[5]] for k in KADAI],
    [1.0, 4.4, 9.0, 2.8], center={0}, keep=False)
SRC_K = "資料：10月 地域課題の整理と3町協議の確認事項"
P(SRC_K, size=8.5)
NOTE("出所と計画本文の反映先は上記の資料に掲げています。")

# ============================================================ 第6節
H1("第6節　ご決定をお願いする事項（3町共通）")
P("広域連合として決めていただく事項です。"
  "「決まらない場合の当方の扱い」は受託者の仮置きであり、"
  "本意見交換会でその内容を諮ります。"
  "**仮置きのままでは計画として確定できません。**")

H2("1　サービス見込量・保険料の算定を止めているもの（%d件）" % len(YUSEN))
P("第2次概算（令和8年10月30日）に向けて、"
  "優先してご決定をお願いしたいものです。", size=10)
CAP("算定を止めている確認事項")
TBL(["No.", "確認事項", "決まらない場合の当方の扱い（仮置き）", "期限"],
    [["No.%d" % c[0], c[3], KITEI.get(c[0], ""), str(c[8])]
     for c in YUSEN],
    [1.4, 5.2, 8.6, 2.0], center={0, 3}, keep=False)
P("資料：業務工程管理表 03_確認事項一覧", size=8.5)

H2("2　計画の記載に関わるもの（%d件）" % len(SONOTA))
CAP("計画の記載に関わる確認事項")
TBL(["No.", "確認事項", "決まらない場合の当方の扱い（仮置き）", "期限"],
    [["No.%d" % c[0], c[3], KITEI.get(c[0], ""), str(c[8])]
     for c in SONOTA],
    [1.4, 5.2, 8.6, 2.0], center={0, 3}, keep=False)
P("資料：業務工程管理表 03_確認事項一覧", size=8.5)
NOTE("介護保険法第117条第2項第1号の記載事項である"
     "日常生活圏域ごとの必要利用定員総数（No.88）は、"
     "記載を落とすことができません。"
     "3町の整備方針が確定するまで［要協議］のまま残します。")

H2("3　受託者が仮置きにより進めた事項（%d件）" % len(SUSUMETA))
P("ご決定をお待ちしているもののうち、"
  "仮置きのとおり計画素案へ反映した事項です。"
  "**反映した内容でよいかをご確認ください。** "
  "ご決定の内容が異なる場合は、その内容により置き直します。", size=10)
CAP("仮置きにより計画素案へ反映した事項")
TBL(["No.", "確認事項", "反映した内容", "計画素案"],
    [["No.%d" % no, BY[no][3], naiyo, saki]
     for no, saki, naiyo in SUSUMETA if no in BY],
    [1.4, 4.6, 8.2, 3.0], center={0}, keep=False)
P("資料：業務工程管理表 03_確認事項一覧、計画素案", size=8.5)
NOTE("混合型特定施設入居者生活介護に係る必要利用定員総数（No.113）は、"
     "介護保険法第117条第3項が勘案事項として挙げるものです。"
     "本計画では定めないこととして記載しています。"
     "区域内の事業所の類型（混合型・介護専用型）は［要確認］であり、"
     "北海道の指定事業所一覧による確認をお願いしています。")

# ============================================================ 第7節
H1("第7節　認知症施策推進計画の扱い")
P("共生社会の実現を推進するための認知症基本法（令和5年法律第65号）は、"
  "第11条で国の認知症施策推進基本計画、第12条で都道府県認知症施策推進計画、"
  "第13条で市町村認知症施策推進計画を定めています。")

CAP("計画の階層と当広域連合の扱い")
TBL(["区分", "根拠", "策定主体", "本計画との関係"],
    [["認知症施策推進基本計画", "認知症基本法第11条", "国",
      "基本計画の内容を踏まえます"],
     ["都道府県認知症施策推進計画", "同法第12条", "北海道",
      "策定の状況は［要確認］です"],
     ["市町村認知症施策推進計画", "同法第13条", "東川町・美瑛町・東神楽町",
      "3町がそれぞれ策定します"],
     ["第10期介護保険事業計画", "介護保険法第117条第3項",
      "大雪地区広域連合",
      "同項により定める事項として認知症施策を掲げ、"
      "3町の市町村計画とは資料2により接続します"]],
    [4.2, 3.2, 4.0, 5.8], first_bold=True)
P("認知症基本法第13条の名宛人は市町村です。"
  "広域連合は地方自治法第284条により設けられ、"
  "介護保険法上の保険者としての事務を共同処理する特別地方公共団体であり、"
  "広域連合の計画に一体のものとして定めても"
  "市町村計画としての効力は生じません。"
  "このため**3町がそれぞれ市町村認知症施策推進計画を策定する**"
  "案により計画素案に反映しています（令和8年9月16日ご了承）。", size=10)
NOTE("条文を受領していないため、計画素案には条番号のみを記載し、"
     "項番号は記載していません［要確認］。")

H2("ご決定をお願いする事項")
_NC_NO = [130, 150, 149]
CAP("認知症施策に関する確認事項")
TBL(["No.", "確認事項", "決まらない場合の当方の扱い（仮置き）"],
    [["No.%d" % n, BY[n][3], KITEI.get(n, "")] for n in _NC_NO
     if n in BY],
    [1.4, 5.6, 10.2], center={0})
P("認知症の人及びその家族等からの意見の聴取は、"
  "記述の追加では満たされません。"
  "主体・方法・時期を3町と調整のうえ定める必要があり、"
  "策定委員会及び運営協議会の構成が決まる前にご判断を要します。", size=10)
NOTE("美瑛町の認知症初期集中支援推進事業は、医師を確保できないため"
     "令和7年度の実績が0件で休止中である旨が実施報告書に記載されています。"
     "市町村計画を3町がそれぞれ策定する場合、"
     "取組の差が計画の差として現れます。")

# ============================================================ 第8節
H1("第8節　町ごとの個別協議でお伺いする事項", brk=False)
P("令和8年10月中旬から下旬の個別協議でお伺いする事項です。"
  "本意見交換会では扱いません。日程の調整をお願いします。", size=10)
for t in TOWNS:
    rows = KOBETSU[t]
    if not rows:
        continue
    CAP("%s　%d件" % (t, len(rows)))
    TBL(["No.", "確認事項", "決まらない場合の当方の扱い（仮置き）", "期限"],
        [["No.%d" % c[0], c[3], KITEI.get(c[0], ""), str(c[8])]
         for c in rows],
        [1.4, 5.2, 8.6, 2.0], center={0, 3})
P("資料：10月 地域課題の整理と3町協議の確認事項（03〜05シート）", size=8.5)

# ============================================================ 第9節
H1("第9節　資料のご提供をお願いしたいもの")
_LACK_MI = [x for x in LACK if str(x[8]) not in ("完了", "受領済")]
_LACK_HI = [x for x in _LACK_MI if str(x[7]) == "高"]
P("資料の提供依頼%d件のうち、まだ受領していないものが%d件あります。"
  "うち優先度の高いものが%d件です。"
  % (len(LACK), len(_LACK_MI), len(_LACK_HI)), size=10)
P("いずれも届かない場合は現に確認できている値のまま確定します。"
  "算定そのものが止まるものはありません。"
  "ただし、国の告示・公布を待つもの（第4節）は"
  "資料のご提供とは別に残ります。", size=10)
CAP("優先度の高い資料　%d件" % len(_LACK_HI))
TBL(["No.", "資料", "用途", "お願い先", "期限"],
    [[str(x[0]), str(x[2]), str(x[4]), str(x[5]), str(x[6])]
     for x in _LACK_HI],
    [1.0, 6.4, 5.0, 2.6, 1.8], center={0, 4}, size=8, keep=False)
P("資料：業務工程管理表 06_資料提供依頼", size=8.5)

# ============================================================ 第10節
H1("第10節　10月以降の進め方")
CAP("工程")
TBL(["時期", "内容", "成果物"],
    [["令和8年10月上旬", "3町合同の意見交換会（本資料）",
      "共通%d件のご決定" % len(KYOTSU)],
     ["令和8年10月中旬から下旬", "町ごとの個別協議",
      "町ごと%d件のご決定" % sum(len(v) for v in KOBETSU.values())],
     ["令和8年10月30日", "サービス見込量 第2次概算",
      "見込量・給付費・保険料の第2次概算"],
     ["令和8年11月13日", "予算編成用の資料", "令和9年度予算に用いる給付費"],
     ["令和8年11月", "構成町意見の反映、第10期計画骨子の整理",
      "計画骨子"],
     ["令和9年1月以降", "意見公募手続、北海道の意見の聴取",
      "計画案"]],
    [4.0, 7.2, 5.8], first_bold=True)
NOTE("介護保険法第117条は、計画を定め又は変更しようとするときに、"
     "あらかじめ被保険者の意見を反映させるために必要な措置を講ずること及び"
     "都道府県の意見を聴くことを定めています。"
     "意見公募手続の時期と方法のご決定をお願いします。")

P("")
P("第2次概算（令和8年10月30日）は、ご決定がない場合も"
  "第6節の仮置きにより組みます。"
  "その場合は、仮置きによったことを第2次概算の資料に明示します。", size=10)

# ============================================================ 自己点検
_ALL = "".join(p.text for p in doc.paragraphs) + "".join(
    c.text for t in doc.tables for r in t.rows for c in r.cells)

chk(1, "確認事項を業務工程管理表から読んでいること",
    "全%d件・未決%d件（3町共通%d件）" % (len(CHECK), len(MIKETSU),
                                        len(KYOTSU)),
    len(CHECK) > 0 and len(KYOTSU) > 0)
chk(2, "確認事項の番号に欠番・重複がないこと",
    "No.1〜No.%d（%d件）" % (max(c[0] for c in CHECK), len(CHECK)),
    sorted(c[0] for c in CHECK) == list(range(1, len(CHECK) + 1)))
_KI_NASHI = [c[0] for c in MIKETSU if not KITEI.get(c[0])]
chk(3, "未決の全件に決まらない場合の扱いがあること",
    "既定値のないもの %s"
    % ("なし" if not _KI_NASHI else "No." + "・No.".join(
        str(x) for x in _KI_NASHI)),
    not _KI_NASHI)
chk(4, "3町共通の件数が優先と計画の記載に分かれていること",
    "算定を止めているもの%d件・計画の記載%d件（計%d件）"
    % (len(YUSEN), len(SONOTA), len(KYOTSU)),
    len(YUSEN) + len(SONOTA) == len(KYOTSU) and len(YUSEN) > 0)
chk(5, "町ごとの件数",
    "／".join("%s%d件" % (t, len(KOBETSU[t])) for t in TOWNS),
    all(len(KOBETSU[t]) >= 0 for t in TOWNS)
    and sum(len(KOBETSU[t]) for t in TOWNS) > 0)
chk(6, "地域の課題を資料から読んでいること",
    "%d件" % len(KADAI), len(KADAI) > 0)
chk(7, "総給付費が算定と一致すること",
    "%s円（3か年計）" % _yen(SOU3), SOU3 > 0 and _yen(SOU3) in _ALL)
chk(8, "算定上の月額基準額が算定と一致すること",
    "%s円" % _n(GETSU), 6000 < GETSU < 7000 and _n(GETSU) in _ALL)
chk(9, "百円未満を四捨五入した基準額を示していること",
    "%s円（第9期と同額）" % _n(KIJUN), _n(KIJUN) in _ALL)
chk(10, "保険料算定表の記号が算定の値によること",
    "A %s円／B %s円／J %s円" % (_yen(G3["A"]), _yen(G3["B"]),
                                _yen(G3["J"])),
    all(_yen(G3[k]) in _ALL for k in ("A", "B", "J")))
chk(11, "算定の前提の区分の件数が合うこと",
    "%d件＝確定%d＋受託者%d＋発注者・3町%d＋国%d"
    % (len(SUEOKI), NZ, NA, NB, NC),
    len(SUEOKI) == NZ + NA + NB + NC)
chk(12, "受託者の側で確定できるものが残っていないこと",
    "区分A %d件" % NA, NA == 0)

# 既定値のとおり計画素案へ反映されていることを、素案の実体から確かめる
_DOC = Document(RP.ROOT + "/output/"
                "第10期介護保険事業計画_協議用素案_令和8年8月.docx")
_DTXT = []
for _el in _DOC.element.body.iter(qn("w:p")):
    _DTXT.append("".join(_t.text or "" for _t in _el.iter(qn("w:t"))))
_DRAFT = "\n".join(_DTXT)

chk(13, "No.88の扱い（3圏域を基盤整備単位とし定員総数は要協議）が"
        "計画素案に反映されていること",
    "必要利用定員総数の表は［要協議］",
    "日常生活圏域ごとの必要利用定員総数" in _DRAFT
    and "3町を単位として" in _DRAFT)
chk(14, "No.113の扱い（混合型特定施設の定員総数を定めない）が"
        "計画素案に反映されていること",
    "第6章第4節4に記載",
    "混合型特定施設入居者生活介護に係る必要利用定員総数" in _DRAFT
    and "定めないこととします" in _DRAFT)
chk(15, "No.106の扱い（補正を行わず効きを併記する）が"
        "計画素案に反映されていること",
    "第6章第6節2に併記",
    "サービス別の趨勢を反映した場合" in _DRAFT
    and "補正を行っていません" in _DRAFT)
chk(16, "No.131の扱い（利用回（日）数と総合事業の量を掲げる）が"
        "計画素案に反映されていること",
    "第6章第2節（4）・第3節2",
    "利用回（日）数" in _DRAFT and "量の見込み" in _DRAFT)
chk(17, "No.112の扱い（総合事業の量の見込みを掲げる）が"
        "計画素案に反映されていること",
    "第6章第3節2",
    "一般介護予防事業" in _DRAFT and "量の見込み" in _DRAFT)
chk(18, "No.105の扱い（施策反映を行わない）が"
        "計画素案に反映されていること",
    "第6章第2節の注に施策の効果は反映していない旨",
    "施策の効果は反映していません" in _DRAFT)
_SS_NG = [no for no, _s, _n2 in SUSUMETA
          if no not in BY or BY[no][7] in SUMI or not KITEI.get(no)]
chk(18.5, "仮置きにより進めた事項が、未決で既定値のあるものであること",
    "%d件／条件に合わないもの %s"
    % (len(SUSUMETA),
       "なし" if not _SS_NG else "No." + "・No.".join(str(x) for x in _SS_NG)),
    not _SS_NG)

# 送付する資料であるため、受託者の内部の仕組み・作業経過を書かない
# （令和8年9月24日ご指示。CLAUDE.md §1）
_NAIBU = ["固定値", "実物から", "章節ごとに", "判定している", "読んで判定",
          "スクリプト", "runpy", ".py", "再実行", "書き写して",
          "個票がなくても", "モジュール"]
_nh = [w for w in _NAIBU if w in _ALL]
chk(19, "受託者の内部の仕組み・作業経過の語が残っていないこと",
    "残り %s" % ("なし" if not _nh else "・".join(_nh)), not _nh)

_NG = ["に由来する", "と整合する", "1件も", "有意差がないため関係がない",
       "全国トップ級"]
_ngh = [w for w in _NG if w in _ALL]
chk(20, "禁止表現が残っていないこと",
    "残り %s" % ("なし" if not _ngh else "・".join(_ngh)), not _ngh)

_PI = [(r"0\d{1,3}-\d{2,4}-\d{4}", "電話番号"),
       (r"\d{2,4}-\d{2,4}-\d{4}", "電話番号"),
       (r"[\w.+-]+@[\w-]+\.[\w.]+", "メールアドレス")]
_pih = [nm for pat, nm in _PI if re.search(pat, _ALL)]
chk(21, "個人情報の形（電話番号・メールアドレス）がないこと",
    "検出 %s" % ("なし" if not _pih else "・".join(_pih)), not _pih)

chk(22, "強調の指定が本文に残っていないこと",
    "** の数 %d" % _ALL.count("**"), "**" not in _ALL)

_OK_CH = re.compile(
    r"[ぁ-んァ-ヴ一-龥々ー０-９0-9A-Za-zＡ-Ｚａ-ｚ"
    r"、。・（）「」『』【】〔〕［］〜％±＋▲△○◯→／()%.,:：;；!！?？"
    r"\-－—―_＿　\s①②③④⑤⑥⑦⑧⑨⑩Ⅰ-Ⅹ°×÷≧≦<>＜＞=＝※\"'\[\]]")
_bad = sorted({c for c in _ALL if not _OK_CH.match(c)})
chk(23, "許容する文字の外の文字が混じっていないこと",
    "検出 %s" % ("なし" if not _bad else "".join(_bad)), not _bad)

_nhead = sum(1 for t in doc.tables
             if t.rows[0]._tr.trPr is not None
             and t.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None)
chk(24, "すべての表に見出し行の繰返しが設定されていること",
    "%d／%d表" % (_nhead, len(doc.tables)), _nhead == len(doc.tables))

chk(25, "日時・場所を記入枠としていること",
    "〔　〕の記入枠", "〔" in _ALL)

# ============================================================ 書き出し
os.makedirs(os.path.dirname(OUT), exist_ok=True)
docx_fix.fix(doc)
doc.save(OUT)

_ng_chk = [c for c in CHECKS if c[3] != "適合"]
print("書き出しました:", OUT)
print("段落 %d ／ 表 %d" % (len(doc.paragraphs), len(doc.tables)))
print("3町共通 %d件（算定を止めているもの%d・計画の記載%d）／町ごと %s"
      % (len(KYOTSU), len(YUSEN), len(SONOTA),
         "・".join("%s%d" % (t, len(KOBETSU[t])) for t in TOWNS)))
print("地域の課題 %d件／資料提供依頼 未受領%d件（うち優先度高%d件）"
      % (len(KADAI), len(_LACK_MI), len(_LACK_HI)))
print("仮置きにより計画素案へ反映した事項 %d件（No.%s）"
      % (len(SUSUMETA),
         "・No.".join(str(x[0]) for x in SUSUMETA)))
print("総給付費 %s円／算定上の月額 %s円（基準額 %s円）"
      % (_yen(SOU3), _n(GETSU), _n(KIJUN)))
print("算定の前提 %d件＝確定%d＋受託者%d＋発注者・3町%d＋国%d"
      % (len(SUEOKI), NZ, NA, NB, NC))
print("自己点検 %d件：適合%d件・不適合%d件"
      % (len(CHECKS), len(CHECKS) - len(_ng_chk), len(_ng_chk)))
for c in _ng_chk:
    print("  不適合:", c[0], c[1], c[2])
if _ng_chk:
    sys.exit(1)
print("すべての点検に適合しました。")
