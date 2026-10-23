# -*- coding: utf-8 -*-
"""
村から提供された給付の台帳を、月別・年度別の集計値に落とす

受領
  令和８年10月22日、村から電子メールで受領した2つの圧縮ファイル。
    「障害サービス - 名前なし」（自立支援給付の台帳 9ファイル）
    「②地域生活支援事業実績」（地域生活支援事業の台帳 2ファイル）
  令和８年10月23日、追加で受領した1つの圧縮ファイル（照会Ｍ-137への回答）。
    「障害児通所実績-名前なし」（障がい児通所支援・障がい児相談支援の台帳 4ファイル）

なぜ原本を格納しないか
  台帳は利用者ごとの記録である（氏名は伏せてあるが、障害支援区分・
  所得区分・利用先の事業所名・月ごとの利用日数と給付費が並ぶ）。
  本村の人口は約2,300人であり、区分と事業所の組合せで個人が分かりうる。
  このため原本はリポジトリに置かず、本スクリプトが出す集計値
  （サービス・年度・月ごとの人数・日数・金額。利用者・事業所の
  名は含まない）だけを source に格納する。
  原本は受領した圧縮ファイルのまま当方の手元で保管する。

使い方
  python3 -I build_kitashiobara_daicho_chuushutsu.py \
      <「障害サービス - 名前なし」を展開した場所> \
      <「②地域生活支援事業実績」を展開した場所> \
      <「障害児通所実績-名前なし」を展開した場所>

  出力: source/村資料_台帳_20261022/台帳の集計.csv

台帳の読み方（自立支援給付）
  ・シートは年度ごと（R6・R7・R8）。行は「実施月」の3月から翌年2月まで。
  ・利用者ごとに列のブロックがあり、ブロックの頭は「利用日数」又は
    「日数」の列（相談支援は「基準額」の列）。
  ・給付費は「当月給付費」＞「給付費」＞「村支出額」の順で取る。
  ・行に「実績」「見込」の札があるシートは「実績」の月だけを数える。
  ・利用者の欄が無いブロックと「合計」「小計」のブロックは数えない。
  ・同じ利用者・同じ値のブロックが二重に置かれている場合は1つに数える。

台帳の読み方（障がい児）
  ・自立支援給付の台帳と様式が違う。列のブロックの頭は「延日数」の列
    （障がい児相談支援は「基準額」の列）で、「利用日数」は支給量の欄である。
  ・利用したかどうかは、通所は延日数が1日以上、相談は給付費が1円以上で決める
    （加算だけが請求されている月は利用に数えない）。
  ・シート名に「見込」とあるシートは、全体を見込みとして実績に数えない
    （放課後等デイサービスの令和８年度は「R8見込」の1枚だけである）。
"""

import csv
import os
import re
import sys

REPO_ROOT = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = f"{REPO_ROOT}/source/村資料_台帳_20261022"
OUT_FILE = f"{OUT_DIR}/台帳の集計.csv"

#   ファイル名の頭の番号でサービスを決める（ファイル名の字は機種依存の
#   ことがあるため、番号とサービス名の両方で照らす）
KYUFU = [
    ("1_", "居宅介護"),
    ("2_", "共同生活援助"),
    ("3_", "生活介護"),
    ("4_", "施設入所支援"),
    ("5_", "短期入所"),
    ("6_", "自立訓練（生活訓練）"),
    ("7_", "就労継続支援Ｂ型"),
    ("8_", "就労定着支援"),
    ("9_", "計画相談支援"),
]
ORDER = [3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2]
MON = ("実施月", "受付月", "提供月")
INFO = ("利用者", "氏名", "施設名", "事業所", "障害程度", "利用時間")


# ===========================================================================
# 読み込み
# ===========================================================================
def sheets(p):
    if p.endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(p)
        for sh in wb.sheets():
            yield sh.name, [[sh.cell_value(r, c) for c in range(sh.ncols)]
                            for r in range(sh.nrows)]
    else:
        from openpyxl import load_workbook
        wb = load_workbook(p, data_only=True)
        for ws in wb.worksheets:
            yield ws.title, [["" if v is None else v for v in r]
                             for r in ws.iter_rows(values_only=True)]


def s(v):
    return str(v).replace("\n", "").replace(" ", "").replace("　", "")


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def hankaku(t):
    return t.translate(str.maketrans("０１２３４５６７８９Ｒ", "0123456789R"))


def month_of(v):
    if isinstance(v, (int, float)) and 1 <= v <= 12 and float(v).is_integer():
        return int(v)
    m = re.fullmatch(r"※?(\d{1,2})(月)?", s(v))
    if m and 1 <= int(m.group(1)) <= 12:
        return int(m.group(1))
    return None


def nendo_of(name):
    m = re.search(r"R\s*(\d+)", hankaku(name))
    return int(m.group(1)) if m else None


# ===========================================================================
# 自立支援給付の台帳
# ===========================================================================
def _label_left(row, c):
    for j in range(c, -1, -1):
        if j < len(row) and s(row[j]) and s(row[j]) not in INFO:
            return s(row[j])
    return ""


def parse_sheet(rows):
    """利用者のブロックを返す: [dict(user, jig, months={月: (日数, 給付費, 札)})]"""
    out = []
    hdrs = [(r, c) for r, row in enumerate(rows) for c, v in enumerate(row)
            if s(v) in MON]
    for k, (hr, mc) in enumerate(hdrs):
        nxt = hdrs[k + 1][0] if k + 1 < len(hdrs) else len(rows)
        info = {}
        for r in range(hr - 1, max(-1, hr - 12), -1):
            for cand in (0, 1):
                kk = s(rows[r][cand]) if cand < len(rows[r]) else ""
                if kk in INFO + ("所得区分", "サービス内容") and kk not in info:
                    info[kk] = r
        hrows = [rows[r] for r in range(max(0, hr - 2), min(len(rows), hr + 2))]
        width = max(len(x) for x in rows[hr:nxt]) if rows[hr:nxt] else 0

        def cells(c):
            return [s(x[c]) for x in hrows if c < len(x) and s(x[c])]

        day_cols = [c for c in range(width)
                    if any(v.startswith("利用日数") or v == "日数"
                           for v in cells(c))]
        base_cols = [c for c in range(width)
                     if any(v.startswith("基準額") for v in cells(c))]
        blocks = day_cols if day_cols else base_cols
        for bi, c0 in enumerate(blocks):
            c1 = blocks[bi + 1] if bi + 1 < len(blocks) else width
            kc = None
            for key in ("当月給付費", "給付費", "村支出額"):
                hit = [c for c in range(c0, c1)
                       if any((key in v) if key != "給付費"
                              else v.startswith(key) for v in cells(c))]
                if hit:
                    kc = hit[0]
                    break

            def at(key):
                r = info.get(key)
                if r is None:
                    return ""
                row = rows[r]
                for j in range(c0, min(c1, len(row))):
                    if s(row[j]) and s(row[j]) not in INFO:
                        return s(row[j])
                return _label_left(row, c0)

            user = at("利用者") or at("氏名")
            if not user or any(w in user for w in ("合計", "小計")) \
                    or "合計" in at("施設名"):
                continue
            months = {}
            for r in range(hr + 1, nxt):
                row = rows[r]
                if mc >= len(row):
                    continue
                mo = month_of(row[mc])
                if mo is None or mo in months:
                    continue
                tags = {s(v) for v in row if s(v) in ("実績", "見込")}
                lab = "見込" if "見込" in tags else (
                    "実績" if "実績" in tags else "")
                d = num(row[c0]) if (day_cols and c0 < len(row)) else None
                kf = num(row[kc]) if (kc is not None and kc < len(row)) else None
                months[mo] = (d, kf, lab)
            out.append(dict(user=user, jig=at("施設名") or at("事業所"),
                            months=months, day=bool(day_cols)))
    has_lab = any(v[2] for u in out for v in u["months"].values())
    for u in out:
        u["has_lab"] = has_lab
    return out


def _used(u, mo):
    v = u["months"].get(mo)
    if not v:
        return False
    d, kf, lab = v
    if lab == "見込" or (u["has_lab"] and lab != "実績"):
        return False
    return (d or 0) > 0 if u["day"] else (kf or 0) > 0


def kyufu_rows(path, jigyo):
    """1ファイル分の月別・年度別の行と、利用先の事業所の名（照合用）。"""
    rows, jig = [], {}
    for na, data in sheets(path):
        y = nendo_of(na)
        us = parse_sheet(data)
        seen, uniq = set(), []
        for u in us:
            key = (u["user"], u["jig"],
                   tuple(sorted((m, v[0], v[1]) for m, v in u["months"].items())))
            if key in seen:
                continue
            seen.add(key)
            uniq.append(u)
        if uniq and uniq[0]["has_lab"]:
            jm = [m for m in ORDER
                  if any(u["months"].get(m, (0, 0, ""))[2] == "実績"
                         for u in uniq)]
        else:
            jm = [m for m in ORDER if any(_used(u, m) for u in uniq)]
        for m in jm:
            hito = [u for u in uniq if _used(u, m)]
            rows.append(dict(
                kind="月", jigyo=jigyo, nendo=y, tsuki=m,
                ninzu=len(hito),
                nissu=sum((u["months"][m][0] or 0) for u in hito if u["day"]),
                kingaku=round(sum((u["months"][m][1] or 0) for u in hito)),
                biko=""))
        jitsu = {u["user"] for u in uniq if any(_used(u, m) for m in jm)}
        rows.append(dict(kind="年", jigyo=jigyo, nendo=y, tsuki="",
                         ninzu=len(jitsu), nissu="", kingaku="",
                         biko=f"実績月{len(jm)}か月"))
        jig[y] = [u["jig"] for u in uniq if u["user"] in jitsu]
    return rows, jig


def _kyoutsuu(a, b):
    """二つの事業所名の最長の共通部分の長さ（法人格の語は除く）。"""
    a = a.replace("社会福祉法人", "")
    b = b.replace("社会福祉法人", "")
    best = 0
    for i in range(len(a)):
        for j in range(i + best + 1, len(a) + 1):
            if a[i:j] in b:
                best = j - i
            else:
                break
    return best


def kasanari(jig_nyusho, jig_seikatsu):
    """施設入所支援の利用者のうち、入所先と同じ事業所で生活介護を
    利用している人数（年度ごと）。事業所名は出力しない。"""
    out = []
    for y in sorted(jig_nyusho):
        sk = jig_seikatsu.get(y, [])
        n = sum(1 for a in jig_nyusho[y]
                if any(_kyoutsuu(a, b) >= 4 for b in sk))
        out.append(dict(kind="重なり", jigyo="施設入所支援×生活介護",
                        nendo=y, tsuki="", ninzu=n, nissu="", kingaku="",
                        biko=f"施設入所{len(jig_nyusho[y])}人のうち"))
    return out


# ===========================================================================
# 障がい児の台帳
# ===========================================================================
KODOMO = [
    ("1_", "放課後等デイサービス"),
    ("2_", "児童発達支援"),
    ("3_", "保育所等訪問支援"),
    ("4_", "障がい児相談支援"),
]


def kodomo_rows(path, jigyo):
    rows = []
    for na, data in sheets(path):
        y = nendo_of(na)
        mikomi_sheet = "見込" in s(na)
        hdr = [r for r, row in enumerate(data) if row and s(row[0]) == "実施月"]
        users, months_all = {}, []
        for hr in hdr:
            head = data[hr]
            soudan = not any(s(v) == "延日数" for v in head)
            starts = [c for c, v in enumerate(head)
                      if s(v) == ("基準額" if soudan else "延日数")]
            lab_r = next((r for r in range(hr - 1, max(-1, hr - 8), -1)
                          if s(data[r][0]) == "利用者"), None)
            for c0 in starts:
                user = s(data[lab_r][c0]) if lab_r is not None and \
                    c0 < len(data[lab_r]) else ""
                if not user or "合計" in user:
                    continue
                for r in range(hr + 1, len(data)):
                    row = data[r]
                    if row and s(row[0]).startswith("合計"):
                        break
                    mo = month_of(row[0]) if row else None
                    if mo is None:
                        continue
                    tags = {s(v) for v in row if s(v) in ("実績", "見込")}
                    if mikomi_sheet or "見込" in tags:
                        continue
                    val = num(row[c0 + 2] if soudan else row[c0]) \
                        if c0 + 2 < len(row) else None
                    used = (val or 0) > 0
                    nissu = 0 if soudan else int(val or 0)
                    kin = num(row[c0 + 2] if soudan else row[c0 + 5]) or 0
                    users.setdefault(user, {})[mo] = (used, nissu, kin)
                    if mo not in months_all:
                        months_all.append(mo)
        jm = [m for m in ORDER if m in months_all]
        for m in jm:
            hito = [u for u, mm in users.items() if mm.get(m, (False,))[0]]
            rows.append(dict(
                kind="月", jigyo=jigyo, nendo=y, tsuki=m, ninzu=len(hito),
                nissu=sum(users[u][m][1] for u in hito),
                kingaku=round(sum(users[u][m][2] for u in hito)), biko=""))
        jitsu = [u for u, mm in users.items() if any(v[0] for v in mm.values())]
        rows.append(dict(kind="年", jigyo=jigyo, nendo=y, tsuki="",
                         ninzu=len(jitsu), nissu="", kingaku="",
                         biko=("見込のみ" if mikomi_sheet
                               else f"実績月{len(jm)}か月")))
    return rows


# ===========================================================================
# 地域生活支援事業の台帳
# ===========================================================================
#   日常生活用具（ストマ）のシートは半年ごと。令和６年度の下半期は
#   「R6.10月～3月分」を令和7年1月に差し替えた「R7.1月～」がある
#   （1人の給付が1月から加わっている）。差し替え後を用いる。
STOMA_SHEET = [
    ("R6.4月～9月分", 6), ("R7.1月～", 6),
    ("R7.4月～9月分", 7), ("R7.10月～3月分", 7),
    ("R8.4月～9月分", 8), ("R8.10月～3月分", 8),
]
STOMA_SKIP = ["R6.10月～3月分"]


def _key(name):
    return hankaku(s(name))


def stoma_rows(path):
    rows = []
    found = set()
    for na, data in sheets(path):
        k = _key(na)
        if k in [_key(x) for x in STOMA_SKIP]:
            found.add(k)
            continue
        hit = [y for nm, y in STOMA_SHEET if _key(nm) == k]
        if not hit:
            raise SystemExit(f"ストマのシートが想定と違う: {na}")
        found.add(k)
        hito, tsuki, kin = 0, 0, 0
        for r in data:
            if len(r) < 8 or s(r[2]) != "ストマ用装具":
                continue
            if s(r[1]):
                hito += 1
            m = re.findall(r"(\d+)", hankaku(s(r[3])).split("月分")[0])
            tsuki += len(m)
            kin += num(r[7]) or 0
        rows.append(dict(kind="用具", jigyo="排泄管理支援用具", nendo=hit[0],
                         tsuki="", ninzu=hito, nissu=tsuki,
                         kingaku=round(kin), biko=s(na)))
    want = {_key(x) for x, _ in STOMA_SHEET} | {_key(x) for x in STOMA_SKIP}
    if found != want:
        raise SystemExit(f"ストマのシートが揃っていない: {sorted(want - found)}")
    return rows


def _getsu(label, year):
    """「R6.4月」「5月」「Ｒ7.1月」から（暦年, 月）を返す。"""
    t = hankaku(s(label))
    m = re.fullmatch(r"(?:R(\d+)\.)?(\d{1,2})月", t)
    if not m:
        return None
    if m.group(1):
        year = int(m.group(1))
    return year, int(m.group(2))


def chiiki_rows(path):
    rows = []
    for na, data in sheets(path):
        name = s(na)
        if name in ("日中一時支援", "地域生活活動支援センター事業"):
            jigyo = ("日中一時支援事業" if name == "日中一時支援"
                     else "地域活動支援センター機能強化事業")
            head = "".join(s(v) for v in data[0])
            m = re.search(r"支給決定(\d+)人", hankaku(head))
            if m:
                rows.append(dict(kind="決定", jigyo=jigyo, nendo="", tsuki="",
                                 ninzu=int(m.group(1)), nissu="", kingaku="",
                                 biko="台帳の見出しの支給決定の人数"))
            year = None
            for r in data[2:]:
                g = _getsu(r[0], year)
                if g is None or num(r[1]) is None:
                    continue
                year, mo = g
                rows.append(dict(kind="月", jigyo=jigyo,
                                 nendo=year if mo >= 4 else year - 1,
                                 tsuki=mo, ninzu=int(num(r[1])),
                                 nissu=int(num(r[2]) or 0),
                                 kingaku=round(num(r[3]) or 0), biko=""))
        elif name == "意思疎通支援事業":
            for r in data[2:]:
                y = nendo_of(s(r[0]))
                if y is None or num(r[1]) is None:
                    continue
                rows.append(dict(kind="年", jigyo="意思疎通支援事業",
                                 nendo=y, tsuki="", ninzu=int(num(r[1])),
                                 nissu="", kingaku=round(num(r[2]) or 0),
                                 biko="申請件数"))
        elif name == "相談支援事業":
            col = [c for c, v in enumerate(data[1]) if "委託料" in s(v)]
            if not col:
                raise SystemExit("相談支援事業の委託料の列が無い")
            year = None
            for r in data[2:]:
                g = _getsu(r[0], year)
                if g is None or num(r[col[0]]) is None:
                    continue
                year, mo = g
                rows.append(dict(kind="月", jigyo="障がい者相談支援事業（委託料）",
                                 nendo=year if mo >= 4 else year - 1,
                                 tsuki=mo, ninzu="", nissu="",
                                 kingaku=round(num(r[col[0]])),
                                 biko="4町村の利用率で按分した本村の分"))
        else:
            raise SystemExit(f"地域生活支援事業のシートが想定と違う: {na}")
    return rows


# ===========================================================================
def main():
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    d_kyufu, d_chiiki, d_kodomo = sys.argv[1], sys.argv[2], sys.argv[3]
    rows, jig = [], {}
    names = sorted(os.listdir(d_kyufu))
    for head, jigyo in KYUFU:
        hit = [f for f in names if f.startswith(head)]
        if len(hit) != 1:
            raise SystemExit(f"台帳のファイルが一意でない: {head}{jigyo} {hit}")
        r, j = kyufu_rows(os.path.join(d_kyufu, hit[0]), jigyo)
        rows += r
        jig[jigyo] = j
    rows += kasanari(jig["施設入所支援"], jig["生活介護"])
    names = sorted(os.listdir(d_kodomo))
    for head, jigyo in KODOMO:
        hit = [f for f in names if f.startswith(head)]
        if len(hit) != 1:
            raise SystemExit(f"台帳のファイルが一意でない: {head}{jigyo} {hit}")
        rows += kodomo_rows(os.path.join(d_kodomo, hit[0]), jigyo)
    for f in sorted(os.listdir(d_chiiki)):
        p = os.path.join(d_chiiki, f)
        if "ストマ" in f:
            rows += stoma_rows(p)
        else:
            rows += chiiki_rows(p)
    os.makedirs(OUT_DIR, exist_ok=True)
    cols = ["kind", "jigyo", "nendo", "tsuki", "ninzu", "nissu",
            "kingaku", "biko"]
    with open(OUT_FILE, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, lineterminator="\n")
        w.writeheader()
        for r in rows:
            w.writerow(r)
    print(f"作成: {OUT_FILE}（{len(rows)}行）")


if __name__ == "__main__":
    main()
