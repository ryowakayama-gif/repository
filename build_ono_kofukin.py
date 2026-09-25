"""小野町 第10期 令和8年度の交付金とりまとめ資料を読み、協議会資料として整理する。

令和8年9月25日受領
  令和８年度保険者機能強化推進交付金・介護保険保険者努力支援交付金
    （都道府県分）評価指標に係る該当状況調査票集計表（推進＋支援）  001732598.xlsx
  令和８年度保険者機能強化推進交付金 及び介護保険保険者努力支援交付金
    に係る評価指標（都道府県分）                                  001732599.pdf

**受領したのは都道府県分である。市町村分は含まれていない。**
小野町の得点分析には市町村分の評価指標と町の得点票が要る。
ここでできるのは、福島県の位置を確かめ、県の支援が期待できる領域と
そうでない領域を分けることである。

**整理の仕方は大雪地区広域連合の案件による。**
交付金の評価指標は、計画の進捗状況のモニタリング結果を外部の関係者を含む
議論の場で検証すること、評価結果を関係者間で共有して自立支援等に資する取組を
検討することを求めている。評価結果を協議会に諮ること自体が要件であるため、
計画本文には掲載せず、毎年度の協議会資料として整理する。
計画本文には、協議会へ報告し検証する手順のみを定める。

**分析の方法も同広域連合による。**
公表資料の明細列ごとに全国該当率（得点した団体の数÷47）を算定し、
「全国の多くが得点しているのに福島県が0点である項目」を洗い出す。
取組の水準が低いのか、事業に着手していないのかを分けられる。

出力
  08_協議会資料/小野町_協議会資料_交付金の評価結果と福島県の位置_YYYYMMDD.docx
  08_協議会資料/小野町_交付金_福島県の得点分析_YYYYMMDD.xlsx
"""

import pathlib
import statistics
import warnings

warnings.filterwarnings("ignore")

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).parent
OUT = ROOT / "小野町_引継ぎ_整理済" / "08_協議会資料"
SRC = pathlib.Path("/root/.claude/uploads/d7dca1d7-b918-515d-8b6d-9309b5e7ef32"
                   "/31df052b-001732598.xlsx")
ASOF = "20260925"
ASOF_JP = "令和8年9月25日"
JP_MIN = "游明朝"
JP_GO = "游ゴシック"

KEN = "福島県"
TOHOKU = ["青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県"]

# 目標別の合計列（配点）。列番号は受領ファイルの並びによる。
MOKUHYO = [
    (40, "推進", "目標Ⅰ　持続可能な地域のあるべき姿をかたちにする", 100),
    (60, "推進", "目標Ⅱ　公正・公平な給付を行う体制を構築する", 100),
    (125, "推進", "目標Ⅲ　介護人材の確保その他のサービス提供基盤の整備", 100),
    (158, "推進", "目標Ⅳ　高齢者が可能な限り自立した日常生活を営む（成果）", 100),
    (159, "推進", "**保険者機能強化推進交付金　合計**", 400),
    (248, "支援", "目標Ⅰ　介護予防／日常生活支援を推進する", 100),
    (283, "支援", "目標Ⅱ　認知症総合支援を推進する", 100),
    (307, "支援", "目標Ⅲ　在宅医療・在宅介護連携の体制を構築する", 100),
    (340, "支援", "目標Ⅳ　高齢者が可能な限り自立した日常生活を営む（成果）", 100),
    (341, "支援", "**介護保険保険者努力支援交付金　合計**", 400),
    (342, "計", "**総合計**", 800),
]
SUBTOTAL_COLS = {c for c, _a, _b, _d in MOKUHYO} | {
    38, 39, 58, 59, 123, 124, 246, 247, 281, 282, 305, 306}

# 各目標は（ⅰ）体制・取組 と（ⅱ）活動 に分かれる。
# 他町村の案件で、この2つは得点の動かし方が違うことを整理している。
#   （ⅰ）体制・取組　仕組みがあるか（計画への記載・要綱・会議体）
#                     → 計画に書けば翌年度の取組として得点になる
#   （ⅱ）活動　　　　実際に行ったか（回数・人数・件数）
#                     → 実績の積上げに時間がかかる
#   成果指標群（目標Ⅳ）　結果が出ているか
#                     → 単年度では動かず、3年程度の遅れで現れる
SHIHYOGUN = [
    ("推進", "目標Ⅰ　持続可能な地域", 38, 39),
    ("推進", "目標Ⅱ　公正・公平な給付", 58, 59),
    ("推進", "目標Ⅲ　介護人材の確保", 123, 124),
    ("支援", "目標Ⅰ　介護予防／日常生活支援", 246, 247),
    ("支援", "目標Ⅱ　認知症総合支援", 281, 282),
    ("支援", "目標Ⅲ　在宅医療・在宅介護連携", 305, 306),
]

# 小野町の計画に直に効く項目。列番号で拾う。
ONO_KANREN = {
    261: "認知症施策に関する市町村支援（イ）",
    9: "地域課題の把握・分析等に関する市町村支援（イ）",
    273: "チームオレンジ設置状況（ア）",
    269: "管内の認知症サポーターステップアップ講座修了者数（ア）",
    176: "介護予防等と保健事業の一体的実施に向けた環境整備（ウ）",
    177: "介護予防等と保健事業の一体的実施に向けた環境整備（エ）",
    178: "介護予防等と保健事業の一体的実施に向けた環境整備（オ）",
    198: "管内の地域包括支援センター事業評価の達成状況（ア）",
    134: "長期的な要介護度の変化（要介護1・2）（ア）",
    316: "長期的な要介護度の変化（要介護1・2）（ア）",
}

HEAD = PatternFill("solid", fgColor="1F3864")
NG = PatternFill("solid", fgColor="FCE4E4")
WARN = PatternFill("solid", fgColor="FFE0B2")
OK = PatternFill("solid", fgColor="E2EFDA")
CALC = PatternFill("solid", fgColor="EAF1FB")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------------------------------------------------------------- 読み取り

def load():
    ws = openpyxl.load_workbook(SRC, data_only=True)["全国集計（都道府県）"]
    rows = {}
    for r in range(14, ws.max_row + 1):
        nm, no = ws.cell(r, 3).value, ws.cell(r, 1).value
        if isinstance(nm, str) and isinstance(no, (int, float)):
            rows[nm.strip()] = r
    return ws, rows


def mokuhyo(ws, rows):
    """目標別の得点・得点率・順位。"""
    fr = rows[KEN]
    out = []
    for c, kbn, lab, hai in MOKUHYO:
        vals = [ws.cell(r, c).value for r in rows.values()
                if isinstance(ws.cell(r, c).value, (int, float))]
        v = ws.cell(fr, c).value
        rank = sorted(vals, reverse=True).index(v) + 1
        out.append({"区分": kbn, "項目": lab, "配点": hai, KEN: v,
                    "得点率": v / hai, "全国平均": statistics.mean(vals),
                    "順位": rank, "団体数": len(vals)})
    return out


def soan_seika():
    """素案の成果指標の件数と、目標値が未設定の件数を数える。

    件数を手で書くと素案の改訂で狂うため、素案そのものから採る。
    """
    cand = sorted((ROOT / "小野町_引継ぎ_整理済" / "03_計画素案").glob("*素案_第*.docx"))
    if not cand:
        return None, None
    doc = Document(str(cand[-1]))
    n = mi = 0
    for t in doc.tables:
        if [c.text.strip() for c in t.rows[0].cells][:1] != ["成果指標（案）"]:
            continue
        for r in t.rows[1:]:
            n += 1
            if r.cells[2].text.strip() == "【協議会で設定】":
                mi += 1
    return n, mi


def shihyogun(ws, rows):
    """（ⅰ）体制・取組 と（ⅱ）活動 の別に得点を取る。"""
    fr = rows[KEN]
    out = []
    for kbn, lb, c1, c2 in SHIHYOGUN:
        row = {"区分": kbn, "目標": lb}
        for key, c in (("体制・取組", c1), ("活動", c2)):
            vals = [ws.cell(r, c).value for r in rows.values()]
            v = ws.cell(fr, c).value
            row[key] = {
                "配点": ws.cell(9, c).value, KEN: v,
                "全国平均": statistics.mean(vals),
                "順位": sorted(vals, reverse=True).index(v) + 1}
        out.append(row)
    return out


def lab(m):
    """「推進目標Ⅲ」のように、交付金の区分をつけた短い名。

    推進・支援のどちらにも目標Ⅰ〜Ⅳがあるため、目標名だけでは判じられない。
    """
    return f"{m['区分']}{m['項目'].split('　')[0]}"


def zero_items(ws, rows, thr=0.60):
    """全国該当率が thr 以上で福島県が0点の項目。大雪の方法による。"""
    fr = rows[KEN]
    out = []
    cur2 = cur3 = cur6 = ""
    for c in range(8, 343):
        for src in (2, 3, 6):
            v = ws.cell(src, c).value
            if v:
                if src == 2:
                    cur2 = str(v).strip()
                elif src == 3:
                    cur3 = str(v).strip()
                else:
                    cur6 = str(v).strip()
        if c in SUBTOTAL_COLS:
            continue
        hai, v, tot = ws.cell(9, c).value, ws.cell(fr, c).value, ws.cell(10, c).value
        if not all(isinstance(x, (int, float)) for x in (hai, v, tot)) or not hai:
            continue
        # 該当率は得点した団体の数で取る。配点を按分する指標があっても
        # 「何団体が得点しているか」がぶれないため。
        vals = [ws.cell(r, c).value for r in rows.values()]
        hit = sum(1 for x in vals if isinstance(x, (int, float)) and x > 0)
        rate = hit / len(rows)
        if v == 0 and rate >= thr:
            out.append({"列": c, "交付金": cur2, "目標": cur3, "評価指標": cur6,
                        "枝番": str(ws.cell(7, c).value or "").strip(),
                        "配点": hai, "全国該当率": rate,
                        "得点した団体": hit, "団体数": len(rows),
                        "全国得点率": tot / (len(rows) * hai),
                        "小野町に効く": ONO_KANREN.get(c, "")})
    out.sort(key=lambda x: -x["全国該当率"])
    return out


def suii(ws, rows):
    """令和7年度と令和8年度の合計得点・順位。"""
    out = []
    for nm, r in rows.items():
        out.append({"都道府県": nm, "R7": ws.cell(r, 6).value,
                    "R7順位": ws.cell(r, 7).value,
                    "R8": ws.cell(r, 342).value,
                    "R8順位": ws.cell(r, 343).value})
    return out


# ---------------------------------------------------------------- 体裁

def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = JP_MIN
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), JP_MIN)
    return doc


def head(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.name = JP_GO
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_GO)
    return p


def body(doc, *ts):
    for t in ts:
        doc.add_paragraph(t)


def note(doc, t):
    p = doc.add_paragraph(t)
    for r in p.runs:
        r.font.size = Pt(9)


def table(doc, header, rows, right_from=99):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, v in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = str(v)
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(8.5)
    for row in rows:
        cs = t.add_row().cells
        for i, v in enumerate(row):
            cs[i].text = "" if v is None else str(v)
            for p in cs[i].paragraphs:
                if i >= right_from:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for r in p.runs:
                    r.font.size = Pt(8.5)


def style_head(ws, row=1):
    for c in ws[row]:
        if c.value is None:
            continue
        c.fill = HEAD
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        c.border = BORDER


def xbody(ws, first=2, wrap=()):
    for row in ws.iter_rows(min_row=first):
        for c in row:
            if c.value is not None:
                c.border = BORDER
            c.font = Font(size=c.font.size or 9, bold=bool(c.font.bold))
            c.alignment = Alignment(vertical="top", wrap_text=(c.column in wrap))


def widths(ws, w):
    for i, v in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = v


def notes(ws, lines):
    ws.append([])
    for t in lines:
        ws.append([t])
        ws.cell(ws.max_row, 1).font = Font(size=9, italic=True)


# ---------------------------------------------------------------- Word

def build_docx(MK, ZI, SU, SG):
    sk_n, sk_mi = soan_seika()
    fk = next(x for x in SU if x["都道府県"] == KEN)
    avg7 = statistics.mean(x["R7"] for x in SU)
    avg8 = statistics.mean(x["R8"] for x in SU)
    sagari = sum(1 for x in SU if x["R8"] < x["R7"])
    ups = sorted(SU, key=lambda x: -(x["R8"] - x["R7"]))
    up_rank = [x["都道府県"] for x in ups].index(KEN) + 1

    doc = new_doc()
    head(doc, "保険者機能強化推進交付金等の評価結果と福島県の位置", 0)
    body(doc, f"小野町高齢者福祉サービス推進協議会　資料　／　{ASOF_JP}")

    head(doc, "1　この資料の位置づけ", 1)
    body(doc,
         "**交付金の評価指標は、評価結果を外部の関係者を含む議論の場で"
         "検証することを求めています。**"
         "評価結果を協議会にお諮りすること自体が要件にあたるため、"
         "計画本文には掲載せず、毎年度の協議会資料として整理します。"
         "計画本文には、協議会へ報告し検証する手順のみを定めます。",
         f"**{ASOF_JP}に、令和8年度の評価指標と全国集計が公表されました。**"
         "ただし受領したのは**都道府県分**です。"
         "**市町村分の評価指標と小野町の得点票は含まれていません。**"
         "小野町自身の得点分析はそれらの受領後に行います。"
         "本資料では、福島県の位置を確かめ、"
         "県の支援が期待できる領域とそうでない領域を分けます。")

    th = sorted([y for y in SU if y["都道府県"] in TOHOKU],
                key=lambda z: z["R8順位"])
    th_rank = [x["都道府県"] for x in th].index(KEN) + 1
    th_txt = ("東北6県では最下位です" if th_rank == len(th) else
              f"東北6県では{th_rank}番目で、"
              f"上位の{th[0]['都道府県']}とは{th[0]['R8'] - fk['R8']:,.0f}点の"
              f"開きがあります")

    head(doc, f"2　福島県は令和7年度{fk['R7順位']}位から"
              f"令和8年度{fk['R8順位']}位へ上がった", 1)
    table(doc, ["区分", "令和7年度", "令和8年度", "差"],
          [("合計得点（800点満点）", f"{fk['R7']:,.0f}点", f"{fk['R8']:,.0f}点",
            f"＋{fk['R8'] - fk['R7']:,.0f}点"),
           ("全国順位", f"{fk['R7順位']}位", f"{fk['R8順位']}位",
            f"{fk['R7順位'] - fk['R8順位']:+d}"),
           ("得点率", f"{fk['R7'] / 800 * 100:.1f}％",
            f"{fk['R8'] / 800 * 100:.1f}％",
            f"＋{(fk['R8'] - fk['R7']) / 800 * 100:.1f}ポイント")], right_from=1)
    body(doc,
         f"**伸び幅＋{fk['R8'] - fk['R7']:.0f}点は47都道府県で{up_rank}位です。**"
         f"全国平均は{avg7:.1f}点から{avg8:.1f}点へ"
         f"＋{avg8 - avg7:.1f}点で、{sagari}団体は前年度を下回っています"
         f"（満点は両年度とも800点）。"
         f"ただし福島県は令和8年度の全国平均{avg8:.1f}点には届いておらず、"
         f"{th_txt}。")
    table(doc, ["県", "令和7年度", "令和8年度", "差", "令和8年度の順位"],
          [(x["都道府県"], f"{x['R7']:,.0f}点", f"{x['R8']:,.0f}点",
            f"{x['R8'] - x['R7']:+,.0f}点", f"{x['R8順位']}位") for x in th],
          right_from=1)

    head(doc, "3　どこで差がついているか", 1)
    table(doc, ["交付金", "目標", "配点", "福島県", "得点率", "全国平均", "順位"],
          [(m["区分"], m["項目"], f"{m['配点']}",
            f"{m[KEN]:,.0f}", f"{m['得点率'] * 100:.1f}％",
            f"{m['全国平均']:.1f}", f"{m['順位']}位") for m in MK],
          right_from=2)
    tan = [m for m in MK if "合計" not in m["項目"]]
    weak = sorted(tan, key=lambda x: -x["順位"])[:3]
    # 成果指標（目標Ⅳ）は推進・支援で同じ値をとる。並べると同じものが二つ出るので
    # 目標名・順位・得点が一致するものは一つに畳む。
    strong, seen = [], set()
    for m in sorted(tan, key=lambda x: x["順位"]):
        key = (m["項目"], m["順位"], m[KEN])
        if key in seen:
            continue
        seen.add(key)
        strong.append(m)
        if len(strong) == 2:
            break
    w0 = weak[0]
    s_seika = next((m for m in strong if "成果" in m["項目"]), None)
    body(doc,
         "**弱いのは"
         + "、".join(f"{lab(m)}（{m['順位']}位）" for m in weak)
         + "です。**"
         f"とくに{lab(w0)}（{w0['項目'].split('　')[1]}）は{w0['順位']}位で、"
         f"全国平均{w0['全国平均']:.1f}点に対し{w0[KEN]:,.0f}点にとどまります。",
         "**強いのは"
         + "、".join(
             f"{m['項目'].split('　')[0]}（成果。推進・支援とも{m['順位']}位）"
             if "成果" in m["項目"] else f"{lab(m)}（{m['順位']}位）"
             for m in strong)
         + "です。**"
         + (f"成果指標（目標Ⅳ）は推進・支援で同じ値をとり、"
            f"全国平均{s_seika['全国平均']:.1f}点に対し{s_seika[KEN]:,.0f}点です。"
            if s_seika else "")
         + "**県内の市町村の実績が良いことを示しており、"
           "小野町の取組もここに効きます。**")

    head(doc, "4　得点の動かし方は3種類に分かれる", 1)
    body(doc,
         "**各目標は（ⅰ）体制・取組 と（ⅱ）活動 に分かれ、"
         "目標Ⅳはこれとは別の成果指標群です。**"
         "他町村の案件で整理したところ、この3つは得点の動かし方が違います。")
    table(doc, ["指標群", "問われること", "得点の動かし方"], [
        ("（ⅰ）体制・取組",
         "仕組みがあるか（計画への記載・要綱・会議体）",
         "**計画に書けば翌年度の取組として得点になる。**"
         "計画の書きぶりが直に効くのはここ"),
        ("（ⅱ）活動", "実際に行ったか（回数・人数・件数）",
         "事業の実施そのもの。実績の積上げに時間がかかる"),
        ("成果指標群（目標Ⅳ）", "結果が出ているか（要介護度の変化等）",
         "**単年度では動かず、3年程度の遅れで現れる**"),
    ])
    body(doc,
         "**評価年度と交付年度は1年ずれます。**"
         "令和8年度の交付金は令和7年度の取組に対する評価です。"
         "**したがって本計画（令和9年度〜令和11年度）に書いたことが"
         "最初に得点に現れるのは令和10年度の交付金です。**"
         "計画期間の3年のうち評価に現れるのは後半の2年分にとどまるため、"
         "体制・取組にあたる事項は計画の初年度から書いておく必要があります。")
    table(doc, ["交付金", "目標", "（ⅰ）体制・取組", "（ⅱ）活動"],
          [(s["区分"], s["目標"],
            f"{s['体制・取組'][KEN]:,.0f}／{s['体制・取組']['配点']}点"
            f"（{s['体制・取組']['順位']}位）",
            f"{s['活動'][KEN]:,.0f}／{s['活動']['配点']}点"
            f"（{s['活動']['順位']}位）") for s in SG], right_from=2)
    mantens = [s for s in SG if s["体制・取組"][KEN] == s["体制・取組"]["配点"]]
    ryoho = [s for s in SG
             if s["体制・取組"]["順位"] >= 35 and s["活動"]["順位"] >= 35]
    katsudo_yowai = [s for s in SG
                     if s["体制・取組"][KEN] / s["体制・取組"]["配点"] >= 0.9
                     and s["活動"][KEN] / s["活動"]["配点"] < 0.5]
    body(doc,
         "**福島県は、体制は整っているが活動が伴っていない、という形をしています。**"
         + (f"（ⅰ）体制・取組が満点なのは"
            + "、".join(f"{s['区分']}{s['目標'].split('　')[0]}" for s in mantens)
            + f"の{len(mantens)}つですが、"
            if mantens else "")
         + (f"（ⅰ）が配点の9割以上でありながら（ⅱ）活動が配点の半分に届かないのは"
            + "、".join(s["区分"] + s["目標"].split("　")[0]
                        for s in katsudo_yowai)
            + f"の{len(katsudo_yowai)}つです。"
            if katsudo_yowai else "")
         + (f"**一方、{'・'.join(s['区分'] + s['目標'].split('　')[0] for s in ryoho)}"
            f"は体制・活動とも35位以下で、仕組みの段階から遅れています。**"
            if ryoho else ""),
         "**小野町にとっては、この差が「県に相談しても実務の支援は返ってこない」"
         "という形で現れます。**"
         "体制の整っている領域（要綱・会議体・計画への位置づけ）は県に照会できますが、"
         "実際の事業の進め方は町が自前で組み立てる前提で計画を書く必要があります。")
    note(doc, "※ 目標Ⅳ（成果指標群）はこの区分がないため掲げていない。"
              "※ これは都道府県分の得点であり、小野町の得点ではない。")

    head(doc, "5　全国の多くが得点していて福島県が0点の項目", 1)
    body(doc,
         f"評価指標の明細ごとに全国該当率を求め、"
         f"**全国の6割以上が得点しているのに福島県が0点である項目**を"
         f"洗い出すと{len(ZI)}件あります。"
         "他町村の案件で用いている方法です。"
         "取組の水準が低いのか、事業に着手していないのかを分けられます。")
    table(doc, ["交付金", "評価指標", "枝番", "配点", "全国該当率", "小野町に効くところ"],
          [("推進" if "推進" in z["交付金"] else "支援",
            z["評価指標"][:30], z["枝番"], f"{z['配点']:.0f}",
            f"{z['全国該当率'] * 100:.1f}％", z["小野町に効く"] or "―")
           for z in ZI[:14]], right_from=3)
    note(doc, f"※ 全{len(ZI)}件のうち全国該当率の高い14件。全件は別添のブックに。")

    head(doc, "6　小野町にとっての意味", 1)
    table(doc, ["領域", "福島県の状況", "小野町の計画への影響"], [
        ("認知症施策",
         "**「認知症施策に関する市町村支援」で全国の93.6％が得点するなか0点。**"
         "努力支援交付金 目標Ⅱ（認知症総合支援）も33位",
         "**県の支援を当てにせず、町が自前で組み立てる前提で書く。**"
         "認知症基本法第13条第1項は、都道府県計画があるときはそれも基本とする"
         "こととしているが、福島県の計画の策定状況は未確認（確認事項C8-3）"),
        ("**認知症の人・家族等の意見聴取**",
         "**市町村分の評価指標に「認知症の人及びその家族等の意見を踏まえた"
         "市町村認知症施策推進計画の策定に着手している」という項目が"
         "あるとされる（配点5点）。**"
         "受領したのは都道府県分であり、本件では未確認",
         "**確認事項C8-2（意見の聴き方）は、認知症基本法第13条第3項が"
         "準用する第12条第3項の努力義務であると同時に、"
         "交付金の得点の条件でもある可能性がある。**"
         "市町村分の受領後に確認し、"
         "意見聴取を実施する案の根拠として協議会にお示しする"),
        ("チームオレンジ",
         "「チームオレンジ設置状況」で全国の70.2％が得点するなか0点",
         "第4章③の認知症施策でチームオレンジを掲げているが、"
         "県内の整備が進んでいない。町単独で進める想定が要る"),
        ("認知症サポーターのステップアップ",
         "「管内の認知症サポーターステップアップ講座修了者数」で0点"
         "（全国該当率68.1％）",
         "**成果指標「認知症サポーター養成講座の受講者数」は"
         "養成数だけでなくステップアップ講座の修了者数で置くほうが"
         "交付金の評価に沿う**"),
        ("介護予防と保健事業の一体的実施",
         "環境整備のウ・エ・オがいずれも0点（全国該当率80〜83％）",
         "高齢者の保健事業と介護予防の一体的実施は国保部門との連携が要る。"
         "**県の広域的な支援が薄いため、町と国保連・後期高齢者医療広域連合との"
         "直接の連携で進める**"),
        ("地域包括支援センターの事業評価",
         "「管内の地域包括支援センター事業評価の達成状況」のア・イ・ウが0点"
         "（全国該当率68.1％）",
         "**センターの事業評価は市町村が行うもの。**"
         "町の評価の実施状況を確認し、第5章の進行管理に位置づける"),
        ("介護人材の確保",
         "目標Ⅲが45位。確保・定着の取組でオが0点（全国該当率91.5％）",
         "**第4章 基本目標2⑥（介護人材の確保）で県の施策に依存した書き方は避ける。**"
         "町内事業所の実態把握（依頼票B22・B23）を先に行う"),
        ("要介護度の変化（成果指標）",
         "推進・支援とも「長期的な要介護度の変化（要介護1・2）」で0点"
         "（全国該当率68.1％）。ただし目標Ⅳ全体は13位",
         "**要介護1・2の維持・改善は成果指標として国が見ている。**"
         "第4章の成果指標に置くかを協議会にお諮りする"),
    ])

    head(doc, "7　成果指向型配分枠（令和8年度に新設）", 1)
    body(doc,
         "**令和8年度から、保険者機能強化推進交付金に成果指向型配分枠"
         "（100点）が設けられました。**"
         "都道府県分の評価指標は「成果指向型の介護予防・健康づくりに関する"
         "取組を行う市町村に対する支援を行っているか」の1項目で、"
         "市町村側にこれに対応する取組があることが前提になっています。",
         "**この100点は、本資料2〜5で見た800点とは別の枠です。**"
         "受領した集計表は800点分（推進400点・支援400点）だけで、"
         "成果指向型配分枠の全国集計は含まれていません。"
         "したがって福島県の530点・26位に、この枠の分は入っていません。")
    table(doc, ["求められていること", "小野町の現状", "対応"], [
        ("地域のデータ分析に基づき課題を設定する",
         "**見える化システムの地域差指数と令和7年度調査で分析済み。**"
         "要介護認定率の地域差指数1.24（県内2位）、"
         "第1号被保険者1人あたり給付月額1.08（県内中位）",
         "素案 第2章3 に記載済み"),
        ("ターゲットとなる対象層を特定する（年齢・状態・性別等）",
         f"**未設定。**成果指標{sk_n}件はいずれも全体を対象としている",
         "**協議会にお諮りする。**"
         "調査では認知機能の低下が65〜69歳で39.1％、"
         "転倒不安が高いなど、対象層を切り出せる材料はある"),
        ("アウトプット指標とアウトカム指標を分けて設定する",
         f"**未分離。**成果指標{sk_n}件は両者が混在している",
         "**第4章の成果指標を2種類に整理し直す。**"
         "たとえば「認知症サポーター養成講座の受講者数」はアウトプット、"
         "「認知症に関する相談窓口の認知度」はアウトカムにあたる"),
        ("具体的な目標値を置く",
         f"**{sk_n}件のうち{sk_mi}件が【協議会で設定】のまま**",
         "**第2回協議会の主題とする。**目標値が確定しないと"
         "成果指向型配分枠の要件を満たせない"),
    ])
    body(doc,
         "**第4章の成果指標の整理は、計画の体裁の問題ではなく"
         "交付金の配分に直に効きます。**"
         "目標値の設定を協議会でご議論いただく際の材料として、"
         "この点を申し添えます。")

    head(doc, "8　協議会にお諮りする事項", 1)
    table(doc, ["No.", "事項", "受託者の案"], [
        ("1", "交付金の評価結果を計画本文に載せるか、協議会資料として扱うか",
         "**協議会資料として毎年度整理する。**"
         "計画本文には協議会へ報告し検証する手順のみを定める"),
        ("2", "第4章の成果指標をアウトプットとアウトカムに分けるか",
         "**分ける。**令和8年度に新設された成果指向型配分枠が両者の分離を"
         "求めている"),
        ("3", "成果指標の対象層を特定するか",
         "**特定する。**調査結果から切り出せる材料がある"),
        ("4", "「要介護1・2の長期的な要介護度の変化」を成果指標に加えるか",
         "**加えることを提案する。**国が見ている指標であり、"
         "福島県は0点だが目標Ⅳ全体は13位で県内市町村の実績は良い"),
        ("5", "認知症サポーターの成果指標を養成数からステップアップ修了者数に"
         "改めるか",
         "**改めることを提案する。**交付金の評価がそちらを見ている"),
    ])

    head(doc, "9　町にご確認いただきたいこと", 1)
    table(doc, ["#", "内容", "優先度"], [
        ("1", "**令和8年度の市町村分の評価指標と、小野町の得点票**", "A"),
        ("2", "令和6・7年度の小野町の得点票（3か年の推移を見るため）", "A"),
        ("3", "地域包括支援センターの事業評価の実施状況", "B"),
        ("4", "チームオレンジの整備状況（見える化では未整備）", "B"),
        ("5", "高齢者の保健事業と介護予防の一体的実施の実施状況", "B"),
        ("6", "福島県の認知症施策推進計画・介護人材確保の施策の状況", "B"),
    ])
    note(doc, "※ 1がそろえば、本資料と同じ方法で小野町自身の得点分析ができます。"
              "全国の多くの市町村が得点していて小野町が0点の項目を洗い出し、"
              "取組の水準が低いのか着手していないのかを分けられます。"
              "**とくに、0点でありながら全国の多くが得点している項目のうち、"
              "新たな事業も予算も要さず計画への記載だけで取れるものを"
              "切り分けることを想定しています。**"
              "（ⅰ）体制・取組の指標は計画の書きぶりで動くためです。")

    OUT.mkdir(parents=True, exist_ok=True)
    p = OUT / f"小野町_協議会資料_交付金の評価結果と福島県の位置_{ASOF}.docx"
    doc.save(p)
    print("出力:", p)


# ---------------------------------------------------------------- Excel

def build_xlsx(MK, ZI, SU, SG):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("00_福島県の位置")
    ws.append(["令和8年度 保険者機能強化推進交付金・介護保険保険者努力支援交付金"
               "（都道府県分）　福島県の位置"])
    ws.cell(1, 1).font = Font(bold=True, size=13)
    ws.append([f"作成日：{ASOF_JP}／出典：厚生労働省の公表資料（都道府県分）"])
    ws.append([])
    ws.append(["交付金", "目標", "配点", "福島県", "得点率", "全国平均", "順位",
               "団体数"])
    hr = ws.max_row
    for m in MK:
        ws.append([m["区分"], m["項目"], m["配点"], m[KEN], m["得点率"],
                   round(m["全国平均"], 1), m["順位"], m["団体数"]])
    style_head(ws, hr)
    for r in range(hr + 1, ws.max_row + 1):
        ws.cell(r, 5).number_format = "0.0%"
        v = ws.cell(r, 7).value
        ws.cell(r, 7).fill = NG if v >= 35 else WARN if v >= 24 else OK
    xbody(ws, first=hr, wrap=(2,))
    widths(ws, [8, 52, 8, 10, 10, 12, 8, 10])
    notes(ws, [
        "**※ 順位の色　赤＝35位以下／橙＝24〜34位／緑＝23位以上**",
        "※ 受領したのは都道府県分である。市町村分の評価指標と"
        "小野町の得点票は含まれていない。",
    ])

    ws = wb.create_sheet("01_令和7→8年度の推移")
    ws.append(["都道府県", "令和7年度", "順位", "令和8年度", "順位", "差", "順位の変化"])
    for x in sorted(SU, key=lambda y: y["R8順位"]):
        ws.append([x["都道府県"], x["R7"], x["R7順位"], x["R8"], x["R8順位"],
                   round(x["R8"] - x["R7"]), x["R7順位"] - x["R8順位"]])
    style_head(ws)
    for r in range(2, ws.max_row + 1):
        for c in (2, 4, 6):
            ws.cell(r, c).number_format = "#,##0"
        if ws.cell(r, 1).value == KEN:
            for c in range(1, 8):
                ws.cell(r, c).fill = WARN
        elif ws.cell(r, 1).value in TOHOKU:
            ws.cell(r, 1).fill = CALC
    xbody(ws)
    widths(ws, [14, 12, 8, 12, 8, 10, 12])
    ws.freeze_panes = "B2"

    ws = wb.create_sheet("02_体制・取組と活動の別")
    ws.append(["（ⅰ）体制・取組 と（ⅱ）活動 の別の得点（都道府県分・福島県）"])
    ws.cell(1, 1).font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["交付金", "目標",
               "（ⅰ）配点", "（ⅰ）福島県", "（ⅰ）全国平均", "（ⅰ）順位",
               "（ⅱ）配点", "（ⅱ）福島県", "（ⅱ）全国平均", "（ⅱ）順位"])
    hr = ws.max_row
    for s in SG:
        a, b = s["体制・取組"], s["活動"]
        ws.append([s["区分"], s["目標"],
                   a["配点"], a[KEN], round(a["全国平均"], 1), a["順位"],
                   b["配点"], b[KEN], round(b["全国平均"], 1), b["順位"]])
    style_head(ws, hr)
    for r in range(hr + 1, ws.max_row + 1):
        for cc in (6, 10):
            v = ws.cell(r, cc).value
            ws.cell(r, cc).fill = NG if v >= 35 else WARN if v >= 24 else OK
    xbody(ws, first=hr, wrap=(2,))
    widths(ws, [8, 30, 10, 12, 12, 8, 10, 12, 12, 8])
    notes(ws, [
        "**※（ⅰ）体制・取組は、計画に書けば翌年度の取組として得点になる。**"
        "（ⅱ）活動は実績の積上げを要する。",
        "※ 目標Ⅳ（成果指標群）はこの区分がないため掲げていない。",
        "**※ 評価年度と交付年度は1年ずれる。**"
        "令和8年度の交付金は令和7年度の取組に対する評価であり、"
        "第10期計画に書いたことが最初に得点に現れるのは令和10年度の交付金である。",
        "※ 順位の色　赤＝35位以下／橙＝24〜34位／緑＝23位以上",
    ])

    ws = wb.create_sheet("03_福島県が0点の項目")
    ws.append(["全国該当率が6割以上で福島県が0点の項目（大雪地区広域連合の方法による）"])
    ws.cell(1, 1).font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["列", "交付金", "目標", "評価指標", "枝番", "配点", "全国該当率",
               "得点した団体", "小野町に効くところ"])
    hr = ws.max_row
    for z in ZI:
        ws.append([z["列"], z["交付金"][:14], z["目標"], z["評価指標"], z["枝番"],
                   z["配点"], z["全国該当率"],
                   f"{z['得点した団体']}/{z['団体数']}", z["小野町に効く"]])
    style_head(ws, hr)
    for r in range(hr + 1, ws.max_row + 1):
        ws.cell(r, 7).number_format = "0.0%"
        ws.cell(r, 7).fill = NG if ws.cell(r, 7).value >= 0.8 else WARN
        if ws.cell(r, 9).value:
            ws.cell(r, 9).fill = CALC
    xbody(ws, first=hr, wrap=(3, 4, 9))
    widths(ws, [6, 16, 34, 40, 6, 6, 12, 12, 44])
    notes(ws, [
        "**※ 全国該当率＝その項目で得点した都道府県の数÷47。**"
        "配点を按分する指標があってもぶれないよう、金額でなく団体の数で取る。",
        "※ 全国の多くが得点しているのに0点である項目は、"
        "取組の水準が低いのではなく事業に着手していない可能性が高い。",
        "**※ これは福島県の分析である。小野町自身の分析には"
        "市町村分の評価指標と町の得点票が要る。**",
    ])

    p = OUT / f"小野町_交付金_福島県の得点分析_{ASOF}.xlsx"
    wb.save(p)
    print("出力:", p)


def selfcheck(ws, rows, MK, SU):
    """読み取った列が本当にその項目かを確かめる。

    列番号で拾っているため、ファイルの並びが変われば黙って別の値を読む。
    合計の突合で気づけるようにしておく。
    """
    ng = []
    if len(rows) != 47:
        ng.append(f"都道府県が{len(rows)}団体（47でない）")
    for c, _k, lb, hai in MOKUHYO:
        if ws.cell(9, c).value != hai:
            ng.append(f"列{c}の配点が{ws.cell(9, c).value}（{hai}のはず／{lb}）")
        tot = sum(ws.cell(r, c).value for r in rows.values())
        if tot != ws.cell(10, c).value:
            ng.append(f"列{c}の合計が47団体の和と不一致")
    # 目標Ⅰ〜Ⅳの和が交付金ごとの合計に一致するか
    for kbn, goukei in (("推進", 159), ("支援", 341)):
        wa = sum(m[KEN] for m in MK
                 if m["区分"] == kbn and "合計" not in m["項目"])
        gk = next(m[KEN] for m in MK if m["区分"] == kbn and "合計" in m["項目"])
        if wa != gk:
            ng.append(f"{kbn}の目標Ⅰ〜Ⅳの和{wa}が合計{gk}と不一致")
    fk = next(x for x in SU if x["都道府県"] == KEN)
    sk = sum(m[KEN] for m in MK if "合計" in m["項目"] and m["区分"] != "計")
    if sk != fk["R8"]:
        ng.append(f"推進＋支援{sk}が総合計{fk['R8']}と不一致")
    for x in SU:
        if not isinstance(x["R7"], (int, float)) or not isinstance(x["R8"], (int, float)):
            ng.append(f"{x['都道府県']}の得点が数値でない")
    print("  自己点検:", "適合" if not ng else "／".join(ng))
    return not ng


def main():
    ws, rows = load()
    MK, ZI, SU = mokuhyo(ws, rows), zero_items(ws, rows), suii(ws, rows)
    SG = shihyogun(ws, rows)
    selfcheck(ws, rows, MK, SU)
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx(MK, ZI, SU, SG)
    build_xlsx(MK, ZI, SU, SG)
    fk = next(x for x in SU if x["都道府県"] == KEN)
    print(f"  福島県 令和7年度{fk['R7']:.0f}点（{fk['R7順位']}位）"
          f"→ 令和8年度{fk['R8']:.0f}点（{fk['R8順位']}位）"
          f"／全国が6割以上得点して0点の項目 {len(ZI)}件")


if __name__ == "__main__":
    main()
