"""小野町 見える化システムの入力用データの整理と、地域支援事業に進めない件の解決策。

令和8年9月25日のご指示
  「現状の資料回収状況で確認できる内容で見える化システムの入力用のデータ整理を
    進めて下さい。合わせてシステム上の地域支援事業の見込量推計に進めない状況に
    ついて解決策を検討して下さい。」

**画面キャプチャ（同日受領）で分かった前提**
  介護保険事業状況報告の設定は、令和6年度＝年報、令和7年度＝月報12か月、
  令和8年度＝月報5月のみ。**令和8年度は1か月分である。**
  推計方法は基本推計、認定率の伸びは令和6年度→令和7年度、
  利用率と1人1月あたり利用回（日）数の伸びはいずれも0。
  施策反映の3画面は「自然体推計に全て戻す」がグレーで、自然体推計のまま。
  左のメニューは「地域支援事業の見込み量推計」以降がグレーで進めない。

**当方の見立て**
  令和8年度が5月の1か月であるため、その月に請求のなかったサービスが0になり、
  ワーニングが解消しないまま次段に進めていない可能性が高い。
  **解決策は、介護保険事業状況報告の設定で令和8年度のチェックを外すこと。**
  これで令和7年度（月報12か月＝完結年度）が最新となり、1か月の汚れが消える。

出力
  11_見える化出力依頼/小野町_地域支援事業の見込量推計に進めない件_YYYYMMDD.docx
  11_見える化出力依頼/小野町_見える化_地域支援事業の入力データ_YYYYMMDD.xlsx
"""

import pathlib
import warnings

warnings.filterwarnings("ignore")

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import build_ono_tanka as T
import ono_mieruka_ws as W
import ono_shizentai as SZ
from build_ono_kaisu import chiiki_plan, chiiki_ryo, sogo_jisseki

ROOT = pathlib.Path(__file__).parent
OUT = ROOT / "小野町_引継ぎ_整理済" / "11_見える化出力依頼"
ASOF = "20260925"
ASOF_JP = "令和8年9月25日"
JP_MIN = "游明朝"
JP_GO = "游ゴシック"

HEAD = PatternFill("solid", fgColor="1F3864")
NG = PatternFill("solid", fgColor="FCE4E4")
WARN = PatternFill("solid", fgColor="FFE0B2")
OK = PatternFill("solid", fgColor="E2EFDA")
CALC = PatternFill("solid", fgColor="EAF1FB")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SYS_D10 = 6466          # システムが表示している第10期（暫定値・月額）
OUR_D10 = 6210          # 当方の標準

# 逆算の結果（build_ono_mieruka_kakunin と同じ算式による）
GYAKU_ARI = 3230764     # 6,466円に合う標準給付費（地域支援事業費あり・千円）
GYAKU_NASHI = 3456325   # 同（地域支援事業費なし・千円）
WS_STD = 3101295        # 受領したワークシートの標準給付費見込額（千円）

# 画面から読み取れた事実。推測と分けて示す。
JIJITSU = [
    ("左のメニュー",
     "**「地域支援事業の見込み量推計」「保険料額の算定」"
     "「推計結果概要の確認」「都道府県への提出」がグレー**",
     "この4段に進めていない"),
    ("施策反映の3画面", "いずれも開ける。"
     "「自然体推計に全て戻す」はグレー",
     "**施策反映は入っていない。**画面までは到達している"),
    ("画面上部", f"**第10期 {SYS_D10:,}円（暫定値）／令和12年度 7,298円**",
     "**保険料は何らかの形で算定されている。**"
     "まったく進んでいないわけではない"),
    ("右上", "「総括表」ボタンは有効に見える",
     "**総括表は出力できる可能性がある**"),
    ("介護保険事業状況報告の設定",
     "**令和8年度は5月のみチェック**（令和7年度は12か月すべて）",
     "**令和8年度は1か月分。当方の見立ての根拠**"),
]

GENIN = [
    ("1", "**令和8年度が5月の1か月であるため、"
          "その月に請求のなかったサービスが0になっている**",
     "**最も可能性が高い。**"
     "見える化ワークシートでも令和8年度に0となった種別が3件ある"
     "（住宅改修費・介護予防住宅改修・介護予防短期入所生活介護）。"
     "0のサービスがワーニングとして残り、次段に進めない形になっている可能性",
     "**介護保険事業状況報告の設定で令和8年度のチェックを外す**"),
    ("2", "在宅サービスの施策反映の2つのサブ画面を通過していない",
     "在宅は「利用者数」と「利用回（日）数」の2画面がある。"
     "両方を開いて次へ進む操作をしないと解放されない形の可能性",
     "**利用回（日）数の画面まで開き、そこから"
     "「地域支援事業の見込み量推計」を押す**"),
    ("3", "ワーニングチェックが未解消",
     "原因1と重なる。他町村の案件ではワーニングチェックの結果が"
     "別途出力されている",
     "**ワーニングチェックの結果を表示し、内容をご共有ください**"),
    ("4", "地域支援事業費の実績が取り込まれていない",
     "**年報 様式4 の令和7年度が全て0で未入力である。**"
     "システムが実績を様式4から取るなら、見込みの基礎が作れない",
     "**令和7年度の決算額を町から受領し、様式4の入力状況を確認する**"),
]

TEJUN = [
    ("1", "**介護保険事業状況報告の設定で、令和8年度の5月のチェックを外す**",
     "**これが本筋。**令和7年度（月報12か月）が最新の完結年度になり、"
     "1か月の汚れが消える。"
     "当方の入力手順（令和8年9月25日版）のステップ2に記載済み"),
    ("2", "「保険料額の更新」を実行する", "設定の変更を反映させる"),
    ("3", "施策反映の3画面を、編集せずに順に通過する",
     "**在宅は「利用者数」と「利用回（日）数」の2画面がある。両方を開く**"),
    ("4", "「地域支援事業の見込み量推計」が押せるようになったか確かめる",
     "押せれば解決。押せなければ手順5へ"),
    ("5", "ワーニングチェックの結果を表示する",
     "**内容をご共有ください。**残っているワーニングが分かれば、"
     "次に何を直せばよいかを特定できる"),
    ("6", "右上の「総括表」ボタンで総括表を出力する",
     "**進めない場合でも、これは出せる可能性がある。**"
     f"出力できれば、{SYS_D10:,}円が何を積み上げた結果かが分かる"),
]

DAIAN = [
    ("計画書の作成", "**支障なし**",
     "当方の算定は完結している（保険料 月額6,210円・"
     "100円未満切上げ6,300円）。素案・見込量とも数値は入っている"),
    ("サービス見込量の推計（仕様書4(2)）", "**支障なし**",
     "年報 様式2 と国保連月報から自前で算定済み"),
    ("地域支援事業の量の見込み（介保法117条2項2号）", "**支障なし**",
     "国保連月報から3事業を算定済み。残り13事業は町の事業実績待ちで、"
     "これはシステムとは別の話"),
    ("自然体推計の出力（総括表）", "**要対応**",
     "**当方の算定との突合に用いる。**"
     "手順6で出力できるかを確かめる"),
    ("都道府県への提出", "**要対応**",
     "**システムを通さないとできない。**期限を町にご確認いただきたい"),
]


# ---------------------------------------------------------------- データ

def chiiki_input():
    """3_地域支援事業費に入れる値。現時点で用意できるものと町の資料が要るもの。"""
    cp = chiiki_plan()
    sg = sogo_jisseki()
    rows = []
    for kbn, name, r6, plan, memo in cp:
        r7 = sg.get("令和7年度", {}).get("事業", {}).get(name, (0, 0))[1]
        rows.append({"区分": kbn, "事業": name, "令和6年度": r6,
                     "令和7年度": r7 if r7 else None,
                     "第10期（各年度）": plan,
                     "出所": memo,
                     "状況": "用意できる" if "月報" in memo else "町の資料"})
    return rows


def ryo_input():
    """地域支援事業の「量」。介護保険法117条2項2号の必須記載事項。"""
    return [{"区分": r[0], "事業": r[1], "単位": r[2], "令和7年度": r[3],
             "第10期": r[4] if len(r) > 4 else None,
             "状況": "用意できる" if r[3] is not None else "町の事業実績"}
            for r in chiiki_ryo()]


def r8_sashikae():
    """令和8年度の実績見込み値を差し替えるための、サービス別の水準。

    システムの令和8年度は令和8年5月の1か月である。
    令和7年度の実績（見える化ワークシート）に差し替えれば1か月の汚れが消える。
    """
    mie, _tot, _ku, _hoken = W.read()
    g = SZ.r8_options()[0]["月報の実績"][0]
    out = []
    for kind in ("介護", "予防"):
        for name, d in mie[kind].items():
            nin = d.get("人数") or {}
            kyu = d.get("給付費") or {}
            n7, n8 = nin.get("令和7年度"), nin.get("令和8年度")
            k7, k8 = kyu.get("令和7年度"), kyu.get("令和8年度")
            if n7 is None and k7 is None:
                continue
            out.append({
                "区分": kind, "サービス": name,
                "人数_令和7年度": n7, "人数_令和8年度": n8,
                "人数_比": (n8 / n7) if (n7 and n8) else None,
                "給付費_令和7年度": k7, "給付費_令和8年度": k8,
                "給付費_比": (k8 / k7) if (k7 and k8) else None,
                "差し替え案_人数": n7,
                "差し替え案_給付費": (k7 * g) if k7 else None,
            })
    return out, g


# ---------------------------------------------------------------- 体裁

def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = JP_MIN
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), JP_MIN)
    return doc


def head(doc, text, level):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = JP_GO
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_GO)


def _emph(p, text):
    for i, part in enumerate(str(text).split("**")):
        if not part:
            continue
        r = p.add_run(part)
        r.font.name = JP_MIN
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_MIN)
        r.bold = bool(i % 2)


def body(doc, *paras):
    for t in paras:
        _emph(doc.add_paragraph(), t)


def note(doc, text):
    p = doc.add_paragraph()
    _emph(p, text)
    for r in p.runs:
        r.font.size = Pt(9)


def table(doc, header_row, rows, right_from=99):
    t = doc.add_table(rows=1, cols=len(header_row))
    t.style = "Table Grid"
    for i, h in enumerate(header_row):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = JP_GO
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_GO)
        r.bold = True
        r.font.size = Pt(9)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            _emph(p, v)
            for r in p.runs:
                r.font.size = Pt(9)
            if i >= right_from:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()


# ---------------------------------------------------------------- Word

def build_docx(CI, RI, SA, g):
    doc = new_doc()
    head(doc, "地域支援事業の見込み量推計に進めない件", 0)
    body(doc, f"小野町　第10期介護保険事業計画　／　{ASOF_JP}")

    head(doc, "0　結論", 1)
    table(doc, ["#", "内容"], [
        ("1", "**まず、介護保険事業状況報告の設定で令和8年度の5月のチェックを"
              "外してください。**"
              "これが本筋です。令和7年度（月報12か月＝完結年度）が最新となり、"
              "**令和8年5月の1か月がもたらす汚れが消えます。**"
              "当方の入力手順（同日版）のステップ2に記載済みのものです"),
        ("2", "**当方の見立ては、令和8年度が1か月であるために"
              "その月に請求のなかったサービスが0になり、"
              "ワーニングが解消しないまま次段に進めていない、というものです。**"
              "見える化ワークシートでも令和8年度に0となった種別が3件あります"),
        ("3", "**進めない間も、計画書の作成に支障はありません。**"
              "当方の算定は完結しており（保険料 月額6,210円）、"
              "サービス見込量も地域支援事業の量も自前で算定済みです。"
              "**システムが要るのは、総括表の出力と都道府県への提出の2つです**"),
        ("4", "**総括表は出力できる可能性があります。**"
              "右上の「総括表」ボタンは有効に見えます。"
              f"出力できれば、システムの{SYS_D10:,}円が何を積み上げた結果かが"
              "分かります（第4章）"),
    ])

    head(doc, "1　画面から読み取れること", 1)
    body(doc, "**推測と分けて、確かめられる事実から記します。**")
    table(doc, ["場所", "見えていること", "意味"],
          [(a, b, c) for a, b, c in JIJITSU])

    head(doc, "2　想定される原因", 1)
    table(doc, ["#", "原因", "根拠", "手だて"],
          [(a, b, c, d) for a, b, c, d in GENIN])
    body(doc,
         "**原因1を最も疑っています。**"
         "令和8年度が5月の1か月であることは設定画面で確定しており、"
         "1か月に請求のなかったサービスは0になります。"
         "**他町村の案件では、同じ状況で12種別が0になっていました。**"
         "小野町のワークシートでも3種別が0です。")

    head(doc, "3　手順", 1)
    table(doc, ["順", "すること", "留意点"],
          [(a, b, c) for a, b, c in TEJUN])
    body(doc,
         "**手順1で解決すれば、それ以上のことは要りません。**"
         "解決しない場合は、手順5のワーニングチェックの内容をお知らせください。"
         "**画面のキャプチャでも構いません。**"
         "残っているワーニングが分かれば、次に何を直せばよいかを特定できます。")

    head(doc, "4　総括表を出していただきたい理由", 1)
    body(doc,
         f"**システムが表示している第10期{SYS_D10:,}円が、"
         "何を積み上げた結果かが分かっていません。**"
         f"当方の算定は{OUR_D10:,}円で、差は{SYS_D10 - OUR_D10:+,}円です。",
         "**逆算すると次のようになります。**"
         "当方の算式に当てはめて、6,466円になる標準給付費を求めたものです。")
    table(doc, ["前提", "6,466円に合う標準給付費", "受領したワークシートの値との比"], [
        ("地域支援事業費を含む（当方と同じ191,749千円）",
         f"**{GYAKU_ARI:,}千円**", f"{GYAKU_ARI / WS_STD:.4f}"),
        ("地域支援事業費を含まない（0円）",
         f"**{GYAKU_NASHI:,}千円**", f"{GYAKU_NASHI / WS_STD:.4f}"),
    ], right_from=1)
    body(doc,
         f"**受領したワークシートの標準給付費見込額は{WS_STD:,}千円で、"
         "どちらとも一致しません。**"
         f"したがって、{SYS_D10:,}円が地域支援事業費を含んでいるのかどうかも、"
         "給付費をいくらで積んだのかも、現時点では分かりません。"
         "**総括表が出れば確定します。**")
    note(doc, "※ 逆算は当方の算式（第1号被保険者負担割合23％・収納率99.35％・"
              "調整交付金見込交付割合6.148％・補正後被保険者数9,797.58人）による。"
              "システムの前提が違えば結果も変わる。")

    head(doc, "5　進めない間にできること／できないこと", 1)
    table(doc, ["区分", "状況", "内容"],
          [(a, b, c) for a, b, c in DAIAN])

    head(doc, "6　入力用データの整理状況", 1)
    body(doc,
         "**現時点の資料で用意できるものを別冊のブックにまとめました。**"
         "町の資料が要るものと分けてあります。")
    yoi = sum(1 for r in CI if r["状況"] == "用意できる")
    machi = sum(1 for r in CI if r["状況"] != "用意できる")
    ryo_yoi = sum(1 for r in RI if r["状況"] == "用意できる")
    ryo_machi = sum(1 for r in RI if r["状況"] != "用意できる")
    table(doc, ["シート", "用意できる", "町の資料が要る", "内容"], [
        ("地域支援事業費", f"**{yoi}件**", f"{machi}件（内訳13欄）",
         "総合事業は国保連月報から事業別に割り付け済み。"
         "**町の資料が要るのは一般介護予防事業（計5,304,726円）と"
         "包括的支援事業・任意事業（計29,335,189円）の2つだが、"
         "ワークシート上は5欄と8欄に分けて入れる必要がある**"),
        ("地域支援事業の量", f"**{ryo_yoi}件**", f"{ryo_machi}件",
         "介護予防・生活支援サービス事業の3事業は月報から算定済み。"
         "一般介護予防事業・包括的支援事業・任意事業は町の事業実績が要る"),
        ("令和8年度の差し替え案", f"**{len(SA)}件**", "―",
         "**手順1（チェックを外す）で足りるが、"
         "外せない場合に備えてサービス別の水準を用意した**"),
    ], right_from=1)

    head(doc, "7　町にお願いしたいこと", 1)
    table(doc, ["#", "内容", "優先度"], [
        ("1", "**介護保険事業状況報告の設定で令和8年度のチェックを外し、"
              "地域支援事業に進めるようになるかをお試しください**", "A"),
        ("2", "**進めない場合、ワーニングチェックの結果**"
              "（画面のキャプチャで可）", "A"),
        ("3", "**総括表の出力**（右上の「総括表」ボタン）", "A"),
        ("4", "包括的支援事業・任意事業費29,335,189円の8欄内訳と、"
              "一般介護予防事業費5,304,726円の5事業別内訳", "A"),
        ("5", "令和7年度の介護保険特別会計決算"
              "（年報 様式4 が全て0で未入力）", "A"),
        ("6", "都道府県への提出の期限", "B"),
    ])

    OUT.mkdir(parents=True, exist_ok=True)
    p = OUT / f"小野町_地域支援事業の見込量推計に進めない件_{ASOF}.docx"
    doc.save(p)
    print("出力:", p)


# ---------------------------------------------------------------- Excel

def style_head(ws, row):
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row, c)
        cell.fill = HEAD
        cell.font = Font(bold=True, color="FFFFFF", size=9, name=JP_GO)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = BORDER


def xbody(ws, first, wrap=()):
    for r in range(first + 1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.font = Font(size=9, name=JP_MIN)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=c in wrap)


def widths(ws, w):
    for i, x in enumerate(w, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = x


def notes(ws, lines):
    ws.append([])
    for t in lines:
        ws.append([t])
        ws.cell(ws.max_row, 1).font = Font(size=9, italic=True, name=JP_MIN)


def build_xlsx(CI, RI, SA, g):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("00_手順")
    ws.append([f"小野町 見える化システム　地域支援事業に進むための手順　（{ASOF_JP}）"])
    ws.cell(1, 1).font = Font(bold=True, size=13, name=JP_GO)
    ws.append([])
    ws.append(["順", "すること", "留意点"])
    hr = ws.max_row
    for a, b, c in TEJUN:
        ws.append([a, b.replace("**", ""), c.replace("**", "")])
    style_head(ws, hr)
    ws.cell(hr + 1, 2).fill = WARN
    xbody(ws, hr, wrap=(2, 3))
    widths(ws, [6, 48, 66])
    notes(ws, [
        "**※ 手順1が本筋。令和8年度は令和8年5月の1か月であり、"
        "その月に請求のなかったサービスが0になる。**",
        "※ 手順1で解決しなければ、手順5のワーニングチェックの内容をご共有ください。",
    ])

    ws = wb.create_sheet("01_地域支援事業費")
    ws.append(["3_地域支援事業費に入れる値"])
    ws.cell(1, 1).font = Font(bold=True, size=12, name=JP_GO)
    ws.append([])
    ws.append(["区分", "事業", "令和6年度（円）", "令和7年度（円）",
               "第10期 各年度（円）", "状況", "出所・備考"])
    hr = ws.max_row
    for r in CI:
        ws.append([r["区分"], r["事業"], r["令和6年度"], r["令和7年度"],
                   r["第10期（各年度）"], r["状況"], r["出所"]])
    style_head(ws, hr)
    for r in range(hr + 1, ws.max_row + 1):
        for c in (3, 4, 5):
            ws.cell(r, c).number_format = "#,##0"
        ws.cell(r, 6).fill = OK if ws.cell(r, 6).value == "用意できる" else NG
    xbody(ws, hr, wrap=(2, 7))
    widths(ws, [14, 30, 16, 16, 18, 12, 54])
    notes(ws, [
        "**※ 緑＝当方で用意できる／赤＝町の資料が要る。**",
        "**※ 第10期は令和6年度決算の据え置き。**"
        "令和7年度は年報 様式4 が未入力で決算額が取れないため、月報で代用している。",
    ])

    ws = wb.create_sheet("02_地域支援事業の量")
    ws.append(["地域支援事業の「量」の見込み（介護保険法117条2項2号・必須記載事項）"])
    ws.cell(1, 1).font = Font(bold=True, size=12, name=JP_GO)
    ws.append([])
    ws.append(["区分", "事業", "単位", "令和7年度", "第10期", "状況"])
    hr = ws.max_row
    for r in RI:
        ws.append([r["区分"], r["事業"], r["単位"], r["令和7年度"],
                   r["第10期"], r["状況"]])
    style_head(ws, hr)
    for r in range(hr + 1, ws.max_row + 1):
        ws.cell(r, 6).fill = OK if ws.cell(r, 6).value == "用意できる" else NG
    xbody(ws, hr, wrap=(2,))
    widths(ws, [16, 34, 12, 14, 14, 14])

    ws = wb.create_sheet("03_令和8年度の差し替え案")
    ws.append(["令和8年度の実績見込み値を差し替える場合のサービス別の水準"])
    ws.cell(1, 1).font = Font(bold=True, size=12, name=JP_GO)
    ws.append([f"※ 手順1（令和8年度のチェックを外す）で足りる。"
               f"外せない場合の備え。給付費は令和7年度×{g:.4f}（月報3か月の前年同期比）"])
    ws.append([])
    ws.append(["区分", "サービス", "人数 令和7年度", "人数 令和8年度(現)",
               "人数 比", "給付費 令和7年度", "給付費 令和8年度(現)",
               "給付費 比", "差し替え案 人数", "差し替え案 給付費"])
    hr = ws.max_row
    for r in SA:
        ws.append([r["区分"], r["サービス"], r["人数_令和7年度"],
                   r["人数_令和8年度"], r["人数_比"], r["給付費_令和7年度"],
                   r["給付費_令和8年度"], r["給付費_比"],
                   r["差し替え案_人数"], r["差し替え案_給付費"]])
    style_head(ws, hr)
    for r in range(hr + 1, ws.max_row + 1):
        for c in (3, 4, 9):
            ws.cell(r, c).number_format = "#,##0.0"
        for c in (6, 7, 10):
            ws.cell(r, c).number_format = "#,##0"
        for c in (5, 8):
            ws.cell(r, c).number_format = "0.000"
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)) and abs(v - 1) > 0.20:
                ws.cell(r, c).fill = NG
            elif isinstance(v, (int, float)) and abs(v - 1) > 0.10:
                ws.cell(r, c).fill = WARN
        for c in (9, 10):
            ws.cell(r, c).fill = CALC
    xbody(ws, hr, wrap=(2,))
    widths(ws, [8, 30, 14, 16, 10, 16, 18, 10, 14, 16])
    ws.freeze_panes = "C2"
    notes(ws, [
        "**※ 比＝令和8年度（現在の1か月の値）÷令和7年度。**"
        "赤＝20％超の振れ／橙＝10％超。**1か月がどれだけ振れるかを示す。**",
        "**※ 差し替え案の人数は令和7年度の実績そのもの、"
        f"給付費は令和7年度×{g:.4f}（国保連月報3か月の前年同期比）。**",
        "※ 要介護度別の値は、現在システムに入っている令和8年度の構成比を保ったまま"
        "サービス別の比で調整するのが実務的である。",
    ])

    p = OUT / f"小野町_見える化_地域支援事業の入力データ_{ASOF}.xlsx"
    wb.save(p)
    print("出力:", p)


def main():
    CI, RI = chiiki_input(), ryo_input()
    SA, g = r8_sashikae()
    OUT.mkdir(parents=True, exist_ok=True)
    build_docx(CI, RI, SA, g)
    build_xlsx(CI, RI, SA, g)
    print(f"  地域支援事業費 用意できる"
          f"{sum(1 for r in CI if r['状況'] == '用意できる')}件／"
          f"町の資料{sum(1 for r in CI if r['状況'] != '用意できる')}件")
    print(f"  地域支援事業の量 用意できる"
          f"{sum(1 for r in RI if r['状況'] == '用意できる')}件／"
          f"町の事業実績{sum(1 for r in RI if r['状況'] != '用意できる')}件")
    furē = [r for r in SA if r["給付費_比"] and abs(r["給付費_比"] - 1) > 0.20]
    print(f"  令和8年度が令和7年度から20％超振れているサービス {len(furē)}件"
          f"／{len(SA)}件")


if __name__ == "__main__":
    main()
