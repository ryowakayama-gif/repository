# -*- coding: utf-8 -*-
"""業務工程管理表の更新（第2回策定委員会資料の確定に向けた点検：素案Ver.2.31・資料v16・問答集v13）
（令和8年10月23日・同日8度目）

  python3 07_ソーススクリプト/update_wbs_R8.10.23h.py

  01_業務内容別の進捗　(3) に追記し、進捗 0.80 → 0.85
  03_確認事項一覧　　　No.189 を起票（資料に残る【町確認】2件）。No.53・85・163 に追記
  04_成果品管理　　　　稿を Ver.2.31／v16／v13 に
  08_作業順位　　　　　48（完了）を加える

⚠ 進捗は作業項目のうち完了したものの割合を受託者が判断して記入する（01 注1）。
再実行しても二重に書かない。
"""
import copy
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月23日 更新⑧（第2回資料の確定に向けた点検）】"
MEMO = "川崎町_第2回策定委員会資料の確定に向けた点検_R8.10.23.md"
OKIKAE = [("川崎町_計画書素案_v2.30_協議用.docx", "川崎町_計画書素案_v2.31_協議用.docx"),
          ("川崎町_計画書案_公表用_v2.30.docx", "川崎町_計画書案_公表用_v2.31.docx"),
          ("川崎町_計画書案_概要版_v2.30.docx", "川崎町_計画書案_概要版_v2.31.docx"),
          ("川崎町_第2回策定委員会資料_R8.11_v15.docx", "川崎町_第2回策定委員会資料_R8.11_v16.docx"),
          ("同_想定問答集_R8.11_v12.docx", "同_想定問答集_R8.11_v13.docx")]

NO189 = [
    "189", "(3)", "R8.10",
    "第2回策定委員会資料に残る【町確認】2件",
    "第2回策定委員会資料 v16 に、町のご確認を要する【町確認】が2件残っています。10月30日（金）の"
    "確定までにご確認をお願いします。\n①３ 表5の注：第9期の在宅介護実態調査は要介護1〜5の認定者"
    "だけを対象としていたか（今回は要支援を含む）。\n②５ 選択３：町外の地域密着型サービス"
    "（定期巡回・随時対応型訪問介護看護、看護小規模多機能型居宅介護。令和7年度に月0.4人・0.2人の"
    "利用の実績）について、本町が事業所を指定しているか（区域外の指定）。なお、仙南圏域の事業所の"
    "所在は公表システムで確かめました（定期巡回2・看護小規模多機能2・小規模多機能4・夜間対応0）。",
    "第2回策定委員会資料・想定問答集（10月30日確定）\n計画素案 5-5",
    "発注者", "確認待ち", "R8.10.30",
    f"{HIZUKE}既定値＝回答がない場合、①は注の【町確認】を落とし出典だけを示す、②は「令和7年度に"
    f"町外の事業所の利用の実績がある」までを書き指定の状況の文を落とす。整理は 01_第10期_最新版成果品/{MEMO}。",
]

TSUIKI = {
    "85": (f"{HIZUKE}資料を v16・想定問答集を v13（全69問）に上げ、素案 Ver.2.31 と食い違う7件を"
           "解いた（重点課題5本を素案 3-4 に置いた、選択１の素案の説明、表37 の千円単位の値、27点の"
           "言い回し、表43 の介護保険法第117条第13項、事業一覧表の件数、選択３の近隣の事業所）。確定に"
           "向けてお願いしたいのは【町確認】2件（No.189）と事務局案のご了承（選択１Ｃ・選択５Ｂ・案Ⅰ・"
           f"頁数）。{MEMO}。　"),
    "53": (f"{HIZUKE}資料 v16（21頁）・想定問答集 v13（全69問）。10月30日に確定、11月10日に事前配布。"
           "紙・電子の別と部数のご指示をお願いします。　"),
    "163": (f"{HIZUKE}介護保険法の条文を e-Gov（令和8年10月1日施行時点）で確かめ、09_元資料/法令/ に"
            "退避した。被保険者の意見＝第117条第12項、都道府県の意見＝第13項、知事への提出＝第14項。"
            "素案 1-4 に項番号を入れ、資料の「第117条第５項」（誤り）を第13項に改めた。令和8年法律"
            "第51号の施行による項の移動は施行時期を確かめて照合し直す。　"),
}


def iro(c):
    c.font = FONT
    c.alignment = Alignment(wrap_text=True, vertical="top")


def mae_ni(c, s):
    if HIZUKE in str(c.value or ""):
        return 0
    c.value = s + str(c.value or "")
    iro(c)
    return 1


def main():
    wb = openpyxl.load_workbook(WBS)
    kekka = []

    ws1 = wb["01_業務内容別の進捗"]
    rows = {str(ws1.cell(y, 1).value or "").strip(): y for y in range(5, ws1.max_row + 1)}
    y = rows["(3)"]
    if mae_ni(ws1.cell(y, 9),
              f"{HIZUKE}第2回の資料を v16・想定問答集を v13 に上げ、素案と食い違う7件を解いた。"
              "確定（10月30日）に残るのは【町確認】2件（No.189）と事務局案のご了承。進捗 0.80→0.85。　"):
        ws1.cell(y, 7).value = 0.85
    kekka.append((ws1.cell(y, 7).value == 0.85, "(3) 0.85"))

    ws3 = wb["03_確認事項一覧"]
    k = 0
    for yy in range(5, ws3.max_row + 1):
        no = str(ws3.cell(yy, 1).value or "").strip()
        if no in TSUIKI:
            k += mae_ni(ws3.cell(yy, 10), TSUIKI[no])
    nos = {str(ws3.cell(yy, 1).value or "").strip() for yy in range(5, ws3.max_row + 1)}
    if "189" not in nos:
        moto = ws3.max_row
        r = moto + 1
        for c in range(1, 11):
            ws3.cell(r, c)._style = copy.copy(ws3.cell(moto, c)._style)
            ws3.cell(r, c).value = NO189[c - 1]
        ws3.row_dimensions[r].height = ws3.row_dimensions[moto].height
    nos = sorted(int(str(ws3.cell(yy, 1).value).strip()) for yy in range(5, ws3.max_row + 1)
                 if str(ws3.cell(yy, 1).value or "").strip().isdigit())
    kekka.append((k in (0, 3), f"確認事項への追記 {k}件"))
    kekka.append((nos == list(range(1, 190)), f"確認事項 1〜189・欠番なし（{len(nos)}件）"))

    ws4 = wb["04_成果品管理"]
    n = 0
    for yy in range(5, ws4.max_row + 1):
        na = str(ws4.cell(yy, 3).value or "")
        nb = na
        for a, b in OKIKAE:
            nb = nb.replace(a, b)
        if nb != na:
            ws4.cell(yy, 3).value = nb
            mae_ni(ws4.cell(yy, 8), f"{HIZUKE}稿を素案 Ver.2.31／公表用・概要版 v2.31／資料 v16／"
                                    "想定問答集 v13 に上げた（確定に向けた整合）。　")
            n += 1
    kekka.append((n in (0, 3), f"成果品の稿 {n}行"))

    ws8 = wb["08_作業順位"]
    aru = {str(ws8.cell(x, 1).value or "").strip() for x in range(5, ws8.max_row + 1)}
    if "48" not in aru:
        moto = ws8.max_row
        r = moto + 1
        for c in range(1, 11):
            ws8.cell(r, c)._style = copy.copy(ws8.cell(moto, c)._style)
        ws8.row_dimensions[r].height = ws8.row_dimensions[moto].height
        for c, v in enumerate([
                "48", "第2回策定委員会資料の確定に向けた点検", "R8.10.23",
                "ご指示による。資料は素案から起こすため、食い違いを素案の側で先に解いてから資料を揃えた",
                "No.85・53・163・189", "素案 Ver.2.31／公表用・概要版 v2.31／資料 v16／問答集 v13",
                "1日", "完了",
                "食い違い7件を解決。介護保険法の条文を e-Gov で確認・退避。仙南圏域の在宅系地域密着型の"
                "事業所を公表システムで確認。残るのは【町確認】2件（No.189）と事務局案のご了承。"], 1):
            ws8.cell(r, c).value = v
    aru = {str(ws8.cell(x, 1).value or "").strip() for x in range(5, ws8.max_row + 1)}
    kekka.append(("48" in aru, "作業順位 48"))

    wb.save(WBS)
    print("業務工程管理表の更新（第2回資料の確定に向けた点検）")
    ng = 0
    for ok, mi in kekka:
        print(("   ○ " if ok else "   × ") + mi)
        ng += not ok
    if ng:
        sys.exit(1)


if __name__ == "__main__":
    main()
