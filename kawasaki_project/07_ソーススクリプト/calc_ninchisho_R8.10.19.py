# -*- coding: utf-8 -*-
"""認知症の人の数の推計（令和8年10月19日）　— 算定の正本

  python3 07_ソーススクリプト/calc_ninchisho_R8.10.19.py

何をするものか
  国の先行研究による**年齢階級別の有病率**を、川崎町の**年齢階級別・
  男女別の人口**に当てて、認知症の人の数と軽度認知障害（ＭＣＩ）の
  人の数を積み上げる。第10期基本指針（案）別表一 事項十一が求める
  「認知症の人の数や推計値」のうち、**推計値**の側を算出する。

  ⚠ **高齢化率（65歳以上人口の割合）だけからは推計できない。**
  有病率は年齢階級によって 1.1％ → 50.3％ と大きく変わるため、
  要るのは65歳以上の「割合」ではなく、65歳以上の「年齢階級別の内訳」
  である。本町は年齢階級別・男女別の人口を町からご提供いただいている
  ため、積み上げを行える。

  ⚠ **65歳以上全体の12.3％を町の65歳以上人口に直接乗じてはならない。**
  12.3％は**性年齢調整後**の全国値であり、町の年齢構成を反映しない。
  本件では参考値（簡便法）として併記し、差が何に由来するかを示す。

  終了コード　0＝検査1〜8がすべて適合／1＝1件でも不適合

出所（いずれも原典を開いて読んだもの）
  ① 年齢階級別の有病率
     厚生労働省（老健局）「認知症および軽度認知障害（ＭＣＩ）の高齢者数
     と有病率の将来推計」1頁「年齢階級別の有病率（2022年時点）」
     https://www.mhlw.go.jp/content/001279920.pdf
     原データ：令和5年度老人保健事業推進費等補助金
     「認知症及び軽度認知障害の有病率調査並びに将来推計に関する研究」
     （研究代表者　九州大学　二宮利治）
     2022年に地域悉皆調査（調査率80％以上）を実施した4地域
     （福岡県久山町・石川県中島町・愛媛県中山町・島根県海士町）による
     ⚠ **同頁の年齢階級別の図はラスター画像であり、本文の文字としては
     取り出せない。** 300dpi で描画して目で読み取った（2度、別々の
     切り出しで読み、一致を確認した）。
  ② 川崎町の年齢階級別・男女別人口
     川崎町「年齢別人口集計表（Ｒ8.6月末時点）」（町ご提供）
     09_元資料/R7実績データ/Ｒ８年齢別人口ｎ.pdf
     ⚠ 同表は1歳刻み。総人口7,648人は**素案 2-2 の総人口と同じ**
     （確認事項No.175）。
  ③ 全国の年齢階級別人口（検査4に用いる）
     総務省統計局「人口推計（2022年10月1日現在）」参考表1
     年齢（5歳階級）別人口―総人口
     https://www.stat.go.jp/data/jinsui/2022np/zuhyou/05k2022-4.xlsx

⚠ **本件は「認知症の人の数」の2つの側のうち推計値の側だけである。**
  基本指針（案）第二 三 ６（五）イは「要支援者や要介護者の中での
  認知症の占める割合**や**先行研究等を参考に」と2つの参照元を並記する。
  前者（認定調査の日常生活自立度Ⅱ以上）は**実数**であり、町からの
  ご提供を待っている（確認事項No.157）。両者は測るものが違い、
  **差が「認定を受けていない認知症の方の規模」になる。**
"""
import os
import sys

# ══════════════════════════════════════════════════════════
#  1　年齢階級別の有病率（2022年時点・単位％）
#     出所①。順は 65-69／70-74／75-79／80-84／85-89／90以上。
#     ⚠ 値はグラフ画像のデータラベルを読み取ったもの。
# ══════════════════════════════════════════════════════════
KAIKYU = ["65〜69歳", "70〜74歳", "75〜79歳", "80〜84歳", "85〜89歳",
          "90歳以上"]

YUBYO = {
    # 指標 → 系列 → 階級ごとの率（％）
    "認知症": {
        "全体": [1.1, 3.1, 7.1, 16.6, 32.8, 50.3],
        "男性": [1.1, 2.8, 6.0, 15.9, 25.2, 36.6],
        "女性": [1.0, 3.1, 7.4, 16.9, 37.2, 55.1],
    },
    "ＭＣＩ": {
        "全体": [6.9, 9.3, 16.5, 22.6, 27.5, 21.2],
        "男性": [10.0, 12.1, 18.7, 22.9, 34.2, 32.5],
        "女性": [4.6, 5.9, 13.8, 21.7, 23.2, 17.0],
    },
}

# 公表されている全国の全体値（出所①1頁の表・2頁の図）
ZENKOKU_KOHYO = {
    "認知症": dict(nin_man=443.2, ritsu=12.3),   # 万人／％（令和4年）
    "ＭＣＩ": dict(nin_man=558.5, ritsu=15.5),
}
ZENKOKU_KOREI_MAN = 3603.0   # 65歳以上高齢者（万人。出所①2頁）
ZENKOKU_GOKEI_RITSU = 27.8   # 認知症＋ＭＣＩ（％。出所①2頁）

# ══════════════════════════════════════════════════════════
#  2　川崎町の年齢別人口（出所②　令和8年6月末現在・1歳刻み）
#     （年齢, 日本人男, 日本人女, 外国人男, 外国人女）
#     ⚠ 65歳以上のみを持つ。総人口・総数は TOTAL に持つ。
# ══════════════════════════════════════════════════════════
MACHI_1SAI = [
    (65, 66, 53, 0, 2), (66, 69, 60, 0, 2), (67, 68, 79, 0, 0),
    (68, 73, 75, 0, 1), (69, 71, 70, 0, 1),
    (70, 70, 71, 0, 0), (71, 79, 85, 0, 1), (72, 82, 79, 0, 1),
    (73, 88, 63, 0, 0), (74, 90, 87, 0, 0),
    (75, 83, 100, 0, 0), (76, 69, 98, 0, 0), (77, 90, 72, 0, 1),
    (78, 100, 73, 0, 0), (79, 52, 45, 1, 0),
    (80, 27, 37, 0, 0), (81, 35, 27, 0, 0), (82, 49, 51, 0, 0),
    (83, 26, 43, 0, 0), (84, 26, 33, 0, 0),
    (85, 29, 46, 0, 0), (86, 32, 48, 0, 1), (87, 19, 45, 0, 0),
    (88, 26, 47, 0, 0), (89, 23, 35, 0, 0),
    (90, 9, 28, 0, 0), (91, 14, 30, 0, 0), (92, 9, 28, 0, 0),
    (93, 8, 19, 0, 0), (94, 10, 14, 0, 0), (95, 3, 10, 0, 0),
    (96, 3, 12, 0, 0), (97, 3, 7, 0, 0), (98, 0, 6, 0, 0),
    (99, 2, 6, 0, 0), (100, 0, 9, 0, 0), (101, 0, 1, 0, 0),
    (102, 0, 2, 0, 0), (103, 1, 0, 0, 0),
]
MACHI_SOJINKO = 7648      # 総人口（出所②の合計行。素案2-2と同じ）
MACHI_JIKEN = "令和8年6月末現在"
MACHI_KOREIKA_SOAN = 42.0  # 素案 2-2 に掲げている高齢化率（％）

# 突合に用いる他の出所の65歳以上人口（時点・母集団が違う）
HIKAKU_KOREI = [
    ("介護保険事業状況報告（第1号被保険者数・令和7年度末）", 3240),
    ("地域包括ケア「見える化」システム 将来推計総括表"
     "（第1号被保険者数・令和8年度）", 3257),
]

# ══════════════════════════════════════════════════════════
#  3　全国の年齢階級別人口（検査4に用いる。出所③　千人）
#     人口推計（2022年10月1日現在）参考表1 総人口
#     ⚠ 90歳以上は 90〜94・95〜99・100歳以上を足す。
# ══════════════════════════════════════════════════════════
ZENKOKU_JINKO_SEN = None   # 読めたときに差し替える（検査4）
ZENKOKU_XLSX = ("/tmp/claude-0/-home-user-repository/"
                "4be8f82c-e2c9-52ac-b2a9-bad5154d0b13/scratchpad/"
                "ninchisho_R8.10.19/jinsui2022_5sai.xlsx")
# ⚠ 上は作業用の一時の場所である。**入力の第一はリポジトリの中**。
ZENKOKU_REPO = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                            "..", "09_元資料", "その他データ",
                            "人口推計R4.10.1_年齢5歳階級別_参考表1.xlsx")
# 読めなかったときに用いる引き写し（上の原典の2022年10月1日欄・千人）
ZENKOKU_HIKIUTSUSHI = [7535, 9376, 7384, 5626, 3775, 1912, 556, 90]
#   65-69／70-74／75-79／80-84／85-89／90-94／95-99／100以上


# ══════════════════════════════════════════════════════════
#  算定
# ══════════════════════════════════════════════════════════
def machi_kaikyu():
    """町の1歳刻みを6階級×男女に畳む。戻り　{階級: (男, 女, 計)}"""
    waku = [(65, 69), (70, 74), (75, 79), (80, 84), (85, 89), (90, 999)]
    out = {}
    for (lo, hi), na in zip(waku, KAIKYU):
        otoko = sum(a + c for (n, a, b, c, d) in MACHI_1SAI if lo <= n <= hi)
        onna = sum(b + d for (n, a, b, c, d) in MACHI_1SAI if lo <= n <= hi)
        out[na] = (otoko, onna, otoko + onna)
    return out


def santei(jinko, shihyo):
    """年齢階級別の積み上げ。

      jinko   {階級: (男, 女, 計)}
      shihyo  "認知症" 又は "ＭＣＩ"

    戻り　dict
      sei      性年齢階級別（男の率×男の人口＋女の率×女の人口）
      zentai   年齢階級別・男女計（全体の率×男女計の人口）
      kani     簡便法（65歳以上計×公表の全体の率）
      meisai   [(階級, 男人口, 女人口, 計人口, 男推計, 女推計, 計推計,
                 全体率での推計)]
    """
    r = YUBYO[shihyo]
    meisai = []
    sei = 0.0
    zentai = 0.0
    for i, na in enumerate(KAIKYU):
        o, n, k = jinko[na]
        mo = o * r["男性"][i] / 100.0
        mn = n * r["女性"][i] / 100.0
        mz = k * r["全体"][i] / 100.0
        sei += mo + mn
        zentai += mz
        meisai.append((na, o, n, k, mo, mn, mo + mn, mz))
    korei = sum(v[2] for v in jinko.values())
    kani = korei * ZENKOKU_KOHYO[shihyo]["ritsu"] / 100.0
    return dict(sei=sei, zentai=zentai, kani=kani, meisai=meisai,
                korei=korei)


def zenkoku_jinko():
    """検査4に用いる全国の年齢階級別人口（千人）。

    ⚠ **リポジトリの中を第一の入力とし**、無ければ作業用の複製、
      それも無ければ引き写しに落ちる。戻り　(値のならび, 出所の説明)
    """
    for p, doko in ((ZENKOKU_REPO, "リポジトリ 09_元資料"),
                    (ZENKOKU_XLSX, "作業用の複製")):
        if not os.path.exists(p):
            continue
        try:
            import openpyxl
            wb = openpyxl.load_workbook(p, data_only=True)
            ws = wb[wb.sheetnames[0]]
            # 見出し行（4行目）で 2022-10-01 の列を探す
            import datetime as dt
            retsu = None
            for c in range(1, ws.max_column + 1):
                v = ws.cell(4, c).value
                if isinstance(v, dt.datetime):
                    v = v.date()
                if isinstance(v, dt.date) and v == dt.date(2022, 10, 1):
                    retsu = c
                    break
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    d = dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
                    if d == dt.date(2022, 10, 1):
                        retsu = c
                        break
            if retsu is None:
                continue
            hoshii = ["65", "70", "75", "80", "85", "90", "95", "100"]
            atai = []
            for mato in hoshii:
                for r_ in range(5, 36):
                    na = str(ws.cell(r_, 2).value or "")
                    na = na.replace("　", " ").strip()
                    if na.replace(" ", "").startswith(mato):
                        atai.append(float(ws.cell(r_, retsu).value))
                        break
            if len(atai) == 8:
                return atai, doko
        except Exception:
            continue
    return list(ZENKOKU_HIKIUTSUSHI), "引き写し（原典の値をスクリプトに持つ）"


def zenkoku_kensan():
    """全国の年齢階級別人口に同じ式を当て、公表の全体値に戻るかを見る。"""
    sen, doko = zenkoku_jinko()
    # 90歳以上＝90-94＋95-99＋100以上
    j = sen[:5] + [sen[5] + sen[6] + sen[7]]
    korei = sum(j)
    out = {}
    for shihyo in ("認知症", "ＭＣＩ"):
        r = YUBYO[shihyo]["全体"]
        nin = sum(j[i] * r[i] / 100.0 for i in range(6))   # 千人
        out[shihyo] = dict(nin_man=nin / 10.0, ritsu=nin / korei * 100.0)
    return out, korei / 10.0, doko


# ══════════════════════════════════════════════════════════
#  本体
# ══════════════════════════════════════════════════════════
def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("認知症の人の数の推計（令和8年10月19日）")
    print("　川崎町 高齢者保健福祉計画・第10期介護保険事業計画")
    print()

    jinko = machi_kaikyu()
    korei = sum(v[2] for v in jinko.values())
    ninchi = santei(jinko, "認知症")
    mci = santei(jinko, "ＭＣＩ")

    # ── 川崎町の年齢階級別人口
    print(f"  ── 川崎町の年齢階級別人口（{MACHI_JIKEN}）")
    print(f"  {'階級':<10}{'男':>8}{'女':>8}{'計':>8}"
          f"{'65歳以上に占める割合':>22}")
    for na in KAIKYU:
        o, n, k = jinko[na]
        print(f"  {na:<10}{o:>8,}{n:>8,}{k:>8,}{k / korei * 100:>20.1f}％")
    print(f"  {'65歳以上計':<10}"
          f"{sum(v[0] for v in jinko.values()):>8,}"
          f"{sum(v[1] for v in jinko.values()):>8,}{korei:>8,}"
          f"{'100.0％':>22}")
    print(f"  総人口 {MACHI_SOJINKO:,}人／高齢化率 "
          f"{korei / MACHI_SOJINKO * 100:.1f}％"
          f"／75歳以上 {sum(jinko[na][2] for na in KAIKYU[2:]):,}人"
          f"（65歳以上の{sum(jinko[na][2] for na in KAIKYU[2:]) / korei * 100:.1f}％）")
    print()

    # ── 推計
    for shihyo, kekka in (("認知症", ninchi), ("ＭＣＩ", mci)):
        r = YUBYO[shihyo]
        print(f"  ── {shihyo}の人の数（推計）")
        print(f"  {'階級':<10}{'有病率 男':>10}{'女':>8}{'全体':>8}"
              f"{'男':>9}{'女':>9}{'計':>9}")
        for i, (na, o, n, k, mo, mn, mk, mz) in enumerate(kekka["meisai"]):
            print(f"  {na:<10}{r['男性'][i]:>9.1f}％{r['女性'][i]:>7.1f}％"
                  f"{r['全体'][i]:>7.1f}％{mo:>9.1f}{mn:>9.1f}{mk:>9.1f}")
        print(f"  {'計':<10}{'':>34}"
              f"{sum(v[4] for v in kekka['meisai']):>9.1f}"
              f"{sum(v[5] for v in kekka['meisai']):>9.1f}"
              f"{kekka['sei']:>9.1f}")
        print(f"   ▶ 性年齢階級別の積み上げ　　{kekka['sei']:,.0f}人"
              f"（65歳以上の{kekka['sei'] / korei * 100:.1f}％）")
        print(f"   ▶ 年齢階級別（男女計）　　　{kekka['zentai']:,.0f}人"
              f"（{kekka['zentai'] / korei * 100:.1f}％）")
        print(f"   △ 簡便法（65歳以上×"
              f"{ZENKOKU_KOHYO[shihyo]['ritsu']}％）　{kekka['kani']:,.0f}人"
              f"　⚠ 参考値。町の年齢構成を反映しない")
        print()

    gokei = ninchi["sei"] + mci["sei"]
    print(f"  ── 認知症とＭＣＩの合計　{gokei:,.0f}人"
          f"（65歳以上の{gokei / korei * 100:.1f}％）")
    print(f"     全国は約28％（令和4年。認知症12.3％＋ＭＣＩ15.5％）")
    print()

    # ══════ 検査1　有病率の表が揃っているか
    print("  ── 検査1　有病率の表（6階級×3系列×2指標＝36件）")
    kensu = sum(len(v) for s in YUBYO.values() for v in s.values())
    judge(kensu == 36, f"有病率が36件そろう（実際 {kensu}件）")
    hani = all(0.0 < x <= 100.0 for s in YUBYO.values()
               for v in s.values() for x in v)
    judge(hani, "いずれの率も0％超100％以下")

    # ══════ 検査2　形（認知症は単調増加／ＭＣＩは90歳以上で低下）
    print("  ── 検査2　階級による形")
    for kei in ("全体", "男性", "女性"):
        v = YUBYO["認知症"][kei]
        judge(all(v[i] < v[i + 1] for i in range(5)),
              f"認知症（{kei}）は階級が上がるごとに高くなる")
    for kei in ("全体", "男性", "女性"):
        v = YUBYO["ＭＣＩ"][kei]
        judge(v[5] < v[4],
              f"ＭＣＩ（{kei}）は90歳以上で85〜89歳より下がる",
              "⚠ 図にするときは注記が要る（単調でない）")

    # ══════ 検査3　町の人口の突合
    print("  ── 検査3　町の年齢階級別人口")
    for na in KAIKYU:
        o, n, k = jinko[na]
        if o + n != k:
            judge(False, f"{na} の男＋女が計と合わない")
    judge(all(jinko[na][0] + jinko[na][1] == jinko[na][2] for na in KAIKYU),
          "各階級で 男＋女＝計")
    judge(korei == 3209, f"65歳以上計が3,209人（実際 {korei:,}人）")
    koreika = korei / MACHI_SOJINKO * 100
    judge(abs(koreika - MACHI_KOREIKA_SOAN) < 0.1,
          f"高齢化率が素案2-2の{MACHI_KOREIKA_SOAN}％と一致"
          f"（算出 {koreika:.1f}％）")
    print("     （参考）他の出所の65歳以上人口との開き")
    for mi, v in HIKAKU_KOREI:
        print(f"       {mi[:46]:<46}{v:>7,}人　差 {korei - v:+,}人"
              f"（{(korei - v) / v * 100:+.1f}％）")

    # ══════ 検査4　全国に同じ式を当てて公表値に戻るか
    print("  ── 検査4　全国に同じ式を当て直す（年齢階級別の率の確かめ）")
    zk, zk_korei_man, doko = zenkoku_kensan()
    print(f"     全国の年齢階級別人口の出所：{doko}")
    print(f"     65歳以上 {zk_korei_man:,.0f}万人"
          f"（公表 {ZENKOKU_KOREI_MAN:,.0f}万人）")
    for shihyo in ("認知症", "ＭＣＩ"):
        a = zk[shihyo]
        k = ZENKOKU_KOHYO[shihyo]
        sa = a["ritsu"] - k["ritsu"]
        judge(abs(sa) <= 1.0,
              f"{shihyo}　算出 {a['ritsu']:.1f}％（{a['nin_man']:.0f}万人）"
              f"／公表 {k['ritsu']}％（{k['nin_man']}万人）　差 {sa:+.1f}ポイント",
              "⚠ 公表値は性年齢調整後。完全一致はしない")
    gokei_zk = zk["認知症"]["ritsu"] + zk["ＭＣＩ"]["ritsu"]
    judge(abs(gokei_zk - ZENKOKU_GOKEI_RITSU) <= 2.0,
          f"認知症＋ＭＣＩ　算出 {gokei_zk:.1f}％／公表 "
          f"{ZENKOKU_GOKEI_RITSU}％")

    # ══════ 検査5　性年齢階級別と男女計の差
    print("  ── 検査5　性年齢階級別と男女計の差")
    for shihyo, kekka in (("認知症", ninchi), ("ＭＣＩ", mci)):
        sa = kekka["sei"] - kekka["zentai"]
        judge(abs(sa) / kekka["sei"] * 100 < 5.0,
              f"{shihyo}　性年齢階級別 {kekka['sei']:.0f}人／男女計 "
              f"{kekka['zentai']:.0f}人　差 {sa:+.0f}人"
              f"（{sa / kekka['zentai'] * 100:+.1f}％）")

    # ══════ 検査6　簡便法との差（年齢構成の効き）
    print("  ── 検査6　簡便法（65歳以上×全国の全体の率）との差")
    for shihyo, kekka in (("認知症", ninchi), ("ＭＣＩ", mci)):
        sa = kekka["sei"] - kekka["kani"]
        muki = "低く" if sa < 0 else "高く"
        print(f"   ・{shihyo}　積み上げ {kekka['sei']:.0f}人／簡便法 "
              f"{kekka['kani']:.0f}人　差 {sa:+.0f}人"
              f"（{sa / kekka['kani'] * 100:+.1f}％）")
        print(f"     ⚠ 積み上げの方が{muki}出る。"
              f"町の65歳以上の年齢の内訳が全国と違うためである")
    judge(True, "簡便法との差を数えた（どちらが正かではなく、"
                "年齢構成の効きを示すもの）")

    # ══════ 検査7　推計値が人口を超えない・階級ごとに人口以下
    print("  ── 検査7　推計値が人口を超えていないか")
    ok = True
    for shihyo, kekka in (("認知症", ninchi), ("ＭＣＩ", mci)):
        for (na, o, n, k, mo, mn, mk, mz) in kekka["meisai"]:
            if mo > o or mn > n or mk > k:
                ok = False
    judge(ok, "いずれの階級でも推計値が人口以下")
    judge(gokei < korei, f"認知症＋ＭＣＩ {gokei:,.0f}人 ＜ "
                         f"65歳以上人口 {korei:,}人")

    # ══════ 検査8　将来値を置いていないこと
    print("  ── 検査8　置いていないもの（未実施として残す）")
    print("   △ **将来（令和12・17・22年度）の推計は行っていない。**")
    print("      年齢階級別の**将来**人口が要る。当方の環境からは")
    print("      国立社会保障・人口問題研究所（www.ipss.go.jp）に接続できず、")
    print("      地域包括ケア「見える化」システムは町が進める範囲である。")
    print("      ⚠ 「未実施」であり「適合」ではない。")
    print("   △ **認定調査の日常生活自立度Ⅱ以上の実数は未受領**"
          "（確認事項No.157）。")
    print("      推計値（本件）と実数は測るものが違い、"
          "差が未認定の方の規模になる。")
    judge(True, "未実施の2件を件数と理由として残した")

    print()
    print("  ── まとめ")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        for m in futeki:
            print(f"      {m}")
        sys.exit(1)
    print("   ○ 検査1〜8　不適合0件")
    print()
    print(f"   ▶ 認知症の人の数（推計）　{ninchi['sei']:,.0f}人"
          f"　65歳以上の{ninchi['sei'] / korei * 100:.1f}％")
    print(f"   ▶ ＭＣＩの人の数（推計）　{mci['sei']:,.0f}人"
          f"　65歳以上の{mci['sei'] / korei * 100:.1f}％")


if __name__ == "__main__":
    main()
