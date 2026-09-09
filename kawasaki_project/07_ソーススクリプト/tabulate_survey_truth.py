# -*- coding: utf-8 -*-
"""アンケート個票（集計データブックの「データ」シート）から、
レビュー用の正解値（単純集計・クロス集計）を機械的に作る。

シート構造
  行1  設問グループ名（A問1 など）が、その設問の先頭列に置かれている
  行2  設問文
  行3  選択肢ラベル（n 列は空）
  行4  '①'…（選択肢）／'n'（有効回答フラグ列）
  行5〜 個票。選択肢セルは '〇'、n 列は 1

使い方
  python3 tabulate_survey_truth.py <xlsx> <出力json> [出力txt]
"""
import json
import sys
from collections import OrderedDict

from decimal import Decimal, ROUND_HALF_UP

import openpyxl

MARK = ("〇", "○", "◯", "1", 1)



def pct1(k, n):
    """百分率を小数第1位で四捨五入する（Python の round は偶数丸めになるため使わない）"""
    if not n:
        return None
    return float(Decimal(k * 100) / Decimal(n)
                 == 0 and 0 or (Decimal(k * 100) / Decimal(n)).quantize(
                     Decimal("0.1"), rounding=ROUND_HALF_UP))


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["データ"]
    ncol = ws.max_column
    nrow = ws.max_row

    # 設問の切り出し：行4が '①'〜 で始まり 'n' で終わる連なりを1設問とする
    qs = []
    cur = None
    for c in range(1, ncol + 1):
        g = ws.cell(1, c).value
        k4 = ws.cell(4, c).value
        if g:
            gname = str(g).strip()
        if k4 is None:
            continue
        k4 = str(k4).strip()
        if k4 == "n":
            if cur:
                cur["ncol"] = c
                qs.append(cur)
                cur = None
            continue
        if cur is None:
            cur = {
                "group": gname,
                "text": str(ws.cell(2, c).value or "").strip(),
                "opts": [],
                "cols": [],
            }
        cur["opts"].append(str(ws.cell(3, c).value or "").strip())
        cur["cols"].append(c)
    if cur:
        cur["ncol"] = None
        qs.append(cur)

    rows = []
    for r in range(5, nrow + 1):
        v = ws.cell(r, 1).value
        if v is None:
            continue
        if str(v).strip() in ("計", "合計", "計）"):   # 末尾の集計行を除く
            continue
        rows.append(r)

    # 属性列（行1に見出し、行4が空）
    attrs = OrderedDict()
    for c in range(1, ncol + 1):
        h = ws.cell(1, c).value
        if h and ws.cell(4, c).value is None and ws.cell(3, c).value is None:
            attrs[str(h).strip()] = c

    return ws, qs, rows, attrs


def tabulate(ws, qs, rows):
    out = []
    for q in qs:
        counts = []
        for c in q["cols"]:
            k = sum(1 for r in rows if ws.cell(r, c).value in MARK)
            counts.append(k)
        if q["ncol"]:
            n = sum(1 for r in rows
                    if ws.cell(r, q["ncol"]).value not in (None, "", 0))
        else:
            n = sum(1 for r in rows
                    if any(ws.cell(r, c).value in MARK for c in q["cols"]))
        total = sum(counts)
        multi = total > n
        out.append({
            "group": q["group"],
            "text": q["text"],
            "n": n,
            "total_selections": total,
            "multi": multi,
            "options": [
                {"label": o, "count": k,
                 "pct": pct1(k, n) if n else None}
                for o, k in zip(q["opts"], counts)
            ],
        })
    return out


def crosstab(ws, qs, rows, attr_col, q):
    """属性列 × 設問 のクロス集計"""
    levels = OrderedDict()
    for r in rows:
        v = ws.cell(r, attr_col).value
        v = str(v).strip() if v is not None else "（無回答）"
        levels.setdefault(v, []).append(r)
    res = {}
    for lv, rs in levels.items():
        if q["ncol"]:
            n = sum(1 for r in rs
                    if ws.cell(r, q["ncol"]).value not in (None, "", 0))
        else:
            n = sum(1 for r in rs
                    if any(ws.cell(r, c).value in MARK for c in q["cols"]))
        cells = []
        for o, c in zip(q["opts"], q["cols"]):
            k = sum(1 for r in rs if ws.cell(r, c).value in MARK)
            cells.append({"label": o, "count": k,
                          "pct": pct1(k, n) if n else None})
        res[lv] = {"n": n, "options": cells}
    return res


def main():
    src, dst = sys.argv[1], sys.argv[2]
    ws, qs, rows, attrs = load(src)
    tab = tabulate(ws, qs, rows)
    data = {"source": src, "records": len(rows),
            "attributes": list(attrs.keys()), "questions": tab}
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)

    if len(sys.argv) > 3:
        lines = [f"個票件数（全体の母数）：{len(rows)}件", f"元ファイル：{src}", ""]
        for t in tab:
            lines.append(f"■ {t['group']}　{t['text']}")
            lines.append(f"　n={t['n']}　選択延べ{t['total_selections']}"
                         f"{'（複数回答）' if t['multi'] else ''}")
            for o in t["options"]:
                lines.append(f"　　{o['label']}：{o['count']}件 "
                             f"{o['pct']}%")
            lines.append("")
        with open(sys.argv[3], "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
    print(f"{len(rows)}件 / 設問{len(tab)}個 → {dst}")
    print("属性列：", list(attrs.keys()))


if __name__ == "__main__":
    main()
