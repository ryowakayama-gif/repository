# -*- coding: utf-8 -*-
"""第2回策定委員会資料 v17 → v18（Red Team 再レビュー（令和8年10月24日）の指摘の是正）

  python3 07_ソーススクリプト/build_2kai_v18.py

再レビュー原本　09_元資料/レビュー/川崎町_第10期計画_RedTeam再レビュー_R8.10.24受領.md
内容の正本　　　協議用素案 Ver.2.33（fix_soan_v2330_saireview.py）

  N01 表26 の算定Ｂ・Ｃの改定率1％・3％（6,235円・5,583円・6,366円・5,714円）は素案になく、
       基金の取崩しを固定した計算とも合わなかった。素案 9-3 の改定率別の文から起こす。
       表27 の報酬改定の行の算定Ｂ・Ｃの差、表28 の注記（「報酬改定・負担割合で差が異なる」は誤り。
       取崩しは一定額の差引きであり、差が変わるのは予定収納率のみ）を直す
  N02 本日のご審議事項⑥「ご意見をいただくことが『評価結果の活用』の要件となります」の言い切り
  N04 表29「介護給付費準備基金の令和８年度末残高」→ 見込額（決算による確定は令和９年度）
  N07 審議事項の囲みが⑦で切れる（⑦に枠がない）。表と表の間に段落がなく Word が1つの表として扱い、
       見出し行が次頁で繰り返されて「す。」だけが残る。表の間に区切りの段落を置き、行を分割しない
  N08 「全国1,741保険者」等 → 市町村（交付金の評価の単位は市町村）
  N09 入所・居住系10人 約211円 → 素案 5-5 の約212円（1円未満四捨五入）
  N10 「同時に進んだ結果と考えられます」→ 影響している可能性。因果は追加の検証を要する
"""
import copy
import importlib.util
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)


def load(name, fn):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(spec)
    sys.modules[name] = m
    spec.loader.exec_module(m)
    return m


V9 = load("build_2kai_v9", "build_2kai_v9.py")
SRC = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v17.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v18.docx"
SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.33_協議用.docx"


def kakikae(p, s):
    p.runs[0].text = V9.clean(s)
    for r in p.runs[1:]:
        r.text = ""


def yen(s):
    return int(re.sub(r"[^\d]", "", s))


def zen(d):
    return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])


def main():
    futeki = []
    KYU = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会資料 v17 → v18（Red Team 再レビューの指摘の是正）")
    S = docx.Document(SOAN)
    sz = zen(S)
    b = yen(re.search(r"パターンB ([\d,]+)円", sz).group(1))
    j10 = int(re.search(r"10人では約(\d+)円動きます", sz).group(1))
    m = re.search(r"改定率1％では算定Ａ([\d,]+)円・算定Ｂ([\d,]+)円・算定Ｃ([\d,]+)円（それぞれ＋(\d+)円・＋(\d+)円・＋(\d+)円）、"
                  r"3％では算定Ａ([\d,]+)円・算定Ｂ([\d,]+)円・算定Ｃ([\d,]+)円（それぞれ＋(\d+)円・＋(\d+)円・＋(\d+)円）", sz)
    judge(m is not None, "素案 9-3 の改定率別の算定Ａ・Ｂ・Ｃの文")
    g = [yen(x) for x in m.groups()]
    i1, s1, i3, s3 = g[0:3], g[3:6], g[6:9], g[9:12]
    judge(True, f"素案から：算定Ｂ {b:,}円・10人 約{j10}円・1％ {i1}（{s1}）・3％ {i3}（{s3}）")

    shutil.copy2(SRC, DST)
    d = docx.Document(DST)
    mae = zen(d)

    def p_rep(atama, kyu, shin, n=1):
        h = [p for p in d.paragraphs if p.text.strip().startswith(atama) and kyu in p.text]
        judge(len(h) == n, f"『{atama[:18]}』の「{kyu[:16]}」（{len(h)}か所）")
        for p in h:
            KYU.append(p.text)
            kakikae(p, p.text.replace(kyu, shin))

    def c_rep(kyu, shin):
        n = 0
        for t in d.tables:
            seen = []
            for r in t.rows:
                for c in r.cells:
                    if c._tc in seen:
                        continue
                    seen.append(c._tc)
                    if kyu in c.text:
                        KYU.append(c.text)
                        V9.set_cell(c, c.text.replace(kyu, shin))
                        n += 1
        return n

    print("  ── N01 表26・表27・表28（報酬改定率別の算定Ａ・Ｂ・Ｃ）")
    t26 = [t for t in d.tables if t.rows[0].cells[0].text.strip() == "介護報酬改定率"][0]
    KYU.append("\n".join(c.text for r in t26.rows for c in r.cells))
    for r in t26.rows[1:4]:
        k = r.cells[0].text.strip()[:2]
        if k == "１％":
            vals = i1
        elif k == "３％":
            vals = i3
        else:
            continue
        judge(yen(r.cells[1].text) == vals[0], f"表26 {k} 算定Ａ {vals[0]:,}円は v17 と同じ")
        for j in (1, 2):
            V9.set_cell(r.cells[1 + j], f"{vals[j]:,}円")

    def sa_bun(s):
        if s[1] == s[2]:
            return f"＋{s[0]}円（算定Ｂ・Ｃでは＋{s[1]}円）"
        return f"＋{s[0]}円（算定Ｂでは＋{s[1]}円、算定Ｃでは＋{s[2]}円）"

    t27 = [t for t in d.tables if [c.text.strip() for c in t.rows[0].cells] == ["要　因", "影　響"]][0]
    n = 0
    for r in t27.rows:
        k = r.cells[0].text.strip()
        if k == "介護報酬改定率が１％だった場合":
            KYU.append(r.cells[1].text)
            V9.set_cell(r.cells[1], sa_bun(s1))
            n += 1
        elif k == "同 ３％だった場合":
            KYU.append(r.cells[1].text)
            V9.set_cell(r.cells[1], sa_bun(s3))
            n += 1
        elif k == "入所・居住系の利用者が10人増えた場合":
            KYU.append(r.cells[1].text)
            V9.set_cell(r.cells[1], f"＋{j10}円")
            n += 1
        elif k == "同 10人減った場合":
            KYU.append(r.cells[1].text)
            V9.set_cell(r.cells[1], f"▲{j10}円")
            n += 1
    judge(n == 4, f"表27 報酬改定・入所・居住系の行（{n}行）")
    KYU28 = ("上の差額は、基金の取崩し・入所・居住系の人数を除き、算定Ａを基準とした値です。取崩しを行う算定Ｂ・Ｃでは、"
             "保険料に掛かる割合で動く要因（報酬改定・負担割合・収納率）の差がこれと異なります。")
    SHIN28 = ("上の差額は、算定Ａ（取崩しなし）を基準とした値です。基金の取崩しは保険料収納必要額から一定額を"
              "差し引くものであるため、報酬改定・負担割合・調整交付金・入所・居住系の人数による差は、算定Ｂ・Ｃでも"
              "ほぼ同じです。予定収納率による差は、取崩し後の額に比例するため、算定Ｂ・Ｃでは小さくなります。")
    judge(c_rep(KYU28, SHIN28) == 1, "表28 の注記")

    print("  ── N09 選択２（入所・居住系10人）")
    t2 = [t for t in d.tables if [c.text.strip() for c in t.rows[0].cells][:3] == ["選択肢", "内　容", "保険料（月額）"]
          and "定員を維持" in t.rows[1].cells[1].text]
    judge(len(t2) == 1, "選択２の表を特定")
    if t2:
        KYU.append("\n".join(c.text for r in t2[0].rows for c in r.cells))
        V9.set_cell(t2[0].rows[1].cells[2], f"{b:,}円")
        V9.set_cell(t2[0].rows[2].cells[2], f"{b - j10:,}円（▲{j10}円）")
        V9.set_cell(t2[0].rows[3].cells[2], f"{b + j10:,}円（＋{j10}円）")
    p_rep("いずれも算定Ｂ（介護給付費準備基金の50％取崩し）の場合の額です。",
          "10人で月額約211円動きます（基金の取崩しの額によらず同じ幅です）。",
          f"10人で月額約{j10}円動きます（1円未満四捨五入。基金の取崩しの額によらず同じ幅です）。")

    print("  ── N02 本日のご審議事項⑥")
    p_rep("⑥　交付金の評価結果のご報告", "ご報告し、ご意見をいただくことが「評価結果の活用」の要件となります",
          "ご報告し、ご意見をいただくことは「評価結果の活用」の要件に当たる取組の一つとなり得ます"
          "（評価は国の該当状況調査への回答と議事の記録によります）")

    print("  ── N04 表29 準備基金の令和８年度末残高")
    n = c_rep("介護給付費準備基金の令和８年度末残高",
              "介護給付費準備基金の令和８年度末残高の見込額（決算による確定は令和９年度）")
    judge(n == 1, f"表29（{n}か所）")
    n = c_rep("町財政担当にご確認をお願い中", "町財政担当にご確認をお願い中。第３回策定委員会では見込額を用います")
    judge(n == 1, f"表29 確定の時期（{n}か所）")

    print("  ── N10 居宅サービスの下回りの説明")
    p_rep("総額が一致していることの内側で", "同時に進んだ結果と考えられます。",
          "同時に進んだことが影響している可能性があります（因果関係は、町内外別の利用の推移等による"
          "追加の検証を要します）。")

    print("  ── N08 「保険者」の数 → 市町村")
    pats = [(r"(\d)保険者", r"\1市町村"), (r"35市町村の平均", "35市町村の平均"),
            (r"380点満点である保険者は", "380点満点である市町村は"),
            (r"同じ64点の保険者は", "同じ64点の市町村は"), (r"同じ得点の保険者は", "同じ得点の市町村は"),
            (r"他の保険者の得点", "他の市町村の得点"), (r"各保険者の評価得点", "各市町村の評価得点"),
            (r"全保険者で按分", "全市町村で按分")]
    n = 0
    seen = set()
    for p in list(d.paragraphs) + [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]:
        if id(p._p) in seen:
            continue
        seen.add(id(p._p))
        s = p.text
        s2 = s
        for a, b_ in pats:
            s2 = re.sub(a, b_, s2)
        if s2 != s:
            KYU.append(s)
            kakikae(p, s2)
            n += 1
    judge(n >= 6, f"{n}段落・セルを改めた")

    print("  ── N07 改頁（審議事項の囲み・表の間）")
    p6 = [p for p in d.paragraphs if p.text.startswith("⑥　交付金の評価結果のご報告")]
    p7 = [p for p in d.paragraphs if p.text.startswith("⑦　策定委員会の回数")]
    judge(len(p6) == 1 and len(p7) == 1, "審議事項⑥・⑦の段落")
    bd6 = p6[0]._p.pPr.find(qn("w:pBdr")) if p6 and p6[0]._p.pPr is not None else None
    judge(bd6 is not None, "⑥に枠がある")
    if bd6 is not None and p7:
        # 枠は段落の書式（字下げ・枠・地色）がすべて同じ段落どうしで1つの囲みになる。⑥の書式を⑦に写す
        old = p7[0]._p.pPr
        if old is not None:
            p7[0]._p.remove(old)
        p7[0]._p.insert(0, copy.deepcopy(p6[0]._p.pPr))
        # ①〜⑦の囲みを同じ頁に置く
        for p in d.paragraphs:
            if re.match(r"^[①②③④⑤⑥]　", p.text):
                kn = p._p.get_or_add_pPr()
                if kn.find(qn("w:keepNext")) is None:
                    kn.append(docx.oxml.OxmlElement("w:keepNext"))
    # 構成の案の表（4案）は頁で分けない（1行だけ次頁に送られる）
    t45 = [t for t in d.tables if [c.text.strip() for c in t.rows[0].cells][:2] == ["構成の案", "回数・前提"]]
    judge(len(t45) == 1, "構成の案の表を特定")
    for r in t45[0].rows[:-1]:
        for c in r.cells:
            for p in c.paragraphs:
                kp = p._p.get_or_add_pPr()
                if kp.find(qn("w:keepNext")) is None:
                    kp.append(docx.oxml.OxmlElement("w:keepNext"))
    body = d.element.body
    kids = list(body.iterchildren())
    n_sep = 0
    for a_, b_ in zip(kids, kids[1:]):
        if a_.tag == qn("w:tbl") and b_.tag == qn("w:tbl"):
            sep = docx.oxml.OxmlElement("w:p")
            pPr = docx.oxml.OxmlElement("w:pPr")
            sp = docx.oxml.OxmlElement("w:spacing")
            sp.set(qn("w:before"), "0")
            sp.set(qn("w:after"), "0")
            sp.set(qn("w:line"), "120")
            sp.set(qn("w:lineRule"), "exact")
            pPr.append(sp)
            rPr = docx.oxml.OxmlElement("w:rPr")
            for tg in ("w:sz", "w:szCs"):
                e = docx.oxml.OxmlElement(tg)
                e.set(qn("w:val"), "4")
                rPr.append(e)
            pPr.append(rPr)
            sep.append(pPr)
            a_.addnext(sep)
            n_sep += 1
    judge(n_sep >= 2, f"隣り合う表の間に区切りの段落 {n_sep}か所")
    n_cs = 0
    for tr in body.iter(qn("w:tr")):
        # 再レビュー N07 で指摘のあった囲み（委託仕様書６(3)の開催回数）に限る。
        # 囲みのすべて・表のすべての行に掛けると、囲みごと次頁へ送られて頁が増える（21頁 → 22頁）
        if not "".join(x.text or "" for x in tr.iter(qn("w:t"))).startswith("※ 委託仕様書６(3)"):
            continue
        trPr = tr.find(qn("w:trPr"))
        if trPr is None:
            trPr = docx.oxml.OxmlElement("w:trPr")
            tcs = tr.findall(qn("w:tc"))
            # trPr は tblPrEx の後・最初の tc の前
            tcs[0].addprevious(trPr)
        if trPr.find(qn("w:cantSplit")) is None:
            cs = docx.oxml.OxmlElement("w:cantSplit")
            trPr.insert(0, cs)
            n_cs += 1
    judge(n_cs == 1, f"開催回数の囲みの行を分割しない（{n_cs}行）")

    print("  ── 段落の書式の要素の順（v16 と同じ）")
    JUN = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
           "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
           "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
           "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
           "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
           "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
           "sectPr", "pPrChange"]
    for pPr in body.iter(qn("w:pPr")):
        ko = list(pPr)
        jun = sorted(ko, key=lambda c: JUN.index(c.tag.split("}")[1]) if c.tag.split("}")[1] in JUN else len(JUN))
        if jun != ko:
            for c_ in jun:
                pPr.append(c_)
    # trPr の子の順（cantSplit は trHeight・tblHeader より前）
    TR = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit",
          "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden", "ins", "del", "trPrChange"]
    for trPr in body.iter(qn("w:trPr")):
        ko = list(trPr)
        jun = sorted(ko, key=lambda c: TR.index(c.tag.split("}")[1]) if c.tag.split("}")[1] in TR else len(TR))
        if jun != ko:
            for c_ in jun:
                trPr.append(c_)
    d.save(DST)

    d2 = docx.Document(DST)
    ato = zen(d2)
    print("  ── 検査1　古い記述")
    for go in ("6,235", "5,583", "6,366", "5,714", "約211円", "＋211円", "▲211円", "＋92円", "＋119円",
               "＋223円", "＋250円", "差がこれと異なります", "要件となります", "結果と考えられます",
               "1,741保険者", "35保険者", "285保険者", "保険者中", "380点満点である保険者",
               "他の保険者の得点", "同じ64点の保険者"):
        judge(go not in ato, f"「{go}」0件")
    print("  ── 検査2　素案と同じ値")
    for v in i1[1:] + i3[1:]:
        judge(f"{v:,}円" in ato and f"{v:,}円" in sz, f"表26 {v:,}円（素案 9-3 にある）")
    judge(f"{b - j10:,}円" in ato and f"{b + j10:,}円" in ato, f"選択２ {b - j10:,}円・{b + j10:,}円（素案の算定Ｂ±{j10}円）")
    judge(f"約{j10}円" in sz, f"約{j10}円は素案 5-5 にある")
    print("  ── 検査3　消えた数は旧の文に現れる")
    kyu = "\n".join(KYU)
    kieta = sorted({mm.group(0) for mm in re.finditer(r"[0-9０-９][0-9０-９,.]*(?:人|名|％|千円|件|頁|日|点|円|位)", mae)
                    if mae.count(mm.group(0)) > ato.count(mm.group(0)) and mm.group(0) not in kyu})
    judge(not kieta, "消えた数はすべて旧の文に現れる", "／".join(kieta[:5]))
    print("  ── 検査4　事業所名・他団体名・⚠")
    for na in ("釜房", "アルパイン", "鶴寿会", "グループホームさわやか", "グループホームふれあい",
               "北塩原", "金ヶ崎", "大雪", "川崎市"):
        judge(na not in ato, f"「{na}」0件")
    judge("⚠" not in ato and "**" not in ato, "⚠・** 0件（第1回の書式）")
    print("  ── 検査5　隣り合う表がない・⑦に枠")
    kids = list(d2.element.body.iterchildren())
    judge(not any(a_.tag == qn("w:tbl") and b_.tag == qn("w:tbl") for a_, b_ in zip(kids, kids[1:])),
          "隣り合う表 0件")
    p7 = [p for p in d2.paragraphs if p.text.startswith("⑦　策定委員会の回数")][0]
    judge(p7._p.pPr.find(qn("w:pBdr")) is not None, "⑦に枠（⑥と同じ）")
    print(f"  {DST}")
    print(f"  ⚠ 町のご確認が要る印【町確認】が {ato.count('【町確認】')}件")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜5　不適合0件")


if __name__ == "__main__":
    main()
