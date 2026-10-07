# -*- coding: utf-8 -*-
"""交付金の3か年（令和6・7・8年度）を枝番の単位で突き合わせ、推移を出す。

出典
  令和6年度 ①令和６年度…/01_（掲載用）令和６年度評価結果（市町村分）【交付見込額確定版】.xlsx
  令和7年度 ①令和７年度…/01_（掲載用）令和７年度評価結果（市町村分）.xlsx
  令和8年度 ①令和８年度…/01_（掲載用）令和８年度評価結果（市町村分）.xlsx

年度により列の並びが異なる（令和7年度から規模区分の列が入る等）ため、
見出し行（2〜9行目）から列の意味を組み立てて突き合わせる。
"""
import glob
import json
import re
import sys
from collections import OrderedDict

import openpyxl

F = {6: glob.glob("09_元資料/交付金評価/①令和６年度*/01_*.xlsx")[0],
     7: glob.glob("09_元資料/交付金評価/①令和７年度*/01_*.xlsx")[0],
     8: glob.glob("09_元資料/交付金評価/①令和８年度*/01_*.xlsx")[0]}
KAWA = 280          # 通し番号（3か年とも同じ）


def norm(s):
    s = str(s or "")
    s = re.sub(r"[\s　（）()、，・]", "", s)
    for a, b in (("１", "1"), ("２", "2"), ("３", "3"), ("４", "4"),
                 ("５", "5"), ("６", "6"), ("７", "7"), ("８", "8"),
                 ("９", "9"), ("０", "0")):
        s = s.replace(a, b)
    return s


def ffill(row):
    out, last = [], None
    for v in row:
        if v is not None and str(v).strip() != "":
            last = v
        out.append(last)
    return out


def load(year):
    wb = openpyxl.load_workbook(F[year], data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    head = {i: list(r) for i, r in
            enumerate(ws.iter_rows(min_row=1, max_row=14, values_only=True), 1)}
    GRP, TGT, SUB = ffill(head[2]), ffill(head[3]), ffill(head[4])
    NO, ITEM = ffill(head[5]), ffill(head[6])
    EDA, EDA2, HAI, AVG = head[7], head[8], head[9], head[11]
    cols = OrderedDict()
    for c in range(len(HAI)):
        if not isinstance(HAI[c], (int, float)):
            continue
        if not (EDA[c] or EDA2[c]):
            continue          # 小計・合計の列は除く
        grp = "推進" if "保険者機能強化" in str(GRP[c]) else "支援"
        tgt = re.sub(r"[　 ].*$", "", str(TGT[c] or ""))
        sub = ("体制" if "体制" in str(SUB[c]) else
               "活動" if "活動" in str(SUB[c]) else
               "成果" if "成果" in str(SUB[c]) else "?")
        base = (grp, tgt, sub, norm(ITEM[c]),
                str(EDA[c] or ""), str(EDA2[c] or ""))
        # 同じ見出しの組合せが複数回現れる（目標Ⅳは推進・支援で同じ指標を
        # 評価するが、見出し行の交付金名が引き継がれないため）。
        # 何回目かを鍵に加えて区別する。
        occ = 1
        while (base + (occ,)) in cols:
            occ += 1
        key = base + (occ,)
        cols[key] = dict(col=c, hai=HAI[c], item=str(ITEM[c] or ""),
                         no=str(NO[c] or ""), avg=AVG[c],
                         eda=str(EDA[c] or ""), eda2=str(EDA2[c] or ""),
                         grp=grp, tgt=tgt, sub=sub)
    kawa = None
    for r in ws.iter_rows(min_row=15, values_only=True):
        if r[0] == KAWA:
            kawa = list(r)
            break
    for k, v in cols.items():
        v["mine"] = kawa[v["col"]] if kawa else None
    return cols


def main():
    D = {y: load(y) for y in (6, 7, 8)}
    for y in (6, 7, 8):
        tot = sum(v["hai"] for v in D[y].values() if "目標Ⅳ" not in v["tgt"])
        mine = sum(v["mine"] for v in D[y].values()
                   if isinstance(v["mine"], (int, float)))
        print(f"令和{y}年度　枝番{len(D[y])}列　"
              f"目標Ⅳを除く配点{tot}　川崎町の枝番合計{mine}",
              file=sys.stderr)
    keys = list(D[8].keys())
    out = []
    for k in keys:
        v8 = D[8][k]
        v7, v6 = D[7].get(k), D[6].get(k)
        out.append(dict(
            grp=v8["grp"], tgt=v8["tgt"], sub=v8["sub"], no=v8["no"],
            item=v8["item"], eda=v8["eda"], eda2=v8["eda2"], hai=v8["hai"],
            r6=(v6 or {}).get("mine"), r7=(v7 or {}).get("mine"),
            r8=v8["mine"],
            a6=(v6 or {}).get("avg"), a7=(v7 or {}).get("avg"),
            a8=v8["avg"],
            hai6=(v6 or {}).get("hai"), hai7=(v7 or {}).get("hai")))
    # 令和8年度にない指標（廃止・改称）も拾う
    gone = [k for k in D[7] if k not in D[8]]
    json.dump(dict(rows=out, gone=[list(k) for k in gone]),
              open("/tmp/claude-0/kf/trend.json", "w"), ensure_ascii=False)
    print(f"突合 {len(out)}枝番／令和7年度にあって令和8年度にない枝番 {len(gone)}",
          file=sys.stderr)


if __name__ == "__main__":
    main()
