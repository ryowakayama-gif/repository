# -*- coding: utf-8 -*-
"""計画素案 Ver.2.13（事業所・住まいの記述の是正）

08_作業順位 の順位26。Ver.2.12 に当てて Ver.2.13 を作る。

  入力　01_第10期_最新版成果品/川崎町_計画書素案_v2.12_根拠整理版.docx
  出力　01_第10期_最新版成果品/川崎町_計画書素案_v2.13_供給体制整理版.docx

是正1　住所地特例24人の時点を令和8年3月末に揃える
  素案の中で時点が2通りになっていた。
    令和7年6月時点 … 2-7（本文・表）・7-1・7-6・9-1
    令和8年3月末　 … 2-3・3-2（本文・表）・3-3・9-1
  2-1 の表が「令和3年度末（2022.3）34人 → 令和7年度末（2026.3）24人」と
  しており、令和7年度末＝令和8年3月末である。同じ24人に2つの時点を置けない。
  ⚠ 令和7年6月時点の24人という別の数値が実在する場合に備え、
    確認事項No.169として起票している。

是正2　1-5 の別表 事項七の対応欄を実態に合わせる
  現在の対応欄は「2-3・9-1・8-3。年齢調整後の認定率は…」ではなく
  「2-7・7-1・9-1。入所率・稼働率は2施設への照会による（確認事項No.29・No.159）」
  であり、**施設・事業所数が揃っているように読める。**
  実際に揃っているのは本町被保険者の受給者数だけであり、
  施設・事業所数は出典がなく、入所率・稼働率・従事者数は未受領である。

是正3　2-7・7-1 の施設数の記述に、出典がないことと落ちがあることを添える
  「特別養護老人ホーム1施設、介護老人保健施設1施設、小規模多機能型居宅介護等の
  地域密着型サービスが中心」という記述には出典がなく、令和7年度に現に給付の
  ある地域密着型介護老人福祉施設（月26.6人）と認知症対応型共同生活介護
  （月25.9人）が落ちている（RedTeamレビューv3 R-3 で既に指摘）。

⚠ 見込量・給付費・保険料の数値は一切変えていない。
"""
import copy
import os
import re
import shutil
import sys
import zipfile

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from fix_soan_v111 import set_el, set_tc                     # noqa: E402

SRC = "01_第10期_最新版成果品/川崎町_計画書素案_v2.12_根拠整理版.docx"
DST = "01_第10期_最新版成果品/川崎町_計画書素案_v2.13_供給体制整理版.docx"

# 是正1　時点の統一
JITEN_FURUI = "令和7年6月時点"
JITEN_ATARASHII = "令和8年3月末"

# 是正2　別表 事項七の対応欄
NANA_FURUI = "2-7・7-1・9-1。入所率・稼働率は2施設への照会による"
NANA_ATARASHII = (
    "2-7・7-1・9-1。揃っているのは本町被保険者の受給者数のみ。"
    "施設・事業所数は出典が確定しておらず、入所率・稼働率・従事者数は未受領"
    "（確認事項No.29・No.50・No.51・No.159）")

# 是正3　施設数の記述に添える注
SHISETSU = "特別養護老人ホーム1施設・介護老人保健施設1施設"
SHISETSU_HYO = "特別養護老人ホーム1施設、介護老人保健施設1施設"
CHU_71 = (
    "⚠ この施設数は、第1回策定委員会でのご説明によるものであり、"
    "町の指定台帳との突合が済んでいません。令和7年度には地域密着型"
    "介護老人福祉施設入所者生活介護（月26.6人）及び認知症対応型共同生活介護"
    "（月25.9人）の給付の実績がありますので、町内の提供体制はこれらを含めて"
    "整理する必要があります。地域密着型サービスは町が指定権者ですので、"
    "町の指定台帳により確定します（確認事項No.29）。")


def main():
    if not os.path.exists(SRC):
        raise SystemExit("入力がない：" + SRC)
    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body

    # ── 雛形（注を加えるため）
    tpl = None
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = "".join(x.text or "" for x in el.iter(qn("w:t"))).strip()
        if s.startswith("⚠ 宮城県自身の評価"):
            tpl = copy.deepcopy(el)
            break
    if tpl is None:
        raise SystemExit("囲みの雛形が見つからない")

    # ── 是正1　時点（本文・表の双方）
    n_jiten = 0
    for el in body.iter(qn("w:t")):
        if el.text and JITEN_FURUI in el.text:
            el.text = el.text.replace(JITEN_FURUI, JITEN_ATARASHII)
            n_jiten += 1
    if n_jiten == 0:
        raise SystemExit(f"「{JITEN_FURUI}」が見つからない")

    # ── 是正2　別表 事項七の対応欄
    n_nana = 0
    for el in body.iterchildren():
        if el.tag != qn("w:tbl"):
            continue
        tb = docx.table.Table(el, doc)
        for r in tb.rows:
            for c in r.cells:
                if NANA_FURUI in c.text:
                    set_tc(c._tc, c.text.replace(NANA_FURUI, NANA_ATARASHII))
                    n_nana += 1
    if n_nana != 1:
        raise SystemExit(f"別表 事項七の対応欄が{n_nana}件（1件のはず）")

    # ── 是正3　7-1 の施設数の記述のあとに注を置く
    n_chu = 0
    for el in list(body.iterchildren()):
        if el.tag != qn("w:p"):
            continue
        s = "".join(x.text or "" for x in el.iter(qn("w:t")))
        if SHISETSU in s and "町内のサービス提供体制は" in s:
            e = copy.deepcopy(tpl)
            set_el(e, CHU_71)
            el.addnext(e)
            n_chu += 1
    if n_chu != 1:
        raise SystemExit(f"7-1 の施設数の記述が{n_chu}件（1件のはず）")

    # ── 2-7 の表の該当セルにも短い注を添える
    n_hyo = 0
    for el in body.iterchildren():
        if el.tag != qn("w:tbl"):
            continue
        tb = docx.table.Table(el, doc)
        for r in tb.rows:
            for c in r.cells:
                if SHISETSU_HYO in c.text and "⚠" not in c.text:
                    set_tc(c._tc, c.text
                           + "　⚠ 出典が確定しておらず、地域密着型介護老人"
                             "福祉施設及び認知症対応型共同生活介護が"
                             "落ちています（確認事項No.29）")
                    n_hyo += 1
    if n_hyo != 1:
        raise SystemExit(f"2-7 の表の該当セルが{n_hyo}件（1件のはず）")

    # ── 版表記
    n_ver = 0
    for el in body.iter(qn("w:t")):
        if el.text and "Ver.2.12" in el.text:
            el.text = el.text.replace("Ver.2.12", "Ver.2.13")
            n_ver += 1
    if n_ver == 0:
        raise SystemExit("本文に Ver.2.12 の表記が見つからない")
    doc.save(DST)
    naoshi_header(DST)

    # ══════════════════════════ 自己点検
    d2 = docx.Document(DST)
    d0 = docx.Document(SRC)
    zen = ("\n".join(q.text for q in d2.paragraphs) + "\n"
           + "\n".join(c.text for t in d2.tables
                       for r in t.rows for c in r.cells))
    ng = []
    if JITEN_FURUI in zen:
        ng.append(f"「{JITEN_FURUI}」が残っている")
    if NANA_FURUI in zen:
        ng.append("別表 事項七の対応欄が古いまま")
    if "揃っているのは本町被保険者の受給者数のみ" not in zen:
        ng.append("別表 事項七の対応欄が改まっていない")
    if CHU_71[:30] not in zen:
        ng.append("7-1 の注が入っていない")
    # 数値を変えていないこと（主な数値が残っていること）
    for v in ("6,822", "998,220", "24.5", "25.0", "26.6", "25.9", "1,741"):
        if v not in zen:
            ng.append(f"主な数値が消えている：{v}")
    # 段落は1つだけ増え、表の数は変わらないこと
    if len(d2.paragraphs) != len(d0.paragraphs) + 1:
        ng.append(f"段落が{len(d0.paragraphs)}→{len(d2.paragraphs)}"
                  "（1つだけ増えるはず）")
    if len(d2.tables) != len(d0.tables):
        ng.append(f"表の数が変わっている（{len(d0.tables)}→{len(d2.tables)}）")
    for w in ("**", "__"):
        if w in zen:
            ng.append(f"強調記号が残っている：{w}")
    z = zipfile.ZipFile(DST)
    hdr = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>",
                             z.read("word/header1.xml").decode("utf-8")))
    if "Ver.2.13" not in hdr:
        ng.append(f"ヘッダーの版表記が更新されていない：{hdr}")

    print("計画素案 Ver.2.13（供給体制整理版）")
    print("保存：", DST)
    print(f"  段落 {len(d2.paragraphs)}（Ver.2.12 は {len(d0.paragraphs)}）"
          f"／表 {len(d2.tables)}")
    print(f"  是正1　住所地特例の時点を揃えた　{n_jiten}か所"
          f"（{JITEN_FURUI} → {JITEN_ATARASHII}）")
    print(f"  是正2　別表 事項七の対応欄　{n_nana}か所")
    print(f"  是正3　7-1 に注を置いた {n_chu}か所／"
          f"2-7 の表に注を添えた {n_hyo}か所")
    print("  ── 自己点検")
    if ng:
        for m in ng:
            print("   ×", m)
        sys.exit(1)
    print("   ○ 住所地特例の時点が1つに揃った")
    print("   ○ 別表 事項七の対応欄が実態に合った")
    print("   ○ 見込量・給付費・保険料の数値は変えていない")
    print("  ⚠ 目次の頁番号は fill_toc_pages_v200.py で作り直してください。")


def naoshi_header(path):
    z = zipfile.ZipFile(path)
    names = z.namelist()
    data = {n: z.read(n) for n in names}
    z.close()
    n_fix = 0
    for n in names:
        if not re.match(r"word/(header|footer)\d*\.xml", n):
            continue
        s = data[n].decode("utf-8")
        s2 = re.sub(r"Ver\.2\.12", "Ver.2.13", s)
        if s2 != s:
            data[n] = s2.encode("utf-8")
            n_fix += 1
    if n_fix == 0:
        raise SystemExit("ヘッダーの版表記（Ver.2.12）が見つからない")
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    shutil.move(tmp, path)


if __name__ == "__main__":
    main()
