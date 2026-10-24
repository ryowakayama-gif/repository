# -*- coding: utf-8 -*-
"""業務工程管理表の更新（第2回委員会資料v9・想定問答集v5）（令和8年10月21日）

  python3 07_ソーススクリプト/update_wbs_R8.10.21.py

なぜ更新するか
  第2回策定委員会（令和8年11月24日）に向けて、委員会資料を v9 に、
  想定問答集を v5 に上げた。あわせて計画素案を Ver.2.22 に上げた。

  ⚠ **委託仕様書の読み違いが1件見つかった。**
  v8 の表38 は「委託仕様書上、策定委員会は3回構成」としていたが、
  仕様書6(3)は**「策定委員会の開催回数は4回開催の予定」**と定めている。
  3回構成は第1回策定委員会の当日にお示しいただいた日程によるもので
  あった。v9 で是正した。

  ⚠ **計画素案に同じ指標の2つの値があった。**
  7-2 の標準給付費（令和7年度）が 1,085,203千円・▲0.7％、
  9-3 が 1,084,587千円・▲0.8％。素案の表15・表109 により
  1,084,587千円 が正しい。Ver.2.22 で是正した。

更新するもの
  03_確認事項一覧　No.49・No.19・No.53・No.83・No.186・No.79
  04_成果品管理　　成果品4（素案 Ver.2.22）・成果品6（資料v9・問答集v5）
  08_作業順位　　　順位36（資料の整理と想定問答集）

⚠ **列8（状態）には書き込まない。** 回答・対応は列10に書く。
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
SOAN = "川崎町_計画書素案_v2.22_給付費整合版.docx"
SHIRYO = "川崎町_第2回策定委員会資料_R8.11_v9.docx"
QA = "川崎町_第2回策定委員会_想定問答集_R8.11_v5.docx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月21日 更新（資料v9・想定問答集v5）】"


def iro(ws, r, n):
    for c in range(1, n + 1):
        ws.cell(r, c).font = FONT
        ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")


TSUIKI = {
    "49": (10,
           f"{HIZUKE}⚠ **委託仕様書の読み違いを是正しました。**"
           f"第2回委員会資料 v8 の表38 は「委託仕様書上、策定委員会は"
           f"3回構成であり」としていましたが、**仕様書6(3)は"
           f"「策定委員会の開催回数は4回開催の予定」**と定めています。"
           f"3回構成は第1回策定委員会の当日にお示しいただいた日程に"
           f"よるものでした。資料 v9 の表38 で是正しています。\n"
           f"⚠ **全4回では納期に間に合いません。**第2回を11月24日として"
           f"逆算すると、全4回・パブコメ30日・印刷3週間で納品は"
           f"令和9年4月4日（**20日超過**）、パブコメ20日・印刷2週間に"
           f"詰めても3月18日（3日超過）です。全3回（第3回で"
           f"パブリックコメントの結果の報告・保険料の決定・答申の内容の"
           f"確定を併せて行う）であれば令和9年3月10日（余裕5日）です。\n"
           f"⚠ **仕様書11**（本仕様の定めのない事項については、必要に応じ"
           f"協議の上実施する）により、回数のご協議をお願いします。"
           f"第2回でお諮りする形で資料を組んでいます。　"),
    "19": (10,
           f"{HIZUKE}⚠ **パブリックコメントの開始が令和8年12月29日に"
           f"当たります。**第2回（11月24日）から計画書（案）の作成と"
           f"宮城県への事前協議に35日を見込むとこうなります。"
           f"年末年始を挟むと実質の周知期間が短くなりますので、"
           f"**令和9年1月4日（月）開始とする案**を資料 v9 の８に"
           f"併記しました（この場合の納品は令和9年3月15日ちょうど）。"
           f"期間（30日間とするか）と開始日のご判断をお願いします。　"),
    "53": (10,
           f"{HIZUKE}開催日が11月24日に確定しましたので、"
           f"**事前配布は11月10日（火）**（14日前）、"
           f"**事前質問の締切は11月17日（火）**（7日前）となります。"
           f"資料 v9・想定問答集 v5 は作成済みです"
           f"（資料は全９版・想定問答集は全61問）。"
           f"紙・電子のいずれによるか、部数をご指示ください。　"),
    "83": (10,
           f"{HIZUKE}⚠ **素案 Ver.2.22 は本文117頁です。**"
           f"資料編（見込13頁）を加えると計130頁となり、仕様書8の"
           f"「100頁程度」との差は**＋30頁**です。"
           f"⚠ **圧縮案は Ver.2.18c（本文106頁）のままであり、"
           f"現在の素案に合わせて作り直す必要があります。**"
           f"「100頁程度」を本文のみと読むか資料編を含めて読むかの"
           f"ご判断とあわせて、作り直しの要否をご指示ください。"
           f"資料 v9 の表35 に3つの版を並べています。　"),
    "186": (10,
            f"{HIZUKE}⚠ **想定問答集 v5 に問を立てました（QB-54）。**"
            f"「保険料の前提になる将来推計は、どれを使っているのか」に"
            f"対し、**正本とするかどうかは町のご判断をお待ちしている**"
            f"旨を答える形にしています。⚠ **10月30日までにご判断を"
            f"いただけないと、保険料の3パターンを資料の確定版に"
            f"載せられません。**　"),
    "79": (10,
           f"{HIZUKE}⚠ **想定問答集 v5 の QB-54 に併記しました。**"
           f"令和8年度の実績見込み値が差し替え前の値に戻っている点に"
           f"ついて、伸び率の基準年が動くため保険料に直接影響する旨を"
           f"答える形にしています。No.186 と同じく10月30日までに"
           f"ご判断をお願いします。　"),
}


def main():
    if not os.path.exists(WBS):
        raise SystemExit("業務工程管理表がない：" + WBS)
    wb = openpyxl.load_workbook(WBS)

    # ── 03_確認事項一覧
    ws = wb["03_確認事項一覧"]
    n = sumi = 0
    atatta = set()
    for x in range(5, ws.max_row + 1):
        no = str(ws.cell(x, 1).value or "").strip()
        if no not in TSUIKI:
            continue
        atatta.add(no)
        retsu, moji = TSUIKI[no]
        mae = str(ws.cell(x, retsu).value or "")
        if HIZUKE in mae:
            sumi += 1
            continue
        ws.cell(x, retsu).value = moji + mae
        iro(ws, x, 10)
        n += 1
    nai = set(TSUIKI) - atatta
    if nai:
        raise SystemExit("確認事項の行が見つからない：No."
                         + "・No.".join(sorted(nai)))

    # ── 04_成果品管理
    ws = wb["04_成果品管理"]
    n_seika = 0
    for x in range(5, ws.max_row + 1):
        no = str(ws.cell(x, 1).value or "").strip()
        na = str(ws.cell(x, 3).value or "")
        if no == "4" and "v2.21_公表情報把握版" in na:
            ws.cell(x, 3).value = na.replace(
                "川崎町_計画書素案_v2.21_公表情報把握版.docx", SOAN)
            ws.cell(x, 8).value = (
                f"{HIZUKE}素案を Ver.2.22 に上げた（全117頁）。"
                f"⚠ **同じ指標に2つの値があったのを是正した。**"
                f"7-2 の標準給付費（令和7年度）が1,085,203千円・▲0.7％、"
                f"9-3 が1,084,587千円・▲0.8％であった。素案の表15"
                f"（保険給付費計1,084,587,307円）と表109（標準給付費"
                f"見込額1,084,587千円）により後者が正しい。"
                f"あわせて、第3回の開催時期を令和9年1月と断定している"
                f"5か所を改めた（11月24日開催ではパブリックコメント30日間を"
                f"挟むと1月には開けない）。　"
                + str(ws.cell(x, 8).value or ""))
            iro(ws, x, 9)
            n_seika += 1
        elif no == "6" and "_v8.docx" in na:
            ws.cell(x, 3).value = (na
                                   .replace("川崎町_第2回策定委員会資料_"
                                            "R8.11_v8.docx", SHIRYO)
                                   .replace("同_想定問答集_R8.11_v4.docx",
                                            "同_想定問答集_R8.11_v5.docx"))
            ws.cell(x, 8).value = (
                f"{HIZUKE}資料を v9（全９版）、想定問答集を v5"
                f"（全61問・8カテゴリ）に上げた。"
                f"⚠ **委託仕様書の読み違いを1件是正した**"
                f"（仕様書6(3)は「4回開催の予定」）。"
                f"⚠ **交付金の「27点」が2つあることを書き分けた**"
                f"（体制・取組指標群の失点27点と、推進 目標Ⅲ の活動"
                f"指標群の27点）。公表情報による町内事業所の状況の表と"
                f"退所先の表を加え、第1回の委員のご発言に基づく人数・"
                f"金額を外した。　"
                + str(ws.cell(x, 8).value or ""))
            iro(ws, x, 9)
            n_seika += 1
    if n_seika != 2:
        raise SystemExit(f"04シートの該当行が{n_seika}件（2件のはず）")

    # ── 08_作業順位　順位36
    ws2 = wb["08_作業順位"]
    jun = [str(ws2.cell(x, 1).value or "").strip()
           for x in range(5, ws2.max_row + 1)]
    n_jun = 0
    if "36" not in jun:
        r = ws2.max_row + 1
        gyo = ["36", "第2回策定委員会の資料の整理と想定問答集",
               "R8.10.21",
               "ご指示による。開催日（令和8年11月24日）が確定し、"
               "ヒアリングを行わない整理と素案 Ver.2.21 が整ったため、"
               "資料と想定問答集を引き直せる。",
               "束のうちNo.49・No.19・No.53・No.83／No.186・No.79の提示",
               "第2回策定委員会資料 v9（全９版・表44）／想定問答集 v5"
               "（全61問）／計画素案 Ver.2.22",
               "1日", "完了",
               "⚠ **委託仕様書の読み違いを1件是正した。** v8 の表38 の"
               "「委託仕様書上、策定委員会は3回構成」は誤りで、"
               "仕様書6(3)は「4回開催の予定」と定めている。"
               "⚠ **計画素案の給付費の二重の値を是正した**"
               "（1,085,203千円→1,084,587千円）。"
               "⚠ **交付金の「27点」が2つあることを書き分けた。**"
               "公表情報による町内事業所の状況の表と退所先の表を資料に"
               "加え、想定問答集に17問を加えた（44問→61問）。"
               "第1回の委員のご発言に基づく人数・金額は資料・問答集の"
               "双方から外した。"]
        for c, v in enumerate(gyo, start=1):
            ws2.cell(r, c).value = v
        iro(ws2, r, len(gyo))
        n_jun = 1

    wb.save(WBS)

    # ══════ 自己点検
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("業務工程管理表の更新（資料v9・想定問答集v5）")
    print(f"  {WBS}")
    print()
    print("  ── 自己点検")
    judge(n + sumi == len(TSUIKI),
          f"確認事項 {len(TSUIKI)}件に書いた（新 {n}・既 {sumi}）")
    judge(n_seika == 2, "成果品管理の2行を更新した")
    judge(n_jun == 1 or "36" in jun, "作業順位36 を記録した")

    wb2 = openpyxl.load_workbook(WBS, data_only=True)
    ws3 = wb2["03_確認事項一覧"]
    jotai = [str(ws3.cell(x, 8).value or "")
             for x in range(5, ws3.max_row + 1)
             if str(ws3.cell(x, 1).value or "").strip() in TSUIKI]
    judge(all(s in ("確認待ち", "決着", "保留", "回答済") for s in jotai),
          "⚠ 列8（状態）を動かしていない", "／".join(sorted(set(jotai))))
    no_all = sorted(int(s) for s in
                    (str(ws3.cell(x, 1).value or "").strip()
                     for x in range(5, ws3.max_row + 1)) if s.isdigit())
    kesu = [i for i in range(1, max(no_all) + 1) if i not in no_all]
    judge(not kesu, f"欠番0（No.1〜{max(no_all)}・{len(no_all)}件）")

    zen = []
    for sh in wb2:
        for row in sh.iter_rows(values_only=True):
            zen += [str(v) for v in row if v is not None]
    s = "\n".join(zen)
    judge("川崎市" not in s and "福岡" not in s,
          "⚠ 川崎市（神奈川県）・福岡県川崎町が混じらない")
    for na in ("釜房みどりの園", "アルパイン", "鶴寿会", "光友会",
               "グループホームさわやか", "グループホームふれあい"):
        judge(na not in s, f"法人・事業所の固有名称「{na}」が入っていない")
    judge(SHIRYO in s, "成果品管理に資料 v9 が載っている")
    judge(QA in s or "想定問答集_R8.11_v5" in s,
          "成果品管理に想定問答集 v5 が載っている")
    judge(SOAN in s, "成果品管理に素案 Ver.2.22 が載っている")

    print()
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        for m in futeki:
            print(f"      {m}")
        sys.exit(1)
    print(f"   ○ 自己点検 不適合0件（確認事項 {len(no_all)}件）")


if __name__ == "__main__":
    main()
