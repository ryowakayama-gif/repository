# -*- coding: utf-8 -*-
"""計画素案 Ver.2.24 → Ver.2.25（町内の軽費老人ホームの廃止の確認結果）
（令和8年10月22日）

  python3 07_ソーススクリプト/fix_soan_v2250_keihi.py

なぜ要るか
  令和8年10月22日に、町内の軽費老人ホームの廃止前後を公表資料で
  調べた結果の提供を受けた（09_元資料/その他データ/軽費老人ホーム廃止_R8.10.22/）。
    ・廃止日は令和6年3月31日（法人の公表による）
    ・廃止前の実入所者数・本町の住民の数・転居先は公表資料では未確認
    ・⚠ 「定員70人＝70人が住まいを失った」と書かない
  これを手元の宮城県「老人福祉施設等入所状況調」と突き合わせた。
    ・県調査は令和6年1月1日まで「休止中」で掲載、4月1日から掲載なし
      → 廃止日 令和6年3月31日 と整合する
    ・⚠ 県調査には**在所者数と退所先が3か月ごとに載っていた。**
      提供資料が「未確認」とした廃止前の入所者数と退所先の一部が埋まる
        令和5年4月1日 13人 → 7月1日 7人 → 10月1日 0人（休止中）
        令和5年1〜9月の退所32人：他の施設28・入院3・在宅1
  ⚠ 素案 Ver.2.24 の 2-7 は「町内の高齢者向けの住まい70人分が失われた」と
    書いており、提供資料が注意する誤読（定員＝居所を失った人数）を招く。
    実際の入所者は令和5年4月1日時点で13人であった。

  ⚠ 施設名・法人名は成果品に書かない（CLAUDE.md）。
  ⚠ 法人の公表ページはこの作業環境から取得できなかった（通信の許可先にない）。
    廃止日は提供資料によるが、県調査の掲載の推移と整合することを確かめた。

数値の正本　09_元資料/その他データ/県入所状況調_川崎町_R8.10.22.csv
終了コード　0＝検査1〜6すべて適合／1＝1件でも不適合
"""
import csv
import os
import re
import shutil
import sys

import docx

MOTO = "01_第10期_最新版成果品/川崎町_計画書素案_v2.24_交付金整理版.docx"
SAKI = "01_第10期_最新版成果品/川崎町_計画書素案_v2.25_軽費老人ホーム確認版.docx"
BAN_MOTO = "Ver.2.24"
BAN = "Ver.2.25"
CSV_IN = "09_元資料/その他データ/県入所状況調_川崎町_R8.10.22.csv"
JURYO = ("09_元資料/その他データ/軽費老人ホーム廃止_R8.10.22/"
         "軽費老人ホーム廃止前後の確認結果_R8.10.22受領.md")
HAISHI = "令和6年3月31日"
# 県調査の時点の並び（和暦の文字列のままでは並ばないため）
JUN = ["R5.4.1", "R5.7.1", "R5.10.1", "R6.1.1", "R6.4.1"]


def n(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def wa(h):
    y, m, d = h[1:].split(".")
    return f"令和{y}年{m}月{d}日"


def keihi():
    """県調査から町内の軽費老人ホームの行を読む。"""
    with open(CSV_IN, encoding="utf-8-sig", newline="") as f:
        rows = [r for r in csv.DictReader(f) if r["種別"] == "軽費"]
    return {r["時点"]: r for r in rows}


def atai():
    K = keihi()
    A = {"定員": n(K["R5.4.1"]["定員"])}
    A["現員"] = {h: n(K[h]["入所現員"]) for h in JUN[:3]}
    A["退所"] = sum(n(K[h]["退所人員"]) for h in JUN[:3])
    A["新入所"] = sum(n(K[h]["新入所人員"]) for h in JUN[:3])
    for k in ("死亡", "入院", "在宅", "他施設"):
        A[k] = sum(n(K[h][k]) for h in JUN[:3])
    A["休止"] = next(h for h in JUN if "休止" in K.get(h, {}).get("施設名", ""))
    A["最後"] = [h for h in JUN if h in K][-1]
    A["K"] = K
    return A


A = atai()
GEN = A["現員"]

HON_213 = (
    f"第3に、町内にあった軽費老人ホーム（定員{A['定員']}人）は、{HAISHI}を"
    f"もって事業を廃止しました。宮城県「老人福祉施設等入所状況調」によると、"
    f"入所者は令和5年4月1日時点で{GEN['R5.4.1']}人、同年7月1日時点で"
    f"{GEN['R5.7.1']}人と減り、同年10月1日時点で{GEN['R5.10.1']}人となり"
    f"休止中とされています。令和5年1月から9月までに退所した{A['退所']}人の"
    f"退所先は、他の施設{A['他施設']}人、入院{A['入院']}人、在宅{A['在宅']}人"
    f"です。第9期計画期間が始まる直前に、町内の高齢者向けの住まいの選択肢が"
    f"定員{A['定員']}人分なくなったことになります。⚠ ただし、定員の"
    f"{A['定員']}人がそのまま住まいを失ったものではありません。令和5年1月より"
    f"前の入所者数、入所されていた方のうち本町の住民の数、移られた先の施設の"
    f"種別と所在の市町村は、公表資料にはありません［要確認］。⚠ 軽費老人ホーム"
    f"は老人福祉法に基づく施設であり、介護サービス情報公表システムには載り"
    f"ません。廃止の理由と入所されていた方への措置は、社会福祉法第64条による"
    f"宮城県への廃止届に記載される事項です（確認事項No.154）。")

HYO4_KYU = ("⚠ 町内の軽費老人ホーム（定員70人）が令和5年10月1日時点で休止中と"
            "なり、令和6年4月1日以降は宮城県「老人福祉施設等入所状況調」に"
            "掲載されていない。")
HYO4_SHIN = (
    f"⚠ 町内の軽費老人ホーム（定員{A['定員']}人）は令和5年10月1日時点で"
    f"休止中（入所者{GEN['R5.10.1']}人）となり、{HAISHI}に廃止された。"
    f"令和5年1月から9月までの退所{A['退所']}人の退所先は他の施設"
    f"{A['他施設']}人・入院{A['入院']}人・在宅{A['在宅']}人（宮城県「老人福祉"
    f"施設等入所状況調」）。")


def dan_kaku(p, s):
    if p.runs:
        p.runs[0].text = s
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(s)


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    print(f"計画素案 {BAN_MOTO} → {BAN}（町内の軽費老人ホームの廃止の確認結果）")
    print()
    for f in (MOTO, CSV_IN, JURYO):
        if not os.path.exists(f):
            print(f"   × 入力がありません：{f}")
            sys.exit(1)

    print("  ── ⓪ 県調査の値の整合")
    K = A["K"]
    sa = all(GEN[JUN[i]] - GEN[JUN[i + 1]]
             == n(K[JUN[i + 1]]["退所人員"]) - n(K[JUN[i + 1]]["新入所人員"])
             for i in range(2))
    judge(sa, "在所者数の差＝退所−新入所（令和5年4月→7月→10月）")
    judge(A["死亡"] + A["入院"] + A["在宅"] + A["他施設"] == A["退所"],
          f"退所理由の合計＝退所{A['退所']}人")
    judge(A["新入所"] == 0, "この間の新入所は0人")
    judge(A["休止"] == "R5.10.1", f"休止中の初出 {wa(A['休止'])}")
    judge(A["最後"] == "R6.1.1" and "R6.4.1" not in K,
          f"最後の掲載 {wa(A['最後'])}・令和6年4月1日は掲載なし"
          f"（廃止日 {HAISHI} と整合）")
    with open(JURYO, encoding="utf-8") as f:
        juryo = f.read()
    judge("2024年3月31日" in juryo and "定員70人" in juryo,
          "受領資料の廃止日・定員と一致")

    shutil.copy2(MOTO, SAKI)
    d = docx.Document(SAKI)
    mae = "\n".join([p.text for p in d.paragraphs]
                    + [c.text for t in d.tables for r in t.rows
                       for c in r.cells])

    print("  ── ① 2-7（p213）")
    p = d.paragraphs[213]
    ok = p.text.startswith("第3に、町内にあった軽費老人ホーム")
    judge(ok, "p213 の冒頭が合う", f"実際『{p.text[:20]}』")
    if ok:
        dan_kaku(p, HON_213)

    print("  ── ② 別表 事項八（表4）")
    t = d.tables[4]
    c = t.rows[8].cells[3]
    ok = HYO4_KYU in c.text
    judge(ok, "表4 行8 列3 に旧の文がある")
    if ok:
        dan_kaku(c.paragraphs[0], c.text.replace(HYO4_KYU, HYO4_SHIN))
        for q in c.paragraphs[1:]:
            for r in q.runs:
                r.text = ""

    print("  ── ③ 版表記")
    n_ban = 0
    parts = [q for sec in d.sections for part in (sec.header, sec.footer)
             for q in part.paragraphs] + list(d.paragraphs)
    for q in parts:
        for r in q.runs:
            if BAN_MOTO in r.text:
                r.text = r.text.replace(BAN_MOTO, BAN)
                n_ban += 1
    judge(n_ban >= 1, f"版表記を {BAN} に改めた（{n_ban}か所）")
    d.save(SAKI)

    d2 = docx.Document(SAKI)
    ato = "\n".join([q.text for q in d2.paragraphs]
                    + [c.text for t in d2.tables for r in t.rows
                       for c in r.cells])

    print("  ── 検査1　⚠ 定員を居所を失った人数と読ませない")
    judge("住まい70人分が失われた" not in ato, "「住まい70人分が失われた」が0件")
    judge("そのまま住まいを失ったものではありません" in ato,
          "定員と入所者の別を書いた")
    print("  ── 検査2　廃止日と県調査の値")
    judge(ato.count(HAISHI) >= 2, f"廃止日（{HAISHI}）が2か所以上")
    for go in (f"令和5年4月1日時点で{GEN['R5.4.1']}人",
               f"退所した{A['退所']}人", f"他の施設{A['他施設']}人"):
        judge(go in ato, f"「{go}」がある")
    print("  ── 検査3　⚠ 未確定のものを推測で埋めていない")
    judge(ato.count("［要確認］") == mae.count("［要確認］"),
          f"［要確認］の件数が同じ（{ato.count('［要確認］')}件）")
    judge("840" not in ato.replace(mae, "") and ato.count("840")
          == mae.count("840"), "延べ840人を使っていない")
    print("  ── 検査4　⚠ 施設名・法人名・連絡先を書いていない")
    # 電話は受領資料にある法人・県の番号に限る（町立病院の公的な番号は別）
    for na in ("蔵王グリーンホーム", "鶴寿会", "0224-84-4853", "0224-84-2262",
               "022-211-2556"):
        judge(na not in ato, f"「{na}」が0件")
    print("  ── 検査5　⚠ 既にあった数が消えていないか")
    kesu = sorted({m.group(0) for m in re.finditer(
        r"[0-9０-９][0-9０-９,.]*(?:人|％|千円|件|名|歳|回)", mae)
        if mae.count(m.group(0)) > ato.count(m.group(0))})
    judge(not kesu, "消えた数が0件", "／".join(kesu[:5]))
    print("  ── 検査6　受託者を主語とする語")
    aru = [w for w in ("受託者", "当社", "ビズアップ", "スクリプト")
           if w in ato]
    judge(not aru, "0件", "／".join(aru))

    print()
    print(f"  {SAKI}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        for m in futeki:
            print(f"      {m}")
        sys.exit(1)
    print(f"   ○ 検査1〜6　不適合0件（［要確認］{ato.count('［要確認］')}件・"
          f"［要協議］{ato.count('［要協議］')}件）")


if __name__ == "__main__":
    main()
