# -*- coding: utf-8 -*-
"""業務工程管理表の更新（認知症の人の数の推計・令和8年10月19日）

  python3 07_ソーススクリプト/update_wbs_R8.10.19b.py

なぜ更新するか
  ご相談（認知症の人数は高齢化率や見える化の数値から確認・推計できないか）
  を受けて、**認知症の人の数の推計値を算出し、計画素案 Ver.2.19 の 6-1 に
  掲げた。** これに伴い台帳を改める。

更新するもの
  03_確認事項一覧　No.157　推計値は算出済みになったため、
                    お願いする内容を**実数（自立度Ⅱ以上）と
                    年齢階級別の将来推計人口**に絞り込む
  04_成果品管理　　成果品3（推計用電子データ）に推計シートを登録
                    成果品4（計画書）を Ver.2.19 に
  08_作業順位　　　順位31 として記録する

⚠ **確認事項の番号は増やさない。** No.157 と同じ問い（認知症の人の数の
  把握）であり、同じ行の中で絞り込む。

⚠ **No.157 を「完了」にしない。** 推計値は出したが、
  ①認定調査の日常生活自立度Ⅱ以上の実数 ②年齢階級別の将来推計人口
  の2つは未受領である。「未実施」を「適合」に丸めない。
"""
import importlib.util
import os
import sys
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
from openpyxl.styles import Alignment, Font

HERE = os.path.dirname(os.path.abspath(__file__))
WBS = "川崎町_業務工程管理表.xlsx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月19日 更新】"

spec = importlib.util.spec_from_file_location(
    "calc_ninchisho", os.path.join(HERE, "calc_ninchisho_R8.10.19.py"))
C = importlib.util.module_from_spec(spec)
sys.modules["calc_ninchisho"] = C
spec.loader.exec_module(C)

SOAN = "川崎町_計画書素案_v2.19_認知症推計版.docx"
SHEET = "川崎町_認知症の人の数の推計_R8.10.19.xlsx"


def maru(x):
    return int(Decimal(str(x)).quantize(Decimal("1"),
                                        rounding=ROUND_HALF_UP))


def no157(nin, mci, korei, kani):
    return (
        f"{HIZUKE}認知症の人の数の**推計値**は算出し、計画素案 6-1 に"
        f"掲げました（認知症{nin}人・65歳以上の{nin / korei * 100:.1f}％／"
        f"軽度認知障害（ＭＣＩ）{mci}人・{mci / korei * 100:.1f}％）。"
        f"年齢階級別・男女別の有病率（厚生労働省。令和4年の悉皆調査による"
        f"国の推計値）を、町ご提供の年齢別人口集計表（令和8年6月末現在・"
        f"1歳刻み）に乗じて積み上げたものです。\n"
        f"⚠ **残るお願いは次の2つです。**\n"
        f"① 要介護認定の認定調査における認知症高齢者の日常生活自立度Ⅱ以上の"
        f"人数（令和6・7年度）。推計値は65歳以上人口の全体を母集団とするのに"
        f"対し、自立度Ⅱ以上は認定を受けておられる方の中での把握であり、"
        f"測るものが違います。**両者の差が、認定を受けておられない"
        f"認知症の方のおおよその規模**になり、早期発見・相談・サポーター"
        f"養成の配り方を考える材料になります。\n"
        f"② 年齢階級別（5歳階級）の将来推計人口（令和12・17・22年）。"
        f"基本指針（案）は「今後どのような推移を辿るのか」の推計を求めて"
        f"いますが、将来の推計には年齢階級別の将来人口が要ります。"
        f"本計画が第1号被保険者数に用いている国立社会保障・人口問題研究所"
        f"「日本の地域別将来推計人口（令和5年推計）」の5歳階級別の値を"
        f"ご提供いただくか、地域包括ケア「見える化」システムからの出力を"
        f"ご検討ください。いただければ同じ式でその日のうちに算出できます。\n"
        f"⚠ **見える化システムからは取れない見込みです。** 令和8年7月15日の"
        f"全国介護保険担当課長会議資料52頁は、別表一の事項ごとの把握方法を"
        f"示しており、事項十一（認知症の人の数）は「各自治体で把握」と"
        f"されています（事項二〜七・十は「見える化システム」と明記）。"
        f"また事項二（人口構造）で見える化システムから取れる人口は"
        f"**年齢三区分**であり、推計に要る5歳階級別の内訳には足りません。\n"
        f"⚠ **高齢化率からは推計できません。** 有病率は年齢階級により"
        f"1.1％（65〜69歳）から50.3％（90歳以上）まで変わるため、"
        f"要るのは65歳以上の「割合」ではなく「年齢階級別の内訳」です。"
        f"65歳以上人口に全国の12.3％を一律に乗じると{kani}人となりますが、"
        f"12.3％は性年齢調整後の全国値で本町の年齢構成を反映せず、"
        f"{kani - nin}人多く出ます。")


def iro(ws, r, n_col):
    for c in range(1, n_col + 1):
        ws.cell(r, c).font = FONT
        ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")


def main():
    if not os.path.exists(WBS):
        raise SystemExit("業務工程管理表がない：" + WBS)
    jinko = C.machi_kaikyu()
    korei = sum(v[2] for v in jinko.values())
    ninchi = C.santei(jinko, "認知症")
    mci = C.santei(jinko, "ＭＣＩ")
    nin, m, kani = (maru(ninchi["sei"]), maru(mci["sei"]),
                    maru(ninchi["kani"]))

    wb = openpyxl.load_workbook(WBS)

    # ══════ 03_確認事項一覧　No.157
    ws = wb["03_確認事項一覧"]
    n = 0
    for x in range(5, ws.max_row + 1):
        if str(ws.cell(x, 1).value or "").strip() == "157":
            ws.cell(x, 5).value = no157(nin, m, korei, kani)
            ws.cell(x, 6).value = "計画素案 6-1（推計値を掲載済み）"
            iro(ws, x, 10)
            n += 1
    if n != 1:
        raise SystemExit(f"No.157 の行が{n}件（1件のはず）")

    # ══════ 04_成果品管理
    ws = wb["04_成果品管理"]
    n_seika = 0
    for x in range(5, ws.max_row + 1):
        no = str(ws.cell(x, 1).value or "").strip()
        na = str(ws.cell(x, 3).value or "")
        if no == "3" and SHEET not in na:
            ws.cell(x, 3).value = na.rstrip() + "\n" + SHEET + "（9シート）"
            ws.cell(x, 8).value = (
                f"{HIZUKE}認知症の人の数の推計（認知症{nin}人・ＭＣＩ{m}人）"
                f"を登録した。基本指針（案）別表一 事項十一が求める"
                f"「認知症の人の数や推計値」のうち推計値の側。"
                f"⚠ **将来の推計値は年齢階級別の将来推計人口が要るため"
                f"未了**（確認事項No.157）。　"
                + str(ws.cell(x, 8).value or ""))
            iro(ws, x, 9)
            n_seika += 1
        elif no == "4" and "v2.18_範囲明記版" in na:
            ws.cell(x, 3).value = na.replace(
                "川崎町_計画書素案_v2.18_範囲明記版.docx",
                "川崎町_計画書素案_v2.19_認知症推計版.docx")
            ws.cell(x, 8).value = (
                f"{HIZUKE}素案を Ver.2.19 に上げた"
                f"（第6章6-1 に認知症の人の数の推計を新設。全114頁）。"
                f"⚠ **圧縮案（v2.18c）と別冊は Ver.2.18 に基づくもので、"
                f"6-1 の新設分を含んでいない。** 圧縮案のご判断"
                f"（確認事項No.83）をいただいた版で合わせる。　"
                + str(ws.cell(x, 8).value or ""))
            iro(ws, x, 9)
            n_seika += 1
    if n_seika != 2:
        raise SystemExit(f"04シートの該当行が{n_seika}件（2件のはず）")

    # ══════ 08_作業順位　順位31
    ws = wb["08_作業順位"]
    if any(str(ws.cell(x, 1).value or "").strip() == "31"
           for x in range(5, ws.max_row + 1)):
        raise SystemExit("順位31 が既にある")
    r = ws.max_row + 1
    while str(ws.cell(r - 1, 1).value or "").strip() != "30":
        r -= 1
        if r < 6:
            raise SystemExit("順位30 の行が見つからない")
    ws.insert_rows(r)
    for c, v in enumerate([
        "31", "認知症の人の数の推計（基本指針 別表一 事項十一）",
        "R8.10.19",
        "ご相談による。有病率は国が公表しており、年齢階級別・男女別の"
        "人口は町ご提供の年齢別人口集計表（令和8年6月末現在・1歳刻み）が"
        "ある。町のご回答を待たずに算出できる。"
        "⚠ 高齢化率だけでは足りず、65歳以上の年齢階級別の内訳が要る。",
        "束のうちNo.157（推計値の側）",
        "認知症の人の数の推計（xlsx・9シート）と計画素案 Ver.2.19",
        "半日", "完了",
        f"令和8年10月19日　認知症{nin}人（65歳以上の"
        f"{ninchi['sei'] / korei * 100:.1f}％）・ＭＣＩ{m}人"
        f"（{mci['sei'] / korei * 100:.1f}％）と推計し、素案 6-1 に"
        f"掲げた。⚠ **Ver.2.18 までの素案には認知症の人の数が"
        f"1か所もなかった**（方針を1-5の対比表に書いていただけ）。"
        f"同じ式を全国の年齢階級別人口に当て直すと国の公表値12.3％に"
        f"戻ることを検算で確かめている。"
        f"⚠ **将来の推計値と自立度Ⅱ以上の実数は未了**（No.157）。"
    ], start=1):
        ws.cell(r, c).value = v
    iro(ws, r, 9)

    wb.save(WBS)
    print(f"更新　{WBS}")

    # ══════════════════════════ 自己点検
    print()
    futeki = []

    def judge(ok, mi):
        print(("   ○ " if ok else "   × ") + mi)
        if not ok:
            futeki.append(mi)

    wb2 = openpyxl.load_workbook(WBS, data_only=True)
    ws = wb2["03_確認事項一覧"]
    nos = [str(ws.cell(x, 1).value or "").strip()
           for x in range(5, ws.max_row + 1)]
    nos = [x for x in nos if x.isdigit()]
    judge(len(nos) == len(set(nos)), f"確認事項に重複がない（{len(nos)}件）")
    kesson = [i for i in range(1, max(int(x) for x in nos) + 1)
              if str(i) not in nos]
    judge(not kesson, f"欠番がない（最大 {max(int(x) for x in nos)}）")
    row157 = [x for x in range(5, ws.max_row + 1)
              if str(ws.cell(x, 1).value or "").strip() == "157"][0]
    t157 = str(ws.cell(row157, 5).value or "")
    judge(f"{nin}人" in t157, f"No.157 に推計値 {nin}人 が入っている")
    judge("各自治体で把握" in t157,
          "No.157 に見える化から取れない旨の根拠が入っている")
    judge(str(ws.cell(row157, 8).value or "").strip() == "確認待ち",
          "No.157 は「確認待ち」のまま（完了に丸めていない）")

    ws = wb2["04_成果品管理"]
    na3 = na4 = ""
    for x in range(5, ws.max_row + 1):
        no = str(ws.cell(x, 1).value or "").strip()
        if no == "3":
            na3 = str(ws.cell(x, 3).value or "")
        if no == "4":
            na4 = str(ws.cell(x, 3).value or "")
    judge(SHEET in na3, "成果品3 に推計シートが登録されている")
    judge(SOAN in na4 and "v2.18_範囲明記版" not in na4,
          "成果品4 が Ver.2.19 を指している")

    ws = wb2["08_作業順位"]
    jun = [str(ws.cell(x, 1).value or "").strip()
           for x in range(5, ws.max_row + 1)]
    jun = [x for x in jun if x.isdigit()]
    judge(jun == [str(i) for i in range(1, len(jun) + 1)],
          f"順位が1から{len(jun)}まで連なっている")

    # 成果品が実在するか
    judge(os.path.exists("01_第10期_最新版成果品/" + SOAN),
          "素案 Ver.2.19 が実在する")
    judge(os.path.exists("05_試算・管理シート/" + SHEET),
          "推計シートが実在する")

    print()
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 自己点検 10件　不適合0件")


if __name__ == "__main__":
    main()
