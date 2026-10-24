# -*- coding: utf-8 -*-
"""計画素案 Ver.2.18（反証で見つかった記述の是正・令和8年10月17日）

  入力　01_第10期_最新版成果品/川崎町_計画書素案_v2.17_中長期補完版.docx
  出力　01_第10期_最新版成果品/川崎町_計画書素案_v2.18_範囲明記版.docx

中長期の算定を独立に反証させたところ、**数値そのものの誤りは出なかったが、
記述の範囲について2件の指摘が当たった。** いずれも本版で直す。

══════ 是正1　5-5 の「入所・居住系」の範囲の書き方（重要）
  旧　「入所・居住系（施設・地域密着型・特定施設）が680,226千円・68.1％」
  ⚠ **括弧を字義どおり読むと数が合わない。**
     施設サービス計493,352千円＋地域密着型計182,194千円
     ＋特定施設6,255千円＝681,801千円であり、1,575千円多い。
     差は地域密着型のうち**在宅のサービス**
     （定期巡回・随時対応型訪問介護看護389千円＋地域密着型通所介護377千円
     ＋看護小規模多機能型居宅介護809千円）である。
  新　含むサービスを名指しする。

══════ 是正2　9-1 の中長期の注に、算定の出所を明記する
  ⚠ **地域包括ケア「見える化」システムの総括表には、サービス別の
     令和17・22年度の給付費が別に入っている。** しかし
     ①令和7年度に給付実績のある7サービスを令和9年度以降0と置いている
       （介護予防訪問入浴介護169千円→0 ほか）
     ②令和11年度から令和27年度まで同じ値が並び、認定者数が伸びても
       給付費が動かない
     という状態であり、**本計画では用いていない。**
     本計画は令和9〜11年度についても同様に用いておらず、
     **全年度を通じて年報の実績に認定者数の伸びを乗じる方法で
     一貫して算定している。**
  ⚠ 総括表のサービス別の値を合計すると令和17年度1,114,654千円となり、
     本計画の1,110,612千円と4,042千円（0.4％）違う。
     **県への事前協議で問われうる点であるため、確認事項として起票する。**

⚠ **数値は1円も変えていない。** 直すのは記述の範囲だけである。
"""
import os
import re
import shutil
import sys
import zipfile

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from fix_soan_v111 import set_el                             # noqa: E402

SRC = "01_第10期_最新版成果品/川崎町_計画書素案_v2.17_中長期補完版.docx"
DST = "01_第10期_最新版成果品/川崎町_計画書素案_v2.18_範囲明記版.docx"

HONBUN = [
    # 是正1　5-5 の範囲
    ("入所・居住系（施設・地域密着型・特定施設）が680,226千円・68.1％を占め、",
     "入所・居住系（施設サービス、認知症対応型共同生活介護、"
     "地域密着型介護老人福祉施設入所者生活介護及び"
     "特定施設入居者生活介護）が680,226千円・68.1％を占め、"),
    # 是正2　9-1 の中長期の注に出所を足す
    ("令和17年度・令和22年度は、第10期計画期間（令和9〜11年度）の"
     "見込量とは位置づけが異なり、計画の数値ではなく長期の見通しです。",
     "令和17年度・令和22年度は、第10期計画期間（令和9〜11年度）の"
     "見込量とは位置づけが異なり、計画の数値ではなく長期の見通しです。"
     "なお本節の見込量は、全年度を通じて、介護保険事業状況報告（年報・"
     "令和7年度）の要介護度別の実績に要介護度別の認定者数の伸びを乗じる"
     "方法で算定しています。"),
]


def hon(el):
    return "".join(x.text or "" for x in el.iter(qn("w:t")))


def main():
    if not os.path.exists(SRC):
        raise SystemExit("入力がない：" + SRC)
    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body

    n_hon = 0
    nokori = []
    for furui, atarashii in HONBUN:
        atta = False
        for el in body.iterchildren():
            if el.tag != qn("w:p"):
                continue
            s = hon(el)
            if furui in s:
                set_el(el, s.replace(furui, atarashii))
                n_hon += 1
                atta = True
                break
        if not atta:
            nokori.append(furui[:34])
    if nokori:
        raise SystemExit("本文が見つからない：" + "／".join(nokori))

    n_ver = 0
    for el in body.iter(qn("w:t")):
        if el.text and "Ver.2.17" in el.text:
            el.text = el.text.replace("Ver.2.17", "Ver.2.18")
            n_ver += 1
    if n_ver == 0:
        raise SystemExit("本文に Ver.2.17 の表記が見つからない")
    for el in body.iterchildren():
        if el.tag == qn("w:p"):
            s = hon(el)
            if "（中長期補完版）" in s:
                set_el(el, s.replace("（中長期補完版）", "（範囲明記版）"))
                break
    doc.save(DST)
    naoshi_header(DST)

    # ══════════════════════════ 自己点検
    d2 = docx.Document(DST)
    d0 = docx.Document(SRC)
    zen = "".join(n.text or "" for n in d2.element.body.iter(qn("w:t")))
    ok, ng = [], []

    # 検査1　字義どおり読むと合わない書き方が消えていること
    if "入所・居住系（施設・地域密着型・特定施設）" in zen:
        ng.append("検査1　5-5 の古い範囲の書き方が残っている")
    else:
        ok.append("検査1　5-5 の「（施設・地域密着型・特定施設）」という"
                  "字義どおりでは合わない書き方は消えている")

    # 検査2　含むサービスが名指しされていること
    NA = ["施設サービス", "認知症対応型共同生活介護",
          "地域密着型介護老人福祉施設入所者生活介護",
          "特定施設入居者生活介護"]
    m = re.search(r"入所・居住系（([^）]*(?:（[^）]*）)?[^）]*)）"
                  r"が680,226千円", zen)
    if m is None:
        # 入れ子の括弧がないため素直に探す
        i = zen.find("入所・居住系（")
        j = zen.find("）が680,226千円")
        naka = zen[i + 7:j] if (i >= 0 and j > i) else ""
    else:
        naka = m.group(1)
    tarinai = [x for x in NA if x not in naka]
    if tarinai:
        ng.append("検査2　範囲に名指しされていないサービス："
                  + "／".join(tarinai))
    else:
        ok.append(f"検査2　入所・居住系に含む{len(NA)}区分が"
                  "すべて名指しされている")

    # 検査3　算定の方法が 9-1 に書かれていること
    if "要介護度別の認定者数の伸びを乗じる" in zen:
        ok.append("検査3　9-1 に算定の方法（年報の実績×認定者数の伸び）が"
                  "書かれている")
    else:
        ng.append("検査3　9-1 に算定の方法が書かれていない")

    # 検査4　数値が1円も動いていないこと
    #   ⚠ 本版は記述だけを直す。数の列は Ver.2.17 と同一でなければならない。
    def kazu_zenbu(d):
        out = []
        for t in d.tables:
            for r in t.rows:
                for c in r.cells:
                    for mm in re.finditer(r"[\d,]+(?:\.\d+)?",
                                          c.text.replace("\n", "")):
                        out.append(mm.group(0))
        return out

    a, b = kazu_zenbu(d0), kazu_zenbu(d2)
    if a != b:
        chigai = [(x, y) for x, y in zip(a, b) if x != y][:6]
        ng.append(f"検査4　表の中の数が変わっている（{len(a)}対{len(b)}）："
                  + "／".join(f"{x}→{y}" for x, y in chigai))
    else:
        ok.append(f"検査4　表の中の数は1つも変わっていない（{len(a)}個を突合）")

    # 検査5　段落・表の数
    if len(d2.tables) != len(d0.tables) or \
            len(d2.paragraphs) != len(d0.paragraphs):
        ng.append(f"検査5　段落・表の数が変わっている")
    else:
        ok.append(f"検査5　段落 {len(d2.paragraphs)}／表 {len(d2.tables)}"
                  "（いずれも変わらず）")

    # 検査6　版表記
    z = zipfile.ZipFile(DST)
    hdr = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>",
                             z.read("word/header1.xml").decode("utf-8")))
    hoka = sorted({v for v in re.findall(r"Ver\.\d+\.\d+", zen)
                   if v != "Ver.2.18"})
    if "Ver.2.18" not in hdr:
        ng.append(f"検査6　ヘッダーの版表記が更新されていない：{hdr}")
    elif hoka:
        ng.append("検査6　現行の版以外の版表記が残っている："
                  + "／".join(hoka))
    else:
        ok.append("検査6　版表記は Ver.2.18 のみ")

    # 検査7　強調記号
    for w in ("**", "__"):
        if w in zen:
            ng.append(f"検査7　強調記号が残っている：{w}")
    if not any("検査7" in x for x in ng):
        ok.append("検査7　強調記号は入っていない")

    print("計画素案 Ver.2.18（範囲明記版）")
    print("保存：", DST)
    print(f"  是正　本文 {n_hon}か所（5-5 の範囲・9-1 の算定の出所）")
    print(f"  版表記 {n_ver}か所")
    print("  ── 自己点検")
    for s in ok:
        print("   ○", s)
    if ng:
        for s in ng:
            print("   ×", s)
        sys.exit(1)


def naoshi_header(path):
    z = zipfile.ZipFile(path)
    names = z.namelist()
    data = {n: z.read(n) for n in names}
    z.close()
    n_fix = 0
    for n in names:
        if not re.match(r"word/(header|footer)\d*\.xml", n):
            continue
        s = data[n].decode("utf-8")
        s2 = re.sub(r"Ver\.2\.17", "Ver.2.18", s)
        if s2 != s:
            data[n] = s2.encode("utf-8")
            n_fix += 1
    if n_fix == 0:
        raise SystemExit("ヘッダーの版表記（Ver.2.17）が見つからない")
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    shutil.move(tmp, path)


if __name__ == "__main__":
    main()
