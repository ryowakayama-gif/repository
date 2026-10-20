# -*- coding: utf-8 -*-
"""業務工程管理表の更新（通信が開いたことによる前進・令和8年10月20日）

  python3 07_ソーススクリプト/update_wbs_R8.10.20.py

なぜ更新するか
  環境の通信の状況を見直したところ、2つが前に進められるようになった。
  ⚠ **いずれも町のご回答を待たずに当方で進められる。**

  ① `laws.e-gov.go.jp` が通るようになった（昨日は403・本日200）
     → No.181（資料編の条文）の条文の原文を当方で取れる
  ② 介護サービス情報公表システムの事業所の一覧が、
     厚生労働省のオープンデータ（ＣＳＶ）として `www.mhlw.go.jp`（許可済み）
     から取れることが分かった
     → 作業順位26（県の公表名簿との突合）と No.29（事業所の一覧）が
       当方で進められる

⚠ **確認事項の番号は増やさない。** いずれも既にある事項の状況の変化である。

⚠ **他団体の固有名称は書かない。** この手は他の案件で先に見つかったものだが、
  台帳には「他の案件で用いられていた手」とだけ記す（CLAUDE.md 3）。
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月20日 更新】"

OPENDATA = ("https://www.mhlw.go.jp/stf/kaigo-kouhyou_opendata.html")

# 列5＝確認をお願いする内容／列10＝回答・対応
#   ⚠ 列8は「状態」である。書き込んではならない。
TSUIKI = {
    "29": (10,
           f"{HIZUKE}⚠ **当方で先に押さえられる見込みが立ちました。**"
           f"介護サービス情報公表システムの画面は自動のアクセスを"
           f"受け付けませんが、厚生労働省が同じ内容を"
           f"**サービス種類別のＣＳＶ（オープンデータ）**として公表しており、"
           f"許可済みのドメインから取れます（{OPENDATA}）。"
           f"本日、市町村コード043249で川崎町の事業所が取れることを"
           f"確かめました（例：認知症対応型共同生活介護2件。"
           f"事業所名・事業所番号・法人番号・法人の名称・所在地・定員の"
           f"24列）。他の案件で用いられていた手です。"
           f"⚠ **指定の年月日・町の指定か県の指定かの別は"
           f"オープンデータに含まれません**ので、その部分は引き続き"
           f"ご確認をお願いします。　"),
    "181": (10,
            f"{HIZUKE}⚠ **条文の原文を当方で取れるようになりました。**"
            f"法令検索（laws.e-gov.go.jp）が、昨日（10月19日）は403でしたが"
            f"本日は200で本文が取れます。資料編の条文抜粋は、"
            f"添付の引用をそのまま使わず原文から起こします。"
            f"名簿（委員・事業所）のご提供は引き続きお願いします。　"),
}

JUNI_33 = [
    "33", "環境の通信の見直しと、公表システムのオープンデータの発見",
    "R8.10.20",
    "ご指示による。環境の通信の設定を見直すにあたり、"
    "どのホストが要るかを用途つきで一覧にした。町のご回答を待たずに行える。",
    "束4のうちNo.29／束11のうちNo.181",
    "許可ドメイン候補の整理（md）と疎通の確認（py・22ホスト）",
    "半日", "完了",
    "令和8年10月20日　⚠ **2つが前に進められるようになった。**"
    "①法令検索（laws.e-gov.go.jp）が通るようになり、資料編の条文を"
    "原文から起こせる（No.181）。"
    "②⚠ **介護サービス情報公表システムの事業所の一覧が、厚生労働省の"
    "オープンデータ（ＣＳＶ）として許可済みのドメインから取れる。**"
    "市町村コード043249で川崎町の事業所を取れることを確かめた"
    "（事業所名・事業所番号・法人番号・法人の名称・所在地・定員の24列）。"
    "作業順位26（県の公表名簿との突合）が、許可を待たずに行える。"
    "⚠ 最優先で足していただきたいのは www.ipss.go.jp"
    "（社人研の市区町村別・5歳階級別の将来推計人口。認知症の人の数の"
    "将来の推計に直に効く）と www.satsuki-jutaku.mlit.go.jp"
    "（サ高住。⚠ 元のドメインを許可しても301で移るため届かない）。"
    "整理は 10_内部資料_非成果品/ に置いた"
    "（他団体の固有名称を含むため成果品の置き場に置かない）。",
]

TSUIKI_26 = (
    f"　{HIZUKE}⚠ **突合の実体を当方で進められるようになった。**"
    f"厚生労働省が介護サービス情報公表システムの事業所の一覧を"
    f"ＣＳＶ（オープンデータ）で公表しており、許可済みのドメインから取れる。"
    f"本日、市町村コード043249で川崎町の事業所が取れることを確かめた。"
    f"⚠ **ただしオープンデータに指定の年月日・指定の主体（町か県か）は"
    f"含まれない**ため、その部分は県の名簿又は町のご確認による。")


def iro(ws, r, n):
    for c in range(1, n + 1):
        ws.cell(r, c).font = FONT
        ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")


def main():
    if not os.path.exists(WBS):
        raise SystemExit("業務工程管理表がない：" + WBS)
    wb = openpyxl.load_workbook(WBS)

    # ── 03_確認事項一覧
    ws = wb["03_確認事項一覧"]
    n = 0
    sumi = 0
    for x in range(5, ws.max_row + 1):
        no = str(ws.cell(x, 1).value or "").strip()
        if no not in TSUIKI:
            continue
        retsu, moji = TSUIKI[no]
        mae = str(ws.cell(x, retsu).value or "")
        if HIZUKE in mae:
            sumi += 1
            continue
        ws.cell(x, retsu).value = moji + mae
        iro(ws, x, 10)
        n += 1
    if n + sumi != len(TSUIKI):
        raise SystemExit(f"該当行が{n + sumi}件（{len(TSUIKI)}件のはず）")

    # ── 08_作業順位　順位26 に追記、順位33 を足す
    ws2 = wb["08_作業順位"]
    n26 = 0
    for x in range(5, ws2.max_row + 1):
        if str(ws2.cell(x, 1).value or "").strip() == "26":
            mae = str(ws2.cell(x, 9).value or "")
            if HIZUKE not in mae:
                ws2.cell(x, 9).value = mae + TSUIKI_26
            iro(ws2, x, 9)
            n26 += 1
    if n26 != 1:
        raise SystemExit(f"順位26 の行が{n26}件（1件のはず）")

    jun = [str(ws2.cell(x, 1).value or "").strip()
           for x in range(5, ws2.max_row + 1)]
    if "33" not in jun:
        r = ws2.max_row + 1
        while str(ws2.cell(r - 1, 1).value or "").strip() != "32":
            r -= 1
            if r < 6:
                raise SystemExit("順位32 の行が見つからない")
        ws2.insert_rows(r)
        for c, v in enumerate(JUNI_33, start=1):
            ws2.cell(r, c).value = v
        iro(ws2, r, 9)

    wb.save(WBS)
    print(f"更新　{WBS}")

    # ══════════════════════════ 自己点検
    print()
    futeki = []

    def judge(ok, mi):
        print(("   ○ " if ok else "   × ") + mi)
        if not ok:
            futeki.append(mi)

    wb2 = openpyxl.load_workbook(WBS, data_only=True)
    ws = wb2["03_確認事項一覧"]
    nos = [str(ws.cell(x, 1).value or "").strip()
           for x in range(5, ws.max_row + 1)]
    nos = [x for x in nos if x.isdigit()]
    judge(len(nos) == len(set(nos)), f"重複がない（{len(nos)}件）")
    saidai = max(int(x) for x in nos)
    judge(not [i for i in range(1, saidai + 1) if str(i) not in nos],
          f"欠番がない（最大 {saidai}）")
    judge(saidai == 186, f"確認事項を増やしていない（最大 {saidai}）")

    hon = {}
    for x in range(5, ws.max_row + 1):
        no = str(ws.cell(x, 1).value or "").strip()
        if no.isdigit():
            hon[no] = "\n".join(str(ws.cell(x, c).value or "")
                                for c in range(1, 11))
    judge("オープンデータ" in hon.get("29", ""),
          "No.29 にオープンデータの件が入っている")
    judge("043249" in hon.get("29", ""),
          "No.29 に川崎町の市町村コードが入っている")
    judge("laws.e-gov.go.jp" in hon.get("181", ""),
          "No.181 に法令検索の件が入っている")
    for no in ("29", "181"):
        x = [r for r in range(5, ws.max_row + 1)
             if str(ws.cell(r, 1).value or "").strip() == no][0]
        if str(ws.cell(x, 8).value or "").strip() != "確認待ち":
            judge(False, f"No.{no} が「確認待ち」でない")
    judge(True, "まだご回答が要る部分を「確認待ち」のままにした")

    ws2 = wb2["08_作業順位"]
    jun = [str(ws2.cell(x, 1).value or "").strip()
           for x in range(5, ws2.max_row + 1)]
    jun = [x for x in jun if x.isdigit()]
    judge(jun == [str(i) for i in range(1, len(jun) + 1)],
          f"順位が1から{len(jun)}まで連なっている")

    # 他団体の固有名称を書いていないこと
    import re
    zen = []
    for sh in ("03_確認事項一覧", "08_作業順位"):
        w = wb2[sh]
        for row in w.iter_rows(values_only=True):
            zen += [str(v) for v in row if v is not None]
    atari = [s[:60] for s in zen
             if re.search(r"(北塩原村|金ケ崎町|金ヶ崎町|大雪)", s)]
    judge(not atari, "台帳に他団体の固有名称を書いていない")

    print()
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 自己点検 9件　不適合0件")


if __name__ == "__main__":
    main()
