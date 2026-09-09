import re,unicodedata,openpyxl,json
from decimal import Decimal,ROUND_HALF_UP
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
def pct1(k,n):
    if not n: return None
    return float((Decimal(k*100)/Decimal(n)).quantize(Decimal('0.1'),rounding=ROUND_HALF_UP))
XL='../R8調査データ/川崎町_在宅介護実態調査_集計データ_20260807.xlsx'
DOC='/home/user/repository/kawasaki_project/09_元資料/R8調査データ/R8.9.9受領版/川崎町_資料編_R8.9.9版.docx'
wb=openpyxl.load_workbook(XL,data_only=True); ws=wb['データ']
rows=[r for r in range(5,ws.max_row+1) if ws.cell(r,1).value is not None and str(ws.cell(r,1).value).strip()!='計']
hdr={}
for c in range(1,ws.max_column+1):
    g=ws.cell(1,c).value
    if g: hdr.setdefault(str(g).strip(),[]).append(c)
def norm(s):
    s=unicodedata.normalize('NFKC',str(s)); s=re.sub(r'[\s　]+','',s)
    s=re.sub(r'[（）()［］\[\]「」・,、。／/？?]','',s)
    return s.replace('～','-').replace('〜','-').replace('~','-')
def qblock(sc):
    cols=[];n=None
    for c in range(sc,ws.max_column+1):
        k=ws.cell(4,c).value
        if k is None: break
        if str(k).strip()=='n': n=c; break
        cols.append(c)
    return cols,n
M=('〇','○'); SEX=hdr['性別'][0]; AGE=hdr['年齢階級'][0]; CARE=hdr['要介護状態区分'][0]
def band(v):
    v=str(v or '')
    if v in ('65歳未満','65～69歳','70～74歳'): return norm('65歳未満～74歳')
    if v in ('75～79歳','80～84歳'): return norm('75～84歳')
    if v in ('85～89歳','90歳以上'): return norm('85～94歳')
def sel(level):
    L=norm(level)
    if L in ('-',''): return rows
    for c in (SEX,AGE,CARE):
        g=[r for r in rows if norm(ws.cell(r,c).value or '')==L]
        if g: return g
    if '／' in level or '/' in level:
        s,b=re.split('[／/]',level,1)
        return [r for r in rows if norm(ws.cell(r,SEX).value or '')==norm(s) and band(ws.cell(r,AGE).value)==norm(b)]
QB={}
for name,cs in hdr.items():
    if not re.match(r'^[AB]問',name): continue
    c=cs[0]
    while c<=ws.max_column:
        sub=str(ws.cell(2,c).value or '').strip(); cols,n=qblock(c)
        if not cols: break
        QB.setdefault(name,[]).append((sub,cols,n))
        c=(n or cols[-1])+1
        v=str(ws.cell(1,c).value or '').strip()
        if v and v!=name: break
SVC={'訪問介護（ホームヘルプ）':'訪問介護','訪問入浴介護':'訪問入浴','訪問看護':'訪問看護','訪問リハビリテーション':'訪問リハ',
 '通所介護（デイサービス）':'通所介護','通所リハビリテーション（デイケア）':'通所リハ','夜間対応型訪問介護':'夜間対応型訪問',
 '定期巡回・随時対応型訪問介護看護':'定期巡回','小規模多機能型居宅介護':'小多機','看護小規模多機能型居宅介護':'看多機',
 '短期入所生活介護・療養介護（ショートステイ）':'ショートステイ','居宅療養管理指導':'居宅療養管理'}
doc=Document(DOC)
def blocks(d):
    for ch in d.element.body.iterchildren():
        if ch.tag==qn('w:p'): yield Paragraph(ch,d)
        elif ch.tag==qn('w:tbl'): yield Table(ch,d)
def ctext(c): return " ".join(p.text.strip() for p in c.paragraphs if p.text.strip())
cap='';bad=[];okc=0;nt=0;unres=[];small=[];h1ok=0;h1bad=[];h1n=0
for b in blocks(doc):
    if isinstance(b,Paragraph):
        if b.text.strip(): cap=b.text.strip()
        continue
    grid=[[ctext(c) for c in r.cells] for r in b.rows]
    m1=re.match(r'^付表1-\d+\s+(A問\d+|B問\d+)\s+(.*)$',cap)
    if m1:
        grp,rest=m1.group(1),re.sub(r'（複数回答）|（[^）]*選択可）','',m1.group(2)).strip()
        bl=[x for x in QB.get(grp,[]) if x[0]==SVC.get(rest)] if grp=='A問8' else QB.get(grp,[])
        if len(bl)!=1: unres.append('付表1 '+cap[:30]); continue
        _,cols,ncol=bl[0]; h1n+=1
        l2c={norm(ws.cell(3,c).value):c for c in cols}
        n_true=sum(1 for x in rows if ws.cell(x,ncol).value not in (None,'',0))
        for r in grid[1:]:
            if len(r)<2: continue
            lab=norm(r[0])
            if lab in ('有効回答数n','有効回答数','合計','計','n','総数'):
                try: v=int(re.sub(r'[^0-9]','',r[1]))
                except: continue
                if v!=n_true: h1bad.append(f"{cap[:26]} 有効回答数 {v}→正解{n_true}")
                else: h1ok+=1
                continue
            c=l2c.get(lab)
            if c is None: continue
            kt=sum(1 for x in rows if ws.cell(x,c).value in M); pt=pct1(kt,n_true)
            e=[]
            mm=re.match(r'^([0-9,]+)',r[1].replace(' ',''))
            if mm:
                kr=int(mm.group(1).replace(',',''))
                if kr!=kt: e.append(f"件数{kr}→{kt}")
            for cell in r[2:]:
                p=re.match(r'^([0-9.]+)\s*[%％]',cell.replace(' ',''))
                if p:
                    pr=float(p.group(1))
                    if pt is not None and abs(pr-pt)>0.001: e.append(f"割合{pr}%→{pt}%")
                    break
            if e: h1bad.append(f"{cap[:24]}「{r[0][:16]}」"+" ".join(e))
            else: h1ok+=1
        continue
    m=re.match(r'^付表2-\d+\s+(A問\d+|B問\d+)\s+(.*)$',cap)
    if not m: continue
    grp,rest=m.group(1),re.sub(r'（\d/\d）|［.*$','',m.group(2)).strip()
    bl=[x for x in QB.get(grp,[]) if x[0]==SVC.get(rest)] if grp=='A問8' else QB.get(grp,[])
    if len(bl)!=1: unres.append('付表2 '+cap[:30]); continue
    _,cols,ncol=bl[0]
    l2c={norm(ws.cell(3,c).value):c for c in cols}
    head=grid[0]; colmap=[l2c.get(norm(h)) for h in head[3:]]
    if any(c is None for c in colmap):
        # 見出しが省略記号で切れている場合は前方一致で救う
        colmap=[]
        for h in head[3:]:
            hn=norm(h).rstrip('…').rstrip('.')
            hit=[c for k,c in l2c.items() if k.startswith(hn[:10])]
            colmap.append(hit[0] if len(hit)==1 else None)
        if any(c is None for c in colmap):
            unres.append(f"付表2 {cap[:30]} 列不明"); continue
    nt+=1
    for r in grid[1:]:
        if len(r)<4: continue
        lev=r[1] or '－'; g=sel(lev)
        if g is None: bad.append(f"{cap[:24]}：区分「{lev}」不明"); continue
        ntr=sum(1 for x in g if ws.cell(x,ncol).value not in (None,'',0))
        try: nr=int(re.sub(r'[^0-9]','',r[2]))
        except: continue
        if nr!=ntr: bad.append(f"{cap[:24]}／{lev}：n {nr}→正解{ntr}")
        else: okc+=1
        for i,cell in enumerate(r[3:]):
            if i>=len(colmap): break
            mm=re.match(r'^([0-9]+)\s*\(?([0-9.]+)%\)?',cell.replace(' ','').replace('件',''))
            if not mm: continue
            kr,pr=int(mm.group(1)),float(mm.group(2))
            kt=sum(1 for x in g if ws.cell(x,colmap[i]).value in M); pt=pct1(kt,ntr)
            e=[]
            if kr!=kt: e.append(f"件数{kr}→{kt}")
            if pt is not None and abs(pr-pt)>0.001: e.append(f"割合{pr}%→{pt}%")
            if e: bad.append(f"{cap[:22]}／{lev}／{head[3+i][:12]}："+" ".join(e))
            else: okc+=1
            if 0<kt<=2 and ntr<10: small.append(f"{cap[:20]}／{lev}(n={ntr})／{head[3+i][:12]}={kt}件")
print(f'付表1：{h1n}表／一致 {h1ok}／不一致 {len(h1bad)}')
for x in h1bad[:15]: print('  ✗',x)
print(f'付表2：{nt}表／一致 {okc}／不一致 {len(bad)}')
for x in bad[:20]: print('  ✗',x)
if unres:
    print(f'\n対応不可 {len(unres)}件'); [print('  ?',u) for u in unres]
print(f'\nn<10 の層で件数1〜2のセル：{len(small)}箇所')
