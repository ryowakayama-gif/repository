# -*- coding: utf-8 -*-
"""町内事業所の実態（公表システム本体からの採録）の管理シート
（令和8年10月20日）

  python3 07_ソーススクリプト/build_jigyosho_jittai_R8.10.20.py

  00_この表について
  01_入所・入居の状況　　定員・在所者数・入所率・待機者数
  02_職員の状況　　　　　⚠ **指定ごとの報告であり合算してはならない**
  03_職種別の従業者　　　指定×職種（実人数と常勤換算）
  04_指定の年月日　　　　指定権者（町か県か）の別
  05_退所先　　　　　　　前年度の退所先（施設系）
  06_残るもの　　　　　　公表システムにも項目がないもの
  07_出所と手順

数値の正本　07_ソーススクリプト/data_jigyosho_jittai_R8.10.20.py
出すもの　　05_試算・管理シート/川崎町_町内事業所の実態_公表システム_R8.10.20.xlsx
"""
import importlib.util
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location(
    "data_jigyosho", os.path.join(HERE, "data_jigyosho_jittai_R8.10.20.py"))
D = importlib.util.module_from_spec(spec)
sys.modules["data_jigyosho"] = D
spec.loader.exec_module(D)

OUT = ("05_試算・管理シート/"
       "川崎町_町内事業所の実態_公表システム_R8.10.20.xlsx")

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="D9E1F2")
WARN = PatternFill("solid", fgColor="FCE4D6")
OK = PatternFill("solid", fgColor="E2EFDA")
WRAP = Alignment(vertical="top", wrap_text=True)


def sheet(wb, name, head, widths):
    ws = wb.create_sheet(name)
    for c, (h, w) in enumerate(zip(head, widths), start=1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True)
        cell.fill = HEAD
        cell.border = BOX
        cell.alignment = WRAP
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    return ws


def put(ws, rows, fill=None):
    for i, r in enumerate(rows, start=2):
        for jj, v in enumerate(r, start=1):
            cell = ws.cell(i, jj, v)
            cell.border = BOX
            cell.alignment = WRAP
        f = fill(r) if fill else None
        if f:
            for jj in range(1, len(r) + 1):
                ws.cell(i, jj).fill = f
    return ws


def main():
    j = D.yomu()
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ordered = sorted(j.values(),
                     key=lambda d: (d["サービスコード"], d["事業所番号"]))

    # ══════ 00
    ws = sheet(wb, "00_この表について", ["項", "内容"], [26, 118])
    put(ws, [
        ["何の表か",
         "宮城県柴田郡川崎町（保険者番号04324）の町内にある介護保険の"
         "指定事業所について、介護サービス情報公表システム（本体）の"
         "詳細ページから採録した実態。令和8年10月20日に取得。"],
        ["なぜ作ったか",
         "入所率・待機者数・職員数・指定の年月日は、同日に取得した"
         "オープンデータ（24列）に含まれないため「事業所ヒアリングで"
         "なければ取れないもの」として整理していたが、公表システム"
         "本体の詳細ページには現に載っていた。ヒアリングを行わない"
         "整理の可否を判断するために、まず採録した。"],
        ["対象",
         f"{len(j)}の指定（8事業所・3法人／延べ14サービス）。"
         "対象はオープンデータの抽出結果から起こしており、"
         "事業所番号を書き写していない。"],
        ["いつ時点の値か",
         "⚠ **事業所ごとに記入日が違う。** "
         + "／".join(sorted(set(d["記入日"] for d in ordered)))
         + "。年1回の報告であり、同じ年度の値ではない。"],
        ["⚠ 合算してはならないもの",
         "職員数は「事業所番号×サービス」ごとの報告であり、同一施設が"
         "複数の指定を持つ場合は同じ職員が重複して報告される"
         "（特別養護老人ホームと地域密着型介護老人福祉施設、"
         "介護老人保健施設と短期入所療養介護・通所リハビリテーション）。"
         "町全体の介護職員数として足し合わせることはできない。"],
        ["⚠ 稼働率ではないもの",
         "入所率は「在所者数÷定員」であり、延べ利用日数による稼働率では"
         "ない。また町外の方の利用も含むため、第9章9-1の本町の被保険者の"
         "利用実績と割ってはならない。"],
        ["計画書に載せないもの",
         "⚠ 町内は8事業所・3法人であるため、事業所ごとの職員数・採用者数・"
         "退職者数を計画書に載せると法人が識別され得る。計画書には"
         "サービスの種類ごとの集計又は数を伴わない記述で載せる。"],
        ["原典",
         "09_元資料/その他データ/公表システム_詳細_R8.10.20/（14ページ）／"
         "09_元資料/その他データ/公表システム_詳細_川崎町_R8.10.20.csv"],
    ])

    # ══════ 01　入所・入居の状況
    ws = sheet(wb, "01_入所・入居の状況",
               ["サービスの種類", "事業所名", "指定権者", "定員",
                "定員の原文", "在所者数", "入所率", "待機者数",
                "平均年齢", "平均在所日数", "記入日", "備考"],
               [26, 24, 8, 7, 14, 8, 9, 9, 9, 11, 11, 40])
    rows = []
    for d in ordered:
        if d["サービスコード"] not in D.NYUSHO and d["定員"] is None:
            continue
        biko = ""
        if d["サービスコード"] == "320":
            biko = "町が指定権者。ユニット単位の定員"
        elif d["サービスコード"] == "540":
            biko = "⚠ 待機者数の項目が公表の様式にない"
        elif d["サービスコード"] in ("150", "160"):
            biko = "⚠ 1日当たりの利用定員。入所率は算定しない"
        elif d["サービスコード"] == "210":
            biko = ("⚠ 空床利用型。定員は併設する地域密着型介護老人福祉"
                    "施設の床を用いる。入所率は算定しない")
        rows.append([d["サービスの種類"], d["事業所名"], d["指定権者"],
                     d["定員"], d["定員の原文"], d["在所者数"],
                     (round(d["入所率"], 1) if d["入所率"] else "―"),
                     (d["待機者数"] if d["待機者数"] is not None else "―"),
                     d["平均年齢"], d["平均在所日数"], d["記入日"], biko])
    put(ws, rows, lambda r: WARN if str(r[11]).startswith("⚠") else None)
    n_nyu = len(rows)

    # ══════ 02　職員の状況
    ws = sheet(wb, "02_職員の状況",
               ["サービスの種類", "事業所名", "実人数の合計",
                "常勤換算の合計", "前年度の採用者数", "前年度の退職者数",
                "職員1人当たりの入所者数", "記入日", "⚠ 注意"],
               [26, 24, 12, 14, 14, 14, 18, 11, 46])
    rows = []
    for d in ordered:
        chui = ""
        if d["事業所番号"] in ("0452280019",) and d["サービスコード"] != "520":
            chui = "⚠ 介護老人保健施設と同一の職員体制。合算しない"
        if d["サービスコード"] == "540":
            chui = "⚠ 併設する特別養護老人ホームと一体の職員体制。合算しない"
        if d["サービスコード"] == "210":
            chui = "⚠ 空床利用型。併設施設と一体の職員体制。合算しない"
        rows.append([d["サービスの種類"], d["事業所名"], d["実人数の合計"],
                     d["常勤換算の合計"], d["採用者数"], d["退職者数"],
                     d["職員1人当たりの入所者数"], d["記入日"], chui])
    put(ws, rows, lambda r: WARN if str(r[8]).startswith("⚠") else None)

    # ══════ 03　職種別の従業者
    ws = sheet(wb, "03_職種別の従業者",
               ["サービスの種類", "事業所名", "職種", "専従・常勤",
                "兼務・常勤", "専従・非常勤", "兼務・非常勤", "合計",
                "常勤換算"],
               [26, 24, 18, 11, 11, 12, 12, 8, 10])
    rows = []
    for d in ordered:
        for na, a, b, c, e, g, h in d["職種別"]:
            rows.append([d["サービスの種類"], d["事業所名"], na,
                         a, b, c, e, g, h])
    put(ws, rows)
    n_shoku = len(rows)

    # ══════ 04　指定の年月日
    ws = sheet(wb, "04_指定の年月日",
               ["サービスコード", "サービスの種類", "事業所名",
                "事業所番号", "指定権者", "指定の年月日",
                "指定の更新年月日（直近）", "事業所の形態", "公表日"],
               [12, 26, 24, 13, 8, 14, 20, 16, 12])
    put(ws, [[d["サービスコード"], d["サービスの種類"], d["事業所名"],
              d["事業所番号"], d["指定権者"], d["指定の年月日"],
              d["指定の更新年月日"] or "―", d["事業所の形態"] or "―",
              d["公表日"]] for d in ordered])

    # ══════ 05　退所先
    ws = sheet(wb, "05_退所先",
               ["サービスの種類", "事業所名", "期間", "自宅等",
                "介護保険施設", "特別養護老人ホーム以外の社会福祉施設",
                "医療機関", "死亡者", "その他", "計", "⚠ 注意"],
               [26, 24, 26, 9, 12, 26, 10, 9, 9, 8, 46])
    rows = []
    for d in ordered:
        t = d["退所先"]
        if not t:
            if d["サービスコード"] in ("540", "210"):
                rows.append([d["サービスの種類"], d["事業所名"],
                             "⚠ 項目なし", "―", "―", "―", "―", "―", "―",
                             "―",
                             "⚠ **公表の様式に退所先の項目がない**"
                             "（未実施。「なし」ではない）"])
            continue
        kei = sum(t.values())
        rows.append([d["サービスの種類"], d["事業所名"],
                     d["退所先の期間"] or "前年度",
                     t.get("自宅等", 0), t.get("介護保険施設", 0),
                     t.get("特別養護老人ホーム以外の社会福祉施設", 0),
                     t.get("医療機関", 0), t.get("死亡者", 0),
                     t.get("その他", 0), kei,
                     "⚠ **退所先は載るが「入所前の居所」は公表の項目にない**"])
    put(ws, rows, lambda r: WARN)
    n_taisho = len([r for r in rows if r[2] != "⚠ 項目なし"])

    # ══════ 06　残るもの
    ws = sheet(wb, "06_残るもの",
               ["件", "残るもの", "なぜ公表システムで埋まらないか",
                "ヒアリングを行わない場合の扱い", "確認事項"],
               [6, 30, 40, 54, 14])
    nokoru = [
        [1, "稼働率（延べ利用日数÷延べ定員日数）",
         "公表の様式に項目がない。載るのは記入日前月の在所者数であり、"
         "月の途中の空床は表れない",
         "入所率（在所者数÷定員）で代える。⚠ 稼働率とは別のものである"
         "ことを計画書に明記する", "No.159"],
        [2, "入所前の居所（在宅か医療機関か）",
         "公表の様式に項目がない（退所先は載る）",
         "⚠ **埋まらない。** 退所先（前年度）で代える。"
         "施設整備の必要量の判断には、退所先と待機者数を用いる",
         "No.159"],
        [3, "休止床の数",
         "公表の様式に項目がない",
         "入所率が100％未満である施設（2件）について、"
         "定員と在所者数の差として示す。休止か空床かは区別しない",
         "No.159"],
        [4, "地域密着型介護老人福祉施設の待機者数",
         "⚠ 公表の様式に待機者数の項目がない（特別養護老人ホーム・"
         "介護老人保健施設には項目がある）",
         "併設する特別養護老人ホームの待機者数（20人）で全体を示す。"
         "町が指定権者であるため、町の指定台帳で補える", "No.29"],
        [5, "指定を受けていない有料老人ホーム・"
            "サービス付き高齢者向け住宅",
         "⚠ 介護保険の指定事業所の仕組みであるため載らない",
         "宮城県の届出状況の公表一覧による。"
         "サ高住情報提供システムは通信が許可されていない", "No.154"],
        [6, "町外の事業所（町民が利用しているもの）の受入余地",
         "町外の事業所を並べると他団体の固有名称が成果品に入る",
         "住所地特例者台帳（24人）と国保連データによる。"
         "⚠ 公表システムでは町外の事業所も採録できるが、"
         "計画書には載せない", "No.12"],
        [7, "経営の状況（収支・人件費率）",
         "公表の様式に項目がない",
         "⚠ **埋まらない。** 計画に書かない。"
         "介護事業経営実態調査（全国）の水準を参照する扱いとする",
         "No.51"],
        [8, "必要な支援・町への要望",
         "公表の様式に項目がない（事業所の自由記述はあるが要望ではない）",
         "⚠ **埋まらない。** 地域ケア会議・地域密着型サービス運営"
         "協議会という既存の場で聴く。第10期の計画期間中の仕組みとして"
         "位置付ける", "No.50"],
    ]
    put(ws, nokoru, lambda r: WARN if "埋まらない" in str(r[3]) else OK)

    # ══════ 07　出所と手順
    ws = sheet(wb, "07_出所と手順", ["順", "手順", "スクリプト・ファイル"],
               [6, 72, 54])
    put(ws, [
        [1, "厚生労働省のオープンデータ（サービス種類別35本のＣＳＶ）から"
            "市町村コード043249を抽出し、町内の14の指定を確定する",
         "07_ソーススクリプト/fetch_kouhyou_opendata_R8.10.20.py"],
        [2, "①の抽出結果（事業所番号×サービスコード）から詳細ページの"
            "場所を起こし、14ページを取得して原典として置く",
         "07_ソーススクリプト/fetch_kouhyou_detail_R8.10.20.py"],
        [3, "詳細ページの <th> の abbr 属性を項目の名として読み、"
            "項目と値のＣＳＶに起こす（⚠ 文字だけで探さない）",
         "09_元資料/その他データ/公表システム_詳細_川崎町_R8.10.20.csv"],
        [4, "ＣＳＶから定員・在所者数・入所率・待機者数・職種別従業者数・"
            "指定の年月日を取り出す（**数値の正本**）",
         "07_ソーススクリプト/data_jigyosho_jittai_R8.10.20.py"],
        [5, "本表を作る", "07_ソーススクリプト/"
            "build_jigyosho_jittai_R8.10.20.py"],
        [6, "計画素案に反映する（Ver.2.21）",
         "07_ソーススクリプト/fix_soan_v2210_hearing_nashi.py"],
    ])

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)

    # ══════ 自己点検
    print("町内事業所の実態（公表システム）の管理シート")
    print(f"  {OUT}")
    print()
    print("  ── 自己点検")
    judge(len(wb.sheetnames) == 8, f"8シート（実際 {len(wb.sheetnames)}）")
    judge(n_nyu >= 6, f"入所・入居の状況が6件以上（実際 {n_nyu}）")
    judge(n_shoku >= 100, f"職種別の行が100行以上（実際 {n_shoku}）")
    judge(n_taisho == 4, f"退所先が取れたのは4件（実際 {n_taisho}）")
    judge(len(nokoru) == 8, f"残るものを8件示している（実際 {len(nokoru)}）")
    # 読み返して値が入っているか
    wb2 = openpyxl.load_workbook(OUT, data_only=True)
    ws2 = wb2["01_入所・入居の状況"]
    teiin = [ws2.cell(r, 4).value for r in range(2, ws2.max_row + 1)]
    judge(sum(x for x in teiin if isinstance(x, (int, float))) >= 173,
          f"定員の合計が173人以上（実際 "
          f"{sum(x for x in teiin if isinstance(x, (int, float))):.0f}人）")
    gh = [ws2.cell(r, 4).value for r in range(2, ws2.max_row + 1)
          if "認知症対応型共同生活介護" == ws2.cell(r, 1).value]
    judge(sum(gh) == 27, f"認知症対応型共同生活介護の定員計が27人"
                         f"（実際 {sum(gh):.0f}人）")
    ws3 = wb2["06_残るもの"]
    judge(any("埋まらない" in str(ws3.cell(r, 4).value)
              for r in range(2, ws3.max_row + 1)),
          "⚠ 「埋まらない」ものを「適合」に丸めていない")
    # ⚠ 他団体・法人名の見張り
    zen = []
    for sh in wb2:
        for row in sh.iter_rows(values_only=True):
            zen += [str(v) for v in row if v is not None]
    s = "\n".join(zen)
    judge("川崎市" not in s and "福岡" not in s,
          "⚠ 川崎市（神奈川県）・福岡県川崎町が混じらない")

    print()
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        for m in futeki:
            print(f"      {m}")
        sys.exit(1)
    print("   ○ 自己点検 不適合0件")


if __name__ == "__main__":
    main()
