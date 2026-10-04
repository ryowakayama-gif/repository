# -*- coding: utf-8 -*-
"""事業の内容の出所一覧（令和8年10月4日）

素案 Ver.2.9 で事業一覧表に加えた「内容」の欄について、
**どの記述がどこから来たのか**を1件ずつ示す。

  ① 第9期計画書（令和6年3月）の本文から機械で取り出したもの
  ② 統合体系表の03シート（新規・拡充の第10期での内容）
  ③ 当方が書き起こしたもの　←**町のご確認をお願いしたいもの**

③は、介護保険法・国の指針が内容を定めているもの（施設サービス等）と、
素案の本文に記述があるもの（町独自の事業）に限っているが、
**実際の運用と食い違っていないかは町でなければ確かめられない。**

  出力　05_試算・管理シート/川崎町_事業の内容の出所_R8.10.4.xlsx

本表は、業務工程管理表の確認事項No.164 に対応する。
"""
import re
import sys

import docx
import docx.table
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, "07_ソーススクリプト")
import fix_soan_v290_naiyo as V29              # noqa: E402

SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.9_事業内容版.docx"
OUT = "05_試算・管理シート/川崎町_事業の内容の出所_R8.10.4.xlsx"

TITLE = PatternFill("solid", fgColor="1F4E78")
HEAD = PatternFill("solid", fgColor="5B9BD5")
LEAD = PatternFill("solid", fgColor="F2F2F2")
F1 = PatternFill("solid", fgColor="E2EFDA")   # 第9期計画書
F2 = PatternFill("solid", fgColor="DAE3F3")   # 統合体系表
F3 = PatternFill("solid", fgColor="FCE4D6")   # 当方が書き起こした


def main():
    auto = V29.from_9ki()
    shin = V29.from_shinki()
    man = V29.NAIYO

    def norm(s):
        return re.sub(r"\s", "", str(s or ""))

    man_n = {norm(k): k for k in man}
    shin_n = {norm(k): k for k in shin}
    auto_n = {norm(k): k for k in auto}

    from docx.oxml.ns import qn
    doc = docx.Document(SOAN)
    rows = []
    sec = None
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            t = "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()
            m = re.match(r"^(\d+-\d+)[　 ]", t)
            if m:
                sec = m.group(1)
        elif el.tag == qn("w:tbl"):
            tb = docx.table.Table(el, doc)
            if [c.text.strip() for c in tb.rows[0].cells][:2] != ["項", "事業"]:
                continue
            for r in tb.rows[1:]:
                name = r.cells[1].text.strip()
                if not name:
                    continue
                k = norm(name)
                if k in man_n:
                    src = "③ 当方が書き起こした"
                    note = ("介護保険法・国の指針が内容を定めているもの、"
                            "又は素案の本文に記述があるもの。"
                            "**実際の運用と食い違っていないかご確認ください。**")
                elif k in shin_n:
                    src = "② 統合体系表（新規・拡充）"
                    note = "統合体系表 03シートの「第10期での内容」による。"
                elif k in auto_n:
                    src = "① 第9期計画書"
                    note = "第9期計画書（令和6年3月）の本文の最初の一文。"
                else:
                    src = "―"
                    note = "出所が辿れない。"
                rows.append((sec, r.cells[0].text.strip()[:30], name,
                             r.cells[2].text.strip(), r.cells[3].text.strip(),
                             src, r.cells[4].text.strip(), note))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01_内容の出所"
    ws.cell(1, 1).value = "事業の内容の出所（素案 Ver.2.9 の事業一覧表）"
    ws.cell(1, 1).font = Font(name="游ゴシック", size=14, bold=True,
                              color="FFFFFF")
    ws.cell(1, 1).fill = TITLE
    ws.cell(2, 1).value = (
        "素案 Ver.2.9 で事業一覧表に加えた「内容」の欄について、"
        "どの記述がどこから来たのかを1件ずつ示します。"
        "**橙色の③は当方が書き起こしたものです。**"
        "介護保険法・国の指針が内容を定めているもの（施設サービス等）と、"
        "素案の本文に記述があるもの（町独自の事業）に限っていますが、"
        "実際の運用と食い違っていないかは町でなければ確かめられません。"
        "③についてご確認をお願いします（確認事項No.164）。")
    ws.cell(2, 1).font = Font(name="游ゴシック", size=9)
    ws.cell(2, 1).fill = LEAD
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
    ws.row_dimensions[2].height = 62
    head = ["節", "項", "事業", "区分", "所管", "内容の出所", "内容",
            "ご確認のお願い"]
    for c, (h, w) in enumerate(zip(head, [6, 24, 40, 6, 16, 22, 70, 46]),
                               start=1):
        cell = ws.cell(4, c)
        cell.value = h
        cell.font = Font(name="游ゴシック", size=9, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    FILL = {"① 第9期計画書": F1, "② 統合体系表（新規・拡充）": F2,
            "③ 当方が書き起こした": F3}
    for i, rec in enumerate(rows):
        for c, v in enumerate(rec, start=1):
            cell = ws.cell(5 + i, c)
            cell.value = v
            cell.font = Font(name="游ゴシック", size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = FILL.get(rec[5], PatternFill())
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:H{4 + len(rows)}"

    ws2 = wb.create_sheet("02_まとめ")
    import collections
    cnt = collections.Counter(r[5] for r in rows)
    ws2.cell(1, 1).value = "まとめ"
    ws2.cell(1, 1).font = Font(name="游ゴシック", size=14, bold=True,
                               color="FFFFFF")
    ws2.cell(1, 1).fill = TITLE
    r = 3
    for k in ("① 第9期計画書", "② 統合体系表（新規・拡充）",
              "③ 当方が書き起こした", "―"):
        if cnt.get(k):
            ws2.cell(r, 1).value = k
            ws2.cell(r, 2).value = cnt[k]
            ws2.cell(r, 1).fill = FILL.get(k, PatternFill())
            r += 1
    ws2.cell(r, 1).value = "計"
    ws2.cell(r, 2).value = len(rows)
    ws2.cell(r, 1).font = Font(name="游ゴシック", size=9, bold=True)
    for c, w in zip("AB", (34, 10)):
        ws2.column_dimensions[c].width = w
    wb.save(OUT)

    print("事業の内容の出所")
    print("  保存：", OUT)
    for k, v in cnt.most_common():
        print(f"   {k}　{v}件")
    print(f"   計 {len(rows)}件")
    print("  ── 自己点検")
    if cnt.get("―"):
        print(f"   × 出所が辿れない行が{cnt['―']}件")
        sys.exit(1)
    print("   ○ すべての行の出所が辿れる")
    print(f"  ⚠ ③の{cnt.get('③ 当方が書き起こした', 0)}件は"
          "町のご確認をお願いします（確認事項No.164）。")


if __name__ == "__main__":
    main()
