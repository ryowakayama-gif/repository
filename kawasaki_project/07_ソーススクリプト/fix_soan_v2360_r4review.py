# -*- coding: utf-8 -*-
"""協議用素案 Ver.2.35 → Ver.2.36（Red Team 再レビュー第4回（令和8年10月24日）の指摘の是正）

  python3 07_ソーススクリプト/fix_soan_v2360_r4review.py

再レビュー原本　09_元資料/レビュー/川崎町_第10期計画_RedTeam再レビュー第4回_R8.10.24受領.md

  F4-01 9-2「第10期給付費の見込み」の令和7年度の欄で、内訳5項目の和（1,084,826千円）と標準給付費見込額
        （1,084,587千円）が239千円違う。原因を年報で確かめた：
          1,084,587千円 は年報 様式3「保険給付支払状況」の計（支払額。審査支払手数料を含まない）
          内訳の総給付費 998,222千円 は年報 様式2 の給付費（給付実績）
          → 差 239千円 ＝ 審査支払手数料 656千円 －（諸費の支払額 998,639千円 － 総給付費 998,222千円 ＝ 417千円）
        2つの集計を1つの列に混ぜていた。令和7年度の欄は令和9〜11年度と同じ定義（5項目の和）に改め、
        様式3 の支払額との違いを出典の注に書く。令和9年度の1千円差は円単位の算定値を合計してから丸めたため。
  F4-01 の波及　「令和6年度1,093,253千円 → 令和7年度1,084,587千円（0.8％減）」も基準が違っていた。
        1,093,253千円 は令和6年度の様式3 の計 1,092,620,753円 に審査支払手数料 632,700円 を加えた額
        （calc_hokenryo_R8.9.8b.py）、1,084,587千円 は手数料を含まない。同じ様式3 の支払額で比べると
        1,092,621千円 → 1,084,587千円 で 0.7％減。7-2・9-3 の本文と 3-4 の表の「0.8％減」を改める。

終了コード　0＝検査すべて適合／1＝1件でも不適合
"""
import re
import shutil
import sys

import docx
import openpyxl

MOTO = "01_第10期_最新版成果品/川崎町_計画書素案_v2.35_協議用.docx"
SAKI = "01_第10期_最新版成果品/川崎町_計画書素案_v2.36_協議用.docx"
NENPO7 = "09_元資料/R7実績データ/年報データ_2025_川崎町.xlsx"
NENPO6 = "09_元資料/R8実績データ/R8.9.8受領版/年報データ_2024_川崎町.xlsx"


def kakikae(p, s):
    p.runs[0].text = s
    for r in p.runs[1:]:
        r.text = ""


def zenbun(d):
    return "\n".join([p.text for p in d.paragraphs]
                     + [c.text for t in d.tables for r in t.rows for c in r.cells])


def subete(d):
    seen = set()
    for p in list(d.paragraphs) + [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]:
        if id(p._p) in seen:
            continue
        seen.add(id(p._p))
        yield p


def yoshiki3(path):
    """年報 様式3「４．保険給付支払状況」の支払義務額累計（円）。"""
    ws = openpyxl.load_workbook(path, data_only=True)["様式３"]
    out = {}
    on = False
    for r in ws.iter_rows(values_only=True):
        v = [x for x in r if x is not None]
        if v and "保険給付支払状況" in str(v[0]):
            on = True
        elif on and v and isinstance(v[-1], (int, float)) and len(v) >= 2:
            out[str(v[0]).strip()] = v[1]
    return out


def sen(x):
    return int((x + 500) // 1000)


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("協議用素案 Ver.2.35 → Ver.2.36（Red Team 再レビュー第4回の指摘の是正）")
    y7, y6 = yoshiki3(NENPO7), yoshiki3(NENPO6)
    shiharai7, shiharai6 = y7["計"], y6["計"]
    shohi7 = y7["介護サービス等諸費"] + y7["介護予防サービス等諸費"]
    judge(sen(shiharai7) == 1084587, f"令和7年度 様式3 の計 {shiharai7:,}円 → {sen(shiharai7):,}千円")
    judge(sen(shiharai6) + sen(632700) == 1093253 or round((shiharai6 + 632700) / 1000) == 1093253,
          f"令和6年度 様式3 の計 {shiharai6:,}円 ＋ 審査支払手数料 632,700円 ＝ {shiharai6 + 632700:,}円（1,093,253千円の出所）")
    ritsu = (sen(shiharai7) - sen(shiharai6)) / sen(shiharai6) * 100
    judge(f"{ritsu:.1f}" == "-0.7", f"同じ様式3 の支払額で {sen(shiharai6):,}千円 → {sen(shiharai7):,}千円（{ritsu:.2f}％）")

    shutil.copy2(MOTO, SAKI)
    d = docx.Document(SAKI)
    mae = zenbun(d)

    def rep(kyu, shin, n_min=1, mi=""):
        n = 0
        for p in subete(d):
            if kyu in p.text and p.runs:
                kakikae(p, p.text.replace(kyu, shin))
                n += 1
        judge(n >= n_min, f"{mi}「{kyu[:26]}」（{n}か所）")
        return n

    print("  ── F4-01 9-2 の表の令和7年度の標準給付費見込額")
    t = [t for t in d.tables if t.rows[0].cells[0].text.strip() == "区分"
         and "令和7年度" in t.rows[0].cells[1].text and "令和9年度" in t.rows[0].cells[2].text][0]
    row = {r.cells[0].text.strip(): r for r in t.rows}
    num = lambda s: int(re.sub(r"[^\d]", "", s))
    uchi = ["総給付費", "特定入所者介護サービス費等", "高額介護サービス費等", "高額医療合算介護サービス費等", "審査支払手数料"]
    wa7 = sum(num(row[k].cells[1].text) for k in uchi)
    kyu7 = num(row["標準給付費見込額"].cells[1].text)
    judge(kyu7 == 1084587 and wa7 - kyu7 == 239, f"令和7年度 内訳の和 {wa7:,}千円・標準給付費見込額 {kyu7:,}千円（差 {wa7 - kyu7}千円）")
    judge(sen(shohi7) - num(row["総給付費"].cells[1].text) == 417 and num(row["審査支払手数料"].cells[1].text) == 656,
          f"差の内訳：審査支払手数料 656千円 －（諸費の支払額 {sen(shohi7):,}千円 － 総給付費 998,222千円 ＝ 417千円）＝ 239千円")
    c = row["標準給付費見込額"].cells[1]
    kakikae(c.paragraphs[0], f"{wa7:,}千円")
    for y, ci in (("令和9年度", 2), ("令和10年度", 3), ("令和11年度", 4)):
        w = sum(num(row[k].cells[ci].text) for k in uchi)
        print(f"     {y}：内訳の和 {w:,}千円・標準給付費見込額 {num(row['標準給付費見込額'].cells[ci].text):,}千円")
    rep("出典：令和7年度実績：介護保険事業状況報告（年報・令和7年度）／令和9〜11年度：本計画の算定値",
        "出典：令和7年度実績：介護保険事業状況報告（年報・令和7年度）／令和9〜11年度：本計画の算定値。"
        f"令和7年度の標準給付費見込額は、上の5項目の和です（総給付費は様式2の給付費。審査支払手数料は審査件数10,929件"
        f"×60円による）。2-4の「保険給付の支払状況」（様式3）の計{sen(shiharai7):,}千円とは、①審査支払手数料を含まないこと、"
        f"②介護・介護予防サービス等諸費の支払額（{sen(shohi7):,}千円）が様式2の給付費（998,222千円）を417千円上回ること"
        f"（支払額と給付実績という集計の違い）により、239千円異なります。令和9〜11年度は、円単位の算定値を合計してから"
        f"千円未満を四捨五入しているため、内訳の和と1千円異なる年度があります。", mi="9-2 出典")

    print("  ── F4-01 の波及　令和6年度→令和7年度の比較の基準")
    a6, a7 = f"{sen(shiharai6):,}", f"{sen(shiharai7):,}"
    rep("標準給付費は令和6年度の1,093,253千円から令和7年度の1,084,587千円へ0.8％減少しています",
        f"保険給付費（介護保険事業状況報告 様式3 の支払額。審査支払手数料を含まない）は令和6年度の{a6}千円から"
        f"令和7年度の{a7}千円へ0.7％減少しています", mi="7-2")
    rep("保険給付費は令和6年度1,093,253千円から令和7年度1,084,587千円へ0.8％減少しており",
        f"保険給付費（様式3 の支払額）は令和6年度{a6}千円から令和7年度{a7}千円へ0.7％減少しており", mi="9-3")
    rep("給付費の平坦化（令和7年度0.8％減）", "給付費の平坦化（令和7年度0.7％減）", mi="3-4 表")

    d.save(SAKI)
    ato = zenbun(docx.Document(SAKI))
    print("  ── 検査1　古い記述")
    for go in ("1,093,253", "0.8％減", "標準給付費は令和6年度"):
        judge(go not in ato, f"「{go}」0件")
    judge(ato.count("1,084,587千円") == mae.count("1,084,587千円") - 1 + 1,
          f"「1,084,587千円」{mae.count('1,084,587千円')}→{ato.count('1,084,587千円')}（9-2 の表から外し、注と比較の文に残す）")
    print("  ── 検査2　保険料の算定（令和9〜11年度）は変わらない")
    for w in ("3,313,925千円", "765,752千円", "6,822円", "6,143円", "5,464円"):
        judge(ato.count(w) == mae.count(w), f"「{w}」{mae.count(w)}→{ato.count(w)}")
    print("  ── 検査3　未確定の印・他団体名・受託者の語は増減なし")
    for w in ("【町確認】", "【委員会協議】", "［要確認］", "［要協議］", "［要内訳］", "大雪", "北塩原", "金ヶ崎",
              "川崎市", "受託者", "当社", "ビズアップ"):
        judge(ato.count(w) == mae.count(w), f"「{w}」{mae.count(w)}→{ato.count(w)}")
    print(f"  {SAKI}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 不適合0件")


if __name__ == "__main__":
    main()
