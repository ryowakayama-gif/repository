# -*- coding: utf-8 -*-
"""計画素案 Ver.2.23 → Ver.2.24（交付金の記述の是正）
（令和8年10月22日）

  python3 07_ソーススクリプト/fix_soan_v2240_kofukin.py

なぜ要るか
  交付金の推移を枝番の単位・得点率で数え直した
  （analyze_kofukin_suii_R8.10.22.py → 川崎町_交付金_推移の再整理_R8.10.22.xlsx）。
  その結果と素案を突き合わせ、次の3つの誤りを見つけた。

  ⚠ 3-3（p265）　「37点の減少は、評価結果の活用が16点から0点となったこと、
    及びアウトカム指標が15点下がったことによります」
    → 16＋15＝31 で 37 に足りない。しかもアウトカム指標は推進・支援の
      **両方**で15点ずつ（計30点）下がっている。全体の内訳は
        成果指標群の変動 ▲30／配点の改定 ▲7／指標の新設 ＋2／
        取組の後退 ▲34／取組の前進 ＋32　＝ ▲37
  ⚠ 3-3（p267）　「県平均との差は146.1ポイントから30.5ポイントへ」
    → 得点の差であるから「点」。
  ⚠ 3-3（p301）・10章（p679）　「第9章9-1の保険料算定における交付金の
    取扱いも、令和8年度の交付実績を据え置いています」
    → **事実と違う。** 保険料算定表（表112）は J＝C＋D－E＋F＋G±H－I で、
      交付金の交付見込額を差し引いていない。保険料の算定は9-3である。
      一方、町の将来推計総括表（令和8年10月19日出力）は
      「保険者機能強化推進交付金等の交付見込額」に3か年4,500千円を置き、
      保険料収納必要額から差し引いている。
    → 差し引いていないことを正しく書き、交付見込額の扱いを
      ［要確認］として残す。**保険料の額そのものは動かさない。**
      効き（月額約40円）は9-4の幅の表に1行加えて示す。

  交付額は令和6年度の原典（交付見込額確定版）にのみある。
    推進1,277千円（うち激変緩和措置443千円）・支援2,461千円（同788千円）
    計3,738千円（うち激変緩和措置1,231千円）
  令和7・8年度の原典には交付額の列がない（確認事項No.39）。

終了コード　0＝検査1〜7すべて適合／1＝1件でも不適合
"""
import copy
import os
import re
import shutil
import sys
from collections import defaultdict
from decimal import ROUND_HALF_UP, Decimal

import docx
import openpyxl

MOTO = "01_第10期_最新版成果品/川崎町_計画書素案_v2.23_県調査反映版.docx"
SAKI = "01_第10期_最新版成果品/川崎町_計画書素案_v2.24_交付金整理版.docx"
BAN_MOTO = "Ver.2.23"
BAN = "Ver.2.24"
SAISEIRI = "05_試算・管理シート/川崎町_交付金_推移の再整理_R8.10.22.xlsx"

# 保険料の算定の前提（表112 と同じ）
J = Decimal(765_752)          # 保険料収納必要額（千円）
RATE = Decimal("0.96")
N = Decimal("9743.1")
KOFU = Decimal(4_500)         # 将来推計総括表の交付見込額（3か年・千円）


def yen(x):
    return int(x.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def geta(j):
    return Decimal(j) * 1000 / RATE / N / 12


KIJUN = yen(geta(J))
KOFU_GO = yen(geta(J - KOFU))
KOFU_SA = KOFU_GO - KIJUN

KOFU_BUN = (
    "本計画では、得点の回復を「増額の可能性がある」ものとして扱い、"
    "金額の見込みは立てていません。なお、令和6年度の交付見込額は、"
    "保険者機能強化推進交付金1,277千円・介護保険保険者努力支援交付金"
    "2,461千円の計3,738千円（うち激変緩和措置分1,231千円）です。"
    "令和7・8年度の交付額は把握していません［要確認］。"
    "第9章9-3の保険料算定（表112）では、交付金の交付見込額を保険料"
    "収納必要額から差し引いていません。見込むか否か及び見込む額は、"
    "第3回策定委員会までに定めます［要確認］（確認事項No.39）。")
KOFU_KYU = ("本計画では、得点の回復を「増額の可能性がある」ものとして扱い、"
            "金額の見込みは立てていません。第9章9-1の保険料算定における"
            "交付金の取扱いも、令和8年度の交付実績を据え置いています"
            "（確認事項No.39）。")


def uchiwake():
    """再整理の表から令和7→8年度の内訳を数え直す（数値の正本）。"""
    wb = openpyxl.load_workbook(SAISEIRI, data_only=True)
    ws = wb["02_低下の内訳_R7→R8"]
    kubun = defaultdict(int)
    komoku = defaultdict(int)
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None or not isinstance(r[10], (int, float)):
            continue
        kubun[r[0]] += r[10]
        komoku[(r[0], r[4])] += r[10]
    ws1 = wb["01_合計の推移"]
    goukei = {r[0]: r for r in ws1.iter_rows(min_row=2, values_only=True)}
    return kubun, komoku, goukei


def dan_kaku(p, s):
    if p.runs:
        p.runs[0].text = s
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(s)


def kaku(cell, s):
    for i, p in enumerate(cell.paragraphs):
        if i == 0:
            dan_kaku(p, s)
        else:
            for r in p.runs:
                r.text = ""


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    print(f"計画素案 {BAN_MOTO} → {BAN}（交付金の記述の是正）")
    print()
    for f in (MOTO, SAISEIRI):
        if not os.path.exists(f):
            print(f"   × 入力がありません：{f}")
            sys.exit(1)

    kubun, komoku, goukei = uchiwake()
    seika = kubun["成果指標群の変動"]
    haiten = kubun["配点の改定（得点率は不変）"]
    shinsetsu = kubun["指標の新設"]
    kotai = kubun["取組の後退（得点率の低下）"]
    zenshin = kubun["取組の前進（得点率の上昇）"]
    tanki = sum(v for (k, na), v in komoku.items()
                if k == "成果指標群の変動" and na.startswith("短期"))
    choki = sum(v for (k, na), v in komoku.items()
                if k == "成果指標群の変動" and na.startswith("長期"))
    kenko = sum(v for (k, na), v in komoku.items()
                if k == "成果指標群の変動" and na.startswith("健康寿命"))
    katsuyo = komoku[("取組の後退（得点率の低下）", "評価結果の活用")]
    tekisei = komoku[("取組の後退（得点率の低下）", "給付費適正化事業の取組状況")]
    zaitaku = komoku[("取組の前進（得点率の上昇）",
                      "在宅医療・介護連携に関する課題・対応策の検討")]
    hokatsu = komoku[("取組の前進（得点率の上昇）",
                      "地域包括支援センター事業評価の達成状況")]
    r7, r8 = goukei["令和7年度"], goukei["令和8年度"]
    gen = r8[1] - r7[1]

    print("  ── ⓪ 再整理の表から内訳を読む")
    judge(seika + haiten + shinsetsu + kotai + zenshin == gen == -37,
          f"内訳の合計が合計の増減に一致（{seika}{haiten:+}{shinsetsu:+}"
          f"{kotai:+}{zenshin:+}＝{gen}）")
    judge(tanki + choki + kenko == seika,
          f"成果指標群の内訳（短期{tanki}・長期{choki:+}・健康寿命{kenko}）")

    DAN = [
        (265, "令和7年度から令和8年度への37点の減少は、",
         f"令和7年度から令和8年度への{-gen}点の減少を、評価指標の枝番の"
         f"単位まで下ろし、各年度の配点に対する得点率で比べると、次のとおり"
         f"です。①アウトカム指標（目標Ⅳ）が推進・支援それぞれ{-seika // 2}点、"
         f"計{-seika}点下がりました。短期的な要介護度の変化で計{-tanki}点、"
         f"健康寿命延伸の状況で計{-kenko}点を失い、長期的な要介護度の変化で"
         f"計{choki}点を得たことによります。アウトカム指標は推進・支援で同一の"
         f"指標であり、同じ得点が両方に計上されます。②国が配点を改めたことに"
         f"より{-haiten}点下がりました（認知症サポーター等を活用した地域支援"
         f"体制の構築が25点満点から20点満点に改められる等。いずれも満点を"
         f"維持しています）。③新設の指標で{shinsetsu}点を得ました。④取組の"
         f"評価が下がった項目で{-kotai}点を失い、⑤上がった項目で{zenshin}点を"
         f"得ました。④のうち「評価結果の活用」{-katsuyo}点と「給付費適正化"
         f"事業の取組状況」のア{-tekisei}点の計{-katsuyo - tekisei}点は、"
         f"令和7年度に満点であったもので、仕組みを定めることで取り戻せます"
         f"（本節の後段）。⑤では在宅医療・介護連携に関する課題・対応策の"
         f"検討（{zaitaku}点）、地域包括支援センター事業評価の達成状況"
         f"（{hokatsu}点）等が上がっています。取組の体制は全国上位である一方、"
         f"その結果を計画に反映し公表する仕組みと、要介護度の変化という成果に"
         f"結びつける仕組みが弱いことを示しており、第10期ではこの2点を計画の"
         f"推進体制に組み込みます。"),
    ]

    shutil.copy2(MOTO, SAKI)
    d = docx.Document(SAKI)
    mae = "\n".join([p.text for p in d.paragraphs]
                    + [c.text for t in d.tables for r in t.rows
                       for c in r.cells])

    print("  ── ① 段落の差し替え（3-3 の37点の内訳）")
    n_dan = 0
    for i, atama, atarashii in DAN:
        p = d.paragraphs[i]
        if not p.text.startswith(atama):
            judge(False, f"p{i} の冒頭が合わない", f"実際『{p.text[:26]}』")
            continue
        dan_kaku(p, atarashii)
        n_dan += 1
    judge(n_dan == len(DAN), f"{len(DAN)}段落を差し替えた（実際 {n_dan}）")

    print("  ── ② 段落の中の語の置き換え")
    # (段落の番号, 旧, 新)
    BUBUN = [
        (267, "県平均との差は146.1ポイントから30.5ポイントへ縮まりました。",
         "県平均との差は146.1点から30.5点へ縮まりました。"),
        (301, KOFU_KYU, KOFU_BUN),
        (679, KOFU_KYU, KOFU_BUN),
        (618, "④第10期の第1号被保険者負担割合が24％と定められた場合は、"
              "月額約309円上がります。",
         "④第10期の第1号被保険者負担割合が24％と定められた場合は、"
         "月額約309円上がります。⑤保険者機能強化推進交付金等の交付見込額は、"
         "本算定では見込んでいません。地域包括ケア「見える化」システムの"
         f"将来推計総括表（令和8年10月19日出力）は3か年で{KOFU:,}千円を"
         f"見込んでおり、これを差し引いた場合は月額約{-KOFU_SA}円下がります"
         "［要確認］。"),
    ]
    n_bu = 0
    for i, kyu, atarashii in BUBUN:
        p = d.paragraphs[i]
        if kyu not in p.text:
            judge(False, f"p{i} に旧の文がない", f"実際『{p.text[:26]}』")
            continue
        dan_kaku(p, p.text.replace(kyu, atarashii))
        n_bu += 1
    judge(n_bu == len(BUBUN), f"{len(BUBUN)}か所を置き換えた（実際 {n_bu}）")

    print("  ── ③ 前提を動かした場合の幅（表115）に交付金の行を加える")
    t = d.tables[115]
    mi = [c.text for c in t.rows[0].cells]
    if mi != ["前提", "本算定の置き方", "動かした場合", "月額基準額", "差"]:
        judge(False, "表115 の見出しが合わない", "／".join(mi))
    else:
        moto_tr = t.rows[-1]._tr
        tr = copy.deepcopy(moto_tr)
        moto_tr.addnext(tr)
        row = t.rows[-1]
        atai = ["保険者機能強化推進交付金等の交付見込額", "見込まない（0円）",
                f"3か年計{KOFU:,}千円（将来推計総括表の値）",
                f"{KOFU_GO:,}円", f"▲{-KOFU_SA}円"]
        for c, s in zip(row.cells, atai):
            kaku(c, s)
        judge(True, f"表115 に1行加えた（{KOFU_GO:,}円・▲{-KOFU_SA}円）")

    print("  ── ④ 版表記")
    n_ban = 0
    parts = [p for sec in d.sections for part in (sec.header, sec.footer)
             for p in part.paragraphs] + list(d.paragraphs)
    for p in parts:
        for r in p.runs:
            if BAN_MOTO in r.text:
                r.text = r.text.replace(BAN_MOTO, BAN)
                n_ban += 1
    judge(n_ban >= 1, f"版表記を {BAN} に改めた（{n_ban}か所）")

    d.save(SAKI)

    d2 = docx.Document(SAKI)
    ato = "\n".join([p.text for p in d2.paragraphs]
                    + [c.text for t in d2.tables for r in t.rows
                       for c in r.cells])

    print("  ── 検査1　⚠ 事実と違う記述が残っていないか")
    judge("令和8年度の交付実績を据え置いています" not in ato,
          "「令和8年度の交付実績を据え置いています」が0件")
    judge("9-1の保険料算定" not in ato, "「9-1の保険料算定」が0件")
    judge(ato.count("差し引いていません") == 2,
          f"3-3・10章の2か所に差し引いていない旨がある"
          f"（{ato.count('差し引いていません')}件）")

    print("  ── 検査2　37点の内訳")
    judge("及びアウトカム指標が15点下がったことによります" not in ato,
          "誤った内訳（16点と15点）が0件")
    for go in (f"計{-seika}点下がりました", f"{-haiten}点下がりました",
               f"新設の指標で{shinsetsu}点", f"{-kotai}点を失い",
               f"{zenshin}点を得ました"):
        judge(go in ato, f"「{go}」がある")

    print("  ── 検査3　点の単位")
    judge("146.1ポイント" not in ato and "30.5ポイント" not in ato,
          "県平均との差を「ポイント」で書いた箇所が0件")
    judge(round(r7[1] - r7[5], 1) == 146.1 and round(r8[1] - r8[5], 1) == 30.5,
          f"県平均との差 令和7年度{r7[1] - r7[5]:.1f}点・令和8年度"
          f"{r8[1] - r8[5]:.1f}点（再整理の表）")

    print("  ── 検査4　交付額と保険料への効き")
    judge(ato.count("計3,738千円") >= 2, "令和6年度の交付見込額3,738千円が2か所")
    judge(KIJUN == 6822, f"表112 の J から算定した基準額 {KIJUN:,}円＝6,822円")
    judge(f"月額約{-KOFU_SA}円下がります" in ato,
          f"9-3 の不確定要素⑤（月額約{-KOFU_SA}円）")
    judge("［要確認］" in ato and ato.count("［要確認］")
          == mae.count("［要確認］") + 5,
          f"［要確認］{mae.count('［要確認］')}→{ato.count('［要確認］')}件"
          "（交付額2・扱い2・効き1）")

    print("  ── 検査5　⚠ 保険料の額を動かしていないか")
    for go in ("6,822円", "6,143円", "5,464円", "765,752千円"):
        judge(mae.count(go) == ato.count(go),
              f"「{go}」の件数が同じ（{ato.count(go)}件）")

    print("  ── 検査6　⚠ 既にあった数が消えていないか")
    kesu = []
    for m in re.finditer(r"[0-9０-９][0-9０-９,.]*(?:人|％|千円|件|名|歳|回)",
                         mae):
        go = m.group(0)
        if mae.count(go) > ato.count(go):
            kesu.append(go)
    judge(not kesu, "消えた数が0件", "／".join(sorted(set(kesu))[:5]))

    print("  ── 検査7　受託者を主語とする語・他団体の名称")
    NG = ("受託者", "当社", "ビズアップ", "スクリプト", "北塩原", "金ヶ崎")
    aru = [w for w in NG if w in ato]
    judge(not aru, "0件", "／".join(aru))

    print()
    print(f"  {SAKI}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        for m in futeki:
            print(f"      {m}")
        sys.exit(1)
    print(f"   ○ 検査1〜7　不適合0件"
          f"（［要確認］{ato.count('［要確認］')}件・"
          f"［要協議］{ato.count('［要協議］')}件）")


if __name__ == "__main__":
    main()
