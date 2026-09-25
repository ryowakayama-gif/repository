# -*- coding: utf-8 -*-
"""業務工程管理表の更新（令和8年9月25日d・報告書再修正版のレビューv8）

  03_確認事項一覧　No.125〜126 を追加（発注者のご判断を要するもの）
  04_成果品管理　No.1（調査報告書）の「確定していない事項」を更新

根拠：`01_第10期_最新版成果品/川崎町_RedTeamレビューv8_再修正版_R8.9.25.md`
"""
import copy

import openpyxl

WB = "川崎町_業務工程管理表.xlsx"


def unmerge_below(ws, at_row):
    kept = []
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= at_row:
            kept.append((m.min_row, m.max_row, m.min_col, m.max_col))
            ws.unmerge_cells(str(m))
    return kept


def remerge(ws, kept, n):
    for r1, r2, c1, c2 in kept:
        ws.merge_cells(start_row=r1 + n, end_row=r2 + n,
                       start_column=c1, end_column=c2)


def copy_style(ws, src_row, dst_row, ncol):
    for c in range(1, ncol + 1):
        ws.cell(dst_row, c)._style = copy.copy(ws.cell(src_row, c)._style)
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def insert(ws, at, tpl, recs, ncol, count_col, old_last):
    n = len(recs)
    kept = unmerge_below(ws, at)
    ws.insert_rows(at, n)
    remerge(ws, kept, n)
    for i, rec in enumerate(recs):
        r = at + i
        copy_style(ws, tpl, r, ncol)
        for c, v in enumerate(rec, start=1):
            ws.cell(r, c).value = v
    new_last = old_last + n
    for r in range(at + n, ws.max_row + 1):
        cell = ws.cell(r, 2)
        if isinstance(cell.value, str) and cell.value.startswith("=COUNTIF"):
            cell.value = cell.value.replace(
                f"{count_col}5:{count_col}{old_last}",
                f"{count_col}5:{count_col}{new_last}")
    return new_last


wb = openpyxl.load_workbook(WB)

# ══════════════════════════════ 03_確認事項一覧（No.124 が 128行）
ws = wb["03_確認事項一覧"]
LAST = 128
NEW = [
    ("125", "(2)(6)", "R8.10",
     "ニーズ調査結果報告書の「（３）第9期調査との比較」を残すか",
     "ニーズ調査結果報告書[P618]〜[P629]の第9期調査との比較について、"
     "報告書本文に「配布数・回収数・回収率の違いなどから、"
     "比較対象とすべきものかについていささか疑問である。"
     "そのため、削除も検討する」との記述が残っています。"
     "①当該比較を残す（注記を書き改める）、"
     "②当該比較を削除する、のいずれとするかをご判断ください。"
     "残す場合は、増減の大きさをもって施策の効果を評価することはできない旨の"
     "注記に書き改めます。"
     "なお掲げられた6指標の増減ポイントの算術はいずれも正しいことを確認済みです。",
     "ニーズ調査結果報告書\n計画素案 第3章・第11章",
     "発注者", "確認待ち", "R8.10",
     "現行の記述は編集途中の判断メモであり、"
     "町名義で公表する報告書には残せない。"
     "あわせて[P619]の「今回」欄／「第9期」欄の説明が逆になっている点も是正する。"),
    ("126", "(2)", "R8.10",
     "在宅介護実態調査報告書の本編と資料編で割合の分母が異なる整理でよいか",
     "再修正版で資料編の割合の分母が「設問ごとの有効回答数」から"
     "「回答者数142（無回答・非該当を含む）」に変更されました。"
     "件数は本編と同一ですが、割合は一致しません"
     "（例：B問4「フルタイム」は本編30.7％（n=75）・資料編16.2％（n=142）。"
     "いずれも23件）。"
     "両文書とも「同じ集計結果を再掲した」と記載しているため、"
     "①この整理のまま、両文書に「件数は同一だが割合の分母が異なる」旨の"
     "相互注記を入れる、②資料編を本編と同じ分母に戻す、"
     "のいずれとするかをご判断ください。",
     "在宅介護実態調査結果報告書（本編・資料編）",
     "発注者", "確認待ち", "R8.10",
     "分母を固定して無回答・非該当を明示する資料編の方式自体は妥当。"
     "付表1は26表356項目すべて新方式で一致することを検算済み。"
     "①を推奨する。"),
]
LAST = insert(ws, LAST + 1, LAST, NEW, 10, "H", LAST)

# ══════════════════════════════ 04_成果品管理　No.1（調査報告書）
ws2 = wb["04_成果品管理"]
ws2.cell(5, 8).value = (
    "【R8.9.25 更新】再修正版3文書をレビュー（v8）。"
    "約6,400項目を個票から再検算し、数値の誤りは37件"
    "（資料編 付表2のn=16の層で割合が偶数丸めとなっているもの。"
    "凡例は「小数第2位を四捨五入」）。"
    "これ以外の数値の不一致は0件。"
    "納品前に必ず処理するもの4件："
    "①本編・ニーズ調査の目次が未更新（1行のみ）、"
    "②3文書とも文書プロパティの最終更新者名が残存、"
    "③ニーズ調査[P629]に編集メモが残存、"
    "④ニーズ調査 表73のキャプションが表の下にある。"
    "ほか数値・説明の修正4件、注記の追加5件。"
    "詳細は「川崎町_RedTeamレビューv8_再修正版_R8.9.25.md」。")

wb.save(WB)
print("保存しました：", WB)
print(f"03_確認事項一覧　No.126まで（{LAST - 4}件）")
print("04_成果品管理　No.1 の確定していない事項を更新しました。")
