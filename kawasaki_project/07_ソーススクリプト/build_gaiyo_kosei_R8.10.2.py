# -*- coding: utf-8 -*-
"""概要版（8頁程度・3,500部）の構成案（令和8年10月2日）

08_作業順位 の順位8。成果品5（仕様書8）は未着手（進捗0.25）。
**本文の章立てと図は固まっているため、構成案と割付け案は当方で作れる。**

  出力　05_試算・管理シート/川崎町_概要版_構成案_R8.10.2.xlsx

仕様書8の仕様　A4版／コート紙／1色刷り／8頁程度／3,500部

3,500部は町の世帯数（約3,000世帯）を上回る数であり、
**全戸配布を前提とした部数**と考えられる。
したがって、読み手は委員ではなく**町民**である。
計画書の要約ではなく、**町民にとって何が変わるのかを示すもの**とする。

  ① 1色刷りであるため、図は濃淡と模様で見分けられるものに限る
  ② 8頁（A4・二つ折りではなく中綴じであれば8頁＝2枚）
  ③ 保険料は最も関心の高い事項であり、1頁を充てる
  ④ 認知症は基本法への対応として新たに章を立てたため、1頁を充てる

**本文に現にある章節と図からのみ組み立てる。**
概要版のために新しい数値を作らない。
"""
import sys

import docx
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, "07_ソーススクリプト")
from data_zuhyo import ZU                      # noqa: E402

SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.8_第9期対比版.docx"
OUT = "05_試算・管理シート/川崎町_概要版_構成案_R8.10.2.xlsx"

TITLE = PatternFill("solid", fgColor="1F4E78")
HEAD = PatternFill("solid", fgColor="5B9BD5")
LEAD = PatternFill("solid", fgColor="F2F2F2")
HI = PatternFill("solid", fgColor="DAE3F3")

# （頁, 見出し, 本文の対応箇所, 載せる図表, 本文に書くこと, 字数の目安）
PAGES = [
    (1, "表紙　川崎町の高齢者保健福祉計画・第10期介護保険事業計画",
     "表紙・1-1・1-3",
     "―（町の風景の写真又は体系図を1点）",
     "計画の名称・計画期間（令和9〜11年度）・策定の趣旨を3行で。"
     "「この計画は、町の介護保険料と高齢者福祉の3年間を定めるものです」"
     "という一文を大きく置く。", 200),
    (2, "川崎町の高齢者のいま",
     "2-1・2-2・2-3",
     "図2-1（年齢階級別推移）／図2-2（高齢化率比較）",
     "高齢化率42.0％（令和8年6月末）。"
     "65歳以上3,209人のうち75歳以上が1,742人。"
     "第1号被保険者は減るが認定者は増える、という形を1枚で示す。", 400),
    (3, "介護サービスの利用と費用",
     "2-3・2-4・2-7",
     "図2-5（受給者の区分別構成）／図2-6（給付費の構成）",
     "月平均450.0人が利用。給付費9.98億円。"
     "施設サービスが49.4％を占めること、"
     "町外施設の利用が常態化していること（住所地特例24人）を書く。", 400),
    (4, "第9期（令和6〜8年度）の振り返り",
     "3-2・3-3・3-4",
     "図3-2（計画値と実績値の対比）",
     "計画に対する実績の達成率98.7％。"
     "居宅が計画を下回り施設が上回ったこと（在宅から施設への移行）。"
     "第10期に向けた課題を5本に絞って示す。", 450),
    (5, "第10期の基本理念と施策の体系",
     "4-1・4-2・4-3",
     "計画の体系図（4-4）",
     "基本理念（第9期踏襲）と7つの柱、4つの重点。"
     "**町民から見て何が変わるのか**を重点ごとに1行で書く。", 450),
    (6, "認知症とともに生きるまちへ",
     "6-1・6-2・6-3",
     "図6-1（KPIの3層構造）",
     "認知症基本法に基づき、第10期から認知症施策推進計画を一体で定めること。"
     "「新しい認知症観」。本人・家族の意見を聴いて作ったこと。"
     "認知症カフェ「喫茶みかん」・チームオレンジ・相談窓口の案内。", 450),
    (7, "介護保険料", "9-3",
     "図9-1（保険料の推移と3パターン）",
     "**最も関心の高い頁。** 第9期6,500円からどう変わるか。"
     "保険料がどのように決まるか（給付費・基金・調整交付金）を"
     "3つの箱で示す。所得段階別の表（13段階）を小さく載せる。"
     "⚠ 金額は第3回策定委員会の決定後に確定する。", 500),
    (8, "計画の進め方と相談先", "10-1・10-2・5-3",
     "―（相談窓口の一覧表）",
     "進行管理（PDCAサイクル）と毎年度の公表。"
     "地域包括支援センター・保健福祉課の連絡先、"
     "介護予防教室・通いの場の案内。"
     "ご意見の寄せ方（意見公募手続）。", 400),
]

RYUI = [
    ("1色刷り", "図は濃淡と模様で見分けられるものに限る。"
     "本文の図は色で見分けるものがあるため、概要版では作り直す"
     "（data_zuhyo.py の数値から、1色刷り用に描き直す）。"),
    ("部数3,500部", "町の世帯数（約3,000世帯）を上回るため、"
     "全戸配布を前提とした部数と考えられる。"
     "読み手は委員ではなく町民であり、専門用語は避ける。"),
    ("保険料の確定時期", "第3回策定委員会（令和9年1月）で決定するため、"
     "概要版の入稿はその後になる。"
     "入稿から納品までの日数（確認事項No.18）と併せて逆算を要する。"),
    ("数値の出所", "概要版のために新しい数値を作らない。"
     "本文に現にある数値と図からのみ組み立てる。"
     "本文を直したら概要版も同じ版で直す。"),
    ("用語", "「地域包括ケアシステム」「日常生活圏域」「アウトカム」など、"
     "本文の用語をそのまま使わない。"
     "必要なものは欄外に短い説明を置く（順位10の用語解説と対にする）。"),
]


def main():
    doc = docx.Document(SOAN)
    secs = set()
    import re
    for p in doc.paragraphs:
        m = re.match(r"^(\d+-\d+)[　 ]", p.text.strip())
        if m:
            secs.add(m.group(1))
    zu = {d["no"] for d in ZU}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01_頁の割付け案"
    ws.cell(1, 1).value = "概要版（8頁程度・3,500部）の構成案"
    ws.cell(1, 1).font = Font(name="游ゴシック", size=14, bold=True,
                              color="FFFFFF")
    ws.cell(1, 1).fill = TITLE
    ws.cell(2, 1).value = (
        "仕様書8　A4版／コート紙／1色刷り／8頁程度／3,500部。"
        "3,500部は町の世帯数（約3,000世帯）を上回るため、"
        "全戸配布を前提とした部数と考えられます。"
        "読み手は委員ではなく町民であり、計画書の要約ではなく"
        "「町民にとって何が変わるのか」を示すものとしています。"
        "本文に現にある章節と図からのみ組み立てており、"
        "概要版のために新しい数値は作っていません。")
    ws.cell(2, 1).font = Font(name="游ゴシック", size=9)
    ws.cell(2, 1).fill = LEAD
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
    ws.row_dimensions[2].height = 64
    head = ["頁", "見出し", "本文の対応箇所", "載せる図表",
            "本文に書くこと", "字数の目安"]
    for c, (h, w) in enumerate(zip(head, [5, 40, 18, 36, 78, 10]), start=1):
        cell = ws.cell(4, c)
        cell.value = h
        cell.font = Font(name="游ゴシック", size=9, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    for i, rec in enumerate(PAGES):
        for c, v in enumerate(rec, start=1):
            cell = ws.cell(5 + i, c)
            cell.value = v
            cell.font = Font(name="游ゴシック", size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if rec[0] == 7:
                cell.fill = HI
    r = 5 + len(PAGES)
    ws.cell(r, 1).value = "計"
    ws.cell(r, 6).value = sum(p[5] for p in PAGES)
    for c in (1, 6):
        ws.cell(r, c).font = Font(name="游ゴシック", size=9, bold=True)
        ws.cell(r, c).fill = HI
    ws.freeze_panes = "A5"

    ws2 = wb.create_sheet("02_作るときの留意")
    ws2.cell(1, 1).value = "概要版を作るときの留意"
    ws2.cell(1, 1).font = Font(name="游ゴシック", size=14, bold=True,
                               color="FFFFFF")
    ws2.cell(1, 1).fill = TITLE
    for c, (h, w) in enumerate(zip(["項目", "内容"], [22, 110]), start=1):
        cell = ws2.cell(3, c)
        cell.value = h
        cell.font = Font(name="游ゴシック", size=9, bold=True, color="FFFFFF")
        cell.fill = HEAD
        ws2.column_dimensions[get_column_letter(c)].width = w
    for i, (a, b) in enumerate(RYUI):
        for c, v in enumerate((a, b), start=1):
            cell = ws2.cell(4 + i, c)
            cell.value = v
            cell.font = Font(name="游ゴシック", size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUT)

    # ══════════════════════════ 自己点検
    ng = []
    for p in PAGES:
        for s in str(p[2]).replace("・", " ").split():
            if re.match(r"^\d+-\d+$", s) and s not in secs:
                ng.append(f"{p[0]}頁　本文にない節を指している：{s}")
        for z in re.findall(r"図\d+-\d+", str(p[3])):
            if z not in zu:
                ng.append(f"{p[0]}頁　本文にない図を指している：{z}")
    print("概要版の構成案")
    print("  保存：", OUT)
    print(f"  {len(PAGES)}頁／字数の目安 計 {sum(p[5] for p in PAGES):,}字")
    print("  ── 自己点検")
    if ng:
        for m in ng:
            print("   ×", m)
        sys.exit(1)
    print("   ○ 指している章節と図は、いずれも本文に現にある")
    print("   ⚠ 1色刷りのため、本文の図はそのまま使えません。"
          "data_zuhyo.py の数値から描き直します。")


if __name__ == "__main__":
    import re
    main()
