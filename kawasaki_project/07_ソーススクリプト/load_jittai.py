# -*- coding: utf-8 -*-
"""在宅介護実態調査の個票を読む。設問ブロックは n 列で区切る。"""
import openpyxl

SRC = ("/home/user/repository/kawasaki_project/04_調査・入力・分析/"
       "R8調査データ/川崎町_在宅介護実態調査_集計データ_20260807.xlsx")


def load():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["データ"]
    NC = ws.max_column

    blocks = []          # dict(key, qname, sub, opts=[(列,選択肢)], ncol)
    qname = None
    sub = None
    cur = None
    for c in range(22, NC + 1):
        if ws.cell(1, c).value:
            qname = str(ws.cell(1, c).value)
        if ws.cell(2, c).value:
            sub = str(ws.cell(2, c).value)
            cur = None
        mark = ws.cell(4, c).value
        lab = ws.cell(3, c).value
        if cur is None:
            cur = dict(qname=qname, sub=sub, opts=[], ncol=None)
            blocks.append(cur)
        if mark == "n":
            cur["ncol"] = c
            cur = None
        elif lab:
            cur["opts"].append((c, str(lab).strip()))
    blocks = [b for b in blocks if b["ncol"]]
    for b in blocks:
        b["key"] = b["qname"] if b["qname"] != "A問8" else f"A問8_{b['sub']}"

    recs = []
    for r in range(5, ws.max_row + 1):
        no = ws.cell(r, 1).value
        if no in (None, "") or str(no) == "計":
            continue
        rec = dict(NO=str(no),
                   性別=ws.cell(r, 6).value,
                   年齢階級=ws.cell(r, 7).value,
                   要介護状態区分=ws.cell(r, 8).value,
                   ans={})
        for b in blocks:
            sel = [lab for c, lab in b["opts"]
                   if ws.cell(r, c).value not in (None, "")]
            rec["ans"][b["key"]] = sel
        recs.append(rec)
    return blocks, recs


AGE3 = {"65歳未満": "74歳以下", "65～69歳": "74歳以下", "70～74歳": "74歳以下",
        "75～79歳": "75～84歳", "80～84歳": "75～84歳",
        "85～89歳": "85歳以上", "90歳以上": "85歳以上"}


def layers(recs):
    """属性軸 -> [(区分名, [rec,...])]"""
    out = {}
    out["性別"] = [(v, [r for r in recs if r["性別"] == v])
                   for v in ["男性", "女性"]]
    ages = ["65歳未満", "65～69歳", "70～74歳", "75～79歳",
            "80～84歳", "85～89歳", "90歳以上"]
    out["年齢階級"] = [(v, [r for r in recs if r["年齢階級"] == v]) for v in ages]
    sx = []
    for s in ["男性", "女性"]:
        for a in ["74歳以下", "75～84歳", "85歳以上"]:
            sx.append((f"{s}／{a}",
                       [r for r in recs if r["性別"] == s
                        and AGE3.get(r["年齢階級"]) == a]))
    out["性別×年齢"] = sx
    cares = ["要支援１", "要支援２", "要介護１", "要介護２", "要介護３",
             "要介護４", "要介護５", "わからない"]
    out["要介護状態区分"] = [(v, [r for r in recs if r["要介護状態区分"] == v])
                             for v in cares]
    return out


if __name__ == "__main__":
    bl, recs = load()
    print("個票", len(recs), "件／集計ブロック", len(bl))
    for b in bl:
        print(f"  {b['key']:28s} 選択肢{len(b['opts'])}")
