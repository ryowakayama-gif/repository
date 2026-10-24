# -*- coding: utf-8 -*-
"""Red Team 再レビュー第4回（令和8年10月24日）への対応：影響範囲を確かめたうえで成果物とエビデンスの一式を ZIP にする

  python3 07_ソーススクリプト/build_redteam_r4_teishutsu_R8.10.24.py

ご指示「レビュー結果の確認と修正進めて下さい。先程同様影響範囲確認の上、成果物とエビデンスをまとめてアウトプットして下さい」
（令和8年10月24日・第4回）。組み立ては前回（build_redteam_r3eikyou_teishutsu_R8.10.24.py）と同じ。
対応表の正本は本スクリプトの TAIOU・SAIHANTEI。点検のログ・検算・差分・PDF は、その場で成果物と原票から作る。
第4回の F4-05 を受け、突合に用いた原票（年報・算定ブック）を 03_エビデンス/原票 に同梱し MD5 を README に載せる。

出力
  11_提出一式/川崎町_RedTeam再レビュー第4回対応_成果物とエビデンス_R8.10.24.zip
  01_第10期_最新版成果品/川崎町_RedTeam再レビュー第4回対応表_R8.10.24.md
  05_試算・管理シート/川崎町_RedTeam再レビュー第4回対応表_R8.10.24.xlsx
終了コード　0＝点検ログがすべて終了コード0・検算が成果物と原票に一致・旧表現0件・ZIP の MD5 が一覧と一致／1＝それ以外
"""
import collections
import csv
import difflib
import hashlib
import importlib.util
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

import docx
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))


def load(name, fn):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


RT1 = load("rt1", "build_redteam_teishutsu_R8.10.24.py")            # md5・zenbun・to_pdf・sc_genpyo
RT2 = load("rt2", "build_redteam_saireview_teishutsu_R8.10.24.py")  # 検算 §1・§2（報酬改定・入所・居住系）
MIKOMI = load("mikomi", "check_mikomi_morasi_R8.10.16.py")         # 素案 9-1 の表の読み取り
NINCHI = load("ninchi", "calc_ninchisho_R8.10.19.py")              # 認知症の人の数の算定の正本

HI = "R8.10.24"
ZIPNAME = f"11_提出一式/川崎町_RedTeam再レビュー第4回対応_成果物とエビデンス_{HI}.zip"
MD_OUT = f"01_第10期_最新版成果品/川崎町_RedTeam再レビュー第4回対応表_{HI}.md"
XLSX_OUT = f"05_試算・管理シート/川崎町_RedTeam再レビュー第4回対応表_{HI}.xlsx"
REVIEW = "09_元資料/レビュー/川崎町_第10期計画_RedTeam再レビュー第4回_R8.10.24受領.md"
NENPO = "09_元資料/R7実績データ/年報データ_2025_川崎町.xlsx"
NENPO6 = "09_元資料/R8実績データ/R8.9.8受領版/年報データ_2024_川崎町.xlsx"
GENPYO = [NENPO, NENPO6, "05_試算・管理シート/川崎町_第10期_計画見込量_R8.9.15.xlsx",
          "05_試算・管理シート/川崎町_第10期_計画見込量_中長期・保険料内訳_R8.10.17.xlsx",
          "05_試算・管理シート/川崎町_策定委員会_逆算工程表_R8.10.20.xlsx"]
RYUHO = "05_試算・管理シート/川崎町_公表用の留保一覧_v2.36.xlsx"
OTOSHI = "05_試算・管理シート/川崎町_公表用で落とした文_v2.36.xlsx"

SEIKA = [  # (区分, ファイル, 第4回レビューの対象の版)
    ("協議用素案", "01_第10期_最新版成果品/川崎町_計画書素案_v2.36_協議用.docx",
     "01_第10期_最新版成果品/川崎町_計画書素案_v2.34_協議用.docx"),
    ("公表用の計画書（案）", "01_第10期_最新版成果品/川崎町_計画書案_公表用_v2.36.docx",
     "01_第10期_最新版成果品/川崎町_計画書案_公表用_v2.34.docx"),
    ("概要版", "01_第10期_最新版成果品/川崎町_計画書案_概要版_v2.36.docx",
     "01_第10期_最新版成果品/川崎町_計画書案_概要版_v2.34.docx"),
    ("第2回策定委員会資料", "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v21.docx",
     "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v19.docx"),
    ("想定問答集", "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v17.docx",
     "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v16.docx"),
]
S = SEIKA[0][1]
P = SEIKA[1][1]
SC = "07_ソーススクリプト/"
TENKEN = [  # (名前, コマンド)
    ("01_素案_自己点検", [SC + "check_soan_selfcheck_R8.9.28.py", S]),
    ("02_素案_保険料の検算", [SC + "check_hokenryo_R8.9.30.py", S]),
    ("03_素案_書体と紙面", [SC + "check_font_R8.9.30.py", S]),
    ("04_素案_8軸の網羅性（記載の網羅性。未実施10事項を含む）", [SC + "check_8jiku_R8.10.4.py", S]),
    ("05_素案_章節の参照の機械検索", [SC + "check_sansho_R8.10.24.py", S]),
    ("06_図表データ管理台帳（素案v2.36）", [SC + "build_zuhyo_daicho.py"]),
    ("07_事業所の受入れ様式（様式の検査。実データの全数照合ではない）", [SC + "check_jigyosho_uke_R8.10.15.py"]),
    ("08_見込量の漏れ", [SC + "check_mikomi_morasi_R8.10.16.py"]),
    ("09_認知症の人の数の整合", [SC + "check_ninchisho_seigo_R8.10.19.py"]),
    ("10_認定者数", [SC + "check_nintei_R8.10.16.py"]),
    ("11_公表用_書体と紙面", [SC + "check_font_R8.9.30.py", P]),
    ("12_委員会資料_自己点検", [SC + "check_2kai_selfcheck_R8.9.28.py"]),
    ("13_守るべき制約の走査", [SC + "check_seiyaku_R8.10.6.py"]),
    ("14_リポジトリの健全性", [SC + "check_repo_kenzensei_R8.10.4.py"]),
]
SCRIPTS = ["fix_soan_v2360_r4review.py", "build_kouhyou_soan.py", "build_gaiyo_v230.py", "build_2kai_v21.py",
           "build_2kai_qa_v17.py", "build_iinkai_gyotei_R8.10.20.py", "check_8jiku_R8.10.4.py",
           "calc_ninchisho_R8.10.19.py", "calc_hokenryo_R8.9.8b.py", "check_mikomi_morasi_R8.10.16.py",
           "build_redteam_r3eikyou_teishutsu_R8.10.24.py", "build_redteam_r4_teishutsu_R8.10.24.py"]

# ══════════════════════════ 対応表（正本）
TAIOU = [
    ("F4-01", "重大", "9-2 の表の令和7年度：内訳5項目の和1,084,826千円と標準給付費見込額1,084,587千円が239千円違う",
     "是正済", "年報で原因を確かめた。1,084,587千円は年報 様式3「保険給付支払状況」の計（支払額。審査支払手数料を含まない）で、"
     "内訳の総給付費998,222千円は様式2 の給付費（給付実績）。差239千円＝審査支払手数料656千円－（介護・介護予防サービス等"
     "諸費の支払額998,639千円－総給付費998,222千円＝417千円）。2つの集計を1つの列に混ぜていたため、令和7年度の欄を令和9〜"
     "11年度と同じ定義（5項目の和1,084,826千円）に改め、様式3 の支払額との違いを出典の注に書いた。令和9年度の1千円差は"
     "円単位の算定値（1,106,838,661円）を合計してから丸めたため。保険料の算定（令和9〜11年度）には波及しない",
     "素案・公表用 9-2", "03_エビデンス/検算.md §7・原票/", ""),
    ("F4-01 波及", "高", "（影響範囲の確認で検出）「標準給付費 令和6年度1,093,253千円→令和7年度1,084,587千円（0.8％減）」",
     "是正済", "1,093,253千円は令和6年度の様式3 の計1,092,620,753円に審査支払手数料632,700円を加えた額で、令和7年度の"
     "1,084,587千円（手数料を含まない）と基準が違っていた。同じ様式3 の支払額で比べ1,092,621千円→1,084,587千円・0.7％減に"
     "改めた（素案 7-2・9-3・3-4 の表、概要版、資料 ３の表・重点課題の表、問答集2か所）。令和8年度見込（町提供・手数料を含む）"
     "の前年比は参考値と注記", "素案 3-4・7-2・9-3／概要版／資料 v21 ３／問答集 v17", "03_エビデンス/検算.md §7・影響範囲の走査.txt", ""),
    ("F4-02", "高", "資料21頁：第3回を「令和9年1月」と書く文と、表・次回予定の「令和9年2月上旬〜中旬」",
     "是正済", "「第1回の当日に示された当初の日程は第3回を令和9年1月とする3回構成。第2回が11月24日となったため、上表は"
     "同日から逆算し第3回を令和9年2月上旬〜中旬としている」と書き分けた", "資料 v21 ８", "03_エビデンス/差分", ""),
    ("F4-03", "高", "パブコメ開始を12月29日→1月4日（6日）遅らせて「3月15日ちょうど」は、案Ⅰの余裕5日と合わない",
     "是正済", "誤りであった。工程の日数（逆算工程表の定数）から計算し直すと、1月4日開始では後続が6日ずつ遅れ、案Ⅰの納品は"
     "3月16日（火）で納期を1日超える。1月4日開始とするには第3回から入稿まで（10日）又は印刷製本を1日以上短縮する必要が"
     "あり、案Ⅱ（印刷2週間）なら3月9日（火）。資料・問答集・逆算工程表（xlsx）を改め、日付単位の工程を検算 §8 に置いた",
     "資料 v21 ８／問答集 v17／05_試算・管理シート/川崎町_策定委員会_逆算工程表_R8.10.20.xlsx",
     "03_エビデンス/検算.md §8", "町の確認に要する日数・印刷会社の入稿条件（町・印刷会社）"),
    ("F4-04", "高", "公表用の留保103件の処理結果がすべて未決定",
     "町のご判断", "台帳に「処理の優先」の列を加え、保険料・見込量（9-1〜9-3）・KPI（10-4）・認知症施策（第6章）の留保を"
     "「1 先行」とした。確定・注記・取下げの決定は町・委員会による（公表用は体裁確認用のまま）",
     "05_試算・管理シート/川崎町_公表用の留保一覧_v2.36.xlsx", "03_エビデンス/川崎町_公表用の留保一覧_v2.36.xlsx",
     "103件の処理の決定（公表・納品の条件）"),
    ("F4-05", "高", "年報原本がZIPになく、37サービスの突合を第三者が再検証できない",
     "是正済", "突合に用いた原票（年報 令和7年度・令和6年度、見込量の算定ブック、中長期・保険料内訳）を 03_エビデンス/原票 に"
     "同梱し、MD5 を README に載せた。対応表の表現は「当方で年報原本と突合（37件一致）。原本を同梱」とする。将来推計総括表"
     "（出力の版が複数）との突合と需要・供給の3ケースは未実施のまま",
     "―", "03_エビデンス/原票/・00_README.md", "総括表との突合・3ケース（第3回まで。No.159・147）"),
    ("F4-06", "中", "「本文118頁」は表紙・挨拶・目次を含む全頁数",
     "是正済", "数え直した：全118頁＝表紙・ご挨拶・目次5頁＋本文113頁、資料編（見込）13頁を含む計131頁。別冊に移した案は全106頁"
     "（本文101頁）・計119頁。「100頁程度」との差を本文のみ（＋13頁／別冊後＋1頁）と冊子全体（＋31頁／＋19頁）で併記した",
     "資料 v21 ８／問答集 v17", "03_エビデンス/差分", "数える範囲の決定（確認事項No.83）"),
    ("F4-07", "中", "KPI の目標が方向のみ・現状値に【町確認】",
     "町のご判断", "目標値・責任者・頻度は策定委員会の協議事項（確認事項No.70・34・109）。通いの場は調査の参加割合と行政実績の"
     "参加率を分けた（第3回 R3-08）。年次で把握できる活動・体制指標の追加は第3回に向けて町と整理する",
     "素案 10-4", "―", "確認事項No.70・34・109"),
    ("F4-08", "中", "8軸ログの未実施事項の説明に「差が未認定の認知症の人の規模」の古い記載",
     "是正済", "check_8jiku と calc_ninchisho の説明を「単純な差を未認定の人の数とは読まない」に改めた（点検ログ04 で確認）。"
     "8軸の51件は記載の網羅性の検査であり、未実施10事項は資料の未受領による（変更なし）",
     "07_ソーススクリプト/check_8jiku_R8.10.4.py", "03_エビデンス/点検ログ/04", ""),
    ("F4-09", "中", "認知症の本人・家族等の意見聴取と反映の実績が未確認",
     "町のご確認待ち", "意見聴取は令和8年10〜11月に実施予定。記録様式は整備済み。実績・反映の記録は町（確認事項No.90・137）",
     "素案 第6章", "―", "確認事項No.90・137"),
    ("5 横断", "改善", "事業一覧134件と施策体系143事業の関係",
     "是正済", "資料に「施策体系の全143事業は、第5・6章の134事業に第8章の9事業を加えたもの（ほかに再掲4件）」と書いた",
     "資料 v21 ８", "03_エビデンス/差分", ""),
]
SAIHANTEI = [  # 第4回レビュー §2 の再判定
    ("R3-01・02・04・06・08・09", "解消", "変更なし"),
    ("R3-03", "管理改善・未確定", "F4-04 のとおり処理の優先の列を加えた。決定は町"),
    ("R3-05", "記述是正・町判断待ち", "F4-02・F4-03 のとおり工程の記述を直した。契約上の手続は町（No.49）"),
    ("R3-07", "概ね解消", "変更なし"),
    ("R3-10", "一部解消", "F4-05 のとおり原票を同梱。総括表・3ケースは未実施"),
    ("R3-11", "町判断待ち", "F4-06 のとおり頁の数え方を分けた。範囲の決定は町（No.83）"),
]
ZANSON = [  # 旧表現の残存検索（第4回と前回までのレビューで挙がった語）
    "1,093,253", "0.8％減", "▲0.8％", "ちょうど", "第３回（令和９年１月）", "本文118頁", "本文を106頁", "本文は106頁",
    "標準給付費は令和6年度", "移行が進んでいます", "この移行が続く", "おおよその規模に当たります", "第12条第4〜7項",
    "年報による確定値", "本算定", "委託仕様書11（本仕様の定めのない", "仕様書11は", "通いの場参加率（65歳以上）",
    "参加率は15.9％", "6,235円", "5,583円", "6,366円", "5,714円", "約211円", "要件となります", "残高の確定",
    "結果と考えられます", "1,741保険者", "35保険者", "998,220", "5,959", "6,382", "25年間", "11つ", "8の事業所",
    "要件を満たします", "概ね半分以下", "生活支援コーディネーター25", "9-4の感度表", "町内14の指定", "第117条第５項",
]


def nenpo_totsugo(ng, csv_path):
    """令和7年度の給付費：素案 9-1 の37サービスの千円の値を、年報 様式２（給付費）の円単位の値と突き合わせる。"""
    so = MIKOMI.yomu_soan(S)
    ws = openpyxl.load_workbook(NENPO, data_only=True)["様式２（給付費）"]
    GROUP = {11, 12, 18, 21, 26, 33, 47, 52}
    rows = {}
    for r in range(13, 52):
        if r in GROUP:
            continue
        nm = next((ws.cell(r, c).value for c in (5, 4, 3) if ws.cell(r, c).value), None)
        if nm:
            rows[str(nm).strip()] = r
    soukei = ws.cell(52, 16).value

    def yen(na):
        if na == "介護予防支援":
            return ws.cell(rows["介護予防支援・居宅介護支援"], 8).value, "介護予防支援・居宅介護支援 H（予防）"
        if na == "居宅介護支援":
            return ws.cell(rows["介護予防支援・居宅介護支援"], 15).value, "介護予防支援・居宅介護支援 O（介護）"
        col, kubun = 15, "O（介護）"
        base = na
        if "（介護予防を含む）" in na:
            base, col, kubun = na.replace("（介護予防を含む）", ""), 16, "P（合計）"
        elif na.startswith("特定介護予防"):
            base, col, kubun = na[len("特定介護予防"):], 8, "H（予防）"
            base = "特定" + base
        elif na.startswith("介護予防"):
            base, col, kubun = na[len("介護予防"):], 8, "H（予防）"
        base = base.replace("特定福祉用具購入費", "特定福祉用具販売")
        hit = [n for n in rows if n == base or n.startswith(base + "（") or n.startswith("複合型サービス(" + base)]
        if not hit:
            return None, "年報に該当の行なし"
        return sum(ws.cell(rows[n], col).value or 0 for n in hit), "＋".join(hit) + " " + kubun

    out = [["サービス（素案 9-1）", "素案の千円", "年報の円", "年報の行・列", "円÷1000 を四捨五入", "一致"]]
    wa_en, wa_sen, itchi = 0, 0, 0
    for na, v in so.items():
        y, w = yen(na)
        ok = y is not None and round(y / 1000) == round(v)
        itchi += ok
        wa_en += y or 0
        wa_sen += round(v)
        out.append([na, f"{round(v):,}", f"{y:,}" if y is not None else "", w,
                    f"{round(y / 1000):,}" if y is not None else "", "○" if ok else "×"])
    out.append(["計", f"{wa_sen:,}", f"{wa_en:,}", f"総計 {soukei:,}円", f"{round(wa_en / 1000):,}", ""])
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as fp:
        csv.writer(fp).writerows(out)
    ok_all = itchi == len(so) and wa_en == soukei
    if not ok_all:
        ng.append("検算§3 年報との突合")
    return [f"## §3 令和7年度の給付費：素案 9-1 の37サービスと年報（様式２ 給付費）の円単位の突合（N11・R3-10）", "",
            f"- 原票：`{NENPO}` の「様式２（給付費）」（介護給付は O 列・予防給付は H 列・介護予防を含むものは P 列）",
            f"- 37サービスのうち、年報の円を千円未満四捨五入した値が素案と一致：{itchi}件／{len(so)}件",
            f"- 37サービスの円の和：{wa_en:,}円（年報の総計 {soukei:,}円と{'一致' if wa_en == soukei else '不一致'}）",
            f"- 素案の千円の和 {wa_sen:,}千円と、総計を丸めた {round(soukei / 1000):,}千円の差 {wa_sen - round(soukei / 1000):+,}千円は、"
            "37の値を別々に丸めたことによる差である（円の和が総計と一致するため、漏れ・重複ではない）",
            "- 明細は 突合_令和7年度給付費_年報様式2.csv", ""], ok_all


EIKYOU = [  # (指摘, 同じ趣旨の記述を拾う形, 残っていてはならない形)
    ("F4-01 給付費の集計の基準（様式2・様式3・手数料）", r"1,08[45],\d{3}千円|1,09[23],\d{3}千円|1,084,587|1,092,62", r"1,093,253|0\.8％減|▲0\.8％"),
    ("F4-02 第3回の時期", r"第[３3]回[^。]{0,12}令和[９9]年[１1２2]月", r"第３回（令和９年１月）"),
    ("F4-03 パブコメ開始と納品日", r"[１1]月[４4]日|ちょうど|納期を", r"ちょうど"),
    ("F4-06 頁数の数え方", r"\d+頁", r"本文118頁|本文を106頁|本文は106頁"),
    ("F4-08 推計値と実数の差", r"未認定|認定を受けて(い|おられ)ない", r"おおよその規模"),
    ("5 事業数の関係", r"134事業|143事業", None),
]


def eikyou_sousa(ng, path):
    """第4回レビューの指摘ごとに、同じ趣旨の記述を5冊で拾い、残っていてはならない形が0件であることを確かめる。"""
    L = ["影響範囲の走査（第4回レビューの指摘と同じ趣旨の記述を、提出の5冊の本文・表から拾ったもの）", "",
         "各行：［冊］前後の文。「残っていてはならない形」が0件であることを確かめ、拾った記述は是正後の文言として目で確かめる。", ""]
    texts = {}
    for k, f, _ in SEIKA:
        texts[k] = list(dict.fromkeys(RT1.zenbun(f)))
    kei_ng = 0
    for mi, hiro, dame in EIKYOU:
        L += [f"## {mi}", f"拾う形：{hiro}", f"残っていてはならない形：{dame or '（なし）'}"]
        for k, _, _ in SEIKA:
            hits = [s_ for s_ in texts[k] if re.search(hiro, s_)]
            n_dame = sum(len(re.findall(dame, s_)) for s_ in texts[k]) if dame else 0
            kei_ng += n_dame
            L.append(f"［{k}］拾った記述 {len(hits)}件・残っていてはならない形 {n_dame}件")
            for s_ in hits:
                m = re.search(hiro, s_)
                L.append("    …" + s_[max(0, m.start() - 50):m.end() + 50].replace("\n", " ") + "…")
        L.append("")
    L.append(f"残っていてはならない形の合計 {kei_ng}件")
    open(path, "w", encoding="utf-8").write("\n".join(L) + "\n")
    if kei_ng:
        ng.append("影響範囲の走査")
    return kei_ng


def kensan(ng, d_ev):
    RT2.S = S
    RT2.SEIKA = SEIKA
    moto = RT2.kensan(ng)
    head = moto.split("## §3")[0]
    L = [head.rstrip(), ""]
    sec3, _ = nenpo_totsugo(ng, os.path.join(d_ev, "突合_令和7年度給付費_年報様式2.csv"))
    L += sec3
    r = NINCHI.santei(NINCHI.machi_kaikyu(), "認知症")
    gyo = [m[6] for m in r["meisai"]]
    wa, kei = sum(round(x) for x in gyo), round(r["sei"])
    sz = "\n".join(RT1.zenbun(S))
    ok4 = f"内訳の和（{wa}人）と計（{kei}人）" in sz
    if not ok4:
        ng.append("検算§4 認知症の注")
    L += ["## §4 認知症の人の数の内訳と計（R3-09）", "",
          "- 算定の正本 `07_ソーススクリプト/calc_ninchisho_R8.10.19.py`（性・年齢階級別の有病率×町の性・年齢階級別人口）",
          "- 年齢階級別（小数）：" + "・".join(f"{x:.2f}" for x in gyo) + f"人、計 {r['sei']:.2f}人",
          f"- 各行を四捨五入した和 {wa}人／計を四捨五入 {kei}人（差 {kei - wa}人）",
          f"- 素案 6-1 の注にこの説明があること：{'○' if ok4 else '×'}", "",
          "## §5 生活支援コーディネーター（N03）", "",
          "- 原票抜粋は 原票抜粋_生活支援コーディネーター.csv。確認事項No.190（町のご確認待ち）", ""]
    sz = "\n".join(RT1.zenbun(S))
    m = re.search(r"給付費([\d,]+)千円÷介護保険事業状況報告の受給者数延べ([\d,]+)人月", sz)
    kyufu, nin = int(m.group(1).replace(",", "")), int(m.group(2).replace(",", ""))
    tanka = kyufu / nin
    rates = [float(x) / 100 for x in re.findall(r"令和\d+年度([\d.]+)％",
                                                [x for x in RT1.zenbun(S) if "見込交付割合は全国の仮置値" in x][0])]
    dm = sum(10 * tanka * 12 * (0.23 + 0.05 - e) for e in rates) * 1000 / 0.96 / 9743.1 / 12
    ok6 = round(dm) == 212 and "10人では約212円" in sz
    if not ok6:
        ng.append("検算§6 29.9万円")
    L += ["## §6 入所・居住系の1人1月あたり給付費（1-2 の留保「29.9万円は概数」）", "",
          f"- 素案 5-5：令和7年度の入所・居住系の給付費 {kyufu:,}千円 ÷ 年報の受給者数 延べ{nin:,}人月 ＝ {tanka:.2f}千円（約{tanka / 10:.1f}万円）",
          "- 受給者数は 様式１の６（15）施設介護サービス受給者数の総数、様式１の７（18）認知症対応型共同生活介護（短期利用以外）・"
          "地域密着型介護老人福祉施設入所者生活介護、様式１の７（16）特定施設入居者生活介護（短期利用以外）の合計（年度の延べ人月）",
          f"- 丸める前の {tanka:.2f}千円で10人の月額を計算すると {dm:.2f}円 → 約{round(dm)}円（29.9万円による211.75円と同じく212円）",
          f"- 素案 5-5 の「約212円」と一致：{'○' if ok6 else '×'}", ""]
    L += kensan7(ng) + kensan8(ng)
    return "\n".join(L)


def yoshiki3(path):
    ws = openpyxl.load_workbook(path, data_only=True)["様式３"]
    out, on = {}, False
    for r in ws.iter_rows(values_only=True):
        v = [x for x in r if x is not None]
        if v and "保険給付支払状況" in str(v[0]):
            on = True
        elif on and v and isinstance(v[-1], (int, float)) and len(v) >= 2:
            out[str(v[0]).strip()] = v[1]
    return out


def kensan7(ng):
    """F4-01：9-2 の表の令和7年度と、令和6年度→令和7年度の比較の基準。"""
    y7, y6 = yoshiki3(NENPO), yoshiki3(NENPO6)
    d = docx.Document(S)
    t = [t for t in d.tables if t.rows[0].cells[0].text.strip() == "区分"
         and "令和7年度" in t.rows[0].cells[1].text and "令和9年度" in t.rows[0].cells[2].text][0]
    row = {r.cells[0].text.strip(): r for r in t.rows}
    num = lambda x: int(re.sub(r"[^\d]", "", x))
    uchi = ["総給付費", "特定入所者介護サービス費等", "高額介護サービス費等", "高額医療合算介護サービス費等", "審査支払手数料"]
    L = ["## §7 令和7年度の標準給付費と保険給付費の集計の基準（F4-01）", "",
         "| 年度 | 内訳5項目の和 | 素案の標準給付費見込額 | 差 |", "|---|---|---|---|"]
    ok = True
    for y, ci in (("令和7年度", 1), ("令和9年度", 2), ("令和10年度", 3), ("令和11年度", 4)):
        w = sum(num(row[k].cells[ci].text) for k in uchi)
        a = num(row["標準給付費見込額"].cells[ci].text)
        L.append(f"| {y} | {w:,}千円 | {a:,}千円 | {w - a:+} |")
        if y == "令和7年度":
            ok &= w == a
        else:
            ok &= abs(w - a) <= 1
    wb = openpyxl.load_workbook("05_試算・管理シート/川崎町_第10期_計画見込量_R8.9.15.xlsx", data_only=True)["⑦保険料算定"]
    r9 = [wb.cell(r, 2).value for r in range(2, 7)]
    shohi = y7["介護サービス等諸費"] + y7["介護予防サービス等諸費"]
    L += ["", f"- 令和7年度は5項目の和に改めた（{'○ 一致' if ok else '× 不一致'}）。",
          f"- 年報 令和7年度 様式3「保険給付支払状況」の計 {y7['計']:,}円（支払額。審査支払手数料を含まない）",
          f"  ＝ 介護サービス等諸費 {y7['介護サービス等諸費']:,} ＋ 介護予防サービス等諸費 {y7['介護予防サービス等諸費']:,}"
          f" ＋ 高額 {y7['高額介護サービス等費']:,} ＋ 高額医療合算 {y7['高額医療合算介護サービス等費']:,}"
          f" ＋ 特定入所者 {y7['特定入所者介護サービス等費']:,}",
          f"- 諸費の支払額 {shohi:,}円 と 様式2 の給付費 998,222,421円 の差 {shohi - 998222421:,}円（約417千円）",
          f"- 旧の差239千円 ＝ 審査支払手数料 656千円 － 417千円",
          f"- 令和9年度の1千円差：算定ブック（⑦保険料算定）の円単位 "
          + "＋".join(f"{x:,.0f}" for x in r9) + f" ＝ {sum(r9):,.0f}円 → {round(sum(r9) / 1000):,}千円（各項目を丸めた和は1,106,840千円）",
          "",
          "| 比較 | 令和6年度 | 令和7年度 | 増減率 |", "|---|---|---|---|",
          f"| 旧（基準が違う） | 1,093,253千円（様式3＋審査支払手数料632,700円） | 1,084,587千円（様式3） | ▲0.8％ |",
          f"| 新（様式3 の支払額で揃える） | {round(y6['計'] / 1000):,}千円 | {round(y7['計'] / 1000):,}千円 | "
          f"{(round(y7['計'] / 1000) - round(y6['計'] / 1000)) / round(y6['計'] / 1000) * 100:.2f}％ |",
          "", "- 保険料の算定は令和9〜11年度の標準給付費見込額（3,313,925千円）によるため、本件の是正は保険料に波及しない"
          "（点検ログ02 の保険料の検算は終了コード0）", ""]
    if not ok:
        ng.append("検算§7")
    return L


def kensan8(ng):
    """F4-03：パブリックコメントの開始日ごとの工程（日付単位）。日数は逆算工程表の定数による。"""
    import datetime as dt
    GY = load("gyotei", "build_iinkai_gyotei_R8.10.20.py")
    Y = "月火水木金土日"
    f = lambda x: f"{x.month}/{x.day}（{Y[x.weekday()]}）"
    L = ["## §8 パブリックコメントの開始日と納品日（F4-03）", "",
         f"日数の正本：`07_ソーススクリプト/build_iinkai_gyotei_R8.10.20.py`（第2回{GY.KAISAI2}・パブコメ{GY.PC_DAYS}日・"
         f"終了から第3回{GY.T_3}日・第3回から入稿{GY.T_NYUKO}日・印刷 案Ⅰ{GY.INSATSU}日／案Ⅱ14日・納期{GY.NOUKI}）", "",
         "| 案 | パブコメ開始 | 終了 | 第3回 | 入稿 | 納品 | 納期との差 |", "|---|---|---|---|---|---|---|"]
    kekka = {}
    for kaishi in (GY.KAISAI2 + dt.timedelta(days=GY.T_PC_START), dt.date(2027, 1, 4)):
        for an, ins in (("案Ⅰ", GY.INSATSU), ("案Ⅱ", 14)):
            owari = kaishi + dt.timedelta(days=GY.PC_DAYS)
            dai3 = owari + dt.timedelta(days=GY.T_3)
            nyuko = dai3 + dt.timedelta(days=GY.T_NYUKO)
            nohin = nyuko + dt.timedelta(days=ins)
            sa = (GY.NOUKI - nohin).days
            kekka[(kaishi.month, an)] = nohin
            L.append(f"| {an} | {f(kaishi)} | {f(owari)} | {f(dai3)} | {f(nyuko)} | {f(nohin)} | "
                     f"{'余裕' + str(sa) + '日' if sa >= 0 else '超過' + str(-sa) + '日'} |")
    s_ = "\n".join(RT1.zenbun(SEIKA[3][1]))
    ok = (kekka[(12, "案Ⅰ")] == dt.date(2027, 3, 10) and kekka[(1, "案Ⅰ")] == dt.date(2027, 3, 16)
          and "３月16日（火）" in s_ and "ちょうど" not in s_)
    L += ["", "- 第3回が土日に当たる場合は前後の平日に置く（日数は暦日）。町の確認・印刷会社の入稿条件は含まない",
          f"- 資料 v21 の記載（1月4日開始で案Ⅰは3月16日・1日超過）と一致：{'○' if ok else '×'}", ""]
    if not ok:
        ng.append("検算§8")
    return L


def main():
    work = tempfile.mkdtemp(prefix="redteam3_")
    root = os.path.join(work, f"川崎町_RedTeam再レビュー第4回対応_{HI}")
    d_seika, d_taiou, d_ev, d_sc = (os.path.join(root, x) for x in
                                    ("01_成果物", "02_レビュー対応", "03_エビデンス", "04_スクリプト"))
    d_log = os.path.join(d_ev, "点検ログ")
    for x in (d_seika, d_taiou, d_ev, d_sc, d_log):
        os.makedirs(x)
    ng = []
    print(f"Red Team 再レビュー第4回対応の一式（{HI}）")

    print("  ── 1 成果物と PDF")
    ichiran = []
    for kubun, f, kyu in SEIKA:
        shutil.copy2(f, d_seika)
        pdf, n = RT1.to_pdf(f, d_seika)
        ichiran.append((kubun, os.path.basename(f), n, RT1.md5(f), os.path.basename(kyu)))
        print(f"   ○ {kubun} {os.path.basename(f)}（PDF {n}頁）")
        if not n:
            ng.append(f"PDF {f}")

    print("  ── 2 点検ログ")
    logs = []
    for na, cmd in TENKEN:
        r = subprocess.run([sys.executable] + cmd, capture_output=True, text=True, timeout=900)
        out = "\n".join(x for x in (r.stdout + r.stderr).splitlines()
                        if "FutureWarning" not in x and "warnings.warn" not in x)
        open(os.path.join(d_log, f"{na}.txt"), "w", encoding="utf-8").write(
            f"$ python3 {' '.join(cmd)}\n終了コード {r.returncode}\n\n{out}\n")
        logs.append((na, r.returncode))
        print(("   ○ " if r.returncode == 0 else "   × ") + f"{na}（終了コード {r.returncode}）")
        if r.returncode:
            ng.append(na)

    print("  ── 3 検算・原票との突合・留保一覧・旧表現の残存検索・差分")
    n_e = eikyou_sousa(ng, os.path.join(d_ev, "影響範囲の走査.txt"))
    print(("   ○ " if n_e == 0 else "   × ") + f"影響範囲の走査（{len(EIKYOU)}指摘×5冊）残っていてはならない形 {n_e}件")
    open(os.path.join(d_ev, "検算.md"), "w", encoding="utf-8").write(kensan(ng, d_ev))
    print(("   ○ " if not [x for x in ng if x.startswith("検算")] else "   × ") + "検算 §1〜§8（成果物・年報・算定ブック・工程の定数と突合）")
    RT1.sc_genpyo(os.path.join(d_ev, "原票抜粋_生活支援コーディネーター.csv"))
    d_gen = os.path.join(d_ev, "原票")
    os.makedirs(d_gen, exist_ok=True)
    for f in GENPYO:
        shutil.copy2(f, d_gen)
    print(f"   ○ 原票 {len(GENPYO)}件を同梱（MD5 は README）")
    for f in (RYUHO, OTOSHI):
        shutil.copy2(f, d_ev)
    wb = openpyxl.load_workbook(RYUHO, read_only=True)
    rr = list(wb["留保の一覧"].iter_rows(values_only=True))
    kara = sum(1 for x in rr[1:] if not (x[5] and x[6] and x[7] and x[8]))
    print(("   ○ " if kara == 0 else "   × ") + f"留保の一覧 {len(rr) - 1}件・優先／確認先／期限／既定の扱いが空の行 {kara}件"
          f"（先行 {sum(1 for x in rr[1:] if str(x[5]).startswith('1'))}件）")
    if kara:
        ng.append("留保の一覧の空欄")
    texts = {k: "\n".join(RT1.zenbun(f)) for k, f, _ in SEIKA}
    zan = ["語\t" + "\t".join(k for k, _, _ in SEIKA)]
    zan_ng = 0
    for w in ZANSON:
        c = [texts[k].count(w) for k, _, _ in SEIKA]
        zan.append(w + "\t" + "\t".join(str(x) for x in c))
        zan_ng += sum(c)
    zan += ["", "公表用・概要版に残さない内部の記号"]
    for w in ("【委員会協議】", "【町確認】", "⚠", "確認事項No"):
        c = [texts[k].count(w) for k in ("公表用の計画書（案）", "概要版")]
        zan.append(f"{w}\t公表用 {c[0]}\t概要版 {c[1]}")
        zan_ng += sum(c)
    ryuho = {w: texts["公表用の計画書（案）"].count(w) for w in ("［要確認］", "［要協議］", "［要内訳］")}
    zan += ["", "公表用に残る留保（扱いは 川崎町_公表用の留保一覧_v2.36.xlsx）"] + [f"{k}\t{v}" for k, v in ryuho.items()]
    open(os.path.join(d_ev, "旧表現の残存検索.txt"), "w", encoding="utf-8").write(
        "第4回と前回までのレビューで挙がった旧の表現が、提出の5冊に残っていないか（件数）\n\n" + "\n".join(zan) +
        f"\n\n旧表現・内部の記号の合計 {zan_ng}件\n")
    print(("   ○ " if zan_ng == 0 else "   × ") + f"旧表現・内部の記号の残存 {zan_ng}件（{len(ZANSON)}語×5冊＋4語×2冊）")
    if zan_ng:
        ng.append("旧表現の残存")
    for kubun, f, kyu in SEIKA:
        a, b = RT1.zenbun(kyu), RT1.zenbun(f)
        diff = list(difflib.unified_diff(a, b, os.path.basename(kyu), os.path.basename(f), n=0, lineterm=""))
        open(os.path.join(d_ev, f"差分_{kubun}.txt"), "w", encoding="utf-8").write(
            f"{kubun}：{os.path.basename(kyu)}（第4回レビューの対象）→ {os.path.basename(f)}\n"
            "段落・表のセルの単位の差分（-＝旧、+＝新）\n\n" + "\n".join(diff) + "\n")
        print(f"   ○ 差分_{kubun}（{sum(1 for x in diff if x.startswith('+') and not x.startswith('+++'))}行が新）")

    print("  ── 4 対応表（md・xlsx）")
    c = collections.Counter(r[3] for r in TAIOU)
    md = ["# Red Team 再レビュー第4回（令和8年10月24日）への対応表", "",
          f"> レビュー原本：`{REVIEW}`（素案 v2.34・公表用 v2.34・概要版 v2.34・資料 v19・問答集 v16 が対象）",
          "> 対応後の版：協議用素案 Ver.2.36／公表用 v2.36／概要版 v2.36／委員会資料 v21／想定問答集 v17",
          "> 指摘の語だけでなく同じ趣旨の記述を5冊で走査し（03_エビデンス/影響範囲の走査.txt）、波及したものも是正した",
          "> 本表の正本は `07_ソーススクリプト/build_redteam_r4_teishutsu_R8.10.24.py` の TAIOU・SAIHANTEI", "",
          "## 1　第4回の指摘（F4-01〜F4-09）と影響範囲の確認で是正したもの", "",
          "| ID | 重要度 | 指摘 | 対応 | 対応の内容 | 箇所 | エビデンス | 残る事項 |", "|---|---|---|---|---|---|---|---|"]
    for row in TAIOU:
        md.append("| " + " | ".join(x.replace("|", "／") for x in row) + " |")
    md += ["", "対応の区分：" + "／".join(f"{k} {v}件" for k, v in c.items()), "",
           "## 2　前回11指摘の再判定への対応", "", "| 前回ID | 第4回の判定 | 今回の対応 |", "|---|---|---|"]
    for row in SAIHANTEI:
        md.append("| " + " | ".join(row) + " |")
    md += ["", "## 3　第4回レビュー §6・§7 の条件との対応", "",
           "| 優先 | 内容 | 状態 |", "|---|---|---|",
           "| 最優先 | 令和7年度標準給付費239千円差の原因と算定への波及 | 是正済。原因は集計の基準の混在。保険料（令和9〜11年度）への波及なし（検算 §7） |",
           "| 最優先 | 第3回開催月の記載統一・1月4日パブコメ案の工程再計算 | 是正済（検算 §8。1月4日開始なら案Ⅰは3月16日・1日超過） |",
           "| 高 | 4回予定→3回の契約上の手続、町の書面決定 | 町の契約担当（No.49） |",
           "| 高 | 公表用留保103件の決裁 | 優先の列を加えた。決定は町・委員会 |",
           "| 高 | 町内外別の利用・供給能力・見込量の3ケース | 第3回まで（No.159・147） |",
           "| 高 | 将来推計・算定ブック・基金・交付金原票の独立突合 | 年報・算定ブックは同梱。総括表・基金決算・交付金原票は未実施 |",
           "| 中 | 生活支援コーディネーター26人（No.190）／認知症本人等の意見聴取（No.90・137）／認定率・自給率等の実績値／頁数の範囲（No.83） | 町のご確認待ち |",
           "", "## 4　当方で確かめられないもの（未実施として残す）", "",
           "- 見える化の将来推計総括表（出力の版が複数）・基金の決算・交付金の該当状況調査の原本と成果物の全数の突合は未実施。",
           "- 令和7年度の諸費の支払額（様式3）と給付実績（様式2）の差417千円の内訳（月のずれ等）は年報からは分けられない。",
           "- 委託契約書は手元にない。全3回への変更の契約上の扱いは町の契約担当のご判断による。",
           "- 「8軸51/51適合」は記載の網羅性の点検であり、同じログの未実施10事項は資料の未受領による。", ""]
    open(MD_OUT, "w", encoding="utf-8").write("\n".join(md))
    shutil.copy2(MD_OUT, d_taiou)
    wbo = openpyxl.Workbook()
    ws = wbo.active
    ws.title = "第4回の指摘"
    ws.append(["ID", "重要度", "指摘", "対応", "対応の内容", "箇所", "エビデンス", "残る事項"])
    for row in TAIOU:
        ws.append(list(row))
    for i, w in enumerate((7, 9, 40, 14, 80, 28, 34, 30), 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws2 = wbo.create_sheet("前回指摘の再判定")
    ws2.append(["前回ID", "第4回の判定", "今回の対応"])
    for row in SAIHANTEI:
        ws2.append(list(row))
    for i, w in enumerate((24, 22, 80), 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for w_ in (ws, ws2):
        for r in w_.iter_rows():
            for cc in r:
                cc.alignment = Alignment(wrap_text=True, vertical="top")
                cc.font = Font(name="游ゴシック", size=9, bold=(cc.row == 1))
                if cc.row == 1:
                    cc.fill = PatternFill("solid", fgColor="D9E1F2")
        w_.freeze_panes = "A2"
    wbo.save(XLSX_OUT)
    shutil.copy2(XLSX_OUT, d_taiou)
    shutil.copy2(REVIEW, d_taiou)
    print(f"   ○ {len(TAIOU)}件（" + "／".join(f"{k}{v}" for k, v in c.items()) + f"）・再判定 {len(SAIHANTEI)}行")

    print("  ── 5 スクリプト")
    for f in SCRIPTS:
        shutil.copy2(SC + f, d_sc)

    print("  ── 6 README（成果品一覧・MD5）")
    rd = [f"# 川崎町 第10期計画　Red Team 再レビュー第4回対応の一式（{HI}）", "",
          "川崎町（宮城県柴田郡）高齢者保健福祉計画・第10期介護保険事業計画", "",
          "## 中身", "", "| フォルダ | 内容 |", "|---|---|",
          "| 01_成果物 | 是正後の5冊（docx）と、当方の環境で作った PDF |",
          "| 02_レビュー対応 | 対応表（md・xlsx）と第4回レビュー原本 |",
          "| 03_エビデンス | 影響範囲の走査・点検ログ・検算 §1〜§8・年報との突合の明細・公表用の留保一覧・旧表現の残存検索・差分（第4回レビューの対象の版から）・原票抜粋 |",
          "| 03_エビデンス/原票 | 突合に用いた原票（年報 令和7年度・令和6年度、見込量の算定ブック、中長期・保険料内訳、逆算工程表） |",
          "| 04_スクリプト | 是正と点検に用いたスクリプト（リポジトリの 07_ソーススクリプト と同じもの） |", "",
          "## 成果品一覧", "", "| 区分 | ファイル | PDF頁数（当方の環境） | MD5 | 比べた旧の版 |", "|---|---|---|---|---|"]
    for kubun, f, n, h, kyu in ichiran:
        rd.append(f"| {kubun} | {f} | {n} | `{h}` | {kyu} |")
    rd += ["", "## 原票（03_エビデンス/原票）", "", "| ファイル | リポジトリの場所 | MD5 |", "|---|---|---|"]
    for f in GENPYO:
        rd.append(f"| {os.path.basename(f)} | {f} | `{RT1.md5(f)}` |")
    rd += ["", "⚠ 頁数は当方の環境（HG丸ｺﾞｼｯｸM-PRO・游ゴシックを IPAゴシックで代用した LibreOffice）による。"
           "配布用の PDF は町の環境（Word）で作り、頁送りを確かめていただく。",
           "", "## 点検ログ", "", "| 点検 | 終了コード |", "|---|---|"]
    for na, rc in logs:
        rd.append(f"| {na} | {rc} |")
    rd += ["", "## 再現の手順", "", "リポジトリの `kawasaki_project/` で次を実行すると、同じ一式ができます。", "", "```",
           "python3 07_ソーススクリプト/fix_soan_v2360_r4review.py",
           "python3 07_ソーススクリプト/fill_toc_pages_v200.py 01_第10期_最新版成果品/川崎町_計画書素案_v2.36_協議用.docx",
           "python3 07_ソーススクリプト/build_kouhyou_soan.py",
           "python3 07_ソーススクリプト/fill_toc_pages_v200.py 01_第10期_最新版成果品/川崎町_計画書案_公表用_v2.36.docx",
           "python3 07_ソーススクリプト/build_gaiyo_v230.py", "python3 07_ソーススクリプト/build_iinkai_gyotei_R8.10.20.py",
           "python3 07_ソーススクリプト/build_2kai_v21.py", "python3 07_ソーススクリプト/build_2kai_qa_v17.py",
           "python3 07_ソーススクリプト/build_redteam_r4_teishutsu_R8.10.24.py", "```", ""]
    open(os.path.join(root, "00_README.md"), "w", encoding="utf-8").write("\n".join(rd))

    print("  ── 7 ZIP")
    os.makedirs(os.path.dirname(ZIPNAME), exist_ok=True)
    with zipfile.ZipFile(ZIPNAME, "w", zipfile.ZIP_DEFLATED) as z:
        for dp, _, fs in os.walk(root):
            for f in sorted(fs):
                full = os.path.join(dp, f)
                z.write(full, os.path.relpath(full, work))
    with zipfile.ZipFile(ZIPNAME) as z:
        names = z.namelist()
        ok = all(hashlib.md5(z.read(f"{os.path.basename(root)}/01_成果物/{f}")).hexdigest() == h
                 for _, f, _, h, _ in ichiran)
    print(("   ○ " if ok else "   × ") + f"{ZIPNAME}（{len(names)}ファイル・{os.path.getsize(ZIPNAME):,}バイト）"
          "・成果物の MD5 が一覧と一致")
    if not ok:
        ng.append("MD5")
    shutil.rmtree(work)
    if ng:
        print(f"   × 不適合 {len(ng)}件：" + "／".join(ng))
        sys.exit(1)
    print("   ○ 不適合0件")


if __name__ == "__main__":
    main()
