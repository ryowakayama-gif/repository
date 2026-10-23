# -*- coding: utf-8 -*-
"""協議用素案 Ver.2.28 → Ver.2.29（宮城県の第10期計画の骨子（案）との対応・地域支援事業費の範囲）
（令和8年10月23日）

  python3 07_ソーススクリプト/fix_soan_v2290_ken.py

ご指示「大雪広域のブランチ確認を進めて下さい」（令和8年10月23日）による。
比較材料のブランチは観点のみ参照し、記載は原典（宮城県・厚生労働省）で確かめて起こす。

反映するもの
  ① 1-2 表1「県の計画」の行
       旧「宮城県高齢者保健福祉計画・介護保険事業支援計画」は県の計画の名称と合わない。
       県は「みやぎ高齢者元気プラン」として、宮城県高齢者福祉計画（老人福祉法第20条の9）と
       介護保険事業支援計画（介護保険法第118条）を一体で定めている。
  ② 1-2 に小見出し・本文・対応表・出典・注記を置く
       第10期みやぎ高齢者元気プラン 骨子（案）（令和8年8月26日 推進委員会 資料2〜4・6）の
       3つの基本的目標・9つの基本課題と、本計画の章節との対応
  ③ 8-2 注記　第7期宮城県介護給付適正化取組方針（令和9〜11年度）の方向（同 資料6）
  ④ 9-1 地域支援事業費の見込み　見える化システムに入力する事業費は
       対象経費支出予定額であること（地域分析の手引き 令和8年9月改訂版 79頁）

原典　09_元資料/その他データ/宮城県_第10期みやぎ高齢者元気プラン_推進委員会R8/
        20260826shiryo.pdf（資料2 p10–13・資料3 p14・資料4 p15–21・資料5 p22–27・資料6 p28–30・資料7 p31）
      09_元資料/その他データ/厚労省_見える化システム等を活用した地域分析の手引き_R8.9改訂版.pdf（79頁）
      09_元資料/法令/認知症基本法_条文_令和5年法律第65号.txt（第13条第1項）
終了コード　0＝検査1〜6すべて適合／1＝1件でも不適合
"""
import copy
import os
import re
import shutil
import sys

import docx
import docx.table as dt
import docx.text.paragraph as dp
from docx.oxml.ns import qn

MOTO = "01_第10期_最新版成果品/川崎町_計画書素案_v2.28_協議用.docx"
SAKI = "01_第10期_最新版成果品/川崎町_計画書素案_v2.29_協議用.docx"
GENTEN = "09_元資料/その他データ/宮城県_第10期みやぎ高齢者元気プラン_推進委員会R8/20260826shiryo.pdf"
TEBIKI = "09_元資料/その他データ/厚労省_見える化システム等を活用した地域分析の手引き_R8.9改訂版.pdf"

KEN_MEI = "第10期みやぎ高齢者元気プラン（宮城県高齢者福祉計画・介護保険事業支援計画）"
MIDASHI = "　▌ 第10期みやぎ高齢者元気プラン（骨子（案））との対応"
HON = ("宮城県は、宮城県高齢者福祉計画（老人福祉法第20条の9）と介護保険事業支援計画"
       "（介護保険法第118条）を一体で定める「みやぎ高齢者元気プラン」の第10期（令和9年度〜"
       "令和11年度）を策定しています。令和8年8月26日のみやぎ高齢者元気プラン推進委員会に"
       "示された骨子（案）は、第9期までの基本理念「高齢者が地域で自分らしい生活を安心して"
       "送れる社会」を継承し、「みんなで支え合う地域づくり」「自分らしい生き方の実現」"
       "「安心できるサービスの提供」の3つの基本的目標を掲げています。高齢者福祉圏域は、"
       "第9期と同じく二次医療圏と同じ4圏域とされ、本町は仙南圏域（2市7町）に属します。"
       "また、第1期宮城県認知症施策推進計画を同プランの1つの章として一体で策定し、"
       "第7期宮城県介護給付適正化取組方針（令和9年度〜令和11年度）を同プランに収めるとしています。"
       "骨子（案）の基本課題と本計画の章節との対応は、次のとおりです。")
HYO = [
    ["県の基本的目標", "県の基本課題（骨子（案））", "本計画の対応箇所"],
    ["みんなで支え合う地域づくり",
     "地域包括ケアシステムの深化（地域包括ケア体制・多職種連携・介護家族の支援）",
     "第5章5-3（柱3 家族介護者への支援）・5-4（柱4）"],
    ["みんなで支え合う地域づくり", "地域支え合いと介護予防・生活支援の推進",
     "第5章5-1（柱1）・5-4（柱4）"],
    ["みんなで支え合う地域づくり", "安全な暮らしの確保（災害・感染症への備え、防犯・交通安全）",
     "第5章5-2（柱2）"],
    ["自分らしい生き方の実現", "生きがいに満ちた生活の実現", "第5章5-3（柱3 社会参加）"],
    ["自分らしい生き方の実現", "自分らしく生きるための権利擁護（権利擁護・高齢者虐待の防止）",
     "第5章5-4（柱4）"],
    ["自分らしい生き方の実現", "認知症の人にやさしいまちづくり（第1期宮城県認知症施策推進計画）",
     "第6章（柱5）"],
    ["安心できるサービスの提供", "介護サービス基盤の計画的な整備",
     "第5章5-5（柱6）、第7章7-1〜7-3、第9章9-1"],
    ["安心できるサービスの提供", "介護人材の確保・養成・定着", "第7章7-4・7-5"],
    ["安心できるサービスの提供",
     "介護サービスの質の確保・向上（第7期宮城県介護給付適正化取組方針）", "第8章8-1・8-2"],
]
HABA = (30, 50, 40)  # 列の幅の比
SHUTTEN = ("出典：宮城県「令和8年度第2回みやぎ高齢者元気プラン推進委員会」資料2・資料3・"
           "資料4・資料6（令和8年8月26日）")
CHUI = ("⚠ 骨子（案）の段階であり、施策の具体的な記載は県の中間案（第3回の推進委員会で示される"
        "予定）によります。県が挙げる第10期の目標指標（案）には、保険者機能強化推進交付金等の"
        "県内市町村平均得点率、在宅等での看取り数、介護職員の人数、特別養護老人ホーム入所定員数"
        "などがあり（目標値は今後検討）、本計画の第3章3-3・第10章10-4の指標と関わります。"
        "認知症基本法第13条第1項により、本計画の第6章（市町村認知症施策推進計画）は、国の"
        "認知症施策推進基本計画に加え、県の認知症施策推進計画が策定されたときはこれを基本とします。"
        "県の中間案が示された段階で、本計画の目標・見込量・認知症施策との整合を確かめます。")

KYU_8_2 = "第10期の区分は国の基本指針の確定を待って整理します（確認事項No.69）。"
SHIN_8_2 = (KYU_8_2 + "なお、宮城県は第7期宮城県介護給付適正化取組方針（令和9年度〜令和11年度）"
            "について、国の基本方針に大きな見直しはないとして、第6期の重点の取組（要介護認定の"
            "適正化、ケアプラン・住宅改修等の点検、縦覧点検・医療情報との突合）を継続する方向を"
            "示しています（令和8年8月26日 みやぎ高齢者元気プラン推進委員会）。")
CHIIKI = ("⚠ 地域包括ケア「見える化」システムに入力する地域支援事業費は、総事業費から寄付金"
          "その他の収入額を差し引いた額のうち、交付金の対象経費として支出する見込みの額"
          "（対象経費支出予定額）とされています（厚生労働省「地域包括ケア『見える化』システム等を"
          "活用した地域分析の手引き」令和8年9月改訂版）。決算額とは範囲が異なり得るため、"
          "システムの入力値と決算を比べるときは、この違いを踏まえます。")


def kakikae(p, s):
    if p.runs:
        p.runs[0].text = s
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(s)


def zenbun(d):
    return "\n".join([p.text for p in d.paragraphs]
                     + [c.text for t in d.tables for r in t.rows for c in r.cells])


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("協議用素案 Ver.2.28 → Ver.2.29（宮城県の第10期計画の骨子（案）との対応）")
    print()
    for f in (MOTO, GENTEN, TEBIKI):
        if not os.path.exists(f):
            print(f"   × 入力がありません：{f}")
            sys.exit(1)

    print("  ── ⓪ 原典の語を確かめる")
    import pymupdf
    gt = "".join(p.get_text() for p in pymupdf.open(GENTEN)).replace("\n", "")
    for go in ("高齢者が地域で自分らしい生活を安心して送れる社会", "仙南圏域（２市７町）",
               "二次医療圏と同じ４圏域", "第１期宮城県認知症施策推進計画",
               "国の基本方針に大きな見直しなし", "インセンティブ交付金県内市町村平均得点率",
               "在宅等での看取り数", "中間案"):
        judge(go in gt, f"県の資料に「{go}」")
    tb = "".join(p.get_text() for p in pymupdf.open(TEBIKI)).replace("\n", "")
    judge("対象経費支出予定額" in tb, "手引きに「対象経費支出予定額」")

    shutil.copy2(MOTO, SAKI)
    d = docx.Document(SAKI)
    mae = zenbun(d)

    def p_hit(atama):
        h = [p for p in d.paragraphs if p.text.strip().startswith(atama.strip())]
        return h[0] if len(h) == 1 else None

    print("  ── ① 1-2 表1「県の計画」の行")
    t1 = [t for t in d.tables
          if [c.text.strip() for c in t.rows[0].cells] == ["区分", "計画名", "整合の方針"]]
    judge(len(t1) == 1, "表1 を特定")
    gyo = [r for r in t1[0].rows if r.cells[0].text.strip() == "県の計画"] if t1 else []
    judge(len(gyo) == 1 and gyo[0].cells[1].text.strip()
          == "宮城県高齢者保健福祉計画・介護保険事業支援計画", "旧の計画名がある")
    if gyo:
        kakikae(gyo[0].cells[1].paragraphs[0], KEN_MEI)
        kakikae(gyo[0].cells[2].paragraphs[0],
                "高齢者福祉圏域（仙南圏域）を単位とする施設整備・人材確保の調整。"
                "骨子（案）との対応は下表")

    print("  ── ② 1-2 に骨子（案）との対応を置く（1-3 の前）")
    mid_tpl = p_hit("　▌ 上位計画・関連計画との関係")
    hon_tpl = p_hit("本計画は、以下の法令・上位計画・関連計画との整合を図りながら策定します。")
    sh_tpl = [p for p in d.paragraphs if p.text.startswith("出典：")][0]
    p13 = p_hit("1-3　計画の期間")
    judge(None not in (mid_tpl, hon_tpl, p13), "雛形の段落と 1-3 を特定")
    chui_tpl = [p for p in d.paragraphs if p.text.startswith("⚠ ")
                and not p.runs[0].bold][0] if True else None
    if t1 and None not in (mid_tpl, hon_tpl, p13):
        anchor = p13._p
        for tpl, naiyo in ((mid_tpl, MIDASHI), (hon_tpl, HON), (t1[0], HYO),
                           (sh_tpl, SHUTTEN), (chui_tpl, CHUI)):
            if isinstance(naiyo, str):
                el = copy.deepcopy(tpl._p)
                anchor.addprevious(el)
                kakikae(dp.Paragraph(el, hon_tpl._parent), naiyo)
                continue
            el = copy.deepcopy(tpl._tbl)
            anchor.addprevious(el)
            trs = el.findall(qn("w:tr"))
            for tr in trs[len(naiyo):]:
                el.remove(tr)
            while len(el.findall(qn("w:tr"))) < len(naiyo):
                el.append(copy.deepcopy(el.findall(qn("w:tr"))[-1]))
            grid = el.find(qn("w:tblGrid"))
            cols = grid.findall(qn("w:gridCol"))
            zentai = sum(int(c.get(qn("w:w"))) for c in cols)
            haba = [int(zentai * x / sum(HABA)) for x in HABA]
            haba[-1] += zentai - sum(haba)
            for c, w in zip(cols, haba):
                c.set(qn("w:w"), str(w))
            for tr in el.findall(qn("w:tr")):
                for tc, w in zip(tr.findall(qn("w:tc")), haba):
                    tc.find(qn("w:tcPr")).find(qn("w:tcW")).set(qn("w:w"), str(w))
            t = dt.Table(el, tpl._parent)
            for ri, g in enumerate(naiyo):
                for ci, v in enumerate(g):
                    cell = t.rows[ri].cells[ci]
                    kakikae(cell.paragraphs[0], v)
                    for q in cell.paragraphs[1:]:
                        q._p.getparent().remove(q._p)
        judge(True, "小見出し・本文・対応表（9行）・出典・注記を置いた")

    print("  ── ③ 8-2 注記（第7期宮城県介護給付適正化取組方針）")
    p = [x for x in d.paragraphs if KYU_8_2 in x.text]
    judge(len(p) == 1, "8-2 の旧の文がある")
    if p:
        kakikae(p[0], p[0].text.replace(KYU_8_2, SHIN_8_2))

    print("  ── ④ 9-1 地域支援事業費の見込み（対象経費支出予定額）")
    p = p_hit("介護予防・日常生活支援総合事業の量は、地域包括ケア「見える化」システムに入力")
    judge(p is not None, "地域支援事業費の見込みの段落を特定")
    if p is not None:
        el = copy.deepcopy(chui_tpl._p)
        p._p.addnext(el)
        kakikae(dp.Paragraph(el, p._parent), CHIIKI)

    d.save(SAKI)
    d2 = docx.Document(SAKI)
    ato = zenbun(d2)

    print("  ── 検査1　古い計画名が残っていない")
    judge("宮城県高齢者保健福祉計画" not in ato, "「宮城県高齢者保健福祉計画」0件")
    print("  ── 検査2　対応表の箇所が素案に現にある章節")
    ari = set(re.findall(r"^(\d+-\d+)　", "\n".join(p.text for p in d2.paragraphs), re.M))
    yobu = set()
    for g in HYO[1:]:
        for a, b in re.findall(r"(\d+)-(\d+)", g[2]):
            yobu.add(f"{a}-{b}")
    nai = sorted(yobu - ari)
    judge(not nai, f"引いた章節 {len(yobu)}件がすべて素案にある", "／".join(nai))
    for sho in ("第6章",):
        judge(any(p.text.startswith("6-1") for p in d2.paragraphs), f"{sho}がある")
    print("  ── 検査3　県の資料にある語だけで書いている")
    for g in HYO[1:]:
        kadai = g[1].split("（")[0]
        judge(kadai.replace("・", "") in gt.replace("･", "").replace("・", ""),
              f"基本課題「{kadai}」が骨子（案）にある")
    print("  ── 検査4　他団体名・受託者の語")
    for w in ("北塩原", "金ヶ崎", "大雪", "川崎市", "受託者", "当社", "ビズアップ", "スクリプト"):
        judge(ato.count(w) == mae.count(w), f"「{w}」が増えていない")
    print("  ── 検査5　未確定の印")
    for w in ("［要確認］", "［要協議］", "［要内訳］"):
        judge(ato.count(w) == mae.count(w), f"{w} {ato.count(w)}件（増減なし）")
    print("  ── 検査6　⚠ 既にあった数が消えていない")
    kieta = sorted({m.group(0) for m in re.finditer(
        r"[0-9０-９][0-9０-９,.]*(?:人|％|千円|件|頁|点|円)", mae)
        if mae.count(m.group(0)) > ato.count(m.group(0))})
    judge(not kieta, "消えた数 0件", "／".join(kieta[:5]))

    print()
    print(f"  {SAKI}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜6　不適合0件")


if __name__ == "__main__":
    main()
