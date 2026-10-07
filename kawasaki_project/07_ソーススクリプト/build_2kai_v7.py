# -*- coding: utf-8 -*-
"""第2回策定委員会資料 v6 → v7（令和8年10月4日）

v6 のあとに分かったことを反映する。

  ① 工程が納期に収まらない　逆算工程表（R8.10.1）の結果を「８．今後の予定」に
     反映し、**開催日と短縮の案のご選択を本日の審議事項に加える。**
     v6 の予定表は「令和9年1〜2月にパブリックコメント」としているが、
     逆算すると11月11日以降の開催では納期（令和9年3月15日）に間に合わない。
  ② 第9期の計画値と実績値の対比（素案 3-3 の図3-2）を「７．」に加える。
     交付金 推進 目標Ⅰ（ⅱ）2「事業計画の進捗状況」（12点・本町0点）が
     求める分析であり、本日ご報告すること自体が要件に当たる。
  ③ 交付金の得点の推移（素案 3-3 の図3-1）を「７．」に差し込む。
  ④ 当社資料の訂正に1件を加える（交付金の令和6年度の得点の食い違い）。
  ⑤ 計画書の頁数の見通しを「８．」に加える（事業内容の欄・圧縮案）。

  python3 07_ソーススクリプト/build_2kai_v7.py

自己点検（check_2kai_selfcheck_R8.9.28.py）に通ることを前提とする。
件数を宣言している箇所と実際の件数を必ず合わせる。
"""
import copy
import os
import re
import shutil
import struct
import sys

import docx
from docx.oxml.ns import qn
from docx.shared import Emu

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v6.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v7.docx"
FIGDIR = "08_図表"
EMU_IN = 914400

# ══════════════════════════ ④ 当社資料の訂正に加える1件
TEISEI7 = ("7",
           "交付金の令和６年度の得点（当社の取りまとめ表の集計が公表値と"
           "食い違っていた）",
           "枝番の合計 428点（体制・取組指標群 340点）",
           "公表値 454点（体制・取組指標群 366点）")

# ══════════════════════════ ② 第9期の計画値と実績値の対比
TAIHI_HEAD = ["区　分", "令和７年度 計画値", "令和７年度 実績値", "達成率"]
TAIHI_ROWS = [
    ["居宅サービス（介護予防を含む）", "353,651千円", "322,680千円", "91.2％"],
    ["地域密着型サービス（介護予防を含む）", "179,884千円", "182,190千円",
     "101.3％"],
    ["施設サービス", "477,485千円", "493,350千円", "103.3％"],
    ["総給付費", "1,011,020千円", "998,220千円", "98.7％"],
]
TAIHI_P1 = ("令和７年度の給付費は、計画10億1,102万円に対し実績９億9,822万円"
            "（達成率98.7％）でした。総額では計画にほぼ沿っていますが、"
            "区分別にみると居宅サービスが計画を8.8％下回る一方、"
            "施設サービスは3.3％上回っています。")
TAIHI_P2 = ("総額が一致していることの内側で、在宅から施設・居住系への移行が"
            "進んでいます。居宅サービスが計画を下回ったのは、町内事業所の縮小"
            "（訪問介護の登録21名に対し実働４名）と、住所地特例24人"
            "（令和８年３月末）に表れる町外施設の利用の常態化が同時に進んだ"
            "結果と考えられます。第１０期のサービス見込量は、"
            "この移行が続くことを前提としています。")
TAIHI_BOX = ("この対比は、保険者機能強化推進交付金 推進 目標Ⅰ（ⅱ）２"
             "「事業計画の進捗状況」（配点12点・本町は令和８年度０点）が"
             "求める分析に当たります。本日ご報告し、ご意見をいただくことが、"
             "そのまま要件を満たすことになります。")

# ══════════════════════════ ① 今後の予定（逆算の結果）
YOTEI_HEAD = ["時　期", "内　容"]
YOTEI_ROWS = [
    ["令和８年11月【本日】", "第２回：計画骨子（案）とご選択いただきたい事項"],
    ["令和８年11月", "選択１のご判断を受けた施策体系の確定"
                     "（統合体系表は作成済み。計画書への書き下ろしは完了）"],
    ["令和８年12月", "計画書（案）の作成。宮城県への事前協議の開始"],
    ["令和８年12月〜令和９年１月", "パブリックコメントの実施"],
    ["令和９年１月", "第３回：保険料基準額と所得段階別保険料の確定"],
    ["令和９年２月", "第４回：計画書（案）・概要版（案）の確定"
                     "（仕様書上の３回には含まれない任意開催）"],
    ["令和９年２月下旬", "成果品の入稿"],
    ["令和９年３月15日", "計画書・概要版の納品（仕様書２・８）"],
]
KURI_P = ("⚠ 開催日から納品までに要する日数を逆算したところ、"
          "**現行の前提では納期に収まりません。**"
          "前提は①パブリックコメント30日間 ②策定委員会 全４回 "
          "③印刷・製本３週間 ④資料の事前配布は開催日の２週間前、です。")
#   列は4つまで（本資料に5列の表の雛形がないため）。
#   案Ｃ（委員会を全３回）は囲みで触れる。
KURI_HEAD = ["第２回の開催日", "案Ａ 現行の前提", "案Ｂ パブコメ20日間",
             "案Ｄ Ｂ＋委員会を全３回"]
KURI_ROWS = [
    ["令和８年11月４日（水）", "±０日", "＋10日", "＋35日"],
    ["令和８年11月11日（水）", "▲７日", "＋３日", "＋28日"],
    ["令和８年11月18日（水）", "▲14日", "▲４日", "＋21日"],
    ["令和８年11月25日（水）", "▲21日", "▲11日", "＋14日"],
]
KURI_BOX = ("事務局案：案Ｂ（パブリックコメントを20日間）。"
            "11月18日までの開催であれば、案Ｃ（委員会を全３回・"
            "第３回で保険料の決定と答申を併せて行う）との組合せ（案Ｄ）で"
            "余裕をもって納期に収まります。"
            "案Ｃのみでも、11月18日の開催で＋11日の余裕となります。"
            "なお本表は、宮城県の意見の聴取に要する期間と、"
            "令和９年３月の町議会に介護保険条例の改正案を諮る日程を"
            "含んでいません。")

# ══════════════════════════ ⑤ 頁数の見通し
PAGE_HEAD = ["版", "本文", "計（資料編13頁を含む）",
             "仕様書８（100頁程度）との差"]
PAGE_ROWS = [
    ["Ver.2.8（事業の内容の欄なし）", "104頁", "117頁", "＋17頁"],
    ["Ver.2.9（事業の内容の欄あり）", "109頁", "122頁", "＋22頁"],
    ["Ver.2.9c（内容の欄あり＋圧縮案）", "102頁", "115頁", "＋15頁"],
]
PAGE_P = ("計画書の事業一覧表（17表・138事業）に「内容」の欄を加えました。"
          "事業の名称だけの一覧では、何をする事業なのかが分からないためです。"
          "欄を加えると５頁増えますが、詳細な分析を別冊に移す圧縮案と"
          "併せると本文102頁となり、欄がなかったときより２頁少なくなります。")
PAGE_BOX = ("仕様書８の「100頁程度」を本文のみと読むのであれば、"
            "圧縮案の本文102頁で収まります。"
            "資料編13頁を含めるかどうかを含めてご判断をお願いします。")

SHINGI_NEW = ("⑦　第２回の開催日と、工程の短縮の案（Ａ〜Ｄ）のご選択（８）。"
              "⚠ 納期に直接に効きます")


def txt(el):
    return "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()


def png_size(b):
    return struct.unpack(">II", b[16:24])


def main():
    if not os.path.exists(SRC):
        raise SystemExit("v6 が見つからない：" + SRC)
    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body
    from fix_soan_v111 import set_el, set_cell, clone_table

    kids = list(body.iterchildren())

    # 雛形
    tpl_p = tpl_box = tpl_head = None
    for el in kids:
        if el.tag != qn("w:p"):
            continue
        s = txt(el)
        if tpl_head is None and s.startswith("◆ "):
            tpl_head = copy.deepcopy(el)
        if tpl_p is None and s.startswith("第１回で最も議論が集中した論点"):
            tpl_p = copy.deepcopy(el)
    for el in kids:
        if el.tag == qn("w:tbl"):
            t = docx.table.Table(el, doc)
            if len(t.rows) == 1 and len(t.columns) == 1:
                tpl_box = copy.deepcopy(el)
                break
    if tpl_p is None or tpl_box is None or tpl_head is None:
        raise SystemExit("雛形が見つからない")

    def P(text, after):
        el = copy.deepcopy(tpl_p)
        set_el(el, text)
        after.addnext(el)
        return el

    def H(text, after):
        el = copy.deepcopy(tpl_head)
        set_el(el, text)
        after.addnext(el)
        return el

    def BOX(text, after):
        el = copy.deepcopy(tpl_box)
        t = docx.table.Table(el, doc)
        set_cell(t.rows[0].cells[0], text.replace("**", ""))
        after.addnext(el)
        return el

    def TBL(head, rows, after):
        """列数の合う雛形を探して複製する。"""
        n = len(head)
        tpl = None
        for el in body.iterchildren():
            if el.tag != qn("w:tbl"):
                continue
            t = docx.table.Table(el, doc)
            if len(t.columns) == n and len(t.rows) >= len(rows) + 1:
                tpl = el
                break
        if tpl is None:
            raise SystemExit(f"{n}列の雛形が見つからない")
        new = clone_table(tpl, [head] + [list(r) for r in rows])
        after.addnext(new)
        return new

    def IMG(png, after, inch=5.4):
        path = os.path.join(FIGDIR, png)
        b = open(path, "rb").read()
        w, h = png_size(b)
        cx = int(inch * EMU_IN)
        doc.add_picture(path, width=Emu(cx), height=Emu(int(cx * h / w)))
        p = doc.paragraphs[-1]._p
        p.getparent().remove(p)
        pPr = tpl_p.find(qn("w:pPr"))
        if pPr is not None:
            old = p.find(qn("w:pPr"))
            if old is not None:
                p.remove(old)
            p.insert(0, copy.deepcopy(pPr))
        tr = tpl_p.find(qn("w:r"))
        rPr = tr.find(qn("w:rPr")) if tr is not None else None
        if rPr is not None:
            for r in p.findall(qn("w:r")):
                if r.find(qn("w:rPr")) is None:
                    r.insert(0, copy.deepcopy(rPr))
        after.addnext(p)
        return p

    def find_p(pred):
        for el in body.iterchildren():
            if el.tag == qn("w:p") and pred(txt(el)):
                return el
        raise SystemExit("段落が見つからない")

    def find_tbl(pred):
        for el in body.iterchildren():
            if el.tag == qn("w:tbl"):
                t = docx.table.Table(el, doc)
                if pred(t):
                    return el, t
        raise SystemExit("表が見つからない")

    # ══════════════════════════ ④ 当社資料の訂正に1件を加える
    el, t = find_tbl(lambda t: len(t.columns) == 4 and len(t.rows) == 7
                     and t.rows[0].cells[1].text.strip() == "訂正の内容")
    data = [[c.text for c in r.cells] for r in t.rows] + [list(TEISEI7)]
    new = clone_table(el, data)
    el.addprevious(new)
    el.getparent().remove(el)
    h = find_p(lambda s: s.startswith("◆ 当社資料の訂正"))
    set_el(h, "◆ 当社資料の訂正（保険料試算の再検算ほかによる７件）")
    p = find_p(lambda s: s.startswith("令和８年９月28日に町へお渡しした"))
    set_el(p, txt(p).replace("次の６件を是正しました",
                             "次の７件を是正しました"))

    # ══════════════════════════ ③ 交付金の得点の推移の図
    anchor = find_p(lambda s: s.startswith("800点は、仕組みが整っているか"))
    IMG("fig3-1_kofukin.png", anchor, 5.4)

    # ══════════════════════════ ② 第9期の計画値と実績値の対比
    #   「◆ 本町の強み」の手前に置く
    anchor = find_p(lambda s: s.startswith("◆ 本町の強み"))
    prev = anchor.getprevious()
    h = H("◆ 第９期の計画値と実績値の対比（事業計画の進捗状況）", prev)
    p1 = P(TAIHI_P1, h)
    img = IMG("fig3-2_keikaku_jisseki.png", p1, 5.4)
    tb = TBL(TAIHI_HEAD, TAIHI_ROWS, img)
    p2 = P(TAIHI_P2, tb)
    BOX(TAIHI_BOX, p2)

    # ══════════════════════════ ① 今後の予定の差し替えと短縮の案
    el, t = find_tbl(lambda t: len(t.columns) == 2 and len(t.rows) == 8
                     and t.rows[0].cells[0].text.strip() == "時　期")
    new = clone_table(el, [YOTEI_HEAD] + [list(r) for r in YOTEI_ROWS])
    el.addprevious(new)
    el.getparent().remove(el)
    # 直後の囲み（仕様書上3回構成）の後ろに、逆算の結果を置く
    box = new.getnext()
    p = P(KURI_P.replace("**", ""), box)
    tb = TBL(KURI_HEAD, KURI_ROWS, p)
    BOX(KURI_BOX, tb)

    # ══════════════════════════ ⑤ 頁数の見通し
    anchor = find_p(lambda s: s.startswith("◆ 今後の予定"))
    prev = anchor.getprevious()
    h = H("◆ 計画書の頁数の見通し", prev)
    p = P(PAGE_P, h)
    tb = TBL(PAGE_HEAD, PAGE_ROWS, p)
    BOX(PAGE_BOX, tb)

    # ══════════════════════════ 本日のご審議事項に⑦を足す
    last = find_p(lambda s: s.startswith("⑥　交付金の評価結果のご報告"))
    P(SHINGI_NEW, last)
    # 次第にも足す
    ji = find_p(lambda s: s.startswith("⑦　計画骨子（案）と今後の予定"))
    set_el(ji, "⑦　計画骨子（案）・頁数の見通し・今後の予定")

    # ══════════════════════════ 版表記
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = "".join(n.text or "" for n in el.iter(qn("w:t")))
        if "第６版" in s or "v6" in s:
            set_el(el, s.replace("第６版", "第７版").replace("v6", "v7"))

    doc.save(DST)

    d2 = docx.Document(DST)
    import zipfile
    z = zipfile.ZipFile(DST)
    print("第2回策定委員会資料 v7")
    print("  保存：", DST)
    print(f"  段落 {len(d2.paragraphs)}／表 {len(d2.tables)}"
          f"／画像 {sum(1 for n in z.namelist() if n.startswith('word/media'))}")
    # ══════════════════════════ 自己点検
    text = "\n".join(q.text for q in d2.paragraphs)
    for t in d2.tables:
        for r in t.rows:
            for c in r.cells:
                text += "\n" + c.text
    ng = []
    for v in ("998,220", "1,011,020", "98.7", "454点", "528点"):
        if v not in text:
            ng.append(f"主な数値が入っていない：{v}")
    if "７件" not in text:
        ng.append("当社資料の訂正の件数が７件になっていない")
    print("  ── 自己点検")
    if ng:
        for m in ng:
            print("   ×", m)
        sys.exit(1)
    print("   ○ 加えた数値と件数が資料に入っている")


if __name__ == "__main__":
    main()
