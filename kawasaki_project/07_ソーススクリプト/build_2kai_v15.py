# -*- coding: utf-8 -*-
"""第2回策定委員会資料 v14 → v15（協議用素案 Ver.2.29 の頁数に揃える）
（令和8年10月23日）

  python3 07_ソーススクリプト/build_2kai_v15.py

協議用素案を Ver.2.29 に上げ（fix_soan_v2290_ken.py。1-2 に宮城県の第10期計画の
骨子（案）との対応表を加えた）、総頁数が117頁 → 118頁になった。
８の頁数（本文の段落と頁数の見通しの表）を素案の実測に揃える。
ほかの内容は素案 Ver.2.29 の変更箇所（1-2・8-2・9-1 の注記）を載せていないため改めない。
"""
import importlib.util
import os
import shutil
import sys

import docx

HERE = os.path.dirname(os.path.abspath(__file__))


def load(name, fn):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(spec)
    sys.modules[name] = m
    spec.loader.exec_module(m)
    return m


V9 = load("build_2kai_v9", "build_2kai_v9.py")
SRC = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v14.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v15.docx"
KYU, HEN = 117, 118
SHIRYO = 13


def main():
    futeki = []

    def judge(ok, mi):
        print(("   ○ " if ok else "   × ") + mi)
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会資料 v14 → v15（頁数を協議用素案 Ver.2.29 に揃える）")
    shutil.copy2(SRC, DST)
    d = docx.Document(DST)
    n = 0
    for p in d.paragraphs:
        if f"現在の計画素案は本文{KYU}頁" in p.text:
            p.runs[0].text = (p.text.replace(f"本文{KYU}頁", f"本文{HEN}頁")
                              .replace(f"計{KYU + SHIRYO}頁", f"計{HEN + SHIRYO}頁")
                              .replace(f"{KYU + SHIRYO - 100}頁上回ります",
                                       f"{HEN + SHIRYO - 100}頁上回ります"))
            for r in p.runs[1:]:
                r.text = ""
            n += 1
    for t in d.tables:
        for r in t.rows:
            if r.cells[0].text.strip() == "現在の計画素案" and r.cells[1].text.strip() == f"{KYU}頁":
                V9.set_cell(r.cells[1], f"{HEN}頁")
                V9.set_cell(r.cells[2], f"{HEN + SHIRYO}頁")
                V9.set_cell(r.cells[3], f"＋{HEN + SHIRYO - 100}頁")
                n += 1
    judge(n == 2, f"本文と表の頁数を {HEN}頁に（{n}か所）")
    d.save(DST)

    d2 = docx.Document(DST)
    ato = "\n".join([p.text for p in d2.paragraphs]
                    + [c.text for t in d2.tables for r in t.rows for c in r.cells])
    judge(f"{KYU}頁" not in ato and f"本文{HEN}頁" in ato and f"計{HEN + SHIRYO}頁" in ato,
          f"頁数 本文{HEN}頁・計{HEN + SHIRYO}頁（{KYU}頁 0件）")
    for na in ("釜房", "アルパイン", "鶴寿会", "北塩原", "金ヶ崎", "大雪", "川崎市"):
        judge(na not in ato, f"「{na}」0件")
    judge("⚠" not in ato and "**" not in ato, "⚠・** 0件（第1回の書式）")
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 不適合0件")


if __name__ == "__main__":
    main()
