# -*- coding: utf-8 -*-
"""業務工程管理表の更新（令和8年9月30日・大雪広域との対比）

  03_確認事項一覧　No.151〜160 を追加（全160件）
  00・01・02・04シート　素案Ver.2.4の状況

根拠：`05_試算・管理シート/川崎町_大雪広域_素案対比表_R8.9.30.xlsx` 07シート
"""
import copy
import sys

import openpyxl

sys.path.insert(0, "07_ソーススクリプト")

WB = "川崎町_業務工程管理表.xlsx"
TAIHI = "05_試算・管理シート/川崎町_大雪広域_素案対比表_R8.9.30.xlsx"


def unmerge_below(ws, at_row):
    kept = []
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= at_row:
            kept.append((m.min_row, m.max_row, m.min_col, m.max_col))
            ws.unmerge_cells(str(m))
    return kept


def remerge(ws, kept, n):
    for r1, r2, c1, c2 in kept:
        ws.merge_cells(start_row=r1 + n, end_row=r2 + n,
                       start_column=c1, end_column=c2)


def copy_style(ws, src_row, dst_row, ncol):
    for c in range(1, ncol + 1):
        ws.cell(dst_row, c)._style = copy.copy(ws.cell(src_row, c)._style)
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def insert(ws, at, tpl, recs, ncol, count_col, old_last):
    n = len(recs)
    kept = unmerge_below(ws, at)
    ws.insert_rows(at, n)
    remerge(ws, kept, n)
    for i, rec in enumerate(recs):
        r = at + i
        copy_style(ws, tpl, r, ncol)
        for c, v in enumerate(rec, start=1):
            ws.cell(r, c).value = v
    new_last = old_last + n
    for r in range(at + n, ws.max_row + 1):
        cell = ws.cell(r, 2)
        if isinstance(cell.value, str) and cell.value.startswith("=COUNTIF"):
            cell.value = cell.value.replace(
                f"{count_col}5:{count_col}{old_last}",
                f"{count_col}5:{count_col}{new_last}")
    return new_last


# 対比表の07シートから確認事項を読む（二重管理を避ける）
tw = openpyxl.load_workbook(TAIHI, data_only=True)["07_確認事項"]
KOUTEI = {  # No → (業務内容, 工程（月）)
    "151": ("(5)(7)", "R8.10"), "152": ("(6)(7)", "R8.10"),
    "153": ("(6)(7)", "R8.11"), "154": ("(4)(6)", "R8.11"),
    "155": ("(4)(6)", "R8.11"), "156": ("(4)(6)", "R8.11"),
    "157": ("(4)(7)", "R8.10"), "158": ("(4)(6)", "R8.11"),
    "159": ("(4)(5)", "R8.10"), "160": ("(4)(7)", "R8.11"),
}
BIKOU = {
    "151": "大雪地区広域連合の協議用素案（令和8年8月）との対比により把握した。"
           "素案Ver.2.4の1-5に制度改正の節を新設し、当該の改正を掲げた。",
    "152": "川崎町は過疎地域に指定されており、指定基準①又は③に該当する可能性がある。"
           "素案Ver.2.4の1-5・7-3・9-1に位置付けた。",
    "153": "介護保険最新情報Vol.1525（令和8年7月14日）による。"
           "素案Ver.2.4の1-5に医療計画等との整合の項を置いた。",
    "154": "第10期基本指針（案）の別表 八。"
           "素案に有料老人ホーム・サービス付き高齢者向け住宅が一度も現れていなかった。",
    "155": "第10期基本指針（案）の別表 四・五。"
           "見える化システムのIDパスは取得しているが直接データを触る想定ではないため、"
           "町又は宮城県による出力をお願いしたい。",
    "156": "第10期基本指針（案）の別表 三。",
    "157": "第10期基本指針（案）の別表 十一。"
           "素案に認知症の人の数がなかった。",
    "158": "第10期基本指針（案）の別表 十。"
           "交付金 支援 目標Ⅲで失点している項目でもある。",
    "159": "第10期基本指針（案）の別表 七。"
           "第2回策定委員会の選択2（施設整備）の判断の根拠となる。"
           "確認事項No.29（事業所一覧）と併せてご確認いただきたい。",
    "160": "社会福祉法等の一部を改正する法律（令和8年法律第51号）による。"
           "介護予防ケアマネジメントの直接委託は地域包括支援センターの"
           "業務負担の軽減に直結する。",
}
NEW = []
for r in range(2, tw.max_row + 1):
    no = str(tw.cell(r, 1).value or "").strip()
    if not no.isdigit():
        continue
    gyomu, tsuki = KOUTEI[no]
    NEW.append((no, gyomu, tsuki,
                str(tw.cell(r, 2).value), str(tw.cell(r, 3).value),
                str(tw.cell(r, 4).value), "発注者", "確認待ち",
                str(tw.cell(r, 5).value), BIKOU[no]))
if len(NEW) != 10:
    raise SystemExit(f"確認事項が{len(NEW)}件しか読めない")

wb = openpyxl.load_workbook(WB)
ws = wb["03_確認事項一覧"]
LAST = 154          # No.150 の行
if str(ws.cell(LAST, 1).value).strip() != "150":
    raise SystemExit(f"No.150 が{LAST}行にない（{ws.cell(LAST, 1).value}）")
LAST = insert(ws, LAST + 1, LAST, NEW, 10, "H", LAST)
print("03_確認事項一覧　No.160 まで（最終行", LAST, "）")

# ══════════════════════════ 00_管理方法
ws = wb["00_管理方法"]
ws.cell(14, 2).value = (
    "令和8年9月30日　本表の記載は同日時点のもの。"
    "計画素案 Ver.2.4（制度改正・日常生活圏域 反映版・649段落125表・全104頁）、"
    "第2回策定委員会資料 v6、想定問答集 v3（全39問）、"
    "施策・事業統合体系表（7章17節37項143事業）、"
    "交付金取りまとめ表（7シート）、計画書別添整理表（5シート）、"
    "認知症の意見聴取記録様式（9シート・様式0〜5）、"
    "認知症の計画記載事項 適合性確認表（6シート）、"
    "大雪広域との素案対比表（8シート）まで作成済み。"
    "点検は、素案の自己点検32件・保険料の検算39項目・"
    "委員会資料の自己点検27件がいずれも適合。")

# ══════════════════════════ 01_業務内容別の進捗
ws = wb["01_業務内容別の進捗"]
ws.cell(11, 5).value = (
    "計画書素案_v2.4_制度改正反映版（649段落125表・全104頁）\n"
    "施策・事業統合体系表_R8.9.25b（6シート・7章17節37項143事業）\n"
    "認知症_意見聴取記録様式_R8.9.30（9シート・様式0〜5・別添22）\n"
    "認知症_計画記載事項_適合性確認表_R8.9.30（6シート）\n"
    "大雪広域_素案対比表_R8.9.30（8シート）\n"
    "交付金_取りまとめ表_R8.9.29（7シート）\n"
    "計画書_別添整理表_R8.9.29（5シート）")

# ══════════════════════════ 02_月別工程管理
ws = wb["02_月別工程管理"]
ws.cell(9, 5).value = (
    "人口推計・事業量推計は、令和7年度の年報を基礎とする独立算定により第1稿を作成。"
    "地域支援事業の費用の見込みはデータ未受領のため保留。"
    "素案はVer.2.4（制度改正・日常生活圏域 反映版・全104頁）まで更新した。"
    "大雪地区広域連合の協議用素案（令和8年8月）と対比し、"
    "介護保険制度改正の主な内容（1-5）・日常生活圏域の設定（1-6）・"
    "24時間対応サービスの確保方策及び介護施設整備に係る基本方針（9-1）・"
    "低所得者支援（9-3）を新設した。"
    "認知症の人及び家族等からの意見聴取は令和8年10〜11月に実施する予定であり、"
    "記録様式（別添22・様式0〜5）を作成済み。")

# ══════════════════════════ 04_成果品管理
ws = wb["04_成果品管理"]
ws.cell(8, 3).value = "川崎町_計画書素案_v2.4_制度改正反映版.docx"
ws.cell(10, 3).value = (
    "川崎町_第2回策定委員会資料_R8.11_v6.docx／"
    "同_想定問答集_R8.11_v3.docx\n"
    "川崎町_第10期_施策事業統合体系表_R8.9.25b.xlsx\n"
    "川崎町_交付金_取りまとめ表_R8.9.29.xlsx\n"
    "川崎町_計画書_別添整理表_R8.9.29.xlsx\n"
    "川崎町_認知症_意見聴取記録様式_R8.9.30.xlsx\n"
    "川崎町_認知症_計画記載事項_適合性確認表_R8.9.30.xlsx\n"
    "川崎町_大雪広域_素案対比表_R8.9.30.xlsx")

wb.save(WB)
print("保存：", WB)
wb2 = openpyxl.load_workbook(WB)
ws2 = wb2["03_確認事項一覧"]
n = sum(1 for r in range(5, ws2.max_row + 1)
        if str(ws2.cell(r, 1).value or "").strip().isdigit())
print("  確認事項", n, "件")
