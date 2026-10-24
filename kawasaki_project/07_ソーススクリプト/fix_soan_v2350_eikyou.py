# -*- coding: utf-8 -*-
"""協議用素案 Ver.2.34 → Ver.2.35（Red Team 再レビュー第3回の指摘の影響範囲の再点検による是正）

  python3 07_ソーススクリプト/fix_soan_v2350_eikyou.py

第3回レビュー（09_元資料/レビュー/川崎町_第10期計画_RedTeam再レビュー第3回_R8.10.24受領.md）を同じ版で
受け直し、指摘の語だけでなく同じ趣旨の記述を5冊で走査した。素案に残っていたもの：

  R3-06 の波及　9-3 の「本算定」（5か所）。保険料の「本算定」は確定の算定と読めるため「本試算」に改める
  R3-08 の波及　KPI 以外の箇所の「通いの場への参加率」。15.9％と同じ指標を指す箇所は「参加割合」に揃え、
               調査の回答者のうちの割合であることを書く
  1-2 の留保　　入所・居住系の1人1月あたり給付費「約29.9万円」の出所を書く
               （令和7年度の入所・居住系の給付費680,226千円 ÷ 年報の受給者数 延べ2,274人月＝299.13千円。
                この値でも10人で211.84円となり、約212円は変わらない）

終了コード　0＝検査すべて適合／1＝1件でも不適合
"""
import re
import shutil
import sys

import docx
import openpyxl

MOTO = "01_第10期_最新版成果品/川崎町_計画書素案_v2.34_協議用.docx"
SAKI = "01_第10期_最新版成果品/川崎町_計画書素案_v2.35_協議用.docx"
NENPO = "09_元資料/R7実績データ/年報データ_2025_川崎町.xlsx"


def kakikae(p, s):
    p.runs[0].text = s
    for r in p.runs[1:]:
        r.text = ""


def zenbun(d):
    return "\n".join([p.text for p in d.paragraphs]
                     + [c.text for t in d.tables for r in t.rows for c in r.cells])


def subete(d):
    seen = set()
    for p in list(d.paragraphs) + [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs]:
        if id(p._p) in seen:
            continue
        seen.add(id(p._p))
        yield p


def jukyu_nengetsu():
    """年報の受給者数（年度の延べ人月）：施設・認知症対応型共同生活介護・地域密着型介護老人福祉施設・特定施設。"""
    wb = openpyxl.load_workbook(NENPO, data_only=True)

    def gyo(sn, na):
        for r in wb[sn].iter_rows(values_only=True):
            v = [x for x in r if x is not None]
            if v and str(v[0]).strip().startswith(na):
                return v[-1]
        raise SystemExit(f"年報 {sn} に {na} がない")
    shisetsu = None
    for r in wb["様式１の６"].iter_rows(values_only=True):
        v = [x for x in r if x is not None]
        if v and str(v[0]).replace("　", "").strip() == "総数" and len(v) == 11 and v[3] is not None:
            shisetsu = v[-1]   # (15) 施設介護サービス受給者数の総数（最後に出る総数）
    gh = gyo("様式１の７（１８）地域密着型", "認知症対応型共同生活介護（短期利用以外）")
    chimitsu = gyo("様式１の７（１８）地域密着型", "地域密着型介護老人福祉施設入所者生活介護")
    tokutei = gyo("様式１の７（１６）居宅介護", "特定施設入居者生活介護（短期利用以外）")
    return shisetsu, gh, chimitsu, tokutei


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("協議用素案 Ver.2.34 → Ver.2.35（第3回レビューの影響範囲の再点検）")
    shutil.copy2(MOTO, SAKI)
    d = docx.Document(SAKI)
    mae = zenbun(d)

    def rep(kyu, shin, n_min=1, mi=""):
        n = 0
        for p in subete(d):
            if kyu in p.text and p.runs:
                kakikae(p, p.text.replace(kyu, shin))
                n += 1
        judge(n >= n_min, f"{mi}「{kyu[:24]}」（{n}か所）")
        return n

    print("  ── R3-06 の波及　9-3 の「本算定」→「本試算」")
    n = rep("本算定", "本試算", mi="9-3")
    judge(n >= 5, f"{n}か所")

    print("  ── R3-08 の波及　通いの場の参加率の言い方")
    rep("調査により把握できる3指標（通いの場への参加率、", "調査により把握できる3指標（通いの場への参加割合（調査の回答者）、",
        mi="10-4")
    rep("n=439。参加率は15.9％", "n=439。参加している割合は15.9％", mi="2-6 表")
    rep("通いの場参加率／特定健診受診率／サポーター活動回数",
        "通いの場への参加割合（調査）・参加者数（町実績）／特定健診受診率／サポーター活動回数", mi="KPI 候補の表")
    n = rep("通いの場への参加率、", "通いの場への参加割合、", n_min=2, mi="KPI の構造の表")

    print("  ── 1-2 の留保　1人1月あたり給付費 約29.9万円の出所")
    sh, gh, cm, tk = jukyu_nengetsu()
    nin = sh + gh + cm + tk
    sz = mae
    kyufu = int(re.search(r"入所・居住系（[^）]+）が([\d,]+)千円", sz).group(1).replace(",", ""))
    tanka = kyufu / nin
    judge(round(tanka / 10, 1) == 29.9, f"{kyufu:,}千円 ÷ 延べ{nin:,}人月（施設{sh:,}・認知症GH{gh}・地域密着特養{cm}・特定施設{tk}）"
                                        f"＝{tanka:.2f}千円（約{tanka / 10:.1f}万円）")
    rates = [float(x) / 100 for x in re.findall(r"令和\d+年度([\d.]+)％",
             [p.text for p in d.paragraphs if "見込交付割合は全国の仮置値による自動計算値" in p.text][0])]
    N = 9743.1
    dm = sum(10 * tanka * 12 * (0.23 + 0.05 - e) for e in rates) * 1000 / 0.96 / N / 12
    judge(round(dm) == 212, f"丸める前の単価でも10人で {dm:.2f}円 → 約{round(dm)}円（変わらない）")
    rep("（1人1月あたり約29.9万円×10人×12か月の3か年分に、",
        f"（1人1月あたり約29.9万円（令和7年度の入所・居住系の給付費{kyufu:,}千円÷介護保険事業状況報告の受給者数"
        f"延べ{nin:,}人月）×10人×12か月の3か年分に、", mi="5-5")

    d.save(SAKI)
    ato = zenbun(docx.Document(SAKI))
    print("  ── 検査1　古い記述")
    for go in ("本算定", "参加率は15.9％", "通いの場への参加率、", "通いの場参加率／"):
        judge(go not in ato, f"「{go}」0件")
    print("  ── 検査2　未確定の印・他団体名・受託者の語は増減なし")
    for w in ("【町確認】", "【委員会協議】", "［要確認］", "［要協議］", "［要内訳］", "大雪", "北塩原", "金ヶ崎",
              "川崎市", "受託者", "当社", "ビズアップ"):
        judge(ato.count(w) == mae.count(w), f"「{w}」{mae.count(w)}→{ato.count(w)}")
    print(f"  {SAKI}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 不適合0件")


if __name__ == "__main__":
    main()
