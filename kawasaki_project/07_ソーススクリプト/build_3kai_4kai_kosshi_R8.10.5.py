# -*- coding: utf-8 -*-
"""第3回・第4回策定委員会資料の骨子案（令和8年10月5日）

08_作業順位 の順位18。

逆算工程表（02シート）では、第3回（保険料の決定・意見公募手続の結果の報告）が
令和9年1月、第4回（計画書（案）の確定・答申）が令和9年2月である。

**第2回の結果を待たずに、諮る事項と資料の構成は骨子として置ける。**
第2回でご選択いただく事項（選択1〜5）の帰結が第3回に流れ込むため、
骨子の段階で「何が決まっていないと開けないか」を先に示しておく。

  出力　03_委員会・説明資料/川崎町_第3回・第4回策定委員会_骨子案_R8.10.5.docx

⚠ **委員会を全3回とするか全4回とするかが決まっていない**（確認事項No.49）。
  第1回では「第3回（令和9年1月）で素案提示及び答申内容の確定」という
  段取りが示された一方、第1回資料の策定スケジュール（案）は
  第4回（令和9年2月・答申書案）を含む4回構成であった。
  **本骨子は4回構成を本文とし、3回構成に寄せる場合の差分を併せて示す。**

⚠ 保険料は第3回で決定する。基金の取崩額（選択5）が決まるまで、
  本骨子は3パターン（算定Ａ・Ｂ・Ｃ）のまま置き、1つに絞らない。

⚠ 意見公募手続の実施支援は委託仕様書の業務内容に規定がない
  （確認事項No.19）。第3回での結果の報告は、実施された場合の骨子である。
"""
import copy
import os
import re
import shutil
import sys
import zipfile

import docx
import openpyxl
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from fix_soan_v111 import clone_table, set_el            # noqa: E402

SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.10_制度改正補完版.docx"
GYOTEI = "05_試算・管理シート/川崎町_策定委員会_逆算工程表_R8.10.1.xlsx"
OUT = "03_委員会・説明資料/川崎町_第3回・第4回策定委員会_骨子案_R8.10.5.docx"

# ══════════════════════════ 第3回　次第（分・内容・諮る事項）
SHIDAI3 = [
    ("開会・事務局説明", "5", "事務局", "―"),
    ("① 第2回でいただいたご意見と反映の状況", "10",
     "選択1〜5のご選択を計画素案にどう反映したかを、"
     "箇所を示してご報告します。",
     "反映の内容のご確認"),
    ("② 意見公募手続の結果のご報告", "20",
     "いただいたご意見の件数・主な内容と、町の考え方（対照表）。"
     "計画の修正に至ったものを分けてお示しします。",
     "町の考え方のご確認／計画に反映する範囲"),
    ("③ 宮城県からいただいた意見と対応", "10",
     "事前協議でいただいた意見と、計画への反映。",
     "対応のご確認"),
    ("④ 介護保険料の決定", "35",
     "算定Ａ（基金の取崩しなし）・算定Ｂ（50％取崩し）・"
     "算定Ｃ（全額取崩し）の3パターンと、それぞれの第1号被保険者への影響。"
     "⚠ 基金の取崩額のご決定をいただきます。",
     "取崩額の決定（第2回の選択5）"),
    ("⑤ 計画書（案）の提示", "20",
     "本文・資料編・概要版の構成と、第2回からの変更点。"
     "⚠ 全3回とする場合は、ここで答申の内容まで確定します。",
     "計画書（案）のご確認"),
    ("⑥ 今後の予定", "10",
     "第4回（答申）・議会への説明・入稿・納品までの日取り。", "―"),
    ("閉会", "5", "事務局", "―"),
]

# ══════════════════════════ 第4回　次第
SHIDAI4 = [
    ("開会・事務局説明", "5", "事務局", "―"),
    ("① 第3回でいただいたご意見と反映の状況", "15",
     "保険料の決定・計画書（案）へのご意見の反映。",
     "反映の内容のご確認"),
    ("② 計画書（案）の最終確認", "40",
     "本文・資料編・概要版の確定版。"
     "⚠ この回を過ぎると、入稿までに直せるのは誤りの訂正のみです。",
     "計画書（案）の確定"),
    ("③ 答申", "20",
     "⚠ 答申を行うか、諮問書・答申書を計画書の別添に掲げるかは"
     "確認事項No.131のご判断によります。",
     "答申の決定"),
    ("④ 今後の予定（議会・入稿・公表）", "10",
     "議会への説明、入稿、公表、見える化システムへの登録。", "―"),
    ("閉会", "5", "事務局", "―"),
]

# ══════════════════════════ 資料の構成（第3回・第4回）
KOSEI3 = [
    ("資料1", "第2回のご意見と反映の状況（一覧）", "受託者", "作成可"),
    ("資料2", "意見公募手続の結果（意見と町の考え方の対照表）", "町・受託者",
     "⚠ 実施後でなければ作れない。様式は作成済み"),
    ("資料3", "宮城県からの意見と対応（一覧）", "町・受託者",
     "⚠ 協議後でなければ作れない。様式は作成済み"),
    ("資料4", "介護保険料の算定（3パターンと影響）", "受託者",
     "作成可（保険料試算ワークブックから）"),
    ("資料5", "計画書（案）本文・資料編・概要版", "受託者",
     "素案 Ver.2.10 を母体とする"),
    ("資料6", "今後の予定", "受託者", "作成可（逆算工程表から）"),
    ("参考", "想定問答集（事務局用）", "受託者", "第2回 v4 を母体とする"),
]

KOSEI4 = [
    ("資料1", "第3回のご意見と反映の状況（一覧）", "受託者", "作成可"),
    ("資料2", "計画書（案）確定版 本文・資料編・概要版", "受託者",
     "第3回の決定を反映した版"),
    ("資料3", "答申書（案）", "町",
     "⚠ 答申を行う場合。確認事項No.131"),
    ("資料4", "今後の予定（議会・入稿・公表）", "受託者", "作成可"),
    ("参考", "想定問答集（事務局用）", "受託者", "第3回から引き継ぐ"),
]

# ══════════════════════════ 何が決まっていないと開けないか
ZENTEI = [
    ("第3回", "第2回の選択1〜5のご選択",
     "選択2・3は見込量に、選択5は保険料に直結する。"
     "決まらないと保険料を決定できない。", "85・49"),
    ("第3回", "意見公募手続の実施（12月）",
     "⚠ 実施支援が本業務の範囲かが決まっていない（確認事項No.19）。"
     "期間が30日間の場合、第2回が11月11日を過ぎると第3回に"
     "結果が間に合わない。", "19・136"),
    ("第3回", "宮城県への事前協議の開始（11月下旬〜12月上旬）",
     "県の運用（要する期間・様式）が分からないため、"
     "日取りを置けない。", "4"),
    ("第3回", "介護給付費準備基金の令和8年度末残高の見込み",
     "取崩額を決定する前提。", "11・33"),
    ("第4回", "答申を行うか／諮問書・答申書を別添に掲げるか",
     "行う場合、答申書（案）の作成と委員長の決裁に日数を要する。",
     "131・129・130"),
    ("第4回", "成果品の入稿期日",
     "印刷・製本に21日間を見込んでいる。"
     "2週間に縮められる場合は第4回を1週間後ろにできる。", "18"),
    ("第3回・第4回", "委員会を全3回とするか全4回とするか",
     "3回構成とする場合、第3回で保険料の決定と答申を併せて行う。"
     "当日の時間が2時間では収まらない見込み。", "49"),
]

# ══════════════════════════ 3回構成に寄せる場合の差分
SANKAI = [
    ("第3回の時間", "2時間 → 2時間30分〜3時間",
     "保険料の決定（35分）と答申（20分）を同じ回で行うため。"),
    ("第3回の次第", "⑤計画書（案）の提示 のあとに ⑥答申 を加える",
     "答申書（案）を事前配布する必要がある。"),
    ("答申書（案）の作成", "第3回の2週間前までに町にてご用意",
     "計画書（案）が固まる前に答申書（案）を書くことになるため、"
     "⚠ 計画の内容に踏み込まない書きぶりとする必要がある。"),
    ("納期までの余裕", "25日の短縮",
     "第4回（答申）を設けないため、入稿を25日早められる。"
     "第2回が11月18日・11月25日になった場合の短縮の案"
     "（案Ｃ・案Ｄ）に当たる。"),
    ("取りうる場面", "第2回の開催が11月18日以降にずれた場合",
     "11月4日・11月11日に開催できるのであれば、"
     "4回構成のままで納期に収まる。"),
]


def honbun(d):
    C = []
    a = C.append
    a(("表題", "第3回・第4回策定委員会資料　骨子案"))
    a(("本文", "令和8年10月5日"))
    a(("囲み",
       "⚠ 本骨子は、第2回策定委員会の開催日が確定していない段階"
       "（確認事項No.85）で、諮る事項と資料の構成を先に置くものです。"
       "開催日と、委員会を全3回とするか全4回とするか（確認事項No.49）が"
       "確定しだい、日取りと次第を改めます。"))

    a(("見出し", "▌1　第3回・第4回の位置づけ"))
    a(("本文",
       "逆算工程表では、第2回策定委員会を令和8年11月11日に開催する場合、"
       "第3回が令和9年1月25日、第4回が令和9年2月19日となります。"
       "第3回で保険料を決定し、第4回で計画書（案）を確定して答申を"
       "いただく組み立てです。"))
    rows = [["年月日", "工程"]]
    for ymd, k in d["kotei"]:
        rows.append([ymd, k])
    a(("T2", rows, [24, 76]))
    a(("本文",
       "第2回の開催が遅れると、この先のすべてが同じだけ後ろにずれます。"
       "第2回を11月18日以降に開催する場合は、工程の短縮（意見公募手続を"
       "20日間とする案、委員会を全3回とする案）を併せてご検討ください。"))

    a(("見出し", "▌2　第3回策定委員会（令和9年1月）の骨子"))
    a(("本文", "保険料の決定と、意見公募手続の結果の報告が中心になります。"))
    rows = [["次第", "分", "内容", "諮る事項"]]
    for x in SHIDAI3:
        rows.append(list(x))
    a(("T4", rows, [26, 6, 44, 24]))
    a(("本文", "資料の構成は次のとおりです。"))
    rows = [["資料", "名称", "作る者", "現在の状態"]]
    for x in KOSEI3:
        rows.append(list(x))
    a(("T4", rows, [10, 36, 14, 40]))
    a(("囲み",
       "⚠ 保険料は、算定Ａ（基金の取崩しなし）・算定Ｂ（50％取崩し）・"
       "算定Ｃ（全額取崩し）の3パターンのまま第3回にお諮りします。"
       "1つに絞ってお示しすることはいたしません。"
       "取崩額のご決定をいただいて、はじめて1つに定まります。"))

    a(("見出し", "▌3　第4回策定委員会（令和9年2月）の骨子"))
    a(("本文",
       "計画書（案）の確定と答申です。"
       "⚠ この回を過ぎると、入稿までに直せるのは誤りの訂正のみになります。"))
    rows = [["次第", "分", "内容", "諮る事項"]]
    for x in SHIDAI4:
        rows.append(list(x))
    a(("T4", rows, [26, 6, 44, 24]))
    a(("本文", "資料の構成は次のとおりです。"))
    rows = [["資料", "名称", "作る者", "現在の状態"]]
    for x in KOSEI4:
        rows.append(list(x))
    a(("T4", rows, [10, 36, 14, 40]))

    a(("見出し", "▌4　何が決まっていないと開けないか"))
    a(("本文",
       "第3回・第4回は、第2回のご選択と、その後の手続の結果を前提とします。"
       "前提が整わないと開催しても決められません。"))
    rows = [["回", "前提", "なぜ要るか", "確認事項No."]]
    for x in ZENTEI:
        rows.append(list(x))
    a(("T4", rows, [12, 26, 48, 14]))

    a(("見出し", "▌5　委員会を全3回とする場合の差分"))
    a(("本文",
       "第1回では「第3回（令和9年1月）で素案提示及び答申内容の確定」という"
       "段取りが示された一方、第1回資料の策定スケジュール（案）は"
       "第4回（令和9年2月・答申書案）を含む4回構成でした"
       "（確認事項No.49）。全3回に寄せる場合の差分を示します。"))
    rows = [["変わるところ", "どう変わるか", "留意いただきたいこと"]]
    for x in SANKAI:
        rows.append(list(x))
    a(("T3", rows, [22, 34, 44]))
    a(("囲み",
       "⚠ 全3回とするのは、第2回の開催が令和8年11月18日以降にずれた場合の"
       "対処としてお考えください。11月4日又は11月11日に開催できるので"
       "あれば、4回構成のままで納期（令和9年3月15日）に収まります。"))

    a(("見出し", "▌6　ご確認をお願いしたいこと"))
    rows = [["", "事項", "確認事項No."]]
    for i, (koto, no) in enumerate(ONEGAI, 1):
        rows.append([str(i), koto, no])
    a(("T3", rows, [8, 70, 22]))
    a(("本文", "以上"))
    return C


ONEGAI = [
    ("第2回策定委員会の開催日", "85"),
    ("委員会を全3回とするか全4回とするか", "49"),
    ("答申を行うか／諮問書・答申書を計画書の別添に掲げるか", "131"),
    ("意見公募手続の実施支援を本業務の範囲とするか", "19"),
    ("宮城県への事前協議に要する期間と様式", "4"),
    ("成果品の入稿期日（印刷・製本に要する日数）", "18"),
]


def yomu():
    wb = openpyxl.load_workbook(GYOTEI)
    ws = wb["02_第2回から納品までの全体工程"]
    kotei = []
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 2).value is None:
            continue
        kotei.append((str(ws.cell(r, 2).value), str(ws.cell(r, 3).value or "")))
    return {"kotei": kotei}


def main():
    for p in (SOAN, GYOTEI):
        if not os.path.exists(p):
            raise SystemExit("入力がない：" + p)
    C = honbun(yomu())

    shutil.copy(SOAN, OUT)
    doc = docx.Document(OUT)
    body = doc.element.body
    tpl = {}
    for el in body.iterchildren():
        if el.tag == qn("w:p"):
            if el.findall(".//" + qn("w:drawing")):
                continue
            s = "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()
            if "表題" not in tpl and s.startswith("第 1 章"):
                tpl["表題"] = copy.deepcopy(el)
            if "見出し" not in tpl and s.startswith("▌"):
                tpl["見出し"] = copy.deepcopy(el)
            if "本文" not in tpl and s.startswith("令和7年度から令和8年度への"):
                tpl["本文"] = copy.deepcopy(el)
            if "囲み" not in tpl and s.startswith("⚠ 宮城県自身の評価"):
                tpl["囲み"] = copy.deepcopy(el)
        elif el.tag == qn("w:tbl"):
            n = len(el.findall(qn("w:tr"))[0].findall(qn("w:tc")))
            tpl.setdefault(f"T{n}", copy.deepcopy(el))
    for k in ("表題", "見出し", "本文", "囲み", "T2", "T3", "T4"):
        if k not in tpl:
            raise SystemExit("雛形が見つからない：" + k)

    sectPr = body.find(qn("w:sectPr"))
    for el in list(body.iterchildren()):
        if el is not sectPr:
            body.remove(el)

    def add(el):
        if sectPr is not None:
            sectPr.addprevious(el)
        else:
            body.append(el)

    n_tbl = 0
    for item in C:
        kind, v = item[0], item[1]
        if kind.startswith("T"):
            add(clone_table(tpl[kind], v, item[2]))
            e = copy.deepcopy(tpl["本文"])
            set_el(e, "")
            add(e)
            n_tbl += 1
        else:
            el = copy.deepcopy(tpl[kind])
            set_el(el, v)
            add(el)
    doc.save(OUT)
    strip_media(OUT)
    naoshi_header(OUT)

    # ══════════════════════════ 自己点検
    d2 = docx.Document(OUT)
    text = "\n".join(q.text for q in d2.paragraphs)
    tbl_text = "\n".join(c.text for t in d2.tables
                         for r in t.rows for c in r.cells)
    zen = text + "\n" + tbl_text
    ng = []
    z = zipfile.ZipFile(OUT)
    if "<w:drawing>" in z.read("word/document.xml").decode("utf-8"):
        ng.append("図が残っている")
    if sum(1 for n in z.namelist() if n.startswith("word/media")):
        ng.append("使っていない画像が残っている")
    if len(d2.tables) != n_tbl:
        ng.append(f"表が{len(d2.tables)}件（{n_tbl}件のはず）")
    for w in ("**", "__"):
        if w in zen:
            ng.append(f"強調記号が残っている：{w}")
    for w in ("大雪", "東川", "東神楽", "上川", "金ヶ崎", "金ケ崎", "川崎市"):
        if w in zen:
            ng.append(f"他団体の名称が入っている：{w}")
    # 保険料を1つに絞っていないこと
    if "6,822" in zen or "6,143" in zen or "5,464" in zen:
        ng.append("骨子に保険料の額を書いている（3パターンのまま置く）")
    if "算定Ａ" not in zen or "算定Ｂ" not in zen or "算定Ｃ" not in zen:
        ng.append("保険料の3パターンが示されていない")
    # 3回構成・4回構成の双方を示していること
    if "全3回" not in zen or "全4回" not in zen:
        ng.append("3回構成と4回構成の双方を示していない")
    # 未確定を断定していないこと
    for w in ("開催します", "決定しました", "確定しました"):
        if w in zen:
            ng.append(f"未確定の事項を断定している疑い：{w}")
    # 確認事項の参照が業務工程管理表にあること
    wb = openpyxl.load_workbook("川崎町_業務工程管理表.xlsx")
    ws = wb["03_確認事項一覧"]
    aru = {str(ws.cell(r, 1).value or "").strip()
           for r in range(5, ws.max_row + 1)}
    for no in sorted(set(re.findall(r"No\.(\d+)", zen))
                     | {n for _, n in ONEGAI}):
        for one in re.split(r"[・、]", str(no)):
            one = one.strip()
            if one and one not in aru:
                ng.append(f"確認事項No.{one}が業務工程管理表にない")

    print("第3回・第4回策定委員会資料　骨子案")
    print("保存：", OUT)
    print(f"  段落 {len(d2.paragraphs)}／表 {len(d2.tables)}"
          f"／字数 {len(text) + len(tbl_text):,}字")
    print(f"  第3回 次第{len(SHIDAI3)}件・資料{len(KOSEI3)}件"
          f"／第4回 次第{len(SHIDAI4)}件・資料{len(KOSEI4)}件"
          f"／前提{len(ZENTEI)}件／3回構成の差分{len(SANKAI)}件")
    print("  ── 自己点検")
    if ng:
        for m in ng:
            print("   ×", m)
        sys.exit(1)
    print("   ○ 保険料は3パターンのまま（額を書いていない）")
    print("   ○ 全3回・全4回の双方を示している（確認事項No.49）")
    print("   ○ 参照している確認事項はすべて業務工程管理表にある")
    print("  ⚠ 開催日・回数が確定しだい、日取りと次第を改めます。")


def strip_media(path):
    z = zipfile.ZipFile(path)
    names = z.namelist()
    data = {n: z.read(n) for n in names}
    z.close()
    used = set(re.findall(r'r:embed="([^"]+)"',
                          data["word/document.xml"].decode("utf-8")))
    rels = "word/_rels/document.xml.rels"
    from lxml import etree
    root = etree.fromstring(data[rels])
    drop = set()
    for rel in list(root):
        tgt = rel.get("Target") or ""
        if tgt.startswith("media/") and rel.get("Id") not in used:
            drop.add("word/" + tgt)
            root.remove(rel)
    data[rels] = etree.tostring(root, xml_declaration=True,
                                encoding="UTF-8", standalone=True)
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            if n not in drop:
                zout.writestr(n, data[n])
    shutil.move(tmp, path)


def naoshi_header(path):
    z = zipfile.ZipFile(path)
    names = z.namelist()
    data = {n: z.read(n) for n in names}
    z.close()
    for n in names:
        if not re.match(r"word/(header|footer)\d*\.xml", n):
            continue
        s = data[n].decode("utf-8")
        s = re.sub(r"川崎町[^<]*計画書素案[^<]*",
                   "川崎町 第3回・第4回策定委員会資料 骨子案", s)
        s = re.sub(r"Ver\.[\d.]+[^<]*", "令和8年10月5日", s)
        data[n] = s.encode("utf-8")
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    shutil.move(tmp, path)


if __name__ == "__main__":
    main()
