# -*- coding: utf-8 -*-
"""第2回策定委員会資料 v8（作業経過の記述の削除）（令和8年10月6日）

08_作業順位 の順位22の一部。

守るべき制約の走査（`check_seiyaku_R8.10.6.py`）により、
v7 の本文に**当方の作業用ファイルの名前**が1か所残っていることが
分かった。委員にお配りする資料に当方の内部の仕組みを書かない
（点検スキル 3「守るべき制約の走査」）。

  入力　03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v7.docx
  出力　03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v8.docx

⚠ 「当社が記載していた誤り」「当社の取りまとめ表の集計が公表値と
  食い違っていた」といった記述は**残す。** 訂正の責任の所在を
  委員にお示しするものであり、伏せるべきものではない。
  走査では「要判断」とし、取り除いていない。

⚠ 数値は一切変えていない。
"""
import os
import re
import shutil
import sys
import zipfile

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from fix_soan_v111 import set_el                             # noqa: E402

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v7.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v8.docx"

NAOSU = [
    ("`07_ソーススクリプト/check_hokenryo_R8.9.30.py`・39項目すべて適合",
     "39項目すべて適合"),
    # ⚠ 表紙の版の表示が［第５版］のまま v6・v7 と版を重ねていた。
    #   本文の「本資料の第５版（v5）から」は履歴の記述であり、残す。
    ("計画骨子（案）とご選択いただきたい事項　［第５版］",
     "計画骨子（案）とご選択いただきたい事項　［第８版］"),
]


def main():
    if not os.path.exists(SRC):
        raise SystemExit("入力がない：" + SRC)
    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body

    n = 0
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = "".join(x.text or "" for x in el.iter(qn("w:t")))
        for furui, atarashii in NAOSU:
            if furui in s:
                set_el(el, s.replace(furui, atarashii))
                n += 1
    if n != len(NAOSU):
        raise SystemExit(f"差し替えが{n}件（{len(NAOSU)}件のはず）")

    doc.save(DST)

    z = zipfile.ZipFile(DST)
    names = z.namelist()
    data = {x: z.read(x) for x in names}
    z.close()
    for x in names:
        if re.match(r"word/(header|footer)\d*\.xml", x):
            t = data[x].decode("utf-8")
            data[x] = t.replace("v7", "v8").encode("utf-8")
    tmp = DST + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for x in names:
            zout.writestr(x, data[x])
    shutil.move(tmp, DST)

    # ══════════════════════════ 自己点検
    d2 = docx.Document(DST)
    d0 = docx.Document(SRC)
    zen = ("\n".join(q.text for q in d2.paragraphs) + "\n"
           + "\n".join(c.text for t in d2.tables
                       for r in t.rows for c in r.cells))
    ng = []
    for w in (".py", "07_ソーススクリプト", "`"):
        if w in zen:
            ng.append(f"作業用ファイルの名前・記号が残っている：{w}")
    if (len(d2.paragraphs), len(d2.tables)) != (len(d0.paragraphs),
                                               len(d0.tables)):
        ng.append("段落数・表数が v7 と違う（本版は文言の差し替えのみ）")
    # 訂正の責任の所在は残っていること
    if "当社" not in zen:
        ng.append("訂正の責任の所在を示す記述まで消している")
    # 表紙の版の表示が揃っていること
    if "［第８版］" not in zen:
        ng.append("表紙の版の表示が［第８版］になっていない")
    if "本資料の第５版（v5）から" not in zen:
        ng.append("履歴の記述（第５版からの是正）まで消している")

    print("第2回策定委員会資料 v8")
    print("保存：", DST)
    print(f"  段落 {len(d2.paragraphs)}／表 {len(d2.tables)}"
          f"（v7 と同じ）／差し替え {n}件")
    print("  ── 自己点検")
    if ng:
        for m in ng:
            print("   ×", m)
        sys.exit(1)
    print("   ○ 作業用ファイルの名前を取り除いた")
    print("   ○ 表紙の版の表示を［第８版］に揃えた"
          "（v6・v7 は［第５版］のままであった）")
    print("   ○ 訂正の責任の所在を示す記述は残している")
    print("   ○ 段落数・表数は v7 と同じ（文言の差し替えのみ）")


if __name__ == "__main__":
    main()
