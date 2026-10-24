# -*- coding: utf-8 -*-
"""第2回策定委員会 想定問答集 v6 → v7（交付金の状況の再整理の反映）
（令和8年10月22日）

  python3 07_ソーススクリプト/build_2kai_qa_v7.py

v6 からの変更の理由
  計画素案を Ver.2.24、委員会資料を v11 に上げた。
  (1) ⚠ QB-36 の「第9章の保険料算定における交付金の取扱いも、令和8年度の
      交付実績を据え置いています」は事実と違う（保険料の試算は交付金を
      見込んでいない）。素案 Ver.2.24 と同じ整理に改める
  (2) ⚠ QB-22・QB-24 の値を素案（表115）に揃える
        改定率1％　65円 → 66円（QB-22 の A と根拠データ。QB-17 は v3 で
          66円に直していたが、QB-22 が直し漏れていた）
        予定収納率99.6％　247円 → 246円（QB-24）
  (3) 3問を加える（全64問 → 全67問。カテゴリH）
      QB-65 37点下がったのはなぜか　QB-66 県内順位が3位から8位に
      下がったのは　QB-67 交付金は保険料に反映されているのか
  (4) ⚠ ヘッダーの「第1回策定委員会 想定問答集」→「第2回」
  (5) 計画素案の参照 Ver.2.23 → Ver.2.24

数値の正本　05_試算・管理シート/川崎町_交付金_推移の再整理_R8.10.22.xlsx
            計画素案 Ver.2.24（保険料の効き）
⚠ 宣言した件数と実際の件数が合っていることを機械で確かめる。
"""
import collections
import copy
import importlib.util
import os
import re
import shutil
import sys

import docx
import openpyxl
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)


def load(name, fn):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(spec)
    sys.modules[name] = m
    spec.loader.exec_module(m)
    return m


S24 = load("fix_soan_v2240", "fix_soan_v2240_kofukin.py")

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v6.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v7.docx"
SU = 67


def ten(v):
    return f"＋{v}点" if v > 0 else f"▲{-v}点"


def ichi():
    wb = openpyxl.load_workbook(S24.SAISEIRI, data_only=True)
    ws = wb["04_目標ごとの位置_R8"]
    return {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}


def cat_of(s):
    m = re.match(r"^カテゴリ ([A-H])　", s)
    return "カテゴリ " + m.group(1) if m else None


def txt(el):
    return "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会 想定問答集 v6 → v7（交付金の状況の再整理）")
    print()
    kubun, komoku, goukei = S24.uchiwake()
    r7, r8 = goukei["令和7年度"], goukei["令和8年度"]
    gen = r8[1] - r7[1]
    seika = kubun["成果指標群の変動"]
    haiten = kubun["配点の改定（得点率は不変）"]
    shinsetsu = kubun["指標の新設"]
    kotai = kubun["取組の後退（得点率の低下）"]
    zenshin = kubun["取組の前進（得点率の上昇）"]

    def sk(atama):
        return sum(v for (k, na), v in komoku.items()
                   if k == "成果指標群の変動" and na.startswith(atama))

    katsuyo = komoku[("取組の後退（得点率の低下）", "評価結果の活用")]
    tekisei = komoku[("取組の後退（得点率の低下）", "給付費適正化事業の取組状況")]
    I = ichi()
    ken_up = round(r8[5] - r7[5], 1)
    zen_up = round(r8[2] - r7[2], 1)

    NAOSU = [
        ("QB-22 A", "A．改定率1％あたり月額約65円です。",
         "改定率1％あたり月額約65円", "改定率1％あたり月額約66円"),
        ("QB-22 根拠", "【根拠データ】 第2回資料 ６（介護報酬改定率別の算定表",
         "改定率1％＝月額＋65円", "改定率1％＝月額＋66円"),
        ("QB-24 A", "A．現年度分だけを見れば令和7年度は99.56％と",
         "月額247円下がります", "月額246円下がります"),
        ("QB-36 A", "A．金額はお示しできません。",
         "第9章の保険料算定における交付金の取扱いも、令和8年度の交付実績を"
         "据え置いています。",
         "第9章の保険料の試算でも、交付金は見込んでいません（QB-67）。"),
    ]

    ADD = [
        ("カテゴリ H",
         f"QB-65　【委員長】　交付金の得点が1年で{-gen}点下がったのはなぜか",
         f"A．枝番の単位まで下ろし、各年度の配点に対する得点率で比べると、"
         f"5つに分かれます。①アウトカム指標（要介護度の変化等）が"
         f"{ten(seika)}です。短期的な要介護度の変化で{ten(sk('短期'))}、"
         f"健康寿命延伸の状況で{ten(sk('健康寿命'))}、長期的な要介護度の変化で"
         f"{ten(sk('長期'))}でした。アウトカム指標は推進・支援の両方に同じ点が"
         f"入るため、1つの指標の変動が2倍に効きます。②国の配点の改定が"
         f"{ten(haiten)}です。認知症サポーター等を活用した地域支援体制の構築は"
         f"25点満点から20点満点に改められましたが、本町は満点を維持して"
         f"います。③新設の指標で{ten(shinsetsu)}です。④取組の評価が下がった"
         f"項目が{ten(kotai)}、⑤上がった項目が{ten(zenshin)}です。④のうち"
         f"「評価結果の活用」{-katsuyo}点と給付費適正化事業のア{-tekisei}点の"
         f"計{-(katsuyo + tekisei)}点は、令和7年度には満点であったもので、"
         f"仕組みを定めることで評価の対象となり得ます。④と⑤はほぼ相殺して"
         f"おり、{-gen}点のうち{-seika}点はアウトカム指標の変動によるものです。",
         "【根拠データ】 第2回資料 ７（37点下がった内訳の表）。令和7・8年度の"
         "全国集計結果を枝番の単位で突き合わせたもの",
         "【関連資料】 計画素案Ver.2.24 第3章3-3、"
         "川崎町_交付金_推移の再整理_R8.10.22.xlsx 02シート"),
        ("カテゴリ H",
         f"QB-66　【有識者】　県内順位が{r7[7]}位から{r8[7]}位に下がったのは、"
         "本町の取組が遅れているからか",
         f"A．主な理由は、県内の保険者の得点が大きく上がったことです。"
         f"宮城県内{r8[8]}保険者の平均は、令和7年度の{r7[5]:.1f}点から"
         f"令和8年度の{r8[5]:.1f}点へ、1年で{ken_up}点上がりました（全国平均の"
         f"上昇は{zen_up}点）。本町は{-gen}点下げましたが、なお県平均を"
         f"{r8[1] - r8[5]:.1f}点上回っています。⚠ 一方で、目標ごとにみると"
         f"弱いところがあります。推進交付金の目標Ⅰ（持続可能な地域の"
         f"あるべき姿をかたちにする）は{I['推進 目標Ⅰ'][1]}点で県内"
         f"{I['推進 目標Ⅰ'][5]}位（県平均{I['推進 目標Ⅰ'][3]}点）、アウトカム"
         f"指標は{I['推進 目標Ⅳ'][1]}点で県内{I['推進 目標Ⅳ'][5]}位です。"
         f"逆に、努力支援交付金の目標Ⅰ（介護予防・日常生活支援）は"
         f"{I['支援 目標Ⅰ'][1]}点で県内{I['支援 目標Ⅰ'][5]}位・全国"
         f"{I['支援 目標Ⅰ'][4]}位です。目標Ⅰの弱さの多くは「評価結果の活用」"
         f"16点によるもので、本計画の第8章・第10章で手当てします。",
         "【根拠データ】 第2回資料 ７（3か年の推移の表）。令和6〜8年度の"
         "全国集計結果（都道府県平均・県内の分布）",
         "【関連資料】 計画素案Ver.2.24 第3章3-3、"
         "川崎町_交付金_推移の再整理_R8.10.22.xlsx 01・04シート"),
        ("カテゴリ H",
         "QB-67　【住民代表】　交付金は、保険料に反映されているのか",
         "A．現時点の試算（取崩しなしで月額6,822円ほか）では見込んでいません。"
         "交付金の交付見込額は保険料収納必要額から差し引くことができ、"
         "地域包括ケア「見える化」システムの将来推計総括表では3か年で"
         f"{S24.KOFU:,}千円が置かれています。これを差し引くと月額約"
         f"{-S24.KOFU_SA}円下がります。令和6年度の交付見込額は3,738千円"
         "（うち激変緩和措置分1,231千円）でしたが、令和7・8年度の交付額を"
         "確認中であり、また交付額は他の保険者の得点にも左右されます。"
         "見込むか否か及び見込む額は、第3回でご協議いただきます。",
         "【根拠データ】 第2回資料 ６（保険料を動かす要因の表）。令和6年度"
         "全国集計結果（交付見込額確定版）。将来推計総括表（令和8年10月19日"
         "出力）5_保険料推計",
         "【関連資料】 計画素案Ver.2.24 第9章9-3（不確定要素⑤・前提を動かした"
         "場合の幅）、確認事項No.39"),
    ]

    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body
    from fix_soan_v111 import set_el

    print("  ── ① 段落の中の語")
    n = 0
    for mi, atama, kyu, atarashii in NAOSU:
        hit = [el for el in body.iterchildren()
               if el.tag == qn("w:p") and txt(el).startswith(atama)]
        if len(hit) != 1 or kyu not in txt(hit[0]):
            judge(False, f"{mi}：当たる段落が{len(hit)}件か旧の語がない")
            continue
        set_el(hit[0], txt(hit[0]).replace(kyu, atarashii))
        n += 1
    judge(n == len(NAOSU), f"{len(NAOSU)}か所を改めた（実際 {n}）")

    n_ver = 0
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = "".join(x.text or "" for x in el.iter(qn("w:t")))
        if "素案Ver.2.23" in s:
            set_el(el, s.replace("Ver.2.23", "Ver.2.24"))
            n_ver += 1
    judge(n_ver >= 30, f"素案の版の参照を {n_ver}か所改めた")

    n_hd = 0
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            for r in p.runs:
                if "第1回策定委員会" in r.text:
                    r.text = r.text.replace("第1回策定委員会", "第2回策定委員会")
                    n_hd += 1
    judge(n_hd >= 1, f"ヘッダーを「第2回」に改めた（{n_hd}か所）")

    kids = list(body.iterchildren())
    tpl_q = tpl_a = None
    for i, el in enumerate(kids):
        if el.tag == qn("w:p") and txt(el).startswith("QB-1　"):
            tpl_q, tpl_a = copy.deepcopy(el), copy.deepcopy(kids[i + 1])
            break

    def last_of(cat):
        cur = last = None
        for el in body.iterchildren():
            if el.tag != qn("w:p"):
                continue
            s = txt(el)
            c = cat_of(s)
            if c:
                cur = c
            if cur == cat and s.startswith("【関連資料】"):
                last = el
        return last

    for cat, q, a, kon, kan in ADD:
        after = last_of(cat)
        for text, tpl in ((kan, tpl_a), (kon, tpl_a), (a, tpl_a), (q, tpl_q)):
            el = copy.deepcopy(tpl)
            set_el(el, text)
            after.addnext(el)
    judge(True, f"{len(ADD)}問を加えた")

    cat_cnt = collections.Counter()
    cur = None
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = txt(el)
        c = cat_of(s)
        if c:
            cur = c
        if re.match(r"^QB-\d+　", s):
            cat_cnt[cur] += 1
    total = sum(cat_cnt.values())
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = txt(el)
        m = re.match(r"^(カテゴリ [A-H]　.+?)（\d+問）$", s)
        if m:
            set_el(el, f"{m.group(1)}（{cat_cnt[cat_of(s)]}問）")
        elif re.search(r"全\d+問", s):
            set_el(el, re.sub(r"全\d+問", f"全{total}問", s))
    from fix_soan_v111 import set_cell
    for t in doc.tables:
        if [c.text.strip() for c in t.rows[0].cells][:3] == \
                ["カテゴリ", "テーマ", "問数"]:
            for r in t.rows[1:]:
                k = "カテゴリ " + r.cells[0].text.strip()
                if k in cat_cnt:
                    set_cell(r.cells[2], str(cat_cnt[k]))
    doc.save(DST)

    d2 = docx.Document(DST)
    zen = "\n".join(p.text for p in d2.paragraphs) + "\n" + "\n".join(
        c.text for t in d2.tables for r in t.rows for c in r.cells)
    hd = "\n".join(p.text for s in d2.sections for p in s.header.paragraphs)
    print("  ── 検査1　件数")
    qs = [int(x) for x in re.findall(r"QB-(\d+)　", zen)]
    judge(len(qs) == total == SU, f"全{SU}問（実際 {len(qs)}・宣言 {total}）")
    judge(sorted(qs) == list(range(1, SU + 1)), "番号に欠番・重複がない")
    judge(f"全{total}問" in zen, f"「全{total}問」と宣言している")
    print("  ── 検査2　⚠ 事実と違う記述・素案と違う値")
    judge("交付実績を据え置いて" not in zen, "「交付実績を据え置いて」が0件")
    judge("65円" not in zen.replace("6,365円", "")
          and "247円" not in zen, "65円・247円が0件（素案は66円・246円）")
    judge(f"月額約{-S24.KOFU_SA}円下がります" in zen,
          f"交付金を見込んだ場合 月額約{-S24.KOFU_SA}円（素案と同じ）")
    judge("Ver.2.23" not in zen, "「Ver.2.23」が0件")
    judge("第1回策定委員会 想定問答集" not in hd, "ヘッダーが第2回")
    print("  ── 検査3　加えた内容")
    for go in (f"{ten(seika)}です", f"{ten(haiten)}です", "1年で78.6点",
               "3,738千円"):
        judge(go in zen, f"「{go}」がある")
    print("  ── 検査4　⚠ 事業所名・他団体名・強調記号")
    for na in ("釜房みどりの園", "アルパイン", "蔵王グリーンホーム",
               "グループホームさわやか", "グループホームふれあい",
               "ながさか", "リリーハイム", "北塩原", "金ヶ崎", "川崎市"):
        judge(na not in zen, f"「{na}」が入っていない")
    judge("**" not in zen, "強調記号（**）が0件")

    print()
    print(f"  {DST}")
    print(f"  全{total}問・" + "／".join(
        f"{k[-1]}{v}問" for k, v in sorted(cat_cnt.items())))
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        for m in futeki:
            print(f"      {m}")
        sys.exit(1)
    print("   ○ 検査1〜4　不適合0件")


if __name__ == "__main__":
    main()
