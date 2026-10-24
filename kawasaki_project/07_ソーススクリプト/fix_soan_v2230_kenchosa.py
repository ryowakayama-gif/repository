# -*- coding: utf-8 -*-
"""計画素案 Ver.2.22 → Ver.2.23（宮城県の入所状況調の反映）
（令和8年10月22日）

  python3 07_ソーススクリプト/fix_soan_v2230_kenchosa.py

なぜ要るか
  宮城県「老人福祉施設等入所状況調」（3か月ごと・施設別）から、
  町内施設の14時点（令和5年4月1日〜令和8年7月1日）を採録した
  （fetch_miyagi_nyusho_R8.10.22.py）。これにより
    ①稼働率・休止床を「把握していない」としていた記述が改まる
      （時点の入所率の推移で代える。延べの稼働率ではない）
    ②退所理由の14時点の計で、在宅復帰と終のすみかの別がはっきりする
    ③⚠ **町内にあった軽費老人ホーム（定員70人）が令和5年10月に休止し、
      令和6年4月以降は一覧に載っていない。** 素案に入っていなかった
    ④入所前の居所は、県の入所指針（令和5年7月25日）第3 が定める
      保険者の責務の中で集計値を求める筋を書く（町のご判断［要協議］）
    ⑤短期入所の稼働は県調査にも公表システムにもない（新たな空白）

  ⚠ 県内他団体の前期計画は確かめていない（ご指示により方針Ｃ）。
    本版の記述はいずれも国・県の原典から起こしている。

数値の正本　09_元資料/その他データ/県入所状況調_川崎町_R8.10.22.csv
"""
import copy
import csv
import os
import re
import shutil
import sys

import docx
import docx.text.paragraph as dp

MOTO = "01_第10期_最新版成果品/川崎町_計画書素案_v2.22_給付費整合版.docx"
SAKI = "01_第10期_最新版成果品/川崎町_計画書素案_v2.23_県調査反映版.docx"
CSV_IN = "09_元資料/その他データ/県入所状況調_川崎町_R8.10.22.csv"
BAN_MOTO, BAN = "Ver.2.22", "Ver.2.23"
SHUTTEN = "宮城県「老人福祉施設等入所状況調」"
JUN = ["R5.4.1", "R5.7.1", "R5.10.1", "R6.1.1", "R6.4.1", "R6.7.1",
       "R6.10.1", "R7.1.1", "R7.4.1", "R7.7.1", "R7.10.1", "R8.1.1",
       "R8.4.1", "R8.7.1"]


def wa(h):
    y, m, d = h.split(".")
    return f"令和{y[1:]}年{m}月{d}日"


def num(s):
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


# ════════════════════════════════════════ 数値（正本から引く）
def atai():
    if not os.path.exists(CSV_IN):
        raise SystemExit("入力がない：" + CSV_IN
                         + "\n先に fetch_miyagi_nyusho_R8.10.22.py を実行")
    with open(CSV_IN, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    by = {}
    for r in rows:
        by.setdefault((r["種別"], r["施設名"].split("（")[0]), {})[r["時点"]] = r

    def ritsu_haba(shu):
        out = []
        for (s, na), d in by.items():
            if s != shu:
                continue
            v = [num(d[h]["入所現員"]) / num(d[h]["定員"]) * 100
                 for h in JUN if h in d and num(d[h]["定員"])
                 and num(d[h]["入所現員"]) is not None]
            out.append((min(v), max(v)))
        return sorted(out, key=lambda x: -x[0])

    def kei(shu, koumoku):
        return sum(num(d[h][koumoku]) or 0 for (s, _), d in by.items()
                   if s == shu for h in JUN if h in d)

    def kibou(shu, h):
        return sum(num(d[h]["入所希望者"]) or 0 for (s, _), d in by.items()
                   if s == shu and h in d)

    keihi = [(h, d[h]) for (s, _), d in by.items() if s == "軽費"
             for h in JUN if h in d]
    return {
        "特養": ritsu_haba("特養"), "老健": ritsu_haba("老健"),
        "ＧＨ": ritsu_haba("ＧＨ"),
        "在宅_老健": kei("老健", "在宅"), "在宅_特養": kei("特養", "在宅"),
        "在宅_ＧＨ": kei("ＧＨ", "在宅"),
        "希望_特養_最初": kibou("特養", JUN[0]),
        "希望_特養_最新": kibou("特養", JUN[-1]),
        "軽費": keihi,
        "軽費定員": num(keihi[0][1]["定員"]) if keihi else None,
        "軽費最後": keihi[-1][0] if keihi else None,
        "時点数": len(set(r["時点"] for r in rows)),
    }


A = atai()


def haba(lst):
    return "と".join(f"{a:.1f}〜{b:.1f}％" for a, b in lst)


KYUSHI = next((h for h, r in A["軽費"] if "休止" in r["施設名"]), None)
NAI = JUN[JUN.index(A["軽費最後"]) + 1] if A["軽費最後"] else None
KENCHOSA = (f"{SHUTTEN}（{wa(JUN[0])}から{wa(JUN[-1])}までの"
            f"3か月ごと{A['時点数']}時点）")

# ════════════════════════════════════════ 段落の差し替え
DAN = [
    (210, "本節に掲載する供給側のデータは次のとおりです。",
     "本節に掲載する供給側のデータは次のとおりです。定員・入所者数・"
     "入所率・待機者数・職種別の従業者数・指定の年月日は、"
     "厚生労働省「介護サービス情報公表システム」（令和7年9月〜令和8年10月"
     "の各事業所の記入日時点）により把握しました。⚠ 公表情報は事業所の"
     "年1回の報告によるため、記入日が事業所ごとに異なります。"
     "あわせて、入所・入居系の施設の入所率と退所理由の推移を"
     + KENCHOSA + "により把握しました。⚠ これは時点の在所者数によるもので"
     "あり、延べ利用日数による稼働率ではありません。"
     "いずれにも項目のないもの（入所前の居所・経営の状況・短期入所の"
     "稼働）は、本計画では把握していないものとして扱い、代わりに用いる"
     "値を明示します。"),
    (211, "町内の供給には2つの特徴があります。",
     None),   # 「2つ」→「3つ」のみ（下で語を置き換える）
    (466, "この提供体制は、厚生労働省「介護サービス情報公表システム」",
     "この提供体制は、厚生労働省「介護サービス情報公表システム」"
     "（令和7年9月〜令和8年10月の各事業所の記入日時点）により確かめました。"
     "入所・入居系の定員は計173人、在所者数は計166人で、入所率は"
     "90.0％から100.0％です。" + KENCHOSA + "でみると、"
     f"介護老人福祉施設は{haba(A['特養'])}、介護老人保健施設は"
     f"{haba(A['老健'])}、認知症対応型共同生活介護は{haba(A['ＧＨ'])}で"
     "推移しており、介護老人福祉施設の1施設には恒常的に空きがあります。"
     "⚠ 入所率は「在所者数÷定員」であり、延べ利用日数による稼働率では"
     "ありません。また、町外の方の利用も含みますので、第9章9-1の本町の"
     "被保険者の利用実績と割ってはなりません。なお、地域密着型サービス"
     "（地域密着型介護老人福祉施設入所者生活介護2施設・認知症対応型"
     "共同生活介護2事業所）及び居宅介護支援は町が指定権者であり、"
     "その他は宮城県が指定権者です。"),
    (567, "⚠ 施設・居住系への入所と退去が、自宅との間で生じているのか",
     "⚠ 施設・居住系への入所と退去が、自宅との間で生じているのか"
     "医療機関との間で生じているのかにより、整備の考え方が変わります。"
     "本町の在宅介護実態調査では、施設等への申込みをしている方が33.8％、"
     "検討している方が14.4％（計48.2％・n=139）でした。"
     + KENCHOSA + "の退所理由を合計すると、在宅への退所は介護老人保健施設"
     f"が{A['在宅_老健']:.0f}人であるのに対し、介護老人福祉施設は"
     f"{A['在宅_特養']:.0f}人、認知症対応型共同生活介護は"
     f"{A['在宅_ＧＨ']:.0f}人です。介護老人保健施設は在宅復帰の機能が"
     "働いており、介護老人福祉施設は終の住まいとして機能しています。"
     "介護老人福祉施設の入所希望者は、"
     f"{wa(JUN[0])}の{A['希望_特養_最初']:.0f}人から{wa(JUN[-1])}の"
     f"{A['希望_特養_最新']:.0f}人へ減っています（⚠ 複数の施設への"
     "申込みが重複して計上された数であり、実人数ではありません）。"
     "⚠ 一方で「入所前にどこにいたか」は、いずれの公表資料にもありません。"
     "宮城県指定介護老人福祉施設入所指針（令和5年7月25日）により、施設は"
     "入所希望者の置かれている状況を確認し、入所希望者名簿を5年間保管する"
     "こととされ、保険者は入所規程の作成に積極的に協力し必要な援助を"
     "行うこととされています。この協力の場において、名簿の集計値"
     "（在宅・医療機関・他の施設の別の人数）の提供を求めるかを検討します"
     "［要協議］。提供を受けられない場合は、退所先と入所希望者数により"
     "整備の必要量を判断します（確認事項No.159）。"),
]

# 2-7 に加える「第3の特徴」（p212 の直後に段落を足す）
DAI3 = ("第3に、町内にあった軽費老人ホーム（定員"
        f"{A['軽費定員']:.0f}人）は、{wa(KYUSHI)}時点で休止中となり、"
        f"{wa(NAI)}以降は{SHUTTEN}に掲載されていません。"
        "第9期計画期間が始まる直前に、町内の高齢者向けの住まい"
        f"{A['軽費定員']:.0f}人分が失われたことになります。"
        "⚠ 軽費老人ホームは老人福祉法に基づく施設であり、"
        "介護サービス情報公表システムには載りません。休止・廃止の経緯と、"
        "入居されていた方の行き先を確認しています［要確認］"
        "（確認事項No.154）。")

# (表の番号, 見出し, 行, 列, 旧の冒頭, 新)
SERU = [
    (4, "事項", 7, 3, "2-7・7-1・9-1。施設・事業所数・利用者数",
     "2-7・7-1・9-1。施設・事業所数・利用者数・入所率・従事者数は"
     "介護サービス情報公表システムにより把握した（令和8年10月）。"
     f"入所率と退所理由の推移は{KENCHOSA}による。⚠ 延べ利用日数による"
     "稼働率は把握しておらず、時点の入所率で代える。⚠ 短期入所の稼働は"
     "いずれの資料にもない（確認事項No.159）"),
    (4, "事項", 8, 2, "未対応", "一部"),
    (4, "事項", 8, 3, "町及び宮城県の名簿により把握する",
     "2-7・7-1。⚠ 町内の軽費老人ホーム（定員"
     f"{A['軽費定員']:.0f}人）が{wa(KYUSHI)}時点で休止中となり、"
     f"{wa(NAI)}以降は{SHUTTEN}に掲載されていない。"
     "特定施設入居者生活介護の指定を受けた住まいは町内に0件"
     "（介護サービス情報公表システム）。指定を受けていない有料老人ホーム・"
     "サービス付き高齢者向け住宅は町及び宮城県の名簿により把握する"
     "（確認事項No.154）。第10期は基本的記載事項に「要介護者等の入居状況」"
     "が追加される"),
    (26, "区分", 2, 2, "入所・入居系の定員は計173人",
     "入所・入居系の定員は計173人、在所者数は計166人、入所率は"
     "90.0％から100.0％です。待機者数は介護老人福祉施設20人・介護老人保健"
     "施設2人（定義は「入所申込者の数－（入所定員－入所者）」）。"
     f"{KENCHOSA}でみた入所率は、介護老人福祉施設{haba(A['特養'])}、"
     f"介護老人保健施設{haba(A['老健'])}、認知症対応型共同生活介護"
     f"{haba(A['ＧＨ'])}です。⚠ 延べ利用日数による稼働率ではありません。"
     "⚠ 短期入所の稼働はいずれの資料にもなく、把握していません"),
    (102, "判断項目", 1, 1, "定員、入所率、待機者数。",
     "定員、入所率とその推移、待機者数、退所理由。"
     "⚠ 延べ利用日数による稼働率・短期入所の稼働・入所前の居所は"
     "公表資料にない"),
    (102, "判断項目", 1, 3, "介護サービス情報公表システム（入所率",
     "介護サービス情報公表システム（入所率・待機者数・退所先）及び"
     f"{SHUTTEN}（3か月ごとの入所率・入所希望者・退所理由）。"
     "入所前の居所は、県の入所指針が定める保険者の協力の場で集計値を"
     "求めることを検討する（確認事項No.159）"),
]


def dan_kaku(p, s):
    if p.runs:
        p.runs[0].text = s
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(s)


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    print(f"計画素案 {BAN_MOTO} → {BAN}（宮城県の入所状況調の反映）")
    print()
    if not os.path.exists(MOTO):
        print(f"   × 前の版がありません：{MOTO}")
        sys.exit(1)
    shutil.copy2(MOTO, SAKI)
    d = docx.Document(SAKI)
    mae = "\n".join([p.text for p in d.paragraphs]
                    + [c.text for t in d.tables for r in t.rows
                       for c in r.cells])
    KYU = []

    # ── 表のセル（段落を足す前に。段落の番号は表に効かない）
    print("  ── ① 表のセル")
    MIDASHI = {i: t.rows[0].cells[0].text.strip()
               for i, t in enumerate(d.tables)}
    n_seru = 0
    for ban, mi, gyo, retsu, kyu, atarashii in SERU:
        if not MIDASHI[ban].startswith(mi):
            judge(False, f"表{ban} の見出しが『{mi}』でない",
                  f"実際『{MIDASHI[ban][:12]}』")
            continue
        c = d.tables[ban].rows[gyo].cells[retsu]
        if not c.text.strip().startswith(kyu):
            judge(False, f"表{ban} 行{gyo} 列{retsu} の旧が合わない",
                  f"実際『{c.text[:30]}』")
            continue
        KYU.append(c.text)
        ps = c.paragraphs
        dan_kaku(ps[0], atarashii)
        for p in ps[1:]:
            for r in p.runs:
                r.text = ""
        n_seru += 1
    judge(n_seru == len(SERU), f"{len(SERU)}セルを差し替えた（実際 {n_seru}）")

    # ── 段落（番号がずれないよう、足す前に差し替える）
    print("  ── ② 段落")
    n_dan = 0
    for i, atama, atarashii in DAN:
        p = d.paragraphs[i]
        if not p.text.startswith(atama):
            judge(False, f"p{i} の冒頭が合わない", f"実際『{p.text[:26]}』")
            continue
        KYU.append(p.text)
        if atarashii is None:          # p211：「2つの特徴」→「3つの特徴」
            dan_kaku(p, p.text.replace("2つの特徴", "3つの特徴", 1))
        else:
            dan_kaku(p, atarashii)
        n_dan += 1
    judge(n_dan == len(DAN), f"{len(DAN)}段落を差し替えた（実際 {n_dan}）")

    # ── 2-7 に第3の特徴を足す（p212 の直後）
    print("  ── ③ 2-7 に「第3の特徴」を足す")
    p212 = d.paragraphs[212]
    ok = p212.text.startswith("第2に、町内の8の事業所を運営しているのは")
    judge(ok, "p212 が「第2に」の段落である", f"実際『{p212.text[:20]}』")
    if ok:
        np = copy.deepcopy(p212._p)
        p212._p.addnext(np)
        dan_kaku(dp.Paragraph(np, p212._parent), DAI3)

    # ── 版表記
    print("  ── ④ 版表記")
    n_ban = 0
    for sec in d.sections:
        for part in (sec.header, sec.footer):
            for p in part.paragraphs:
                for r in p.runs:
                    if BAN_MOTO in r.text:
                        r.text = r.text.replace(BAN_MOTO, BAN)
                        n_ban += 1
    for p in d.paragraphs:
        for r in p.runs:
            if BAN_MOTO in r.text:
                r.text = r.text.replace(BAN_MOTO, BAN)
                n_ban += 1
    judge(n_ban >= 1, f"版表記を {BAN} に改めた（{n_ban}か所）")

    d.save(SAKI)

    d2 = docx.Document(SAKI)
    ato = "\n".join([p.text for p in d2.paragraphs]
                    + [c.text for t in d2.tables for r in t.rows
                       for c in r.cells])

    print("  ── 検査1　県調査の値が入っているか")
    for go in (SHUTTEN, f"{A['在宅_老健']:.0f}人", haba(A["特養"]),
               f"{A['希望_特養_最初']:.0f}人から"):
        judge(go in ato, f"「{go[:24]}」がある")

    print("  ── 検査2　⚠ 軽費老人ホームの廃止が入っているか")
    judge(ato.count("軽費老人ホーム") >= 2,
          f"「軽費老人ホーム」が2-7と別表八にある（{ato.count('軽費老人ホーム')}件）")
    judge("3つの特徴" in ato and "2つの特徴" not in ato,
          "2-7 の「2つの特徴」を「3つの特徴」に改めた")
    judge("第3に、町内にあった軽費老人ホーム" in ato, "第3の特徴の段落がある")

    print("  ── 検査3　⚠ 稼働率に丸めていないか・残るものを書いているか")
    judge("把握していません" in ato or "把握していないものとして" in ato,
          "把握していないものを書いている")
    judge(ato.count("延べ利用日数による稼働率") >= 3,
          "「延べ利用日数による稼働率」ではない旨が3か所以上")
    judge("短期入所の稼働" in ato, "短期入所の稼働が空白である旨がある")
    judge("稼働率・休止床の数は公表の項目になく" not in ato,
          "旧の「稼働率・休止床の数は公表の項目になく」が0件")

    print("  ── 検査4　入所前の居所の扱い")
    judge("入所希望者名簿を5年間保管" in ato, "入所指針の記述がある")
    judge(ato.count("［要協議］") > mae.count("［要協議］"),
          f"［要協議］が増えた（{mae.count('［要協議］')}→{ato.count('［要協議］')}）")
    judge(ato.count("［要確認］") == mae.count("［要確認］") + 1,
          f"［要確認］が1件増えた（軽費の経緯）"
          f"（{mae.count('［要確認］')}→{ato.count('［要確認］')}）")

    print("  ── 検査5　⚠ 既にあった数が消えていないか")
    kyu_zen = "\n".join(KYU)
    kieta = []
    for m in re.finditer(r"[0-9０-９][0-9０-９,.]*(?:人|％|千円|件|名|歳|回)",
                         mae):
        go = m.group(0)
        if mae.count(go) > ato.count(go) and go not in kyu_zen:
            kieta.append(go)
    kieta = sorted(set(kieta))
    judge(not kieta, "消えた数はすべて差し替えた旧の文に現れている",
          "／".join(kieta[:5]))

    print("  ── 検査6　⚠ 事業所名・他団体名が入っていないか")
    for na in ("釜房みどりの園", "アルパイン", "蔵王グリーンホーム",
               "グループホームさわやか", "グループホームふれあい",
               "ながさか", "リリーハイム", "ケアハウスやまぶき", "八宮荘"):
        judge(na not in ato, f"「{na}」が入っていない")
    judge("川崎市" not in ato and "福岡" not in ato,
          "川崎市（神奈川県）・福岡県川崎町が混じらない")

    print("  ── 検査7　受託者を主語とする語・強調記号")
    aru = [w for w in ("受託者", "当社", "ビズアップ", "スクリプト")
           if w in ato]
    judge(not aru, "受託者を主語とする語が0件", "／".join(aru))
    judge("**" not in ato, "強調記号（**）が0件")

    print()
    print(f"  {SAKI}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        for m in futeki:
            print(f"      {m}")
        sys.exit(1)
    print(f"   ○ 検査1〜7　不適合0件（［要確認］{ato.count('［要確認］')}件・"
          f"［要協議］{ato.count('［要協議］')}件）")


if __name__ == "__main__":
    main()
