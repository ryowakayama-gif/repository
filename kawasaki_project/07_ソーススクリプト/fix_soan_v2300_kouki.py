# -*- coding: utf-8 -*-
"""協議用素案 Ver.2.29 → Ver.2.30（後期高齢者数の古い値の是正）
（令和8年10月23日）

  python3 07_ソーススクリプト/fix_soan_v2300_kouki.py

概要版を素案から組み直す際に見つけた食い違いの是正。
  2-1 の表（介護保険事業状況報告）：75歳以上 令和3年度末1,510人 → 令和7年度末1,742人（＋15.4％）
  3-4 の課題①：「75歳以上が4年で+10.9%（1,510→1,675人）」
  5-4 の本文：「後期高齢者75歳以上1,675人の増加局面」
1,675人は令和7年6月の月報の値（create_data_summary.py）で、2-1 を令和7年度末に
改めた際に、3-4・5-4 が古い値のまま残っていた。2-1 の値（年報で確かめた値。
check_nintei_R8.10.16.py）に揃える。値は 2-1 の表から読む。

終了コード　0＝検査1〜4すべて適合／1＝1件でも不適合
"""
import re
import shutil
import sys

import docx

MOTO = "01_第10期_最新版成果品/川崎町_計画書素案_v2.29_協議用.docx"
SAKI = "01_第10期_最新版成果品/川崎町_計画書素案_v2.30_協議用.docx"


def kakikae(p, s):
    p.runs[0].text = s
    for r in p.runs[1:]:
        r.text = ""


def zenbun(d):
    return "\n".join([p.text for p in d.paragraphs]
                     + [c.text for t in d.tables for r in t.rows for c in r.cells])


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("協議用素案 Ver.2.29 → Ver.2.30（後期高齢者数の古い値の是正）")
    shutil.copy2(MOTO, SAKI)
    d = docx.Document(SAKI)
    mae = zenbun(d)

    # 2-1 の表から読む
    t = [t for t in d.tables if t.rows[0].cells[0].text.strip() == "区分"
         and "令和7年度末" in t.rows[0].cells[2].text]
    gyo = [r for r in t[0].rows if r.cells[0].text.startswith("（再掲）後期高齢者")] if t else []
    judge(len(gyo) == 1, "2-1 の表の後期高齢者の行を特定")
    r3 = gyo[0].cells[1].text.strip()
    r7 = gyo[0].cells[2].text.strip()
    ritsu = gyo[0].cells[4].text.strip().replace("+", "＋")
    judge((r3, r7) == ("1,510人", "1,742人"), f"2-1 の値 {r3}→{r7}（{ritsu}）")
    r7n = r7.rstrip("人")

    print("  ── ① 3-4 課題①")
    KYU1 = "75歳以上が4年で+10.9%（1,510→1,675人）。"
    SHIN1 = f"75歳以上が4年で{ritsu}（{r3.rstrip('人')}→{r7}。令和3年度末→令和7年度末）。"
    n = 0
    for tb in d.tables:
        for r in tb.rows:
            for c in r.cells:
                if KYU1 in c.text:
                    kakikae(c.paragraphs[0], c.text.replace(KYU1, SHIN1))
                    n += 1
    judge(n == 1, f"3-4 の表を改めた（{n}か所）")

    print("  ── ② 5-4 本文")
    KYU2 = "後期高齢者75歳以上1,675人の増加局面"
    SHIN2 = f"後期高齢者（75歳以上。令和7年度末{r7}）の増加局面"
    hit = [p for p in d.paragraphs if KYU2 in p.text]
    judge(len(hit) == 1, "5-4 の旧の文がある")
    for p in hit:
        kakikae(p, p.text.replace(KYU2, SHIN2))

    d.save(SAKI)
    ato = zenbun(docx.Document(SAKI))
    print("  ── 検査1　古い値が残っていない")
    judge("1,675" not in ato and "10.9%" not in ato, "「1,675」「10.9%」0件")
    print("  ── 検査2　2-1 と同じ値")
    judge(ato.count(r7n) >= mae.count(r7n) + 2, f"「{r7n}」が2か所増えた")
    print("  ── 検査3　未確定の印")
    for w in ("［要確認］", "［要協議］", "［要内訳］"):
        judge(ato.count(w) == mae.count(w), f"{w} 増減なし")
    print("  ── 検査4　ほかの数が消えていない")
    kieta = sorted({m.group(0) for m in re.finditer(
        r"[0-9０-９][0-9０-９,.]*(?:人|％|%|千円|件|頁|点|円)", mae)
        if mae.count(m.group(0)) > ato.count(m.group(0))
        and m.group(0) not in KYU1 + KYU2})   # 旧の文の中の数（部分一致を含む）は除く
    judge(not kieta, "消えた数 0件（旧の2つを除く）", "／".join(kieta[:5]))
    print(f"  {SAKI}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 不適合0件")


if __name__ == "__main__":
    main()
