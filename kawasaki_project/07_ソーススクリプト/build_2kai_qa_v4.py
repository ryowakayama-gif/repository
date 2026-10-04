# -*- coding: utf-8 -*-
"""第2回策定委員会 想定問答集 v3 → v4（令和8年10月4日）

  ① **件数の食い違いを是正する。**
     v3 は「全35問」「カテゴリH（5問）」と書いているが、実際は39問・H は9問。
     QB-36〜39 を加えたときに、宣言した件数が追随していなかった。
  ② 素案の版の参照を更新する（Ver.1.25 → Ver.2.9）。29か所。
  ③ 第2回資料 v7 で加えた事項に対応する問答を5問加える。
     工程の切迫（開催日・パブリックコメントの期間）、
     第9期の計画値と実績値の対比、計画書の頁数、事業の内容の欄。

  python3 07_ソーススクリプト/build_2kai_qa_v4.py

**宣言した件数と実際の件数が合っていることを機械で確かめる。**
合わなければ終了コード1で終わる。
"""
import copy
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v3.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v4.docx"

VER_OLD, VER_NEW = "計画素案Ver.1.25", "計画素案Ver.2.9"

# ══════════════════════════ 加える問答
#   （カテゴリの見出しの冒頭, 問, 答, 根拠データ, 関連資料）
ADD = [
    ("カテゴリ F",
     "QB-40　【委員長】　計画書が100頁を超えるのは問題ないのか",
     "A．委託仕様書は「100頁程度」としています。現在の素案は本文109頁・"
     "資料編13頁の計122頁です。詳細な分析（交付金の枝番の分析、保険料の"
     "感度分析、認知症施策の追加候補）を別冊に移す圧縮案では本文102頁・"
     "計115頁となります。「100頁程度」を本文のみと読むのであれば、"
     "圧縮案の本文102頁で収まります。本文のみと読むか資料編を含めるかを、"
     "町とご相談のうえ定めたいと考えています。",
     "【根拠データ】 第2回資料 ８（計画書の頁数の見通しの表）。"
     "素案Ver.2.9 本文109頁、v2.9c 圧縮案 本文102頁、別冊「詳細分析」9頁",
     "【関連資料】 計画素案Ver.2.9、確認事項No.83"),
    ("カテゴリ F",
     "QB-41　【民生委員】　どの事業が何をするものなのか、計画書を見て分かるのか",
     "A．分かるようにしました。これまでの素案は事業の一覧表に「項・事業・"
     "区分・所管」しか載せておらず、事業の名称だけでは何をするものかが"
     "分かりませんでした。17の表・138事業すべてに「内容」の欄を加えています。"
     "内容は、第9期計画書の記述（57件）、統合体系表の新規・拡充の内容"
     "（33件）、介護保険法や国の指針が定める内容と素案の本文に記述がある"
     "もの（48件）から採っており、当方で作文したものはありません。"
     "48件については、実際の運用と食い違いがないか町にご確認いただきます。",
     "【根拠データ】 計画素案Ver.2.9 第5章・第6章の事業一覧表（17表138事業）",
     "【関連資料】 川崎町_事業の内容の出所_R8.10.4.xlsx、確認事項No.164"),
    ("カテゴリ G",
     "QB-42　【委員長】　本日の開催が遅れると、何が起きるのか",
     "A．納期に収まらなくなります。計画書・概要版の納品は令和9年3月15日で、"
     "そこから逆算すると、印刷・製本に3週間、第4回（答申）に2月中旬、"
     "パブリックコメントに30日間、計画書（案）の作成と県への事前協議に"
     "1か月余りを要します。現在の前提のままでは、本日の開催が11月11日以降に"
     "なると間に合いません。パブリックコメントを20日間にする、策定委員会を"
     "全3回にするなどの短縮が必要になります。本日、開催日と短縮の案を"
     "併せてご判断いただきたい理由がここにあります。",
     "【根拠データ】 第2回資料 ８（工程の短縮の案の表）。"
     "11月4日 ±0日／11月11日 ▲7日／11月18日 ▲14日／11月25日 ▲21日",
     "【関連資料】 川崎町_策定委員会_逆算工程表_R8.10.1.xlsx、"
     "確認事項No.85・No.49・No.19"),
    ("カテゴリ G",
     "QB-43　【住民代表】　パブリックコメントを20日間に短くしてよいのか",
     "A．国の「意見公募手続」の指針は30日間以上を原則としていますが、"
     "これは行政手続法が定める国の手続についてのものです。市町村が計画の"
     "案について行う意見公募は、町の実施要領の定めによります。"
     "本町の実施要領を確認したうえで、期間を定めたいと考えています。"
     "期間を短くする場合でも、広報・ホームページ・窓口での周知は"
     "同じように行い、いただいた意見は第3回又は第4回の策定委員会で"
     "対応方針とあわせてお示しします。",
     "【根拠データ】 第2回資料 ８（工程の短縮の案の表・案Ｂ）",
     "【関連資料】 確認事項No.19（パブリックコメントの取扱い）・"
     "No.136（実施要領）"),
    ("カテゴリ H",
     "QB-44　【有識者】　第9期は、計画どおりに進んだのか",
     "A．令和7年度でみると、給付費は計画10億1,102万円に対し実績9億9,822万円"
     "（達成率98.7％）で、総額ではほぼ計画どおりでした。ただし内わけでは、"
     "居宅サービスが計画を8.8％下回る一方、施設サービスは3.3％上回っています。"
     "総額が一致していることの内側で、在宅から施設・居住系への移行が"
     "進んでいます。第10期の見込量は、この移行が続くことを前提としています。"
     "なお、この計画値と実績値の対比は、交付金 推進 目標Ⅰ（ⅱ）2"
     "「事業計画の進捗状況」（12点・本町0点）が求める分析に当たります。",
     "【根拠データ】 第2回資料 ７（第9期の計画値と実績値の対比の表）。"
     "居宅 353,651→322,680千円／地域密着型 179,884→182,190千円／"
     "施設 477,485→493,350千円",
     "【関連資料】 計画素案Ver.2.9 第3章3-3（図3-2）、"
     "川崎町_第9期_計画値実績対比表_R8.10.1.xlsx"),
]


RE_CAT = re.compile(r"^カテゴリ\s*([A-H])")


def cat_of(s):
    """「カテゴリ F　計画骨子…」→「カテゴリ F」。

    見出しの文字数で切ると全角空白まで拾ってしまうため、
    記号で取り出す。
    """
    m = RE_CAT.match(s)
    return "カテゴリ " + m.group(1) if m else None


def txt(el):
    return "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()


def main():
    if not os.path.exists(SRC):
        raise SystemExit("v3 が見つからない：" + SRC)
    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body
    from fix_soan_v111 import set_el, set_cell

    # ══════════════════════════ ② 版の参照を更新
    n_ver = 0
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = "".join(n.text or "" for n in el.iter(qn("w:t")))
        if VER_OLD in s:
            set_el(el, s.replace(VER_OLD, VER_NEW))
            n_ver += 1
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                if VER_OLD in c.text:
                    set_cell(c, c.text.replace(VER_OLD, VER_NEW))
                    n_ver += 1

    # ══════════════════════════ ③ 問答を加える
    #   各カテゴリの最後の問答の後ろに置く
    kids = list(body.iterchildren())
    tpl_q = tpl_a = None
    for i, el in enumerate(kids):
        if el.tag != qn("w:p"):
            continue
        s = txt(el)
        if tpl_q is None and s.startswith("QB-1　"):
            tpl_q = copy.deepcopy(el)
            tpl_a = copy.deepcopy(kids[i + 1])
    if tpl_q is None:
        raise SystemExit("問答の雛形が見つからない")

    def last_of(cat):
        """そのカテゴリの最後の要素を返す。"""
        cur, last = None, None
        for el in body.iterchildren():
            if el.tag != qn("w:p"):
                continue
            s = txt(el)
            c = cat_of(s)
            if c:
                cur = c
            if cur == cat and s.startswith("【関連資料】"):
                last = el
        if last is None:
            raise SystemExit("カテゴリが見つからない：" + cat)
        return last

    n_add = 0
    for cat, q, a, kon, kan in ADD:
        after = last_of(cat)
        for text, tpl in ((kan, tpl_a), (kon, tpl_a), (a, tpl_a), (q, tpl_q)):
            el = copy.deepcopy(tpl)
            set_el(el, text)
            after.addnext(el)
        n_add += 1

    # ══════════════════════════ ① 件数の是正
    import collections
    full = []
    for el in body.iterchildren():
        if el.tag == qn("w:p"):
            full.append(txt(el))
    cat_cnt = collections.Counter()
    cur = None
    for s in full:
        c = cat_of(s)
        if c:
            cur = c
        if re.match(r"^QB-\d+　", s):
            cat_cnt[cur] += 1
    total = sum(cat_cnt.values())

    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = txt(el)
        m = re.match(r"^(カテゴリ [A-H]　.+?)（\d+問）$", s)
        if m:
            set_el(el, f"{m.group(1)}（{cat_cnt[cat_of(s)]}問）")
        if "全35問" in s:
            set_el(el, s.replace("全35問", f"全{total}問"))
    # 構成の表
    for t in doc.tables:
        h = [c.text.strip() for c in t.rows[0].cells]
        if h[:3] == ["カテゴリ", "テーマ", "問数"]:
            for r in t.rows[1:]:
                k = "カテゴリ " + r.cells[0].text.strip()
                if k in cat_cnt:
                    set_cell(r.cells[2], str(cat_cnt[k]))
    # 版
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = "".join(n.text or "" for n in el.iter(qn("w:t")))
        if "第３版" in s or "v3" in s:
            set_el(el, s.replace("第３版", "第４版").replace("v3", "v4"))

    doc.save(DST)

    # ══════════════════════════ 自己点検
    d2 = docx.Document(DST)
    text = "\n".join(q.text for q in d2.paragraphs)
    qs = re.findall(r"QB-(\d+)　", text)
    cat2 = collections.Counter()
    cur = None
    for p in d2.paragraphs:
        s = p.text.strip()
        c = cat_of(s)
        if c:
            cur = c
        if re.match(r"^QB-\d+　", s):
            cat2[cur] += 1
    ng = []
    if len(qs) != len(set(qs)):
        ng.append("QB の番号に重複がある")
    if f"全{len(qs)}問" not in text:
        ng.append(f"総数の宣言が全{len(qs)}問になっていない")
    for p in d2.paragraphs:
        m = re.match(r"^(カテゴリ ([A-H])　.+?)（(\d+)問）$", p.text.strip())
        if m and int(m.group(3)) != cat2["カテゴリ " + m.group(2)]:
            ng.append(f"カテゴリ{m.group(2)}の宣言{m.group(3)}問と実際"
                      f"{cat2['カテゴリ ' + m.group(2)]}問が合わない")
    for t in d2.tables:
        h = [c.text.strip() for c in t.rows[0].cells]
        if h[:3] == ["カテゴリ", "テーマ", "問数"]:
            s = sum(int(r.cells[2].text) for r in t.rows[1:]
                    if r.cells[2].text.strip().isdigit())
            if s != len(qs):
                ng.append(f"構成の表の合計{s}問と実際{len(qs)}問が合わない")
    if VER_OLD in text:
        ng.append("古い版の参照が残っている")

    print("第2回策定委員会 想定問答集 v4")
    print("  保存：", DST)
    print(f"  問答 {len(qs)}問（v3 は39問・加えた {n_add}問）")
    print("   " + "／".join(f"{k[-1]} {v}問" for k, v in sorted(cat2.items())))
    print(f"  素案の版の参照を更新 {n_ver}か所（{VER_OLD} → {VER_NEW}）")
    print("  ── 自己点検")
    if ng:
        for m in ng:
            print("   ×", m)
        sys.exit(1)
    print("   ○ 宣言した件数と実際の件数が一致し、古い版の参照もない")


if __name__ == "__main__":
    main()
