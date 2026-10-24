# -*- coding: utf-8 -*-
"""計画素案 Ver.2.26 → Ver.2.27（書体を第9期計画に揃え、協議用素案とする）
（令和8年10月23日）

  python3 07_ソーススクリプト/fix_soan_v2270_font.py

ご指示（令和8年10月23日・確認事項No.187）
  ・書体は第9期と同じ HG丸ゴシックM-PRO に揃える
  ・見出しと図表の題はそのまま（「1-1」「▌」「図3-1」）
  ・他の案件と同様に、協議用素案と公表用の計画で分ける運用とする

本スクリプトは協議用素案を作る。公表用は build_kouhyou_soan.py が
**この協議用素案を読んで組み替える**（内容の正本は協議用素案の1か所）。

(1) 書体　文書の既定・すべてのスタイル・すべての run・ヘッダー・フッターの
    游ゴシックを HG丸ｺﾞｼｯｸM-PRO に替える
    ⚠ Windows に登録されている書体名は半角カナの「HG丸ｺﾞｼｯｸM-PRO」である
      （第9期計画の PDF の埋め込み名は HGMaruGothicMPRO）。全角で書くと
      Word が見つけられず代わりの書体を当てる
    ⚠ 文字の大きさ・行送りは変えない（ご指示は書体のみ。本文10pt・行送り16pt）
    ⚠ 図（png）の中の文字は作図の書体のまま（この作業環境に同書体がない）
(2) 名称　表紙「計画書素案（令和8年10月時点）」→「協議用素案（令和8年10月時点）」
          ヘッダー「…（素案）」→「…（協議用素案）」
(3) 細い列　公表用の紙面の目視で、事業一覧の17表の「区分」（8.5mm）に「継続」
          「拡充」が収まらず縦に折れていた。サービス見込量の2表の「単位」
          （13.6mm）も「千円／年」が収まらない。4字以下の語が収まる幅
          （字の大きさ×字数＋セルの左右の余白3.8mm）に足し、同じ表の
          最も広い列から同じだけ引く（表の幅は変えない）

終了コード　0＝検査1〜5すべて適合／1＝1件でも不適合
"""
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

MOTO = "01_第10期_最新版成果品/川崎町_計画書素案_v2.26_版表記整理版.docx"
SAKI = "01_第10期_最新版成果品/川崎町_計画書素案_v2.27_協議用.docx"
KYU_FONT = "游ゴシック"
FONT = "HG丸ｺﾞｼｯｸM-PRO"
HYOSHI_KYU = "計画書素案（令和8年10月時点）"
HYOSHI = "協議用素案（令和8年10月時点）"
HEADER = "川崎町高齢者保健福祉計画・第10期介護保険事業計画（協議用素案）"
ATTR = ("ascii", "hAnsi", "eastAsia", "cs")


def kakikae_font(el):
    n = 0
    for rf in el.iter(qn("w:rFonts")):
        for a in ATTR:
            v = rf.get(qn("w:" + a))
            if v and KYU_FONT in v:
                rf.set(qn("w:" + a), FONT)
                n += 1
    return n


def hosoi_retsu(d):
    """4字以下の語が収まらない列を広げ、最も広い列から引く。"""
    n = 0
    for t in d.tables:
        cols = t._tbl.tblGrid.findall(qn("w:gridCol"))
        ws = [int(c.get(qn("w:w"))) for c in cols]
        k = len(ws)
        hirogeta = False
        for ci in range(k):
            need = 0
            for tr in t._tbl.findall(qn("w:tr"))[1:]:
                tcs = tr.findall(qn("w:tc"))
                if len(tcs) != k:
                    continue
                for p in tcs[ci].iter(qn("w:p")):
                    s = "".join(x.text or "" for x in p.iter(qn("w:t"))).strip()
                    if not s or len(s) > 4:
                        continue
                    sz = max([int(x.get(qn("w:val"))) / 2
                              for x in p.iter(qn("w:sz"))] or [8.5])
                    mm = (sum(0.5 if ch.isascii() else 1.0 for ch in s)
                          * sz * 0.3528 + 3.8 + 0.5)
                    need = max(need, int(mm / 25.4 * 1440))
            if need > ws[ci]:
                sa = need - ws[ci]
                hiro = max(range(k), key=lambda j: ws[j])
                ws[ci] += sa
                ws[hiro] -= sa
                hirogeta = True
        if not hirogeta:
            continue
        for c, w in zip(cols, ws):
            c.set(qn("w:w"), str(w))
        for tr in t._tbl.findall(qn("w:tr")):
            ichi = 0
            for tc in tr.findall(qn("w:tc")):
                gs = tc.tcPr.find(qn("w:gridSpan"))
                span = int(gs.get(qn("w:val"))) if gs is not None else 1
                tc.tcPr.find(qn("w:tcW")).set(qn("w:w"), str(sum(ws[ichi:ichi + span])))
                ichi += span
        n += 1
    return n


def zenbun(d):
    out = [p.text for p in d.paragraphs]
    out += [c.text for t in d.tables for r in t.rows for c in r.cells]
    for s in d.sections:
        for part in (s.header, s.footer):
            out += [p.text for p in part.paragraphs]
    return "\n".join(out)


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("計画素案 Ver.2.26 → Ver.2.27（書体を第9期に揃え、協議用素案とする）")
    print()
    if not os.path.exists(MOTO):
        print(f"   × 前の版がありません：{MOTO}")
        sys.exit(1)
    shutil.copy2(MOTO, SAKI)
    d = docx.Document(SAKI)
    mae = zenbun(d)
    sz_mae = [s.get(qn("w:val")) for s in d.element.iter(qn("w:sz"))]

    print("  ── ① 書体")
    n = kakikae_font(d.styles.element)
    n += kakikae_font(d.element)
    for s in d.sections:
        for part in (s.header, s.footer):
            n += kakikae_font(part._element)
    judge(n > 0, f"rFonts の属性 {n}か所を {FONT} に")

    print("  ── ② 協議用素案の名称")
    hit = [p for p in d.paragraphs if p.text.strip() == HYOSHI_KYU]
    judge(len(hit) == 1, "表紙の名称が1か所")
    for p in hit:
        p.runs[0].text = HYOSHI
        for r in p.runs[1:]:
            r.text = ""
    n_hd = 0
    for s in d.sections:
        for p in s.header.paragraphs:
            if "（素案）" in p.text:
                p.runs[0].text = HEADER
                for r in p.runs[1:]:
                    r.text = ""
                n_hd += 1
    judge(n_hd >= 1, f"ヘッダー {n_hd}か所")
    print("  ── ③ 細い列")
    n_h = hosoi_retsu(d)
    judge(n_h >= 19, f"広げた表 {n_h}（事業一覧17・サービス見込量2 ほか）")
    d.save(SAKI)

    d2 = docx.Document(SAKI)
    ato = zenbun(d2)
    print("  ── 検査1　游ゴシックが残っていない")
    import zipfile
    with zipfile.ZipFile(SAKI) as z:
        xml = "".join(z.read(n).decode("utf-8") for n in z.namelist()
                      if n.startswith("word/") and n.endswith(".xml")
                      and "theme" not in n)
    judge(KYU_FONT not in xml, "word/*.xml（テーマを除く）に游ゴシックが0件")
    judge(xml.count(FONT) > 100, f"{FONT} の指定 {xml.count(FONT)}か所")
    print("  ── 検査2　文字の大きさを変えていない")
    sz_ato = [s.get(qn("w:val")) for s in d2.element.iter(qn("w:sz"))]
    judge(sz_mae == sz_ato, f"w:sz {len(sz_ato)}か所が同じ")
    print("  ── 検査3　名称")
    judge(HYOSHI in ato and HEADER in ato and HYOSHI_KYU not in ato,
          "表紙・ヘッダーが協議用素案")
    print("  ── 検査4　文言・数値は名称のほかは変えていない")
    judge(mae.replace(HYOSHI_KYU, HYOSHI).replace(
        "川崎町高齢者保健福祉計画・第10期介護保険事業計画（素案）", HEADER) == ato,
        "全文が一致（名称の置き換えを除く）")
    print("  ── 検査5　見出し・図表の題はそのまま")
    for rx, mi in ((r"^\d+-\d+　", "節「1-1」"), (r"^図\d+-\d+", "図「図3-1」")):
        a = sum(1 for p in d.paragraphs if re.match(rx, p.text.strip()))
        b = sum(1 for p in d2.paragraphs if re.match(rx, p.text.strip()))
        judge(a == b and a > 0, f"{mi} {b}件（不変）")

    print()
    print(f"  {SAKI}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        for m in futeki:
            print(f"      {m}")
        sys.exit(1)
    print("   ○ 検査1〜5　不適合0件")


if __name__ == "__main__":
    main()
