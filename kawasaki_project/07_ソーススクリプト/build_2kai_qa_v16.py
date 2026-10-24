# -*- coding: utf-8 -*-
"""第2回策定委員会 想定問答集 v15 → v16（Red Team 再レビュー第3回（令和8年10月24日）の指摘の是正）

  python3 07_ソーススクリプト/build_2kai_qa_v16.py

資料 v19・素案 Ver.2.34 と同じ言い回しに揃える。問の数（全69問）は変えない。
  R3-01 QB-44「在宅から施設・居住系への移行が進んでいます」「この移行が続くことを前提」の断定
  R3-05 QB-58 開催回数の変更の根拠に委託仕様書11（定めのない事項の協議）を援用していた
"""
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v15.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v16.docx"
SU = 69


def txt(el):
    return "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会 想定問答集 v15 → v16（Red Team 再レビュー第3回の指摘の是正）")
    NAOSU = [
        ("総額が一致していることの内側で、在宅から施設・居住系への移行が進んでいます。第10期の見込量は、この移行が続く"
         "ことを前提としています。",
         "この乖離には、供給体制や町外施設の利用状況が影響した可能性がありますが、在宅から施設等への移行の実態及び原因は"
         "確かめられていません（町内外別の利用者数や経年の推移等による追加の検証を要します）。第10期の見込量は令和7年度の"
         "利用実績を出発点としているため、この利用の構成が続くことを前提としています。"),
        ("仕様書11は「本仕様の定めのない事項については、必要に応じ協議の上実施する」としていますので、ご協議のうえ"
         "お決めいただきたく存じます。",
         "回数は委託仕様書6(3)に「4回開催の予定」と定められた事項ですので、全3回とすることは予定の変更に当たります。"
         "町の契約担当で契約上の手続（変更契約又は書面による合意の要否）をご確認いただき、委員会のご意見を踏まえて"
         "町が書面で決定します。"),
    ]
    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    from fix_soan_v111 import set_el
    els = list(doc.element.body.iter(qn("w:p")))
    for kyu, shin in NAOSU:
        hit = [el for el in els if kyu in txt(el)]
        judge(len(hit) == 1, f"「{kyu[:24]}」（{len(hit)}か所）")
        for el in hit:
            set_el(el, txt(el).replace(kyu, shin))
    doc.save(DST)

    d2 = docx.Document(DST)
    zen = "\n".join(txt(el) for el in d2.element.body.iter(qn("w:p")))
    print("  ── 検査1　件数")
    ns = [int(x) for x in re.findall(r"QB-(\d+)　", zen)]
    judge(sorted(ns) == list(range(1, SU + 1)) and f"全{SU}問" in zen, f"QB-1〜{SU}・欠番なし")
    print("  ── 検査2　古い記述")
    for go in ("移行が進んでいます", "この移行が続く", "仕様書11は"):
        judge(go not in zen, f"「{go}」0件")
    print("  ── 検査3　事業所名・他団体名（受託者の語は事務局用の本書では v15 から増やさない）")
    for na in ("釜房", "アルパイン", "鶴寿会", "グループホームさわやか", "グループホームふれあい",
               "北塩原", "金ヶ崎", "大雪", "川崎市"):
        judge(na not in zen, f"「{na}」0件")
    zen0 = "\n".join(txt(el) for el in docx.Document(SRC).element.body.iter(qn("w:p")))
    for na in ("受託者", "当社", "ビズアップ"):
        judge(zen.count(na) == zen0.count(na), f"「{na}」{zen0.count(na)}→{zen.count(na)}")
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜3　不適合0件")


if __name__ == "__main__":
    main()
