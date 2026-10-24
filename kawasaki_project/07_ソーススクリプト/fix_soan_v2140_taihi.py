# -*- coding: utf-8 -*-
"""計画素案 Ver.2.14（他メンバー素案との対比による是正）

他メンバー作成の第10期計画素案（令和8年10月16日受領）と当方の正本を
突き合わせて分かった**当方の欠落・修正**を当てる。
Ver.2.13 に当てて Ver.2.14 を作る。

  入力　01_第10期_最新版成果品/川崎町_計画書素案_v2.13_供給体制整理版.docx
  出力　01_第10期_最新版成果品/川崎町_計画書素案_v2.14_見込量補完版.docx

是正1　9-1 の見込量の表に落ちていた15サービスを加える（**いちばん重い**）
  ⚠ 介護保険法第117条第2項第1号は「各年度における介護給付等対象サービスの
  **種類ごとの**量の見込み」を定めることを求めている。
  当方の表は22サービスのみで、**総給付費の8.5％（84,370千円）分の
  サービスが載っていなかった。** 落ちていたのは
    介護給付　短期入所療養介護（老健）・特定福祉用具購入費・住宅改修費・
              居宅介護支援
    介護予防給付　11サービス（介護予防訪問入浴介護・介護予防訪問看護・
              介護予防訪問リハビリテーション・介護予防通所リハビリテーション・
              介護予防短期入所生活介護・介護予防短期入所療養介護・
              介護予防福祉用具貸与・特定介護予防福祉用具購入費・
              介護予防住宅改修・介護予防認知症対応型共同生活介護・
              介護予防支援）
  ⚠ **数値は持っていた。** 計画見込量のワークブックの
  ③介護予防サービス・④介護サービスのシートから読む。
  総給付費998,222千円と保険料の算定には影響しない（記載の欠落である）。

是正2　9-3 の所得段階表の判定基準額を80.9万円に改める
  当方の素案の中で3通りになっていた。
    2-2 の所得段階別被保険者数の表　年金80.9万円以下（正しい）
    1-5 の制度改正の表　80.9万円→82.65万円の引上げ（政令は改正済み）
    9-3 の所得段階別保険料の表　**年金収入等80万円以下（古い）**
  他メンバー作成の素案も80.9万円としており、注に
  「基準額（80.9万円）は令和７年度以降の金額です」と書いている。

是正3　5-5 の「地域密着型介護老人福祉施設入居者生活介護」を正す
  ⚠ 法令上の名称は「地域密着型介護老人福祉施設**入所者**生活介護」である
  （介護保険法第8条第22項）。「特定施設入居者生活介護」は「入居者」で
  正しいため取り違えやすい。
  **この誤りは第9期計画書（3件）に始まり、統合体系表（D134）を経て
  当方の素案と他メンバーの素案の双方に伝播していた。**
  統合体系表は別に直す（fix_taikei_D134_R8.10.16.py）。

是正4　認知症対応型共同生活介護の範囲の注
  9-1 の表3は「認知症対応型共同生活介護」給付費80,059千円・月24.5人
  （介護予防を含まない）、表10は「認知症対応型共同生活介護
  （介護予防を含む）」月25.9人としており、**同じ節の中で範囲が違った。**
  是正1で介護予防給付を別建てにしたため、内訳の注を添える。

⚠ **保険料・総給付費・標準給付費の数値は一切変えていない。**
⚠ **他メンバー素案からの内容の取り込み（条文の抜粋・総人口の推移・
  件数の列・調査結果）は本版では行わない。** 出典の揃え方と頁数が
  町のご判断（確認事項No.175・No.177）を要するためである。
"""
import copy
import importlib.util
import os
import re
import shutil
import sys
import zipfile

import docx
import openpyxl
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from fix_soan_v111 import clone_table, set_el, set_tc        # noqa: E402

SRC = "01_第10期_最新版成果品/川崎町_計画書素案_v2.13_供給体制整理版.docx"
DST = "01_第10期_最新版成果品/川崎町_計画書素案_v2.14_見込量補完版.docx"
WB = "05_試算・管理シート/川崎町_第10期_計画見込量_R8.9.15.xlsx"

# 是正1　加えるサービス
#   （シート, サービス名, 素案に載せる名称, 区分）
KUWAERU = [
    ("④介護サービス", "短期入所療養介護（老健）",
     "短期入所療養介護（介護老人保健施設）", "居宅"),
    ("④介護サービス", "特定福祉用具購入費", "特定福祉用具購入費", "居宅"),
    ("④介護サービス", "住宅改修費", "住宅改修", "居宅"),
    ("④介護サービス", "介護予防支援・居宅介護支援", "居宅介護支援", "居宅"),
    ("③介護予防サービス", "介護予防訪問入浴介護",
     "介護予防訪問入浴介護", "予防"),
    ("③介護予防サービス", "介護予防訪問看護", "介護予防訪問看護", "予防"),
    ("③介護予防サービス", "介護予防訪問リハビリテーション",
     "介護予防訪問リハビリテーション", "予防"),
    ("③介護予防サービス", "介護予防通所リハビリテーション",
     "介護予防通所リハビリテーション", "予防"),
    ("③介護予防サービス", "介護予防短期入所生活介護",
     "介護予防短期入所生活介護", "予防"),
    ("③介護予防サービス", "介護予防短期入所療養介護（老健）",
     "介護予防短期入所療養介護（介護老人保健施設）", "予防"),
    ("③介護予防サービス", "介護予防福祉用具貸与",
     "介護予防福祉用具貸与", "予防"),
    ("③介護予防サービス", "特定介護予防福祉用具購入費",
     "特定介護予防福祉用具購入費", "予防"),
    ("③介護予防サービス", "介護予防住宅改修", "介護予防住宅改修", "予防"),
    ("③介護予防サービス", "介護予防認知症対応型共同生活介護",
     "介護予防認知症対応型共同生活介護", "予防"),
    ("③介護予防サービス", "介護予防支援", "介護予防支援", "予防"),
]

# 是正2
KIJUN_FURUI = [
    ("世帯非課税で年金収入等80万円以下", "世帯非課税で年金収入等80.9万円以下"),
    ("世帯非課税／年金収入等80万円超120万円以下",
     "世帯非課税／年金収入等80.9万円超120万円以下"),
    ("世帯課税・本人非課税／年金収入等80万円以下",
     "世帯課税・本人非課税／年金収入等80.9万円以下"),
    ("世帯課税・本人非課税／年金収入等80万円超",
     "世帯課税・本人非課税／年金収入等80.9万円超"),
]
CHU_93 = (
    "⚠ 判定に用いる年金収入等の基準額は、令和7年度以降は80.9万円です"
    "（表の区分はこれによります）。第10期からは政令の改正により"
    "82.65万円に引き上げられます（1-5）。条例の改正の手続を要します。")

# 是正3
NA_FURUI = "地域密着型介護老人福祉施設入居者生活介護"
NA_ATARASHII = "地域密着型介護老人福祉施設入所者生活介護"

# 是正4
CHU_91 = (
    "⚠ 本節の表は、介護給付と介護予防給付を分けて掲げています"
    "（介護保険法第117条第2項第1号は介護給付等対象サービスの種類ごとの"
    "量の見込みを定めることを求めています）。認知症対応型共同生活介護は"
    "介護給付が月24.5人、介護予防給付が月1.4人で、合わせて月25.9人です。"
    "必要利用定員総数の表（本節）は介護予防給付を含む月25.9人によります。")


def yomu_wb():
    """ワークブックから サービス名→(令和7・9・10・11年度) を読む。"""
    wb = openpyxl.load_workbook(WB, data_only=True)
    out = {}
    for sh in ("③介護予防サービス", "④介護サービス"):
        ws = wb[sh]
        for r in range(2, ws.max_row + 1):
            na = str(ws.cell(r, 1).value or "").strip()
            ko = str(ws.cell(r, 2).value or "").strip()
            if not na or not ko:
                continue
            vs = [ws.cell(r, c).value for c in range(3, 7)]
            if all(isinstance(v, (int, float)) for v in vs):
                out.setdefault((sh, na), {})[ko] = vs
    return out


def ke(v, tan):
    """表に書く形に整える。"""
    if tan == "千円":
        return f"{v:,.0f}"
    return f"{v:.1f}"


def main():
    if not os.path.exists(SRC):
        raise SystemExit("入力がない：" + SRC)
    wbv = yomu_wb()
    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body

    # ── 雛形
    tpl_chu = None
    tpl_tbl = None
    for el in body.iterchildren():
        if el.tag == qn("w:p") and tpl_chu is None:
            s = "".join(x.text or "" for x in el.iter(qn("w:t"))).strip()
            if s.startswith("⚠ 宮城県自身の評価"):
                tpl_chu = copy.deepcopy(el)
    cur = ""
    hyo4 = None     # 9-1 の居宅の表
    hyo10 = None    # 9-1 の必要利用定員総数の表
    for el in body.iterchildren():
        if el.tag == qn("w:p"):
            t = "".join(x.text or "" for x in el.iter(qn("w:t")))
            m = re.match(r"^(\d+-\d+)[　 ]", t.strip())
            if m:
                cur = m.group(1)
        elif el.tag == qn("w:tbl") and cur == "9-1":
            tb = docx.table.Table(el, doc)
            hd = [c.text.replace("\n", "").strip() for c in tb.rows[0].cells]
            if hd and hd[0] == "サービス種類" and "令和7年度実績" in " | ".join(hd):
                first = tb.rows[1].cells[0].text.strip()
                if first == "訪問介護":
                    hyo4 = el
                    tpl_tbl = copy.deepcopy(el)
            if "町内の定員" in " | ".join(hd):
                hyo10 = el
    if tpl_chu is None or hyo4 is None or tpl_tbl is None:
        raise SystemExit("雛形・9-1 の居宅の表が見つからない")

    # ══════ 是正1　介護予防給付の表を新設し、居宅の表に4件を加える
    # (a) 居宅の表の末尾に介護給付の4件を加える
    tb4 = docx.table.Table(hyo4, doc)
    n_gyo = 0
    for sh, na, hyoki, kubun in KUWAERU:
        if kubun != "居宅":
            continue
        d = wbv.get((sh, na))
        if d is None:
            raise SystemExit(f"ワークブックにない：{sh}／{na}")
        koms = [k for k in d if "給付費" in k]
        tmpl_row = tb4.rows[-1]._tr
        for i, ko in enumerate(["給付費（千円）"]
                               + [k for k in d if "給付費" not in k]):
            k = koms[0] if "給付費" in ko else ko
            vs = d[k]
            tan = "千円" if "給付費" in ko else ""
            tr = copy.deepcopy(tmpl_row)
            tb4._tbl.append(tr)
            row = tb4.rows[-1]
            cs = row.cells
            set_tc(cs[0]._tc, hyoki if i == 0 else "")
            set_tc(cs[1]._tc, "給付費" if "給付費" in ko else
                   ("回（日）数" if "回" in k or "日" in k else "利用者数"))
            set_tc(cs[2]._tc, "千円／年" if "給付費" in ko else
                   ("回" if "回" in k else ("日" if "日" in k else "人／月")))
            for j, v in enumerate(vs):
                set_tc(cs[3 + j]._tc,
                       ke(v, "千円" if "給付費" in ko else ""))
            # 令和17・22年度の欄は空ける（中長期の推計は別途）
            for j in (7, 8):
                if j < len(cs):
                    set_tc(cs[j]._tc, "―")
            n_gyo += 1

    # (b) 介護予防給付の表を新設して居宅の表の後ろに置く
    yobo = [(sh, na, hyoki) for sh, na, hyoki, k in KUWAERU if k == "予防"]
    rows = [["サービス種類", "項目", "単位", "令和7年度実績", "令和9年度",
             "令和10年度", "令和11年度", "令和17年度", "令和22年度"]]
    n_yobo = 0
    for sh, na, hyoki in yobo:
        d = wbv.get((sh, na))
        if d is None:
            raise SystemExit(f"ワークブックにない：{sh}／{na}")
        kyu = [k for k in d if "給付費" in k][0]
        rows.append([hyoki, "給付費", "千円／年"]
                    + [ke(v, "千円") for v in d[kyu]] + ["―", "―"])
        for ko, vs in d.items():
            if "給付費" in ko:
                continue
            tan = "回" if "回" in ko else ("日" if "日" in ko else "人／月")
            na2 = "回（日）数" if ("回" in ko or "日" in ko) else "利用者数"
            rows.append(["", na2, tan] + [ke(v, "") for v in vs]
                        + ["―", "―"])
        n_yobo += 1
    tbl_yobo = clone_table(tpl_tbl, rows)
    hyo4.addnext(tbl_yobo)
    # 表のキャプションと注
    cap = copy.deepcopy(tpl_chu)
    set_el(cap, "⚠ 介護予防給付のサービスの見込量（令和17・22年度の推計は"
                "別途）。介護給付と分けて掲げています。")
    tbl_yobo.addnext(cap)

    # ══════ 是正4　9-1 に範囲の注を置く
    e = copy.deepcopy(tpl_chu)
    set_el(e, CHU_91)
    hyo4.addprevious(e)

    # ══════ 是正2　9-3 の所得段階表の基準額
    n_kijun = 0
    for el in body.iterchildren():
        if el.tag != qn("w:tbl"):
            continue
        tb = docx.table.Table(el, doc)
        for r in tb.rows:
            for c in r.cells:
                for a, b in KIJUN_FURUI:
                    if a in c.text:
                        set_tc(c._tc, c.text.replace(a, b))
                        n_kijun += 1
    if n_kijun == 0:
        raise SystemExit("9-3 の所得段階表の基準額が見つからない")

    # ══════ 是正3　名称の誤り
    n_na = 0
    for el in body.iter(qn("w:t")):
        if el.text and NA_FURUI in el.text:
            el.text = el.text.replace(NA_FURUI, NA_ATARASHII)
            n_na += 1
    if n_na == 0:
        raise SystemExit(f"「{NA_FURUI}」が見つからない")

    # ══════ 版表記
    n_ver = 0
    for el in body.iter(qn("w:t")):
        if el.text and "Ver.2.13" in el.text:
            el.text = el.text.replace("Ver.2.13", "Ver.2.14")
            n_ver += 1
    if n_ver == 0:
        raise SystemExit("本文に Ver.2.13 の表記が見つからない")
    doc.save(DST)

    # 9-3 の注を所得段階表の後ろに置く（保存後に再度開いて位置を決める）
    doc = docx.Document(DST)
    body = doc.element.body
    tpl2 = None
    for el in body.iterchildren():
        if el.tag == qn("w:p"):
            s = "".join(x.text or "" for x in el.iter(qn("w:t"))).strip()
            if s.startswith("⚠ 宮城県自身の評価"):
                tpl2 = copy.deepcopy(el)
                break
    n_chu93 = 0
    for el in list(body.iterchildren()):
        if el.tag != qn("w:tbl"):
            continue
        tb = docx.table.Table(el, doc)
        if any("基準額に対する割合" in c.text
               for c in tb.rows[0].cells) and n_chu93 == 0:
            e = copy.deepcopy(tpl2)
            set_el(e, CHU_93)
            el.addnext(e)
            n_chu93 += 1
    if n_chu93 != 1:
        raise SystemExit(f"9-3 の所得段階表が{n_chu93}件（1件のはず）")
    doc.save(DST)
    naoshi_header(DST)

    # ══════════════════════════ 自己点検
    d2 = docx.Document(DST)
    d0 = docx.Document(SRC)
    zen = "".join(n.text or "" for n in d2.element.body.iter(qn("w:t")))
    ng = []
    if NA_FURUI in zen:
        ng.append(f"「{NA_FURUI}」が残っている")
    for a, _b in KIJUN_FURUI:
        if a in zen:
            ng.append(f"所得段階の基準額が古いまま：{a}")
    for _sh, _na, hyoki, _k in KUWAERU:
        if hyoki not in zen:
            ng.append(f"加えたサービスが本文にない：{hyoki}")
    for v in ("6,822", "998,222", "3,313,925", "152,500", "9,743.1"):
        if v not in zen:
            ng.append(f"主な数値が消えている：{v}")
    if len(d2.tables) != len(d0.tables) + 1:
        ng.append(f"表の数が{len(d0.tables)}→{len(d2.tables)}（1つ増えるはず）")
    for w in ("**", "__"):
        if w in zen:
            ng.append(f"強調記号が残っている：{w}")
    z = zipfile.ZipFile(DST)
    hdr = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>",
                             z.read("word/header1.xml").decode("utf-8")))
    if "Ver.2.14" not in hdr:
        ng.append(f"ヘッダーの版表記が更新されていない：{hdr}")

    print("計画素案 Ver.2.14（見込量補完版）")
    print("保存：", DST)
    print(f"  段落 {len(d2.paragraphs)}（Ver.2.13 は {len(d0.paragraphs)}）"
          f"／表 {len(d2.tables)}（Ver.2.13 は {len(d0.tables)}）")
    print(f"  是正1　見込量の表に加えたサービス　"
          f"介護給付 {n_gyo and len([x for x in KUWAERU if x[3] == '居宅'])}件"
          f"／介護予防給付 {n_yobo}件（表を1つ新設）")
    print(f"  是正2　所得段階の判定基準額を80.9万円に改めた　{n_kijun}か所"
          f"／注を1か所")
    print(f"  是正3　地域密着型介護老人福祉施設入所者生活介護に正した　"
          f"{n_na}か所")
    print("  是正4　9-1 に介護給付と介護予防給付の範囲の注を置いた")
    print("  ── 自己点検")
    if ng:
        for m in ng:
            print("   ×", m)
        sys.exit(1)
    print("   ○ 落ちていた15サービスが本文に入った")
    print("   ○ 所得段階の判定基準額が2-2・1-5 と揃った")
    print("   ○ 法令上の名称に正した")
    print("   ○ 保険料・総給付費・標準給付費の数値は変えていない")
    print("  ⚠ 目次の頁番号は fill_toc_pages_v200.py で作り直してください。")
    print("  ⚠ 他メンバー素案からの内容の取り込み（条文の抜粋・総人口の推移・"
          "件数の列・調査結果）は本版では行っていません"
          "（確認事項No.175・No.177のご判断を待ちます）。")


def naoshi_header(path):
    z = zipfile.ZipFile(path)
    names = z.namelist()
    data = {n: z.read(n) for n in names}
    z.close()
    n_fix = 0
    for n in names:
        if not re.match(r"word/(header|footer)\d*\.xml", n):
            continue
        s = data[n].decode("utf-8")
        s2 = re.sub(r"Ver\.2\.13", "Ver.2.14", s)
        if s2 != s:
            data[n] = s2.encode("utf-8")
            n_fix += 1
    if n_fix == 0:
        raise SystemExit("ヘッダーの版表記（Ver.2.13）が見つからない")
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    shutil.move(tmp, path)


if __name__ == "__main__":
    main()
