# -*- coding: utf-8 -*-
"""保険料の県内比較分析から、他団体への言及を除く（令和8年10月19日）

  python3 07_ソーススクリプト/fix_kennai_hikaku_R8.10.19.py

なぜ要るか
  CLAUDE.md 第3節　**他団体の固有名称を成果品に残さない。**
  `01_第10期_最新版成果品/川崎町_保険料_県内比較分析・要因分析.docx` の
  本文1か所に、他団体（別の案件）での方針への言及がありました。

  ⚠ **この文書そのものは川崎町の分析です。**他団体名が現れるのは
  1か所の注記だけですので、文書を移すのではなく、その1文を書き改めます。
  （他団体が主題の資料は 10_内部資料_非成果品/ へ移しました。）

  言いたいことは「断定的な水準表現を避ける」という書き方の方針であり、
  それは他団体の名を出さずに書けます。

  終了コード　0＝書き改めて点検に通った／1＝不適合
"""
import os
import re
import sys

import docx

DOC = "01_第10期_最新版成果品/川崎町_保険料_県内比較分析・要因分析.docx"

KYU = ("※ この区別は金ヶ崎町案件で『県内最低水準の保険料』という表現を"
       "将来推計との整合リスクから保留とした方針と同一の考え方。"
       "断定的な水準表現は避け、構造的根拠とともに記述する。")
SHIN = ("※ 『県内最低水準の保険料』のような断定的な水準表現は、"
        "将来推計との整合が崩れたときに説明が立たなくなるため用いない。"
        "水準そのものではなく、その水準になっている構造的な根拠を"
        "示す書き方とする。")

NG = re.compile(r"(金ケ崎町|金ヶ崎町|大雪地区広域連合|大雪広域|東川町|"
                r"東神楽町|上川町|美瑛町|ジャパン総研|福祉工房|"
                r"東京コンサルティング|川崎市)")


def main():
    if not os.path.exists(DOC):
        raise SystemExit("文書がない：" + DOC)
    d = docx.Document(DOC)

    atta = 0
    for p in d.paragraphs:
        if KYU in p.text:
            ts = list(p._p.iter(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
            if not ts:
                continue
            zen = "".join(t.text or "" for t in ts)
            ts[0].text = zen.replace(KYU, SHIN)
            ts[0].set("{http://www.w3.org/XML/1998/namespace}space",
                      "preserve")
            for t in ts[1:]:
                t.text = ""
            atta += 1

    if atta == 0:
        # 既に直っているなら、そのまま点検だけ行う
        print("書き改める箇所はありませんでした（既に直っています）")
    else:
        d.save(DOC)
        print(f"書き改めました　{DOC}（{atta}か所）")

    # ══════════════════════════ 点検
    print()
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    d2 = docx.Document(DOC)
    zen = [p.text for p in d2.paragraphs]
    for t in d2.tables:
        for r in t.rows:
            for c in r.cells:
                zen.append(c.text)
    s = "\n".join(zen)

    print("  ── 点検")
    atari = [x[:70] for x in zen if NG.search(x)]
    judge(not atari, "他団体・他社の固有名称が1つもない",
          f"→ {atari[:2]}" if atari else "")
    judge(SHIN[:20] in s, "書き改めた文が入っている")
    judge("県内" in s and "保険料" in s,
          "もとの論旨（県内比較・保険料）は残っている")
    judge(len(d2.paragraphs) == len(d.paragraphs),
          f"段落の数が変わっていない（{len(d2.paragraphs)}）")

    print()
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 不適合0件")


if __name__ == "__main__":
    main()
