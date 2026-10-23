# -*- coding: utf-8 -*-
"""サービス見込量の表に落ちがないかの検査（令和8年10月16日）

**他メンバー作成の第10期計画素案（令和8年10月16日受領）と当方の正本を
突き合わせたところ、当方の素案 9-1 のサービス見込量の表に載っている
サービスが22件しかなく、総給付費の8.5％（84,368千円）分のサービスが
表から落ちていることが分かった。**

  介護保険法第117条第2項第1号は「各年度における介護給付等対象サービスの
  **種類ごとの**量の見込み」を定めることを求めている。
  落ちているのは、居宅介護支援・介護予防支援・短期入所療養介護・
  住宅改修・特定福祉用具購入費及び**介護予防給付の13サービス**である。

  ⚠ **数値そのものは持っている。** 計画見込量のワークブック
  （05_試算・管理シート/川崎町_第10期_計画見込量_R8.9.15.xlsx）の
  ③介護予防サービス・④介護サービスのシートに全サービスが入っている。
  **素案の表に書き出していなかっただけである。**

  ⚠ 総給付費（998,222千円）は正しい。保険料の算定にも影響しない。
  **種類ごとの記載の欠落であり、法定記載事項の充足の問題である。**

本検査は、ワークブックに入っているサービスが素案の表にすべて現れているかを
機械で確かめる。**同じ落ちを二度と出さないための検査である。**

  python3 07_ソーススクリプト/check_mikomi_morasi_R8.10.16.py [素案のパス]

検査
  1　ワークブックのサービスが素案の見込量の表にすべてあるか
  2　素案の表の令和7年度の給付費が、ワークブックの値と一致するか
  3　素案の表に載っているサービスの給付費の合計が、総給付費と合うか
  4　素案に載っていないサービスの金額と、総給付費に対する割合
"""
import os
import re
import sys

import docx
import openpyxl
from docx.oxml.ns import qn

SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.30_協議用.docx"
WB = "05_試算・管理シート/川崎町_第10期_計画見込量_R8.9.15.xlsx"

# ワークブックのシートと、素案での扱い
#   「合算」＝素案では介護給付に合算して載せている（落ちではない）
#   ⚠ 当方の④シートの項目名「介護予防支援・居宅介護支援」は不正確である。
#     ④は介護給付の居宅介護支援であり、介護予防支援は③にある。
#   ⚠ **当方の素案は、介護予防給付をどこにも合算していない。**
#     9-1 の表3の「認知症対応型共同生活介護」は給付費80,059千円・
#     月24.5人で、**介護予防分（4,349千円・月1.4人）を含まない。**
#     一方、同じ 9-1 の表10（必要利用定員総数）は
#     「認知症対応型共同生活介護（介護予防を含む）月25.9人」としており、
#     **同じ節の中で同じ名称のサービスの範囲が違う**（24.5と25.9）。
#     これも是正を要する（Ver.2.14）。
GASSAN = {}
# 合計の行（サービスではない）
GOKEI_GYO = ("予防 給付費 合計", "介護 給付費 合計", "総給付費",
             "予防 在宅サービス", "介護 在宅サービス", "予防 居住系サービス",
             "介護 居住系サービス", "介護 施設サービス")


def yomu_wb():
    """ワークブックから サービス名→令和7年度の給付費（千円）を読む。"""
    wb = openpyxl.load_workbook(WB, data_only=True)
    out, sheet = {}, {}
    for sh in ("③介護予防サービス", "④介護サービス"):
        ws = wb[sh]
        for r in range(2, ws.max_row + 1):
            na = str(ws.cell(r, 1).value or "").strip()
            ko = str(ws.cell(r, 2).value or "").strip()
            v = ws.cell(r, 3).value
            if not na or "給付費" not in ko:
                continue
            if na in GOKEI_GYO:
                continue
            if isinstance(v, (int, float)):
                out[na] = float(v)
                sheet[na] = sh
    ws = wb["⑤総給付費"]
    sogo = None
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip() == "総給付費":
            sogo = float(ws.cell(r, 2).value)
    return out, sheet, sogo


def yomu_soan(path):
    """素案 9-1 の見込量の表から サービス名→令和7年度の給付費を読む。"""
    d = docx.Document(path)
    cur, out = "", {}
    for el in d.element.body.iterchildren():
        if el.tag == qn("w:p"):
            t = "".join(n.text or "" for n in el.iter(qn("w:t")))
            m = re.match(r"^(\d+-\d+)[　 ]", t.strip())
            if m:
                cur = m.group(1)
        elif el.tag == qn("w:tbl") and cur == "9-1":
            tb = docx.table.Table(el, d)
            hd = [c.text.replace("\n", "").strip() for c in tb.rows[0].cells]
            if "令和7年度実績" not in " | ".join(hd):
                continue
            try:
                ci = [i for i, h in enumerate(hd) if "令和7年度実績" in h][0]
                ki = [i for i, h in enumerate(hd) if h == "項目"][0]
            except IndexError:
                continue
            na = None
            for r in tb.rows[1:]:
                cs = [c.text.replace("\n", "").strip() for c in r.cells]
                if len(cs) <= max(ci, ki):
                    continue
                if cs[0]:
                    na = cs[0]
                if na and "給付費" in cs[ki]:
                    s = cs[ci].replace(",", "")
                    if re.fullmatch(r"-?\d+(?:\.\d+)?", s):
                        out.setdefault(na, float(s))
    return out


# ワークブックの名称と素案の表記の読み替え
#   ⚠ ワークブックの「介護予防支援・居宅介護支援」は、介護給付の
#     居宅介護支援である（介護予防支援は③シートにある）。
#     **ワークブックの項目名が不正確であり、素案では「居宅介護支援」とする。**
YOMIKAE = {
    "介護予防支援・居宅介護支援": "居宅介護支援",
    "短期入所療養介護（老健）": "短期入所療養介護（介護老人保健施設）",
    "介護予防短期入所療養介護（老健）":
        "介護予防短期入所療養介護（介護老人保健施設）",
}


def seiki(na):
    """素案の表記とワークブックの表記を突き合わせるための正規化。"""
    s = YOMIKAE.get(na, na)
    s = s.replace("（介護予防を含む）", "").replace("（老健）", "")
    s = s.replace("費", "").strip()
    return s


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else SOAN
    for p in (path, WB):
        if not os.path.exists(p):
            raise SystemExit("入力がない：" + p)
    wbv, sheet, sogo = yomu_wb()
    sv = yomu_soan(path)

    # ── 検査1　ワークブックのサービスが素案にあるか
    soan_key = {seiki(k): k for k in sv}
    nai, aru = [], []
    for k, v in sorted(wbv.items(), key=lambda x: -x[1]):
        kk = seiki(k)
        if kk in soan_key:
            aru.append((k, soan_key[kk], v))
        elif k in GASSAN and seiki(GASSAN[k]) in soan_key:
            aru.append((k, soan_key[seiki(GASSAN[k])] + "（合算）", v))
        else:
            nai.append((k, sheet[k], v))

    # ── 検査2　値の一致（合算のものは除く）
    chigai = []
    for k, sk, v in aru:
        if "（合算）" in sk:
            continue
        a = sv.get(sk)
        if a is not None and abs(a - v) > 1.0:
            chigai.append((k, v, a, a - v))

    # ── 検査3・4　合計
    soan_kei = sum(sv.values())
    ochi = sum(v for _, _, v in nai)

    ng = []
    if nai:
        ng.append(f"検査1　素案の見込量の表に載っていないサービスが"
                  f"{len(nai)}件ある（{ochi:,.0f}千円・総給付費の"
                  f"{ochi / sogo * 100:.1f}％）")
    for k, v, a, d in chigai:
        ng.append(f"検査2　{k}　ワークブック{v:,.0f}千円と"
                  f"素案{a:,.0f}千円が{d:+,.0f}千円食い違う")
    if abs(soan_kei + ochi - sogo) > 2.0:
        ng.append(f"検査3　素案の表の合計{soan_kei:,.0f}千円＋落ち"
                  f"{ochi:,.0f}千円＝{soan_kei + ochi:,.0f}千円が、"
                  f"総給付費{sogo:,.0f}千円と合わない")

    print("サービス見込量の表に落ちがないかの検査")
    print("  素案：", path.split("/")[-1])
    print("  ワークブック：", WB.split("/")[-1])
    print(f"  ワークブックのサービス {len(wbv)}件"
          f"／素案の表のサービス {len(sv)}件")
    print(f"  素案の表の令和7年度 給付費の合計 {soan_kei:,.0f}千円"
          f"／総給付費 {sogo:,.0f}千円"
          f"（差 {sogo - soan_kei:,.0f}千円・{(sogo - soan_kei) / sogo * 100:.1f}％）")
    if nai:
        print(f"  ── 素案の表に載っていないサービス {len(nai)}件")
        for k, sh, v in nai:
            print(f"   × {k:<34}{v:>10,.0f}千円　（{sh}）")
    if aru:
        print(f"  ── 素案の表に載っていて値が一致したもの "
              f"{len(aru) - len(chigai)}件")
    print("  ── 自己点検")
    if ng:
        for m in ng:
            print("   ×", m)
        print("  ⚠ 介護保険法第117条第2項第1号は「種類ごとの量の見込み」を"
              "定めることを求めています。落ちているサービスを"
              "素案 9-1 の表に書き出してください。"
              "数値はワークブックに入っています。")
        sys.exit(1)
    print("   ○ ワークブックのサービスはすべて素案の見込量の表にある")
    print("   ○ 素案の表の値はワークブックと一致している")
    print("   ○ 素案の表の合計が総給付費と合う")


if __name__ == "__main__":
    main()
