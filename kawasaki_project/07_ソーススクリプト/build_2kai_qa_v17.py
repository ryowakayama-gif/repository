# -*- coding: utf-8 -*-
"""第2回策定委員会 想定問答集 v16 → v17（Red Team 再レビュー第4回（令和8年10月24日）の指摘の是正）

  python3 07_ソーススクリプト/build_2kai_qa_v17.py

資料 v21・素案 Ver.2.36 と同じ値・言い回しに揃える。問の数（全69問）は変えない。
  F4-01 の波及　「標準給付費 令和6年度1,093,253千円→令和7年度1,084,587千円（0.8％減）」は基準が違っていた
               → 様式3 の支払額 1,092,621千円→1,084,587千円（0.7％減）
  F4-03 「１月４日開始なら納品は３月15日ちょうど」→ 案Ⅰでは３月16日・納期を１日超える（資料 v21 から読む）
  F4-06 「本文118頁」→ 全118頁（表紙・目次5頁と本文113頁）。別冊に移した案は全106頁・本文101頁
"""
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v16.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v17.docx"
SHIRYO = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v21.docx"
SU = 69


def txt(el):
    return "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()


def hankaku(s):
    return s.translate(str.maketrans("０１２３４５６７８９", "0123456789"))


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会 想定問答集 v16 → v17（Red Team 再レビュー第4回の指摘の是正）")
    s = docx.Document(SHIRYO)
    sz = "\n".join([p.text for p in s.paragraphs] + [c.text for t in s.tables for r in t.rows for c in r.cells])
    t = [t for t in s.tables if [c.text.strip() for c in t.rows[0].cells] == ["年　度", "保険給付費（支払額）", "前年比"]][0]
    r6, r7 = t.rows[1].cells[1].text.strip(), t.rows[2].cells[1].text.strip()
    gen = t.rows[2].cells[2].text.strip().lstrip("▲")
    m = re.search(r"案Ⅰでは納品が令和９年(.+?)となり、納期を(.)日超えます", sz)
    nohin, koe = hankaku(m.group(1)), hankaku(m.group(2))
    m2 = re.search(r"全(\d+)頁（表紙・ご挨拶・目次の(.)頁と本文(\d+)頁", sz)
    zen, mae, hon = m2.group(1), hankaku(m2.group(2)), m2.group(3)
    m3 = re.search(r"全体を(\d+)頁程度（本文(\d+)頁程度）", sz)
    zen_b, hon_b = m3.group(1), m3.group(2)
    judge(True, f"資料 v21 から：{r6}→{r7}（▲{gen}）・１月４日開始の納品 {nohin}（{koe}日超過）・"
                f"全{zen}頁＝前付{mae}頁＋本文{hon}頁・別冊案 全{zen_b}頁＝本文{hon_b}頁")

    NAOSU = [
        ("標準給付費は令和6年度1,093,253千円から令和7年度1,084,587千円へ0.8％減っています。",
         f"保険給付費（介護保険事業状況報告 様式3 の支払額。審査支払手数料を含まない）は令和6年度{r6}から令和7年度{r7}へ"
         f"{gen}減っています。"),
        ("標準給付費 令和6年度1,093,253千円→令和7年度1,084,587千円（▲0.8％）",
         f"保険給付費（様式3 の支払額）令和6年度{r6}→令和7年度{r7}（▲{gen}）"),
        ("この場合の納品は令和9年3月15日ちょうどになります。",
         f"この場合、後続の工程が6日ずつ遅れ、案Ⅰ（印刷3週間）では納品が令和9年{nohin}となり納期を{koe}日超えます。"
         "1月4日開始とするには、第3回から入稿まで又は印刷製本を1日以上短縮する必要があります。"),
        ("現在の素案は本文118頁・資料編（見込13頁）の計131頁です。",
         f"現在の素案は全{zen}頁（表紙・ご挨拶・目次{mae}頁と本文{hon}頁）で、資料編（見込13頁）を含めると計131頁です。"),
        ("を別冊に移すと、本文は106頁程度・計119頁程度となります。「100頁程度」を本文のみと読む場合でも、なお6頁程度の圧縮を要します。",
         f"を別冊に移すと、全{zen_b}頁程度（本文{hon_b}頁程度）・資料編を含め計119頁程度となります。「100頁程度」を本文のみと"
         "読む場合はほぼ収まり、冊子全体と読む場合はなお19頁程度の圧縮を要します。"),
        ("素案 本文118頁、別冊に移した場合 本文106頁程度",
         f"素案 全{zen}頁（本文{hon}頁）、別冊に移した場合 全{zen_b}頁程度（本文{hon_b}頁程度）"),
        ("A．超えます。現在の素案は本文118頁で、資料編（見込13頁）を加えると計131頁です。詳細な分析を別冊に移すと本文は106頁程度になります。",
         f"A．超えます。現在の素案は全{zen}頁（表紙・ご挨拶・目次{mae}頁と本文{hon}頁）で、資料編（見込13頁）を加えると計131頁"
         f"です。詳細な分析を別冊に移すと全{zen_b}頁程度（本文{hon_b}頁程度）になります。"),
        ("計画素案 本文118頁", f"計画素案 全{zen}頁（本文{hon}頁）"),
    ]
    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    from fix_soan_v111 import set_el
    els = list(doc.element.body.iter(qn("w:p")))
    for kyu, shin in NAOSU:
        hit = [el for el in els if kyu in txt(el)]
        judge(len(hit) >= 1, f"「{kyu[:26]}」（{len(hit)}か所）")
        for el in hit:
            set_el(el, txt(el).replace(kyu, shin))
    doc.save(DST)

    d2 = docx.Document(DST)
    zen_t = "\n".join(txt(el) for el in d2.element.body.iter(qn("w:p")))
    print("  ── 検査1　件数")
    ns = [int(x) for x in re.findall(r"QB-(\d+)　", zen_t)]
    judge(sorted(ns) == list(range(1, SU + 1)) and f"全{SU}問" in zen_t, f"QB-1〜{SU}・欠番なし")
    print("  ── 検査2　古い記述")
    for go in ("1,093,253", "0.8％", "ちょうど", "本文118頁", "本文は106頁", "本文106頁"):
        judge(go not in zen_t, f"「{go}」0件")
    print("  ── 検査3　事業所名・他団体名（受託者の語は事務局用の本書では v16 から増やさない）")
    for na in ("釜房", "アルパイン", "鶴寿会", "北塩原", "金ヶ崎", "大雪", "川崎市"):
        judge(na not in zen_t, f"「{na}」0件")
    zen0 = "\n".join(txt(el) for el in docx.Document(SRC).element.body.iter(qn("w:p")))
    for na in ("受託者", "当社", "ビズアップ"):
        judge(zen_t.count(na) == zen0.count(na), f"「{na}」{zen0.count(na)}→{zen_t.count(na)}")
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜3　不適合0件")


if __name__ == "__main__":
    main()
