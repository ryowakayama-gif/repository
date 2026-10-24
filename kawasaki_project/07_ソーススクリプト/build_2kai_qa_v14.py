# -*- coding: utf-8 -*-
"""第2回策定委員会 想定問答集 v13 → v14（Red Team レビュー（令和8年10月24日）の指摘の是正）

  python3 07_ソーススクリプト/build_2kai_qa_v14.py

資料 v17・素案 Ver.2.32 と同じ値・言い回しに揃える。問の数（全69問）は変えない。
  A04 報酬改定・負担割合・収納率の差額は算定Ａ（取崩しなし）の値。基準を書き、報酬改定は
      算定Ｂ・Ｃの差も書く（資料 v17 の表26 から計算）
  A06 「そのままイの要件を満たします」「記録し計画に書くだけで要件を満たす可能性」→ 評価の対象となり得る
  B03 「国は基金残高の概ね半分以下を取崩水準として示しており」→ 事務局案の考え方として書く
  B15 「町内14の指定」→「町内8事業所の14の指定」
"""
import re
import shutil
import sys
import os

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v13.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v14.docx"
SHIRYO = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v17.docx"
SU = 69


def txt(el):
    return "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()


def yen(s):
    return int(re.sub(r"[^\d]", "", s))


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会 想定問答集 v13 → v14（Red Team レビューの指摘の是正）")
    s = docx.Document(SHIRYO)
    t26 = [t for t in s.tables if t.rows[0].cells[0].text.strip() == "介護報酬改定率"][0]
    g = {r.cells[0].text.strip()[:2]: [yen(c.text) for c in r.cells[1:4]] for r in t26.rows[1:4]}
    d1 = [a - b for a, b in zip(g["１％"], g["０％"])]
    d3 = [a - b for a, b in zip(g["３％"], g["０％"])]
    judge(d1[0] == 66 and d3[0] == 196, f"資料 表26 から 1％ Ｂ＋{d1[1]}・Ｃ＋{d1[2]}／3％ Ｂ＋{d3[1]}・Ｃ＋{d3[2]}")

    NAOSU = [
        ("A．基金の取崩方針により",
         "改定率1％で月額66円、負担割合が24％となれば月額309円上がります。",
         "改定率1％で月額66円、負担割合が24％となれば月額309円上がります（いずれも基金の取崩しを"
         f"行わない算定Ａの場合。事務局案の算定Ｂでは、改定率1％で月額{d1[1]}円です）。"),
        ("A．改定率1％あたり月額約66円です。",
         "A．改定率1％あたり月額約66円です。3％の改定であれば約196円上がります。",
         "A．基金の取崩しを行わない算定Ａでは、改定率1％あたり月額約66円、3％で約196円上がります。"
         f"事務局案の算定Ｂでは1％で約{d1[1]}円・3％で約{d3[1]}円、算定Ｃでは1％で約{d1[2]}円・"
         f"3％で約{d3[2]}円です。"),
        ("【根拠データ】 第2回資料 ６（介護報酬改定率別の算定表",
         "改定率1％＝月額＋66円、3％＝＋196円",
         f"改定率1％＝月額＋66円、3％＝＋196円（算定Ａ）。算定Ｂ＋{d1[1]}円・＋{d3[1]}円、"
         f"算定Ｃ＋{d1[2]}円・＋{d3[2]}円"),
        ("A．現年度分だけを見れば",
         "99.6％を採ると保険料は月額246円下がりますが、",
         "99.6％を採ると保険料は月額246円下がりますが（算定Ａの場合）、"),
        ("A．配点としては大きくありません。",
         "すでに町が行っていることを記録し計画に書くだけで要件を満たす可能性があります。",
         "すでに町が行っていることを記録し、国の該当状況調査で回答できる状態にすれば、評価の対象と"
         "なり得ます。"),
        ("A．", "したがって、本日この資料で交付金の評価結果をご報告し、ご意見をいただくことが、"
               "そのままイの要件を満たします。",
         "したがって、本日この資料で交付金の評価結果をご報告し、ご意見をいただくことは、イの要件に"
         "当たる取組の一つとなり得ます（評価は国の該当状況調査への回答と、本日の議事の記録によります）。"),
        ("A．", "国は基金残高の概ね半分以下を取崩水準として示しており、事務局案は第9期と同じ考え方による"
               "50％取崩し（6,143円）です。",
         "事務局案は、基金残高の見通しと第11期以降の保険料の平準化を踏まえ、第9期と同じ考え方による"
         "50％取崩し（6,143円）を比較案の一つとして示すものです。"),
    ]
    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body
    from fix_soan_v111 import set_el
    for atama, kyu, shin in NAOSU:
        hit = [el for el in body.iterchildren()
               if el.tag == qn("w:p") and txt(el).startswith(atama) and kyu in txt(el)]
        judge(len(hit) == 1, f"『{atama[:16]}』の「{kyu[:22]}」")
        for el in hit:
            set_el(el, txt(el).replace(kyu, shin))
    n = 0
    for el in body.iterchildren():
        if el.tag == qn("w:p") and ("町内14の指定" in txt(el) or "・14の指定すべて" in txt(el)):
            set_el(el, txt(el).replace("町内14の指定", "町内8事業所の14の指定")
                   .replace("・14の指定すべて", "・8事業所の14の指定すべて"))
            n += 1
    judge(n == 4, f"14の指定 → 8事業所の14の指定（{n}か所）")
    doc.save(DST)

    d2 = docx.Document(DST)
    zen = "\n".join(p.text for p in d2.paragraphs) + "\n" + "\n".join(
        c.text for t in d2.tables for r in t.rows for c in r.cells)
    print("  ── 検査1　件数")
    ns = [int(x) for x in re.findall(r"QB-(\d+)　", zen)]
    judge(sorted(ns) == list(range(1, SU + 1)) and f"全{SU}問" in zen, f"QB-1〜{SU}・欠番なし")
    print("  ── 検査2　古い記述")
    for go in ("そのままイの要件を満たします", "書くだけで要件を満たす", "概ね半分以下", "町内14の指定"):
        judge(go not in zen, f"「{go}」0件")
    print("  ── 検査3　事業所名・他団体名")
    for na in ("釜房", "アルパイン", "鶴寿会", "グループホームさわやか", "グループホームふれあい",
               "北塩原", "金ヶ崎", "大雪", "川崎市"):
        judge(na not in zen, f"「{na}」0件")
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜3　不適合0件")


if __name__ == "__main__":
    main()
