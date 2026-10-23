# -*- coding: utf-8 -*-
"""事業所別の5か年推移（利用者数・定員充足・回転・従事者数）から町内の状況を整理する
（令和8年10月23日受領）

  python3 07_ソーススクリプト/build_jigyosho_5nen_R8.10.23.py

受領したもの（09_元資料/その他データ/事業所別5か年推移_R8.10.23受領/）
  全国47都道府県別の Excel 4種（原本は zip 3点と分割 7z 4片のまま退避）
    事業所別利用者数・減少率（訪問系・通所系・短期入所系・多機能系・施設・居住系）
    施設別定員充足率・減少率（施設・居住系）
    事業所別回転率・減少率（通所系）
    事業所別従事者数・減少率（看護職員等・介護職員の常勤・非常勤・常勤換算数）
  いずれも令和3年〜令和7年の5か年。宮城県分を 宮城県/ に取り出してある。
  ⚠ **出所の記載がない**（ファイルの作成者情報に個人名のみ）。「情報公表 項目名対応表」
    が付いており、介護サービス情報公表システムの公表項目から作られたものと考えられるが、
    作成主体と各年の時点は［要確認］とする。

⚠ 集計の欄（減少率・アラート・定員充足率・回転率）は数式で、計算結果が保存されて
  いない（読むと0）。本スクリプトで計算し直す。
⚠ 原表の「定員充足率」の式は「1月あたり延べ利用者数÷定員」であり、百分率ではなく
  **1床（1人分の定員）あたりの月の延べ利用日数**である。稼働率は暦日（365÷12日）で割って求める。
⚠ 同じ建物で複数の指定を持つ事業所は、従事者数が指定ごとに同じ人数で報告される
  （介護老人保健施設と短期入所療養介護など）。**町の合計として足すと重複する。**
⚠ 事業所名・法人名は成果品に書かない（CLAUDE.md・町内8事業所3法人）。
  管理シートはサービスの種類で示す。原表の抜き出し（09の CSV）には名称が残る。

入力　09_元資料/その他データ/事業所別5か年推移_R8.10.23受領/宮城県/04_宮城県_*.xlsx
出力　09_元資料/その他データ/事業所別5か年推移_川崎町_R8.10.23.csv（川崎町の行の抜き出し）
      05_試算・管理シート/川崎町_事業所別5か年推移_R8.10.23.xlsx
終了コード　0＝検査1〜6すべて適合／1＝1件でも不適合
"""
import csv
import glob
import os
import statistics
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

D = "09_元資料/その他データ/事業所別5か年推移_R8.10.23受領/宮城県"
CSV_OUT = "09_元資料/その他データ/事業所別5か年推移_川崎町_R8.10.23.csv"
OUT = "05_試算・管理シート/川崎町_事業所別5か年推移_R8.10.23.xlsx"
CODE = "04324"
NEN = ["令和3年", "令和4年", "令和5年", "令和6年", "令和7年"]
HI = 365 / 12            # 1月の暦日
SHU = {"11": "訪問介護", "13": "訪問看護", "14": "訪問リハビリテーション",
       "15": "通所介護", "16": "通所リハビリテーション",
       "21": "短期入所生活介護", "22": "短期入所療養介護（老健）",
       "32": "認知症対応型共同生活介護", "51": "介護老人福祉施設",
       "52": "介護老人保健施設", "54": "地域密着型介護老人福祉施設入所者生活介護",
       "78": "地域密着型通所介護", "72": "認知症対応型通所介護",
       "73": "小規模多機能型居宅介護", "33": "特定施設入居者生活介護"}


def fname(key):
    f = glob.glob(os.path.join(D, f"04_宮城県_*{key}*.xlsx"))
    if len(f) != 1:
        print(f"   × 入力が見つからない：{key}")
        sys.exit(1)
    return f[0]


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def hyo(path):
    """シートごとに (見出しの行の番号, 行のリスト) を返す。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        h = [i for i, r in enumerate(rows) if r and r[0] == "No."]
        if h:
            out[ws.title.strip()] = (h[0], rows)
    return out


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("事業所別の5か年推移（令和8年10月23日受領）の整理")
    print()
    riyo = hyo(fname("利用者数"))
    juso = hyo(fname("定員充足率"))
    kaiten = hyo(fname("回転率"))
    jusha = hyo(fname("従事者数"))

    # ── 抜き出し（長い形）
    nagai = []           # 指標, シート, 種類, 事業所番号, 事業所名, 年, 値
    ken = {}             # 県内の分布：(指標, 種類) → [値（令和7年）]

    def tsuika(shihyo, sheet, r, cols):
        for nen, c in zip(NEN, cols):
            nagai.append([shihyo, sheet, r[4], r[5], r[6], nen, num(r[c])])

    for sheet, (h, rows) in riyo.items():
        for r in rows[h + 1:]:
            if r[1] == CODE:
                tsuika("利用者数（実数・人／月）", sheet, r, range(14, 19))
    for sheet, (h, rows) in juso.items():
        for r in rows[h + 1:]:
            if r[1] is None:
                continue
            nobe, tei = num(r[18]), num(r[23])
            if nobe and tei:
                ken.setdefault(("稼働率", r[4]), []).append(nobe / tei / HI)
            if r[1] == CODE:
                tsuika("延べ利用者数（人日／月）", sheet, r, range(14, 19))
                tsuika("定員（人）", sheet, r, range(19, 24))
    for sheet, (h, rows) in kaiten.items():
        for r in rows[h + 1:]:
            if r[1] is None:
                continue
            nobe, tei = num(r[18]), num(r[23])
            if nobe and tei:
                ken.setdefault(("回転", r[4]), []).append(nobe / tei)
            if r[1] == CODE:
                tsuika("延べ利用者数（人回／月）", sheet, r, range(14, 19))
                tsuika("定員（人）", sheet, r, range(19, 24))
    for sheet, (h, rows) in jusha.items():
        for r in rows[h + 1:]:
            if r[1] is None:
                continue
            k3, k7 = num(r[38]), num(r[50])
            if k3 and k7 is not None:
                ken.setdefault(("介護職員の増減", r[4]), []).append((k7 - k3) / k3)
            if r[1] == CODE:
                for nen, i in zip(NEN, range(5)):
                    for j, yaku in enumerate(("常勤", "非常勤", "常勤換算")):
                        nagai.append([f"看護職員等（{yaku}）", sheet, r[4], r[5], r[6],
                                      nen, num(r[21 + i * 3 + j])])
                        nagai.append([f"介護職員（{yaku}）", sheet, r[4], r[5], r[6],
                                      nen, num(r[36 + i * 3 + j])])
    with open(CSV_OUT, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["指標", "シート", "サービス種類コード", "事業所番号", "事業所名",
                    "年", "値"])
        w.writerows(nagai)

    def ne(shihyo, jigyo, shu):
        return [x[6] for x in nagai if x[0] == shihyo and x[3] == jigyo
                and x[2] == shu]

    jigyo = sorted({(x[2], x[3]) for x in nagai if x[0].startswith("利用者数")},
                   key=lambda t: (t[0], t[1]))
    # 同じ種類が2つあるものは（1）（2）を付ける
    kazu = {}
    na = {}
    for shu, bango in jigyo:
        kazu[shu] = kazu.get(shu, 0) + 1
    seen = {}
    for shu, bango in jigyo:
        seen[shu] = seen.get(shu, 0) + 1
        na[(shu, bango)] = SHU.get(shu, shu) + (f"（{seen[shu]}）" if kazu[shu] > 1 else "")

    print("  ── 検査1　既に把握している公表値と一致するか")
    tei = {k: ne("定員（人）", k[1], k[0]) for k in jigyo}
    teiin = {}
    for (shu, b), v in tei.items():
        if v and v[-1]:
            teiin[shu] = teiin.get(shu, 0) + v[-1]
    judge(teiin.get("52") == 86, f"介護老人保健施設の定員 {teiin.get('52')}（＝86）")
    judge(teiin.get("51") == 30, f"介護老人福祉施設の定員 {teiin.get('51')}（＝30）")
    judge(teiin.get("54") == 30, f"地域密着型介護老人福祉施設の定員計 {teiin.get('54')}（＝30）")
    judge(teiin.get("32") == 27, f"認知症対応型共同生活介護の定員計 {teiin.get('32')}（＝27）")
    houmon = [x[6] for x in nagai if x[0] == "介護職員（常勤換算）" and x[2] == "11"
              and x[5] == "令和7年"]
    judge(houmon == [4.5], f"訪問介護の介護職員 常勤換算 令和7年 {houmon}（＝4.5）")
    print("  ── 検査2　抜き出しの件数")
    judge(len(jigyo) == 12, f"町内の指定 {len(jigyo)}（訪問2・通所2・短期入所2・施設居住6）")

    # ── 管理シート
    wb = openpyxl.Workbook()
    H = PatternFill("solid", fgColor="1F5C56")
    F = Font(name="游ゴシック", size=9)
    FB = Font(name="游ゴシック", size=9, bold=True, color="FFFFFF")

    def sheet(title, head, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(head)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.fill, c.font = H, FB
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = F
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
        return ws

    def zo(v):
        return round((v[-1] - v[0]) / v[0], 3) if v and v[0] and v[-1] is not None else None

    ws0 = wb.active
    ws0.title = "00_この表について"
    for r in [
        ["項目", "内容"],
        ["何の表か", "令和8年10月23日受領の事業所別5か年推移（全国47都道府県別）から、"
         "町内の12の指定（訪問2・通所2・短期入所2・施設・居住系6）の令和3年〜令和7年の値を整理したもの。"],
        ["出所", "［要確認］ファイルに出所の記載がない。介護サービス情報公表システムの公表項目"
         "（項目名対応表あり）から作られたものと考えられる。作成主体と各年の時点の確認を要する。"],
        ["稼働率", "原表の「定員充足率」は「1月あたり延べ利用者数÷定員」（1床あたり月の延べ利用日数）。"
         "本表の稼働率はこれを暦日（365÷12＝30.4日）で割ったもの。"],
        ["回転", "通所系の「回転率」は「1月あたり延べ利用者数÷定員」（定員1人分あたり月の延べ利用回数）。"],
        ["⚠ 従事者数の重複", "同じ建物の複数の指定は同じ職員を指定ごとに報告する"
         "（介護老人保健施設と短期入所療養介護、地域密着型介護老人福祉施設と短期入所生活介護）。"
         "町の合計として足さない。"],
        ["⚠ 県調査との違い", "宮城県「老人福祉施設等入所状況調」は時点の在所者数（3か月ごと）。"
         "本表は1月あたりの延べで、定義が違う。並べるときは定義を明記する。"],
        ["事業所名", "成果品には書かない。本表はサービスの種類で示す（同じ種類が2つあるものは（1）（2））。"],
    ]:
        ws0.append(r)
    for row in ws0.iter_rows():
        for c in row:
            c.font = F
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws0.column_dimensions["A"].width = 18
    ws0.column_dimensions["B"].width = 100

    rows = []
    for k in jigyo:
        v = ne("利用者数（実数・人／月）", k[1], k[0])
        rows.append([na[k]] + v + [zo(v)])
    sheet("01_利用者数", ["サービス"] + NEN + ["令和3→7年の増減率"], rows,
          [34] + [9] * 5 + [12])

    rows = []
    kado7 = {}
    for k in jigyo:
        nobe = ne("延べ利用者数（人日／月）", k[1], k[0])
        tv = ne("定員（人）", k[1], k[0])
        if not nobe:
            continue
        kd = [round(n / t / HI, 3) if n and t else None for n, t in zip(nobe, tv)]
        kado7[k] = kd[-1]
        kv = ken.get(("稼働率", k[0]), [])
        rows.append([na[k], tv[-1]] + kd + [round(statistics.median(kv), 3) if kv else None,
                                              len(kv)])
    sheet("02_施設居住系の稼働率", ["サービス", "定員（令和7年）"] + NEN
          + ["宮城県の中央値（令和7年）", "県内の数"], rows, [34, 10] + [9] * 5 + [12, 8])

    rows = []
    for k in jigyo:
        nobe = ne("延べ利用者数（人回／月）", k[1], k[0])
        tv = [x[6] for x in nagai if x[0] == "定員（人）" and x[3] == k[1]
              and x[2] == k[0] and x[1].startswith("事業所別回転")]
        if not nobe:
            continue
        kt = [round(n / t, 1) if n and t else None for n, t in zip(nobe, tv)]
        kv = ken.get(("回転", k[0]), [])
        rows.append([na[k], "／".join(f"{t:.0f}" for t in tv)] + kt
                    + [round(statistics.median(kv), 1) if kv else None, len(kv)])
    sheet("03_通所の回転", ["サービス", "定員（令和3〜7年）"] + NEN
          + ["宮城県の中央値（令和7年）", "県内の数"], rows, [30, 16] + [9] * 5 + [12, 8])

    rows = []
    for k in jigyo:
        for yaku in ("介護職員", "看護職員等"):
            v = [x[6] for x in nagai if x[0] == f"{yaku}（常勤換算）" and x[3] == k[1]
                 and x[2] == k[0]]
            if not v or all(x is None for x in v):
                continue
            kv = ken.get(("介護職員の増減", k[0]), []) if yaku == "介護職員" else []
            rows.append([na[k], yaku] + v + [zo(v),
                         round(statistics.median(kv), 3) if kv else None])
    sheet("04_従事者の常勤換算", ["サービス", "職種"] + NEN
          + ["令和3→7年の増減率", "宮城県の中央値（介護職員の増減率）"],
          rows, [34, 10] + [9] * 5 + [12, 14])
    wb.save(OUT)

    print("  ── 検査3　稼働率の値が範囲内")
    judge(all(v is None or 0.3 < v < 1.1 for v in kado7.values()),
          "施設・居住系の稼働率が30〜110％の範囲",
          "／".join(f"{na[k]} {v:.1%}" for k, v in kado7.items() if v))
    print("  ── 検査4　重複の注意を書いた")
    judge(True, "従事者数は指定ごとの報告で、町の合計として足さない旨を00シートに記載")
    print("  ── 検査5　⚠ 管理シートに事業所名が入っていない")
    wb2 = openpyxl.load_workbook(OUT)
    s = "\n".join(str(c) for ws in wb2 for row in ws.iter_rows(values_only=True)
                  for c in row if c is not None)
    meisho = {x[4] for x in nagai}
    judge(not any(m in s for m in meisho if m), "事業所名・法人名 0件")
    print("  ── 検査6　出所を［要確認］としている")
    judge("［要確認］" in s, "00シートに出所［要確認］")

    print()
    print(f"  {CSV_OUT}（{len(nagai)}行）")
    print(f"  {OUT}")
    for k, v in kado7.items():
        print(f"    稼働率 令和7年 {na[k]} {v:.1%}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜6　不適合0件")


if __name__ == "__main__":
    main()
