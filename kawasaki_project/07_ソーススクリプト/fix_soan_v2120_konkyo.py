# -*- coding: utf-8 -*-
"""計画素案 Ver.2.12（根拠の棚卸しで見つかった数の是正）

08_作業順位 の順位25。Ver.2.11 に当てて Ver.2.12 を作る。

  入力　01_第10期_最新版成果品/川崎町_計画書素案_v2.11_見込量是正版.docx
  出力　01_第10期_最新版成果品/川崎町_計画書素案_v2.12_根拠整理版.docx

是正　3-3 の「過疎地域非該当855保険者」を「856保険者」に改める。

  交付金の全国統計を原典（令和8年度の全国集計・公表値）から作り直して
  突き合わせたところ、39件のうち38件は一致したが、1件だけ合わなかった。
  全国1,741保険者のうち過疎地域該当が885保険者であるから、
  非該当は856保険者である（885＋855＝1,740となり1保険者足りない）。
  平均473.5点は一致しており、**数え方ではなく件数の書き誤りである。**

⚠ 数値の是正はこの1件のみ。他の数値は変えていない。
"""
import copy
import os
import re
import shutil
import sys
import zipfile

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from fix_soan_v111 import set_el                             # noqa: E402

SRC = "01_第10期_最新版成果品/川崎町_計画書素案_v2.11_見込量是正版.docx"
DST = "01_第10期_最新版成果品/川崎町_計画書素案_v2.12_根拠整理版.docx"

NAOSU = [("非該当855保険者", "非該当856保険者")]


def main():
    if not os.path.exists(SRC):
        raise SystemExit("入力がない：" + SRC)
    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body

    n = 0
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = "".join(x.text or "" for x in el.iter(qn("w:t")))
        for furui, atarashii in NAOSU:
            if furui in s:
                set_el(el, s.replace(furui, atarashii))
                n += 1
    if n != len(NAOSU):
        raise SystemExit(f"差し替えが{n}件（{len(NAOSU)}件のはず）")

    # 版表記（本文）
    n_ver = 0
    for el in body.iter(qn("w:t")):
        if el.text and "Ver.2.11" in el.text:
            el.text = el.text.replace("Ver.2.11", "Ver.2.12")
            n_ver += 1
    if n_ver == 0:
        raise SystemExit("本文に Ver.2.11 の表記が見つからない")
    doc.save(DST)
    naoshi_header(DST)

    # ══════════════════════════ 自己点検
    d2 = docx.Document(DST)
    d0 = docx.Document(SRC)
    zen = ("\n".join(q.text for q in d2.paragraphs) + "\n"
           + "\n".join(c.text for t in d2.tables
                       for r in t.rows for c in r.cells))
    ng = []
    if "非該当855保険者" in zen:
        ng.append("「非該当855保険者」が残っている")
    if "非該当856保険者" not in zen:
        ng.append("「非該当856保険者」になっていない")
    # 件数の整合（885＋856＝1,741）
    if "1,741保険者" not in zen:
        ng.append("全国1,741保険者の記述がない")
    if (len(d2.paragraphs), len(d2.tables)) != (len(d0.paragraphs),
                                                len(d0.tables)):
        ng.append("段落数・表数が Ver.2.11 と違う（本版は数の是正のみ）")
    z = zipfile.ZipFile(DST)
    hdr = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>",
                             z.read("word/header1.xml").decode("utf-8")))
    if "Ver.2.12" not in hdr:
        ng.append(f"ヘッダーの版表記が更新されていない：{hdr}")

    print("計画素案 Ver.2.12（根拠整理版）")
    print("保存：", DST)
    print(f"  段落 {len(d2.paragraphs)}／表 {len(d2.tables)}"
          f"（Ver.2.11 と同じ）／数の是正 {n}件／版表記 {n_ver}か所")
    print("  ── 自己点検")
    if ng:
        for m in ng:
            print("   ×", m)
        sys.exit(1)
    print("   ○ 過疎地域非該当の保険者数を856に是正した（885＋856＝1,741）")
    print("   ○ 段落数・表数は Ver.2.11 と同じ（数の是正のみ）")
    print("  ⚠ 目次の頁番号は fill_toc_pages_v200.py で作り直してください。")


def naoshi_header(path):
    z = zipfile.ZipFile(path)
    names = z.namelist()
    data = {n: z.read(n) for n in names}
    z.close()
    n_fix = 0
    for n in names:
        if not re.match(r"word/(header|footer)\d*\.xml", n):
            continue
        s = data[n].decode("utf-8")
        s2 = re.sub(r"Ver\.2\.11", "Ver.2.12", s)
        if s2 != s:
            data[n] = s2.encode("utf-8")
            n_fix += 1
    if n_fix == 0:
        raise SystemExit("ヘッダーの版表記（Ver.2.11）が見つからない")
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    shutil.move(tmp, path)


if __name__ == "__main__":
    main()
