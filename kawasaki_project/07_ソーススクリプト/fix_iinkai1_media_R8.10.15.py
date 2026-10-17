# -*- coding: utf-8 -*-
"""第1回策定委員会資料の画像の拡張子の是正（令和8年10月15日）

08_作業順位 の順位26の作業中に見つかったもの。

`03_委員会・説明資料/川崎町_第1回策定委員会資料.docx` は、
**python-docx で開けない状態**であった。原因は次のとおり。

  ・word/media/3d9085….undefined という名の画像が1点入っている
  ・中身は PNG（先頭が \\x89PNG）であり、画像そのものは壊れていない
  ・[Content_Types].xml は png/jpeg/jpg/bmp/gif/svg/odttf を宣言しているが、
    undefined を宣言していない
  ・そのため、拡張子から中身の種類を決められず、読み込みが止まる
  ・この画像は本文から rId9 で1回参照されている（捨ててよいものではない）

Word は寛容に開くため、お渡しした資料そのものに見た目の支障はないと見ている。
ただし**当方の点検スクリプトがこの1件だけ素通りしていた**ため、
拡張子を .png に改めて、他の成果品と同じように点検できる状態にする。

  入力　03_委員会・説明資料/川崎町_第1回策定委員会資料.docx
  出力　同（上書き。元のものは 09_元資料/旧版図表_R8.9.30/ と同じ形で退避）

⚠ 中身は一切変えていない。画像の名を改め、参照先を合わせただけである。
⚠ この資料は令和8年9月2日の第1回策定委員会でお配りしたものであり、
  以後 v3〜v8統合版 に引き継がれている。本件は記録として残す版の是正である。
"""
import os
import re
import shutil
import sys
import zipfile

import docx

SRC = "03_委員会・説明資料/川崎町_第1回策定委員会資料.docx"
TAIHI = "09_元資料/旧版成果品/川崎町_第1回策定委員会資料_拡張子是正前.docx"
WARUI = ".undefined"

# 画像の先頭のバイトから種類を決める
MAGIC = [
    (b"\x89PNG\r\n\x1a\n", "png"),
    (b"\xff\xd8\xff", "jpeg"),
    (b"GIF87a", "gif"),
    (b"GIF89a", "gif"),
    (b"BM", "bmp"),
]


def shurui(data):
    for mg, ext in MAGIC:
        if data.startswith(mg):
            return ext
    return None


def main():
    if not os.path.exists(SRC):
        raise SystemExit("入力がない：" + SRC)

    z = zipfile.ZipFile(SRC)
    names = z.namelist()
    data = {n: z.read(n) for n in names}
    z.close()

    warui = [n for n in names
             if n.startswith("word/media/") and n.endswith(WARUI)]
    if not warui:
        print("是正の要るものがありません（既に是正済みです）")
        return
    if len(warui) != 1:
        raise SystemExit(f"拡張子の不明な画像が{len(warui)}件（1件のはず）")
    old = warui[0]
    ext = shurui(data[old])
    if ext is None:
        raise SystemExit("画像の種類を先頭のバイトから決められない")
    new = old[: -len(WARUI)] + "." + ext
    moto = data[old]

    # 元のものを退避（追跡されていない成果品を作らない）
    os.makedirs(os.path.dirname(TAIHI), exist_ok=True)
    if not os.path.exists(TAIHI):
        shutil.copy(SRC, TAIHI)

    # rels の Target を合わせる
    rels_name = "word/_rels/document.xml.rels"
    rels = data[rels_name].decode("utf-8")
    furui_t = "media/" + old.split("/")[-1]
    atarashii_t = "media/" + new.split("/")[-1]
    if furui_t not in rels:
        raise SystemExit("rels に画像の参照がない：" + furui_t)
    rels = rels.replace(furui_t, atarashii_t)
    data[rels_name] = rels.encode("utf-8")

    # Content_Types に png の宣言があることを確かめる
    ct = data["[Content_Types].xml"].decode("utf-8")
    if f'Extension="{ext}"' not in ct:
        raise SystemExit(f"[Content_Types].xml に {ext} の宣言がない")

    # 書き出す（画像の名だけを差し替える）
    tmp = SRC + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(new if n == old else n, data[n])
    shutil.move(tmp, SRC)

    # ══════════════════════════ 自己点検
    ng = []
    try:
        d = docx.Document(SRC)
    except Exception as e:
        print("   × 是正後も python-docx で開けない：", e)
        sys.exit(1)
    z2 = zipfile.ZipFile(SRC)
    n2 = z2.namelist()
    if any(x.endswith(WARUI) for x in n2):
        ng.append("拡張子の不明な画像が残っている")
    if new not in n2:
        ng.append(f"{new} がない")
    if z2.read(new) != moto:
        ng.append("画像の中身が変わっている")
    xml = z2.read("word/document.xml").decode("utf-8")
    rid = re.search(r'Id="([^"]+)"[^>]*Target="' + re.escape(atarashii_t) + '"',
                    z2.read("word/_rels/document.xml.rels").decode("utf-8"))
    if rid is None:
        ng.append("rels に新しい Target がない")
    elif xml.count(rid.group(1)) != 1:
        ng.append(f"本文からの参照が{xml.count(rid.group(1))}回（1回のはず）")
    n_img = len(re.findall(r"<w:drawing>", xml))
    # 退避した元のものと、段落数・表数が同じであること
    d0 = zipfile.ZipFile(TAIHI)
    x0 = d0.read("word/document.xml").decode("utf-8")
    for tag, nm in (("<w:p>", "段落"), ("<w:tbl>", "表"), ("<w:drawing>", "図")):
        if xml.count(tag) != x0.count(tag):
            ng.append(f"{nm}の数が変わっている"
                      f"（{x0.count(tag)}→{xml.count(tag)}）")

    print("第1回策定委員会資料の画像の拡張子の是正")
    print("保存：", SRC)
    print(f"  {old.split('/')[-1]}")
    print(f"  → {new.split('/')[-1]}（{len(moto):,}バイト・先頭のバイトから"
          f"{ext.upper()}と判定）")
    print(f"  退避：{TAIHI}")
    print(f"  段落 {len(d.paragraphs)}／表 {len(d.tables)}／図 {n_img}")
    print("  ── 自己点検")
    if ng:
        for m in ng:
            print("   ×", m)
        sys.exit(1)
    print("   ○ python-docx で開けるようになった")
    print("   ○ 画像の中身は変わっていない（バイトが一致）")
    print("   ○ 段落・表・図の数は変わっていない")
    print("  ⚠ 中身は一切変えていません。画像の名を改め、参照先を合わせただけです。")


if __name__ == "__main__":
    main()
