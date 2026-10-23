# -*- coding: utf-8 -*-
"""第2回策定委員会資料 v15 → v16（確定に向けて協議用素案 Ver.2.31 に揃える）
（令和8年10月23日）

  python3 07_ソーススクリプト/build_2kai_v16.py

ご指示「第2回策定委員会資料の確定に向けて進めて下さい」による。
資料は素案から起こし、数値を別に持たない（CLAUDE.md §3）。v15 と素案の食い違いを、
素案 Ver.2.31（fix_soan_v2310_kakutei.py）で素案の側を整えたうえで、資料を素案に揃える。

  (1) ４ 重点課題（５本）の表13：素案 3-4 の重点課題の表から起こす
      （v15 の「85歳以上の認定率58.67％」は素案になく第1回資料の値。3本目の名称も素案に揃える）
  (2) ５ 選択１：「計画素案は６つの基本目標により体系化」は現在の素案と合わない
      （素案は事務局案Ｃの7つの柱で組んでいる）。「当初の計画素案」と改め、現在の素案がＣ案で
      組まれていることを書く。表16のＢ・表17の見出し・本文の「計画素案に記載のない」も同じ
  (3) ７ 表37（第9期の計画値と実績値）：実績を図3-2（data_zuhyo.py）の千円単位の値に
      （v15 は10千円単位の値で、選択２の998,222千円と2つの値が並んでいた）
  (4) ７ 見出し「取り戻せる27点」：素案と同じく「評価の対象となり得る」の趣旨に
  (5) ８ 表39 第１章：宮城県の第10期みやぎ高齢者元気プラン（骨子（案））との対応（素案 1-2）
  (6) ８ 事業一覧表の件数：「17表・138事業」は体系の143事業と紛れる。内訳（134事業と再掲4件）を書く
  (7) ８ 表43：都道府県の意見の聴取は介護保険法第117条第13項（v15 は「第５項」と誤っていた）
  (8) ５ 選択３：「近隣市町村の指定事業所の有無は確認できておりません【町確認】」を、素案 5-5 の
      仙南圏域の事業所の文（公表システム・令和8年6月末）から起こす。町の指定の状況は【町確認】に残す
"""
import importlib.util
import os
import re
import shutil
import sys

import docx

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)


def load(name, fn):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(spec)
    sys.modules[name] = m
    spec.loader.exec_module(m)
    return m


V9 = load("build_2kai_v9", "build_2kai_v9.py")
from data_zuhyo import ZU                                   # noqa: E402

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v15.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v16.docx"
SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.31_協議用.docx"
ZEN = "０１２３４５６７８９"


def zen(s):
    return "".join(ZEN[int(c)] if c.isdigit() else c for c in s)


def kakikae(p, s):
    p.runs[0].text = V9.clean(s)
    for r in p.runs[1:]:
        r.text = ""


def soan_juten():
    d = docx.Document(SOAN)
    for t in d.tables:
        if [c.text for c in t.rows[0].cells][:2] == ["重点課題", "根拠となる事実"]:
            return [[c.text for c in r.cells] for r in t.rows]
    sys.exit("素案 3-4 の重点課題の表が見つからない")


def main():
    futeki = []
    KYU = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会資料 v15 → v16（確定に向けて素案 Ver.2.31 に揃える）")
    J = soan_juten()
    f32 = next(z for z in ZU if z["no"] == "図3-2")
    shutil.copy2(SRC, DST)
    d = docx.Document(DST)
    mae = "\n".join([p.text for p in d.paragraphs]
                    + [c.text for t in d.tables for r in t.rows for c in r.cells])

    def p_rep(atama, kyu, shin, n=1):
        h = [p for p in d.paragraphs if p.text.strip().startswith(atama) and kyu in p.text]
        judge(len(h) == n, f"『{atama[:18]}』の「{kyu[:18]}」（{len(h)}か所）")
        for p in h:
            KYU.append(p.text)
            kakikae(p, p.text.replace(kyu, shin))

    print("  ── (1) 表13 重点課題（素案 3-4 から）")
    t13 = [t for t in d.tables if [c.text.strip() for c in t.rows[0].cells][:2] == ["優先", "重点課題"]]
    judge(len(t13) == 1 and len(t13[0].rows) == len(J), "表13 を特定（5行）")
    if t13:
        for r, g in zip(t13[0].rows[1:], J[1:]):
            KYU.append("\n".join(c.text for c in r.cells))
            no, na = g[0].split("　", 1)
            V9.set_cell(r.cells[0], zen(no))
            V9.set_cell(r.cells[1], na)
            V9.set_cell(r.cells[2], g[1])
            V9.set_cell(r.cells[3], g[2].split("／", 1)[1])
    p_rep("第１回では分野別に５分野・課題11本をお示ししましたが",
          "５本に集約しました。",
          "①調査又は実績で裏付けのある事実があること、②第１０期の３年間で打ち手があること、"
          "の２つを基準に５本に集約しました（計画素案3-4）。")

    print("  ── (2) 選択１の説明（現在の素案はＣ案で組んでいる）")
    p_rep("第９期の基本理念は「誰もが健やかに暮らせるまちづくり」です。",
          "計画素案は課題を起点とした６つの基本目標により体系化しており、章立てが異なります。",
          "当初の計画素案は課題を起点とした６つの基本目標により体系化しており、章立てが異なって"
          "いました。現在の計画素案は、下の事務局案Ｃにより組んでいます（７つの柱）。")
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                if c.text.startswith("計画素案の６つの基本目標により体系化する"):
                    KYU.append(c.text)
                    V9.set_cell(c, c.text.replace("計画素案の６つの基本目標", "当初の計画素案の６つの基本目標"))
                if c.text.strip() == "計画素案の基本目標":
                    V9.set_cell(c, "当初の計画素案の基本目標")
    p_rep("町ご提示の体系（案）と計画素案の基本目標との対応は", "計画素案の基本目標",
          "当初の計画素案の基本目標")
    p_rep("町案にあって計画素案に記載のない事業が10件", "町案にあって計画素案に記載のない事業が10件、計画素案にあって",
          "町案にあって当初の計画素案に記載のない事業が10件、当初の計画素案にあって")
    p_rep("第９期計画の施策・事業（約110施策）を全数点検し", "町案にあって計画素案に記載のない事業が10件",
          "町案にあって当初の計画素案に記載のない事業が10件")
    p_rep("第９期計画の施策・事業（約110施策）を全数点検し", "計画素案にあって町案に対応する項がない施策",
          "当初の計画素案にあって町案に対応する項がない施策")

    print("  ── (3) 表37（図3-2 の千円単位の値）")
    t37 = [t for t in d.tables if [c.text.strip() for c in t.rows[0].cells][:2] == ["区　分", "令和７年度 計画値"]]
    judge(len(t37) == 1, "表37 を特定")
    if t37:
        for r, (_, kei, jis) in zip(t37[0].rows[1:], f32["rows"]):
            KYU.append("\n".join(c.text for c in r.cells))
            V9.set_cell(r.cells[1], f"{kei:,}千円")
            V9.set_cell(r.cells[2], f"{jis:,}千円")
            V9.set_cell(r.cells[3], f"{jis / kei * 100:.1f}％")

    print("  ── (4) 27点の見出しと言い回し")
    for p in d.paragraphs:
        if "取り戻せる27点" in p.text:
            KYU.append(p.text)
            kakikae(p, p.text.replace("取り戻せる27点", "取り戻しをめざす27点"))
    for t in d.tables:
        if [c.text.strip() for c in t.rows[0].cells][-1:] == ["取り戻せるか"]:
            V9.set_cell(t.rows[0].cells[-1], "取り戻せるか（見込み）")

    print("  ── (5)(6) 表39 第１章・事業一覧表の件数")
    for t in d.tables:
        for r in t.rows:
            if r.cells[0].text.strip() == "第１章" and r.cells[2].text.strip() == "長期総合計画との関係を明記":
                KYU.append(r.cells[2].text)
                V9.set_cell(r.cells[2], "長期総合計画との関係と、宮城県の第１０期みやぎ高齢者元気プラン"
                                        "（骨子（案））との対応を明記")
                judge(True, "表39 第１章")
    p_rep("計画書の事業一覧表（17表・138事業）", "（17表・138事業）",
          "（第５章・第６章の17表。134事業と再掲４件）")

    print("  ── (7) 表43 介護保険法の項")
    n = 0
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                if "介護保険法第117条第５項" in c.text:
                    KYU.append(c.text)
                    V9.set_cell(c, c.text.replace("介護保険法第117条第５項", "介護保険法第117条第13項"))
                    n += 1
    judge(n == 1, f"表43 の項番号（{n}か所）")
    print("  ── (8) 選択３ 近隣市町村の事業所（素案 5-5 から）")
    sp = [p.text for p in docx.Document(SOAN).paragraphs if "仙南圏域（2市7町）には" in p.text]
    m = re.search(r"(仙南圏域（2市7町）には、.+?。)", sp[0]) if sp else None
    judge(bool(m), "素案 5-5 の仙南圏域の文を読む")
    if m:
        bun = zen(m.group(1))[:-1]     # 資料の書き方（全角数字）に。末尾の「。」を除く
        p_rep("在宅介護実態調査では、定期巡回・随時対応型訪問介護看護の利用が",
              "近隣市町村の指定事業所の有無は確認できておりません【町確認】。",
              f"宮城県の高齢者福祉圏域である{bun}（介護サービス情報公表システム、令和８年６月末時点）。"
              "町外の地域密着型サービスの利用には本町による指定が要り、その状況をご確認いただいて"
              "います【町確認】。")

    # (8') 段落の書式（w:pPr）の子の順を Word の定め（スキーマの順）に並べ直す。
    #     v13 で第1回の書式に揃えた際に pStyle が先頭にない段落ができ、docx の検証で
    #     不適合になっていた（v13〜v15 に持ち越し）。
    JUN = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
           "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
           "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
           "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
           "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
           "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
           "sectPr", "pPrChange"]
    from docx.oxml.ns import qn
    narabe = 0
    for pPr in d.element.body.iter(qn("w:pPr")):
        ko = list(pPr)
        jun = sorted(ko, key=lambda c: JUN.index(c.tag.split("}")[1])
                     if c.tag.split("}")[1] in JUN else len(JUN))
        if jun != ko:
            for c in jun:
                pPr.append(c)
            narabe += 1
    judge(True, f"段落の書式の要素の順を直した（{narabe}段落）")
    d.save(DST)

    d2 = docx.Document(DST)
    ato = "\n".join([p.text for p in d2.paragraphs]
                    + [c.text for t in d2.tables for r in t.rows for c in r.cells])
    soan = docx.Document(SOAN)
    st = "\n".join([p.text for p in soan.paragraphs]
                   + [c.text for t in soan.tables for r in t.rows for c in r.cells])
    print("  ── 検査1　古い記述が残っていない")
    for go in ("58.67", "第117条第５項", "322,680", "998,220", "（17表・138事業）", "取り戻せる27点"):
        judge(go not in ato, f"「{go}」0件")
    judge(ato.count("計画素案は課題を起点とした") == ato.count("当初の計画素案は課題を起点とした"),
          "「計画素案は課題を起点とした」はすべて「当初の」付き")
    print("  ── 検査2　素案と同じ値")
    for g in J[1:]:
        judge(g[1] in ato and g[1] in st, f"重点課題「{g[0].split('　')[1]}」の根拠が素案と同じ")
    # 表37 の実績は素案 2-4 の表（円単位）を千円未満四捨五入した値（図3-2 と同じ）
    t24 = [t for t in soan.tables if [c.text for c in t.rows[0].cells][:4]
           == ["サービス区分", "要支援計", "要介護計", "合計（円）"]][0]
    en = {r.cells[0].text.split("\n")[0]: int(r.cells[3].text.replace(",", "")) for r in t24.rows[1:]}
    for k, (_, _, jis) in zip(("居宅サービス計", "地域密着型計", "施設サービス計", "総計"), f32["rows"]):
        judge(round(en[k] / 1000) == jis and f"{jis:,}千円" in ato,
              f"表37 の{k} {jis:,}千円 ＝ 素案 2-4 の {en[k]:,}円")
    print("  ── 検査3　消えた数は旧の文に現れる")
    kyu = "\n".join(KYU)
    kieta = sorted({m.group(0) for m in re.finditer(
        r"[0-9０-９][0-9０-９,.]*(?:人|名|％|千円|件|頁|日|点|円|位)", mae)
        if mae.count(m.group(0)) > ato.count(m.group(0)) and m.group(0) not in kyu})
    judge(not kieta, "消えた数はすべて旧の文に現れる", "／".join(kieta[:5]))
    print("  ── 検査4　事業所名・他団体名・⚠")
    for na in ("釜房", "アルパイン", "鶴寿会", "グループホームさわやか", "グループホームふれあい",
               "北塩原", "金ヶ崎", "大雪", "川崎市"):
        judge(na not in ato, f"「{na}」0件")
    judge("⚠" not in ato and "**" not in ato, "⚠・** 0件（第1回の書式）")
    print(f"  {DST}")
    print(f"  ⚠ 町のご確認が要る印【町確認】が {ato.count('【町確認】')}件残っています（確定の前に解く）")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜4　不適合0件")


if __name__ == "__main__":
    main()
