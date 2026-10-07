# -*- coding: utf-8 -*-
"""意見公募手続（パブリックコメント）一式（令和8年10月5日）

08_作業順位 の順位17。

計画素案 1-4 に「法令が定める手続」の表を置き、意見公募手続を掲げたが、
**実物の様式がない。** 逆算工程表では令和8年12月に開始する。
期間（30日／20日）が決まらなくても、実施要領・様式は雛形として作れる。

  出力　05_試算・管理シート/川崎町_意見公募手続_一式_R8.10.5.xlsx
          00_この表について／01_実施要領（案）／02_工程（30日・20日）／
          03_意見と町の考え方の対照表（公表用の様式）／
          04_受付の記録（内部用の様式）／05_関係する確認事項
        03_委員会・説明資料/川崎町_意見公募手続_意見提出用紙_R8.10.5.docx

⚠ **意見公募手続の実施支援は、委託仕様書の業務内容に規定がない。**
  本業務の範囲とするか、仕様書11「その他」による協議事項とするかは
  確認事項No.19のご判断による。**ご判断を待つ間に手が止まらないよう、
  様式を先に用意しておくもの**であり、実施をお約束するものではない。

⚠ 期間は**町の実施要領（例規）の定めによる。** 本一式は30日間を既定値とし、
  20日間とする場合の工程も併せて示す。**町に意見公募手続の要綱・要領が
  あれば、そちらが優先する**（確認事項No.136）。

⚠ 第9期は令和6年1月31日から2月9日に実施している（10日間）。
  前例として示すものであり、第10期の期間を縛るものではない。
"""
import copy
import datetime as dt
import importlib.util
import os
import re
import shutil
import sys
import zipfile

import docx
import openpyxl
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from fix_soan_v111 import clone_table, set_el            # noqa: E402


def load(name, fname):
    spec = importlib.util.spec_from_file_location(name,
                                                  os.path.join(HERE, fname))
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


GY = load("gyotei2", "build_iinkai_gyotei_R8.10.1.py")

SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.10_制度改正補完版.docx"
OUT_X = "05_試算・管理シート/川崎町_意見公募手続_一式_R8.10.5.xlsx"
OUT_D = "03_委員会・説明資料/川崎町_意見公募手続_意見提出用紙_R8.10.5.docx"

W = "月火水木金土日"
KIKAN = [30, 20]            # 期間の案（日）
HAJIME = dt.date(2026, 12, 16)   # 開始日の仮置き（候補2による）


def ymd(d):
    return "%d年%d月%d日（%s）" % (d.year - 2018, d.month, d.day,
                                   W[d.weekday()])


# ══════════════════════════ 01　実施要領（案）
YORYO = [
    ("1　目的",
     "高齢者保健福祉計画・第10期介護保険事業計画（令和9年度から令和11年度まで）"
     "の策定にあたり、計画の案を公表して広く意見を求め、"
     "寄せられた意見を考慮して計画を定めるため。",
     "介護保険法第117条が定める、被保険者の意見を反映させるために"
     "必要な措置に当たる。"),
    ("2　意見を求める案件",
     "川崎町高齢者保健福祉計画・第10期介護保険事業計画（案）",
     "計画書（案）本文と概要版を公表する。資料編は閲覧に供する。"),
    ("3　公表の方法",
     "①町のホームページへの掲載　②保健福祉課の窓口での閲覧"
     "　③保健センター・役場の窓口での閲覧　④図書館等での閲覧",
     "⚠ 実際の閲覧場所は町でお決めください。"
     "第9期の実施例があればそれに倣います（確認事項No.136）。"),
    ("4　意見を提出できる方",
     "①町内に住所を有する方　②町内に事務所又は事業所を有する方"
     "　③町内の事務所又は事業所に勤務する方　④町内の学校に在学する方"
     "　⑤当該案件に利害関係を有する方",
     "⚠ 範囲は町の要綱・要領の定めによります。"),
    ("5　意見の提出の方法",
     "意見提出用紙（様式）により、①窓口への持参　②郵送　③ファクシミリ"
     "　④電子メール　⑤町ホームページの入力フォーム のいずれかによる。",
     "氏名（名称）・住所・連絡先の記載を求める。"
     "⚠ 匿名の意見の扱いは町でお決めください。"),
    ("6　意見の提出の期間",
     "⚠ 30日間を既定値とする（02シート）。20日間とする場合の工程も"
     "併せて示している。",
     "期間は町の実施要領（例規）の定めによる。"
     "第9期は令和6年1月31日から2月9日に実施している（10日間）。"),
    ("7　意見の取扱い",
     "寄せられた意見は、意見の概要と町の考え方を取りまとめ、"
     "策定委員会に報告したうえで公表する。"
     "個人を特定できる情報は公表しない。",
     "公表の様式は03シート。第3回策定委員会（令和9年1月）で報告する。"),
    ("8　意見に対する個別の回答",
     "個別の回答は行わない。意見の概要と町の考え方の公表をもって"
     "回答に代える。",
     "⚠ 個別に回答する運用とする場合は、工程に回答の期間を要します。"),
    ("9　公表の時期と方法",
     "意見の提出の期間の終了後、計画の決定までの間に、"
     "町のホームページへの掲載その他3の方法により公表する。",
     "計画書の資料編に、実施の結果（期間・件数・主な意見）を掲げる案"
     "（確認事項No.134ほか・別添整理表）。"),
]

# ══════════════════════════ 03　対照表（公表用）の記入例
TAISHO_REI = [
    ["（記入例）", "9-3 保険料", "保険料の上げ幅が大きい。"
     "基金をもっと取り崩せないか。",
     "介護給付費準備基金は、第10期の3年間の給付費の増に備えて"
     "保有するものです。取崩額は策定委員会のご審議を経て決定しました。"
     "いただいたご意見を踏まえ、取崩しの考え方を 9-3 に加えます。",
     "計画に反映した", "9-3 に注記を加える"],
    ["", "", "", "", "", ""],
    ["", "", "", "", "", ""],
    ["", "", "", "", "", ""],
]

KUBUN = [
    ("計画に反映した", "意見を踏まえ、計画の案を修正した。"),
    ("計画に既に盛り込まれている", "意見の趣旨は、計画の案に既に含まれている。"),
    ("今後の参考とする", "計画の修正には至らないが、"
                        "事業の実施にあたって参考とする。"),
    ("計画に反映できない", "法令・制度の定め、又は財政上の理由により"
                          "計画に盛り込むことができない。理由を示す。"),
    ("その他", "計画の内容に関わらないご意見・ご質問。"),
]

# ══════════════════════════ 04　受付の記録（内部用）
UKETSUKE_REI = [
    ["（記入例）", "R8.12.20", "郵送", "町内在住", "9-3 保険料", "1",
     "対照表 No.1 へ", ""],
    ["", "", "", "", "", "", "", ""],
    ["", "", "", "", "", "", "", ""],
]

KAKUNIN = ["19", "136", "134", "162", "83", "49", "85"]


# ══════════════════════════ 意見提出用紙（docx）
def youshi():
    """(種類, 中身) の並び。種類は 表題／見出し／本文／囲み／T2〜T5。"""
    C = []
    a = C.append
    a(("表題", "川崎町高齢者保健福祉計画・第10期介護保険事業計画（案）"
               "　意見提出用紙"))
    a(("本文", "川崎町高齢者保健福祉計画・第10期介護保険事業計画（案）に"
               "ついて、ご意見をお寄せください。"))
    a(("見出し", "▌ご記入にあたって"))
    a(("本文", "・この用紙でなくても、同じ内容をご記入いただければ"
               "どのような様式でも結構です。"))
    a(("本文", "・ご意見は、意見の概要と町の考え方を取りまとめて公表します。"
               "個人を特定できる情報は公表しません。"))
    a(("本文", "・ご意見に対する個別の回答は行いません。"))
    a(("本文", "・お名前・ご住所の記載のないものは、お取り扱いできない"
               "場合があります。"))
    a(("見出し", "▌ご記入ください"))
    a(("T2", [["ふりがな／お名前（団体の場合は名称と代表者名）", "［　　　　］"],
              ["ご住所", "［　　　　］"],
              ["ご連絡先（電話・ファクシミリ・電子メール）", "［　　　　］"],
              ["該当する区分（いずれかに○）",
               "町内に住所がある／町内に事務所・事業所がある／"
               "町内の事務所・事業所に勤務している／町内の学校に在学している／"
               "この計画に利害関係がある"],
              ["ご意見の該当箇所（章・節・頁）", "［　　　　］"]],
       [36, 64]))
    a(("見出し", "▌ご意見"))
    # ⚠ 自由にご記入いただく欄は、真ん中に線が入らないよう**1列の表**にする
    a(("T1", [["ご意見の内容（この欄にご記入ください。"
               "足りない場合は別紙を添えてください。）"],
              [""], [""], [""], [""], [""]], [100]))
    a(("見出し", "▌提出の方法・期限"))
    a(("T2", [["提出の期限", "［　年　月　日（　）　まで　必着］"],
              ["提出先", "川崎町 保健福祉課"],
              ["持参・郵送", "［住所　　　　］"],
              ["ファクシミリ", "［　　　　］"],
              ["電子メール", "［　　　　］"],
              ["町ホームページの入力フォーム", "［　　　　］"]],
       [30, 70]))
    a(("囲み",
       "⚠ ［　］の箇所は、実施の決定後に町にてご記入ください。"
       "提出の期間・提出先・閲覧場所は、町の実施要領の定めによります。"))
    return C


def build_docx():
    shutil.copy(SOAN, OUT_D)
    doc = docx.Document(OUT_D)
    body = doc.element.body
    tpl = {}
    for el in body.iterchildren():
        if el.tag == qn("w:p"):
            if el.findall(".//" + qn("w:drawing")):
                continue
            s = "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()
            if "表題" not in tpl and s.startswith("第 1 章"):
                tpl["表題"] = copy.deepcopy(el)
            if "見出し" not in tpl and s.startswith("▌"):
                tpl["見出し"] = copy.deepcopy(el)
            if "本文" not in tpl and s.startswith("令和7年度から令和8年度への"):
                tpl["本文"] = copy.deepcopy(el)
            if "囲み" not in tpl and s.startswith("⚠ 宮城県自身の評価"):
                tpl["囲み"] = copy.deepcopy(el)
        elif el.tag == qn("w:tbl"):
            n = len(el.findall(qn("w:tr"))[0].findall(qn("w:tc")))
            tpl.setdefault(f"T{n}", copy.deepcopy(el))
    for k in ("表題", "見出し", "本文", "囲み", "T1", "T2"):
        if k not in tpl:
            raise SystemExit("雛形が見つからない：" + k)
    sectPr = body.find(qn("w:sectPr"))
    for el in list(body.iterchildren()):
        if el is not sectPr:
            body.remove(el)

    def add(el):
        if sectPr is not None:
            sectPr.addprevious(el)
        else:
            body.append(el)

    n_tbl = 0
    for item in youshi():
        kind, v = item[0], item[1]
        if kind.startswith("T"):
            add(clone_table(tpl[kind], v, item[2]))
            e = copy.deepcopy(tpl["本文"])
            set_el(e, "")
            add(e)
            n_tbl += 1
        else:
            el = copy.deepcopy(tpl[kind])
            set_el(el, v)
            add(el)
    doc.save(OUT_D)
    strip_media(OUT_D)
    naoshi_header(OUT_D, "川崎町 意見公募手続 意見提出用紙")
    return n_tbl


def strip_media(path):
    z = zipfile.ZipFile(path)
    names = z.namelist()
    data = {n: z.read(n) for n in names}
    z.close()
    used = set(re.findall(r'r:embed="([^"]+)"',
                          data["word/document.xml"].decode("utf-8")))
    rels = "word/_rels/document.xml.rels"
    from lxml import etree
    root = etree.fromstring(data[rels])
    drop = set()
    for rel in list(root):
        tgt = rel.get("Target") or ""
        if tgt.startswith("media/") and rel.get("Id") not in used:
            drop.add("word/" + tgt)
            root.remove(rel)
    data[rels] = etree.tostring(root, xml_declaration=True,
                                encoding="UTF-8", standalone=True)
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            if n not in drop:
                zout.writestr(n, data[n])
    shutil.move(tmp, path)


def naoshi_header(path, midashi):
    z = zipfile.ZipFile(path)
    names = z.namelist()
    data = {n: z.read(n) for n in names}
    z.close()
    for n in names:
        if not re.match(r"word/(header|footer)\d*\.xml", n):
            continue
        s = data[n].decode("utf-8")
        s = re.sub(r"川崎町[^<]*計画書素案[^<]*", midashi, s)
        s = re.sub(r"Ver\.[\d.]+[^<]*", "（案）", s)
        data[n] = s.encode("utf-8")
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    shutil.move(tmp, path)


def yomu_kakunin(nos):
    wb = openpyxl.load_workbook("川崎町_業務工程管理表.xlsx")
    ws = wb["03_確認事項一覧"]
    out, want = [], list(nos)
    for r in range(5, ws.max_row + 1):
        v = str(ws.cell(r, 1).value or "").strip()
        if v in want:
            out.append([v, str(ws.cell(r, 4).value or ""),
                        str(ws.cell(r, 8).value or ""),
                        str(ws.cell(r, 10).value or "")])
    out.sort(key=lambda x: want.index(x[0]))
    miss = set(want) - {x[0] for x in out}
    if miss:
        raise SystemExit("確認事項が見つからない：" + "／".join(sorted(miss)))
    return out


def main():
    if not os.path.exists(SOAN):
        raise SystemExit("入力がない：" + SOAN)
    wb = openpyxl.Workbook()

    ws = GY.sheet(
        wb, "00_この表について", "意見公募手続（パブリックコメント）一式",
        "計画素案 1-4 に法令が定める手続として意見公募手続を掲げましたが、"
        "実物の様式がありませんでした。逆算工程表では令和8年12月に"
        "開始します。期間が決まらなくても様式は用意できますので先に作りました。\n"
        "⚠ 意見公募手続の実施支援は、委託仕様書の業務内容に規定がありません。"
        "本業務の範囲とするか、仕様書11「その他」による協議事項とするかは"
        "確認事項No.19のご判断によります。ご判断を待つ間に手が止まらないよう"
        "様式を先に用意するものであり、実施をお約束するものではありません。\n"
        "⚠ 町に意見公募手続の要綱・要領がある場合は、そちらが優先します"
        "（確認事項No.136）。本一式は当方の案です。",
        ["シート", "内容"], [30, 96], first=True)
    GY.put(ws, [
        ["01_実施要領（案）", f"目的から公表の時期まで{len(YORYO)}項目。"],
        ["02_工程", "期間を30日間とする場合と20日間とする場合の日取り。"],
        ["03_意見と町の考え方の対照表",
         "公表する様式。区分5つの定義を併せて置いています。"],
        ["04_受付の記録", "内部で意見を受け付けた順に記録する様式。"],
        ["05_関係する確認事項", f"{len(KAKUNIN)}件。"],
        ["意見提出用紙（別ファイル・docx）",
         "住民の方にお配りする様式。"
         "03_委員会・説明資料/川崎町_意見公募手続_意見提出用紙_R8.10.5.docx"],
    ])

    ws = GY.sheet(
        wb, "01_実施要領（案）", "意見公募手続 実施要領（案）",
        "⚠ 本案は当方が一般的な運用に倣って書いたものです。"
        "町に要綱・要領がある場合はそちらによります（確認事項No.136）。\n"
        "⚠ 「備考」欄に、町にお決めいただく必要がある箇所を書いています。",
        ["項目", "案", "備考"], [22, 60, 56])
    GY.put(ws, [list(x) for x in YORYO],
           [GY.WARN if "⚠" in x[1] + x[2] else None for x in YORYO])

    ws = GY.sheet(
        wb, "02_工程", "意見公募手続の工程（期間30日・20日）",
        f"第2回策定委員会を令和8年11月11日（候補2）に開催する場合の"
        f"工程です。開始日は{ymd(HAJIME)}を仮置きしています。\n"
        "⚠ 30日間は国の指針が原則とする期間です。市町村の計画の意見公募は"
        "町の実施要領の定めによります。20日間とする案は、"
        "第2回の開催が遅れた場合に納期へ収めるための短縮の案です"
        "（第2回策定委員会資料の案Ｂ）。",
        ["期間の案", "公表・開始", "終了", "取りまとめ",
         "第3回策定委員会での報告", "備考"], [14, 22, 22, 22, 24, 40])
    rows, fills = [], []
    for n in KIKAN:
        owari = HAJIME + dt.timedelta(days=n - 1)
        matome = owari + dt.timedelta(days=10)
        houkoku = matome + dt.timedelta(days=4)
        rows.append([f"{n}日間", ymd(HAJIME), ymd(owari), ymd(matome),
                     ymd(houkoku),
                     "国の指針が原則とする期間。" if n == 30
                     else "短縮の案（案Ｂ）。町の実施要領の定めによる。"])
        fills.append(GY.OKF if n == 30 else GY.WARN)
    GY.put(ws, rows, fills)

    ws = GY.sheet(
        wb, "03_意見と町の考え方の対照表", "意見と町の考え方の対照表（公表用）",
        "いただいたご意見の概要と、町の考え方を1件ずつ対応させて公表する"
        "様式です。個人を特定できる情報は載せません。\n"
        "⚠ 1行目は記入例です。お使いの際は削除してください。\n"
        "⚠ 同じ趣旨のご意見は、まとめて1件として扱うことができます。"
        "その場合は件数を括弧で添えてください。",
        ["No.", "該当箇所", "ご意見の概要", "町の考え方",
         "区分", "計画の修正"], [8, 18, 40, 56, 20, 24])
    r = GY.put(ws, TAISHO_REI, [GY.LEAD] + [None] * (len(TAISHO_REI) - 1))
    r += 2
    ws.cell(r, 1).value = "【区分】"
    ws.cell(r, 1).font = openpyxl.styles.Font(name="游ゴシック", size=10,
                                              bold=True)
    GY.put(ws, [[k, s, "", "", "", ""] for k, s in KUBUN], r0=r + 1)

    ws = GY.sheet(
        wb, "04_受付の記録", "意見の受付の記録（内部用）",
        "受け付けた順に記録し、対照表のどの番号に対応するかを残します。"
        "この表は公表しません。\n"
        "⚠ 氏名・住所は、この表には書かずに提出用紙の原本で管理する運用を"
        "お勧めします（個人情報の取扱い）。",
        ["No.", "受付日", "提出の方法", "提出者の区分", "該当箇所",
         "意見の件数", "対照表の番号", "備考"],
        [8, 14, 14, 18, 20, 12, 16, 30])
    GY.put(ws, UKETSUKE_REI, [GY.LEAD] + [None] * (len(UKETSUKE_REI) - 1))

    kak = yomu_kakunin(KAKUNIN)
    ws = GY.sheet(
        wb, "05_関係する確認事項", "意見公募手続に関わる確認事項",
        "業務工程管理表 03_確認事項一覧 から読んでいます。"
        "内容はそちらが正本です。",
        ["No.", "確認事項", "状態", "決着しない場合の当方の扱い"],
        [8, 40, 12, 62])
    GY.put(ws, kak, [GY.OKF if x[2] == "完了" else None for x in kak])

    wb.save(OUT_X)
    n_tbl = build_docx()

    # ══════════════════════════ 自己点検
    ng = []
    wb2 = openpyxl.load_workbook(OUT_X)
    if len(wb2.worksheets) != 6:
        ng.append(f"シートが{len(wb2.worksheets)}枚（6枚のはず）")
    zen = "\n".join(str(c.value) for w_ in wb2.worksheets
                    for row in w_.iter_rows() for c in row
                    if c.value is not None)
    d2 = docx.Document(OUT_D)
    zen += "\n" + "\n".join(q.text for q in d2.paragraphs)
    zen += "\n" + "\n".join(c.text for t in d2.tables
                            for r_ in t.rows for c in r_.cells)
    for w in ("**", "__"):
        if w in zen:
            ng.append(f"強調記号が残っている：{w}")
    for w in ("大雪", "東川", "東神楽", "上川", "金ヶ崎", "金ケ崎", "川崎市"):
        if w in zen:
            ng.append(f"他団体の名称が入っている：{w}")
    # 仕様書外であることを明記していること
    if "業務内容に規定がありません" not in zen:
        ng.append("仕様書に規定がない旨の断りがない")
    # 期間を断定していないこと
    for w in ("30日間とします", "30日間と定める"):
        if w in zen:
            ng.append(f"期間を断定している：{w}")
    # 未確定の箇所は［　］で書くこと（他の括弧を使っていないこと）
    if re.search(r"【[^】]*記入[^】]*】", zen):
        ng.append("未確定箇所に【　】を使っている（［　］で書く）")
    if len(d2.tables) != n_tbl:
        ng.append(f"用紙の表が{len(d2.tables)}件（{n_tbl}件のはず）")
    z = zipfile.ZipFile(OUT_D)
    if "<w:drawing>" in z.read("word/document.xml").decode("utf-8"):
        ng.append("意見提出用紙に図が残っている")
    if sum(1 for n in z.namelist() if n.startswith("word/media")):
        ng.append("使っていない画像が残っている")

    print("意見公募手続（パブリックコメント）一式")
    print("保存：", OUT_X)
    print(f"  シート{len(wb2.worksheets)}枚／実施要領{len(YORYO)}項目"
          f"／期間の案{len(KIKAN)}件／区分{len(KUBUN)}件"
          f"／確認事項{len(kak)}件")
    print("保存：", OUT_D)
    print(f"  段落{len(d2.paragraphs)}／表{len(d2.tables)}")
    print("  ── 自己点検")
    if ng:
        for m in ng:
            print("   ×", m)
        sys.exit(1)
    print("   ○ 仕様書に規定がないことを明記している（確認事項No.19）")
    print("   ○ 期間を断定していない（町の実施要領の定めによる）")
    print("   ○ 未確定箇所は［　］で書いている")
    print("  ⚠ 実施をお約束するものではありません。"
          "ご判断を待つ間に手が止まらないよう様式を先に用意したものです。")


if __name__ == "__main__":
    main()
