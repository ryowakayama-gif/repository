# -*- coding: utf-8 -*-
"""業務工程管理表の更新（No.187 のご判断・協議用素案と公表用の運用）
（令和8年10月23日・同日2度目）

  python3 07_ソーススクリプト/update_wbs_R8.10.23b.py

⚠ 列8（状態）は動かさない。回答・対応（列10）に記す。
"""
import copy
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
SOAN_OLD = "川崎町_計画書素案_v2.26_版表記整理版.docx"
SOAN = ("川崎町_計画書素案_v2.27_協議用.docx（協議用素案）／"
        "川崎町_計画書案_公表用_v2.27.docx（公表用の計画書（案））")
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月23日 更新②（No.187 ご判断・2つの版）】"


def iro(ws, r, n):
    for c in range(1, n + 1):
        ws.cell(r, c).font = FONT
        ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")


def main():
    wb = openpyxl.load_workbook(WBS)
    ws = wb["03_確認事項一覧"]
    n = 0
    for x in range(5, ws.max_row + 1):
        if (str(ws.cell(x, 1).value or "").strip() == "187"
                and HIZUKE not in str(ws.cell(x, 10).value or "")):
            ws.cell(x, 10).value = (
                f"{HIZUKE}ご指示：①書体は第9期と同じ HG丸ゴシックM-PRO に揃える"
                f"（大きさ・行送りは変えない）②見出しと図表の題はそのまま"
                f"③協議用素案と公表用の計画書（案）に分けて運用する。"
                f"素案 Ver.2.27（協議用）と公表用 v2.27 に反映した。　"
                + str(ws.cell(x, 10).value or ""))
            iro(ws, x, 10)
            n += 1

    ws = wb["04_成果品管理"]
    n_seika = 0
    for x in range(5, ws.max_row + 1):
        na = str(ws.cell(x, 3).value or "")
        if str(ws.cell(x, 1).value or "").strip() == "4" and SOAN_OLD in na:
            ws.cell(x, 3).value = na.replace(SOAN_OLD, SOAN)
            ws.cell(x, 8).value = (
                f"{HIZUKE}協議用素案（Ver.2.27・116頁）と公表用の計画書（案）"
                f"（113頁・体裁確認用）に分けた。内容の正本は協議用素案で、公表用は"
                f"build_kouhyou_soan.py が組む。書体を HG丸ゴシックM-PRO に。"
                f"⚠ 事業一覧17表の「区分」等、短い語が縦に折れていた27表の列を広げた。　"
                + str(ws.cell(x, 8).value or ""))
            iro(ws, x, 9)
            n_seika += 1

    ws2 = wb["08_作業順位"]
    if not any(str(ws2.cell(x, 1).value or "").strip() == "42"
               for x in range(5, ws2.max_row + 1)):
        moto = ws2.max_row
        r = moto + 1
        for c in range(1, 11):
            ws2.cell(r, c)._style = copy.copy(ws2.cell(moto, c)._style)
        ws2.row_dimensions[r].height = ws2.row_dimensions[moto].height
        for c, v in enumerate([
                "42", "書体を第9期に揃え、協議用素案と公表用の計画書（案）に分ける",
                "R8.10.23", "ご指示による（No.187）。",
                "No.187", "協議用素案 Ver.2.27／公表用の計画書（案）v2.27／"
                "公表用で落とした文（xlsx）／運用メモ（md）", "1日", "完了",
                "HG丸ｺﾞｼｯｸM-PRO に（Windows の登録名は半角カナ）。公表用は協議用から"
                "組み替え、⚠・確認事項No・【町確認】を落とす／揃える。⚠ ［要確認］"
                "［要協議］を含む文は落とさない（公表用 ［要確認］70・［要協議］13）。"
                "落とした文48件を記録。短い語が縦に折れていた27表の列を広げた。"], 1):
            ws2.cell(r, c).value = v
    wb.save(WBS)

    wb2 = openpyxl.load_workbook(WBS, data_only=True)
    s = "\n".join(str(v) for sh in wb2 for row in sh.iter_rows(values_only=True)
                  for v in row if v is not None)
    jun = [str(wb2["08_作業順位"].cell(x, 1).value or "").strip()
           for x in range(5, wb2["08_作業順位"].max_row + 1)]
    ok = (n in (0, 1) and n_seika in (0, 1) and jun.count("42") == 1
          and "公表用_v2.27" in s
          and not any(w in s for w in ("蔵王グリーンホーム", "鶴寿会", "北塩原",
                                       "金ヶ崎", "大雪", "川崎市")))
    print("業務工程管理表の更新（No.187・2つの版）")
    print(("   ○ " if ok else "   × ") + f"No.187 {n}件・成果品4 {n_seika}件・作業順位42")
    if not ok:
        sys.exit(1)


if __name__ == "__main__":
    main()
