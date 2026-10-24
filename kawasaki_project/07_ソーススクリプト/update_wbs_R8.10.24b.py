# -*- coding: utf-8 -*-
"""業務工程管理表の更新（Red Team 再レビュー（令和8年10月24日）への対応）

  python3 07_ソーススクリプト/update_wbs_R8.10.24b.py

  00_管理方法　　　基準日の記載を再レビュー対応後に（件数は台帳から数える）
  03_確認事項一覧　No.190 を起票（生活支援コーディネーター26人の区分・配置）。No.85・148 に追記
  04_成果品管理　　稿を 素案 Ver.2.33／公表用・概要版 v2.33／資料 v18／問答集 v15 に
  08_作業順位　　　50（完了）を加える
再実行しても二重に書かない。
"""
import collections
import copy
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月24日 更新②（Red Team 再レビューへの対応）】"
TAIOU = "川崎町_RedTeam再レビュー対応表_R8.10.24.md"
ZIP = "11_提出一式/川崎町_RedTeam再レビュー対応_成果物とエビデンス_R8.10.24.zip"
OKIKAE = [("川崎町_計画書素案_v2.32_協議用.docx", "川崎町_計画書素案_v2.33_協議用.docx"),
          ("川崎町_計画書案_公表用_v2.32.docx", "川崎町_計画書案_公表用_v2.33.docx"),
          ("川崎町_計画書案_概要版_v2.32.docx", "川崎町_計画書案_概要版_v2.33.docx"),
          ("川崎町_第2回策定委員会資料_R8.11_v17.docx", "川崎町_第2回策定委員会資料_R8.11_v18.docx"),
          ("同_想定問答集_R8.11_v14.docx", "同_想定問答集_R8.11_v15.docx")]
NO190 = [
    "190", "(3)", "R8.10",
    "生活支援コーディネーター26人の区分と配置（第1層・第2層）",
    "計画素案の生活支援コーディネーター26人（令和7年度見込）は、町の実績データ（データ入力フォーマット"
    "02シート）の「生活支援サポーター（SC）_人数」と、国の令和8年度評価（該当状況調査）の"
    "「生活支援コーディネーター数」がいずれも26人であることによります。人数は一致しますが、"
    "同じ方々を指すか（生活支援体制整備事業の生活支援コーディネーターとして配置されているか）を"
    "お示しください。あわせて、第1層・第2層の配置の内訳をお願いします。",
    "計画素案 5-4・10-4（KPI）\n概要版（KPI の行）",
    "発注者", "確認待ち", "R8.11.10",
    f"{HIZUKE}Red Team 再レビュー N03。素案 Ver.2.33 で KPI と 5-4 の表に【町確認】を付した。"
    "既定値＝回答がない場合、KPI の現状値は「確認中」とし数値を載せない（公表用は［要確認］、"
    "概要版は載せない）。",
]
TSUIKI = {
    "85": (f"{HIZUKE}再レビューを受け、資料を v18・想定問答集を v15 に上げた。報酬改定1％・3％の算定Ｂ・Ｃを"
           "素案 9-3 から起こし直した（1％ Ｂ6,208・Ｃ5,529円、3％ Ｂ6,339・Ｃ5,659円。旧の6,235円等は誤り）。"
           "選択２は5,931円・6,355円（±212円）。ご審議事項⑥の言い切り、表29 の基金残高（見込額）、"
           f"改頁（21頁の「す。」）を直した。当方の環境の PDF は22頁。{TAIOU}。　"),
    "148": (f"{HIZUKE}報酬改定の差は、基金の取崩しが一定額の差引きであるため算定Ａ・Ｂ・Ｃでほぼ同じ"
            "（1％ ＋66・＋65・＋65円、3％ ＋196・＋196・＋195円）。前回の算定Ｂ＋92円・算定Ｃ＋119円は"
            "誤った表から計算していた。入所・居住系±10人は211.7円を四捨五入して212円。　"),
}


def iro(c):
    c.font = FONT
    c.alignment = Alignment(wrap_text=True, vertical="top")


def mae_ni(c, s):
    if HIZUKE in str(c.value or ""):
        return 0
    c.value = s + str(c.value or "")
    iro(c)
    return 1


def main():
    wb = openpyxl.load_workbook(WBS)
    kekka = []
    ws3 = wb["03_確認事項一覧"]
    nos = {str(ws3.cell(y, 1).value or "").strip() for y in range(5, ws3.max_row + 1)}
    if "190" not in nos:
        moto = ws3.max_row
        r = moto + 1
        for c in range(1, 11):
            ws3.cell(r, c)._style = copy.copy(ws3.cell(moto, c)._style)
            ws3.cell(r, c).value = NO190[c - 1]
        ws3.row_dimensions[r].height = ws3.row_dimensions[moto].height
    k = 0
    for y in range(5, ws3.max_row + 1):
        no = str(ws3.cell(y, 1).value or "").strip()
        if no in TSUIKI:
            k += mae_ni(ws3.cell(y, 10), TSUIKI[no])
    kekka.append((k in (0, len(TSUIKI)), f"確認事項への追記 {k}件"))
    nos = sorted(int(str(ws3.cell(y, 1).value).strip()) for y in range(5, ws3.max_row + 1)
                 if str(ws3.cell(y, 1).value or "").strip().isdigit())
    kekka.append((nos == list(range(1, 191)), f"確認事項 1〜190・欠番なし（{len(nos)}件）"))
    k3 = collections.Counter(str(ws3.cell(y, 8).value).strip() for y in range(5, ws3.max_row + 1)
                             if str(ws3.cell(y, 1).value or "").strip().isdigit())

    ws0 = wb["00_管理方法"]
    for r in ws0.iter_rows():
        for c in r:
            v = str(c.value or "")
            if v.startswith("令和8年10月24日　本表の記載は同日時点のもの。Red Team レビュー（"):
                c.value = ("令和8年10月24日　本表の記載は同日時点のもの。Red Team 再レビュー（同日）を受け、"
                           "協議用素案 Ver.2.33・公表用の計画書（案）v2.33・概要版 v2.33・第2回策定委員会資料 v18・"
                           "想定問答集 v15（全69問）に是正した。新規の指摘11件のうち是正済9・一部是正1・"
                           "町のご確認待ち1（No.190）。点検ログ14件はすべて終了コード0。"
                           f"確認事項は{sum(k3.values())}件のうち{k3['完了']}件が完了。一式は {ZIP}。\n"
                           "（前の記載）" + v)
                iro(c)
                kekka.append((True, "基準日の記載を再レビュー対応後に"))

    ws4 = wb["04_成果品管理"]
    n = 0
    for y in range(5, ws4.max_row + 1):
        na = str(ws4.cell(y, 3).value or "")
        nb = na
        for a, b in OKIKAE:
            nb = nb.replace(a, b)
        if nb != na:
            ws4.cell(y, 3).value = nb
            mae_ni(ws4.cell(y, 8), f"{HIZUKE}Red Team 再レビューの指摘を是正した版に上げた（{TAIOU}）。　")
            n += 1
    kekka.append((n in (0, 3), f"成果品の稿 {n}行"))

    ws8 = wb["08_作業順位"]
    aru = {str(ws8.cell(x, 1).value or "").strip() for x in range(5, ws8.max_row + 1)}
    if "50" not in aru:
        moto = ws8.max_row
        r = moto + 1
        for c in range(1, 11):
            ws8.cell(r, c)._style = copy.copy(ws8.cell(moto, c)._style)
        ws8.row_dimensions[r].height = ws8.row_dimensions[moto].height
        for c, v in enumerate([
                "50", "Red Team 再レビュー（令和8年10月24日）への対応と成果物・エビデンスの一式", "R8.10.24",
                "ご指示による。新規の指摘11件（N01〜N11）と旧指摘の再判定を最新の版で確かめ、是正した",
                "No.85・148・190", "素案 Ver.2.33／公表用・概要版 v2.33／資料 v18／問答集 v15／ZIP",
                "1日", "完了",
                "報酬改定の算定Ｂ・Ｃを表101 から独立に計算し直し（前回の表26 の値は誤り）、素案に書いて資料・"
                "問答集を起こし直した。交付金の「保険者」→市町村、212円、基金の見込額、改頁。"
                "対応表・点検ログ・検算・留保一覧・差分・PDF を ZIP に同梱。"], 1):
            ws8.cell(r, c).value = v
    aru = {str(ws8.cell(x, 1).value or "").strip() for x in range(5, ws8.max_row + 1)}
    kekka.append(("50" in aru, "作業順位 50"))

    wb.save(WBS)
    print("業務工程管理表の更新（Red Team 再レビューへの対応）")
    ng = 0
    for ok, mi in kekka:
        print(("   ○ " if ok else "   × ") + mi)
        ng += not ok
    if ng:
        sys.exit(1)


if __name__ == "__main__":
    main()
