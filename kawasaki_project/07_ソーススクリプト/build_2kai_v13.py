# -*- coding: utf-8 -*-
"""第2回策定委員会資料 v12 → v13（第1回の書式に揃え、修正履歴を除く）
（令和8年10月22日）

  python3 07_ソーススクリプト/build_2kai_v13.py

ご指示（令和8年10月22日）
  (1) 計画素案は策定委員会に未提出であるため、**修正履歴の記載は不要**
  (2) 表の中の文言が中央揃えで見えにくい → **左揃え**に
  (3) **書式は第1回の資料を踏襲する**

第1回の資料（川崎町_第1回策定委員会資料_R8.9.2_v8統合版.docx）と比べて
違っていたもの
  ・⚠ 表紙の帯の網掛けが solid → clear に変わり、白抜きの題名が
    **見えなくなっていた**（題名・計画名が白地に白の文字）
  ・表紙の町のロゴ（画像）がない
  ・表の本文の行に1行おきの網掛け（EFF6F4）がない
  ・注記の記号が ⚠（31件）。第1回は ※・● で、⚠ は0件
  ・2章が1頁目から始まっている（第1回は次第とロゴで1頁、2章から改頁）

除く修正履歴
  ・「◆ 当社資料の訂正（11件）」の節（本文と表）
  ・素案の版（Ver.1.18／Ver.1.20／Ver.2.16／Ver.2.18c／Ver.2.25）と
    資料の版（［第12版］・第１版・第２版）
  ・審議事項①の「本資料の訂正11件」
  ・事業所への照会をやめた経緯、報告書の再修正の経緯、確定値への一本化の経緯
  ※ 第1回にご説明した内容の訂正（5件）は、委員にお示しした内容の訂正で
    あるため残す（第1回資料の条番号の誤りは、本資料の版の記述を除いて残す）

あわせて直す食い違い（いずれも古い値の残り）
  ・調査結果報告書の検算　本文 約6,400項目／表 約6,100項目 → 約6,400項目
    （6,100 は再修正版の前の検算。再修正版は約6,400項目で全数検算）
  ・圧縮後の本文　表 106頁／囲み 102頁 → 106頁
    （102頁は Ver.2.9c の頃の見込み。現行の圧縮案は106頁）
"""
import copy
import importlib.util
import io
import os
import re
import shutil
import sys

import docx
import docx.text.paragraph as dp
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))


def load(name, fn):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(spec)
    sys.modules[name] = m
    spec.loader.exec_module(m)
    return m


V9 = load("build_2kai_v9", "build_2kai_v9.py")

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v12.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v13.docx"
DAI1 = "03_委員会・説明資料/川崎町_第1回策定委員会資料_R8.9.2_v8統合版.docx"
SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.25_軽費老人ホーム確認版.docx"
SOAN_HEN = 118        # 素案の本文の頁数（fill_toc_pages_v200.py の総頁数）
ZEBRA = "EFF6F4"
MIDORI = "1F5C56"

# ── 段落の書き換え（冒頭で当てる。None は削除）
DAN = [
    ("計画骨子（案）とご選択いただきたい事項　［第12版］",
     "計画骨子（案）とご選択いただきたい事項"),
    ("第１回策定委員会でご説明した内容のうち、次の５件を訂正いたします。",
     "第１回策定委員会でご説明した内容のうち、次の５件を訂正いたします。"
     "誤った前提のままご議論いただくことのないよう、冒頭でお示しします。"
     "５件目は、第１回資料に記載していた認知症基本法の条番号の誤りです。"),
    ("◆ 当社資料の訂正（", None),
    ("令和８年９月28日に町へお渡しした本資料の第５版（v5）から", None),
    ("第１回で最も議論が集中した論点です。",
     "第１回で最も議論が集中した論点です。第１０期計画では、サービス提供"
     "体制の維持を新たな重点課題として位置付けることを提案します。供給側の"
     "状況は、厚生労働省「介護サービス情報公表システム」により町内14の指定"
     "すべてについて把握しました（令和８年10月）。定員・在所者数・入所率・"
     "待機者数・職種別の従業者数・指定の年月日・退所先が得られています。"),
    ("令和８年６〜７月に実施した２調査の結果報告書について、",
     "令和８年６〜７月に実施した２調査の結果報告書を確定しました。調査票の"
     "個票から全数の検算を行い（約6,400項目）、第１回でお示しした数値からの"
     "変更はありません。施設入所の検討状況の読み方と、在宅系サービスの利用の"
     "低さの解釈は、選択２・選択３でご説明します。"),
    ("上表は、令和７年度（完結年度）の介護保険事業状況報告（年報）を基礎",
     "上表は、令和７年度（完結年度）の介護保険事業状況報告（年報）を基礎と"
     "して算定した値です。利用率・利用回（日）数・単価は令和７年度で固定して"
     "おり、令和９年度の介護報酬改定は織り込んでいません。"),
    ("計画書の事業一覧表（17表・138事業）に「内容」の欄を加えました。",
     "計画書の事業一覧表（17表・138事業）には、何をする事業なのかが分かる"
     f"よう「内容」の欄を設けています。現在の計画素案は本文{SOAN_HEN}頁で、"
     f"資料編（見込13頁）を加えると計{SOAN_HEN + 13}頁となり、仕様書８の"
     f"「100頁程度」を{SOAN_HEN + 13 - 100}頁上回ります。詳細な分析を別冊に"
     "移すことで、本文を106頁程度とする案を検討しています。"),
    ("①　訂正事項のご確認（第１回のご説明の訂正５件・本資料の訂正11件）",
     "①　訂正事項のご確認（第１回のご説明の訂正５件）"),
    ("⑤　事業所への照会を行わず、介護サービス情報公表システムにより",
     "⑤　介護サービス情報公表システム等の公表資料により供給側を把握する"
     "整理についてのご意見（３）"),
]

# ── 表のセル（表の見出し, 行, 列, 旧の冒頭, 新）。None の行は行ごと削除
SERU = [
    ("区　分", 5, 3, "計画書の法的根拠の記載。素案はVer.1.18で是正済み",
     "計画書の法的根拠の記載"),
    ("確定した事項", 2, 1, "個票（一般高齢者576件・認定者142件）から約6,100項目",
     "個票（一般高齢者576件・認定者142件）から約6,400項目を再集計し、報告書の"
     "値との不一致は０件。χ²検定９件も再計算し完全に一致"),
    ("確定した事項", 4, None, "残る確認事項", None),
    ("⚠ 事業所への照会は行いません。", 0, 0, "⚠ 事業所への照会は行いません。",
     "<表6>"),
    ("町ご提示の体系（案）", 0, 1, "計画素案Ver.1.20の基本目標",
     "計画素案の基本目標"),
    ("最も大きな変動要因は", 0, 0, "最も大きな変動要因は",
     "最も大きな変動要因は介護給付費準備基金の取崩額で、全額と取崩しなしとで"
     "月額1,358円の差があります。次いで第１号被保険者負担割合（24％で＋309円）"
     "と調整交付金の見込交付割合（5.0％で＋274円）です。所得段階別第１号被保険者"
     "数は、令和７年度末の実績（13段階）によっています。"),
    ("版", 0, 0, "版", "区　分"),
    ("版", 1, None, "Ver.2.16", None),
    ("版", 2, 0, "Ver.2.25", "現在の計画素案"),
    ("版", 3, 0, "圧縮案 Ver.2.18c", "詳細な分析を別冊に移した場合（見込み）"),
    ("版", 3, 1, "106頁", "106頁程度"),
    ("版", 3, 2, "119頁", "119頁程度"),
    ("版", 3, 3, "＋19頁", "＋19頁程度"),
    ("仕様書８の「100頁程度」を本文のみと読む", 0, 0, "仕様書８の",
     "仕様書８の「100頁程度」を本文のみと読む場合でも、別冊に移した後の本文は"
     "106頁程度であり、なお６頁程度の圧縮を要します。資料編13頁を含めるか"
     "どうかを含めてご判断をお願いします。"),
]

HYO6 = (
    "● 供給側の状況は、厚生労働省「介護サービス情報公表システム」により、"
    "町内14の指定すべてについて把握しました（令和８年10月）。 定員・在所者数・"
    "入所率・待機者数・職種別の従業者数・前年度の採用者数及び退職者数・指定の"
    "年月日・退所先が公表されています。")

# 当社 → 事務局（計画は町が定めるもの。資料は町の事務局の資料）
TOSHA = [
    ("当社の計画素案は課題を起点とした", "計画素案は課題を起点とした"),
    ("当社はＣをお勧めします", "事務局はＣをお勧めします"),
]


def kakikae(p, s):
    if p.runs:
        p.runs[0].text = V9.clean(s)
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(V9.clean(s))


def sakujo(el):
    el.getparent().remove(el)


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会資料 v12 → v13（第1回の書式・修正履歴を除く）")
    print()
    for f in (SRC, DAI1, SOAN):
        if not os.path.exists(f):
            print(f"   × 入力がありません：{f}")
            sys.exit(1)
    shutil.copy2(SRC, DST)
    d = docx.Document(DST)
    d1 = docx.Document(DAI1)
    mae = "\n".join([p.text for p in d.paragraphs]
                    + [c.text for t in d.tables for r in t.rows
                       for c in r.cells])
    KYU = []

    print("  ── ① 表のセル（表の番号が動く前に行う）")
    MID = [t.rows[0].cells[0].text.strip() for t in d.tables]
    kesu_gyo = []
    n = 0
    for mi, gyo, retsu, kyu, atarashii in SERU:
        hit = [i for i, m in enumerate(MID) if m.startswith(mi)
               and gyo < len(d.tables[i].rows)
               and d.tables[i].rows[gyo].cells[retsu or 0].text.strip()
               .startswith(kyu)]
        if len(hit) != 1:
            judge(False, f"『{mi}』の表が{len(hit)}件")
            continue
        t = d.tables[hit[0]]
        if retsu is None:
            if t.rows[gyo].cells[0].text.strip().startswith(kyu):
                KYU.append("\n".join(c.text for c in t.rows[gyo].cells))
                kesu_gyo.append(t.rows[gyo]._tr)
                n += 1
            else:
                judge(False, f"『{mi}』行{gyo} の旧が合わない")
            continue
        c = t.rows[gyo].cells[retsu]
        if not c.text.strip().startswith(kyu):
            judge(False, f"『{mi}』行{gyo} 列{retsu} の旧が合わない",
                  f"実際『{c.text[:20]}』")
            continue
        KYU.append(c.text)
        if atarashii == "<表6>":
            p0 = c.paragraphs[0]
            p0.runs[0].text = "● 供給側の状況は、厚生労働省「介護サービス情報公表"\
                "システム」により、町内14の指定すべてについて把握しました"\
                "（令和８年10月）。 "
            rest = HYO6.split("（令和８年10月）。 ", 1)[1]
            if len(p0.runs) > 1:
                p0.runs[1].text = rest
                for r in p0.runs[2:]:
                    r.text = ""
            else:
                p0.add_run(rest)
        else:
            V9.set_cell(c, atarashii)
        n += 1
    for tr in kesu_gyo:
        sakujo(tr)
    judge(n == len(SERU), f"{len(SERU)}か所を改めた（実際 {n}）")

    print("  ── ② 段落（修正履歴の削除・書き換え）")
    n = 0
    kesu_el = []
    for atama, atarashii in DAN:
        hit = [p for p in d.paragraphs if p.text.strip().startswith(atama)]
        if len(hit) != 1:
            judge(False, f"『{atama[:20]}』の段落が{len(hit)}件")
            continue
        KYU.append(hit[0].text)
        if atarashii is None:
            kesu_el.append(hit[0]._p)
        else:
            kakikae(hit[0], atarashii)
        n += 1
    # 当社資料の訂正の表（No／訂正の内容／訂正前／訂正後）
    for t in d.tables:
        if [c.text.strip() for c in t.rows[0].cells][:4] == \
                ["No", "訂正の内容", "訂正前", "訂正後"]:
            KYU.append("\n".join(c.text for r in t.rows for c in r.cells))
            kesu_el.append(t._tbl)
    judge(len(kesu_el) == 3, f"当社資料の訂正の節を除いた（見出し・本文・表 {len(kesu_el)}）")
    for el in kesu_el:
        sakujo(el)
    judge(n == len(DAN), f"{len(DAN)}段落を改めた（実際 {n}）")

    print("  ── ③ 当社 → 事務局")
    n = 0
    for p in d.paragraphs:
        for kyu, atarashii in TOSHA:
            if kyu in p.text:
                KYU.append(p.text)
                kakikae(p, p.text.replace(kyu, atarashii))
                n += 1
    judge(n == len(TOSHA), f"{len(TOSHA)}か所（実際 {n}）")

    print("  ── ④ ⚠ を第1回の記号に（行頭は ※、文中は除く）")
    n_kome = n_jo = 0

    def kigo(p):
        nonlocal n_kome, n_jo
        atama = True
        for r in p.runs:
            t = r.text
            if not t:
                continue
            if atama and t.lstrip().startswith("⚠"):
                t = t.replace("⚠ ", "※ ", 1).replace("⚠", "※", 1)
                n_kome += 1
            k = t.count("⚠")
            if k:
                t = t.replace("⚠ ", "").replace("⚠", "")
                n_jo += k
            r.text = t
            atama = False

    for p in d.paragraphs:
        kigo(p)
    for t in d.tables:
        seen = set()
        for row in t.rows:
            for c in row.cells:
                if c._tc in seen:
                    continue
                seen.add(c._tc)
                for p in c.paragraphs:
                    kigo(p)
    judge(True, f"行頭 ⚠→※ {n_kome}件・文中の ⚠ を除いた {n_jo}件")

    print("  ── ⑤ 表紙（帯の網掛けを solid に戻す・ロゴ）")
    n = 0
    for p in d.paragraphs[:4]:
        shd = p._p.pPr.find(qn("w:shd")) if p._p.pPr is not None else None
        if shd is not None and shd.get(qn("w:color")) == "205C56":
            shd.set(qn("w:val"), "solid")
            n += 1
    judge(n == 4, f"帯の4段落を solid に（{n}）")
    # 第1回の次第の後の空き2段落とロゴの段落を写す
    b1 = list(d1.element.body.iterchildren())
    rid = re.findall(r'r:embed="(\w+)"', b1[24].xml)[0]
    blob = d1.part.related_parts[rid].blob
    hako = [t for t in d.tables
            if t.rows[0].cells[0].text.startswith("本日は、分析のご報告ではなく")]
    judge(len(hako) == 1, "次第の下の囲みがある")
    mae_el = hako[0]._tbl
    for el in (b1[22], b1[23]):
        e = copy.deepcopy(el)
        mae_el.addnext(e)
        mae_el = e
    logo = copy.deepcopy(b1[22])
    mae_el.addnext(logo)
    lp = dp.Paragraph(logo, d.paragraphs[0]._parent)
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.add_run().add_picture(io.BytesIO(blob), width=d1.inline_shapes[0].width)
    judge(len(d.inline_shapes) >= 3, f"ロゴを入れた（画像 {len(d.inline_shapes)}点）")
    # 2章から改頁（第1回と同じく次第とロゴで1頁目）
    h2 = [p for p in d.paragraphs if p.text.startswith("２．")]
    if h2:
        ppr = h2[0]._p.get_or_add_pPr()
        if ppr.find(qn("w:pageBreakBefore")) is None:
            ppr.insert(0, OxmlElement("w:pageBreakBefore"))
    judge(len(h2) == 1, "2章の見出しの前で改頁")

    print("  ── ⑥ 表：本文は左揃え・1行おきの網掛け（第1回と同じ）")
    n_hidari = n_zebra = 0
    for t in d.tables:
        midori = False
        tc0 = t.rows[0].cells[0]._tc.tcPr
        if tc0 is not None and tc0.find(qn("w:shd")) is not None:
            midori = tc0.find(qn("w:shd")).get(qn("w:fill")) == MIDORI
        for ri, row in enumerate(t.rows):
            if midori and ri == 0:
                continue                      # 見出し行は中央のまま
            seen = set()
            for c in row.cells:
                if c._tc in seen:
                    continue
                seen.add(c._tc)
                for p in c.paragraphs:
                    if p.text.strip() in ("□", "―", "─"):
                        continue
                    if p.alignment != WD_ALIGN_PARAGRAPH.LEFT:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        n_hidari += 1
                if midori and ri % 2 == 0:
                    shd = c._tc.get_or_add_tcPr().find(qn("w:shd"))
                    if shd is not None and shd.get(qn("w:fill")) == "FFFFFF":
                        shd.set(qn("w:fill"), ZEBRA)
                        n_zebra += 1
    judge(n_hidari > 0, f"左揃えにした段落 {n_hidari}・網掛けしたセル {n_zebra}")

    d.save(DST)

    d2 = docx.Document(DST)
    ato = "\n".join([p.text for p in d2.paragraphs]
                    + [c.text for t in d2.tables for r in t.rows
                       for c in r.cells])
    print("  ── 検査1　修正履歴が残っていないか")
    for go in ("Ver.", "第12版", "第５版", "第１版", "第２版", "（v5）", "訂正前",
               "訂正後", "本資料の訂正", "当社", "是正済み", "一本化",
               "再修正版", "事業所への照会"):
        judge(go not in ato, f"「{go}」が0件")
    judge("第１回のご説明の訂正５件" in ato, "第1回のご説明の訂正5件は残っている")
    print("  ── 検査2　第1回の書式")
    judge("⚠" not in ato, "⚠ が0件（第1回と同じ）")
    nuri = [p._p.pPr.find(qn("w:shd")).get(qn("w:val"))
            for p in d2.paragraphs[:4]]
    judge(nuri == ["solid"] * 4, "表紙の帯が solid（題名が見える）")
    judge(len(d2.inline_shapes) == len(docx.Document(SRC).inline_shapes) + 1,
          "ロゴが1点増えた")
    naka = 0
    for t in d2.tables:
        tc0 = t.rows[0].cells[0]._tc.tcPr
        midori = (tc0 is not None and tc0.find(qn("w:shd")) is not None
                  and tc0.find(qn("w:shd")).get(qn("w:fill")) == MIDORI)
        for ri, row in enumerate(t.rows):
            if midori and ri == 0:
                continue
            for c in row.cells:
                for p in c.paragraphs:
                    if (p.alignment == WD_ALIGN_PARAGRAPH.CENTER
                            and p.text.strip() not in ("□", "―", "─", "")):
                        naka += 1
    judge(naka == 0, f"表の本文に中央揃えの文言が0件（実際 {naka}）")
    print("  ── 検査3　食い違いの是正")
    judge("6,100項目" not in ato and ato.count("6,400項目") >= 2,
          "検算は約6,400項目に統一")
    judge("102頁" not in ato and "106頁程度" in ato, "圧縮後の本文は106頁程度に統一")
    print("  ── 検査4　⚠ 既にあった数が消えていないか")
    kyu_zen = "\n".join(KYU)
    kieta = sorted({m.group(0) for m in re.finditer(
        r"[0-9０-９][0-9０-９,.]*(?:人|名|％|千円|件|頁|日|点|円)", mae)
        if mae.count(m.group(0)) > ato.count(m.group(0))
        and m.group(0) not in kyu_zen})
    judge(not kieta, "消えた数はすべて除いた・書き換えた旧の文に現れている",
          "／".join(kieta[:6]))
    print("  ── 検査5　施設名・法人名・他団体名・強調記号")
    for na in ("蔵王グリーンホーム", "鶴寿会", "釜房みどりの園", "アルパイン",
               "ながさか", "リリーハイム", "北塩原", "金ヶ崎", "川崎市"):
        judge(na not in ato, f"「{na}」が入っていない")
    judge("**" not in ato, "強調記号（**）が0件")

    print()
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        for m in futeki:
            print(f"      {m}")
        sys.exit(1)
    print("   ○ 検査1〜5　不適合0件")


if __name__ == "__main__":
    main()
