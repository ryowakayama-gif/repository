# -*- coding: utf-8 -*-
import openpyxl, warnings
warnings.filterwarnings("ignore")
D="/home/user/repository/kawasaki_project/09_元資料/R8実績データ/R8.9.11受領版/"
N=D+"【川崎町】第10期_将来推計総括表_R8.9.11出力_自然体復帰後.xlsx"
wb=openpyxl.load_workbook(N,data_only=True); ws=wb["2_サービス別給付費"]
KYOJU={"特定施設入居者生活介護","認知症対応型共同生活介護","地域密着型特定施設入居者生活介護"}
SHISETSU={"介護老人福祉施設","介護老人保健施設","介護医療院","地域密着型介護老人福祉施設入所者生活介護"}
cur=None; data={}
for r in range(67,142):
    b,c=(str(ws.cell(r,i).value or "").strip() for i in (2,3))
    vs=[ws.cell(r,x).value for x in range(4,10)]
    if b: cur=b
    if not c or not cur: continue
    data.setdefault(cur,{})[c]=vs
tot7=tot8=0.0; out=[]
for nm,d in data.items():
    if nm in KYOJU or nm in SHISETSU or nm=="未使用": continue
    g=d.get("給付費（千円）")
    if not g or g[1] is None: continue
    q=d.get("回数（回）") or d.get("日数（日）") or d.get("人数（人）")
    g7,g8=g[1] or 0,g[2] or 0; q7,q8=(q[1] or 0),(q[2] or 0)
    ra=(g7/q7)/(g8/q8) if (q7 and q8 and g8) else None
    out.append((nm,g7,g8,ra)); tot7+=g7; tot8+=g8
for nm,g7,g8,ra in sorted(out,key=lambda x:-x[1]):
    print(f"{nm:<30}{g7:>10,.0f}{g8:>10,.0f}"+(f"{ra:>8.3f}" if ra else f"{'—':>8}"))
print(f"{'在宅（介護）計':<30}{tot7:>10,.0f}{tot8:>10,.0f}{tot7/tot8:>8.3f}")
z=sum(wb["1_推計値サマリ"].cell(62,c).value for c in (7,8,9))
d=z*(tot7/tot8-1)
print(f"\n在宅（介護）R9〜R11計={z:,.0f}千円  増≒{d:,.0f}千円  月額≒{d*1000*0.239063/106817:+.1f}円")
