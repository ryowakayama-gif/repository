# -*- coding: utf-8 -*-
"""第2回策定委員会資料 v13 → v14（見える化システムの事業所分析の反映）
（令和8年10月23日）

  python3 07_ソーススクリプト/build_2kai_v14.py

協議用素案を Ver.2.28 に上げた（fix_soan_v2280_jigyosho.py）。素案を直したら
同じ内容を載せている成果品も同じ版で直す（CLAUDE.md）。
  (1) ３の表5の注記：延べ利用日数による稼働率・短期入所の利用・職員数の推移は
      見える化システムの事業所分析で把握した旨を加え、「いずれの資料にもないもの」
      を４件 → ３件（入所前の居所・経営の状況・町への要望）に改める
  (2) ３に5年間の推移の表（素案2-7の表から起こす。4列）と、地域密着型の空き床の段落
  (3) ８の頁数を協議用素案の実測（本文117頁）に揃える（v13 は118頁）

数値は素案 Ver.2.28 の表から読む（委員会資料は数値を別に持たない）。
"""
import copy
import importlib.util
import os
import re
import shutil
import sys

import docx
import docx.table as dt
import docx.text.paragraph as dp

HERE = os.path.dirname(os.path.abspath(__file__))


def load(name, fn):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(spec)
    sys.modules[name] = m
    spec.loader.exec_module(m)
    return m


V9 = load("build_2kai_v9", "build_2kai_v9.py")
SRC = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v13.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v14.docx"
SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.28_協議用.docx"
HEN = 117
SHUTTEN = "地域包括ケア「見える化」システム（事業所分析・各年9月）"


def soan_hyo():
    d = docx.Document(SOAN)
    for t in d.tables:
        if t.rows[1].cells[0].text.startswith("介護老人福祉施設（定員30人）の稼働率"):
            return [[c.text for c in r.cells] for r in t.rows]
    print("   × 素案の5年間の推移の表が見つからない")
    sys.exit(1)


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会資料 v13 → v14（見える化システムの事業所分析）")
    print()
    S = soan_hyo()
    # 4列に：区分・令和3年・令和7年・県内の中央値
    HYO = [[r[0], r[1], r[3], r[4]] for r in S]
    HYO[0] = ["区　分", "令和３年", "令和７年", "宮城県内の中央値（令和７年）"]
    gyo = {r[0]: r for r in S}
    mitsu = [r for r in S if r[0].startswith("地域密着型")][0]
    kai = [r for r in S if r[0].startswith("施設・居住系の介護職員")][0]
    p100 = [r for r in S if r[0].startswith("同 定員100人")][0]
    tan = [r for r in S if r[0].startswith("短期入所")][0]
    m7 = float(mitsu[3].rstrip("％")) / 100
    aki = round(30 * (1 - m7))

    shutil.copy2(SRC, DST)
    d = docx.Document(DST)
    mae = "\n".join([p.text for p in d.paragraphs]
                    + [c.text for t in d.tables for r in t.rows for c in r.cells])
    KYU = []

    print("  ── ① 表5 の注記")
    t5 = [t for t in d.tables
          if t.rows[0].cells[0].text.startswith("● 供給側の状況は")]
    judge(len(t5) == 1, "表5 を特定")
    c = t5[0].rows[0].cells[0]
    p1, p2 = c.paragraphs[1], c.paragraphs[2]
    a1 = "これは時点の在所者数によるもので、延べ利用日数による稼働率ではありません。"
    ok = a1 in p1.text and p2.text.startswith("※ いずれの資料にもないものは４件です。")
    judge(ok, "旧の注記がある")
    if ok:
        KYU += [p1.text, p2.text]
        p1.runs[0].text = p1.text.replace(
            a1, a1 + f"延べ利用日数による稼働率・短期入所の利用・職員数の推移は、{SHUTTEN}で"
                     "令和３年から令和７年まで把握しました（下の４つ目の表）。")
        for r in p1.runs[1:]:
            r.text = ""
        p2.runs[0].text = ("※ いずれの資料にもないものは３件です。 入所前の居所・経営の状況・"
                           "町への要望。入所前の居所は、宮城県指定介護老人福祉施設入所指針が定める"
                           "保険者の協力の場で集計値を求めることを検討し、経営の状況は計画に書かず、"
                           "町への要望は地域ケア会議等の既存の場で聴きます。")
        for r in p2.runs[1:]:
            r.text = ""

    print("  ── ② 5年間の推移の表と段落（入所率の推移の表の後）")
    t_suii = [t for t in d.tables
              if any("入所率（" in x.text for x in t.rows[0].cells)]
    judge(len(t_suii) == 1, "入所率の推移の表を特定")
    if t_suii:
        tn = V9.clone_table_after(t_suii[0], t_suii[0]._tbl, len(HYO))
        for ri, g in enumerate(HYO):
            for ci, v in enumerate(g):
                V9.set_cell(tn.rows[ri].cells[ci], v)
        # 直後に段落（本文の雛形は次の段落を採る）
        tpl = [p for p in d.paragraphs if p.text.startswith("これまで、給付費が伸びない理由が")]
        judge(len(tpl) == 1, "本文の雛形の段落を特定")
        np_ = copy.deepcopy(tpl[0]._p)
        tn._tbl.addnext(np_)
        para = dp.Paragraph(np_, tpl[0]._parent)
        s = (f"上の表は{SHUTTEN}によります。延べ利用日数でみた稼働率は、介護老人福祉施設"
             f"{gyo[S[1][0]][3]}・介護老人保健施設{gyo[S[2][0]][3]}と高い一方、"
             f"地域密着型介護老人福祉施設入所者生活介護は令和５年に定員を20人から30人に"
             f"増やした後も{mitsu[3]}（令和７年）にとどまり、約{aki}人分が埋まっていません。"
             f"短期入所の利用者は月{tan[1]}から{tan[3]}へほぼ半減し、施設・居住系の介護職員"
             f"（常勤換算）は{kai[1]}から{kai[3]}、定員100人あたりで{p100[1]}から{p100[3]}に"
             f"減っています。地域密着型の空きが職員不足による受入れの制限によるものかを"
             f"確かめ、第１０期の見込量の扱いを第３回でご協議いただきます。")
        para.runs[0].text = V9.clean(s)
        for r in para.runs[1:]:
            r.text = ""

    print("  ── ③ 頁数")
    n = 0
    for p in d.paragraphs:
        if "現在の計画素案は本文118頁" in p.text:
            KYU.append(p.text)
            p.runs[0].text = (p.text.replace("本文118頁", f"本文{HEN}頁")
                              .replace("計131頁", f"計{HEN + 13}頁")
                              .replace("31頁上回ります", f"{HEN + 13 - 100}頁上回ります"))
            for r in p.runs[1:]:
                r.text = ""
            n += 1
    for t in d.tables:
        for r in t.rows:
            if r.cells[0].text.strip() == "現在の計画素案" and r.cells[1].text.strip() == "118頁":
                KYU.append("\n".join(c.text for c in r.cells))
                V9.set_cell(r.cells[1], f"{HEN}頁")
                V9.set_cell(r.cells[2], f"{HEN + 13}頁")
                V9.set_cell(r.cells[3], f"＋{HEN + 13 - 100}頁")
                n += 1
    judge(n == 2, f"本文と表の頁数を {HEN}頁に（{n}か所）")
    d.save(DST)

    d2 = docx.Document(DST)
    ato = "\n".join([p.text for p in d2.paragraphs]
                    + [c.text for t in d2.tables for r in t.rows for c in r.cells])
    print("  ── 検査1　素案と同じ値")
    soan_t = "\n".join(c for r in S for c in r)
    for v in (mitsu[3], kai[1], kai[3], tan[1], tan[3]):
        judge(v in ato and v in soan_t, f"「{v}」（素案の表と同じ）")
    print("  ── 検査2　古い記述")
    judge("短期入所の稼働" not in ato, "「短期入所の稼働」0件")
    judge("いずれの資料にもないものは４件" not in ato and "いずれの資料にもないものは３件" in ato,
          "いずれの資料にもないものは３件")
    judge("118頁" not in ato and f"{HEN}頁" in ato, f"頁数 {HEN}頁")
    print("  ── 検査3　消えた数は旧の文に現れる")
    kyu = "\n".join(KYU)
    kieta = sorted({m.group(0) for m in re.finditer(
        r"[0-9０-９][0-9０-９,.]*(?:人|名|％|千円|件|頁|日|点|円)", mae)
        if mae.count(m.group(0)) > ato.count(m.group(0)) and m.group(0) not in kyu})
    judge(not kieta, "消えた数はすべて旧の文に現れる", "／".join(kieta[:5]))
    print("  ── 検査4　事業所名・他団体名・⚠")
    for na in ("釜房", "アルパイン", "鶴寿会", "グループホームさわやか", "グループホームふれあい",
               "北塩原", "金ヶ崎", "大雪", "川崎市"):
        judge(na not in ato, f"「{na}」0件")
    judge("⚠" not in ato and "**" not in ato, "⚠・** 0件（第1回の書式）")

    print()
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜4　不適合0件")


if __name__ == "__main__":
    main()
