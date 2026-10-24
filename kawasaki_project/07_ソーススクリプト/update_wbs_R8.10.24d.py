# -*- coding: utf-8 -*-
"""業務工程管理表の更新（Red Team 再レビュー第3回の影響範囲の再点検（令和8年10月24日））

  python3 07_ソーススクリプト/update_wbs_R8.10.24d.py

  00_管理方法　　　基準日の記載を第3回対応後に（件数は台帳から数える）
  03_確認事項一覧　No.85（保険料の「本試算」・29.9万円の出所）に追記
  04_成果品管理　　稿を 素案 Ver.2.35／公表用・概要版 v2.35／資料 v20 に（問答集 v16 はそのまま）
  08_作業順位　　　52（完了）を加える
再実行しても二重に書かない。
"""
import collections
import copy
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月24日 更新④（第3回レビューの影響範囲の再点検）】"
TAIOU = "川崎町_RedTeam再レビュー第3回_影響範囲再点検の対応表_R8.10.24.md"
ZIP = "11_提出一式/川崎町_RedTeam再レビュー第3回_影響範囲再点検_成果物とエビデンス_R8.10.24.zip"
OKIKAE = [("川崎町_計画書素案_v2.34_協議用.docx", "川崎町_計画書素案_v2.35_協議用.docx"),
          ("川崎町_計画書案_公表用_v2.34.docx", "川崎町_計画書案_公表用_v2.35.docx"),
          ("川崎町_計画書案_概要版_v2.34.docx", "川崎町_計画書案_概要版_v2.35.docx"),
          ("川崎町_第2回策定委員会資料_R8.11_v19.docx", "川崎町_第2回策定委員会資料_R8.11_v20.docx")]
TSUIKI = {
    "85": (f"{HIZUKE}第3回レビューを同じ版で受け直し、同じ趣旨の記述を5冊で走査した。保険料の「本算定」（素案 9-3 の5か所・"
           "資料 表26）は確定の算定と読めるため「本試算」に改めた。入所・居住系の1人1月あたり給付費約29.9万円の出所"
           "（給付費680,226千円÷年報の受給者数延べ2,274人月＝299.13千円）を素案 5-5 に書き、丸める前の値でも10人で"
           "約212円（211.84円）であることを確かめた。素案 Ver.2.35・資料 v20。　"),
}


def iro(c):
    c.font = FONT
    c.alignment = Alignment(wrap_text=True, vertical="top")


def mae_ni(c, s):
    if HIZUKE in str(c.value or ""):
        return 0
    c.value = s + str(c.value or "")
    iro(c)
    return 1


def main():
    wb = openpyxl.load_workbook(WBS)
    kekka = []
    ws3 = wb["03_確認事項一覧"]
    k = 0
    for y in range(5, ws3.max_row + 1):
        no = str(ws3.cell(y, 1).value or "").strip()
        if no in TSUIKI:
            k += mae_ni(ws3.cell(y, 10), TSUIKI[no])
    kekka.append((k in (0, len(TSUIKI)), f"確認事項への追記 {k}件"))
    k3 = collections.Counter(str(ws3.cell(y, 8).value).strip() for y in range(5, ws3.max_row + 1)
                             if str(ws3.cell(y, 1).value or "").strip().isdigit())

    ws0 = wb["00_管理方法"]
    for r in ws0.iter_rows():
        for c in r:
            v = str(c.value or "")
            if v.startswith("令和8年10月24日　本表の記載は同日時点のもの。Red Team 再レビュー第3回（同日）を受け"):
                c.value = ("令和8年10月24日　本表の記載は同日時点のもの。Red Team 再レビュー第3回の影響範囲を5冊で走査し、"
                           "残っていた同じ趣旨の記述（保険料の「本算定」、通いの場の「参加率」、29.9万円の出所）を、"
                           "協議用素案 Ver.2.35・公表用の計画書（案）v2.35・概要版 v2.35・第2回策定委員会資料 v20 で是正した"
                           "（想定問答集 v16 は是正の対象なし）。点検ログ14件はすべて終了コード0。"
                           f"確認事項は{sum(k3.values())}件のうち{k3['完了']}件が完了。一式は {ZIP}。\n"
                           "（前の記載）" + v)
                iro(c)
                kekka.append((True, "基準日の記載を影響範囲の再点検後に"))

    ws4 = wb["04_成果品管理"]
    n = 0
    for y in range(5, ws4.max_row + 1):
        na = str(ws4.cell(y, 3).value or "")
        nb = na
        for a, b in OKIKAE:
            nb = nb.replace(a, b)
        if nb != na:
            ws4.cell(y, 3).value = nb
            mae_ni(ws4.cell(y, 8), f"{HIZUKE}第3回レビューの影響範囲の再点検で是正した版に上げた（{TAIOU}）。　")
            n += 1
    kekka.append((n in (0, 3), f"成果品の稿 {n}行"))

    ws8 = wb["08_作業順位"]
    aru = {str(ws8.cell(x, 1).value or "").strip() for x in range(5, ws8.max_row + 1)}
    if "52" not in aru:
        moto = ws8.max_row
        r = moto + 1
        for c in range(1, 11):
            ws8.cell(r, c)._style = copy.copy(ws8.cell(moto, c)._style)
        ws8.row_dimensions[r].height = ws8.row_dimensions[moto].height
        for c, v in enumerate([
                "52", "Red Team 再レビュー第3回の影響範囲の再点検と成果物・エビデンスの一式", "R8.10.24",
                "ご指示による。受領したレビューは第3回と同じもの（MD5 一致）。指摘と同じ趣旨の記述を5冊で走査した",
                "No.85", "素案 Ver.2.35／公表用・概要版 v2.35／資料 v20／ZIP",
                "1日", "完了",
                "残っていたもの：保険料の「本算定」6か所、通いの場の「参加率」5か所。29.9万円の出所を年報の受給者数から示した。"
                "影響範囲の走査（7指摘×5冊）・検算 §6 を ZIP に同梱"], 1):
            ws8.cell(r, c).value = v
    aru = {str(ws8.cell(x, 1).value or "").strip() for x in range(5, ws8.max_row + 1)}
    kekka.append(("52" in aru, "作業順位 52"))

    wb.save(WBS)
    print("業務工程管理表の更新（第3回レビューの影響範囲の再点検）")
    ng = 0
    for ok, mi in kekka:
        print(("   ○ " if ok else "   × ") + mi)
        ng += not ok
    if ng:
        sys.exit(1)


if __name__ == "__main__":
    main()
