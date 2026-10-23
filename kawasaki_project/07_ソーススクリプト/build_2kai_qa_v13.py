# -*- coding: utf-8 -*-
"""第2回策定委員会 想定問答集 v12 → v13（確定に向けて素案 Ver.2.31・資料 v16 に揃える）
（令和8年10月23日）

  python3 07_ソーススクリプト/build_2kai_qa_v13.py

  (1) 第9期の計画値と実績値の根拠データ：実績を千円単位の値に（資料 表37・素案 図3-2 と同じ）
  (2) 事業一覧表の件数：「17の表・138事業」→ 17の表（134事業と再掲4件）
  (3) 27点：「計画に書くことで取れる」「確実に取り戻せる」→ 評価の対象となり得る（素案 3-3 と同じ）
      資料の見出し「取り戻しをめざす27点」に揃える
  (4) 1問を加える（全68問 → 全69問。カテゴリF）
      QB-69 宮城県の計画とはどう関係するのか（問の番号は全体の通し番号）（素案 1-2 の対応表による）
"""
import collections
import copy
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from data_zuhyo import ZU                                   # noqa: E402

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v12.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v13.docx"
SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.31_協議用.docx"
SU = 69


def txt(el):
    return "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()


def cat_of(s):
    m = re.match(r"^カテゴリ ([A-H])　", s)
    return "カテゴリ " + m.group(1) if m else None


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会 想定問答集 v12 → v13（確定に向けて素案 Ver.2.31・資料 v16 に揃える）")
    f32 = next(z for z in ZU if z["no"] == "図3-2")["rows"]
    soan = docx.Document(SOAN)
    st = "\n".join([p.text for p in soan.paragraphs]
                   + [c.text for t in soan.tables for r in t.rows for c in r.cells])
    t12 = [t for t in soan.tables if [c.text for c in t.rows[0].cells][:3] == ["区分", "計画名", "整合の方針"]][0]
    ken = [r.cells[1].text for r in t12.rows if r.cells[0].text == "県の計画"][0]
    hon = [p.text for p in soan.paragraphs if p.text.startswith("宮城県は、宮城県高齢者福祉計画")][0]
    m = re.search(r"本町は(仙南圏域（\d市\d町）)", hon)
    judge(bool(m) and ken.startswith("第10期みやぎ高齢者元気プラン"), "素案 1-2 から県の計画名と圏域を読む")
    ken_mei = ken.split("（")[0]
    ken_ken = m.group(1) if m else ""

    NAOSU = [
        ("【根拠データ】 第2回資料 ７（第9期の計画値と実績値の対比の表）",
         "居宅 353,651→322,680千円／地域密着型 179,884→182,190千円／施設 477,485→493,350千円",
         f"居宅 {f32[0][1]:,}→{f32[0][2]:,}千円／地域密着型 {f32[1][1]:,}→{f32[1][2]:,}千円／"
         f"施設 {f32[2][1]:,}→{f32[2][2]:,}千円（年報の円単位の額を千円未満四捨五入。計{f32[3][2]:,}千円）"),
        ("A．分かるようにしました。これまでの素案は事業の一覧表に",
         "17の表・138事業すべてに", "17の表（134事業と再掲4件）すべてに"),
        ("【根拠データ】 計画素案 第5章・第6章の事業一覧表（17表138事業）",
         "（17表138事業）", "（17表・134事業と再掲4件。体系全体は第8章の9事業を含め143事業）"),
        ("【根拠データ】 第2回資料 ７（交付金の評価と取り戻せる27点）",
         "取り戻せる27点", "取り戻しをめざす27点"),
        ("A．配点としては大きくありません。ただし、この27点は",
         "仕組みを定めて計画に書くことで取れる部分です。",
         "仕組みを定めて実施することにより評価の対象となり得る部分です（得点は国の該当状況調査への"
         "回答と実績によります）。"),
        ("A．配点としては大きくありません。ただし、この27点は",
         "確実に取り戻せるのはこの27点であり", "取り戻せる見込みが最も高いのはこの27点であり"),
    ]
    ADD = ("カテゴリ F",
           "QB-69　【学識経験者】　宮城県の計画とはどう関係するのか",
           f"A．宮城県は「{ken_mei}」（宮城県高齢者福祉計画と介護保険事業支援計画を一体で定めるもの）を"
           "策定中で、令和8年8月26日の推進委員会に骨子（案）が示されています。基本理念を継承し、"
           "「みんなで支え合う地域づくり」「自分らしい生き方の実現」「安心できるサービスの提供」の"
           f"3つの基本的目標を掲げています。本町は{ken_ken}に属します。計画素案1-2に、県の9つの"
           "基本課題と本計画の章節の対応表を置きました。県の認知症施策推進計画は同プランの1つの章として"
           "策定されるため、県の中間案が示された段階で、本計画の第6章・目標・見込量との整合を確かめます。"
           "また、計画の案を確定する前に宮城県の意見を聴くことが介護保険法第117条第13項で定められて"
           "おり、令和8年12月から事前協議を行います。",
           "【根拠データ】 計画素案 第1章1-2（第10期みやぎ高齢者元気プラン（骨子（案））との対応）・"
           "1-4（法令が定める手続）。宮城県「令和8年度第2回みやぎ高齢者元気プラン推進委員会」資料2〜6",
           "【関連資料】 第2回資料 ８（計画骨子（案）の章構成・今後の予定）")

    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body
    from fix_soan_v111 import set_el, set_cell
    n = 0
    for atama, kyu, atarashii in NAOSU:
        hit = [el for el in body.iterchildren()
               if el.tag == qn("w:p") and txt(el).startswith(atama) and kyu in txt(el)]
        judge(len(hit) == 1, f"『{atama[:22]}』の「{kyu[:16]}」")
        for el in hit:
            set_el(el, txt(el).replace(kyu, atarashii))
            n += 1

    # 加える問（カテゴリF の最後）
    kids = list(body.iterchildren())
    tpl_q = tpl_a = None
    for i, el in enumerate(kids):
        if el.tag == qn("w:p") and txt(el).startswith("QB-1　"):
            tpl_q, tpl_a = copy.deepcopy(el), copy.deepcopy(kids[i + 1])
            break
    cur = last = None
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        c = cat_of(txt(el))
        if c:
            cur = c
        if cur == ADD[0]:
            if txt(el).startswith("【関連資料】"):
                last = el
    q = ADD[1]
    for text, tpl in ((ADD[4], tpl_a), (ADD[3], tpl_a), (ADD[2], tpl_a), (q, tpl_q)):
        el = copy.deepcopy(tpl)
        set_el(el, text)
        last.addnext(el)

    cnt = collections.Counter()
    cur = None
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = txt(el)
        c = cat_of(s)
        if c:
            cur = c
        if re.match(r"^Q[A-H]-\d+　", s):
            cnt[cur] += 1
    total = sum(cnt.values())
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = txt(el)
        mm = re.match(r"^(カテゴリ [A-H]　.+?)（\d+問）$", s)
        if mm:
            set_el(el, f"{mm.group(1)}（{cnt[cat_of(s)]}問）")
        elif re.search(r"全\d+問", s):
            set_el(el, re.sub(r"全\d+問", f"全{total}問", s))
    for t in doc.tables:
        if [c.text.strip() for c in t.rows[0].cells][:3] == ["カテゴリ", "テーマ", "問数"]:
            for r in t.rows[1:]:
                k = "カテゴリ " + r.cells[0].text.strip()
                if k in cnt:
                    set_cell(r.cells[2], str(cnt[k]))
    doc.save(DST)

    d2 = docx.Document(DST)
    zen = "\n".join(p.text for p in d2.paragraphs) + "\n" + "\n".join(
        c.text for t in d2.tables for r in t.rows for c in r.cells)
    print("  ── 検査1　件数")
    judge(total == SU and f"全{SU}問" in zen, f"全{SU}問（{'／'.join(f'{k[-1]}{v}' for k, v in sorted(cnt.items()))}）")
    ns = [int(x) for x in re.findall(r"QB-(\d+)　", zen)]
    judge(sorted(ns) == list(range(1, SU + 1)), f"QB-1〜{SU} 欠番・重複なし（{len(ns)}問）")
    print("  ── 検査2　古い記述")
    for go in ("322,680", "182,190", "493,350", "138事業", "取り戻せる27点", "計画に書くことで取れる",
               "確実に取り戻せる", "第117条第５項"):
        judge(go not in zen, f"「{go}」0件")
    print("  ── 検査3　素案と同じ名称")
    judge(ken_mei in zen and ken_mei in st, f"「{ken_mei}」")
    print("  ── 検査4　事業所名・他団体名")
    for na in ("釜房", "アルパイン", "鶴寿会", "グループホームさわやか", "グループホームふれあい",
               "北塩原", "金ヶ崎", "大雪", "川崎市"):
        judge(na not in zen, f"「{na}」0件")
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜4　不適合0件")


if __name__ == "__main__":
    main()
