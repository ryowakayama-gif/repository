# -*- coding: utf-8 -*-
"""保険料試算の内部整合の検算

レビュー指摘7（保険料試算に数値不整合がある）への対応として、
素案に印字されている値だけを用いて次を確かめる。

  検算1　地域支援事業費の年度別の合計が3か年計に一致するか
  検算2　標準給付費見込額＋地域支援事業費が合計（A＋B）に一致するか
  検算3　年度別の①＝A＋B、C＝①×23％、及びCの内訳の合計が計に一致するか
  検算4　年度別のJ＝C＋D－E、及びJの内訳の合計が計に一致するか
  検算5　月額基準額（算定Ａ）が J÷収納率÷③÷12 に一致するか
  検算6　3パターンの月額基準額が基金取崩額から算定した値に一致するか
  検算7　基金の月額換算が算定式による直接換算に一致するか
  検算8　前提を動かした場合の幅の「差」が基準額の差引きに一致するか
  検算9　13段階別の月額が算定Ｂ×乗率に一致するか

使い方：
  python3 07_ソーススクリプト/check_hokenryo_R8.9.30.py \
      01_第10期_最新版成果品/川崎町_計画書素案_v2.6_図表整理版.docx
"""
import re
import sys
from decimal import ROUND_HALF_UP, Decimal

import docx

RATE = Decimal("0.96")
N = Decimal("9743.1")
K9 = 6_500


def yen(v):
    return int(Decimal(str(v)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def num(s):
    m = re.search(r"([\d,]+(?:\.\d+)?)", s.replace("千円", "").replace("円", ""))
    return Decimal(m.group(1).replace(",", "")) if m else None


def tables(doc):
    out = []
    for t in doc.tables:
        out.append([[c.text.strip() for c in r.cells] for r in t.rows])
    return out


def find(tbls, head):
    for t in tbls:
        if t and t[0][:len(head)] == head:
            return t
    return None


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else (
        "01_第10期_最新版成果品/川崎町_計画書素案_v2.6_図表整理版.docx")
    doc = docx.Document(path)
    tbls = tables(doc)
    body = "\n".join(p.text for p in doc.paragraphs)
    ok, ng = [], []

    def chk(name, cond, detail):
        (ok if cond else ng).append(f"{name}　{detail}")

    # 表101
    t101 = find(tbls, ["記号", "項目", "第10期計"])
    if t101 is None:
        raise SystemExit("表101（保険料算定表）が見つからない")
    row = {r[0]: r for r in t101[1:]}
    A = num(row["A"][2])
    B = num(row["B"][2])
    B_y = [num(row["B"][i]) for i in (3, 4, 5)]
    TOT = num(row["①"][2])
    C = num(row["C"][2])
    D = num(row["D"][2])
    E = num(row["E"][2])
    J = num(row["J"][2])

    chk("検算1", sum(B_y) == B,
        f"地域支援事業費の内訳 {'+'.join(f'{v:,.0f}' for v in B_y)}"
        f"＝{sum(B_y):,.0f}千円／表記 {B:,.0f}千円")
    chk("検算2", A + B == TOT,
        f"A＋B＝{A + B:,.0f}千円／表記 {TOT:,.0f}千円")
    A_y = [num(row["A"][i]) for i in (3, 4, 5)]
    one_y = [num(row["①"][i]) for i in (3, 4, 5)]
    C_y = [num(row["C"][i]) for i in (3, 4, 5)]
    D_y = [num(row["D"][i]) for i in (3, 4, 5)]
    E_y = [num(row["E"][i]) for i in (3, 4, 5)]
    J_y = [num(row["J"][i]) for i in (3, 4, 5)]
    chk("検算3", all(a + b == o for a, b, o in zip(A_y, B_y, one_y)),
        "年度別の ①＝A＋B "
        + "／".join(f"{a + b:,.0f}対{o:,.0f}" for a, b, o in zip(A_y, B_y, one_y)))
    want_C = [(o * Decimal("0.23")).quantize(Decimal("1"),
                                            rounding=ROUND_HALF_UP)
              for o in one_y]
    chk("検算3", want_C == C_y,
        "年度別の C＝①×23％ "
        + "／".join(f"{w:,.0f}対{g:,.0f}" for w, g in zip(want_C, C_y)))
    chk("検算3", sum(C_y) == C,
        f"C の内訳の合計 {sum(C_y):,.0f}千円／表記 {C:,.0f}千円")
    chk("検算4", all(c + d - e == j
                   for c, d, e, j in zip(C_y, D_y, E_y, J_y)),
        "年度別の J＝C＋D－E "
        + "／".join(f"{c + d - e:,.0f}対{j:,.0f}"
                   for c, d, e, j in zip(C_y, D_y, E_y, J_y)))
    chk("検算4", sum(J_y) == J,
        f"J の内訳の合計 {sum(J_y):,.0f}千円／表記 {J:,.0f}千円")

    def geta(j, rate=RATE, n=N):
        return Decimal(j) * 1000 / rate / n / 12

    pa_hyo = num(row.get("", [None, None, None])[2]) if "" in row else None
    for r in t101[1:]:
        if "月額基準額" in r[1]:
            pa_hyo = num(r[2])
    pa = yen(geta(J))
    chk("検算5", pa_hyo == pa,
        f"J÷96.0％÷9,743.1人÷12＝{pa:,}円／表記 {pa_hyo:,.0f}円")

    # 3パターン
    t103 = find(tbls, ["パターン", "基金取崩方針", "特徴", "月額基準額",
                       "第11期影響"])
    kikin = Decimal(152_500)
    want = {"A": yen(geta(J)), "B": yen(geta(J - kikin / 2)),
            "C": yen(geta(J - kikin))}
    for r in t103[1:]:
        got = num(r[3])
        sa = re.search(r"第9期比\s*([+＋▲])([\d,]+)円", r[3])
        chk("検算6", got == want[r[0]],
            f"パターン{r[0]} 算定 {want[r[0]]:,}円／表記 {got:,.0f}円")
        if sa:
            sign = -1 if sa.group(1) == "▲" else 1
            d = sign * int(sa.group(2).replace(",", ""))
            chk("検算6", d == want[r[0]] - K9,
                f"パターン{r[0]} 第9期比 算定 {want[r[0]] - K9:+,}円／表記 {d:+,}円")

    # 基金の月額換算
    m = re.search(r"152,500千円は保険料に換算すると月額約([\d,]+)円", body)
    if m:
        got = int(m.group(1).replace(",", ""))
        chk("検算7", got == yen(geta(kikin)),
            f"基金の月額換算 算定 {yen(geta(kikin)):,}円／本文 {got:,}円")
    t102 = find(tbls, ["項目", "金額・方針"])
    for r in (t102 or [])[1:]:
        if "基金残高（見込み）" in r[0]:
            got = re.search(r"月額約([\d,]+)円分", r[1])
            if got:
                v = int(got.group(1).replace(",", ""))
                chk("検算7", v == yen(geta(kikin)),
                    f"基金の月額換算（表102）算定 {yen(geta(kikin)):,}円／表記 {v:,}円")

    # 幅の表
    t104 = find(tbls, ["前提", "本算定の置き方", "動かした場合", "月額基準額", "差"])
    for r in (t104 or [])[1:]:
        g = num(r[3])
        sa = re.search(r"([+＋▲])([\d,]+)円", r[4])
        if g is None or sa is None:
            continue
        sign = -1 if sa.group(1) == "▲" else 1
        d = sign * int(sa.group(2).replace(",", ""))
        chk("検算8", d == int(g) - pa,
            f"{r[2][:22]} 基準額{int(g):,}円－{pa:,}円＝{int(g) - pa:+,}円／表記 {d:+,}円")

    # 13段階
    t108 = find(tbls, ["段階", "対象", "基準額に対する割合", "第9期 月額",
                       "第10期 月額（算定Ｂ）"])
    pb = want["B"]
    for r in (t108 or [])[1:]:
        mul = re.match(r"([\d.]+)", r[2])
        if not mul:
            continue
        want_v = yen(Decimal(pb) * Decimal(mul.group(1)))
        got = num(r[4])
        chk("検算9", got == want_v,
            f"{r[0]} {mul.group(1)}×{pb:,}円＝{want_v:,}円／表記 {got:,.0f}円")
        g2 = re.search(r"軽減後([\d,]+)円", r[2])
        g3 = re.search(r"軽減後([\d,]+)円", r[4])
        if g2 and g3:
            w = yen(Decimal(pb) * Decimal(re.search(r"軽減後([\d.]+)",
                                                    r[2]).group(1)))
            v = int(g3.group(1).replace(",", ""))
            chk("検算9", v == w,
                f"{r[0]} 軽減後 算定 {w:,}円／表記 {v:,}円")

    print(f"保険料の検算 {path}")
    for s in ng:
        print("  ✗", s)
    print(f"  適合 {len(ok)}件／不適合 {len(ng)}件")
    if ng:
        sys.exit(1)


if __name__ == "__main__":
    main()
