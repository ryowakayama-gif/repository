# -*- coding: utf-8 -*-
"""町内施設の入所状況の推移（県の入所状況調）の管理シート
（令和8年10月22日）

  python3 07_ソーススクリプト/build_nyusho_suii_R8.10.22.py

  00_この表について
  01_入所率の推移　　　　3か月ごと14期・施設別（定員・入所現員・入所率）
  02_入所希望者の推移　　待機者。⚠ 複数施設への申込みが重複して計上される
  03_入退所の推移　　　　新入所・退所・平均待機期間・平均利用期間
  04_退所理由の推移　　　死亡・入院・在宅・他施設
  05_休止・廃止　　　　　⚠ 軽費老人ホームの休止から廃止までの経過
  06_公表システムとの対比　2つの出所で同じことを言えるか
  07_残るもの　　　　　　この調査でも埋まらないもの

入力　09_元資料/その他データ/県入所状況調_川崎町_R8.10.22.csv
出すもの　05_試算・管理シート/川崎町_施設の入所状況の推移_県調査_R8.10.22.xlsx
"""
import csv
import os
import sys
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CSV_IN = "09_元資料/その他データ/県入所状況調_川崎町_R8.10.22.csv"
OUT = ("05_試算・管理シート/"
       "川崎町_施設の入所状況の推移_県調査_R8.10.22.xlsx")
JUN = ["R5.4.1", "R5.7.1", "R5.10.1", "R6.1.1", "R6.4.1", "R6.7.1",
       "R6.10.1", "R7.1.1", "R7.4.1", "R7.7.1", "R7.10.1", "R8.1.1",
       "R8.4.1", "R8.7.1"]
WA = {"R5": "令和5年", "R6": "令和6年", "R7": "令和7年", "R8": "令和8年"}

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="D9E1F2")
WARN = PatternFill("solid", fgColor="FCE4D6")
OK = PatternFill("solid", fgColor="E2EFDA")
WRAP = Alignment(vertical="top", wrap_text=True)


def wa(h):
    y, m, d = h.split(".")
    return f"{WA[y]}{m}月{d}日"


def num(s):
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def sheet(wb, name, head, widths):
    ws = wb.create_sheet(name)
    for c, (h, w) in enumerate(zip(head, widths), start=1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True)
        cell.fill = HEAD
        cell.border = BOX
        cell.alignment = WRAP
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "B2"
    return ws


def put(ws, rows, fill=None):
    for i, r in enumerate(rows, start=2):
        for j, v in enumerate(r, start=1):
            cell = ws.cell(i, j, v)
            cell.border = BOX
            cell.alignment = WRAP
        f = fill(r) if fill else None
        if f:
            for j in range(1, len(r) + 1):
                ws.cell(i, j).fill = f


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    if not os.path.exists(CSV_IN):
        raise SystemExit("入力がない：" + CSV_IN
                         + "\n先に fetch_miyagi_nyusho_R8.10.22.py を実行")
    with open(CSV_IN, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))

    # 施設 → 時点 → 行
    shi = OrderedDict()
    for r in rows:
        shi.setdefault((r["種別"], r["施設名"].split("（")[0]), {})[r["時点"]] = r
    kai = list(shi.keys())

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ══════ 00
    ws = sheet(wb, "00_この表について", ["項", "内容"], [26, 118])
    put(ws, [
        ["何の表か",
         "宮城県「老人福祉施設等入所状況調」から、宮城県柴田郡川崎町"
         "（保険者番号04324）にある施設の推移を取り出したもの。"
         f"{wa(JUN[0])}から{wa(JUN[-1])}までの3か月ごと14時点。"],
        ["なぜ作ったか",
         "⚠ 稼働率・休止床の数は「公表の項目にない」として残していた。"
         "県が3か月ごとに施設別の入所状況を公表していることが分かり、"
         "時点の在所率としてであれば推移が追える。"
         "介護サービス情報公表システムは年1回の報告の1時点であるのに対し、"
         "こちらは3か月ごとの時系列である。"],
        ["出所",
         "https://www.pref.miyagi.jp/soshiki/chouju/nyusyojyoukyou.html　"
         "原典は 09_元資料/その他データ/県入所状況調_R8.10.22取得/ に"
         "退避している（14時点＋留意事項・よくある質問・報告様式・入所指針）。"],
        ["⚠ 定員の定義",
         "報告月1日現在の**長期（短期分を除く）**の定員。"
         "⚠ 短期入所の稼働は分からない。"],
        ["⚠ 入所希望者の定義",
         "報告月1日現在の待機者の人数。⚠ **複数の施設へ申し込めるため"
         "重複して計上されている**（県の注記）。"
         "町内2施設の数を足し合わせると実人数より多くなる。"],
        ["⚠ 介護サービス情報公表システムとの違い",
         "県調査は広域型と地域密着型を**合算**している"
         "（介護老人福祉施設30人＋地域密着型10人＝40人）。"
         "公表システムは指定ごとに分かれる。混ぜない。"],
        ["⚠ 稼働率ではない",
         "入所率（入所現員÷定員）は時点の値であり、"
         "延べ利用日数による稼働率ではない。14時点を並べることで"
         "水準と変動を見る。"],
        ["⚠ 公表値の食い違い",
         "県の留意事項（7）は「退所者の人数と退所理由の合計人数は"
         "等しくなる」としているが、74行のうち1行で合わない"
         f"（{wa('R6.10.1')}・認知症対応型共同生活介護の1事業所）。"
         "⚠ 当方の読み取りではなく公表値そのものの食い違いである。"],
    ])

    # ══════ 01　入所率の推移
    ws = sheet(wb, "01_入所率の推移",
               ["種別", "施設", "項目"] + [wa(h) for h in JUN],
               [12, 22, 10] + [13] * len(JUN))
    rows1 = []
    for (shu, na) in kai:
        for kou in ("定員", "入所現員", "入所率"):
            line = [shu, na, kou]
            for h in JUN:
                r = shi[(shu, na)].get(h)
                if not r:
                    line.append("―")
                    continue
                tei, gen = num(r["定員"]), num(r["入所現員"])
                if kou == "定員":
                    line.append(tei if tei is not None else "⚠ 休止")
                elif kou == "入所現員":
                    line.append(gen if gen is not None else "⚠ 休止")
                else:
                    line.append(round(gen / tei * 100, 1)
                                if tei and gen is not None else "―")
            rows1.append(line)
    put(ws, rows1, lambda r: OK if r[2] == "入所率" else None)

    # ══════ 02　入所希望者
    ws = sheet(wb, "02_入所希望者の推移",
               ["種別", "施設"] + [wa(h) for h in JUN] + ["⚠ 注意"],
               [12, 22] + [13] * len(JUN) + [40])
    rows2 = []
    for (shu, na) in kai:
        line = [shu, na]
        for h in JUN:
            r = shi[(shu, na)].get(h)
            line.append(num(r["入所希望者"]) if r else "―")
        line.append("⚠ 複数施設への申込みが重複して計上されています")
        rows2.append(line)
    put(ws, rows2, lambda r: WARN)

    # ══════ 03　入退所
    ws = sheet(wb, "03_入退所の推移",
               ["種別", "施設", "項目"] + [wa(h) for h in JUN],
               [12, 22, 18] + [13] * len(JUN))
    rows3 = []
    for (shu, na) in kai:
        for kou, key in (("新たな入所者", "新入所人員"),
                         ("平均待機期間（月）", "平均待機期間月"),
                         ("退所者", "退所人員"),
                         ("平均利用期間（月）", "平均利用期間月")):
            line = [shu, na, kou]
            for h in JUN:
                r = shi[(shu, na)].get(h)
                line.append(num(r[key]) if r else "―")
            rows3.append(line)
    put(ws, rows3)

    # ══════ 04　退所理由
    ws = sheet(wb, "04_退所理由の推移",
               ["種別", "施設", "退所理由"] + [wa(h) for h in JUN] + ["計"],
               [12, 22, 10] + [12] * len(JUN) + [8])
    rows4 = []
    for (shu, na) in kai:
        for riyu in ("死亡", "入院", "在宅", "他施設"):
            line = [shu, na, riyu]
            kei = 0
            for h in JUN:
                r = shi[(shu, na)].get(h)
                v = num(r[riyu]) if r else None
                line.append(v if v is not None else "―")
                kei += v or 0
            line.append(kei)
            rows4.append(line)
    put(ws, rows4, lambda r: OK if r[2] == "在宅" else None)

    # ══════ 05　休止・廃止
    ws = sheet(wb, "05_休止・廃止",
               ["時点", "施設", "定員", "入所現員", "一覧での書かれ方"],
               [16, 26, 10, 12, 50])
    rows5 = []
    kyushi_na = [k for k in kai if k[0] == "軽費"]
    for (shu, na) in kyushi_na:
        for h in JUN:
            r = shi[(shu, na)].get(h)
            if r:
                rows5.append([wa(h), r["施設名"],
                              num(r["定員"]) if num(r["定員"]) is not None
                              else "（空欄）",
                              num(r["入所現員"]) if num(r["入所現員"])
                              is not None else "（空欄）",
                              "休止中と記載" if "休止" in r["施設名"]
                              else "通常の記載"])
            else:
                rows5.append([wa(h), na, "―", "―",
                              "⚠ **一覧に載っていない**（廃止）"])
    put(ws, rows5, lambda r: WARN)

    # ══════ 06　公表システムとの対比
    ws = sheet(wb, "06_公表システムとの対比",
               ["項目", "県の入所状況調", "介護サービス情報公表システム",
                "どちらを使うか"], [22, 42, 42, 46])
    put(ws, [
        ["時点", f"3か月ごと（{wa(JUN[0])}〜{wa(JUN[-1])}の14時点）",
         "事業所の年1回の報告の1時点（記入日が事業所ごとに異なる）",
         "⚠ **推移は県調査、指定ごとの内わけは公表システム**"],
        ["単位", "施設単位（広域型と地域密着型を合算）",
         "指定単位（事業所番号×サービスコード）",
         "必要利用定員総数は公表システム、稼働の推移は県調査"],
        ["定員", "長期のみ（短期分を除く）",
         "指定ごとの定員（短期入所も別に載る）",
         "短期入所は公表システム"],
        ["待機者", "入所希望者（⚠ 重複計上）",
         "「入所申込者の数－（入所定員－入所者）」（特養・老健のみ）",
         "⚠ **定義が違う。並べない。**それぞれ定義を明記して使う"],
        ["退所先", "死亡・入院・在宅・他施設（3か月ごと）",
         "自宅等・介護保険施設・医療機関・死亡ほか（前年度又は3か月）",
         "推移は県調査、区分の細かさは公表システム"],
        ["職員数", "⚠ 項目なし", "職種別の従業者数と常勤換算",
         "公表システム"],
        ["対象の広さ", "⚠ **老人福祉法の施設**（養護・軽費を含む）",
         "⚠ **介護保険の指定事業所**（養護・軽費は載らない）",
         "⚠ **高齢者向け住まいは県調査でしか見えない**"],
    ])

    # ══════ 07　残るもの
    ws = sheet(wb, "07_残るもの",
               ["件", "残るもの", "県調査で埋まるか", "代わりに用いる値",
                "確認事項"], [5, 26, 32, 48, 12])
    nokoru = [
        [1, "稼働率（延べ利用日数÷延べ定員日数）",
         "△ 延べではない。時点の入所率が3か月ごとに14期",
         "入所率の推移（01シート）で水準と変動を示す。"
         "⚠ 稼働率とは別のものである旨を明記する", "No.159"],
        [2, "入所前の居所（在宅か医療機関か）",
         "× 項目にない。退所理由（在宅・入院）はある",
         "⚠ **県の入所指針により、施設は入所希望者の置かれている状況を"
         "把握し名簿を5年保管することとされている。** 保険者は入所規程の"
         "作成に協力する責務がある（指針 第3）。新たな調査ではなく、"
         "この関与の中で集計値の提供を求める", "No.159"],
        [3, "休止床の数",
         "△ 施設単位の休止は分かる（05シートに実例）。床単位は不明",
         "定員と入所現員の差の**恒常性**（14期）で代える", "No.159"],
        [4, "経営の状況（収支・人件費率）", "× 項目にない",
         "⚠ **計画に書かない。** 介護事業経営実態調査（全国・"
         "サービス種類別の収支差率）の水準を参照するにとどめる", "No.51"],
        [5, "町への要望・必要な支援", "× 項目にない",
         "地域ケア会議・地域密着型サービス運営協議会という既存の場で聴く。"
         "計画期間中の仕組みとして第7章に位置付け済み", "No.50"],
        [6, "短期入所の稼働",
         "× 定員・入所現員とも長期（短期分を除く）",
         "⚠ **新たに分かった空白。** 公表システムの利用者数で代える",
         "No.159"],
    ]
    put(ws, nokoru,
        lambda r: WARN if str(r[2]).startswith("×") else OK)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)

    # ══════ 自己点検と、メモに使う数の取り出し
    print("町内施設の入所状況の推移（県の入所状況調）")
    print(f"  {OUT}")
    print()
    print("  ── 入所率の推移（％）")
    for (shu, na) in kai:
        v = []
        for h in JUN:
            r = shi[(shu, na)].get(h)
            if not r:
                v.append(None)
                continue
            tei, gen = num(r["定員"]), num(r["入所現員"])
            v.append(gen / tei * 100 if tei and gen is not None else None)
        aru = [x for x in v if x is not None]
        if aru:
            print(f"     {shu:<4}{na:<20}"
                  f"{min(aru):5.1f}〜{max(aru):5.1f}　"
                  f"（最新 {v[-1]:.1f}）" if v[-1] is not None else
                  f"     {shu:<4}{na:<20}{min(aru):5.1f}〜{max(aru):5.1f}")
    print("  ── 入所希望者の推移（人）")
    for (shu, na) in kai:
        v = [num(shi[(shu, na)][h]["入所希望者"])
             for h in JUN if h in shi[(shu, na)]]
        v = [x for x in v if x is not None]
        if v:
            print(f"     {shu:<4}{na:<20}{min(v):.0f}〜{max(v):.0f}　"
                  f"（最初 {v[0]:.0f} → 最新 {v[-1]:.0f}）")
    print("  ── 退所理由「在宅」の計（14期）")
    for (shu, na) in kai:
        kei = sum(num(shi[(shu, na)][h]["在宅"]) or 0
                  for h in JUN if h in shi[(shu, na)])
        print(f"     {shu:<4}{na:<20}{kei:.0f}人")

    print()
    print("  ── 自己点検")
    judge(len(wb.sheetnames) == 8, f"8シート（実際 {len(wb.sheetnames)}）")
    judge(len(kai) == 6, f"施設6件（実際 {len(kai)}）",
          "／".join(n for _, n in kai))
    judge(len(rows1) == len(kai) * 3, f"01シート {len(kai) * 3}行")
    judge(len(nokoru) == 6, "残るもの6件")
    wb2 = openpyxl.load_workbook(OUT, data_only=True)
    zen = []
    for sh in wb2:
        for row in sh.iter_rows(values_only=True):
            zen += [str(v) for v in row if v is not None]
    s = "\n".join(zen)
    judge("川崎市" not in s and "福岡" not in s,
          "⚠ 川崎市（神奈川県）・福岡県川崎町が混じらない")
    # ⚠ 他団体の施設名が混じっていないか
    for na in ("ながさか", "リリーハイム", "ケアハウスやまぶき", "八宮荘"):
        judge(na not in s, f"⚠ 他の市町村の施設「{na}」が入っていない")
    judge("延べ利用日数" in s, "⚠ 稼働率との違いを書いている")

    print()
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        for m in futeki:
            print(f"      {m}")
        sys.exit(1)
    print("   ○ 自己点検 不適合0件")


if __name__ == "__main__":
    main()
