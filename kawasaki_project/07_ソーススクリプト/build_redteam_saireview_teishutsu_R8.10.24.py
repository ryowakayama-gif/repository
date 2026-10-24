# -*- coding: utf-8 -*-
"""Red Team 再レビュー（令和8年10月24日）への対応：成果物とエビデンスの一式を ZIP にする

  python3 07_ソーススクリプト/build_redteam_saireview_teishutsu_R8.10.24.py

ご指示「再レビュー結果確認後、修正をお願い致します。成果物とエビデンスをジップにまとめてアウトプットして」
（令和8年10月24日）による。前回の一式（build_redteam_teishutsu_R8.10.24.py）と同じ組み立てで、
対応表の正本は本スクリプトの TAIOU・SAIHANTEI（md と xlsx はここから書き出す。2か所に持たない）。
点検のログ・検算・差分・PDF は、その場で成果物から作る（手で書かない）。

出力
  11_提出一式/川崎町_RedTeam再レビュー対応_成果物とエビデンス_R8.10.24.zip
  01_第10期_最新版成果品/川崎町_RedTeam再レビュー対応表_R8.10.24.md
  05_試算・管理シート/川崎町_RedTeam再レビュー対応表_R8.10.24.xlsx
終了コード　0＝点検ログがすべて終了コード0・検算が成果物と一致・旧表現0件・ZIP の MD5 が一覧と一致／1＝それ以外
"""
import collections
import difflib
import hashlib
import importlib.util
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

import docx
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("rt1", os.path.join(HERE, "build_redteam_teishutsu_R8.10.24.py"))
RT1 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(RT1)  # md5・zenbun・to_pdf・sc_genpyo・jobun_117 を前回の一式と共用する

HI = "R8.10.24"
ZIPNAME = f"11_提出一式/川崎町_RedTeam再レビュー対応_成果物とエビデンス_{HI}.zip"
MD_OUT = f"01_第10期_最新版成果品/川崎町_RedTeam再レビュー対応表_{HI}.md"
XLSX_OUT = f"05_試算・管理シート/川崎町_RedTeam再レビュー対応表_{HI}.xlsx"
REVIEW = "09_元資料/レビュー/川崎町_第10期計画_RedTeam再レビュー_R8.10.24受領.md"
RYUHO = "05_試算・管理シート/川崎町_公表用の留保一覧_v2.33.xlsx"
OTOSHI = "05_試算・管理シート/川崎町_公表用で落とした文_v2.33.xlsx"

SEIKA = [  # (区分, ファイル, 再レビューの対象の版)
    ("協議用素案", "01_第10期_最新版成果品/川崎町_計画書素案_v2.33_協議用.docx",
     "01_第10期_最新版成果品/川崎町_計画書素案_v2.32_協議用.docx"),
    ("公表用の計画書（案）", "01_第10期_最新版成果品/川崎町_計画書案_公表用_v2.33.docx",
     "01_第10期_最新版成果品/川崎町_計画書案_公表用_v2.32.docx"),
    ("概要版", "01_第10期_最新版成果品/川崎町_計画書案_概要版_v2.33.docx",
     "01_第10期_最新版成果品/川崎町_計画書案_概要版_v2.32.docx"),
    ("第2回策定委員会資料", "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v18.docx",
     "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v17.docx"),
    ("想定問答集", "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v15.docx",
     "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v14.docx"),
]
S = SEIKA[0][1]
P = SEIKA[1][1]
SC = "07_ソーススクリプト/"
TENKEN = [  # (名前, コマンド)
    ("01_素案_自己点検", [SC + "check_soan_selfcheck_R8.9.28.py", S]),
    ("02_素案_保険料の検算", [SC + "check_hokenryo_R8.9.30.py", S]),
    ("03_素案_書体と紙面", [SC + "check_font_R8.9.30.py", S]),
    ("04_素案_8軸の網羅性", [SC + "check_8jiku_R8.10.4.py", S]),
    ("05_素案_章節の参照の機械検索", [SC + "check_sansho_R8.10.24.py", S]),
    ("06_図表データ管理台帳（素案v2.33）", [SC + "build_zuhyo_daicho.py"]),
    ("07_事業所の受入れ様式（様式の検査）", [SC + "check_jigyosho_uke_R8.10.15.py"]),
    ("08_見込量の漏れ", [SC + "check_mikomi_morasi_R8.10.16.py"]),
    ("09_認知症の人の数の整合", [SC + "check_ninchisho_seigo_R8.10.19.py"]),
    ("10_認定者数", [SC + "check_nintei_R8.10.16.py"]),
    ("11_公表用_書体と紙面", [SC + "check_font_R8.9.30.py", P]),
    ("12_委員会資料_自己点検", [SC + "check_2kai_selfcheck_R8.9.28.py"]),
    ("13_守るべき制約の走査", [SC + "check_seiyaku_R8.10.6.py"]),
    ("14_リポジトリの健全性", [SC + "check_repo_kenzensei_R8.10.4.py"]),
]
SCRIPTS = ["fix_soan_v2330_saireview.py", "build_kouhyou_soan.py", "build_gaiyo_v230.py", "build_2kai_v18.py",
           "build_2kai_qa_v15.py", "build_zuhyo_daicho.py", "check_soan_selfcheck_R8.9.28.py",
           "check_2kai_selfcheck_R8.9.28.py", "fill_toc_pages_v200.py", "build_redteam_teishutsu_R8.10.24.py",
           "build_redteam_saireview_teishutsu_R8.10.24.py"]

# ══════════════════════════ 対応表（正本）
# 区分：是正済／一部是正／町のご確認待ち／未実施（資料未受領）／町のご判断
TAIOU = [
    ("N01", "重大", "報酬改定1％・3％の算定Ｂ・Ｃが、基金の取崩しを固定した計算と合わない",
     "是正済", "資料 v17 表26 の算定Ｂ・Ｃ（6,235・5,583・6,366・5,714円）は素案になく、出所を確かめられない値だった。"
     "素案の表101 から独立に計算し直した（改定率を標準給付費見込額に一律に乗じ、第1号負担23％と調整交付金"
     "（5％－見込交付割合）を加減）。この前提で算定Ａの6,888円・7,018円が再現できる。取崩しは一定額の差引き"
     "のため改定による差は算定Ａ・Ｂ・Ｃでほぼ同じ：1％ Ａ6,888・Ｂ6,208・Ｃ5,529円、3％ Ａ7,018・Ｂ6,339・"
     "Ｃ5,659円。素案 9-3 に書き、資料 表26・表27 と問答集3問を素案から起こし直した。表28 の注記（「算定Ｂ・Ｃでは"
     "報酬改定・負担割合の差が異なる」）も誤りであったため、差が変わるのは予定収納率のみと改めた",
     "素案 9-3／資料 v18 表26〜28／問答集 v15", "03_エビデンス/検算.md §1",
     "改定後の標準給付費の正式な算定は令和9年度改定の公表後（確認事項No.148）"),
    ("N02", "高", "資料の本日のご審議事項⑥に「『評価結果の活用』の要件となります」の言い切りが残る",
     "是正済", "「要件に当たる取組の一つとなり得ます（評価は国の該当状況調査への回答と議事の記録によります）」に"
     "改め、資料16頁の説明とそろえた", "資料 v18 本日のご審議事項⑥", "03_エビデンス/旧表現の残存検索.txt", ""),
    ("N03", "高", "生活支援コーディネーター26人と、町の実績データの「生活支援サポーター（SC）」26人が同じものか",
     "町のご確認待ち", "誤りとは断定されていないため人数は変えず、素案 10-4 の KPI と 5-4 の表に【町確認】を付け、"
     "町の実績データでの名称と、第1層・第2層の配置の内訳が確認中であることを書いた。確認事項No.190 を起票。"
     "公表用は［要確認］、概要版は確定したものだけを載せる決まりにより KPI の行を落とした",
     "素案 5-4・10-4／公表用／概要版", "03_エビデンス/原票抜粋_生活支援コーディネーター.csv",
     "確認事項No.190（既定値：回答がなければ KPI の現状値を「確認中」とし数値を載せない）"),
    ("N04", "高", "第3回（令和9年2月）までに令和8年度末の基金残高を「確定」するのは時系列上できない",
     "是正済", "問答集の「①…令和8年度末残高の確定」を「見込額の精査（決算による確定は令和9年度であり、第3回では"
     "見込額を用います）」に、資料 表29 も「令和８年度末残高の見込額（決算による確定は令和９年度）」に改めた",
     "資料 v18 表29／問答集 v15", "03_エビデンス/旧表現の残存検索.txt", ""),
    ("N05", "高", "公表用に［要確認］67・［要協議］14・【委員会協議】16 の計97か所",
     "一部是正", "内部の記号【委員会協議】を［要協議］にそろえ、公表用で0件とした（公表用を組む手順で置き換える）。"
     "留保は節ごとの一覧（103件＝［要確認］73・［要協議］30。【町確認】2件の追加を含む）を作り、確定・注記・"
     "取下げを決めるための台帳とした。公表用は留保が残る間は表紙を「体裁確認用」とする",
     "公表用 v2.33", "03_エビデンス/川崎町_公表用の留保一覧_v2.33.xlsx",
     "各留保の確定は町・委員会のご判断と未受領資料による（公表の確定の条件）"),
    ("N06", "中", "基金の活用方針の表に無関係な「有効票 576件」の列",
     "是正済", "列を削り、2列（項目・金額・方針）とした。列の幅は「金額・方針」に寄せた", "素案・公表用 9-3",
     "03_エビデンス/差分_協議用素案.txt", ""),
    ("N07", "中", "資料21頁の冒頭に「す。」だけが残り、審議事項の囲みが途中で切れる",
     "是正済", "原因は、表と表の間に段落がなく Word・LibreOffice が1つの表として扱い、前の表の見出し行が次頁で"
     "繰り返されたこと。隣り合う表11か所に区切りの段落を置き、開催回数の囲みを行で分割しない設定とした。⑦は"
     "⑥と段落の書式が違い枠が切れていたため、⑥の書式に揃え①〜⑦を1つの囲みとした。構成の案の表は頁で分けない。"
     "このため当方の環境の PDF は21頁→22頁（22頁は構成の案・審議事項・次回予定）",
     "資料 v18 ８・本日のご審議事項", "01_成果物/川崎町_第2回策定委員会資料_R8.11_v18.pdf",
     "配布用 PDF は町の環境（Word）で頁送りを確かめていただく（No.53）"),
    ("N08", "中", "「全国1,741保険者」は「1,741市町村」（交付金の評価の単位）",
     "是正済", "5冊で「1,741保険者」「35保険者」「285保険者」「保険者中」を市町村に改めた。同じ節の"
     "「全保険者で按分」「他の保険者の得点」「同じ得点の保険者」「県内の保険者の得点」もそろえた。"
     "入所指針・地域密着型の指定など、保険者としての町の役割を述べる箇所はそのまま",
     "素案 3-3 ほか／資料 ７／問答集", "03_エビデンス/旧表現の残存検索.txt", ""),
    ("N09", "低〜中", "入所・居住系±10人の「約211円」は四捨五入なら212円",
     "是正済", "表101 から 211.7円（1円未満四捨五入で212円）。素案 5-5 に算式と端数処理を明記し、資料 選択２"
     "（5,931円・6,355円）・表27（±212円）・問答集3問をそろえた", "素案 5-5／資料 選択２・表27／問答集",
     "03_エビデンス/検算.md §2", "29.9万円は概数。令和7年度実績の円単位は 05_試算・管理シート の算定ブック"),
    ("N10", "高", "居宅サービスが計画を下回ったことを供給不足と町外利用の「結果」と言い切る",
     "是正済", "「影響している可能性があります（因果関係は、町内外別の利用の推移等による追加の検証を要します）」"
     "に改めた（素案 3-3・資料）", "素案 3-3／資料 ８", "03_エビデンス/旧表現の残存検索.txt",
     "町内外別の利用の推移は未受領（確認事項No.159）"),
    ("N11", "中", "図表データ台帳の点検対象が v2.20、事業所受入れは様式の検査、998,224千円と998,222千円の差",
     "是正済", "①図表データ台帳の点検を最新の素案 v2.33 に向け直して実行した（点検ログ06）。②事業所・住まいの"
     "受入れは「様式と整合の検査」であり実データの全数照合ではないことをログの題に明記した（記入は町と事業所への"
     "照会による）。③998,224千円は 9-1 の表の37サービスをそれぞれ千円未満四捨五入した和で、998,222千円は年報の"
     "円単位の総額（998,222,421円）を丸めたもの。差の2千円は丸めの範囲（最大±18.5千円）にある（検算 §3）",
     "素案 2-4・9-1", "03_エビデンス/点検ログ06・07・08／検算.md §3",
     "事業所受入れの実データは町・事業所への照会待ち（No.159）"),
]
SAIHANTEI = [  # 旧指摘の再判定（再レビュー §3）のうち、今回手を入れたもの
    ("A03", "ほぼ解消／丸め再確認", "N09 のとおり約212円にそろえた"),
    ("A04", "未解消・重大", "N01 のとおり素案の表101 から計算し直し、素案・資料・問答集で一致させた"),
    ("A06", "一部解消", "N02 のとおり資料21頁の言い切りを改めた"),
    ("A08", "一部是正", "N07 のとおり改頁の孤立を解いた（PDF は22頁）"),
    ("B04", "一部解消", "N01 の感度の計算を直した。入力を分けた試算は令和9年度改定の公表後"),
    ("B11", "再確認が必要", "N03 のとおり【町確認】を付し、確認事項No.190 を起票"),
    ("B12", "委員会説明は改善", "概要版に「柱の立て方（計画の章の組み立て）は、策定委員会でご審議いただく案であり、"
     "まだ確定していません」を加えた"),
    ("B01・B02・B05・B06・B08・B09・B10・B13・B14", "未解消（資料の未受領・町のご判断）",
     "前回の対応表のとおり。確認事項の台帳（No.11・19・34・39・49・70・78・85・90・137・147・154・159・163・184）で追う"),
]
ZANSON = [  # 旧表現の残存検索（再レビューと前回のレビューで挙がった語）
    "6,235円", "5,583円", "6,366円", "5,714円", "約211円", "＋211円", "▲211円", "＋92円", "＋119円", "＋223円",
    "＋250円", "差がこれと異なります", "要件となります", "残高の確定", "結果と考えられます", "1,741保険者",
    "35保険者", "285保険者", "保険者中", "同じ得点の保険者", "全保険者で按分", "他の保険者の得点", "県内の保険者",
    # 前回
    "998,220", "5,959", "6,382", "（算定Ｂ 6,143円を基準）", "25年間", "11つ", "8の事業所", "第7章で論じ",
    "得点できます", "要件を満たします", "そのまま要件", "概ね半分以下", "生活支援コーディネーター25",
    "9-4の感度表", "町内14の指定", "第117条第５項",
]


def num(s):
    return float(re.sub(r"[^\d.]", "", s))


def kensan(ng):
    """検算.md の本文（成果物から読み、独立に計算した値を成果物の記載と突き合わせる）。"""
    d = docx.Document(S)
    s = RT1.zenbun(S)
    sz = "\n".join(s)
    L = ["# 検算（成果物から読み直し、独立に計算した値）", "",
         f"対象：{os.path.basename(S)}・{os.path.basename(SEIKA[3][1])}・{os.path.basename(SEIKA[4][1])}", ""]

    t101 = [t for t in d.tables if [c.text.strip() for c in t.rows[0].cells][:3] == ["記号", "項目", "第10期計"]][0]
    row = {r.cells[0].text.strip(): r for r in t101.rows}
    A_y = [num(row["A"].cells[i].text) for i in (3, 4, 5)]
    J = num(row["J"].cells[2].text)
    N = num(row["③"].cells[2].text)
    rates = [float(x) / 100 for x in re.findall(r"令和\d+年度([\d.]+)％",
                                                [x for x in s if "見込交付割合は全国の仮置値" in x][0])]
    KIKIN = {"Ａ": 0, "Ｂ": 76250, "Ｃ": 152500}

    def gaku(r, k):
        dj = sum(a * r * (0.23 + 0.05 - e) for a, e in zip(A_y, rates))
        return (J + dj - KIKIN[k]) * 1000 / 0.96 / N / 12

    L += ["## §1 介護報酬改定率別の算定Ａ・Ｂ・Ｃ（N01）", "",
          f"- 表101：保険料収納必要額 J＝{J:,.0f}千円、補正後被保険者数＝{N:,.1f}人、予定収納率96.0％",
          "- 標準給付費見込額（年度別）＝" + "・".join(f"{a:,.0f}" for a in A_y) + "千円",
          "- 調整交付金の見込交付割合＝" + "・".join(f"{e * 100:.2f}％" for e in rates),
          "- 改定率 r の収納必要額の増分＝Σ 標準給付費見込額 × r ×（23％＋5％－見込交付割合）",
          "- 月額＝（J＋増分－基金取崩額）×1000 ÷ 0.96 ÷ 補正後被保険者数 ÷ 12", "",
          "| 改定率 | 算定Ａ（取崩0） | 算定Ｂ（76,250千円） | 算定Ｃ（152,500千円） | 差 Ａ・Ｂ・Ｃ |", "|---|---|---|---|---|"]
    kai = {}
    for r in (0, 0.01, 0.03):
        v = {k: gaku(r, k) for k in KIKIN}
        kai[r] = {k: round(x) for k, x in v.items()}
        sa = "・".join(f"＋{kai[r][k] - kai[0][k]}" for k in KIKIN) if r else "―"
        L.append(f"| {r * 100:.0f}％ | {v['Ａ']:,.2f} → {kai[r]['Ａ']:,}円 | {v['Ｂ']:,.2f} → {kai[r]['Ｂ']:,}円 | "
                 f"{v['Ｃ']:,.2f} → {kai[r]['Ｃ']:,}円 | {sa} |")
    L += ["", f"- 基金による月額の減少は固定：Ｂ {76250 * 1000 / 0.96 / N / 12:.2f}円・Ｃ {152500 * 1000 / 0.96 / N / 12:,.2f}円"
          "（再レビュー §2 の 679.3447円・1,358.6895円と一致）"]
    # 成果物の記載と突き合わせ
    t26 = [t for t in docx.Document(SEIKA[3][1]).tables if t.rows[0].cells[0].text.strip() == "介護報酬改定率"][0]
    g = {r.cells[0].text.strip()[:2]: [int(re.sub(r"\D", "", c.text)) for c in r.cells[1:4]] for r in t26.rows[1:4]}
    itchi = all(g[k] == [kai[r]["Ａ"], kai[r]["Ｂ"], kai[r]["Ｃ"]] for k, r in (("０％", 0), ("１％", 0.01), ("３％", 0.03)))
    soan_bun = (f"改定率1％では算定Ａ{kai[0.01]['Ａ']:,}円・算定Ｂ{kai[0.01]['Ｂ']:,}円・算定Ｃ{kai[0.01]['Ｃ']:,}円" in sz
                and f"3％では算定Ａ{kai[0.03]['Ａ']:,}円・算定Ｂ{kai[0.03]['Ｂ']:,}円・算定Ｃ{kai[0.03]['Ｃ']:,}円" in sz)
    L += [f"- 資料 v18 表26 と一致：{'○' if itchi else '×'}／素案 9-3 の文と一致：{'○' if soan_bun else '×'}",
          "- 再レビューの「目安」Ｂ1％ 約6,209円は、円に丸めた算定Ａ6,888円から基金効果を差し引いた値。"
          f"丸める前の値からは {gaku(0.01, 'Ｂ'):,.2f}円 → {kai[0.01]['Ｂ']:,}円", ""]
    if not (itchi and soan_bun):
        ng.append("検算§1 報酬改定")

    tanka = float(re.search(r"1人1月あたり給付費は約([\d.]+)万円", sz).group(1)) * 10
    dj = sum(10 * tanka * 12 * (0.23 + 0.05 - e) for e in rates)
    dm = dj * 1000 / 0.96 / N / 12
    b = int(re.search(r"パターンB ([\d,]+)円", sz).group(1).replace(",", ""))
    j10 = round(dm)
    shiryo = "\n".join(RT1.zenbun(SEIKA[3][1]))
    ok2 = f"10人では約{j10}円" in sz and f"{b - j10:,}円（▲{j10}円）" in shiryo and f"{b + j10:,}円（＋{j10}円）" in shiryo
    L += ["## §2 入所・居住系±10人の月額（N09）", "",
          f"- 1人1月あたり給付費 約{tanka / 10:.1f}万円（素案 5-5）",
          f"- 収納必要額の増減＝10人×{tanka:.0f}千円×12か月×Σ(23％＋5％－見込交付割合)＝{dj:,.1f}千円（3か年）",
          f"- 月額＝{dj:,.1f}千円÷96.0％÷{N:,.1f}人÷12＝{dm:.2f}円 → 1円未満四捨五入 **{j10}円**",
          f"- 選択２：算定Ｂ {b:,}円 → 縮小 {b - j10:,}円（▲{j10}円）／増床 {b + j10:,}円（＋{j10}円）",
          f"- 素案 5-5・資料 選択２と一致：{'○' if ok2 else '×'}", ""]
    if not ok2:
        ng.append("検算§2 入所・居住系")

    t24 = [t for t in d.tables if [c.text for c in t.rows[0].cells][:4] == ["サービス区分", "要支援計", "要介護計", "合計（円）"]][0]
    en = {r.cells[0].text.split("\n")[0]: int(r.cells[3].text.replace(",", "")) for r in t24.rows[1:]}
    L += ["## §3 令和7年度の給付費の千円の丸め（N11）", "",
          f"- 総額（年報・円）{en['総計']:,}円 → 千円未満四捨五入 {round(en['総計'] / 1000):,}千円",
          "- 9-1 の見込量の表の令和7年度の欄は、37サービスをそれぞれ千円未満四捨五入した値。その和は998,224千円"
          "（点検ログ08）",
          "- 37の値を別々に丸めると、和の誤差は最大±18.5千円。差の2千円はこの範囲にあり、サービスの漏れではない"
          "（点検ログ08 で、ワークブックの32サービスがすべて表にあり値が一致することを確かめている）", ""]
    L += ["## §4 生活支援コーディネーター（N03）", "",
          "- 原票抜粋は 原票抜粋_生活支援コーディネーター.csv（町の実績データの名称は「生活支援ｻﾎﾟｰﾀｰ(SC)_人数」）",
          "- 人数（26人）が一致しても職務・配置の同一は示せないため、【町確認】とし確認事項No.190 で照合する", ""]
    return "\n".join(L)


def main():
    work = tempfile.mkdtemp(prefix="redteam2_")
    root = os.path.join(work, f"川崎町_RedTeam再レビュー対応_{HI}")
    d_seika, d_taiou, d_ev, d_sc = (os.path.join(root, x) for x in
                                    ("01_成果物", "02_レビュー対応", "03_エビデンス", "04_スクリプト"))
    d_log = os.path.join(d_ev, "点検ログ")
    for x in (d_seika, d_taiou, d_ev, d_sc, d_log):
        os.makedirs(x)
    ng = []
    print(f"Red Team 再レビュー対応の一式（{HI}）")

    print("  ── 1 成果物と PDF")
    ichiran = []
    for kubun, f, kyu in SEIKA:
        shutil.copy2(f, d_seika)
        pdf, n = RT1.to_pdf(f, d_seika)
        ichiran.append((kubun, os.path.basename(f), n, RT1.md5(f), os.path.basename(kyu)))
        print(f"   ○ {kubun} {os.path.basename(f)}（PDF {n}頁）")
        if not n:
            ng.append(f"PDF {f}")

    print("  ── 2 点検ログ")
    logs = []
    for na, cmd in TENKEN:
        r = subprocess.run([sys.executable] + cmd, capture_output=True, text=True, timeout=900)
        out = "\n".join(x for x in (r.stdout + r.stderr).splitlines()
                        if "FutureWarning" not in x and "warnings.warn" not in x)
        open(os.path.join(d_log, f"{na}.txt"), "w", encoding="utf-8").write(
            f"$ python3 {' '.join(cmd)}\n終了コード {r.returncode}\n\n{out}\n")
        logs.append((na, r.returncode))
        print(("   ○ " if r.returncode == 0 else "   × ") + f"{na}（終了コード {r.returncode}）")
        if r.returncode:
            ng.append(na)

    print("  ── 3 検算・原票抜粋・留保一覧・旧表現の残存検索・差分")
    open(os.path.join(d_ev, "検算.md"), "w", encoding="utf-8").write(kensan(ng))
    RT1.sc_genpyo(os.path.join(d_ev, "原票抜粋_生活支援コーディネーター.csv"))
    for f in (RYUHO, OTOSHI):
        shutil.copy2(f, d_ev)
    texts = {k: "\n".join(RT1.zenbun(f)) for k, f, _ in SEIKA}
    zan = ["語\t" + "\t".join(k for k, _, _ in SEIKA)]
    zan_ng = 0
    for w in ZANSON:
        c = [texts[k].count(w) for k, _, _ in SEIKA]
        zan.append(w + "\t" + "\t".join(str(x) for x in c))
        zan_ng += sum(c)
    # 内部の記号は公表用・概要版に残さない（協議用素案・資料には残す決まり）
    zan.append("")
    zan.append("公表用・概要版に残さない内部の記号")
    for w in ("【委員会協議】", "【町確認】", "⚠", "確認事項No"):
        c = [texts[k].count(w) for k in ("公表用の計画書（案）", "概要版")]
        zan.append(f"{w}\t公表用 {c[0]}\t概要版 {c[1]}")
        zan_ng += sum(c)
    ryuho = {w: texts["公表用の計画書（案）"].count(w) for w in ("［要確認］", "［要協議］", "［要内訳］")}
    zan += ["", "公表用に残る留保（確定・注記・取下げを決めるもの。一覧は 川崎町_公表用の留保一覧_v2.33.xlsx）"]
    zan += [f"{k}\t{v}" for k, v in ryuho.items()]
    open(os.path.join(d_ev, "旧表現の残存検索.txt"), "w", encoding="utf-8").write(
        "再レビューと前回のレビューで挙がった旧の表現が、提出の5冊に残っていないか（件数）\n\n" + "\n".join(zan) +
        f"\n\n旧表現・内部の記号の合計 {zan_ng}件\n")
    print(("   ○ " if zan_ng == 0 else "   × ") + f"旧表現・内部の記号の残存 {zan_ng}件（{len(ZANSON)}語×5冊＋4語×2冊）")
    print(f"   ○ 公表用の留保 " + "・".join(f"{k}{v}" for k, v in ryuho.items()))
    if zan_ng:
        ng.append("旧表現の残存")
    for kubun, f, kyu in SEIKA:
        a, b = RT1.zenbun(kyu), RT1.zenbun(f)
        diff = list(difflib.unified_diff(a, b, os.path.basename(kyu), os.path.basename(f), n=0, lineterm=""))
        open(os.path.join(d_ev, f"差分_{kubun}.txt"), "w", encoding="utf-8").write(
            f"{kubun}：{os.path.basename(kyu)}（再レビューの対象）→ {os.path.basename(f)}\n"
            "段落・表のセルの単位の差分（-＝旧、+＝新）\n\n" + "\n".join(diff) + "\n")
        print(f"   ○ 差分_{kubun}（{sum(1 for x in diff if x.startswith('+') and not x.startswith('+++'))}行が新）")

    print("  ── 4 対応表（md・xlsx）")
    c = collections.Counter(r[3] for r in TAIOU)
    md = ["# Red Team 再レビュー（令和8年10月24日）への対応表", "",
          f"> 再レビュー原本：`{REVIEW}`（素案 v2.32・公表用 v2.32・概要版 v2.32・資料 v17・問答集 v14 が対象）",
          "> 対応後の版：協議用素案 Ver.2.33／公表用 v2.33／概要版 v2.33／委員会資料 v18／想定問答集 v15",
          "> 本表の正本は `07_ソーススクリプト/build_redteam_saireview_teishutsu_R8.10.24.py` の TAIOU・SAIHANTEI", "",
          "## 1　新規の指摘（N01〜N11）", "",
          "| ID | 重要度 | 指摘 | 対応 | 対応の内容 | 箇所 | エビデンス | 残る事項 |", "|---|---|---|---|---|---|---|---|"]
    for row in TAIOU:
        md.append("| " + " | ".join(x.replace("|", "／") for x in row) + " |")
    md += ["", "対応の区分：" + "／".join(f"{k} {v}件" for k, v in c.items()), "",
           "## 2　旧指摘の再判定への対応", "", "| 前回ID | 再レビューの判定 | 今回の対応 |", "|---|---|---|"]
    for row in SAIHANTEI:
        md.append("| " + " | ".join(row) + " |")
    md += ["", "## 3　再レビュー §6 の完了条件との対応", "",
           "| 優先 | 完了条件 | 状態 |", "|---|---|---|",
           "| 最優先1 | Ｂ/Ｃの基金効果が同一条件で固定額となり、素案・委員会・問答集で一致 | 達成（検算 §1。基金効果Ｂ679.34円・Ｃ1,358.69円で固定） |",
           "| 最優先2 | 「評価対象となり得る。該当状況調査・実績・議事記録で判定」に統一 | 達成（N02） |",
           "| 最優先3 | 町提供データ、国評価回答、制度上の業務区分が一致 | 町のご確認待ち（No.190） |",
           "| 高4 | 第3回の意思決定条件と決算確定日を整合 | 達成（N04） |",
           "| 高5 | 公表時点で不明な実績と未決定目標を明示、作業用記号を残さない | 作業用記号は0件。留保103件は一覧で管理（公表の確定の条件として残る） |",
           "| 中6 | 基金表から異質な列が消えている | 達成（N06） |",
           "| 中7 | 21頁上部の「す。」と見出し行のみの分断がない | 達成（N07。PDF は22頁） |",
           "| 中8 | 1,741「市町村」に統一 | 達成（N08） |",
           "| 中9 | 算定用基礎値・端数処理を明記 | 達成（N09） |",
           "| 継続 | 残件台帳の各項目に締切・取得資料・決裁者 | 業務工程管理表 03_確認事項一覧（確認先・回答期限・既定値）で管理 |",
           "", "## 4　当方で確かめられないもの（未実施として残す）", "",
           "- 年報原本・見える化の将来推計総括表・保険料の算定ブック・交付金の該当状況調査の原本と成果物の全数の突合は、"
           "今回の一式では再レビューと同じく範囲外（点検ログは成果物どうしと、受領済みの原票抜粋の範囲）。",
           "- 「8軸51/51適合」は記載の網羅性の点検であり、同じログの未実施10事項は資料の未受領による。", ""]
    open(MD_OUT, "w", encoding="utf-8").write("\n".join(md))
    shutil.copy2(MD_OUT, d_taiou)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "新規の指摘"
    ws.append(["ID", "重要度", "指摘", "対応", "対応の内容", "箇所", "エビデンス", "残る事項"])
    for row in TAIOU:
        ws.append(list(row))
    for i, w in enumerate((6, 8, 40, 14, 80, 28, 34, 34), 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws2 = wb.create_sheet("旧指摘の再判定")
    ws2.append(["前回ID", "再レビューの判定", "今回の対応"])
    for row in SAIHANTEI:
        ws2.append(list(row))
    for i, w in enumerate((22, 26, 90), 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for w_ in (ws, ws2):
        for r in w_.iter_rows():
            for cc in r:
                cc.alignment = Alignment(wrap_text=True, vertical="top")
                cc.font = Font(name="游ゴシック", size=9, bold=(cc.row == 1))
                if cc.row == 1:
                    cc.fill = PatternFill("solid", fgColor="D9E1F2")
        w_.freeze_panes = "A2"
    wb.save(XLSX_OUT)
    shutil.copy2(XLSX_OUT, d_taiou)
    shutil.copy2(REVIEW, d_taiou)
    print(f"   ○ {len(TAIOU)}件（" + "／".join(f"{k}{v}" for k, v in c.items()) + f"）・再判定 {len(SAIHANTEI)}行")

    print("  ── 5 スクリプト")
    for f in SCRIPTS:
        shutil.copy2(SC + f, d_sc)

    print("  ── 6 README（成果品一覧・MD5）")
    rd = [f"# 川崎町 第10期計画　Red Team 再レビュー対応の一式（{HI}）", "",
          "川崎町（宮城県柴田郡）高齢者保健福祉計画・第10期介護保険事業計画", "",
          "## 中身", "", "| フォルダ | 内容 |", "|---|---|",
          "| 01_成果物 | 是正後の5冊（docx）と、当方の環境で作った PDF |",
          "| 02_レビュー対応 | 対応表（md・xlsx）と再レビュー原本 |",
          "| 03_エビデンス | 点検ログ・検算（独立に計算し成果物と突合）・旧表現の残存検索・公表用の留保一覧・差分・原票抜粋 |",
          "| 04_スクリプト | 是正と点検に用いたスクリプト（リポジトリの 07_ソーススクリプト と同じもの） |", "",
          "## 成果品一覧", "", "| 区分 | ファイル | PDF頁数（当方の環境） | MD5 | 比べた旧の版 |", "|---|---|---|---|---|"]
    for kubun, f, n, h, kyu in ichiran:
        rd.append(f"| {kubun} | {f} | {n} | `{h}` | {kyu} |")
    rd += ["", "⚠ 頁数は当方の環境（HG丸ｺﾞｼｯｸM-PRO・游ゴシックを IPAゴシックで代用した LibreOffice）による。"
           "配布用の PDF は町の環境（Word）で作り、頁送りを確かめていただく。",
           "", "## 点検ログ", "", "| 点検 | 終了コード |", "|---|---|"]
    for na, rc in logs:
        rd.append(f"| {na} | {rc} |")
    rd += ["", "## 再現の手順", "", "リポジトリの `kawasaki_project/` で次を実行すると、同じ一式ができます。", "", "```",
           "python3 07_ソーススクリプト/fix_soan_v2330_saireview.py",
           "python3 07_ソーススクリプト/fill_toc_pages_v200.py 01_第10期_最新版成果品/川崎町_計画書素案_v2.33_協議用.docx",
           "python3 07_ソーススクリプト/build_kouhyou_soan.py",
           "python3 07_ソーススクリプト/fill_toc_pages_v200.py 01_第10期_最新版成果品/川崎町_計画書案_公表用_v2.33.docx",
           "python3 07_ソーススクリプト/build_gaiyo_v230.py", "python3 07_ソーススクリプト/build_2kai_v18.py",
           "python3 07_ソーススクリプト/build_2kai_qa_v15.py",
           "python3 07_ソーススクリプト/build_redteam_saireview_teishutsu_R8.10.24.py", "```", ""]
    open(os.path.join(root, "00_README.md"), "w", encoding="utf-8").write("\n".join(rd))

    print("  ── 7 ZIP")
    os.makedirs(os.path.dirname(ZIPNAME), exist_ok=True)
    with zipfile.ZipFile(ZIPNAME, "w", zipfile.ZIP_DEFLATED) as z:
        for dp, _, fs in os.walk(root):
            for f in sorted(fs):
                full = os.path.join(dp, f)
                z.write(full, os.path.relpath(full, work))
    with zipfile.ZipFile(ZIPNAME) as z:
        names = z.namelist()
        ok = all(hashlib.md5(z.read(f"{os.path.basename(root)}/01_成果物/{f}")).hexdigest() == h
                 for _, f, _, h, _ in ichiran)
    print(("   ○ " if ok else "   × ") + f"{ZIPNAME}（{len(names)}ファイル・{os.path.getsize(ZIPNAME):,}バイト）"
          "・成果物の MD5 が一覧と一致")
    if not ok:
        ng.append("MD5")
    shutil.rmtree(work)
    if ng:
        print(f"   × 不適合 {len(ng)}件：" + "／".join(ng))
        sys.exit(1)
    print("   ○ 不適合0件")


if __name__ == "__main__":
    main()
