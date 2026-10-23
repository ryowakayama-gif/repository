# -*- coding: utf-8 -*-
"""第2回策定委員会 想定問答集 v10 → v11（見える化システムの事業所分析の反映）
（令和8年10月23日）

  python3 07_ソーススクリプト/build_2kai_qa_v11.py

協議用素案 Ver.2.28・委員会資料 v14 と同じ内容に揃える。
  (1) QB-3：いずれの資料にもないもの 4件→3件。事業所分析で把握した旨を加える
  (2) 入所率の問：延べ利用日数による稼働率（事業所分析）を加え、
      「床単位で休止か空床か」の記述を、地域密着型の空きの把握に改める
  (3) 頁数の問（QB-40・QB-56）：118頁 → 117頁
  (4) 1問を加える（全67問 → 全68問。カテゴリB）
      QB-68 増やした地域密着型特養の定員は埋まっているのか
数値は素案 Ver.2.28 の表から読む。
"""
import collections
import copy
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v10.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v11.docx"
SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.28_協議用.docx"
SHUTTEN = "地域包括ケア「見える化」システム（事業所分析・各年9月）"
HEN = 117
SU = 68


def soan_hyo():
    d = docx.Document(SOAN)
    for t in d.tables:
        if t.rows[1].cells[0].text.startswith("介護老人福祉施設（定員30人）の稼働率"):
            return [[c.text for c in r.cells] for r in t.rows]
    sys.exit("素案の表が見つからない")


def txt(el):
    return "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()


def cat_of(s):
    m = re.match(r"^カテゴリ ([A-H])　", s)
    return "カテゴリ " + m.group(1) if m else None


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会 想定問答集 v10 → v11（見える化システムの事業所分析）")
    print()
    S = soan_hyo()
    g = {r[0][:6]: r for r in S}
    tok = S[1][3]
    rok = S[2][3]
    mit = [r for r in S if r[0].startswith("地域密着型")][0]
    tan = [r for r in S if r[0].startswith("短期入所")][0]
    kai = [r for r in S if r[0].startswith("施設・居住系の介護職員")][0]
    aki = round(30 * (1 - float(mit[3].rstrip("％")) / 100))

    NAOSU = [
        ("A．介護サービス情報公表システムにより把握しました。町内14の",
         "⚠ いずれの資料にもないものは4件（入所前の居所・経営の状況・町への要望・短期入所の稼働）です。",
         f"さらに、延べ利用日数による稼働率・短期入所の利用・職員数の推移を、{SHUTTEN}で"
         f"令和3年から令和7年まで把握しました。⚠ いずれの資料にもないものは3件"
         f"（入所前の居所・経営の状況・町への要望）です。"),
        ("A．⚠ 入所率は「記入日前月の在所者数÷定員」であり、",
         "⚠ 施設単位の休止は県の調査で分かりますが、床単位で休止なのか空床なのかの別は、"
         "いずれの資料にもありません。",
         f"延べ利用日数でみた稼働率（{SHUTTEN}）は、令和7年で介護老人福祉施設{tok}・"
         f"介護老人保健施設{rok}、地域密着型介護老人福祉施設{mit[3]}です。⚠ 地域密着型は"
         f"令和5年に定員を30人に増やした後、約{aki}人分が埋まっていません（QB-68）。"),
        ("A．委託仕様書は「100頁程度」としています。現在の素案は本文118頁",
         "本文118頁・資料編（見込13頁）の計131頁", f"本文{HEN}頁・資料編（見込13頁）の計{HEN + 13}頁"),
        ("【根拠データ】 第2回資料 ８（計画書の頁数の見通しの表）。素案 本文118頁",
         "素案 本文118頁", f"素案 本文{HEN}頁"),
        ("QB-56　【住民代表】　計画書が118頁では", "118頁", f"{HEN}頁"),
        ("A．超えます。現在の素案は本文118頁で", "本文118頁で、資料編（見込13頁）を加えると計131頁",
         f"本文{HEN}頁で、資料編（見込13頁）を加えると計{HEN + 13}頁"),
        ("【根拠データ】 第2回資料 ８（頁数の見通しの表）。計画素案 本文118頁", "118頁", f"{HEN}頁"),
    ]
    ADD = ("カテゴリ B",
           "QB-68　【介護事業者】　増やした地域密着型特養の定員は埋まっているのか",
           f"A．埋まっていません。{SHUTTEN}によると、町内の地域密着型介護老人福祉施設"
           f"（2施設）は令和5年に定員を20人から30人に増やしましたが、延べ利用日数でみた"
           f"稼働率は{mit[1]}（令和3年）から{mit[3]}（令和7年）に下がっており、約{aki}人分の"
           f"受入れの余地があります。宮城県内の同じサービスの中央値は{mit[4]}です。⚠ これが"
           f"職員不足による受入れの制限によるものかは確かめていません。施設・居住系の介護職員"
           f"（常勤換算）は{kai[1]}から{kai[3]}に減っており、短期入所の利用者も月{tan[1]}から"
           f"{tan[3]}へほぼ半減しています。空きの理由を確かめたうえで、第10期の見込量と"
           f"必要利用定員総数（月27.2人）を据え置くか、町内の空き床の活用を見込量に反映するかを、"
           f"第3回でご協議いただきます。",
           f"【根拠データ】 第2回資料 ３（5年間の推移の表）。{SHUTTEN}。稼働率＝延べ利用者数÷"
           f"（定員×30日）（厚生労働省 地域分析の手引き 令和8年9月改訂版）",
           "【関連資料】 計画素案 第2章2-7・第9章9-1、確認事項No.159・No.50")

    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body
    from fix_soan_v111 import set_el

    n = 0
    for atama, kyu, atarashii in NAOSU:
        hit = [el for el in body.iterchildren()
               if el.tag == qn("w:p") and txt(el).startswith(atama)]
        if len(hit) != 1 or kyu not in txt(hit[0]):
            judge(False, f"『{atama[:24]}』{len(hit)}件・旧の語 {kyu[:16]}")
            continue
        set_el(hit[0], txt(hit[0]).replace(kyu, atarashii))
        n += 1
    judge(n == len(NAOSU), f"{len(NAOSU)}か所を改めた（実際 {n}）")

    kids = list(body.iterchildren())
    tpl_q = tpl_a = None
    for i, el in enumerate(kids):
        if el.tag == qn("w:p") and txt(el).startswith("QB-1　"):
            tpl_q, tpl_a = copy.deepcopy(el), copy.deepcopy(kids[i + 1])
            break
    cur = last = None
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        c = cat_of(txt(el))
        if c:
            cur = c
        if cur == ADD[0] and txt(el).startswith("【関連資料】"):
            last = el
    for text, tpl in ((ADD[4], tpl_a), (ADD[3], tpl_a), (ADD[2], tpl_a), (ADD[1], tpl_q)):
        el = copy.deepcopy(tpl)
        set_el(el, text)
        last.addnext(el)

    cnt = collections.Counter()
    cur = None
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = txt(el)
        c = cat_of(s)
        if c:
            cur = c
        if re.match(r"^QB-\d+　", s):
            cnt[cur] += 1
    total = sum(cnt.values())
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = txt(el)
        m = re.match(r"^(カテゴリ [A-H]　.+?)（\d+問）$", s)
        if m:
            set_el(el, f"{m.group(1)}（{cnt[cat_of(s)]}問）")
        elif re.search(r"全\d+問", s):
            set_el(el, re.sub(r"全\d+問", f"全{total}問", s))
    from fix_soan_v111 import set_cell
    for t in doc.tables:
        if [c.text.strip() for c in t.rows[0].cells][:3] == ["カテゴリ", "テーマ", "問数"]:
            for r in t.rows[1:]:
                k = "カテゴリ " + r.cells[0].text.strip()
                if k in cnt:
                    set_cell(r.cells[2], str(cnt[k]))
    doc.save(DST)

    d2 = docx.Document(DST)
    zen = "\n".join(p.text for p in d2.paragraphs) + "\n" + "\n".join(
        c.text for t in d2.tables for r in t.rows for c in r.cells)
    print("  ── 検査1　件数")
    qs = [int(x) for x in re.findall(r"QB-(\d+)　", zen)]
    judge(sorted(qs) == list(range(1, SU + 1)) and total == SU, f"全{SU}問・欠番なし")
    judge(f"全{SU}問" in zen, f"「全{SU}問」と宣言")
    print("  ── 検査2　素案・資料と同じ値")
    for v in (mit[3], tok, rok, kai[1], kai[3]):
        judge(v in zen, f"「{v}」")
    judge("短期入所の稼働" not in zen and "118頁" not in zen, "古い記述（短期入所の稼働・118頁）0件")
    print("  ── 検査3　事業所名・他団体名")
    for na in ("釜房", "アルパイン", "鶴寿会", "グループホームさわやか", "グループホームふれあい",
               "北塩原", "金ヶ崎", "大雪", "川崎市"):
        judge(na not in zen, f"「{na}」0件")
    print()
    print(f"  {DST}　全{total}問・" + "／".join(f"{k[-1]}{v}" for k, v in sorted(cnt.items())))
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜3　不適合0件")


if __name__ == "__main__":
    main()
