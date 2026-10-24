# -*- coding: utf-8 -*-
"""業務工程管理表の更新（素案の版の表記の整理・見え方の点検、素案Ver.2.26・問答集v10）
（令和8年10月23日）

  python3 07_ソーススクリプト/update_wbs_R8.10.23.py

  ・素案の紙面から版の表記を除いた（表紙は「計画書素案（令和8年10月時点）」）
  ・行送りの lineRule・表の列幅・表の中の長い文の揃えを直した
  ・書体と見出しの形を第9期計画に揃えるかを確認事項 No.187 として起こす

⚠ 列8（状態）は新しい確認事項にだけ「確認待ち」を入れる。既存の行は動かさない。
"""
import copy
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
SOAN_OLD = "川崎町_計画書素案_v2.25_軽費老人ホーム確認版.docx"
SOAN = "川崎町_計画書素案_v2.26_版表記整理版.docx"
QA_OLD = "同_想定問答集_R8.11_v9.docx"
QA = "同_想定問答集_R8.11_v10.docx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月23日 更新（版の表記・見え方の点検）】"
NO = "187"


def iro(ws, r, n):
    for c in range(1, n + 1):
        ws.cell(r, c).font = FONT
        ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")


def tsuika(ws, atai):
    moto = ws.max_row
    r = moto + 1
    for c in range(1, len(atai) + 1):
        ws.cell(r, c)._style = copy.copy(ws.cell(moto, c)._style)
        ws.cell(r, c).value = atai[c - 1]
    ws.row_dimensions[r].height = ws.row_dimensions[moto].height
    return r


def main():
    wb = openpyxl.load_workbook(WBS)

    ws = wb["03_確認事項一覧"]
    if not any(str(ws.cell(x, 1).value or "").strip() == NO
               for x in range(5, ws.max_row + 1)):
        r = tsuika(ws, [
            NO, "(8)", "R8.10", "計画書の書体と見出しの形を第9期計画に揃えるか",
            "第9期計画は本文を HG丸ゴシックM-PRO 11pt（行送り約19.6pt）とし、"
            "節を「第１節」、表・図の題を「■」（番号なし）、出典を「資料：」と"
            "しています。計画素案は游ゴシック 10pt（行送り16pt）、節を「1-1」、"
            "図に「図3-1」等の番号を付けています。①書体と大きさ（Ａ現状116頁／"
            "Ｂ書体だけ揃える116頁前後／Ｃ大きさ・行送りも揃える約126頁）、"
            "②見出し・図表の題の形を揃える時期、をご判断ください。"
            "詳細は 川崎町_計画書の見え方の点検_R8.10.23.md。",
            "計画書（成果品4）の紙面・頁数", "発注者", "確認待ち", "R8.11",
            "既定値＝①Ｂ（書体だけ HG丸ゴシックM-PRO に揃え、大きさは10pt）、"
            "②公表版（計画書の案）の段階で「第１節」「■」「資料：」に揃える。"
            "協議の間は現在の形のままとする。"])
        iro(ws, r, 10)

    ws = wb["04_成果品管理"]
    n_seika = 0
    for x in range(5, ws.max_row + 1):
        no = str(ws.cell(x, 1).value or "").strip()
        na = str(ws.cell(x, 3).value or "")
        if no == "4" and SOAN_OLD in na:
            ws.cell(x, 3).value = na.replace(SOAN_OLD, SOAN)
            ws.cell(x, 8).value = (
                f"{HIZUKE}素案を Ver.2.26 に上げた（全116頁。日本語の書体を"
                f"游ゴシックに字幅の近い書体に寄せて実測。これまでの118頁は置き換え"
                f"前の値）。紙面から版の表記を除き、表紙を「計画書素案（令和8年"
                f"10月時点）」とした。⚠ 行送り5,621段落に lineRule がなく町の Word で"
                f"頁数が変わり得たため明示した。⚠ 129表の列幅が実寸でなく、制度改正の"
                f"表で文字が縦に折れていたため実寸で固定した。表の中の長い文441段落を"
                f"左揃えに。文言・数値は変えていない。　"
                + str(ws.cell(x, 8).value or ""))
            iro(ws, x, 9)
            n_seika += 1
        elif no == "6" and QA_OLD in na:
            ws.cell(x, 3).value = na.replace(QA_OLD, QA)
            ws.cell(x, 8).value = (
                f"{HIZUKE}想定問答集を v10 に上げた（素案の版の参照48か所を外した）。　"
                + str(ws.cell(x, 8).value or ""))
            iro(ws, x, 9)
            n_seika += 1

    ws2 = wb["08_作業順位"]
    if not any(str(ws2.cell(x, 1).value or "").strip() == "41"
               for x in range(5, ws2.max_row + 1)):
        tsuika(ws2, [
            "41", "素案の版の表記の整理と、計画の見え方の点検（第9期計画との対比）",
            "R8.10.23",
            "ご指示による。同じ型の素案を作っている他の案件の紙面の点検の観点を借り、"
            "第9期計画の PDF の書体・紙面を実測して突き合わせた。",
            f"No.{NO}（新規）", "計画素案 Ver.2.26／想定問答集 v10／見え方の点検（md）",
            "1日", "完了",
            "版の表記8か所とヘッダーを除いた。⚠ **行送りの lineRule なし5,621段落・"
            "列幅が実寸でない129表**を直した（町の Word で頁数と目次がずれ得た／"
            "制度改正の表で文字が縦に折れていた）。表の中の長い文441段落を左揃えに。"
            "第9期は HG丸ゴシックM-PRO 11pt・「第１節」「■」「資料：」。"
            "揃えるかを No.187 でご判断いただく。"])
    wb.save(WBS)

    wb2 = openpyxl.load_workbook(WBS, data_only=True)
    ws3 = wb2["03_確認事項一覧"]
    nos = sorted(int(s) for s in (str(ws3.cell(x, 1).value or "").strip()
                                  for x in range(5, ws3.max_row + 1))
                 if s.isdigit())
    s = "\n".join(str(v) for sh in wb2 for row in sh.iter_rows(values_only=True)
                  for v in row if v is not None)
    jun = [str(wb2["08_作業順位"].cell(x, 1).value or "").strip()
           for x in range(5, wb2["08_作業順位"].max_row + 1)]
    ok = [
        (nos == list(range(1, max(nos) + 1)) and max(nos) == 187,
         f"確認事項 欠番0・No.187 まで（{len(nos)}件）"),
        (n_seika in (0, 2), f"成果品4・6 を更新（{n_seika}件）"),
        (jun.count("41") == 1, "作業順位41 が1行"),
        (SOAN in s and "想定問答集_R8.11_v10" in s, "素案 Ver.2.26・問答集 v10 が載っている"),
        (not any(w in s for w in ("蔵王グリーンホーム", "鶴寿会", "北塩原", "金ヶ崎",
                                  "大雪", "川崎市")), "施設名・法人名・他団体名なし"),
    ]
    print("業務工程管理表の更新（版の表記・見え方の点検）")
    for g, mi in ok:
        print(("   ○ " if g else "   × ") + mi)
    if not all(g for g, _ in ok):
        sys.exit(1)


if __name__ == "__main__":
    main()
