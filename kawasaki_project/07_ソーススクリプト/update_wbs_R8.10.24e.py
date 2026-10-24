# -*- coding: utf-8 -*-
"""業務工程管理表の更新（Red Team 再レビュー第4回（令和8年10月24日）への対応）

  python3 07_ソーススクリプト/update_wbs_R8.10.24e.py

  00_管理方法　　　基準日の記載を第4回対応後に
  03_確認事項一覧　No.85（給付費の集計の基準）・No.49（1月4日開始の工程）・No.83（頁の数え方）に追記
  04_成果品管理　　稿を 素案 Ver.2.36／公表用・概要版 v2.36／資料 v21／問答集 v17 に
  08_作業順位　　　53（完了）を加える
再実行しても二重に書かない。
"""
import collections
import copy
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月24日 更新⑤（Red Team 再レビュー第4回への対応）】"
TAIOU = "川崎町_RedTeam再レビュー第4回対応表_R8.10.24.md"
ZIP = "11_提出一式/川崎町_RedTeam再レビュー第4回対応_成果物とエビデンス_R8.10.24.zip"
OKIKAE = [("川崎町_計画書素案_v2.35_協議用.docx", "川崎町_計画書素案_v2.36_協議用.docx"),
          ("川崎町_計画書案_公表用_v2.35.docx", "川崎町_計画書案_公表用_v2.36.docx"),
          ("川崎町_計画書案_概要版_v2.35.docx", "川崎町_計画書案_概要版_v2.36.docx"),
          ("川崎町_第2回策定委員会資料_R8.11_v20.docx", "川崎町_第2回策定委員会資料_R8.11_v21.docx"),
          ("同_想定問答集_R8.11_v16.docx", "同_想定問答集_R8.11_v17.docx")]
TSUIKI = {
    "85": (f"{HIZUKE}F4-01。9-2 の表の令和7年度で、標準給付費見込額（1,084,587千円＝年報 様式3 の支払額。審査支払手数料を"
           "含まない）と内訳5項目の和（1,084,826千円。総給付費は様式2 の給付実績）が239千円違っていた。令和7年度の欄を5項目の和に"
           "改め、様式3 との違いを注に書いた。令和6年度→令和7年度の比較も基準が違っていた（1,093,253千円は手数料を含む）ため、"
           "様式3 の支払額で 1,092,621千円→1,084,587千円・0.7％減に改めた。保険料（令和9〜11年度）への波及なし。　"),
    "49": (f"{HIZUKE}F4-03。パブリックコメントを令和9年1月4日開始とすると後続が6日遅れ、案Ⅰの納品は3月16日（火）で納期を1日"
           "超える（旧の「3月15日ちょうど」は誤り）。1月4日開始なら第3回から入稿まで又は印刷製本を1日以上短縮するか、案Ⅱ"
           "（印刷2週間・3月9日）。資料 v21・問答集 v17・逆算工程表を改めた。　"),
    "83": (f"{HIZUKE}F4-06。素案は全118頁＝表紙・ご挨拶・目次5頁＋本文113頁、資料編（見込）13頁を含め計131頁。別冊に移した案は"
           "全106頁（本文101頁）・計119頁。「100頁程度」との差は本文のみで＋13頁（別冊後＋1頁）、冊子全体で＋31頁（＋19頁）。"
           "数える範囲を町と書面で確かめたい。　"),
}


def iro(c):
    c.font = FONT
    c.alignment = Alignment(wrap_text=True, vertical="top")


def mae_ni(c, s):
    if HIZUKE in str(c.value or ""):
        return 0
    c.value = s + str(c.value or "")
    iro(c)
    return 1


def main():
    wb = openpyxl.load_workbook(WBS)
    kekka = []
    ws3 = wb["03_確認事項一覧"]
    k = 0
    for y in range(5, ws3.max_row + 1):
        no = str(ws3.cell(y, 1).value or "").strip()
        if no in TSUIKI:
            k += mae_ni(ws3.cell(y, 10), TSUIKI[no])
    kekka.append((k in (0, len(TSUIKI)), f"確認事項への追記 {k}件"))
    k3 = collections.Counter(str(ws3.cell(y, 8).value).strip() for y in range(5, ws3.max_row + 1)
                             if str(ws3.cell(y, 1).value or "").strip().isdigit())

    ws0 = wb["00_管理方法"]
    for r in ws0.iter_rows():
        for c in r:
            v = str(c.value or "")
            if v.startswith("令和8年10月24日　本表の記載は同日時点のもの。Red Team 再レビュー第3回の影響範囲を5冊で走査し"):
                c.value = ("令和8年10月24日　本表の記載は同日時点のもの。Red Team 再レビュー第4回を受け、協議用素案 Ver.2.36・"
                           "公表用の計画書（案）v2.36・概要版 v2.36・第2回策定委員会資料 v21・想定問答集 v17（全69問）に是正した。"
                           "第4回の指摘9件と影響範囲の確認で見つけた1件のうち是正済7・町のご判断2・町のご確認待ち1（ほか横断1件是正）。"
                           "点検ログ14件はすべて終了コード0。"
                           f"確認事項は{sum(k3.values())}件のうち{k3['完了']}件が完了。一式は {ZIP}。\n"
                           "（前の記載）" + v)
                iro(c)
                kekka.append((True, "基準日の記載を第4回対応後に"))

    ws4 = wb["04_成果品管理"]
    n = 0
    for y in range(5, ws4.max_row + 1):
        na = str(ws4.cell(y, 3).value or "")
        nb = na
        for a, b in OKIKAE:
            nb = nb.replace(a, b)
        if nb != na:
            ws4.cell(y, 3).value = nb
            mae_ni(ws4.cell(y, 8), f"{HIZUKE}Red Team 再レビュー第4回の指摘を是正した版に上げた（{TAIOU}）。　")
            n += 1
    kekka.append((n in (0, 3), f"成果品の稿 {n}行"))

    ws8 = wb["08_作業順位"]
    aru = {str(ws8.cell(x, 1).value or "").strip() for x in range(5, ws8.max_row + 1)}
    if "53" not in aru:
        moto = ws8.max_row
        r = moto + 1
        for c in range(1, 11):
            ws8.cell(r, c)._style = copy.copy(ws8.cell(moto, c)._style)
        ws8.row_dimensions[r].height = ws8.row_dimensions[moto].height
        for c, v in enumerate([
                "53", "Red Team 再レビュー第4回（令和8年10月24日）への対応と成果物・エビデンスの一式", "R8.10.24",
                "ご指示による。第4回の指摘（F4-01〜F4-09）を最新の版で確かめ、影響範囲を5冊と工程表で確かめて是正した",
                "No.49・83・85", "素案 Ver.2.36／公表用・概要版 v2.36／資料 v21／問答集 v17／逆算工程表／ZIP",
                "1日", "完了",
                "標準給付費の239千円差は年報の様式2（給付実績）と様式3（支払額・手数料を含まない）の混在。令和6年度との比較も同じ"
                "原因で基準が違っていた（0.8％→0.7％）。1月4日開始の納品は3月16日（1日超過）に訂正。頁は全体と本文を分けた。"
                "原票（年報・算定ブック）を ZIP に同梱"], 1):
            ws8.cell(r, c).value = v
    aru = {str(ws8.cell(x, 1).value or "").strip() for x in range(5, ws8.max_row + 1)}
    kekka.append(("53" in aru, "作業順位 53"))

    wb.save(WBS)
    print("業務工程管理表の更新（Red Team 再レビュー第4回への対応）")
    ng = 0
    for ok, mi in kekka:
        print(("   ○ " if ok else "   × ") + mi)
        ng += not ok
    if ng:
        sys.exit(1)


if __name__ == "__main__":
    main()
