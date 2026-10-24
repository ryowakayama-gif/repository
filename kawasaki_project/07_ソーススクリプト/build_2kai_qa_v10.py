# -*- coding: utf-8 -*-
"""第2回策定委員会 想定問答集 v9 → v10（素案の版の表記を外す）
（令和8年10月23日）

  python3 07_ソーススクリプト/build_2kai_qa_v10.py

v9 からの変更の理由
  計画素案の紙面から版の表記を除いた（fix_soan_v2260_hyoki.py）。
  素案は表紙で「計画書素案（令和8年10月時点）」と示し、版の番号はファイル名で
  管理する。問答集の【関連資料】の「計画素案Ver.2.25」も「計画素案」とする。
  問の数は変えない（全67問）。
"""
import re
import shutil
import sys
import os

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v9.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v10.docx"
RE_VER = re.compile(r"(計画)?素案\s?Ver\.\s?[0-9.]+")


def main():
    from fix_soan_v111 import set_el
    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    n = 0
    for el in doc.element.body.iter(qn("w:p")):
        s = "".join(x.text or "" for x in el.iter(qn("w:t")))
        if RE_VER.search(s):
            set_el(el, RE_VER.sub("計画素案", s).replace("計画計画素案", "計画素案"))
            n += 1
    doc.save(DST)

    d2 = docx.Document(DST)
    zen = "\n".join(p.text for p in d2.paragraphs) + "\n" + "\n".join(
        c.text for t in d2.tables for r in t.rows for c in r.cells)
    qs = [int(x) for x in re.findall(r"QB-(\d+)　", zen)]
    ok = [
        (n >= 40, f"素案の版の参照を外した（{n}か所）"),
        ("Ver." not in zen, "「Ver.」が0件"),
        ("計画計画" not in zen, "語の重なりがない"),
        (sorted(qs) == list(range(1, 68)), "全67問・欠番なし"),
        (not any(w in zen for w in ("蔵王グリーンホーム", "鶴寿会", "北塩原",
                                    "金ヶ崎", "川崎市", "大雪")),
         "施設名・法人名・他団体名が入っていない"),
    ]
    print("第2回策定委員会 想定問答集 v9 → v10（素案の版の表記を外す）")
    for good, mi in ok:
        print(("   ○ " if good else "   × ") + mi)
    print(f"  {DST}")
    if not all(g for g, _ in ok):
        sys.exit(1)


if __name__ == "__main__":
    main()
