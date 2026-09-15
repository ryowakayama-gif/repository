# -*- coding: utf-8 -*-
"""新推計パターン設定値一式の⑦シートに、所得段階別第1号被保険者数（13段階）を投入する。

令和7年度末の実績構成比（年報・様式1 所得段階別・列20）を、
令和9〜11年度の第1号被保険者数見込みに適用して段階別人数を割り付ける。
"""
import openpyxl
from openpyxl.styles import Font, PatternFill

BASE = "/home/user/repository/kawasaki_project"
XLSX = f"{BASE}/05_試算・管理シート/川崎町_新推計パターン_設定値一式_R8.9.11.xlsx"

RITSU = [0.455, 0.685, 0.690, 0.90, 1.00, 1.20, 1.30, 1.50, 1.70,
         1.90, 2.10, 2.30, 2.40]
R7 = [383, 298, 305, 343, 679, 410, 427, 232, 85, 26, 17, 7, 28]
TOTALS = {"R9": 3231, "R10": 3203, "R11": 3177}
SRC = sum(R7)


def allocate(total):
    """最大剰余法で構成比どおりに割り付ける。"""
    raw = [n / SRC * total for n in R7]
    base = [int(x) for x in raw]
    rest = total - sum(base)
    order = sorted(range(13), key=lambda i: raw[i] - base[i], reverse=True)
    for i in order[:rest]:
        base[i] += 1
    return base


ALLOC = {y: allocate(t) for y, t in TOTALS.items()}

wb = openpyxl.load_workbook(XLSX)
ws = wb["⑦保険料額の算定"]

ws.cell(3, 1).value = ("【1】所得段階別第1号被保険者数（標準段階区分）"
                       "　★令和7年度末の実績（年報）を反映しました")
ws.cell(3, 1).font = Font(bold=True)

for i in range(13):
    r = 5 + i
    for c, y in ((2, "R9"), (3, "R10"), (4, "R11")):
        cell = ws.cell(r, c, ALLOC[y][i])
        cell.number_format = "#,##0"
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
    ws.cell(r, 5).value = (f"乗率 {RITSU[i]:.3f}　"
                           f"令和7年度末実績 {R7[i]:,}人（構成比 {R7[i] / SRC * 100:.1f}％）")

for c, y in ((2, "R9"), (3, "R10"), (4, "R11")):
    ws.cell(18, c).value = sum(ALLOC[y])

HOSEI = sum(n * r for n, r in zip(R7, RITSU)) / SRC
ws.cell(20, 1).value = ("※ 出所：介護保険事業状況報告（年報・令和7年度）"
                        "様式1「所得段階別」列20「年度末現在被保険者数」。"
                        "令和7年度末の実績構成比を各年度の第1号被保険者数見込みに割り付けています。")
ws.cell(21, 1).value = ("　 これまで第9段階で頭打ち（第10〜13段階が0人）になっていたため、"
                        "所得段階別加入割合による補正係数が 0.964047 と低く出ていました。")
ws.cell(22, 1).value = (f"　 実績を反映すると補正係数は {HOSEI:.6f} となり、"
                        "保険料基準額は月額352円下がります（取崩なしで6,822円）。")

wb.save(XLSX)

print("■ ⑦シート【1】所得段階別第1号被保険者数（更新後）")
print(f"{'段階':<8}{'乗率':>7}{'R9':>8}{'R10':>8}{'R11':>8}")
for i in range(13):
    print(f"第{i + 1}段階{'':<3}{RITSU[i]:>7.3f}"
          f"{ALLOC['R9'][i]:>8,}{ALLOC['R10'][i]:>8,}{ALLOC['R11'][i]:>8,}")
print(f"{'合計':<8}{'':>7}{sum(ALLOC['R9']):>8,}"
      f"{sum(ALLOC['R10']):>8,}{sum(ALLOC['R11']):>8,}")
print(f"\n補正係数　{HOSEI:.6f}")
