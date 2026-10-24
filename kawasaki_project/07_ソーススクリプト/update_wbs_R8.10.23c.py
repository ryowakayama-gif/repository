# -*- coding: utf-8 -*-
"""業務工程管理表の更新（事業所別5か年推移の受領と確認）
（令和8年10月23日・同日3度目）

  python3 07_ソーススクリプト/update_wbs_R8.10.23c.py

⚠ 列8（状態）は動かさない。⚠ 事業所名・法人名は書かない。
"""
import copy
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月23日 更新③（事業所別5か年推移）】"
MEMO = "川崎町_事業所別5か年推移の確認結果_R8.10.23.md"

TSUIKI = {
    "159": (f"{HIZUKE}令和8年10月23日受領の事業所別5か年推移（令和3〜7年）で、"
            f"施設・居住系6か所の延べ利用日数による稼働率が分かりました"
            f"（介護老人福祉施設97.4％・介護老人保健施設90.7％ほか）。"
            f"⚠ 令和5年に増床した地域密着型介護老人福祉施設（定員10→20）の"
            f"稼働率が令和7年で69.4％（県内の中央値93.6％）です。"
            f"短期入所の利用者は2か所とも5年でほぼ半減しています。"
            f"⚠ データの出所が記載されていないため、作成主体と時点のご確認を"
            f"お願いします［要確認］。詳細は {MEMO}。　"),
    "50": (f"{HIZUKE}事業所別5か年推移で、介護職員（常勤換算）が介護老人保健施設"
           f"30.0→25.5人（▲15％）、介護老人福祉施設21.2→18.3人（令和5年は14.0人）"
           f"でした（県内の中央値は▲2％前後）。第1回でご報告のあった1ユニット分の"
           f"人員削減と時期が近いものの、同じものかは確かめていません。　"),
}


def iro(ws, r, n):
    for c in range(1, n + 1):
        ws.cell(r, c).font = FONT
        ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")


def main():
    wb = openpyxl.load_workbook(WBS)
    ws = wb["03_確認事項一覧"]
    n = 0
    for x in range(5, ws.max_row + 1):
        no = str(ws.cell(x, 1).value or "").strip()
        if no in TSUIKI and HIZUKE not in str(ws.cell(x, 10).value or ""):
            ws.cell(x, 10).value = TSUIKI[no] + str(ws.cell(x, 10).value or "")
            iro(ws, x, 10)
            n += 1
    ws2 = wb["08_作業順位"]
    if not any(str(ws2.cell(x, 1).value or "").strip() == "43"
               for x in range(5, ws2.max_row + 1)):
        moto = ws2.max_row
        r = moto + 1
        for c in range(1, 11):
            ws2.cell(r, c)._style = copy.copy(ws2.cell(moto, c)._style)
        ws2.row_dimensions[r].height = ws2.row_dimensions[moto].height
        for c, v in enumerate([
                "43", "事業所別5か年推移（利用者数・稼働・回転・従事者数）の確認",
                "R8.10.23", "受領による（全国47都道府県分。原本を09へ退避）。",
                "No.159（稼働率・短期入所）／No.50（介護人材）",
                "事業所別5か年推移の整理（xlsx・5シート）／川崎町の行の抜き出し（csv）／"
                "確認結果（md）", "半日", "完了",
                "町内12の指定の令和3〜7年の値を抜き出し、既知の定員・常勤換算と一致を確認。"
                "⚠ 集計の欄は数式で値が保存されておらず計算し直した。⚠ 原表の「定員充足率」"
                "は延べ÷定員（百分率でない）。地域密着型特養の増床分の稼働69.4％、"
                "短期入所の半減、老健・特養の介護職員の減少を把握。⚠ 出所の記載なし"
                "［要確認］。素案への反映は出所の確認後を推奨。"], 1):
            ws2.cell(r, c).value = v
    wb.save(WBS)

    wb2 = openpyxl.load_workbook(WBS, data_only=True)
    s = "\n".join(str(v) for sh in wb2 for row in sh.iter_rows(values_only=True)
                  for v in row if v is not None)
    ng = [w for w in ("釜房", "アルパイン", "鶴寿会", "グループホームさわやか", "グループホームふれあい",
                      "北塩原", "金ヶ崎", "大雪", "川崎市") if w in s]
    jun = [str(wb2["08_作業順位"].cell(x, 1).value or "").strip()
           for x in range(5, wb2["08_作業順位"].max_row + 1)]
    ok = n in (0, 2) and jun.count("43") == 1 and not ng
    print("業務工程管理表の更新（事業所別5か年推移）")
    print(("   ○ " if ok else "   × ") + f"確認事項 {n}件・作業順位43・名称の混入 {ng}")
    if not ok:
        sys.exit(1)


if __name__ == "__main__":
    main()
