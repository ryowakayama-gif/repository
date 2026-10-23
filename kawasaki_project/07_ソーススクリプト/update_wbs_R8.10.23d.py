# -*- coding: utf-8 -*-
"""業務工程管理表の更新（見える化システムの出力のお願い・事業所分析の出所の確定）
（令和8年10月23日・同日4度目）

  python3 07_ソーススクリプト/update_wbs_R8.10.23d.py

  ・06_資料提供依頼一覧に、見える化システムからご出力をお願いする資料11件
    （No.40〜50）を加える
  ・⚠ No.39 が注記の後ろ（53行目）に足されており、件数の式（I5:I40）と表題
    （38件）も古いままだった。表の続きに並べ直し、式と表題を直す
  ・確認事項 No.159・No.50 に、出所の確定（見える化システムの事業所分析）と、
    9月の30日による計算し直しを追記する

⚠ 列8（状態）は動かさない（新しい依頼の「状態」は「未着手」とする）。
"""
import copy
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
MEMO = "川崎町_見える化システムからご出力をお願いする資料_R8.10.23.md"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月23日 更新④（出所の確定）】"

# No, 領域, 資料, 内容, 確定する事項, 入手先, 希望時期, 優先度, 状態, 備考
SHIN = [
    (40, "F", "見える化 将来推計の推計結果概要（総括表）の再出力",
     "予定収納率・調整交付金の見込交付割合・交付金の交付見込額・所得段階別被保険者数を改めた後の総括表",
     "素案9-3の保険料。確認事項No.39・75・78・174", "発注者（見える化システム）",
     "R8.11上旬", "高", "未着手", f"手引き80頁。{MEMO} No.2"),
    (41, "F", "見える化「介護保険事業状況報告の設定」画面の写し",
     "令和8年度の実績見込みの算定に用いる月（令和8年5月から設定した月まで）",
     "令和8年度の実績見込み。確認事項No.79・80", "発注者（見える化システム）",
     "R8.11上旬", "高", "未着手", f"手引き77〜78頁。{MEMO} No.3"),
    (42, "B", "認知症高齢者の日常生活自立度別の要介護認定者の状況",
     "自立度Ⅱ以上の人数（見える化の要介護者の状況の指標又は要介護認定データ）",
     "素案6-1の［要確認］。確認事項No.157", "発注者（見える化システム／認定データ）",
     "R8.11上旬", "高", "未着手", f"手引き22頁。{MEMO} No.4"),
    (43, "F", "要介護認定適正化事業の業務分析データ",
     "要介護認定の適正化の実施状況（国保連合会介護給付適正化システム。自治体ID要）",
     "給付適正化3事業。確認事項No.127・59・122", "発注者", "R8.11上旬", "高", "未着手",
     f"手引き84頁。{MEMO} No.5"),
    (44, "B", "見える化 調整済み認定率", "性・年齢調整済み認定率（令和5年9月末・令和8年9月末）",
     "第9期KPI・別表 四。確認事項No.35・155", "発注者（見える化システム）", "R8.11", "中",
     "未着手", f"手引き9〜20頁。{MEMO} No.6"),
    (45, "D", "見える化 介護サービス自給率", "町内の事業所による給付の割合（時系列・地域比較）",
     "町外施設の利用の大きさ。確認事項No.155", "発注者（見える化システム）", "R8.11", "中",
     "未着手", f"手引き58〜60頁。{MEMO} No.7"),
    (46, "F", "見える化 将来推計 伸びの比較グラフ", "認定率・利用率・利用回（日）数の伸び（Excel）",
     "伸びの選択。確認事項No.79", "発注者（見える化システム）", "R8.11", "中", "未着手",
     f"手引き77頁。{MEMO} No.8"),
    (47, "A", "見える化 レーダーチャート", "需要・供給・医療介護連携の18指標の全国偏差値",
     "素案第2章・第3章の本町の位置", "発注者（見える化システム）", "R8.11", "中", "未着手",
     f"手引き48〜57頁。{MEMO} No.9"),
    (48, "A", "見える化 地域分析・検討シート（比較地域を設定したもの）",
     "近隣町・同規模保険者との比較", "素案第2章の比較", "発注者（見える化システム）",
     "R8.11", "中", "未着手", f"手引き85〜87頁。{MEMO} No.10"),
    (49, "A", "見える化 人口メッシュ（地図）", "7地区の高齢者の分布と事業所からの距離",
     "訪問・通所が困難な地区。素案1-6", "発注者（見える化システム）", "R8.11", "中",
     "未着手", f"手引き62〜66頁。{MEMO} No.11"),
    (50, "F", "見える化 地域支援事業費の入力の画面の写し",
     "入力されている事業費と、その額の作り方（対象経費支出予定額か決算額か）",
     "素案9-2・9-3の地域支援事業費。確認事項No.77・168", "発注者（見える化システム）",
     "R8.11", "低", "未着手", f"手引き79頁。{MEMO} No.12"),
]

TSUIKI = {
    "159": (f"{HIZUKE}受領した事業所別5か年推移の出所は地域包括ケア「見える化」システムの"
            f"事業所分析機能であるとご教示いただきました（各年9月の値。手引き66頁）。"
            f"稼働率を9月の30日で計算し直しました（地域密着型介護老人福祉施設（増床分）"
            f"70.3％・県内の中央値94.9％、介護老人福祉施設98.8％、介護老人保健施設91.9％）。"
            f"出所の［要確認］は解けました。　"),
    "50": (f"{HIZUKE}事業所別5か年推移の出所は見える化システムの事業所分析機能"
           f"（従事者数は介護サービス情報公表システムの各年9月時点）と確定しました。　"),
}


def main():
    wb = openpyxl.load_workbook(WBS)
    ws = wb["06_資料提供依頼一覧"]

    # ── 並びの確認
    last = max(x for x in range(5, ws.max_row + 1)
               if str(ws.cell(x, 1).value or "").strip().isdigit()
               and x < 44)
    no39 = [x for x in range(44, ws.max_row + 1)
            if str(ws.cell(x, 1).value or "").strip() == "39"]
    sudeni = {str(ws.cell(x, 1).value or "").strip() for x in range(5, ws.max_row + 1)}
    if "40" in sudeni:
        print("   ○ 既に No.40 以降がある（再実行）。何もしない")
        return
    if last != 42 or len(no39) != 1:
        print(f"   × 想定と並びが違う（表の最後 {last}・No.39 の行 {no39}）")
        sys.exit(1)

    # ── 下の部分（集計・注記）を写し取る
    shita = []
    for x in range(43, ws.max_row + 1):
        if x == no39[0]:
            continue
        shita.append([(ws.cell(x, c).value, copy.copy(ws.cell(x, c)._style))
                      for c in range(1, 11)])
    r39 = [ws.cell(no39[0], c).value for c in range(1, 11)]
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 43:
            ws.unmerge_cells(str(rng))
    for x in range(43, ws.max_row + 1):
        for c in range(1, 11):
            ws.cell(x, c).value = None

    # ── 表の続き：No.39 と No.40〜50
    gyo = [r39] + [list(r) for r in SHIN]
    style = [copy.copy(ws.cell(42, c)._style) for c in range(1, 11)]
    x = 43
    for g in gyo:
        for c in range(1, 11):
            ws.cell(x, c).value = g[c - 1]
            ws.cell(x, c)._style = copy.copy(style[c - 1])
        ws.row_dimensions[x].height = ws.row_dimensions[42].height
        x += 1
    owari = x - 1

    # ── 集計・注記を後ろに
    for row in shita:
        for c, (v, st) in enumerate(row, 1):
            if isinstance(v, str) and v.startswith("=COUNTIF(I5:I40"):
                v = v.replace("I5:I40", f"I5:I{owari}")
            ws.cell(x, c).value = v
            ws.cell(x, c)._style = st
        if row[0][0] == "対象外" and row[1][0] is None:
            ws.cell(x, 2).value = f'=COUNTIF(I5:I{owari},"対象外")'
        x += 1
    t = str(ws.cell(1, 1).value or "")
    ws.cell(1, 1).value = t.replace("（38件）", f"（{len(gyo) + 38}件）")

    # ── 確認事項への追記
    ws3 = wb["03_確認事項一覧"]
    n = 0
    for y in range(5, ws3.max_row + 1):
        no = str(ws3.cell(y, 1).value or "").strip()
        if no in TSUIKI and HIZUKE not in str(ws3.cell(y, 10).value or ""):
            ws3.cell(y, 10).value = TSUIKI[no] + str(ws3.cell(y, 10).value or "")
            ws3.cell(y, 10).font = FONT
            ws3.cell(y, 10).alignment = Alignment(wrap_text=True, vertical="top")
            n += 1
    wb.save(WBS)

    # ── 点検
    wb2 = openpyxl.load_workbook(WBS)
    w = wb2["06_資料提供依頼一覧"]
    nos = [int(str(w.cell(y, 1).value).strip()) for y in range(5, w.max_row + 1)
           if str(w.cell(y, 1).value or "").strip().isdigit()]
    shiki = [w.cell(y, 2).value for y in range(5, w.max_row + 1)
             if isinstance(w.cell(y, 2).value, str) and w.cell(y, 2).value.startswith("=COUNTIF")]
    ok = [
        (nos == list(range(1, 51)), f"依頼No.1〜50 が順に並ぶ（{len(nos)}件・最後 {nos[-1]}）"),
        (all(f"I5:I{owari}" in s for s in shiki) and len(shiki) >= 4,
         f"件数の式の範囲が I5:I{owari}（{len(shiki)}式）"),
        ("（50件）" in str(w.cell(1, 1).value), f"表題 {w.cell(1, 1).value}"),
        (n == 2, f"確認事項 No.159・No.50 に追記（{n}件）"),
    ]
    print("業務工程管理表の更新（見える化の出力のお願い）")
    for g, mi in ok:
        print(("   ○ " if g else "   × ") + mi)
    if not all(g for g, _ in ok):
        sys.exit(1)


if __name__ == "__main__":
    main()
