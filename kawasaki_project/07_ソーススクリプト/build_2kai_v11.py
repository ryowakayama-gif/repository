# -*- coding: utf-8 -*-
"""第2回策定委員会資料 v10 → v11（交付金の状況の再整理の反映）
（令和8年10月22日）

  python3 07_ソーススクリプト/build_2kai_v11.py

v10 からの変更の理由
  計画素案を Ver.2.24 に上げた（fix_soan_v2240_kofukin.py）。
  素案を直したら、同じ内容を載せている成果品も同じ版で直す（CLAUDE.md）。
  (1) ７の冒頭に「37点下がった内訳」の表と「3か年の推移と県内の動き」の
      表を加える。v10 は 565点→528点（▲37点）とだけ書き、内訳がなかった
      ⚠ 国が配点を下げただけの▲7点は取組の後退ではない
      ⚠ 宮城県平均が1年で78.6点上がったことが、県内順位3位→8位の主因
  (2) ６の保険料の要因の表に「交付金の交付見込額を見込んだ場合 ▲40円」を、
      確定していない事項の表に交付金の扱いを加える
      （素案 9-3 の不確定要素⑤・表115 と同じ値）
  (3) ⚠ ６の要因の表の値を素案（表115）に揃える
        介護報酬改定率1％　＋65円 → ＋66円（同じ資料の改定率別の表でも
          6,888円－6,822円＝66円であり、v10 は1つの資料の中で食い違っていた）
        予定収納率99.6％　▲247円 → ▲246円
  (4) 頁数の表と本文の素案の版を Ver.2.24（118頁）に改める

数値の正本　05_試算・管理シート/川崎町_交付金_推移の再整理_R8.10.22.xlsx
            計画素案 Ver.2.24（保険料の効き）
"""
import copy
import importlib.util
import os
import re
import shutil
import sys

import docx

HERE = os.path.dirname(os.path.abspath(__file__))


def load(name, fn):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(spec)
    sys.modules[name] = m
    spec.loader.exec_module(m)
    return m


V9 = load("build_2kai_v9", "build_2kai_v9.py")             # 表の操作の道具
S24 = load("fix_soan_v2240", "fix_soan_v2240_kofukin.py")  # 内訳・効き

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v10.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v11.docx"


def ten(v):
    return f"＋{v}点" if v > 0 else f"▲{-v}点"


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会資料 v10 → v11（交付金の状況の再整理）")
    print()
    for f in (SRC, S24.SAKI, S24.SAISEIRI):
        if not os.path.exists(f):
            print(f"   × 入力がありません：{f}")
            sys.exit(1)

    kubun, komoku, goukei = S24.uchiwake()
    K = {
        "成果": kubun["成果指標群の変動"],
        "配点": kubun["配点の改定（得点率は不変）"],
        "新設": kubun["指標の新設"],
        "後退": kubun["取組の後退（得点率の低下）"],
        "前進": kubun["取組の前進（得点率の上昇）"],
    }

    def ko(k, na):
        return komoku[(k, na)]

    def seika(atama):
        return sum(v for (k, na), v in komoku.items()
                   if k == "成果指標群の変動" and na.startswith(atama))

    n_kotai = len({na for (k, na) in komoku if k.startswith("取組の後退")})
    n_zen = len({na for (k, na) in komoku if k.startswith("取組の前進")})
    katsuyo = ko("取組の後退（得点率の低下）", "評価結果の活用")
    tekisei = ko("取組の後退（得点率の低下）", "給付費適正化事業の取組状況")
    G = {r[0]: r for r in goukei.values()}
    r6, r7, r8 = G["令和6年度"], G["令和7年度"], G["令和8年度"]
    gen = r8[1] - r7[1]
    judge(sum(K.values()) == gen, f"内訳の合計 {sum(K.values())}＝増減 {gen}")

    HYO_UCHI = [
        ["区　分", "増　減", "主な項目", "取り戻せるか"],
        ["成果指標群（アウトカム）", ten(K["成果"]),
         f"短期的な要介護度の変化{ten(seika('短期'))}・健康寿命延伸の状況"
         f"{ten(seika('健康寿命'))}、長期的な要介護度の変化"
         f"{ten(seika('長期'))}（推進・支援の両方に同じ点が入ります）",
         "要介護度の変化という成果によるもので、仕組みを整えるだけでは"
         "戻りません"],
        ["国の配点の改定", ten(K["配点"]),
         "認知症サポーター等を活用した地域支援体制の構築"
         "（25点満点→20点満点）ほか",
         "取組の後退ではありません（満点を維持）"],
        ["指標の新設", ten(K["新設"]),
         "介護予防・生活支援の体制整備（新設の枝番）", "―"],
        ["取組の評価が下がった項目", ten(K["後退"]),
         f"評価結果の活用{ten(katsuyo)}・給付費適正化事業の取組状況ア"
         f"{ten(tekisei)}ほか{n_kotai - 2}項目",
         f"うち{-(katsuyo + tekisei)}点は仕組みを定めることで評価の対象と"
         "なり得ます（下の失点の表）"],
        ["取組の評価が上がった項目", ten(K["前進"]),
         "在宅医療・介護連携に関する課題・対応策の検討"
         f"{ten(ko('取組の前進（得点率の上昇）', '在宅医療・介護連携に関する課題・対応策の検討'))}"
         "・地域包括支援センター事業評価の達成状況"
         f"{ten(ko('取組の前進（得点率の上昇）', '地域包括支援センター事業評価の達成状況'))}"
         f"ほか{n_zen - 2}項目", "―"],
        ["計", ten(gen), f"{r7[1]}点 → {r8[1]}点", ""],
    ]

    def jun(r):
        return f"{r[1]}点（{r[3]:,}位）"

    HYO_SUII = [
        ["年　度", "本町（全国順位）", "全国平均",
         f"宮城県平均（本町の県内順位／{r8[8]}保険者）"],
        ["令和６年度", jun(r6), f"{r6[2]:.1f}点", f"{r6[5]:.1f}点（{r6[7]}位）"],
        ["令和７年度", jun(r7), f"{r7[2]:.1f}点", f"{r7[5]:.1f}点（{r7[7]}位）"],
        ["令和８年度", jun(r8), f"{r8[2]:.1f}点", f"{r8[5]:.1f}点（{r8[7]}位）"],
    ]
    ken_up = round(r8[5] - r7[5], 1)
    zen_up = round(r8[2] - r7[2], 1)

    MIDASHI = "◆ 37点下がった内訳と県内の動き"
    HON1 = (f"{-gen}点の減少を、枝番の単位まで下ろし、各年度の配点に対する"
            "得点率で比べると、次の５つに分かれます。⚠ 国が配点を下げた"
            "だけの項目は、点が下がっても取組の後退ではありません。")
    HON2 = (f"同じ期間に、宮城県内{r8[8]}保険者の平均は令和７年度の"
            f"{r7[5]:.1f}点から令和８年度の{r8[5]:.1f}点へ、１年で{ken_up}点"
            f"上がりました（全国平均の上昇は{zen_up}点）。本町の県内順位が"
            f"{r7[7]}位から{r8[7]}位になったのは、本町が{-gen}点下げたことに"
            "加え、県内が大きく上がったためです。県平均との差は"
            f"{r7[1] - r7[5]:.1f}点から{r8[1] - r8[5]:.1f}点に縮まりました。"
            "⚠ 交付額は、得点に第１号被保険者数を乗じた値で国の予算を按分"
            "して決まるため、他の保険者の得点が上がると、本町の得点が同じでも"
            "交付額は減ることがあります。令和６年度の本町の交付見込額は"
            "計3,738千円（うち激変緩和措置分1,231千円）です。令和７・８年度の"
            "交付額は確認中です。第10期の保険料の試算では交付金を見込んで"
            "いません（見込んだ場合の効きは６の表）。")

    shutil.copy2(SRC, DST)
    d = docx.Document(DST)
    mae = "\n".join([p.text for p in d.paragraphs]
                    + [c.text for t in d.tables for r in t.rows
                       for c in r.cells])
    KYU = []
    P = d.paragraphs

    print("  ── ① 段落")
    DAN = [
        (4, "計画骨子（案）とご選択いただきたい事項　［第10版］",
         lambda s: s.replace("［第10版］", "［第11版］")),
        (70, "保険者機能強化推進交付金・介護保険保険者努力支援交付金は、",
         lambda s: s.replace("からは37点下がりました。",
                             "からは37点下がりました（内訳は下の表）。")),
        (90, "計画書の事業一覧表（17表・138事業）に",
         lambda s: s.replace("Ver.2.23", "Ver.2.24")
         .replace("本文117頁", "本文118頁").replace("計130頁", "計131頁")
         .replace("差は＋30頁", "差は＋31頁")),
    ]
    n = 0
    for i, atama, f in DAN:
        p = P[i]
        if not p.text.strip().startswith(atama):
            judge(False, f"p{i} の冒頭が合わない", f"実際『{p.text[:24]}』")
            continue
        s = f(p.text)
        if s == p.text:
            judge(False, f"p{i} が変わらない")
            continue
        KYU.append(p.text)
        p.runs[0].text = V9.clean(s)
        for r in p.runs[1:]:
            r.text = ""
        n += 1
    judge(n == len(DAN), f"{len(DAN)}段落を改めた（実際 {n}）")

    print("  ── ② 表のセル")
    MID = {i: t.rows[0].cells[0].text.strip() for i, t in enumerate(d.tables)}
    SERU = [
        (26, "要", 1, 1, "＋65円", "＋66円"),
        (26, "要", 5, 1, "▲247円", "▲246円"),
        (38, "版", 2, 0, "Ver.2.23（現在・",
         "Ver.2.24（現在・認知症の人の数の推計、公表情報による供給の把握、"
         "県の入所状況調の推移、交付金の推移の内訳を追加）"),
        (38, "版", 2, 1, "117頁", "118頁"),
        (38, "版", 2, 2, "130頁", "131頁"),
        (38, "版", 2, 3, "＋30頁", "＋31頁"),
    ]
    n = 0
    for ban, mi, gyo, retsu, kyu, atarashii in SERU:
        if not MID[ban].startswith(mi):
            judge(False, f"表{ban} の見出しが合わない", f"実際『{MID[ban][:16]}』")
            continue
        c = d.tables[ban].rows[gyo].cells[retsu]
        if not c.text.strip().startswith(kyu):
            judge(False, f"表{ban} 行{gyo} 列{retsu} の旧が合わない",
                  f"実際『{c.text[:24]}』")
            continue
        KYU.append(c.text)
        V9.set_cell(c, atarashii)
        n += 1
    judge(n == len(SERU), f"{len(SERU)}セルを改めた（実際 {n}）")

    print("  ── ③ ６の表に交付金の行を足す")
    t26, t28 = d.tables[26], d.tables[28]
    ok = (MID[26].startswith("要") and MID[28].startswith("事"))
    judge(ok, "表26 が要因の表・表28 が確定していない事項の表である")
    if ok:
        for t, gyo in (
                (t26, ["交付金の交付見込額（３か年4,500千円）を見込んだ場合",
                       f"▲{-S24.KOFU_SA}円"]),
                (t28, ["保険者機能強化推進交付金等の交付見込額を保険料の算定に"
                       "見込むか（令和７・８年度の交付額を含む）",
                       "第３回策定委員会でご協議"])):
            tr = copy.deepcopy(t.rows[-1]._tr)
            t.rows[-1]._tr.addnext(tr)
            for c, v in zip(t.rows[-1].cells, gyo):
                V9.set_cell(c, v)

    print("  ── ④ ７の冒頭に内訳と推移を足す")
    p70, p71, p72 = P[70], P[71], P[72]
    judge(p71.text.startswith("◆ ３つの指標群"), "p71 が小見出しである")
    t31 = d.tables[31]
    judge(t31.rows[0].cells[0].text.startswith("指標群"), "表31 が指標群の表である")
    els = []
    h = copy.deepcopy(p71._p)
    els.append((h, MIDASHI))
    b1 = copy.deepcopy(p72._p)
    els.append((b1, HON1))
    ta = copy.deepcopy(t31._tbl)
    els.append((ta, HYO_UCHI))
    b2 = copy.deepcopy(p72._p)
    els.append((b2, HON2))
    tb = copy.deepcopy(t31._tbl)
    els.append((tb, HYO_SUII))
    mae_el = p70._p
    import docx.table as dt
    import docx.text.paragraph as dp
    for el, naiyo in els:
        mae_el.addnext(el)
        mae_el = el
        if isinstance(naiyo, str):
            para = dp.Paragraph(el, p70._parent)
            para.runs[0].text = V9.clean(naiyo)
            for r in para.runs[1:]:
                r.text = ""
        else:
            t = dt.Table(el, t31._parent)
            while len(t.rows) > len(naiyo):
                el.remove(t.rows[-1]._tr)
            while len(t.rows) < len(naiyo):
                el.append(copy.deepcopy(t.rows[-1]._tr))
            for ri, gyo in enumerate(naiyo):
                for ci, v in enumerate(gyo):
                    V9.set_cell(t.rows[ri].cells[ci], v)
    judge(len(d.tables) == 47, f"表が45 → 47（実際 {len(d.tables)}）")

    d.save(DST)

    d2 = docx.Document(DST)
    ato = "\n".join([p.text for p in d2.paragraphs]
                    + [c.text for t in d2.tables for r in t.rows
                       for c in r.cells])
    soan = docx.Document(S24.SAKI)
    soan_t = "\n".join([p.text for p in soan.paragraphs]
                       + [c.text for t in soan.tables for r in t.rows
                          for c in r.cells])

    print("  ── 検査1　内訳と推移")
    for go in (MIDASHI, "▲30点", "▲7点", "＋2点", "▲34点", "＋32点", "▲37点",
               "497.5点", "78.6点", "30.5点", "3,738千円"):
        judge(go in ato, f"「{go}」がある")
    judge("146.1ポイント" not in ato and "30.5ポイント" not in ato,
          "点の差を「ポイント」と書いていない")
    print("  ── 検査2　⚠ 素案と同じ値か（保険料の効き）")
    t115 = [t for t in soan.tables
            if [c.text for c in t.rows[0].cells][:3]
            == ["前提", "本算定の置き方", "動かした場合"]][0]
    soan_sa = {r.cells[2].text: r.cells[4].text for r in t115.rows[1:]}
    judge(soan_sa.get("1％") == "＋66円" and "＋66円" in ato
          and "＋65円" not in ato, "改定率1％＝＋66円（素案と同じ）")
    judge(soan_sa.get("99.6％（令和7年度の現年度分実績）") == "▲246円"
          and "▲246円" in ato and "▲247円" not in ato,
          "収納率99.6％＝▲246円（素案と同じ）")
    judge(f"▲{-S24.KOFU_SA}円" in ato and f"▲{-S24.KOFU_SA}円" in soan_t,
          f"交付金を見込んだ場合 ▲{-S24.KOFU_SA}円（素案と同じ）")
    judge("令和8年度の交付実績を据え置いて" not in ato,
          "事実と違う記述（交付実績の据え置き）が0件")
    print("  ── 検査3　版と頁")
    judge("Ver.2.23" not in ato and "Ver.2.24" in ato, "素案の版が Ver.2.24")
    judge("［第11版］" in ato, "版が［第11版］")
    print("  ── 検査4　⚠ 既にあった数が消えていないか")
    kyu_zen = "\n".join(KYU)
    kieta = sorted(set(
        m.group(0) for m in re.finditer(
            r"[0-9０-９][0-9０-９,.]*(?:人|名|％|千円|件|頁|日|点|円)", mae)
        if mae.count(m.group(0)) > ato.count(m.group(0))
        and m.group(0) not in kyu_zen))
    judge(not kieta, "消えた数はすべて差し替えた旧の文に現れている",
          "／".join(kieta[:5]))
    print("  ── 検査5　⚠ 事業所名・他団体名・強調記号")
    for na in ("釜房みどりの園", "アルパイン", "蔵王グリーンホーム",
               "グループホームさわやか", "グループホームふれあい",
               "ながさか", "リリーハイム", "北塩原", "金ヶ崎"):
        judge(na not in ato, f"「{na}」が入っていない")
    judge("川崎市" not in ato and "福岡" not in ato,
          "川崎市・福岡県川崎町が混じらない")
    judge("**" not in ato, "強調記号（**）が0件")

    print()
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        for m in futeki:
            print(f"      {m}")
        sys.exit(1)
    print("   ○ 検査1〜5　不適合0件")


if __name__ == "__main__":
    main()
