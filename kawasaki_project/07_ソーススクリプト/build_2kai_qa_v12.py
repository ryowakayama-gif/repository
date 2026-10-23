# -*- coding: utf-8 -*-
"""第2回策定委員会 想定問答集 v11 → v12（協議用素案 Ver.2.29 の頁数に揃える）
（令和8年10月23日）

  python3 07_ソーススクリプト/build_2kai_qa_v12.py

協議用素案 Ver.2.29 で総頁数が117頁 → 118頁になった。頁数の問（QB-40・QB-56）と
根拠データの頁数を揃える。委員会資料 v15 と同じ値。問の数（全68問）は変えない。
"""
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v11.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v12.docx"
KYU, HEN = 117, 118
SHIRYO = 13
SU = 68


def txt(el):
    return "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()


def main():
    futeki = []

    def judge(ok, mi):
        print(("   ○ " if ok else "   × ") + mi)
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会 想定問答集 v11 → v12（頁数を協議用素案 Ver.2.29 に揃える）")
    from fix_soan_v111 import set_el
    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    n = 0
    for el in doc.element.body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = txt(el)
        if f"{KYU}頁" in s:
            set_el(el, s.replace(f"{KYU}頁", f"{HEN}頁")
                   .replace(f"計{KYU + SHIRYO}頁", f"計{HEN + SHIRYO}頁"))
            n += 1
    judge(n == 5, f"頁数の記述 5か所を改めた（実際 {n}）")
    doc.save(DST)

    d2 = docx.Document(DST)
    zen = "\n".join(p.text for p in d2.paragraphs) + "\n" + "\n".join(
        c.text for t in d2.tables for r in t.rows for c in r.cells)
    judge(f"{KYU}頁" not in zen and f"{KYU + SHIRYO}頁" not in zen, f"{KYU}頁・{KYU + SHIRYO}頁 0件")
    qs = [int(x) for x in re.findall(r"QB-(\d+)　", zen)]
    judge(sorted(qs) == list(range(1, SU + 1)) and f"全{SU}問" in zen, f"全{SU}問・欠番なし")
    for na in ("釜房", "アルパイン", "鶴寿会", "北塩原", "金ヶ崎", "大雪", "川崎市"):
        judge(na not in zen, f"「{na}」0件")
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 不適合0件")


if __name__ == "__main__":
    main()
