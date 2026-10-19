# -*- coding: utf-8 -*-
"""将来推計総括表の2版を**全シート・全セル**で突き合わせる（令和8年10月19日）

  python3 07_ソーススクリプト/compare_sokatsu_zen_R8.10.19.py <旧版> <新版> [出力txt]

なぜ要るか
  `compare_sokatsu.py` は「2_サービス別給付費」しか見ない。
  ⚠ **令和8年10月19日に受領した版は、認定者数・保険料・第1号被保険者数の
  年齢区分が動いていた。** 給付費だけを見ていると気づけない。
  版を受け取るたびに、**どのシートのどの行が動いたか**を全件で数える。

  終了コード　0＝差が無い／1＝差がある（差があること自体は異常ではない。
  「気づかずに素案へ持ち込む」ことが異常である）

⚠ 見出し（行の名前）は、1〜3列目のうち文字が入っている最も右の列から採る。
  総括表は階層のある表で、項目名の入る列が行により違うためである。
"""
import os
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")

# 数の差をどこから「動いた」と見るか（丸め・浮動小数の揺れを除く）
SHIKII_ZETTAI = 0.5      # 絶対差がこれ以下なら同じとみなす
SHIKII_WARIAI = 0.0005   # 相対差がこれ以下なら同じとみなす


def midashi(ws, r):
    """行の名前。1〜3列目のうち文字のある最も右の列。無ければ上の行から。"""
    for rr in range(r, max(r - 12, 0), -1):
        for c in (3, 2, 1):
            v = ws.cell(rr, c).value
            if isinstance(v, str) and v.strip():
                na = v.strip().replace("\n", " ")
                return (na if rr == r else f"（{na}）の下 {r - rr}行目")
    return f"{r}行目"


def yomu(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for ws in wb:
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v is not None and v != "":
                    out[(ws.title, r, c)] = v
    return wb, out


def onaji(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) \
            and not isinstance(a, bool) and not isinstance(b, bool):
        sa = abs(a - b)
        if sa <= SHIKII_ZETTAI:
            return True
        moto = max(abs(a), abs(b))
        return moto > 0 and sa / moto <= SHIKII_WARIAI
    return str(a).strip() == str(b).strip()


def main():
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    kyu, shin = sys.argv[1], sys.argv[2]
    for p in (kyu, shin):
        if not os.path.exists(p):
            raise SystemExit("ファイルがない：" + p)

    wb_k, a = yomu(kyu)
    wb_s, b = yomu(shin)
    ws_s = {w.title: w for w in wb_s}
    ws_k = {w.title: w for w in wb_k}

    dashi = []

    def p(s=""):
        dashi.append(s)
        print(s)

    p("将来推計総括表の全セル突合")
    p(f"  旧 {os.path.basename(kyu)}")
    p(f"  新 {os.path.basename(shin)}")
    p()
    p(f"  シート　旧 {len(wb_k.sheetnames)}／新 {len(wb_s.sheetnames)}")
    nai_s = [x for x in wb_k.sheetnames if x not in wb_s.sheetnames]
    nai_k = [x for x in wb_s.sheetnames if x not in wb_k.sheetnames]
    if nai_s:
        p(f"  ⚠ 新版に無いシート：{nai_s}")
    if nai_k:
        p(f"  ⚠ 旧版に無いシート：{nai_k}")
    p(f"  値の入ったセル　旧 {len(a):,}／新 {len(b):,}")
    p()

    kagi = sorted(set(a) | set(b))
    chigai = []
    for k in kagi:
        va, vb = a.get(k), b.get(k)
        if va is None and vb is None:
            continue
        if va is None or vb is None or not onaji(va, vb):
            chigai.append((k, va, vb))

    # シートごとにまとめる
    by_sheet = {}
    for (sh, r, c), va, vb in chigai:
        by_sheet.setdefault(sh, []).append(((r, c), va, vb))

    p(f"  ── 動いたセル　{len(chigai):,}件")
    for sh in wb_s.sheetnames:
        ms = by_sheet.get(sh, [])
        p(f"   {sh}：{len(ms)}件")
    p()

    for sh in wb_s.sheetnames:
        ms = by_sheet.get(sh, [])
        if not ms:
            continue
        p(f"  ══════ {sh}（{len(ms)}件）")
        # 行ごとにまとめる
        by_row = {}
        for (r, c), va, vb in ms:
            by_row.setdefault(r, []).append((c, va, vb))
        for r in sorted(by_row):
            w = ws_s.get(sh) or ws_k.get(sh)
            na = midashi(w, r)
            kumi = []
            for c, va, vb in sorted(by_row[r]):
                def f(v):
                    if isinstance(v, float):
                        return (f"{v:,.4f}".rstrip("0").rstrip(".")
                                if abs(v) < 1e7 else f"{v:,.0f}")
                    if isinstance(v, int):
                        return f"{v:,}"
                    return "（空）" if v is None else str(v)[:28]
                kumi.append(f"列{c} {f(va)} → {f(vb)}")
            p(f"   {r:>4}　{na[:40]:<40}" + "／".join(kumi[:6])
              + ("　ほか" if len(kumi) > 6 else ""))
        p()

    if len(sys.argv) > 3:
        with open(sys.argv[3], "w", encoding="utf-8") as f:
            f.write("\n".join(dashi) + "\n")
        print(f"  書き出し　{sys.argv[3]}")

    sys.exit(1 if chigai else 0)


if __name__ == "__main__":
    main()
