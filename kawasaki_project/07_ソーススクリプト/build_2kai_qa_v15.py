# -*- coding: utf-8 -*-
"""第2回策定委員会 想定問答集 v14 → v15（Red Team 再レビュー（令和8年10月24日）の指摘の是正）

  python3 07_ソーススクリプト/build_2kai_qa_v15.py

資料 v18・素案 Ver.2.33 と同じ値・言い回しに揃える。問の数（全69問）は変えない。
  N01 報酬改定の算定Ｂ・Ｃの差（＋92円・＋223円・＋119円・＋250円）は誤った表26 から計算していた。
      資料 v18 の表26（素案 9-3 から起こしたもの）から計算し直す。取崩しは一定額の差引きのため、
      改定による差は算定Ａ・Ｂ・Ｃでほぼ同じ
  N04 「介護給付費準備基金の令和8年度末残高の確定」→ 見込額の精査（決算による確定は令和9年度）
  N08 「1,741保険者」「35保険者」「県内の保険者の得点」等 → 市町村
  N09 入所・居住系10人 約211円 → 約212円（素案 5-5）
"""
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v14.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会_想定問答集_R8.11_v15.docx"
SHIRYO = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v18.docx"
SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.33_協議用.docx"
SU = 69


def txt(el):
    return "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()


def yen(s):
    return int(re.sub(r"[^\d]", "", s))


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会 想定問答集 v14 → v15（Red Team 再レビューの指摘の是正）")
    s = docx.Document(SHIRYO)
    t26 = [t for t in s.tables if t.rows[0].cells[0].text.strip() == "介護報酬改定率"][0]
    g = {r.cells[0].text.strip()[:2]: [yen(c.text) for c in r.cells[1:4]] for r in t26.rows[1:4]}
    d1 = [a - b for a, b in zip(g["１％"], g["０％"])]
    d3 = [a - b for a, b in zip(g["３％"], g["０％"])]
    judge(d1[0] == 66 and d3[0] == 196, f"資料 v18 表26 から 1％ Ａ＋{d1[0]}・Ｂ＋{d1[1]}・Ｃ＋{d1[2]}／"
                                         f"3％ Ａ＋{d3[0]}・Ｂ＋{d3[1]}・Ｃ＋{d3[2]}")
    sz = "\n".join(p.text for p in docx.Document(SOAN).paragraphs)
    j10 = int(re.search(r"10人では約(\d+)円動きます", sz).group(1))
    judge(j10 == 212, f"素案 5-5 から 10人 約{j10}円")

    NAOSU = [
        ("（いずれも基金の取崩しを行わない算定Ａの場合。事務局案の算定Ｂでは、改定率1％で月額92円です）",
         f"（いずれも基金の取崩しを行わない算定Ａの場合。基金の取崩しは一定額を差し引くものであるため、"
         f"事務局案の算定Ｂでも改定率1％で月額{d1[1]}円とほぼ同じです）"),
        ("事務局案の算定Ｂでは1％で約92円・3％で約223円、算定Ｃでは1％で約119円・3％で約250円です。",
         f"事務局案の算定Ｂでは1％で約{d1[1]}円・3％で約{d3[1]}円、算定Ｃでは1％で約{d1[2]}円・"
         f"3％で約{d3[2]}円です。基金の取崩しは一定額を差し引くものであるため、改定による差は"
         "算定Ａ・Ｂ・Ｃでほぼ同じです。"),
        ("算定Ｂ＋92円・＋223円、算定Ｃ＋119円・＋250円",
         f"算定Ｂ＋{d1[1]}円・＋{d3[1]}円、算定Ｃ＋{d1[2]}円・＋{d3[2]}円"),
        ("10人増減すると保険料は月額約211円動く",
         f"10人増減すると保険料は月額約{j10}円動く（1円未満四捨五入）"),
        ("10人増えると保険料は月額約211円上がります。", f"10人増えると保険料は月額約{j10}円上がります。"),
        ("入所・居住系の利用者10人で月額約211円", f"入所・居住系の利用者10人で月額約{j10}円"),
        ("それまでに、①介護給付費準備基金の令和8年度末残高の確定、②調整交付金の見込交付割合（宮城県への照会）、"
         "③令和9年度の介護報酬改定の告示、④第1号被保険者負担割合の政令、の4件が確定する必要があります。",
         "それまでに、①介護給付費準備基金の令和8年度末残高の見込額の精査（決算による確定は令和9年度であり、"
         "第3回では見込額を用います）、②調整交付金の見込交付割合（宮城県への照会）、③令和9年度の介護報酬改定の"
         "告示、④第1号被保険者負担割合の政令、の4件を確かめておく必要があります。"),
        # N08
        ("宮城県内35保険者の平均", "宮城県内35市町村の平均"),
        ("全国1,741保険者の96.9％", "全国1,741市町村の96.9％"),
        ("全国で14.6％・4.8％の保険者しか得点していない", "全国で14.6％・4.8％の市町村しか得点していない"),
        ("各保険者の評価得点", "各市町村の評価得点"),
        ("全保険者で按分", "全市町村で按分"),
        ("他の保険者の得点", "他の市町村の得点"),
        ("県内の保険者の得点が大きく上がった", "県内の市町村の得点が大きく上がった"),
    ]
    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body
    from fix_soan_v111 import set_el
    els = [el for el in body.iter(qn("w:p"))]
    for kyu, shin in NAOSU:
        hit = [el for el in els if kyu in txt(el)]
        judge(len(hit) >= 1, f"「{kyu[:24]}」（{len(hit)}か所）")
        for el in hit:
            set_el(el, txt(el).replace(kyu, shin))
    doc.save(DST)

    d2 = docx.Document(DST)
    zen = "\n".join(txt(el) for el in d2.element.body.iter(qn("w:p")))
    print("  ── 検査1　件数")
    ns = [int(x) for x in re.findall(r"QB-(\d+)　", zen)]
    judge(sorted(ns) == list(range(1, SU + 1)) and f"全{SU}問" in zen, f"QB-1〜{SU}・欠番なし")
    print("  ── 検査2　古い記述")
    for go in ("92円", "223円", "119円", "250円", "約211円", "残高の確定", "1,741保険者", "35保険者",
               "保険者しか得点", "全保険者", "他の保険者", "県内の保険者"):
        judge(go not in zen, f"「{go}」0件")
    print("  ── 検査3　資料・素案と同じ値")
    judge(f"約{j10}円" in zen, f"約{j10}円")
    print("  ── 検査4　事業所名・他団体名")
    for na in ("釜房", "アルパイン", "鶴寿会", "グループホームさわやか", "グループホームふれあい",
               "北塩原", "金ヶ崎", "大雪", "川崎市"):
        judge(na not in zen, f"「{na}」0件")
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜4　不適合0件")


if __name__ == "__main__":
    main()
