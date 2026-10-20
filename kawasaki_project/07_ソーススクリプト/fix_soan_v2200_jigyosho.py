# -*- coding: utf-8 -*-
"""計画素案 Ver.2.20（町内の事業所の確定・令和8年10月20日）

  入力　01_第10期_最新版成果品/川崎町_計画書素案_v2.19_認知症推計版.docx
  出力　01_第10期_最新版成果品/川崎町_計画書素案_v2.20_事業所確定版.docx

なぜ要るか
  作業順位26（事業所の突合）を、介護サービス情報公表システムの
  **オープンデータ**により実施した。その結果を素案に反映する。
  突合の結果は
  05_試算・管理シート/川崎町_事業所・住まい_突合結果_R8.10.20.xlsx。

══════ 直すもの
  ① ⚠ **誤りが1件あった。** 2-7 の表と 7-1 の本文が
     「小規模多機能型居宅介護等の地域密着型サービスが中心」としているが、
     **町内に小規模多機能型居宅介護はない。**
  ② ［要確認］が4件解ける
     ・小規模多機能型居宅介護（9-1 の0のサービスの表）
     ・介護医療院（同上）
     ・地域密着型介護老人福祉施設入所者生活介護の町内の定員（7-1）
     ・在宅系の地域密着型サービスの町内の事業所の有無（6-4 ほか）
  ③ 2-7 に、町外の事業所への依存と担い手の集中を加える

⚠ **認知症対応型共同生活介護の町内の定員は［要確認］を残す。**
  2事業所のうち1事業所が定員を報告しておらず、9人までしか分からない。
  「分からないもの」を埋めない。

⚠ **本計画に個々の法人名・事業所名は掲げない。** 数と集まり方のみを示す。

⚠ **数値はスクリプトに書かない。** 突合の結果（オープンデータの抽出）
  から読む。
"""
import copy
import csv
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from fix_soan_v111 import set_el, set_tc                     # noqa: E402

SRC = "01_第10期_最新版成果品/川崎町_計画書素案_v2.19_認知症推計版.docx"
DST = "01_第10期_最新版成果品/川崎町_計画書素案_v2.20_事業所確定版.docx"
CSV_MEISAI = ("09_元資料/その他データ/"
              "公表システム_オープンデータ_川崎町_R8.6末.csv")
KYU, SHIN = "Ver.2.19", "Ver.2.20"

SHUTTEN = ("厚生労働省「介護サービス情報公表システム」オープンデータ"
           "（令和8年6月末時点）")

# 町内に事業所が無いのに令和7年度の実績があるサービス（突合の04シート）
#   ⚠ 値は素案 9-1 から読むため、ここには名前だけを持つ。
GAIBU = ["訪問入浴介護", "訪問看護", "福祉用具貸与",
         "特定施設入居者生活介護", "定期巡回・随時対応型訪問介護看護",
         "看護小規模多機能型居宅介護", "地域密着型通所介護"]


def yomu():
    if not os.path.exists(CSV_MEISAI):
        raise SystemExit("突合の抽出がない：" + CSV_MEISAI)
    with open(CSV_MEISAI, encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    svc, hojin, ban = {}, set(), set()
    for r in rows:
        k = r["サービスコード"]
        t = (r.get("定員") or "0").strip()
        n = int(float(t)) if t.replace(".", "").isdigit() else 0
        o = svc.setdefault(k, {"件数": 0, "定員": 0, "未報告": 0})
        o["件数"] += 1
        if n > 0:
            o["定員"] += n
        else:
            o["未報告"] += 1
        hojin.add(r["法人番号"])
        ban.add(r["事業所番号"])
    return rows, svc, hojin, ban


def hon(el):
    return "".join(x.text or "" for x in el.iter(qn("w:t")))


def main():
    if not os.path.exists(SRC):
        raise SystemExit("入力がない：" + SRC)
    rows, svc, hojin, ban = yomu()

    g540 = svc.get("540", {"件数": 0, "定員": 0, "未報告": 0})
    g320 = svc.get("320", {"件数": 0, "定員": 0, "未報告": 0})
    g510 = svc.get("510", {"件数": 0, "定員": 0})
    g520 = svc.get("520", {"件数": 0, "定員": 0})
    g210 = svc.get("210", {"件数": 0, "定員": 0})

    # ── 置き換えるもの（旧, 新）
    TAIRITSU_26 = (
        "特別養護老人ホーム1施設、介護老人保健施設1施設、"
        "小規模多機能型居宅介護等の地域密着型サービスが中心"
        "　⚠ 出典が確定しておらず、地域密着型介護老人福祉施設及び"
        "認知症対応型共同生活介護が落ちています（確認事項No.29）")
    SHIN_26 = (
        f"町内には{len(ban)}の事業所があり、"
        f"延べ{len(rows)}のサービスを担っています（{SHUTTEN}）。"
        f"介護老人福祉施設1（定員{g510['定員']}人）、"
        f"介護老人保健施設1（定員{g520['定員']}人）、"
        f"地域密着型介護老人福祉施設入所者生活介護"
        f"{g540['件数']}（定員計{g540['定員']}人）、"
        f"認知症対応型共同生活介護{g320['件数']}、"
        f"短期入所生活介護1（定員{g210['定員']}人）、"
        f"訪問介護・通所介護・訪問リハビリテーション・"
        f"通所リハビリテーション・短期入所療養介護・居宅介護支援です。"
        f"　⚠ 町内に小規模多機能型居宅介護はありません。"
        f"　⚠ 定員・入所者数・入所率は、認知症対応型共同生活介護の"
        f"1事業所を含め未報告のものがあり、照会票で確かめます"
        f"（確認事項No.29・No.159）")

    HONBUN = [
        # ① 7-1 の本文の誤り
        ("町内のサービス提供体制は、特別養護老人ホーム1施設・"
         "介護老人保健施設1施設・小規模多機能型居宅介護等の"
         "地域密着型サービスが中心です。",
         f"町内のサービス提供体制は、介護老人福祉施設1施設・"
         f"介護老人保健施設1施設・"
         f"地域密着型介護老人福祉施設入所者生活介護{g540['件数']}施設・"
         f"認知症対応型共同生活介護{g320['件数']}事業所を中心とし、"
         f"これらを含む{len(ban)}の事業所が町内にあります"
         f"（{SHUTTEN}）。"),
        # ② 在宅系の地域密着型の［要確認］
        ("町内及び近隣市町村に指定事業所が存在するかを確認できておらず"
         "［要確認］、存在しない場合、利用が0件であることは供給側の事実であり、"
         "利用勧奨では対応できません。",
         f"{SHUTTEN}により、これらのサービスの事業所は町内にないこと"
         f"を確かめました（夜間対応型訪問介護及び認知症対応型通所介護、"
         f"小規模多機能型居宅介護も町内にありません）。"
         f"定期巡回・随時対応型訪問介護看護と看護小規模多機能型居宅介護に"
         f"わずかな実績があるのは、町外の事業所の利用によるものです。"
         f"近隣市町村における事業所の有無は確かめていません［要確認］。"
         f"利用勧奨では対応できない点は変わりません。"),
    ]

    # (表の先頭の列の見出し, 行の1列目, 置き換える列の番号, 旧, 新)
    #   ⚠ **セルの文字だけで探さない。** 「［要確認］」だけのセルは
    #     素案に複数あるため、表と行を指定して当てる。
    SERU = [
        ("0の区分", "小規模多機能型居宅介護（介護予防を含む）", 2,
         "［要確認］町内に事業所があるとの整理があり、実績との食い違いを"
         "確認中",
         f"供給なし（{SHUTTEN}により、町内に事業所がないことを"
         f"確かめました）"),
        ("0の区分", "介護医療院", 2, "［要確認］供給なし又は利用なし",
         f"供給なし（{SHUTTEN}により、町内に事業所がないことを"
         f"確かめました）"),
        ("町内の定員", "地域密着型介護老人福祉施設入所者生活介護", -1,
         "［要確認］",
         f"{g540['定員']}人（{g540['件数']}施設・{SHUTTEN}）"),
        ("町内の定員", "認知症対応型共同生活介護", -1,
         "［要確認］\n第9期計画値は月27.0人",
         f"{g320['件数']}事業所（{SHUTTEN}）\n"
         f"⚠ 定員は［要確認］。{g320['未報告']}事業所が定員を報告して"
         f"いないため、分かっているのは{g320['定員']}人までです\n"
         f"第9期計画値は月27.0人"),
    ]

    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body

    # ── 本文の置き換え
    n_hon = 0
    nokori = []
    for furui, atarashii in HONBUN:
        atta = False
        for el in body.iterchildren():
            if el.tag != qn("w:p"):
                continue
            s = hon(el)
            if furui in s:
                set_el(el, s.replace(furui, atarashii))
                atta = True
                n_hon += 1
                break
        if not atta:
            nokori.append(furui[:34])

    # ── 表のセルの置き換え（表と行を指定して当てる）
    n_seru = 0
    for midashi, gyo, retsu, furui, atarashii in SERU:
        atta = False
        for t in doc.tables:
            h = [c.text.strip() for c in t.rows[0].cells]
            if not any(midashi in x for x in h):
                continue
            for r in t.rows[1:]:
                c = r.cells
                na = re.sub(r"\s+", "", c[0].text)
                if not na.startswith(re.sub(r"\s+", "", gyo)):
                    continue
                tgt = c[retsu]
                ima = re.sub(r"\s+", "", tgt.text)
                if ima != re.sub(r"\s+", "", furui):
                    continue
                set_tc(tgt._tc, atarashii)
                atta = True
                n_seru += 1
                break
            if atta:
                break
        if not atta:
            nokori.append(f"{midashi}／{gyo}")

    # ── 2-7 の表（事業所・施設）
    n_26 = 0
    for tc in body.iter(qn("w:tc")):
        if hon(tc).strip() == TAIRITSU_26.strip():
            set_tc(tc, SHIN_26)
            n_26 += 1
            break
    if n_26 == 0:
        nokori.append("2-7 の表（事業所・施設）")

    # ── 2-7 に、町外への依存と担い手の集中を書き足す
    oku = None
    for el in body.iterchildren():
        if el.tag == qn("w:p") and "　▌ 供給制約を示す事象" in hon(el):
            oku = el
            break
    tpl = None
    for el in body.iterchildren():
        if el.tag == qn("w:p") and hon(el).startswith(
                "施設・事業所は役場周辺に偏在しており"):
            tpl = el
            break
    n_tsuika = 0
    if oku is not None and tpl is not None:
        for mono in [
            f"町内の供給には2つの特徴があります。第1に、町民が利用して"
            f"いるサービスのうち7種類は、町内に事業所がありません。"
            f"{'、'.join(GAIBU)}がこれに当たり、いずれも町外の事業所の"
            f"利用によるものです（{SHUTTEN}と第9章9-1の令和7年度の実績の"
            f"突合による）。町域の供給だけでは町民の利用を説明できない"
            f"構造が、ここにも表れています。",
            f"第2に、町内の{len(ban)}の事業所を運営しているのは"
            f"{len(hojin)}法人です。1つの法人が複数の種類のサービスを"
            f"担っており、1法人の撤退がサービスの種類ごと失われることを"
            f"意味します。第7章7-1の提供体制の確保では、この点を"
            f"踏まえた検討が要ります。",
        ]:
            e = copy.deepcopy(tpl)
            set_el(e, mono)
            oku.addprevious(e)
            n_tsuika += 1
    else:
        nokori.append("2-7 への書き足しの位置")

    # ── 版表記
    n_ver = 0
    for el in body.iter(qn("w:t")):
        if el.text and KYU in el.text:
            el.text = el.text.replace(KYU, SHIN)
            n_ver += 1
    for sec in doc.sections:
        for hf in (sec.header, sec.footer, sec.even_page_header,
                   sec.first_page_header):
            for el in hf._element.iter(qn("w:t")):
                if el.text and KYU in el.text:
                    el.text = el.text.replace(KYU, SHIN)
                    n_ver += 1

    if nokori:
        raise SystemExit("当てられなかった箇所：" + "／".join(nokori))

    doc.save(DST)
    print(f"作成　{DST}")
    print(f"  本文 {n_hon}／表のセル {n_seru}／2-7の表 {n_26}／"
          f"書き足し {n_tsuika}／版表記 {n_ver}")

    # ══════════════════════════ 点検
    print()
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    d2 = docx.Document(DST)
    zen = [p.text for p in d2.paragraphs]
    for t in d2.tables:
        for r in t.rows:
            for c in r.cells:
                zen.append(c.text)
    s = "\n".join(zen)
    d0 = docx.Document(SRC)
    zen0 = [p.text for p in d0.paragraphs]
    for t in d0.tables:
        for r in t.rows:
            for c in r.cells:
                zen0.append(c.text)
    s0 = "\n".join(zen0)

    print("  ── 検査1　誤りが直ったか")
    judge("小規模多機能型居宅介護等の地域密着型サービスが中心" not in s,
          "「小規模多機能型居宅介護等…が中心」の記述が無くなった")
    judge("町内に小規模多機能型居宅介護はありません" in s,
          "町内に無いことが書かれている")

    print("  ── 検査2　［要確認］が減ったか")
    n0 = s0.count("［要確認］")
    n1 = s.count("［要確認］")
    judge(n1 < n0, f"［要確認］が {n0}件 → {n1}件（▲{n0 - n1}）")
    judge("［要確認］" in s,
          "⚠ 分かっていないものは［要確認］のまま残っている",
          "（認知症対応型共同生活介護の定員・近隣市町村の事業所ほか）")

    print("  ── 検査3　数が突合の結果と合うか")
    judge(f"{g540['定員']}人" in s,
          f"地域密着型介護老人福祉施設の定員 {g540['定員']}人")
    judge(f"{len(ban)}の事業所" in s, f"町内の事業所 {len(ban)}")
    judge(f"{len(hojin)}法人" in s, f"法人 {len(hojin)}")
    judge(f"{g320['定員']}人までです" in s,
          f"認知症対応型共同生活介護は{g320['定員']}人までと書いている")

    print("  ── 検査4　既にあった数値を動かしていないか")
    from collections import Counter
    c0 = Counter(re.findall(r"\d[\d,]*", s0))
    c2 = Counter(re.findall(r"\d[\d,]*", s))
    kieta = {k: v - c2.get(k, 0) for k, v in c0.items() if v > c2.get(k, 0)}
    kieta.pop("2.19", None)
    kieta.pop("19", None)
    # ⚠ **意図して消したものだけを、理由つきで許す。**
    #   旧 6-4 の「存在しない場合、利用が0件であることは供給側の事実」は、
    #   実際には定期巡回0.4人／月・看多機0.2人／月の実績があり、
    #   「0件」という仮定そのものが成り立たないため落とした。
    KIETE_YOI = {"0": 1}
    for k, v in KIETE_YOI.items():
        if kieta.get(k) == v:
            kieta.pop(k)
    judge(not kieta, "元からあった数が、意図したもの以外に消えていない",
          f"（消えた {dict(list(kieta.items())[:4])}）" if kieta else "")

    print("  ── 検査5　制約")
    judge("法人の名称" not in s and "鶴寿会" not in s
          and "光友会" not in s and "社会福祉協議会　訪問" not in s,
          "個々の法人名・事業所名を本文に掲げていない")
    for mi, pat in (("受託者を主語とする語",
                     r"(受託者|当社|弊社|ビズアップ|スクリプト)"),
                    ("強調記号の残り", r"(\*\*|__)"),
                    ("他団体の固有名称",
                     r"(大雪|東川町|東神楽町|金ケ崎町|金ヶ崎町|北塩原村)"),
                    ("川崎市との混同", r"川崎市")):
        atari = [x[:50] for x in zen if re.search(pat, x)]
        judge(not atari, f"{mi} が無い", f"→ {atari[:1]}" if atari else "")

    print("  ── 検査6　版表記")
    nok = re.findall(r"Ver\.\d+\.\d+[a-z]?", s)
    judge(all(v == SHIN for v in nok),
          f"本文の版表記がすべて {SHIN}（{len(nok)}か所）")
    hdr = "\n".join(p.text for sec in d2.sections
                    for p in sec.header.paragraphs)
    judge(KYU not in hdr and SHIN in hdr, "ヘッダーの版表記も改めた")

    print()
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        for m in futeki:
            print(f"      {m}")
        sys.exit(1)
    print("   ○ 検査1〜6　不適合0件")


if __name__ == "__main__":
    main()
