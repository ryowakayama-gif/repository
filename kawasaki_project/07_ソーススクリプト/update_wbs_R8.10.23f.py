# -*- coding: utf-8 -*-
"""業務工程管理表の更新（宮城県の第10期計画の骨子（案）・見える化の定めの確認：
素案Ver.2.29・資料v15・問答集v12）（令和8年10月23日・同日6度目）

  python3 07_ソーススクリプト/update_wbs_R8.10.23f.py

  03_確認事項一覧　No.188 を起票。No.50・69・77・79・80・168 に追記
  04_成果品管理　　版を Ver.2.29／v15／v12 に。概要版が古い素案によることを記す
  06_資料提供依頼一覧　No.51・No.52 を加える（50件 → 52件）
  08_作業順位　　　45（完了）・46（未着手）を加える

⚠ 列8（状態）は動かさない（No.188 の新設を除く）。再実行しても二重に書かない。
"""
import copy
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月23日 更新⑥（県の第10期計画の骨子（案）・見える化の定め）】"
OKIKAE = [("川崎町_計画書素案_v2.28_協議用.docx", "川崎町_計画書素案_v2.29_協議用.docx"),
          ("川崎町_計画書案_公表用_v2.28.docx", "川崎町_計画書案_公表用_v2.29.docx"),
          ("川崎町_第2回策定委員会資料_R8.11_v14.docx", "川崎町_第2回策定委員会資料_R8.11_v15.docx"),
          ("同_想定問答集_R8.11_v11.docx", "同_想定問答集_R8.11_v12.docx")]
MEMO = "川崎町_県の第10期計画骨子と見える化の定めの確認_R8.10.23.md"

NO188 = [
    "188", "(5)(7)", "R8.11",
    "県へ推計データを提出する際の入力（第10期計画策定における取組の入力・必要入所（利用）定員、"
    "入居定員総数調査票）",
    "地域包括ケア「見える化」システムの将来推計機能では、都道府県へ推計データを提出する際に、"
    "地域分析結果の資料、「改正介護保険法を踏まえた第10期計画策定における取組の入力」（選択式）、"
    "「第10期必要入所（利用）定員、入居定員総数調査票」の登録を行うこととされています"
    "（利用マニュアル【システム操作編② 将来推計】第11.0版・令和8年9月30日に加わった画面）。"
    "取組の入力は、中長期のサービス見込量の推計、サービス提供体制の確保（中山間・人口減少地域の"
    "把握を含む）、医療・介護連携の推進について問うものです。素案の章節（9-1 中長期の見込み・"
    "1-6 日常生活圏域と7地区・7-2 供給の制約・1-5 医療計画との整合）から回答案を作成して"
    "お渡しできます。①操作編②のご提供（資料提供依頼No.51）、②調査票の様式のご提供（No.52）、"
    "③県への提出の時期、をご教示ください。\n⚠ 当方は操作編②の原本を受領しておらず、"
    "上の画面の構成は原本で確かめていません。",
    "見える化システムの県への提出（第2回目以降）\n計画素案 9-1（必要利用定員総数）",
    "発注者", "確認待ち", "R8.11",
    f"{HIZUKE}既定値＝操作編②のご提供を受けて設問を確かめ、素案から回答案を作成してお渡しする。"
    "調査票の受領までは素案 9-1 の必要利用定員総数（月27.2人）を正とする。"
    f"整理は 01_第10期_最新版成果品/{MEMO} §4。",
]

TSUIKI = {
    "50": (f"{HIZUKE}事業所分析で、利用者数が5年続けて前年を下回る指定は0件、介護職員"
           "（常勤換算）が令和3年→令和7年で2割以上減った指定は1件（短期入所療養介護 ▲25.0％。"
           "老健と同じ職員を指定ごとに報告しているもの）でした。また、宮城県のみやぎ高齢者元気"
           "プラン推進委員会（令和8年7月14日）の議事録に、町内の訪問介護・居宅介護支援の事業者の"
           "経営が厳しく、ほかに担い手がないとの委員の発言があります（素案には書かない）。　"),
    "69": (f"{HIZUKE}宮城県は第7期宮城県介護給付適正化取組方針（令和9〜11年度）について、"
           "国の基本方針に大きな見直しはないとして、第6期の重点（要介護認定の適正化／ケアプラン・"
           "住宅改修等の点検／縦覧点検・医療情報との突合）を継続する方向を示しました（令和8年"
           "8月26日 推進委員会 資料6）。協議用素案 Ver.2.29 の8-2 の注記に加えました。"
           "区分の確定は引き続き国の基本指針によります。　"),
    "77": (f"{HIZUKE}見える化システムに入力する地域支援事業費は「対象経費支出予定額」"
           "（総事業費から寄付金その他の収入額を差し引いた額のうち、交付金の対象経費として支出する"
           "見込みの額）とされています（地域分析の手引き 令和8年9月改訂版 79頁）。③の決算との差"
           "（12％）の一部は、この範囲の違いによる可能性があります。協議用素案 Ver.2.29 の9-1に"
           "注記しました。　"),
    "168": (f"{HIZUKE}入力値は「対象経費支出予定額」とされています（手引き79頁）。"
            "伸びの根拠のご確認の際、決算額ではなくこの範囲の額であるかもあわせてご確認ください。"
            "協議用素案 Ver.2.29 の9-1に注記しました。　"),
    "79": (f"{HIZUKE}手引き76頁により、令和8年度の実績見込みは、令和8年5月から保険者が設定した"
           "月までの令和8年度月報の累計値から算出されます。設定する月が少ないほど、見込みが少数の月の"
           "値に左右されます。新しい推計パターンでは、できるだけ新しい月まで設定いただくことを"
           "お勧めします（推奨の月数を定めた操作編②は未受領・資料提供依頼No.51）。　"),
    "80": (f"{HIZUKE}「令和8年5月ベース」の表示は、令和8年度の月報を5月の1か月だけ設定している"
           "ことによる可能性があります（手引き76頁）。「介護保険事業状況報告の設定」の画面の写しを"
           "いただければ確かめます。　"),
}

SHIRYO = [
    ["51", "F", "見える化 利用マニュアル【システム操作編② 将来推計】第11.0版",
     "令和8年度の実績見込みに用いる月・県への提出時の入力（取組の入力・調査票）・"
     "地域支援事業費の入力の定め",
     "確認事項No.79・80・188", "発注者（見える化システム「操作に困ったときは」）", "R8.11", "高",
     "未着手", f"手引き75・77頁の脚注によりシステム内からダウンロードできる。{MEMO}"],
    ["52", "D", "第10期必要入所（利用）定員、入居定員総数調査票（市区町村版）",
     "厚生労働省が別途配布する調査票の様式（県経由）", "素案9-1の必要利用定員総数。確認事項No.188",
     "発注者（宮城県長寿社会政策課経由）", "R8.11", "中", "未着手",
     f"見える化システムで県へ提出する際に登録する調査票。{MEMO}"],
]

JUNI = [
    ["45", "比較材料ブランチ（観点のみ）の確認と宮城県の第10期計画骨子（案）の反映", "R8.10.23",
     "ご指示による。観点は原典（宮城県 推進委員会資料・手引き）で確かめてから反映。",
     "No.50・69・77・79・80・168・188",
     "協議用素案 Ver.2.29／公表用 v2.29／資料 v15／問答集 v12", "半日", "完了",
     "1-2 に県の骨子（案）との対応表。8-2・9-1 に注記。資料・問答集は頁数（118頁）のみ。"
     "No.188 起票、資料提供依頼No.51・52。"],
    ["46", "概要版を協議用素案の最新版から組み直す", "", "概要版（R8.10.5）は素案 Ver.2.20 から"
     "組んだもので、Ver.2.21以降の改め（書体・事業所分析・県計画との対応ほか）が入っていない。",
     "No.187", "計画素案 概要版", "半日", "未着手",
     "build_gaiyo_R8.10.4.py／R8.10.5.py の SOAN を最新の協議用素案に上げ、数値を素案から読み直す。"],
]


def iro(c):
    c.font = FONT
    c.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    wb = openpyxl.load_workbook(WBS)
    kekka = []

    # ── 04 成果品管理
    ws = wb["04_成果品管理"]
    n = 0
    for x in range(5, ws.max_row + 1):
        no = str(ws.cell(x, 1).value or "").strip()
        na = str(ws.cell(x, 3).value or "")
        nb = na
        for a, b in OKIKAE:
            nb = nb.replace(a, b)
        if HIZUKE in str(ws.cell(x, 8).value or ""):
            continue
        moji = None
        if nb != na:
            ws.cell(x, 3).value = nb
            moji = (f"{HIZUKE}協議用素案を Ver.2.29（118頁）に上げ、公表用も組み直した（116頁）。"
                    "1-2 に宮城県の第10期みやぎ高齢者元気プラン骨子（案）との対応表を置き、"
                    "県の計画の名称を改めた。8-2（第7期県適正化取組方針）・9-1（地域支援事業費の"
                    "範囲）に注記。　") if no == "4" else (
                    f"{HIZUKE}資料を v15、想定問答集を v12 に上げた（素案の頁数118頁に揃えた"
                    "のみ。全68問）。　")
        elif no == "5":
            moji = (f"{HIZUKE}⚠ 概要版は素案 Ver.2.20 から組んだもので、Ver.2.21 以降の改めが"
                    "入っていない。送付の前に最新の協議用素案から組み直す（作業順位46）。　")
        if moji:
            ws.cell(x, 8).value = moji + str(ws.cell(x, 8).value or "")
            iro(ws.cell(x, 8))
            n += 1
    kekka.append((n in (0, 3), f"成果品 {n}件"))

    # ── 03 確認事項
    ws3 = wb["03_確認事項一覧"]
    k = 0
    for y in range(5, ws3.max_row + 1):
        no = str(ws3.cell(y, 1).value or "").strip()
        if no in TSUIKI and HIZUKE not in str(ws3.cell(y, 10).value or ""):
            ws3.cell(y, 10).value = TSUIKI[no] + str(ws3.cell(y, 10).value or "")
            iro(ws3.cell(y, 10))
            k += 1
    kekka.append((k in (0, len(TSUIKI)), f"確認事項への追記 {k}件"))
    nos = {str(ws3.cell(y, 1).value or "").strip() for y in range(5, ws3.max_row + 1)}
    if "188" not in nos:
        moto = ws3.max_row
        r = moto + 1
        for c in range(1, 11):
            ws3.cell(r, c)._style = copy.copy(ws3.cell(moto, c)._style)
            ws3.cell(r, c).value = NO188[c - 1]
        ws3.row_dimensions[r].height = ws3.row_dimensions[moto].height
    nos = [int(str(ws3.cell(y, 1).value).strip()) for y in range(5, ws3.max_row + 1)
           if str(ws3.cell(y, 1).value or "").strip().isdigit()]
    kekka.append((sorted(nos) == list(range(1, 189)), f"確認事項 1〜188・欠番なし（{len(nos)}件）"))

    # ── 06 資料提供依頼一覧
    ws6 = wb["06_資料提供依頼一覧"]
    gyo = {str(ws6.cell(x, 1).value or "").strip(): x for x in range(5, ws6.max_row + 1)}
    if "51" not in gyo:
        last = gyo["50"]
        shita = [[(ws6.cell(x, c).value, copy.copy(ws6.cell(x, c)._style)) for c in range(1, 11)]
                 for x in range(last + 1, ws6.max_row + 1)]
        for rng in list(ws6.merged_cells.ranges):
            if rng.min_row > last:
                ws6.unmerge_cells(str(rng))
        for x in range(last + 1, ws6.max_row + 1):
            for c in range(1, 11):
                ws6.cell(x, c).value = None
        x = last + 1
        for g in SHIRYO:
            for c in range(1, 11):
                ws6.cell(x, c).value = g[c - 1]
                ws6.cell(x, c)._style = copy.copy(ws6.cell(last, c)._style)
            ws6.row_dimensions[x].height = ws6.row_dimensions[last].height
            x += 1
        owari = x - 1
        for row in shita:
            for c, (v, st) in enumerate(row, 1):
                if isinstance(v, str) and v.startswith(f"=COUNTIF(I5:I{last}"):
                    v = v.replace(f"I5:I{last}", f"I5:I{owari}")
                ws6.cell(x, c).value = v
                ws6.cell(x, c)._style = st
            x += 1
        t = str(ws6.cell(1, 1).value or "")
        ws6.cell(1, 1).value = t.replace("（50件）", "（52件）")
    shiki = [ws6.cell(y, 2).value for y in range(5, ws6.max_row + 1)
             if isinstance(ws6.cell(y, 2).value, str) and ws6.cell(y, 2).value.startswith("=COUNTIF")]
    kekka.append(("（52件）" in str(ws6.cell(1, 1).value)
                  and all("I5:I56" in s for s in shiki) and len(shiki) == 5,
                  f"資料提供依頼 52件・式 {len(shiki)}件が I5:I56"))

    # ── 08 作業順位
    ws8 = wb["08_作業順位"]
    aru = {str(ws8.cell(x, 1).value or "").strip() for x in range(5, ws8.max_row + 1)}
    for g in JUNI:
        if g[0] in aru:
            continue
        moto = ws8.max_row
        r = moto + 1
        for c in range(1, 11):
            ws8.cell(r, c)._style = copy.copy(ws8.cell(moto, c)._style)
        ws8.row_dimensions[r].height = ws8.row_dimensions[moto].height
        for c, v in enumerate(g, 1):
            ws8.cell(r, c).value = v
    aru = {str(ws8.cell(x, 1).value or "").strip() for x in range(5, ws8.max_row + 1)}
    kekka.append(({"45", "46"} <= aru, "作業順位 45・46"))

    wb.save(WBS)
    print("業務工程管理表の更新（県の第10期計画の骨子（案）・見える化の定め）")
    ng = 0
    for ok, mi in kekka:
        print(("   ○ " if ok else "   × ") + mi)
        ng += not ok
    if ng:
        sys.exit(1)


if __name__ == "__main__":
    main()
