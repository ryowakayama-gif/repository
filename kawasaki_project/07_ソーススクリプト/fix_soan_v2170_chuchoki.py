# -*- coding: utf-8 -*-
"""計画素案 Ver.2.17（中長期の見込量の補完・作業順位29）

  入力　01_第10期_最新版成果品/川崎町_計画書素案_v2.16_内容整備版.docx
  出力　01_第10期_最新版成果品/川崎町_計画書素案_v2.17_中長期補完版.docx

══════ なぜ直すか
  素案 9-2 の総給付費は 令和17年度 1,110,612千円・令和22年度 1,218,805千円
  であるが、9-1 のサービス別の表の令和17年度の給付費を足すと
  1,015,734千円しかなく、**94,878千円（総給付費の8.5％）足りなかった。**
  Ver.2.14 で 9-1 に加えた15サービス（居宅4・介護予防11）について、
  中長期の列を「─」のまま空けたことによる、表の間の範囲の食い違いである。

⚠ **Ver.2.16 で置いた注の記述が誤っていた。**
  「令和17年度及び令和22年度は、サービスごとの長期推計を行っていないため
  「―」としています」と書いたが、**これらの値は既に算定されており、
  9-2 の総給付費には最初から入っていた。**
  「推計していない」のではなく「表に書き写していなかった」のである。
  本版で値を入れ、注を書き直す。

══════ 是正1　9-1 の「─」を埋める（15サービス×3項目）
  算定は 07_ソーススクリプト/calc_chuchoki_R8.10.17.py に置く（算定の正本）。
  式は第10期計画期間と同じ。
    見込量(年,要介護度) ＝ 令和7年度実績(要介護度)
                          × 認定者数(年,要介護度) ÷ 認定者数(令和8年度,要介護度)
  ⚠ **行の引当が正しいことは、算定した令和7年度・令和11年度の
    給付費・利用者数・回（日）数が、素案の表の値と一致することで
    確かめている**（37項目すべて一致。一致しなければ算定を返さない）。

══════ 是正2　介護予防の表の注を書き直す
══════ 是正3　9-1 に、サービス別の和と 9-2 の総給付費が合うことの注を置く

⚠ **保険料・給付費の総額は変えていない。** 9-2 の総給付費は元から
  15サービスを含んでいたため、**本版で総給付費は1円も動かない。**
  動くのは 9-1 の表の空欄が埋まることだけである。
"""
import copy
import importlib.util
import os
import re
import shutil
import sys
import zipfile

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from fix_soan_v111 import set_el, set_tc                     # noqa: E402

SRC = "01_第10期_最新版成果品/川崎町_計画書素案_v2.16_内容整備版.docx"
DST = "01_第10期_最新版成果品/川崎町_計画書素案_v2.17_中長期補完版.docx"


def load(name, fname):
    spec = importlib.util.spec_from_file_location(
        name, os.path.join(HERE, fname))
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


CC = load("calc_chuchoki_fix", "calc_chuchoki_R8.10.17.py")

# 是正2　介護予防の表の注（Ver.2.16 の誤った記述を書き直す）
CHU_YOBO_FURUI = (
    "⚠ 上表は介護予防給付のサービスの見込量です。介護保険法第117条第2項"
    "第1号は介護給付等対象サービスの種類ごとの量の見込みを定めることを"
    "求めているため、介護給付と分けて掲げています。令和17年度及び"
    "令和22年度は、サービスごとの長期推計を行っていないため"
    "「―」としています。")
CHU_YOBO_ATARASHII = (
    "⚠ 上表は介護予防給付のサービスの見込量です。介護保険法第117条第2項"
    "第1号は介護給付等対象サービスの種類ごとの量の見込みを定めることを"
    "求めているため、介護給付と分けて掲げています。令和17年度及び"
    "令和22年度は、第10期計画期間と同じ算定式（令和7年度実績に要介護度別の"
    "認定者数の伸びを乗じるもの）によります。")


def seiri(s):
    return (s or "").replace("\n", "").strip()


def hon(el):
    return "".join(x.text or "" for x in el.iter(qn("w:t")))


def mikomi_hyo(doc):
    """見込量の4表を見出しとの対応で拾う"""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    KUBUN = ("▌ 居宅サービスの見込量", "▌ 介護予防サービスの見込量",
             "▌ 地域密着型サービスの見込量", "▌ 施設サービスの見込量")
    out = {}
    midashi = None
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            s = seiri(Paragraph(el, doc).text)
            if s.startswith("▌") and "見込量" in s:
                midashi = s
        elif el.tag == qn("w:tbl"):
            if midashi in KUBUN and midashi not in out:
                out[midashi] = Table(el, doc)
            midashi = None
    return out


def main():
    if not os.path.exists(SRC):
        raise SystemExit("入力がない：" + SRC)

    # ══════ 算定（行の引当を素案で確かめてから使う）
    kekka = CC.santei()
    soan_g, soan_k = CC.yomu_soan_kyufu(SRC)
    mi = CC.kensho(kekka, soan_g, soan_k)
    if mi:
        raise SystemExit("素案に突合できないサービスがある：" + "／".join(mi))

    shutil.copy(SRC, DST)
    doc = docx.Document(DST)
    body = doc.element.body
    tbls = mikomi_hyo(doc)

    # ══════ 是正1　「─」を埋める
    #   項目ごとの入れ方　給付費＝千円の整数（カンマ区切り）
    #                     利用者数・回（日）数＝小数1位
    def moji(ko, v):
        if v is None:
            return "─"
        if ko == "給付費":
            return f"{int(v):,}"
        return f"{v:.1f}"

    n_cell = 0
    n_sv = 0
    nokori = []
    for na, rec in kekka.items():
        if na.startswith("__"):
            continue
        atta = False
        for mi_name, t in tbls.items():
            h = [seiri(c.text) for c in t.rows[0].cells]
            try:
                i17 = [i for i, x in enumerate(h) if "令和17年度" in x][0]
                i22 = [i for i, x in enumerate(h) if "令和22年度" in x][0]
            except IndexError:
                continue
            cur = ""
            for r in t.rows[1:]:
                cs = r.cells
                v0 = seiri(cs[0].text)
                if v0:
                    cur = v0
                if cur != na:
                    continue
                ko = seiri(cs[1].text)
                if ko == "給付費":
                    v17, v22 = rec["令和17年度"], rec["令和22年度"]
                elif ko == "利用者数":
                    v17, v22 = rec["令和17年度人数"], rec["令和22年度人数"]
                elif ko == "回（日）数":
                    v17, v22 = rec["令和17年度量"], rec["令和22年度量"]
                else:
                    continue
                # ⚠ 既に値が入っている欄は触らない（「─」だけを埋める）
                if seiri(cs[i17].text) not in ("─", "―", "-", ""):
                    continue
                set_tc(cs[i17]._tc, moji(ko, v17))
                set_tc(cs[i22]._tc, moji(ko, v22))
                n_cell += 2
                atta = True
        if atta:
            n_sv += 1
        else:
            nokori.append(na)
    if nokori:
        raise SystemExit("素案の表に見つからないサービス：" + "／".join(nokori))
    if n_sv != 15:
        raise SystemExit(f"埋めたサービスが{n_sv}件（15件のはず）")

    # ══════ 是正2　介護予防の表の注を書き直す
    n_chu = 0
    for el in body.iterchildren():
        if el.tag != qn("w:p"):
            continue
        s = hon(el)
        if CHU_YOBO_FURUI in s:
            set_el(el, s.replace(CHU_YOBO_FURUI, CHU_YOBO_ATARASHII))
            n_chu += 1
            break
    if n_chu == 0:
        raise SystemExit("介護予防の表の注（Ver.2.16 の記述）が見つからない")

    # ══════ 是正3　和が合うことの注を置く
    #   ⚠ 読者が 9-1 の4表を縦に足して 9-2 と突き合わせられるようにする。
    kisai17 = kisai22 = 0
    for t in tbls.values():
        h = [seiri(c.text) for c in t.rows[0].cells]
        try:
            i17 = [i for i, x in enumerate(h) if "令和17年度" in x][0]
            i22 = [i for i, x in enumerate(h) if "令和22年度" in x][0]
        except IndexError:
            continue
        for r in t.rows[1:]:
            c = [seiri(x.text) for x in r.cells]
            if c[1] != "給付費":
                continue
            for src, add in ((c[i17], "17"), (c[i22], "22")):
                m = re.match(r"^([\d,]+)$", src)
                if not m:
                    continue
                v = int(m.group(1).replace(",", ""))
                if add == "17":
                    kisai17 += v
                else:
                    kisai22 += v
    CHU_WA = (
        f"⚠ 上の4つの表の給付費を縦に足すと、令和17年度は"
        f"{kisai17:,}千円、令和22年度は{kisai22:,}千円となり、"
        f"第9章9-2 の総給付費と一致します"
        f"（千円未満の丸めにより数千円の差が生じます）。"
        f"令和17年度・令和22年度は、第10期計画期間（令和9〜11年度）の"
        f"見込量とは位置づけが異なり、計画の数値ではなく長期の見通しです。")

    # 雛形（⚠ の囲み）を探す
    tpl = None
    for el in body.iterchildren():
        if el.tag == qn("w:p") and hon(el).strip().startswith("⚠ 上表は介護予防"):
            tpl = copy.deepcopy(el)
            break
    if tpl is None:
        raise SystemExit("注の雛形が見つからない")
    # 施設サービスの表の後ろの「出典：」の段落の前に置く
    oku = None
    for el in body.iterchildren():
        if el.tag == qn("w:p") and hon(el).strip().startswith(
                "出典：令和7年度実績：介護保険事業状況報告"):
            oku = el
            break
    if oku is None:
        raise SystemExit("9-1 の出典の段落が見つからない")
    e = copy.deepcopy(tpl)
    set_el(e, CHU_WA)
    oku.addprevious(e)

    # ══════ 版表記
    n_ver = 0
    for el in body.iter(qn("w:t")):
        if el.text and "Ver.2.16" in el.text:
            el.text = el.text.replace("Ver.2.16", "Ver.2.17")
            n_ver += 1
    if n_ver == 0:
        raise SystemExit("本文に Ver.2.16 の表記が見つからない")
    for el in body.iterchildren():
        if el.tag == qn("w:p"):
            s = hon(el)
            if "（内容整備版）" in s:
                set_el(el, s.replace("（内容整備版）", "（中長期補完版）"))
                break
    doc.save(DST)
    naoshi_header(DST)

    # ══════════════════════════ 自己点検
    d2 = docx.Document(DST)
    d0 = docx.Document(SRC)
    zen = "".join(n.text or "" for n in d2.element.body.iter(qn("w:t")))
    ok, ng = [], []

    # 検査1　9-1 の給付費の和が 9-2 の総給付費と合う
    tb2 = mikomi_hyo(d2)
    wa = {"令和17年度": 0, "令和22年度": 0, "令和11年度": 0}
    for t in tb2.values():
        h = [seiri(c.text) for c in t.rows[0].cells]
        idx = {}
        for y in wa:
            hit = [i for i, x in enumerate(h) if y in x]
            if hit:
                idx[y] = hit[0]
        for r in t.rows[1:]:
            c = [seiri(x.text) for x in r.cells]
            if c[1] != "給付費":
                continue
            for y, i in idx.items():
                m = re.match(r"^([\d,]+)$", c[i])
                if m:
                    wa[y] += int(m.group(1).replace(",", ""))
    # 9-2 の総給付費
    soukei = {}
    for t in d2.tables:
        h = [seiri(c.text) for c in t.rows[0].cells]
        if h and h[0] == "区分" and any("令和17年度" in x for x in h):
            for r in t.rows[1:]:
                c = [seiri(x.text) for x in r.cells]
                if c[0] == "総給付費":
                    for i, x in enumerate(h):
                        m = re.search(r"([\d,]+)", c[i]) if i else None
                        if m and ("令和" in x):
                            soukei[x.replace("実績", "").strip()] = int(
                                m.group(1).replace(",", ""))
    for y in ("令和11年度", "令和17年度", "令和22年度"):
        s = soukei.get(y)
        if s is None:
            ng.append(f"検査1　9-2 に{y}の総給付費がない")
        elif abs(wa[y] - s) <= 3:
            ok.append(f"検査1　{y}　9-1 の4表の給付費の和{wa[y]:,}千円"
                      f"＝9-2 の総給付費{s:,}千円（差{wa[y] - s:+,}千円・"
                      f"千円未満の丸めによる）")
        else:
            ng.append(f"検査1　{y}　9-1 の和{wa[y]:,}千円と"
                      f"9-2 の総給付費{s:,}千円が合わない"
                      f"（差{wa[y] - s:+,}千円）")

    # 検査2　誤った注が消えていること
    if "サービスごとの長期推計を行っていない" in zen:
        ng.append("検査2　Ver.2.16 の誤った注が残っている")
    else:
        ok.append("検査2　「サービスごとの長期推計を行っていない」という"
                  "誤った記述は消えている")

    # 検査3　埋めた15サービスの値が算定と一致すること
    chigai = []
    for na, rec in kekka.items():
        if na.startswith("__"):
            continue
        for t in tb2.values():
            h = [seiri(c.text) for c in t.rows[0].cells]
            hit = [i for i, x in enumerate(h) if "令和17年度" in x]
            if not hit:
                continue
            i17 = hit[0]
            cur = ""
            for r in t.rows[1:]:
                c = [seiri(x.text) for x in r.cells]
                if c[0]:
                    cur = c[0]
                if cur != na or c[1] != "給付費":
                    continue
                m = re.match(r"^([\d,]+)$", c[i17])
                if not m:
                    chigai.append(f"{na} が埋まっていない")
                elif int(m.group(1).replace(",", "")) != rec["令和17年度"]:
                    chigai.append(f"{na} 素案{c[i17]}／算定{rec['令和17年度']}")
    if chigai:
        ng.append("検査3　" + "／".join(chigai[:6]))
    else:
        ok.append(f"検査3　埋めた15サービスの令和17年度の給付費は"
                  f"算定と一致する")

    # 検査4　総給付費・保険料が動いていないこと
    for v in ("1,110,612", "1,218,805", "998,222", "3,313,925", "6,822",
              "9,743.1", "152,500"):
        if v not in zen:
            ng.append(f"検査4　主な数値が消えている：{v}")
    if not any("検査4" in x for x in ng):
        ok.append("検査4　総給付費・標準給付費見込額・保険料の数値は"
                  "動いていない")

    # 検査5　表の数・段落の数
    if len(d2.tables) != len(d0.tables):
        ng.append(f"検査5　表の数が{len(d0.tables)}→{len(d2.tables)}")
    elif len(d2.paragraphs) != len(d0.paragraphs) + 1:
        ng.append(f"検査5　段落が{len(d0.paragraphs)}→{len(d2.paragraphs)}"
                  "（1つだけ増えるはず）")
    else:
        ok.append(f"検査5　表 {len(d2.tables)}（変わらず）／"
                  f"段落 {len(d2.paragraphs)}（注を1つ加えた）")

    # 検査6　版表記
    z = zipfile.ZipFile(DST)
    hdr = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>",
                             z.read("word/header1.xml").decode("utf-8")))
    hoka = sorted({v for v in re.findall(r"Ver\.\d+\.\d+", zen)
                   if v != "Ver.2.17"})
    if "Ver.2.17" not in hdr:
        ng.append(f"検査6　ヘッダーの版表記が更新されていない：{hdr}")
    elif hoka:
        ng.append("検査6　現行の版以外の版表記が残っている："
                  + "／".join(hoka))
    else:
        ok.append("検査6　版表記は Ver.2.17 のみ")

    print("計画素案 Ver.2.17（中長期補完版）")
    print("保存：", DST)
    print(f"  段落 {len(d2.paragraphs)}／表 {len(d2.tables)}")
    print(f"  是正1　9-1 の「─」を埋めた　{n_sv}サービス・{n_cell}セル")
    print(f"  是正2　介護予防の表の注を書き直した（Ver.2.16 の誤り）")
    print(f"  是正3　和が合うことの注を1か所置いた")
    print(f"  版表記 {n_ver}か所")
    print("  ── 自己点検")
    for s in ok:
        print("   ○", s)
    if ng:
        for s in ng:
            print("   ×", s)
        sys.exit(1)


def naoshi_header(path):
    z = zipfile.ZipFile(path)
    names = z.namelist()
    data = {n: z.read(n) for n in names}
    z.close()
    n_fix = 0
    for n in names:
        if not re.match(r"word/(header|footer)\d*\.xml", n):
            continue
        s = data[n].decode("utf-8")
        s2 = re.sub(r"Ver\.2\.16", "Ver.2.17", s)
        if s2 != s:
            data[n] = s2.encode("utf-8")
            n_fix += 1
    if n_fix == 0:
        raise SystemExit("ヘッダーの版表記（Ver.2.16）が見つからない")
    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
    shutil.move(tmp, path)


if __name__ == "__main__":
    main()
