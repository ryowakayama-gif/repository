# -*- coding: utf-8 -*-
"""第2回策定委員会資料 v19 → v20（第3回レビューの影響範囲の再点検による是正）

  python3 07_ソーススクリプト/build_2kai_v20.py

内容の正本　協議用素案 Ver.2.35（fix_soan_v2350_eikyou.py）

  R3-06 の波及　表26 の行の名「０％（令和7年度の単価で固定・本算定）」の「本算定」は、確定の算定と読める。
               素案 9-3 の「本試算」に揃える
"""
import importlib.util
import os
import shutil
import sys

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)


def load(name, fn):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(spec)
    sys.modules[name] = m
    spec.loader.exec_module(m)
    return m


V9 = load("build_2kai_v9", "build_2kai_v9.py")
SRC = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v19.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v20.docx"
SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.35_協議用.docx"


def zen(d):
    return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会資料 v19 → v20（第3回レビューの影響範囲の再点検）")
    sz = zen(docx.Document(SOAN))
    judge("本算定" not in sz and "本試算" in sz, "素案 v2.35 は「本試算」")
    shutil.copy2(SRC, DST)
    d = docx.Document(DST)
    mae = zen(d)
    n = 0
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                if "本算定" in c.text:
                    V9.set_cell(c, c.text.replace("本算定", "本試算"))
                    n += 1
    judge(n == 1, f"表26 の行の名（{n}か所）")
    d.save(DST)
    d2 = docx.Document(DST)
    ato = zen(d2)
    print("  ── 検査")
    judge("本算定" not in ato, "「本算定」0件")
    judge(len(ato) == len(mae), f"字数は変わらない（{len(mae):,}→{len(ato):,}）")
    kids = list(d2.element.body.iterchildren())
    judge(not any(a_.tag == qn("w:tbl") and b_.tag == qn("w:tbl") for a_, b_ in zip(kids, kids[1:])),
          "隣り合う表 0件（v18 の区切りを保つ）")
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 不適合0件")


if __name__ == "__main__":
    main()
