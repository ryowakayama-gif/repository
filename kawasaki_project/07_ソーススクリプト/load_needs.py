# -*- coding: utf-8 -*-
"""ニーズ調査の個票を読む。設問ブロックは n 列で区切る。"""
import openpyxl

SRC = ("/home/user/repository/kawasaki_project/04_調査・入力・分析/"
       "R8調査データ/川崎町_介護予防日常生活圏域ニーズ調査_集計データ_20260807.xlsx")
MARK = ("〇", "○", "◯", "1", 1)


def load():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["データ"]
    NC, NR = ws.max_column, ws.max_row
    blocks = []
    cur = None
    gname = None
    for c in range(22, NC + 1):
        if ws.cell(1, c).value:
            gname = str(ws.cell(1, c).value).strip()
        k4 = ws.cell(4, c).value
        if k4 is None:
            continue
        k4 = str(k4).strip()
        if k4 == "n":
            if cur:
                cur["ncol"] = c
                blocks.append(cur)
                cur = None
            continue
        if k4 == "自由記述":
            continue
        if cur is None:
            cur = dict(group=gname, text=str(ws.cell(2, c).value or "").strip(),
                       opts=[], cols=[], ncol=None)
        if ws.cell(2, c).value and not cur["opts"]:
            cur["text"] = str(ws.cell(2, c).value).strip()
        cur["opts"].append(str(ws.cell(3, c).value or "").strip())
        cur["cols"].append(c)
    if cur:
        blocks.append(cur)

    rows = [r for r in range(5, NR + 1)
            if ws.cell(r, 1).value is not None
            and str(ws.cell(r, 1).value).strip() not in ("計", "合計")]

    recs = []
    for r in rows:
        a = {}
        for i, b in enumerate(blocks):
            sel = [o for o, c in zip(b["opts"], b["cols"])
                   if ws.cell(r, c).value in MARK]
            a[i] = sel
        recs.append(a)
    return blocks, recs


if __name__ == "__main__":
    b, recs = load()
    print("個票", len(recs), "件／設問ブロック", len(b))
    for i, q in enumerate(b):
        print(f"{i:3d} {q['group'] or '':6s} {q['text'][:44]}  選択肢{len(q['opts'])}")
