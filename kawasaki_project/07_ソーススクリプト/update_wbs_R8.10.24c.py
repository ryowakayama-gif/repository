# -*- coding: utf-8 -*-
"""業務工程管理表の更新（Red Team 再レビュー第3回（令和8年10月24日）への対応）

  python3 07_ソーススクリプト/update_wbs_R8.10.24c.py

  00_管理方法　　　基準日の記載を第3回対応後に（件数は台帳から数える）
  03_確認事項一覧　No.49（回数の契約上の手続）・No.83（頁数の対象範囲）・No.109（KPI の分母）・
                   No.157（推計値と実数の差の読み方）に追記
  04_成果品管理　　稿を 素案 Ver.2.34／公表用・概要版 v2.34／資料 v19／問答集 v16 に
  08_作業順位　　　51（完了）を加える
再実行しても二重に書かない。
"""
import collections
import copy
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月24日 更新③（Red Team 再レビュー第3回への対応）】"
TAIOU = "川崎町_RedTeam再レビュー第3回対応表_R8.10.24.md"
ZIP = "11_提出一式/川崎町_RedTeam再レビュー第3回対応_成果物とエビデンス_R8.10.24.zip"
OKIKAE = [("川崎町_計画書素案_v2.33_協議用.docx", "川崎町_計画書素案_v2.34_協議用.docx"),
          ("川崎町_計画書案_公表用_v2.33.docx", "川崎町_計画書案_公表用_v2.34.docx"),
          ("川崎町_計画書案_概要版_v2.33.docx", "川崎町_計画書案_概要版_v2.34.docx"),
          ("川崎町_第2回策定委員会資料_R8.11_v18.docx", "川崎町_第2回策定委員会資料_R8.11_v19.docx"),
          ("同_想定問答集_R8.11_v15.docx", "同_想定問答集_R8.11_v16.docx")]
TSUIKI = {
    "49": (f"{HIZUKE}全3回とすることは委託仕様書6(3)の「4回開催の予定」の変更に当たるため、仕様書11（定めのない事項の"
           "協議）を根拠とせず、町の契約担当で契約上の手続（変更契約又は書面による合意の要否）をご確認いただき、"
           "町が書面で決定する形に資料 v19・問答集 v16（QB-58）を改めた。委員会には工程の見通しについてご意見をいただく。"
           "既定値＝第2回の前までに手続の判断がない場合、資料は案Ⅰを事務局案として示したまま、回数は未決定と扱う。　"),
    "83": (f"{HIZUKE}Red Team 再レビュー第3回 R3-11。素案は本文118頁・資料編（見込）13頁の計131頁。仕様書8の"
           "「100頁程度」の対象範囲（本文のみか、資料編を含むか。概要版は別）を町と書面で確かめたい。　"),
    "109": (f"{HIZUKE}R3-08。通いの場の参加割合15.9％（n=439）は一般高齢者調査の回答者のうちの割合であるため、KPI の"
            "名称に分母を書いた。行政実績の参加率（参加者数÷65歳以上人口）は、地域支援事業のデータの受領後に別の"
            "指標として置く。　"),
    "157": (f"{HIZUKE}R3-02。推計値（374人）と自立度Ⅱ以上の実数は把握の方法と対象が異なるため、単純な差を"
            "要介護認定を受けていない認知症の方の数とは読まない、と素案 6-1 を改めた（Ver.2.34）。"
            "内訳の和372人と計374人の差（四捨五入）も注に書いた。　"),
}


def iro(c):
    c.font = FONT
    c.alignment = Alignment(wrap_text=True, vertical="top")


def mae_ni(c, s):
    if HIZUKE in str(c.value or ""):
        return 0
    c.value = s + str(c.value or "")
    iro(c)
    return 1


def main():
    wb = openpyxl.load_workbook(WBS)
    kekka = []
    ws3 = wb["03_確認事項一覧"]
    k = 0
    for y in range(5, ws3.max_row + 1):
        no = str(ws3.cell(y, 1).value or "").strip()
        if no in TSUIKI:
            k += mae_ni(ws3.cell(y, 10), TSUIKI[no])
    kekka.append((k in (0, len(TSUIKI)), f"確認事項への追記 {k}件"))
    k3 = collections.Counter(str(ws3.cell(y, 8).value).strip() for y in range(5, ws3.max_row + 1)
                             if str(ws3.cell(y, 1).value or "").strip().isdigit())

    ws0 = wb["00_管理方法"]
    for r in ws0.iter_rows():
        for c in r:
            v = str(c.value or "")
            if v.startswith("令和8年10月24日　本表の記載は同日時点のもの。Red Team 再レビュー（同日）を受け"):
                c.value = ("令和8年10月24日　本表の記載は同日時点のもの。Red Team 再レビュー第3回（同日）を受け、"
                           "協議用素案 Ver.2.34・公表用の計画書（案）v2.34・概要版 v2.34・第2回策定委員会資料 v19・"
                           "想定問答集 v16（全69問）に是正した。第3回の指摘11件のうち是正済8・一部是正2・町のご判断1。"
                           "令和7年度の給付費37サービスを年報の円単位で突合（37件一致）。点検ログ14件はすべて終了コード0。"
                           f"確認事項は{sum(k3.values())}件のうち{k3['完了']}件が完了。一式は {ZIP}。\n"
                           "（前の記載）" + v)
                iro(c)
                kekka.append((True, "基準日の記載を第3回対応後に"))

    ws4 = wb["04_成果品管理"]
    n = 0
    for y in range(5, ws4.max_row + 1):
        na = str(ws4.cell(y, 3).value or "")
        nb = na
        for a, b in OKIKAE:
            nb = nb.replace(a, b)
        if nb != na:
            ws4.cell(y, 3).value = nb
            mae_ni(ws4.cell(y, 8), f"{HIZUKE}Red Team 再レビュー第3回の指摘を是正した版に上げた（{TAIOU}）。　")
            n += 1
    kekka.append((n in (0, 3), f"成果品の稿 {n}行"))

    ws8 = wb["08_作業順位"]
    aru = {str(ws8.cell(x, 1).value or "").strip() for x in range(5, ws8.max_row + 1)}
    if "51" not in aru:
        moto = ws8.max_row
        r = moto + 1
        for c in range(1, 11):
            ws8.cell(r, c)._style = copy.copy(ws8.cell(moto, c)._style)
        ws8.row_dimensions[r].height = ws8.row_dimensions[moto].height
        for c, v in enumerate([
                "51", "Red Team 再レビュー第3回（令和8年10月24日）への対応と成果物・エビデンスの一式", "R8.10.24",
                "ご指示による。第3回の指摘11件（R3-01〜R3-11）と先行指摘の再判定を最新の版で確かめ、是正した",
                "No.49・83・109・157", "素案 Ver.2.34／公表用・概要版 v2.34／資料 v19／問答集 v16／ZIP",
                "1日", "完了",
                "施設への移行・未認定の認知症の方の数の断定を改め、法令の項、回数の決め方、保険料が試算であること、"
                "通いの場の分母を直した。令和7年度の給付費37サービスを年報の円単位で突合し、998,224千円と998,222千円の差が"
                "丸めであることを確かめた。公表用の留保の台帳に確認先・期限・既定の扱いの列を置いた（処理は町）。"], 1):
            ws8.cell(r, c).value = v
    aru = {str(ws8.cell(x, 1).value or "").strip() for x in range(5, ws8.max_row + 1)}
    kekka.append(("51" in aru, "作業順位 51"))

    wb.save(WBS)
    print("業務工程管理表の更新（Red Team 再レビュー第3回への対応）")
    ng = 0
    for ok, mi in kekka:
        print(("   ○ " if ok else "   × ") + mi)
        ng += not ok
    if ng:
        sys.exit(1)


if __name__ == "__main__":
    main()
