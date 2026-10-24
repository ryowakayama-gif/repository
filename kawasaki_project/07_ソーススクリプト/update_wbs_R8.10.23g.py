# -*- coding: utf-8 -*-
"""業務工程管理表の更新（概要版を最新の協議用素案から組み直し・素案 Ver.2.30）
（令和8年10月23日・同日7度目）

  python3 07_ソーススクリプト/update_wbs_R8.10.23g.py

  00_管理方法　　　　基準日を令和8年10月23日に進める（件数は台帳から数える）
  01_業務内容別の進捗　(7)(8) に追記。(8) の進捗 0.60 → 0.65
  02_月別工程管理　　令和8年10月に追記
  04_成果品管理　　　計画書・概要版の稿を改める。概要版の進捗 0.75 → 0.85
  08_作業順位　　　　46 を完了に。47（素案の後期高齢者数の是正）を加える

⚠ 進捗は作業項目のうち完了したものの割合を受託者が判断して記入する（01 注1）。
再実行しても二重に書かない。
"""
import collections
import copy
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月23日 更新⑦（概要版の組み直し）】"
SOAN = "川崎町_計画書素案_v2.30_協議用.docx"
KOUHYOU = "川崎町_計画書案_公表用_v2.30.docx"
GAIYO = "川崎町_計画書案_概要版_v2.30.docx"


def iro(c):
    c.font = FONT
    c.alignment = Alignment(wrap_text=True, vertical="top")


def mae_ni(c, s):
    if HIZUKE in str(c.value or ""):
        return 0
    c.value = s + str(c.value or "")
    iro(c)
    return 1


def main():
    wb = openpyxl.load_workbook(WBS)
    kekka = []

    # ── 件数（台帳から）
    ws3 = wb["03_確認事項一覧"]
    k3 = collections.Counter(str(ws3.cell(y, 8).value).strip() for y in range(5, ws3.max_row + 1)
                             if str(ws3.cell(y, 1).value or "").strip().isdigit())
    ws6 = wb["06_資料提供依頼一覧"]
    k6 = collections.Counter(str(ws6.cell(y, 9).value).strip() for y in range(5, ws6.max_row + 1)
                             if str(ws6.cell(y, 1).value or "").strip().isdigit())
    n3, n6 = sum(k3.values()), sum(k6.values())

    # ── 00 基準日
    ws0 = wb["00_管理方法"]
    for r in ws0.iter_rows():
        for c in r:
            v = str(c.value or "")
            if v.startswith("令和8年10月4日　本表の記載は同日時点のもの。"):
                c.value = (
                    "令和8年10月23日　本表の記載は同日時点のもの。協議用素案 Ver.2.30（118頁）と、"
                    "公表用の計画書（案）v2.30（116頁）・概要版 v2.30（8頁・図8点・表6件）、"
                    "第2回策定委員会資料 v15・想定問答集 v12（全68問）まで作成済み。点検は、素案の"
                    "自己点検32件・保険料の検算40項目・書体の点検16件・図表の自己点検10件・委員会資料の"
                    "自己点検27件・概要版の自己点検（検査1〜8）がいずれも適合（8軸の網羅性の未実施10件は"
                    f"別に数えている）。確認事項は{n3}件のうち{k3['完了']}件が完了、"
                    f"{n3 - k3['完了']}件が未完了（確認待ち{k3['確認待ち']}件ほか）。資料提供依頼は"
                    f"{n6}件のうち{k6['完了']}件を受領。\n（前の基準日 令和8年10月4日の記載）" + v)
                iro(c)
                kekka.append((True, "基準日を令和8年10月23日に"))

    # ── 01 業務内容別の進捗
    ws1 = wb["01_業務内容別の進捗"]
    rows = {str(ws1.cell(y, 1).value or "").strip(): y for y in range(5, ws1.max_row + 1)}
    n = mae_ni(ws1.cell(rows["(7)"], 9),
               f"{HIZUKE}協議用素案を Ver.2.30 に上げた。概要版を組み直す際に、3-4 の課題①と 5-4 の"
               "後期高齢者数が令和7年6月の月報の値（1,675人・＋10.9％）のまま残り、2-1（令和7年度末"
               "1,742人・＋15.4％）と食い違っていたことが分かり、2-1 に揃えた。　")
    if n and ws1.cell(rows["(8)"], 7).value in (0.6, "0.6"):
        ws1.cell(rows["(8)"], 7).value = 0.65
    n += mae_ni(ws1.cell(rows["(8)"], 9),
                f"{HIZUKE}③概要版を最新の協議用素案（Ver.2.30）から組み直した（{GAIYO}・8頁）。"
                "数値・名称はすべて素案の docx から読み、素案にない数・未確定の値（［要確認］等）は"
                "載せない。従前の版で4つの重点の名称が素案と違っていたこと、ヘッダーに版が出ていた"
                "ことを是正。書式は公表用（HG丸ｺﾞｼｯｸM-PRO）。残るのは保険料の確定後の差替え・"
                "町のご確認・入稿（No.18）。進捗 0.60→0.65。　")
    kekka.append((ws1.cell(rows["(8)"], 7).value == 0.65, f"(7)(8) に追記（{n}件）・(8) 0.65"))

    # ── 02 月別工程
    ws2 = wb["02_月別工程管理"]
    for y in range(1, ws2.max_row + 1):
        if str(ws2.cell(y, 1).value or "").strip() == "令和8年10月":
            mae_ni(ws2.cell(y, 5),
                   f"{HIZUKE}協議用素案 Ver.2.30・公表用 v2.30・概要版 v2.30（8頁）まで作成。"
                   "第2回策定委員会資料 v15・想定問答集 v12。　")
            kekka.append((True, "月別工程 令和8年10月に追記"))

    # ── 04 成果品
    ws4 = wb["04_成果品管理"]
    for y in range(5, ws4.max_row + 1):
        no = str(ws4.cell(y, 1).value or "").strip()
        if no == "4" and HIZUKE not in str(ws4.cell(y, 8).value or ""):
            ws4.cell(y, 3).value = (str(ws4.cell(y, 3).value)
                                    .replace("川崎町_計画書素案_v2.29_協議用.docx", SOAN)
                                    .replace("川崎町_計画書案_公表用_v2.29.docx", KOUHYOU))
            mae_ni(ws4.cell(y, 8), f"{HIZUKE}協議用素案 Ver.2.30（3-4・5-4 の後期高齢者数を 2-1 に"
                                   "揃えた）。頁数は118頁・公表用116頁で変わらない。　")
        if no == "5" and HIZUKE not in str(ws4.cell(y, 8).value or ""):
            ws4.cell(y, 3).value = f"{GAIYO}（8頁・3,752字・表6件・図8点）"
            ws4.cell(y, 6).value = 0.85
            mae_ni(ws4.cell(y, 8),
                   f"{HIZUKE}最新の協議用素案 Ver.2.30 から組み直した（build_gaiyo_v230.py）。"
                   "数値・名称は素案の docx から読む（概要版に数値を持たない）。素案にない数（全国・県の"
                   "高齢化率の本文の記載）と未確定のもの（サポーターの累計・チームオレンジの整備年度）を"
                   "落とした。4つの重点・7つの柱は素案 4-2・4-3 と同じ名称。保険料は確定前である旨を明記。"
                   "残るのは保険料の確定後の差替え・町のご確認・入稿（No.18）。旧稿（R8.10.5）は履歴として残す。　")
    kekka.append((any(str(ws4.cell(y, 1).value).strip() == "5" and ws4.cell(y, 6).value == 0.85
                      for y in range(5, ws4.max_row + 1)), "成果品 計画書・概要版"))

    # ── 08 作業順位
    ws8 = wb["08_作業順位"]
    aru = {}
    for y in range(5, ws8.max_row + 1):
        aru[str(ws8.cell(y, 1).value or "").strip()] = y
    if "46" in aru:
        y = aru["46"]
        ws8.cell(y, 3).value = "R8.10.23"
        ws8.cell(y, 6).value = f"{GAIYO}（8頁）"
        ws8.cell(y, 8).value = "完了"
        if HIZUKE not in str(ws8.cell(y, 9).value or ""):
            ws8.cell(y, 9).value = (f"{HIZUKE}build_gaiyo_v230.py を新設。図8点（柱と重点の図は素案の"
                                    "名称で作り直し）。自己点検 検査1〜8 適合・docx 検証適合。")
    if "47" not in aru:
        moto = ws8.max_row
        r = moto + 1
        for c in range(1, 11):
            ws8.cell(r, c)._style = copy.copy(ws8.cell(moto, c)._style)
        ws8.row_dimensions[r].height = ws8.row_dimensions[moto].height
        for c, v in enumerate([
                "47", "協議用素案の後期高齢者数の古い値の是正（3-4・5-4）", "R8.10.23",
                "概要版の組み直しで発見。2-1（令和7年度末1,742人）と3-4・5-4（令和7年6月1,675人）が食い違い",
                "―", f"協議用素案 Ver.2.30／公表用 v2.30", "1時間", "完了",
                "fix_soan_v2300_kouki.py。値は 2-1 の表から読む。頁数・未確定の印は変わらない。"
                "委員会資料・問答集はこの値を載せておらず改めない。"], 1):
            ws8.cell(r, c).value = v
    aru = {str(ws8.cell(y, 1).value or "").strip(): y for y in range(5, ws8.max_row + 1)}
    kekka.append((ws8.cell(aru["46"], 8).value == "完了" and "47" in aru, "作業順位 46 完了・47"))

    wb.save(WBS)
    print("業務工程管理表の更新（概要版の組み直し）")
    ng = 0
    for ok, mi in kekka:
        print(("   ○ " if ok else "   × ") + mi)
        ng += not ok
    print(f"  確認事項 {n3}件：" + "／".join(f"{k}{v}" for k, v in k3.most_common()))
    print(f"  資料提供依頼 {n6}件：" + "／".join(f"{k}{v}" for k, v in k6.most_common()))
    if ng:
        sys.exit(1)


if __name__ == "__main__":
    main()
