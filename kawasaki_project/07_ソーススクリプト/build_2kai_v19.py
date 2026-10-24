# -*- coding: utf-8 -*-
"""第2回策定委員会資料 v18 → v19（Red Team 再レビュー第3回（令和8年10月24日）の指摘の是正）

  python3 07_ソーススクリプト/build_2kai_v19.py

再レビュー原本　09_元資料/レビュー/川崎町_第10期計画_RedTeam再レビュー第3回_R8.10.24受領.md
内容の正本　　　協議用素案 Ver.2.34（fix_soan_v2340_r3review.py）

  R3-01 ８「在宅から施設・居住系への移行が進んでいます」の断定 → 乖離の事実と、原因が未確定であること
  R3-04 表（第10章の行）「第13条第3項が準用する第12条第4〜7項」→ 第12条第3項から第7項まで
        （意見聴取は第12条第3項）
  R3-05 開催回数を全３回とする根拠に委託仕様書11（定めのない事項の協議）を援用していた。６(3)に定めのある
        事項の変更であるため、契約上の手続（変更契約又は書面による合意の要否）を町の契約担当で確かめ、
        町が書面で決める形に改める。委員会にはご意見をいただく
  R3-06 「第１０期の算定（令和７年度の年報による確定値）」→ 現時点の試算（未確定のものを明記）
  R3-07 頁数の見通しの囲み（仕様書８の「100頁程度」）が20頁と21頁に分かれる → 行を分割しない
"""
import importlib.util
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)


def load(name, fn):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(spec)
    sys.modules[name] = m
    spec.loader.exec_module(m)
    return m


V9 = load("build_2kai_v9", "build_2kai_v9.py")
SRC = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v18.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v19.docx"
SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.34_協議用.docx"


def kakikae(p, s):
    p.runs[0].text = V9.clean(s)
    for r in p.runs[1:]:
        r.text = ""


def zen(d):
    return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])


def main():
    futeki = []
    KYU = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会資料 v18 → v19（Red Team 再レビュー第3回の指摘の是正）")
    sz = zen(docx.Document(SOAN))
    judge("在宅から施設等への移行の実態及び原因は確かめられていません" in sz, "素案 3-3 は v2.34 の言い回し")
    shutil.copy2(SRC, DST)
    d = docx.Document(DST)
    mae = zen(d)

    def p_rep(kyu, shin, n=1):
        h = [p for p in d.paragraphs if kyu in p.text]
        judge(len(h) == n, f"本文「{kyu[:24]}」（{len(h)}か所）")
        for p in h:
            KYU.append(p.text)
            kakikae(p, p.text.replace(kyu, shin))

    def c_rep(kyu, shin, n=1):
        k = 0
        for t in d.tables:
            seen = []
            for r in t.rows:
                for c in r.cells:
                    if c._tc in seen:
                        continue
                    seen.append(c._tc)
                    if kyu in c.text:
                        KYU.append(c.text)
                        V9.set_cell(c, c.text.replace(kyu, shin))
                        k += 1
        judge(k == n, f"表「{kyu[:24]}」（{k}か所）")

    print("  ── R3-01 居宅・施設の乖離の読み方")
    p_rep("総額が一致していることの内側で、在宅から施設・居住系への移行が進んでいます。居宅サービスが計画を下回ったのは、"
          "町内の在宅サービスの供給余力が乏しいこと（訪問介護の従業者は常勤換算4.5人）と、住所地特例24人（令和８年３月末）"
          "に表れる町外施設の利用の常態化が同時に進んだことが影響している可能性があります（因果関係は、町内外別の利用の"
          "推移等による追加の検証を要します）。第１０期のサービス見込量は、この移行を前提に置くかどうかで変わります。",
          "総額は計画に近い一方、居宅サービスは計画を下回り、施設サービスは上回りました。この乖離には、町内の在宅"
          "サービスの供給余力が乏しいこと（訪問介護の従業者は常勤換算4.5人）と、住所地特例24人（令和８年３月末）に"
          "表れる町外施設の利用の常態化が影響した可能性がありますが、在宅から施設等への移行の実態及び原因は"
          "確かめられていません（町内外別の利用者数や経年の推移等による追加の検証を要します）。第１０期のサービス"
          "見込量は、令和７年度の利用実績を出発点としているため、この利用の構成が続くことを前提としています。")

    print("  ── R3-04 認知症基本法の条項")
    c_rep("認知症基本法第13条第3項が準用する第12条第4〜7項", "認知症基本法第13条第3項が準用する第12条第3項から第7項まで")

    print("  ── R3-05 策定委員会の回数の決め方")
    c_rep("全４回とする場合は納期に間に合わないため、回数を全３回とすることについて、委託仕様書11（本仕様の定めのない"
          "事項は協議の上実施する）によりご協議をお願いします。",
          "全４回とする場合は納期に間に合いません。回数を全３回とすることは、委託仕様書６(3)に定められた予定の変更に"
          "当たるため、町の契約担当で契約上の手続（変更契約又は書面による合意の要否）をご確認いただき、町が書面で"
          "決定します。本日は、工程の見通しについてご意見をいただきます。")
    p_rep("回数のご判断をお願いします。",
          "回数は、契約上の手続を確かめたうえで町が書面で決定します。本日は工程の見通しについてご意見をお願いします。")
    p_rep("⑦　策定委員会の回数（全３回とするか全４回とするか）と、パブリックコメントの開始日のご選択（８）。納期に直接に効きます",
          "⑦　パブリックコメントの開始日のご選択と、策定委員会の回数（全３回とするか全４回とするか）についてのご意見（８）。"
          "回数は契約上の手続を確かめたうえで町が書面で決定します")

    print("  ── R3-06 保険料の算定の見出し（試算であること）")
    p_rep("◆ 第１０期の算定（令和７年度の年報による確定値）",
          "◆ 第１０期保険料の現時点の試算（令和７年度の介護保険事業状況報告等を基礎に算定。基金取崩額・令和９年度の"
          "報酬改定等は未確定）")
    p_rep("このうち給付費の水準は、令和７年度の年報を基礎とする算定により確定しました。",
          "このうち給付費の水準は、令和７年度の年報（確定値）を基礎として算定しました。第１０期の保険料は、下の要因が"
          "定まるまでは試算です。")

    print("  ── R3-07 頁数の見通しの囲みを頁で分けない")
    n = 0
    for tr in d.element.body.iter(qn("w:tr")):
        tx = "".join(x.text or "" for x in tr.iter(qn("w:t")))
        if len(tr.findall(qn("w:tc"))) == 1 and tx.startswith("仕様書８の「100頁程度」を本文のみと読む場合でも"):
            trPr = tr.find(qn("w:trPr"))
            if trPr is None:
                trPr = docx.oxml.OxmlElement("w:trPr")
                tr.findall(qn("w:tc"))[0].addprevious(trPr)
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.insert(0, docx.oxml.OxmlElement("w:cantSplit"))
            n += 1
    judge(n == 1, f"頁数の見通しの囲み（{n}行）")

    # trPr の子の順（cantSplit は trHeight・tblHeader より前）・段落の書式の要素の順（v16 と同じ）
    TR = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit",
          "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden", "ins", "del", "trPrChange"]
    for trPr in d.element.body.iter(qn("w:trPr")):
        ko = list(trPr)
        jun = sorted(ko, key=lambda c: TR.index(c.tag.split("}")[1]) if c.tag.split("}")[1] in TR else len(TR))
        if jun != ko:
            for c_ in jun:
                trPr.append(c_)
    JUN = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
           "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
           "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
           "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
           "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
           "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
           "sectPr", "pPrChange"]
    for pPr in d.element.body.iter(qn("w:pPr")):
        ko = list(pPr)
        jun = sorted(ko, key=lambda c: JUN.index(c.tag.split("}")[1]) if c.tag.split("}")[1] in JUN else len(JUN))
        if jun != ko:
            for c_ in jun:
                pPr.append(c_)
    d.save(DST)

    d2 = docx.Document(DST)
    ato = zen(d2)
    print("  ── 検査1　古い記述")
    for go in ("移行が進んでいます", "この移行を前提", "第12条第4〜7項", "委託仕様書11（本仕様の定めのない",
               "年報による確定値", "算定により確定しました", "回数のご判断をお願いします"):
        judge(go not in ato, f"「{go}」0件")
    print("  ── 検査2　消えた数は旧の文に現れる")
    kyu = "\n".join(KYU)
    kieta = sorted({mm.group(0) for mm in re.finditer(r"[0-9０-９][0-9０-９,.]*(?:人|名|％|千円|件|頁|日|点|円|位)", mae)
                    if mae.count(mm.group(0)) > ato.count(mm.group(0)) and mm.group(0) not in kyu})
    judge(not kieta, "消えた数はすべて旧の文に現れる", "／".join(kieta[:5]))
    print("  ── 検査3　事業所名・他団体名・⚠・受託者の語")
    for na in ("釜房", "アルパイン", "鶴寿会", "グループホームさわやか", "グループホームふれあい",
               "北塩原", "金ヶ崎", "大雪", "川崎市", "受託者", "当社", "ビズアップ"):
        judge(na not in ato, f"「{na}」0件")
    judge("⚠" not in ato and "**" not in ato, "⚠・** 0件（第1回の書式）")
    kids = list(d2.element.body.iterchildren())
    judge(not any(a_.tag == qn("w:tbl") and b_.tag == qn("w:tbl") for a_, b_ in zip(kids, kids[1:])),
          "隣り合う表 0件（v18 の区切りを保つ）")
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜3　不適合0件")


if __name__ == "__main__":
    main()
