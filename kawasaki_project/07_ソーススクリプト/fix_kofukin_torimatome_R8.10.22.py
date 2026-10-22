# -*- coding: utf-8 -*-
"""交付金の取りまとめ表の分類の是正（令和8年10月22日）

  python3 07_ソーススクリプト/fix_kofukin_torimatome_R8.10.22.py

なぜ要るか
  取りまとめ表（R8.9.29）の「02_推移の総括」は、項目の得点を**素点**で
  比べている。⚠ 国が配点を下げただけの項目が「直近で落ちた」に入る。
    支援 目標Ⅱ「認知症サポーター等を活用した地域支援体制の構築」
      令和7年度 25点満点で25点 → 令和8年度 20点満点で20点（満点を維持）
      ア〜オの5枝が各5点から各4点に改められたことによる
  得点率で比べ直した結果（analyze_kofukin_suii_R8.10.22.py）、
  **素点で「落ちた」とされ得点率では落ちていない項目は、この1件だけ**であった。

  ⚠ 元の作成スクリプト（build_kofukin_torimatome_R8.9.29.py）は
    スクラップパッドの中間ファイルを入力にしており作り直せないため、
    出来上がった表に当てて是正版を作る。

入力　05_試算・管理シート/川崎町_交付金_取りまとめ表_R8.9.29.xlsx
出力　05_試算・管理シート/川崎町_交付金_取りまとめ表_R8.10.22.xlsx
"""
import shutil
import sys

import openpyxl
from openpyxl.styles import PatternFill

SRC = "05_試算・管理シート/川崎町_交付金_取りまとめ表_R8.9.29.xlsx"
DST = "05_試算・管理シート/川崎町_交付金_取りまとめ表_R8.10.22.xlsx"
OK = PatternFill("solid", fgColor="E2EFDA")
KOUMOKU = "認知症サポーター等を活用した地域支援体制の構築"


def main():
    shutil.copy2(SRC, DST)
    wb = openpyxl.load_workbook(DST)
    ws = wb["02_推移の総括"]
    n = 0
    for r in range(2, ws.max_row + 1):
        if (str(ws.cell(r, 1).value or "").startswith("直近で落ちた")
                and ws.cell(r, 5).value == KOUMOKU):
            ws.cell(r, 1).value = "配点の改定（満点を維持）"
            ws.cell(r, 6).value = 20
            ws.cell(r, 11).value = (
                "⚠ 取組の低下ではない。ア〜オの5枝の配点が各5点から各4点に"
                "改められ、令和8年度も20点満点で20点。素点で比べたため"
                "「直近で落ちた」に入っていた（令和8年10月22日に是正）")
            for c in range(1, 12):
                ws.cell(r, c).fill = OK
            n += 1
    ws0 = wb["00_この表について"]
    r = ws0.max_row + 1
    ws0.cell(r, 1).value = "令和8年10月22日の是正"
    ws0.cell(r, 2).value = (
        "02_推移の総括は素点で比べているため、国が配点を下げただけの項目が"
        "「直近で落ちた」に入っていた（1件）。得点率で比べ直し、"
        "「配点の改定（満点を維持）」に改めた。得点率での比較は "
        "川崎町_交付金_推移の再整理_R8.10.22.xlsx による")
    wb.save(DST)

    ok = n == 1
    wb2 = openpyxl.load_workbook(DST, data_only=True)
    ochita = [row for row in wb2["02_推移の総括"].iter_rows(min_row=2,
              values_only=True) if str(row[0] or "").startswith("直近で落ちた")
              and row[4] == KOUMOKU]
    print("交付金の取りまとめ表の分類の是正")
    print(f"   {'○' if ok else '×'} 該当行を1件改めた（実際 {n}）")
    print(f"   {'○' if not ochita else '×'} 「直近で落ちた」に認知症サポーターの"
          "項目が残っていない")
    print(f"  {DST}")
    if not ok or ochita:
        sys.exit(1)


if __name__ == "__main__":
    main()
