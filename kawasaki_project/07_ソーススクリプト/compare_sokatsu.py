# -*- coding: utf-8 -*-
"""将来推計総括表の2版を突き合わせ、設定変更が見込量に効いたかを測る。

使い方
  python3 compare_sokatsu.py <旧版xlsx> <新版xlsx> [出力txt]
"""
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")
Y = ["令和6年度", "令和7年度", "令和8年度", "令和9年度", "令和10年度", "令和11年度"]


def rows(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["2_サービス別給付費"]
    out, section, cur = {}, "", None
    for r in range(8, ws.max_row + 1):
        a, b, c = (ws.cell(r, i).value for i in (1, 2, 3))
        if a and str(a).strip():
            section = str(a).strip()
        if not (c and str(c).strip()):
            continue
        k = str(c).strip()
        vs = [ws.cell(r, x).value for x in range(4, 10)]
        vs = [v if isinstance(v, (int, float)) else None for v in vs]
        if "給付費" in k:
            nm = str(b).strip() if b and str(b).strip() else section
            cur = {"name": nm, "給付費": vs, "人数": None, "量": None,
                   "total": section.startswith("合計")}
            out[r] = cur
        elif cur is not None and "人数" in k:
            cur["人数"] = vs
        elif cur is not None and ("回数" in k or "日数" in k):
            cur["量"] = vs
    return out


def settings(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["4_施策反映の解説"]
    out = []
    for r in range(9, 45):
        a = [str(ws.cell(r, c).value) for c in range(1, 13)
             if ws.cell(r, c).value is not None and str(ws.cell(r, c).value).strip()]
        if len(a) >= 2 and ("伸び" in a[0] or "実績値" in a[0] or "手法" in a[0]):
            out.append((a[0].strip(), a[1].strip()))
    return out


def premium(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["5_保険料推計"]
    d = {}
    for r in range(1, ws.max_row + 1):
        lab = None
        for c in range(1, 6):
            v = ws.cell(r, c).value
            if v and str(v).strip():
                lab = str(v).strip()
                break
        v = ws.cell(r, 6).value
        if lab and isinstance(v, (int, float)):
            d.setdefault((lab, r >= 65), v)
    return d


A, B = sys.argv[1], sys.argv[2]
L = ["将来推計総括表の2版の突合",
     f"旧：{A.split('/')[-1]}", f"新：{B.split('/')[-1]}", "=" * 76, ""]

L.append("■ 推計方法の設定")
sa, sb = settings(A), settings(B)
for (ka, va), (kb, vb) in zip(sa, sb):
    mark = "　" if va == vb else "★変更"
    L.append(f"  {mark} {ka[:34]:<34} {va[:26]:<26} → {vb[:26]}")

ra, rb = rows(A), rows(B)
L.append("")
L.append("■ 令和9〜11年度の見込値のうち、変化した値の数")
cnt = {"人数": [0, 0], "量": [0, 0], "給付費": [0, 0]}
for r, a in ra.items():
    b = rb.get(r)
    if not b:
        continue
    for f in ("人数", "量", "給付費"):
        if not a[f] or not b[f]:
            continue
        for j in (3, 4, 5):
            x, y = a[f][j], b[f][j]
            if not isinstance(x, (int, float)) or not isinstance(y, (int, float)):
                continue
            cnt[f][0] += 1
            if abs(x - y) > 1e-9:
                cnt[f][1] += 1
for f, (t, d) in cnt.items():
    L.append(f"  {f:<8}{t:>5}値中 {d:>3}値が変化")

L.append("")
L.append("■ 給付費が変わったサービス（令和9年度・千円）")
L.append(f"  {'サービス':<32}{'R7実績':>10}{'R8基準':>10}{'旧R9':>10}{'新R9':>10}{'差':>9}")
for r, a in ra.items():
    b = rb.get(r)
    if not b or a["total"]:
        continue
    d = (b["給付費"][3] or 0) - (a["給付費"][3] or 0)
    if abs(d) < 0.5:
        continue
    L.append(f"  {a['name'][:32]:<32}{(a['給付費'][1] or 0):>10,.0f}"
             f"{(a['給付費'][2] or 0):>10,.0f}{(a['給付費'][3] or 0):>10,.0f}"
             f"{(b['給付費'][3] or 0):>10,.0f}{d:>9,.0f}")

L.append("")
L.append("■ 令和8年度が0のまま残っているサービス")
z = 0
for r, a in ra.items():
    b = rb.get(r)
    if not b or a["total"]:
        continue
    g = a["給付費"]
    if (g[2] or 0) == 0 and max(g[0] or 0, g[1] or 0) > 0:
        z += 1
        L.append(f"  {a['name'][:34]:<34}R7={g[1] or 0:>9,.0f}  "
                 f"旧R9={g[3] or 0:>6,.0f}  新R9={b['給付費'][3] or 0:>6,.0f}")
L.append(f"  計 {z}件")

L.append("")
L.append("■ 保険料の算定要素")
pa, pb = premium(A), premium(B)
for k in ("保険料基準額（月額）", "標準給付費見込額", "総給付費", "地域支援事業費",
          "第1号被保険者負担分相当額", "調整交付金相当額", "調整交付金見込額",
          "保険料収納必要額", "第1号被保険者数", "所得段階別加入割合補正係数",
          "調整交付金見込交付割合"):
    for late in (True, False):
        if (k, late) in pa and (k, late) in pb:
            x, y = pa[(k, late)], pb[(k, late)]
            m = "　" if abs(x - y) < 1e-9 else "★"
            L.append(f"  {m}{k:<28}{x:>18,.2f}{y:>18,.2f}{y-x:>14,.2f}")
            break

out = sys.argv[3] if len(sys.argv) > 3 else "compare_sokatsu.txt"
with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(L) + "\n")
print("\n".join(L))
print(f"\n→ {out}")
