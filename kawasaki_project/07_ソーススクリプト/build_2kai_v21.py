# -*- coding: utf-8 -*-
"""第2回策定委員会資料 v20 → v21（Red Team 再レビュー第4回（令和8年10月24日）の指摘の是正）

  python3 07_ソーススクリプト/build_2kai_v21.py

再レビュー原本　09_元資料/レビュー/川崎町_第10期計画_RedTeam再レビュー第4回_R8.10.24受領.md
内容の正本　　　協議用素案 Ver.2.36（fix_soan_v2360_r4review.py）
工程の日数の正本 build_iinkai_gyotei_R8.10.20.py の定数（T_PC_START・PC_DAYS・T_3・T_NYUKO・INSATSU）

  F4-01 の波及　３ 表「標準給付費」令和6年度1,093,253千円（審査支払手数料を含む）と令和7年度1,084,587千円
               （含まない）を並べていた。年報 様式3 の支払額に揃える（1,092,621千円→1,084,587千円・▲0.7％）。
               重点課題の表の「0.8％減」も素案に揃える
  F4-02 開催回数の囲み「第３回（令和９年１月）」が現在の工程（２月上旬〜中旬）と食い違って読める
               → 当初の日程であったことと、現在の工程案を書き分ける
  F4-03 「１月４日開始なら納品は３月15日ちょうど」は誤り。12月29日→１月４日は６日の後ろ倒しで、案Ⅰの余裕
               ５日を超える（３月16日・納期を１日超過）。工程の日数から計算し直して書く
  F4-06 「本文118頁」は表紙・ご挨拶・目次の５頁を含む全頁数。本文（113頁）と全体を分けて数え直す
               （別冊に移した案は全106頁・本文101頁。いずれも当方の環境での PDF の頁数）
  ５　　 事業一覧の134事業と施策体系の143事業の関係（第８章の９事業）を書く
"""
import datetime as dt
import importlib.util
import os
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)


def load(name, fn):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(spec)
    sys.modules[name] = m
    spec.loader.exec_module(m)
    return m


V9 = load("build_2kai_v9", "build_2kai_v9.py")
GY = load("build_iinkai_gyotei", "build_iinkai_gyotei_R8.10.20.py")
SRC = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v20.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v21.docx"
SOAN = "01_第10期_最新版成果品/川崎町_計画書素案_v2.36_協議用.docx"
# 頁数（当方の環境の PDF。fill_toc_pages_v200.py の「総頁数・本文の開始」による）
ZEN, MAE = 118, 5            # 素案 Ver.2.36：全118頁・本文の開始6頁目
ZEN_B, MAE_B = 106, 5        # 別冊に移した案（Ver.2.18c 圧縮案）：全106頁・本文の開始6頁目
SHIRYOHEN = 13               # 資料編（見込み）
YOTEI = 100                  # 仕様書8「100頁程度」
YOBI = "月火水木金土日"
ZEN_SU = "０１２３４５６７８９"


def zs(n):
    return str(n).translate(str.maketrans("0123456789", ZEN_SU))


def kakikae(p, s):
    p.runs[0].text = V9.clean(s)
    for r in p.runs[1:]:
        r.text = ""


def zen(d):
    return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])


def nohin(pc_start, ins):
    """パブコメの開始日から納品日（案Ⅰ・Ⅱの3回構成）。"""
    return pc_start + dt.timedelta(days=GY.PC_DAYS + GY.T_3 + GY.T_NYUKO + ins)


def hizuke(d):
    return f"{zs(d.month)}月{d.day}日（{YOBI[d.weekday()]}）"


def main():
    futeki = []
    KYU = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会資料 v20 → v21（Red Team 再レビュー第4回の指摘の是正）")
    sz = zen(docx.Document(SOAN))
    m = re.search(r"保険給付費（介護保険事業状況報告 様式3 の支払額。審査支払手数料を含まない）は令和6年度の([\d,]+)千円から"
                  r"令和7年度の([\d,]+)千円へ([\d.]+)％減少", sz)
    judge(m is not None, "素案 7-2 の比較の文（様式3 の支払額）")
    r6, r7, gen = m.group(1), m.group(2), m.group(3)
    shutil.copy2(SRC, DST)
    d = docx.Document(DST)
    mae = zen(d)

    def p_rep(kyu, shin, n=1):
        h = [p for p in d.paragraphs if kyu in p.text]
        judge(len(h) == n, f"本文「{kyu[:24]}」（{len(h)}か所）")
        for p in h:
            KYU.append(p.text)
            kakikae(p, p.text.replace(kyu, shin))

    def c_rep(kyu, shin, n=1):
        k = 0
        for t in d.tables:
            seen = []
            for r in t.rows:
                for c in r.cells:
                    if c._tc in seen:
                        continue
                    seen.append(c._tc)
                    if kyu in c.text:
                        KYU.append(c.text)
                        V9.set_cell(c, c.text.replace(kyu, shin))
                        k += 1
        judge(k == n, f"表「{kyu[:24]}」（{k}か所）")

    print("  ── F4-01 の波及　給付費の推移の表")
    t = [t for t in d.tables if [c.text.strip() for c in t.rows[0].cells] == ["年　度", "標準給付費", "前年比"]]
    judge(len(t) == 1, "給付費の推移の表を特定")
    if t:
        t = t[0]
        KYU.append("\n".join(c.text for r in t.rows for c in r.cells))
        V9.set_cell(t.rows[0].cells[1], "保険給付費（支払額）")
        V9.set_cell(t.rows[1].cells[1], f"{r6}千円")
        V9.set_cell(t.rows[2].cells[2], f"▲{gen}％")
        V9.set_cell(t.rows[3].cells[2], "（参考）＋2.1％")
    p_rep("資料：令和6年度・令和7年度は介護保険事業状況報告（年報）による確定値、令和8年度は町提供実績データ（令和8年9月1日受領）"
          "による見込みです。",
          "資料：令和6年度・令和7年度は介護保険事業状況報告（年報）様式3の保険給付の支払額（審査支払手数料を含まない）、"
          "令和8年度は町提供実績データ（令和8年9月1日受領）の保険給付費見込額（審査支払手数料を含む）です。令和8年度の"
          "前年比は集計の範囲が異なるため参考値です。")
    c_rep("給付費の平坦化（令和7年度0.8％減）", f"給付費の平坦化（令和7年度{gen}％減）")

    print("  ── F4-02 第３回の時期（当初の日程と現在の工程案）")
    c_rep("第１回策定委員会（令和８年９月２日）の当日にお示しいただいた日程は、第３回（令和９年１月）で答申の内容を確定する"
          "３回構成でした。上表は、11月24日の開催から納期まで逆算したものです。",
          "第１回策定委員会（令和８年９月２日）の当日にお示しいただいた当初の日程は、第３回を令和９年１月として答申の内容を"
          "確定する３回構成でした。第２回の開催日が11月24日となったため、上表は同日から納期まで逆算し、第３回を令和９年２月"
          "上旬〜中旬としています。")

    print("  ── F4-03 パブリックコメントを１月４日開始とした場合の納品日")
    pc0 = GY.KAISAI2 + dt.timedelta(days=GY.T_PC_START)
    pc1 = dt.date(2027, 1, 4)
    zure = (pc1 - pc0).days
    n0 = nohin(pc0, GY.INSATSU)
    n1 = nohin(pc1, GY.INSATSU)
    n1b = nohin(pc1, 14)
    yoyu0 = (GY.NOUKI - n0).days
    koe = (n1 - GY.NOUKI).days
    judge(pc0 == dt.date(2026, 12, 29) and n0 == dt.date(2027, 3, 10) and yoyu0 == 5,
          f"案Ⅰ：パブコメ開始 {pc0}・納品 {n0}・余裕 {yoyu0}日（資料の表と一致）")
    judge(zure == 6 and koe == 1, f"１月４日開始：{zure}日の後ろ倒しで案Ⅰの納品 {n1}（納期を{koe}日超過）・案Ⅱ {n1b}")
    c_rep("パブリックコメントの開始を令和９年１月４日とする場合は、納品が令和９年３月15日ちょうどになります。",
          f"パブリックコメントの開始を令和９年１月４日（12月29日から{zs(zure)}日後）とする場合は、後続の工程が{zs(zure)}日ずつ"
          f"遅れ、案Ⅰでは納品が令和９年{hizuke(n1)}となり、納期を{zs(koe)}日超えます。１月４日開始とするには、第３回から入稿まで"
          f"（{GY.T_NYUKO}日）又は印刷製本（３週間）を{zs(koe)}日以上短縮する必要があります（案Ⅱ（印刷２週間）では"
          f"{hizuke(n1b)}納品・余裕{zs((GY.NOUKI - n1b).days)}日）。")
    c_rep("令和９年１月４日（月）開始とする案を併記します",
          f"令和９年１月４日（月）開始とする案を併記します（その場合、案Ⅰの納品は{hizuke(n1)}で納期を{zs(koe)}日超えます）")

    print("  ── F4-06 頁数の数え方（全体・本文・資料編）")
    hon, hon_b = ZEN - MAE, ZEN_B - MAE_B
    kei, kei_b = ZEN + SHIRYOHEN, ZEN_B + SHIRYOHEN
    p_rep(f"現在の計画素案は本文{ZEN}頁で、資料編（見込13頁）を加えると計131頁となり、仕様書８の「100頁程度」を31頁上回ります。"
          "詳細な分析を別冊に移すことで、本文を106頁程度とする案を検討しています。",
          f"現在の計画素案は全{ZEN}頁（表紙・ご挨拶・目次の{zs(MAE)}頁と本文{hon}頁。当方の環境でのPDFの頁数）で、資料編"
          f"（見込{SHIRYOHEN}頁）を加えると計{kei}頁となります。仕様書８の「100頁程度」に対し、本文のみで数えると{hon - YOTEI}頁、"
          f"資料編を含む冊子全体で数えると{kei - YOTEI}頁上回ります。詳細な分析を別冊に移すことで、全体を{ZEN_B}頁程度"
          f"（本文{hon_b}頁程度）とする案を検討しています。")
    p_rep("計画書の事業一覧表（第５章・第６章の17表。134事業と再掲４件）",
          "計画書の事業一覧表（第５章・第６章の17表。134事業と再掲４件。施策体系の全143事業は、これに第８章の９事業を"
          "加えたものです）")
    tp = [t for t in d.tables if t.rows[0].cells[0].text.strip() == "区　分" and "本文" in t.rows[0].cells[1].text]
    judge(len(tp) == 1, "頁数の見通しの表を特定")
    if tp:
        tp = tp[0]
        KYU.append("\n".join(c.text for r in tp.rows for c in r.cells))
        V9.set_cell(tp.rows[0].cells[1], f"本文（表紙・目次の{zs(MAE)}頁を除く）")
        V9.set_cell(tp.rows[0].cells[2], f"計（表紙・目次{zs(MAE)}頁と資料編{SHIRYOHEN}頁を含む）")
        V9.set_cell(tp.rows[0].cells[3], "仕様書８（100頁程度）との差（本文／計）")
        V9.set_cell(tp.rows[1].cells[1], f"{hon}頁")
        V9.set_cell(tp.rows[1].cells[2], f"{kei}頁")
        V9.set_cell(tp.rows[1].cells[3], f"＋{hon - YOTEI}頁／＋{kei - YOTEI}頁")
        V9.set_cell(tp.rows[2].cells[1], f"{hon_b}頁程度")
        V9.set_cell(tp.rows[2].cells[2], f"{kei_b}頁程度")
        V9.set_cell(tp.rows[2].cells[3], f"＋{hon_b - YOTEI}頁程度／＋{kei_b - YOTEI}頁程度")
    c_rep("仕様書８の「100頁程度」を本文のみと読む場合でも、別冊に移した後の本文は106頁程度であり、なお６頁程度の圧縮を要します。"
          "資料編13頁を含めるかどうかを含めてご判断をお願いします。",
          f"仕様書８の「100頁程度」を本文のみと読む場合は、別冊に移した後の本文{hon_b}頁程度でほぼ収まります。表紙・目次・資料編を"
          f"含む冊子全体と読む場合は{kei_b}頁程度で、なお{kei_b - YOTEI}頁程度の圧縮を要します。数える範囲（本文・資料編・"
          "概要版）を町と書面で確かめたうえで定めます。")

    # 段落の書式の要素の順（v16 と同じ）
    JUN = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
           "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
           "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
           "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
           "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
           "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
           "sectPr", "pPrChange"]
    for pPr in d.element.body.iter(qn("w:pPr")):
        ko = list(pPr)
        jun = sorted(ko, key=lambda c: JUN.index(c.tag.split("}")[1]) if c.tag.split("}")[1] in JUN else len(JUN))
        if jun != ko:
            for c_ in jun:
                pPr.append(c_)
    d.save(DST)

    d2 = docx.Document(DST)
    ato = zen(d2)
    print("  ── 検査1　古い記述")
    for go in ("1,093,253", "0.8％", "ちょうど", "第３回（令和９年１月）", "本文118頁", "本文を106頁", "本文は106頁"):
        judge(go not in ato, f"「{go}」0件")
    print("  ── 検査2　素案と同じ値")
    judge(f"{r6}千円" in ato and f"{r6}千円" in sz, f"令和6年度 {r6}千円（素案 7-2 にある）")
    print("  ── 検査3　消えた数は旧の文に現れる")
    kyu = "\n".join(KYU)
    kieta = sorted({mm.group(0) for mm in re.finditer(r"[0-9０-９][0-9０-９,.]*(?:人|名|％|千円|件|頁|日|点|円|位)", mae)
                    if mae.count(mm.group(0)) > ato.count(mm.group(0)) and mm.group(0) not in kyu})
    judge(not kieta, "消えた数はすべて旧の文に現れる", "／".join(kieta[:5]))
    print("  ── 検査4　事業所名・他団体名・⚠・受託者の語・表の区切り")
    for na in ("釜房", "アルパイン", "鶴寿会", "北塩原", "金ヶ崎", "大雪", "川崎市", "受託者", "当社", "ビズアップ"):
        judge(na not in ato, f"「{na}」0件")
    judge("⚠" not in ato and "**" not in ato, "⚠・** 0件")
    kids = list(d2.element.body.iterchildren())
    judge(not any(a_.tag == qn("w:tbl") and b_.tag == qn("w:tbl") for a_, b_ in zip(kids, kids[1:])), "隣り合う表 0件")
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜4　不適合0件")


if __name__ == "__main__":
    main()
