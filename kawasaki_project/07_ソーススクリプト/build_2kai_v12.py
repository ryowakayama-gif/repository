# -*- coding: utf-8 -*-
"""第2回策定委員会資料 v11 → v12（町内の軽費老人ホームの廃止の確認結果）
（令和8年10月22日）

  python3 07_ソーススクリプト/build_2kai_v12.py

v11 からの変更の理由
  計画素案を Ver.2.25 に上げた（fix_soan_v2250_keihi.py）。同じ版で直す。
  (1) 表6 の軽費老人ホームの行：廃止日（令和6年3月31日）、入所者の推移
      （13人→7人→0人）と退所先（他の施設28・入院3・在宅1）を書き、
      ⚠ 定員70人が住まいを失ったのではないことを明示する
  (2) 版 ［第11版］→［第12版］、素案の版 Ver.2.24 → Ver.2.25（118頁は不変）

数値の正本　09_元資料/その他データ/県入所状況調_川崎町_R8.10.22.csv
            （fix_soan_v2250_keihi.py が読む）
"""
import importlib.util
import os
import re
import shutil
import sys

import docx

HERE = os.path.dirname(os.path.abspath(__file__))


def load(name, fn):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, fn))
    m = importlib.util.module_from_spec(spec)
    sys.modules[name] = m
    spec.loader.exec_module(m)
    return m


V9 = load("build_2kai_v9", "build_2kai_v9.py")
S25 = load("fix_soan_v2250", "fix_soan_v2250_keihi.py")
A, GEN = S25.A, S25.GEN

SRC = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v11.docx"
DST = "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v12.docx"

KYU_ATAMA = "⚠ 町内にあった軽費老人ホーム（定員70人）は、"
MIDASHI = (f"⚠ 町内にあった軽費老人ホーム（定員{A['定員']}人）は、"
           f"{S25.HAISHI}に廃止されました。 ")
HON = (f"県の入所状況調では、入所者は令和5年4月1日{GEN['R5.4.1']}人→7月1日"
       f"{GEN['R5.7.1']}人→10月1日{GEN['R5.10.1']}人（休止中）と減り、"
       f"令和5年1月から9月までに退所した{A['退所']}人のうち{A['他施設']}人は"
       f"他の施設へ移られています（入院{A['入院']}人・在宅{A['在宅']}人）。"
       f"定員の{A['定員']}人が住まいを失ったものではありません。老人福祉法の"
       "施設であり、介護サービス情報公表システムには載りません。本町の住民の"
       "数と移られた先の内訳は、宮城県への廃止届の記載内容を確認しています。")


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu)
                                                   if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("第2回策定委員会資料 v11 → v12（軽費老人ホームの廃止の確認結果）")
    print()
    shutil.copy2(SRC, DST)
    d = docx.Document(DST)
    mae = "\n".join([p.text for p in d.paragraphs]
                    + [c.text for t in d.tables for r in t.rows
                       for c in r.cells])
    KYU = []

    print("  ── ① 表6 の軽費老人ホームの行")
    c = d.tables[6].rows[0].cells[0]
    hit = [p for p in c.paragraphs if p.text.startswith(KYU_ATAMA)]
    judge(len(hit) == 1, f"当たる段落が1件（実際 {len(hit)}）")
    if len(hit) == 1:
        p = hit[0]
        KYU.append(p.text)
        p.runs[0].text = MIDASHI      # 1つ目の run は太字の見出し
        if len(p.runs) > 1:
            p.runs[1].text = HON
            for r in p.runs[2:]:
                r.text = ""
        else:
            p.add_run(HON)

    print("  ── ② 版")
    n = 0
    for p in d.paragraphs:
        s = p.text
        t = (s.replace("［第11版］", "［第12版］")
             .replace("Ver.2.24", "Ver.2.25"))
        if t != s:
            KYU.append(s)
            p.runs[0].text = V9.clean(t)
            for r in p.runs[1:]:
                r.text = ""
            n += 1
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                if "Ver.2.24（現在" in cell.text:
                    KYU.append(cell.text)
                    V9.set_cell(cell, cell.text.replace("Ver.2.24", "Ver.2.25")
                                .replace("交付金の推移の内訳を追加",
                                         "交付金の推移の内訳、軽費老人ホームの"
                                         "廃止の確認結果を追加"))
                    n += 1
    judge(n >= 3, f"版を改めた（{n}か所）")
    d.save(DST)

    d2 = docx.Document(DST)
    ato = "\n".join([p.text for p in d2.paragraphs]
                    + [c.text for t in d2.tables for r in t.rows
                       for c in r.cells])
    print("  ── 検査1　素案と同じ値")
    soan = docx.Document(S25.SAKI)
    st = "\n".join(p.text for p in soan.paragraphs)
    for go, sgo in ((S25.HAISHI, S25.HAISHI),
                    (f"{A['退所']}人のうち{A['他施設']}人", f"退所した{A['退所']}人"),
                    (f"4月1日{GEN['R5.4.1']}人", f"4月1日時点で{GEN['R5.4.1']}人")):
        judge(go in ato and sgo in st, f"「{go}」（素案にも同じ値）")
    judge("住まいを失ったものではありません" in ato, "定員と入所者の別を書いた")
    judge("Ver.2.24" not in ato and "［第12版］" in ato, "版が［第12版］・Ver.2.25")
    print("  ── 検査2　⚠ 既にあった数が消えていないか")
    kyu_zen = "\n".join(KYU)
    kieta = sorted({m.group(0) for m in re.finditer(
        r"[0-9０-９][0-9０-９,.]*(?:人|名|％|千円|件|頁|日|点|円)", mae)
        if mae.count(m.group(0)) > ato.count(m.group(0))
        and m.group(0) not in kyu_zen})
    judge(not kieta, "消えた数はすべて差し替えた旧の文に現れている",
          "／".join(kieta[:5]))
    print("  ── 検査3　⚠ 施設名・法人名・他団体名")
    for na in ("蔵王グリーンホーム", "鶴寿会", "釜房みどりの園", "アルパイン",
               "ながさか", "リリーハイム", "北塩原", "金ヶ崎", "川崎市"):
        judge(na not in ato, f"「{na}」が入っていない")
    judge("**" not in ato, "強調記号（**）が0件")

    print()
    print(f"  {DST}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 検査1〜3　不適合0件")


if __name__ == "__main__":
    main()
