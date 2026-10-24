# -*- coding: utf-8 -*-
"""業務工程管理表の更新（交付金の状況の再整理と、素案Ver.2.24・資料v11・問答集v7）
（令和8年10月22日・同日3度目）

  python3 07_ソーススクリプト/update_wbs_R8.10.22c.py

  ・交付金の推移を枝番の単位・得点率で数え直した（37点の内訳）
  ・素案の事実と違う記述（保険料算定は交付実績を据え置いている）を直した
  ・交付金を保険料に見込むかを No.39 に追記した（総括表は4,500千円を見込む）

⚠ 列8（状態）には書き込まない。⚠ 同日3度目のため印を分ける。
"""
import copy
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
SOAN_OLD = "川崎町_計画書素案_v2.23_県調査反映版.docx"
SOAN = "川崎町_計画書素案_v2.24_交付金整理版.docx"
SHIRYO_OLD = "川崎町_第2回策定委員会資料_R8.11_v10.docx"
SHIRYO = "川崎町_第2回策定委員会資料_R8.11_v11.docx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月22日 更新③（交付金の再整理・素案Ver.2.24）】"


def iro(ws, r, n):
    for c in range(1, n + 1):
        ws.cell(r, c).font = FONT
        ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")


TSUIKI = {
    "39": (10, f"{HIZUKE}⚠ 素案（Ver.2.23 まで）は「保険料算定は令和8年度の"
               f"交付実績を据え置いている」と書いていたが、保険料算定表は"
               f"交付金を差し引いておらず、事実と違っていた（Ver.2.24 で是正）。"
               f"一方、令和8年10月19日出力の将来推計総括表は「保険者機能強化"
               f"推進交付金等の交付見込額」に3か年4,500千円を置き、差し引いて"
               f"いる。見込めば月額約40円下がる（6,822円→6,782円）。"
               f"令和6年度の交付見込額は3,738千円（推進1,277・支援2,461、"
               f"うち激変緩和措置分1,231千円）。令和7・8年度の交付額と、"
               f"保険料に見込むか否か・見込む額のご判断をお願いします。"
               f"既定値＝見込まない（素案のまま）とし、総括表との食い違いを"
               f"第3回の資料で明示する。　"),
    "58": (10, f"{HIZUKE}37点の減少を枝番の単位・得点率で分けると、成果指標群"
               f"▲30・配点の改定▲7・新設＋2・取組の後退▲34・前進＋32。"
               f"「評価結果の活用」16点は取組の後退34点のうち最大の項目"
               f"（単独で最大の要因ではない。最大は成果指標群）。　"),
}


def main():
    wb = openpyxl.load_workbook(WBS)
    ws = wb["03_確認事項一覧"]
    n = 0
    for x in range(5, ws.max_row + 1):
        no = str(ws.cell(x, 1).value or "").strip()
        if no in TSUIKI and HIZUKE not in str(ws.cell(x, 10).value or ""):
            retsu, moji = TSUIKI[no]
            ws.cell(x, retsu).value = moji + str(ws.cell(x, retsu).value or "")
            iro(ws, x, 10)
            n += 1

    ws = wb["04_成果品管理"]
    n_seika = 0
    for x in range(5, ws.max_row + 1):
        no = str(ws.cell(x, 1).value or "").strip()
        na = str(ws.cell(x, 3).value or "")
        if no == "4" and SOAN_OLD in na:
            ws.cell(x, 3).value = na.replace(SOAN_OLD, SOAN)
            ws.cell(x, 8).value = (
                f"{HIZUKE}素案を Ver.2.24 に上げた（全118頁）。⚠ 3-3 の37点の"
                f"内訳の誤り（16点＋15点としていた）と、3-3・10章の事実と違う"
                f"記述（保険料算定は交付実績を据え置いている）を直した。9-3 の"
                f"不確定要素と前提を動かした場合の幅に交付金の行（▲40円）を"
                f"加えた。保険料の額は動かしていない。［要確認］15→20件。　"
                + str(ws.cell(x, 8).value or ""))
            iro(ws, x, 9)
            n_seika += 1
        elif no == "6" and SHIRYO_OLD in na:
            ws.cell(x, 3).value = (na.replace(SHIRYO_OLD, SHIRYO)
                                   .replace("同_想定問答集_R8.11_v6.docx",
                                            "同_想定問答集_R8.11_v7.docx"))
            ws.cell(x, 8).value = (
                f"{HIZUKE}資料を v11（全11版・表47）、想定問答集を v7"
                f"（全67問）に上げた。7 に37点の内訳と3か年の推移（宮城県平均が"
                f"1年で78.6点上昇）を加えた。⚠ 6 の要因の表の値を素案に揃えた"
                f"（改定率1％ ＋65円→＋66円、収納率99.6％ ▲247円→▲246円）。"
                f"⚠ 問答集のヘッダーが「第1回」のままであったのを直した。　"
                + str(ws.cell(x, 8).value or ""))
            iro(ws, x, 9)
            n_seika += 1

    ws2 = wb["08_作業順位"]
    aru = any(str(ws2.cell(x, 1).value or "").strip() == "38"
              for x in range(5, ws2.max_row + 1))
    if not aru:
        moto = ws2.max_row
        r = moto + 1
        for c in range(1, 11):
            ws2.cell(r, c)._style = copy.copy(ws2.cell(moto, c)._style)
        ws2.row_dimensions[r].height = ws2.row_dimensions[moto].height
        atai = [
            "38", "交付金の状況の再整理と策定委員会資料への反映", "R8.10.22",
            "ご指示による。同じ原典（全国集計結果）を枝の単位まで解いた他の"
            "案件の点検の観点を借り、本町の記述に同じ誤りがないかを確かめた。"
            "町のご回答を待たずに行える。",
            "No.39（交付金を保険料に見込むか）／No.58",
            "交付金の推移の再整理（xlsx・5シート）／取りまとめ表 R8.10.22／"
            "計画素案 Ver.2.24／資料 v11／想定問答集 v7／交付金の状況の"
            "再整理（md）",
            "1日", "完了",
            "⚠ **素案に事実と違う記述が2か所あった**（保険料算定は交付実績を"
            "据え置いている→実際は交付金を見込んでいない）。⚠ 37点の内訳を"
            "成果▲30・配点改定▲7・新設＋2・後退▲34・前進＋32と正した。"
            "宮城県平均が1年で78.6点上がったことが県内順位3位→8位の主因。"
            "総括表は交付金4,500千円を見込んでおり、素案と食い違う"
            "（月額約40円）。No.39 でご判断を仰ぐ。",
        ]
        for c, v in enumerate(atai, 1):
            ws2.cell(r, c).value = v
    wb.save(WBS)

    futeki = []

    def judge(ok, mi):
        print(("   ○ " if ok else "   × ") + mi)
        if not ok:
            futeki.append(mi)

    print("業務工程管理表の更新（交付金の再整理・素案Ver.2.24）")
    judge(n in (0, 2), f"確認事項2件に書いた（{n}件）")
    judge(n_seika in (0, 2), f"成果品管理の成果品4・6 を更新した（{n_seika}件）")
    wb2 = openpyxl.load_workbook(WBS, data_only=True)
    ws3 = wb2["03_確認事項一覧"]
    jotai = {str(ws3.cell(x, 8).value or "") for x in range(5, ws3.max_row + 1)
             if str(ws3.cell(x, 1).value or "").strip() in TSUIKI}
    judge(jotai <= {"確認待ち", "決着", "保留", "回答済"},
          "⚠ 列8（状態）を動かしていない")
    nos = sorted(int(s) for s in (str(ws3.cell(x, 1).value or "").strip()
                                  for x in range(5, ws3.max_row + 1))
                 if s.isdigit())
    judge(nos == list(range(1, max(nos) + 1)), f"欠番0（{len(nos)}件）")
    jun = [str(wb2["08_作業順位"].cell(x, 1).value or "").strip()
           for x in range(5, wb2["08_作業順位"].max_row + 1)]
    judge(jun.count("38") == 1, "作業順位38 が1行だけある")
    s = "\n".join(str(v) for sh in wb2 for row in sh.iter_rows(values_only=True)
                  for v in row if v is not None)
    for na in ("釜房みどりの園", "アルパイン", "蔵王グリーンホーム",
               "ながさか", "リリーハイム", "川崎市", "北塩原", "金ヶ崎"):
        judge(na not in s, f"「{na}」が入っていない")
    judge(SOAN in s and SHIRYO in s and "想定問答集_R8.11_v7" in s,
          "成果品管理に素案 Ver.2.24・資料 v11・問答集 v7 が載っている")
    if futeki:
        sys.exit(1)
    print("   ○ 自己点検 不適合0件")


if __name__ == "__main__":
    main()
