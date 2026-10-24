# -*- coding: utf-8 -*-
"""協議用素案 Ver.2.32 → Ver.2.33（Red Team 再レビュー（令和8年10月24日）の指摘の是正）

  python3 07_ソーススクリプト/fix_soan_v2330_saireview.py

再レビュー原本　09_元資料/レビュー/川崎町_第10期計画_RedTeam再レビュー_R8.10.24受領.md

  N01 9-3　介護報酬改定率別の算定Ｂ・Ｃを素案に書く（資料が素案にない値（6,235円等）を持ち、
       しかも基金の取崩しを固定した計算と合わなかった）。改定率を標準給付費見込額に一律に乗じる
       （この前提で算定Ａの6,888円・7,018円を再現できることを確かめる）。取崩しは保険料収納必要額から
       一定額を差し引くため、改定による増減の幅は算定Ａ・Ｂ・Ｃでほぼ同じ
  N09 5-5　入所・居住系10人の月額　約211円 → 約212円（211.7円。1円未満四捨五入）と算式を明記
  N06 9-3　介護給付費準備基金の表に紛れ込んだ「有効票（576件）」の列を削る
  N08 全体　「全国1,741保険者」等 → 市町村（交付金の評価の単位は市町村。1,741＝全国の市町村数）。
       同じ節の「全保険者で按分」「他の保険者の得点」「同じ得点の保険者」も市町村に揃える
  N10 3-3　「同時に進んだ結果と考えられます」→ 影響している可能性。因果は追加の検証を要する
  N03 5-4・10-4　生活支援コーディネーター26人：町の実績データの名称は「生活支援サポーター（SC）」で
       あり、第1層・第2層の配置の内訳は町に確かめる【町確認】（確認事項No.190）

終了コード　0＝検査すべて適合／1＝1件でも不適合
"""
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

MOTO = "01_第10期_最新版成果品/川崎町_計画書素案_v2.32_協議用.docx"
SAKI = "01_第10期_最新版成果品/川崎町_計画書素案_v2.33_協議用.docx"
KIKIN = {"Ａ": 0, "Ｂ": 76250, "Ｃ": 152500}


def kakikae(p, s):
    p.runs[0].text = s
    for r in p.runs[1:]:
        r.text = ""


def set_cell(c, s):
    kakikae(c.paragraphs[0], s)
    for q in c.paragraphs[1:]:
        q._p.getparent().remove(q._p)


def zenbun(d):
    return "\n".join([p.text for p in d.paragraphs]
                     + [c.text for t in d.tables for r in t.rows for c in r.cells])


def num(s):
    return float(re.sub(r"[^\d.]", "", s))


def main():
    futeki = []

    def judge(ok, mi, setsu=""):
        print(("   ○ " if ok else "   × ") + mi + (("　" + setsu) if setsu else ""))
        if not ok:
            futeki.append(mi)

    print("協議用素案 Ver.2.32 → Ver.2.33（Red Team 再レビューの指摘の是正）")
    shutil.copy2(MOTO, SAKI)
    d = docx.Document(SAKI)
    mae = zenbun(d)
    P = d.paragraphs

    print("  ── N01 介護報酬改定率別の算定Ａ・Ｂ・Ｃ（表101 から計算）")
    t101 = [t for t in d.tables if [c.text.strip() for c in t.rows[0].cells][:3] == ["記号", "項目", "第10期計"]][0]
    row = {r.cells[0].text.strip(): r for r in t101.rows}
    A_y = [num(row["A"].cells[i].text) for i in (3, 4, 5)]
    J = num(row["J"].cells[2].text)
    N = num(row["③"].cells[2].text)
    rates = [float(x) / 100 for x in re.findall(r"令和\d+年度([\d.]+)％",
             [p.text for p in P if "見込交付割合は全国の仮置値による自動計算値" in p.text][0])]

    def gaku(r, kikin):
        dj = sum(a * r * (0.23 + 0.05 - e) for a, e in zip(A_y, rates))
        return (J + dj - kikin) * 1000 / 0.96 / N / 12

    tbl = {}
    for r in (0, 0.01, 0.03):
        tbl[r] = {k: gaku(r, v) for k, v in KIKIN.items()}
    judge(round(tbl[0]["Ａ"]) == 6822 and round(tbl[0]["Ｂ"]) == 6143 and round(tbl[0]["Ｃ"]) == 5464,
          "改定0％ Ａ6,822・Ｂ6,143・Ｃ5,464 を再現")
    judge(round(tbl[0.01]["Ａ"]) == 6888 and round(tbl[0.03]["Ａ"]) == 7018,
          f"改定率を標準給付費見込額に乗じると算定Ａ 1％{round(tbl[0.01]['Ａ'])}・3％{round(tbl[0.03]['Ａ'])}（感度の表の6,888・7,018と一致）")
    B1, C1, B3, C3 = (round(tbl[0.01]["Ｂ"]), round(tbl[0.01]["Ｃ"]), round(tbl[0.03]["Ｂ"]), round(tbl[0.03]["Ｃ"]))
    sa = {k: (round(tbl[0.01][k]) - round(tbl[0][k]), round(tbl[0.03][k]) - round(tbl[0][k])) for k in KIKIN}
    bun = (f"介護報酬の改定率は、標準給付費見込額に一律に乗じて試算しています。基金の取崩しは保険料収納必要額から"
           f"一定額を差し引くものであるため、改定による増減の幅は算定Ａ・Ｂ・Ｃでほぼ同じです。改定率1％では"
           f"算定Ａ{round(tbl[0.01]['Ａ']):,}円・算定Ｂ{B1:,}円・算定Ｃ{C1:,}円（それぞれ＋{sa['Ａ'][0]}円・"
           f"＋{sa['Ｂ'][0]}円・＋{sa['Ｃ'][0]}円）、3％では算定Ａ{round(tbl[0.03]['Ａ']):,}円・算定Ｂ{B3:,}円・"
           f"算定Ｃ{C3:,}円（それぞれ＋{sa['Ａ'][1]}円・＋{sa['Ｂ'][1]}円・＋{sa['Ｃ'][1]}円）です。")
    h = [p for p in P if "③令和9年度の介護報酬改定の内容は未確定です。" in p.text]
    judge(len(h) == 1, "9-3 不確定要素の段落")
    for p in h:
        kakikae(p, p.text.replace("③令和9年度の介護報酬改定の内容は未確定です。",
                                  "③令和9年度の介護報酬改定の内容は未確定です。" + bun))

    print("  ── N09 入所・居住系10人の月額（1円未満四捨五入）")
    h = [p for p in P if "10人では約211円動きます（基金の取崩しの額によりません）。" in p.text]
    judge(len(h) == 1, "5-5 の段落")
    tanka = num(h[0].text.split("1人1月あたり給付費は約")[1].split("万円")[0]) * 10
    dj = sum(10 * tanka * 12 * (0.23 + 0.05 - e) for e in rates)
    dm = dj * 1000 / 0.96 / N / 12
    j10 = round(dm)
    judge(j10 == 212, f"10人で {dm:.1f}円 → 四捨五入 {j10}円")
    for p in h:
        kakikae(p, p.text.replace(
            "10人では約211円動きます（基金の取崩しの額によりません）。",
            f"10人では約{j10}円動きます（1人1月あたり約29.9万円×10人×12か月の3か年分に、第1号被保険者の負担"
            f"（23％）と調整交付金（相当額5％－見込交付割合）を加減し、予定収納率と補正後被保険者数で割った値。"
            f"1円未満四捨五入。基金の取崩しの額によりません）。"))

    print("  ── N06 準備基金の表の「有効票」の列")
    t = [t for t in d.tables if [c.text.strip() for c in t.rows[0].cells] == ["項目", "金額・方針", "有効票"]]
    judge(len(t) == 1, "準備基金の表を特定")
    if t:
        tb = t[0]._tbl
        grid = tb.find(qn("w:tblGrid"))
        cols = grid.findall(qn("w:gridCol"))
        w_del = int(cols[2].get(qn("w:w")))
        grid.remove(cols[2])
        cols[1].set(qn("w:w"), str(int(cols[1].get(qn("w:w"))) + w_del))
        for tr in tb.findall(qn("w:tr")):
            tcs = tr.findall(qn("w:tc"))
            w1 = tcs[1].find(qn("w:tcPr")).find(qn("w:tcW"))
            w1.set(qn("w:w"), str(int(w1.get(qn("w:w"))) + int(tcs[2].find(qn("w:tcPr")).find(qn("w:tcW")).get(qn("w:w")))))
            tr.remove(tcs[2])
        judge([c.text.strip() for c in t[0].rows[0].cells] == ["項目", "金額・方針"], "列を削った（2列）")

    print("  ── N08 「保険者」の数 → 市町村")
    n = 0
    pats = [(r"(\d)保険者", r"\1市町村"), (r"35市町村の市町村平均", "35市町村の平均"),
            (r"380点満点である保険者は", "380点満点である市町村は"),
            # 交付金の評価・按分の単位も市町村（同じ節の言い回し）
            (r"同じ得点の保険者は", "同じ得点の市町村は"), (r"同じ64点の保険者は", "同じ64点の市町村は"),
            (r"各保険者の評価得点", "各市町村の評価得点"), (r"全保険者で按分", "全市町村で按分"),
            (r"他の保険者の得点", "他の市町村の得点")]
    targets = list(P) + [p for t_ in d.tables for r in t_.rows for c in r.cells for p in c.paragraphs]
    seen = set()
    for p in targets:
        if id(p._p) in seen:
            continue
        seen.add(id(p._p))
        s = p.text
        s2 = s
        for a, b in pats:
            s2 = re.sub(a, b, s2)
        if s2 != s:
            kakikae(p, s2)
            n += 1
    judge(n >= 10, f"{n}段落・セルを改めた")

    print("  ── N10 3-3 因果の言い切り")
    KYU10 = "同時に進んだ結果と考えられます。"
    h = [p for p in P if KYU10 in p.text]
    judge(len(h) == 1, "3-3 の段落")
    for p in h:
        kakikae(p, p.text.replace(KYU10, "同時に進んだことが影響している可能性があります（因果関係は、町内外別の"
                                         "利用の推移等による追加の検証を要します）。"))

    print("  ── N03 生活支援コーディネーターの区分【町確認】")
    n = 0
    for t_ in d.tables:
        for r in t_.rows:
            for c in r.cells:
                if c.text.strip() == "26人（令和7年度見込）":
                    set_cell(c, "26人（令和7年度見込）【町確認】")
                    n += 1
                elif "実：生活支援コーディネーター26人（令和7年度見込。町の実績データ）" in c.text and "生活支援サポーター（SC）" not in c.text:
                    set_cell(c, c.text.replace(
                        "実：生活支援コーディネーター26人（令和7年度見込。町の実績データ）",
                        "実：生活支援コーディネーター26人（令和7年度見込。町の実績データでの名称は「生活支援サポーター"
                        "（SC）」。第1層・第2層の配置の内訳は確認中【町確認】）"))
                    n += 1
    judge(n == 2, f"10-4 KPI と 5-4 の表（{n}か所）")

    d.save(SAKI)
    ato = zenbun(docx.Document(SAKI))
    print("  ── 検査1　古い記述")
    for go in ("約211円", "保険者中", "1,741保険者", "35保険者", "285保険者", "結果と考えられます",
               "同じ得点の保険者", "全保険者で按分", "他の保険者の得点", "各保険者の評価得点"):
        judge(go not in ato, f"「{go}」0件")
    judge("有効票" not in "".join(c.text for t_ in docx.Document(SAKI).tables
                                 if t_.rows[0].cells[0].text.strip() == "項目"
                                 and "金額・方針" in t_.rows[0].cells[1].text for r in t_.rows for c in r.cells),
          "準備基金の表に「有効票」0件")
    print("  ── 検査2　未確定の印")
    judge(ato.count("【町確認】") == mae.count("【町確認】") + 2, "【町確認】＋2（生活支援コーディネーター）")
    for w in ("［要確認］", "［要協議］", "［要内訳］"):
        judge(ato.count(w) == mae.count(w), f"{w} 増減なし")
    print("  ── 検査3　他団体名・受託者の語")
    for w in ("大雪", "北塩原", "金ヶ崎", "川崎市", "受託者", "当社", "ビズアップ"):
        judge(ato.count(w) == mae.count(w), f"「{w}」増減なし")
    print(f"  {SAKI}")
    if futeki:
        print(f"   × 不適合 {len(futeki)}件")
        sys.exit(1)
    print("   ○ 不適合0件")


if __name__ == "__main__":
    main()
