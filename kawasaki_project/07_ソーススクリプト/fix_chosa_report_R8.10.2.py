# -*- coding: utf-8 -*-
"""調査結果報告書3文書の納品前処理（令和8年10月2日）

08_作業順位 の順位4。RedTeamレビューv8（令和8年9月25日）が
「納品前に必ず処理いただきたいもの」として挙げた4件は、
いずれも当方の作業で直せる（ご確認を待たない）。

  ① 目次が更新されていない（A ニーズ調査・C 本編）
     Word の目次フィールドのキャッシュが1行しか入っていない。
     `word/settings.xml` に `<w:updateFields w:val="true"/>` を置き、
     **Word で開いたときに目次が自動で更新される**ようにする。
     （頁番号は組版によるため、Word 又は Word 互換のソフトで
     開いた時点で確定する。当方の環境では確定できない。）
  ② `cp:lastModifiedBy`「山内 香」の残存（A・B・C）
     町名義で公表する文書であるため削除する。
  ③ ニーズ調査の編集メモ「削除も検討する」の残存
     レビューv8 の修正案の文に書き換える。
  ④ ニーズ調査 表73 のキャプションが表の下にある
     他の表はすべて表の上にあるため、上へ移す。

  あわせて、レビューv8【9】の「空の見出し」があれば取り除く。

  入力　令和8年9月25日受領の再修正版3文書
        **09_元資料/R8調査データ/R8.9.25受領版/（リポジトリの中）**
        ⚠ 受領当日のアップロード領域はセッションごとに作り直されるため、
        そこを入力にしていると次のセッションで作り直せない。
        原本はリポジトリに退避してあり、こちらを第一の入力とする。
  出力　01_第10期_最新版成果品/川崎町_ニーズ調査結果報告書_R8.10.2版.docx
        01_第10期_最新版成果品/川崎町_在宅介護実態調査結果報告書_R8.10.2版.docx
        01_第10期_最新版成果品/川崎町_在宅介護実態調査結果報告書_資料編_R8.10.2版.docx

**数値には一切触れない。** レビューv8 が挙げた数値・説明の修正4件と
注記の追加5件は、分母の整理（確認事項No.98・No.110・No.126）の
ご判断を要するため本スクリプトでは扱わない。
"""
import os
import re
import shutil
import sys
import zipfile

import docx
from docx.oxml.ns import qn

OUT_DIR = "01_第10期_最新版成果品"
# 受領した原本の置き場（リポジトリの中。これが第一の入力である）
KEEP_DIR = "09_元資料/R8調査データ/R8.9.25受領版"
# 受領した当日のアップロード領域。**セッションごとに作り直されるため、
# 次のセッションでは消えている。** KEEP_DIR に原本がないときだけ使う。
UPLOAD_DIR = "/root/.claude/uploads/4be8f82c-e2c9-52ac-b2a9-bad5154d0b13"

# （アップロード時のファイル名, 保存するファイル名, 略号, 段落数, 表数）
DOCS = [
    ("f2ba7be5-__________________.docx",
     "川崎町_ニーズ調査結果報告書_R8.10.2版.docx", "A", 659, 103),
    ("333220e2-______________________142______.docx",
     "川崎町_在宅介護実態調査結果報告書_資料編_R8.10.2版.docx", "B", 412, 57),
    ("8de2134d-________________________.docx",
     "川崎町_在宅介護実態調査結果報告書_R8.10.2版.docx", "C", 455, 51),
]


def keep_name(dst):
    """成果品の名から、09_元資料 に残した原本の名を作る。"""
    return os.path.join(KEEP_DIR, dst.replace("_R8.10.2版", "_R8.9.25受領版"))


def src_of(up, dst):
    """原本のパスと、どこから採ったかを返す。

    ⚠ **リポジトリの中（09_元資料）を第一の入力とする。**
    アップロード領域はセッションごとに作り直されるため、そこだけを
    入力にしていると、次のセッションでは成果品を作り直せなくなる。
    """
    keep = keep_name(dst)
    if os.path.exists(keep):
        return keep, "09_元資料（リポジトリ）"
    up = os.path.join(UPLOAD_DIR, up)
    if os.path.exists(up):
        return up, "アップロード領域（⚠ リポジトリへ退避します）"
    return None, None

MEMO_OLD = "※上記は第9期と第10期の配布数・回収数・回収率の違いなどから、"
MEMO_NEW = ("※ 第9期調査とは配布数・回収数・回収率が異なるため、"
            "上記の比較は傾向の参考にとどめ、"
            "増減の大きさをもって施策の効果を評価することはできない。")

UPDATE_FIELDS = '<w:updateFields w:val="true"/>'


def set_el(el, text):
    """段落の文字を入れ替える（最初の run の書式を残す）。"""
    rs = el.findall(qn("w:r"))
    if not rs:
        raise SystemExit("run のない段落は書き換えられない")
    for r in rs[1:]:
        el.remove(r)
    ts = rs[0].findall(qn("w:t"))
    for t in ts[1:]:
        rs[0].remove(t)
    if not ts:
        raise SystemExit("w:t のない run")
    ts[0].text = text
    ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def fix_zip(path, update_fields):
    """settings.xml に目次の自動更新を置き、作成者名を取り除く。"""
    tmp = path + ".tmp"
    done = {"fields": False, "author": False}
    with zipfile.ZipFile(path) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/settings.xml" and update_fields:
                s = data.decode("utf-8")
                if "updateFields" not in s:
                    s = re.sub(r"(<w:settings[^>]*>)", r"\1" + UPDATE_FIELDS,
                               s, count=1)
                    done["fields"] = True
                data = s.encode("utf-8")
            elif item.filename == "docProps/core.xml":
                s = data.decode("utf-8")
                s2 = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>",
                            "<cp:lastModifiedBy></cp:lastModifiedBy>", s)
                if s2 != s:
                    done["author"] = True
                data = s2.encode("utf-8")
            info = zipfile.ZipInfo(item.filename, item.date_time)
            info.compress_type = item.compress_type
            info.external_attr = item.external_attr
            zout.writestr(info, data)
    shutil.move(tmp, path)
    return done


def fix_body(path):
    """編集メモの書き換え・表73のキャプションの移動・空の見出しの削除。"""
    doc = docx.Document(path)
    body = doc.element.body
    kids = list(body.iterchildren())
    n_memo = n_cap = n_empty = 0

    for el in kids:
        if el.tag != qn("w:p"):
            continue
        t = "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()
        if t.startswith(MEMO_OLD):
            set_el(el, MEMO_NEW)
            n_memo += 1

    # 表73 のキャプションを、直前の表の上へ移す
    kids = list(body.iterchildren())
    for i, el in enumerate(kids):
        if el.tag != qn("w:p"):
            continue
        t = "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()
        if t.startswith("表73　") and i > 0 and kids[i - 1].tag == qn("w:tbl"):
            tbl = kids[i - 1]
            el.getparent().remove(el)
            tbl.addprevious(el)
            n_cap += 1
            break

    # 空の見出し（目次を更新すると空行として現れる）
    for el in list(body.iterchildren()):
        if el.tag != qn("w:p"):
            continue
        t = "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()
        if t:
            continue
        pr = el.find(qn("w:pPr"))
        if pr is None:
            continue
        st = pr.find(qn("w:pStyle"))
        if st is not None and "Heading" in str(st.get(qn("w:val"))):
            el.getparent().remove(el)
            n_empty += 1

    doc.save(path)
    return n_memo, n_cap, n_empty


def main():
    os.makedirs(KEEP_DIR, exist_ok=True)
    ng = []
    for src, dst, mark, n_p, n_t in DOCS:
        s, whence = src_of(src, dst)
        if s is None:
            ng.append(f"{mark} 原本がない：{keep_name(dst)}")
            continue
        # 受領した原本を 09_元資料 に残す（既にあればそれを使っている）
        keep = keep_name(dst)
        if not os.path.exists(keep):
            shutil.copy(s, keep)
        print(f"  {mark} 原本の出所：{whence}")
        out = os.path.join(OUT_DIR, dst)
        shutil.copy(s, out)

        d = docx.Document(out)
        if (len(d.paragraphs), len(d.tables)) != (n_p, n_t):
            ng.append(f"{mark} 規模が合わない："
                      f"{len(d.paragraphs)}段落{len(d.tables)}表"
                      f"（レビューv8 は{n_p}段落{n_t}表）")
            continue

        n_memo, n_cap, n_empty = fix_body(out)
        z = fix_zip(out, update_fields=(mark in ("A", "C")))
        print(f"  {mark} {dst}")
        print(f"     目次の自動更新 {'置いた' if z['fields'] else '―'}"
              f"／作成者名 {'取り除いた' if z['author'] else '―'}"
              f"／編集メモ {n_memo}件／表73のキャプション {n_cap}件"
              f"／空の見出し {n_empty}件")

    # ══════════════════════════ 自己点検
    print("  ── 自己点検")
    for src, dst, mark, n_p, n_t in DOCS:
        out = os.path.join(OUT_DIR, dst)
        if not os.path.exists(out):
            continue
        z = zipfile.ZipFile(out)
        core = z.read("docProps/core.xml").decode("utf-8")
        lm = re.findall(r"<cp:lastModifiedBy>(.*?)</cp:lastModifiedBy>", core)
        if any(x.strip() for x in lm):
            ng.append(f"{mark} 作成者名が残っている：{lm}")
        st = z.read("word/settings.xml").decode("utf-8")
        if mark in ("A", "C") and "updateFields" not in st:
            ng.append(f"{mark} 目次の自動更新が置かれていない")
        d = docx.Document(out)
        for q in d.paragraphs:
            if "削除も検討" in q.text:
                ng.append(f"{mark} 編集メモが残っている")
                break
        # 表73 のキャプションが表の上にあるか
        if mark == "A":
            kids = list(d.element.body.iterchildren())
            for i, el in enumerate(kids):
                if el.tag != qn("w:p"):
                    continue
                t = "".join(n.text or "" for n in el.iter(qn("w:t"))).strip()
                if t.startswith("表73　"):
                    nxt = kids[i + 1] if i + 1 < len(kids) else None
                    if nxt is None or nxt.tag != qn("w:tbl"):
                        ng.append("A 表73のキャプションの次が表でない")
                    break
    # 原本がリポジトリに残っていること（次のセッションでも作り直せる）
    for src, dst, mark, n_p, n_t in DOCS:
        if not os.path.exists(keep_name(dst)):
            ng.append(f"{mark} 原本が 09_元資料 に残っていない："
                      f"{keep_name(dst)}")
    if ng:
        for m in ng:
            print("   ×", m)
        sys.exit(1)
    print("   ○ 作成者名の残存なし／目次の自動更新あり（A・C）／"
          "編集メモの残存なし／表73のキャプションは表の上")
    print("   ○ 原本3点は 09_元資料/R8調査データ/R8.9.25受領版/ にある"
          "（アップロード領域が消えても作り直せる）")
    print("  ⚠ 目次の頁番号は Word で開いた時点で確定します"
          "（当方の環境では確定できません）。")


if __name__ == "__main__":
    print("調査結果報告書3文書の納品前処理")
    main()
