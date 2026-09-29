# -*- coding: utf-8 -*-
"""第2回策定委員会資料の自己点検（資料内の二値・計画素案との横断整合）

大雪広域連合の案件（ブランチ claude/大雪広域介護保険事業計画・f938242）が
外部レビューを受けて新設した自己点検を、委員会資料にも置いたもの。
計画素案については check_soan_selfcheck_R8.9.28.py が同じ役割を担う。

  検査1　交付金の数値が出典（001732614.xlsx）から再現できるか
  検査2　同じ事項に2つの値が資料内に置かれていないか
  検査3　計画素案と食い違う値がないか（横断）
  検査4　資料の中で件数を宣言している箇所と、実際の件数が合っているか

不適合があれば終了コード1を返す。

使い方：
  python3 07_ソーススクリプト/check_2kai_selfcheck_R8.9.28.py \
      03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v6.docx \
      01_第10期_最新版成果品/川崎町_計画書素案_v2.6_図表整理版.docx
"""
import json
import re
import sys

import docx
import openpyxl

XLSX = ("09_元資料/交付金評価/③令和８年度交付金評価指標等（市町村分・公表版）/"
        "001732614_令和８年度全国集計（市町村）.xlsx")
GROUPS = "04_調査・入力・分析/R8.9.25交付金再解析/交付金列定義_R8.9.28.json"
TAIKEI = ("05_試算・管理シート/"
          "川崎町_第10期_施策事業統合体系表_R8.9.25b.xlsx")
KAWASAKI = 280


def text_of(path):
    doc = docx.Document(path)
    out = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                out.append(cell.text)
    return doc, "\n".join(out)


def main(shiryo, soan):
    ws = openpyxl.load_workbook(XLSX, data_only=True,
                                read_only=True)["全国集計（市町村）"]
    rows = [r for r in ws.iter_rows(min_row=15, max_row=1797, values_only=True)
            if r[0] is not None and isinstance(r[335], (int, float))]
    kawa = [r for r in rows if r[0] == KAWASAKI][0]
    g = json.load(open(GROUPS, encoding="utf-8"))

    def gs(r, k):
        return sum(r[i] for i in g[k] if isinstance(r[i], (int, float)))

    def rank(v, vals):
        return sorted(vals, reverse=True).index(v) + 1

    doc, body = text_of(shiryo)
    _, sbody = text_of(soan)
    ng, ok = [], 0

    # ══════════ 検査1　交付金の数値
    tot = [r[335] for r in rows]
    want = [("令和8年度の得点", f"{kawa[335]}点"),
            ("全国順位", f"{rank(kawa[335], tot)}位"),
            ("全国平均", f"{sum(tot) / len(tot):.1f}点"),
            ("保険者数", f"{len(rows):,}保険者")]
    for name, key in (("体制・取組指標群", "TAI"), ("活動指標群", "KATSU"),
                      ("成果指標群", "SEI")):
        vals = [gs(r, key) for r in rows]
        mine = gs(kawa, key)
        want.append((f"{name}の得点", f"{mine}点"))
        want.append((f"{name}の全国順位", f"{rank(mine, vals):,}位"))
        want.append((f"{name}の全国平均", f"{sum(vals) / len(vals):.1f}点"))
    for name, val in want:
        if val.replace(",", "") in body.replace(",", ""):
            ok += 1
        else:
            ng.append(f"検査1　{name}「{val}」が資料に見当たらない")

    # ══════════ 検査2　資料内の二値
    PAIR = [
        ("準備基金152,500千円の保険料換算",
         re.compile(r"月額(?:約)?([\d,]+)円分?に相当")),
        ("入所・居住系の1人1月あたり給付費",
         re.compile(r"入所・居住系[^。]{0,40}?約(\d+\.\d)万円")),
        ("入所・居住系±10人の保険料への影響",
         re.compile(r"[＋▲]?(\d+)円[）]?\s*$|10人で月額約(\d+)円")),
        ("訂正事項の件数", re.compile(r"訂正事項(?:（|)([０-９\d５]+)件")),
        ("総事業数", re.compile(r"全[７7]章\d+節\d+項(\d+)事業")),
    ]
    for name, pat in PAIR:
        vals = {m.groups() for m in pat.finditer(body)}
        vals = {v for v in vals if any(x for x in v)}
        if len(vals) > 1:
            ng.append(f"検査2　{name} に2つ以上の値がある: {sorted(vals)}")
        elif vals:
            ok += 1

    # ══════════ 検査3　計画素案との横断
    CROSS = [
        ("基金の保険料換算", r"月額約([\d,]+)円分に相当"),
        ("算定A（取崩なし）", r"A? ?6,822円|6,822円"),
        ("総事業数", r"全[７7]章(\d+)節(\d+)項(\d+)事業"),
        ("調整交付金5.0％の影響", r"5\.0％[^。]{0,20}?[＋]?(\d+)円"),
        ("第1号被保険者負担割合24％の影響", r"24％[^。]{0,24}?[＋]?(\d+)円"),
    ]
    for name, pat in CROSS:
        a = {m.groups() for m in re.finditer(pat, body)}
        b = {m.groups() for m in re.finditer(pat, sbody)}
        if a and b and not (a & b):
            ng.append(f"検査3　{name} が計画素案と食い違う: 資料{sorted(a)} / 素案{sorted(b)}")
        elif a and b:
            ok += 1

    # ══════════ 検査4　宣言した件数と実際の件数
    m = re.search(r"訂正事項（([０-９\d]+)件）", body)
    if m:
        n = int(m.group(1).translate(str.maketrans("０１２３４５６７８９",
                                                   "0123456789")))
        # 訂正事項の表（区分／第1回のご説明／正しい内容／影響）の行数
        real = None
        for t in doc.tables:
            h = [c.text.strip() for c in t.rows[0].cells]
            if h[:2] == ["区　分", "第１回のご説明"]:
                real = len(t.rows) - 1
        if real is None:
            ng.append("検査4　訂正事項の表が見つからない")
        elif real != n:
            ng.append(f"検査4　訂正事項は「{n}件」と書いているが表は{real}行")
        else:
            ok += 1
    sentaku = len(re.findall(r"◆ 選択[１-５1-5]　", body))
    m = re.search(r"ご選択いただきたい事項（([０-９\d]+)件）", body)
    if m:
        n = int(m.group(1).translate(str.maketrans("０１２３４５６７８９",
                                                   "0123456789")))
        if sentaku != n:
            ng.append(f"検査4　選択は「{n}件」と書いているが本文には{sentaku}件")
        else:
            ok += 1

    # 町案から取り込む事業の件数と、列挙した事業の数が合っているか
    m = re.search(r"記載のない事業が(\d+)件（([^）]+)）", body)
    if m:
        n = int(m.group(1))
        real = len([x for x in m.group(2).split("、") if x.strip()])
        if n != real:
            ng.append(f"検査4　町案から取り込む事業は「{n}件」と書いているが"
                      f"、列挙は{real}件")
        else:
            ok += 1
    cnt = {int(x) for x in re.findall(r"記載のない事業が(\d+)件", body)}
    if len(cnt) > 1:
        ng.append(f"検査2　町案から取り込む事業の件数に2つの値がある: {sorted(cnt)}")
    # 統合体系表の「05_体系案との異同」シートと突き合わせる
    try:
        tw = openpyxl.load_workbook(TAIKEI, data_only=True,
                                    read_only=True)["05_体系案との異同"]
        real = sum(1 for r in tw.iter_rows(min_row=2, values_only=True)
                   if r and str(r[0]).strip() == "素案に取り込む事業")
        if cnt and real not in cnt:
            ng.append(f"検査3　町案から取り込む事業が資料は{sorted(cnt)}件、"
                      f"統合体系表の異同シートは{real}件")
        elif cnt:
            ok += 1
    except Exception as e:      # noqa: BLE001
        ng.append(f"検査3　統合体系表を読めない: {e}")

    print(f"自己点検 {shiryo}")
    print(f"  適合 {ok}件／不適合 {len(ng)}件")
    for m in ng:
        print("  ×", m)
    return 1 if ng else 0


if __name__ == "__main__":
    a = sys.argv[1] if len(sys.argv) > 1 else \
        "03_委員会・説明資料/川崎町_第2回策定委員会資料_R8.11_v6.docx"
    b = sys.argv[2] if len(sys.argv) > 2 else \
        "01_第10期_最新版成果品/川崎町_計画書素案_v2.6_図表整理版.docx"
    sys.exit(main(a, b))
