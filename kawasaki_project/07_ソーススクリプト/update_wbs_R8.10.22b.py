# -*- coding: utf-8 -*-
"""業務工程管理表の更新（方針Ｃの決定と、素案Ver.2.23・資料v10・問答集v6）
（令和8年10月22日・同日2度目）

  python3 07_ソーススクリプト/update_wbs_R8.10.22b.py

  ・県内他団体の前期計画は確かめずに進める（方針Ｃ・ご指示）
  ・県の入所状況調の結果を、素案 Ver.2.23・委員会資料 v10・
    想定問答集 v6 に同じ版で反映した

⚠ 列8（状態）には書き込まない。⚠ 同日2度目のため印を分ける。
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font

WBS = "川崎町_業務工程管理表.xlsx"
SOAN_OLD = "川崎町_計画書素案_v2.22_給付費整合版.docx"
SOAN = "川崎町_計画書素案_v2.23_県調査反映版.docx"
SHIRYO_OLD = "川崎町_第2回策定委員会資料_R8.11_v9.docx"
SHIRYO = "川崎町_第2回策定委員会資料_R8.11_v10.docx"
FONT = Font(name="游ゴシック", size=9)
HIZUKE = "【令和8年10月22日 更新②（方針Ｃ・素案Ver.2.23）】"


def iro(ws, r, n):
    for c in range(1, n + 1):
        ws.cell(r, c).font = FONT
        ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")


TSUIKI = {
    "154": (10, f"{HIZUKE}軽費老人ホームの休止・廃止を素案 Ver.2.23 の"
                f"2-7（第3の特徴）と別表 事項八に書きました（別表 八は"
                f"「未対応」→「一部」）。経緯と入居されていた方の行き先は"
                f"［要確認］として本文に残しています。　"),
    "159": (10, f"{HIZUKE}県の入所状況調による入所率・退所理由の推移を"
                f"素案 Ver.2.23 の 2-7・7-1 に入れました。入所前の居所は、"
                f"入所指針が定める保険者の協力の場で集計値を求めるかを"
                f"［要協議］として本文に置いています。ご判断をお願いします。　"),
}


def main():
    wb = openpyxl.load_workbook(WBS)
    ws = wb["03_確認事項一覧"]
    n = 0
    for x in range(5, ws.max_row + 1):
        no = str(ws.cell(x, 1).value or "").strip()
        if no in TSUIKI and HIZUKE not in str(ws.cell(x, 10).value or ""):
            retsu, moji = TSUIKI[no]
            ws.cell(x, retsu).value = moji + str(ws.cell(x, retsu).value or "")
            iro(ws, x, 10)
            n += 1

    ws = wb["04_成果品管理"]
    n_seika = 0
    for x in range(5, ws.max_row + 1):
        no = str(ws.cell(x, 1).value or "").strip()
        na = str(ws.cell(x, 3).value or "")
        if no == "4" and SOAN_OLD in na:
            ws.cell(x, 3).value = na.replace(SOAN_OLD, SOAN)
            ws.cell(x, 8).value = (
                f"{HIZUKE}素案を Ver.2.23 に上げた（全117頁）。県の入所状況調"
                f"（3か月ごと14時点）による入所率・退所理由の推移と、⚠ 町内の"
                f"軽費老人ホーム（定員70人）の休止・廃止を 2-7・7-1・別表に"
                f"入れた。［要確認］14→15件・［要協議］11→12件。　"
                + str(ws.cell(x, 8).value or ""))
            iro(ws, x, 9)
            n_seika += 1
        elif no == "6" and SHIRYO_OLD in na:
            ws.cell(x, 3).value = (na.replace(SHIRYO_OLD, SHIRYO)
                                   .replace("同_想定問答集_R8.11_v5.docx",
                                            "同_想定問答集_R8.11_v6.docx"))
            ws.cell(x, 8).value = (
                f"{HIZUKE}資料を v10（全10版・表45）、想定問答集を v6"
                f"（全64問）に上げた。素案 Ver.2.23 と同じ版で、入所率の推移の"
                f"表・軽費老人ホームの休止・廃止・残る4件の整理を入れた。"
                f"⚠ v9 の表6 の日付の書き誤り（元号と西暦の混在）を正した。　"
                + str(ws.cell(x, 8).value or ""))
            iro(ws, x, 9)
            n_seika += 1

    ws2 = wb["08_作業順位"]
    jun = [str(ws2.cell(x, 1).value or "").strip()
           for x in range(5, ws2.max_row + 1)]
    for x in range(5, ws2.max_row + 1):
        if str(ws2.cell(x, 1).value or "").strip() == "37":
            mae = str(ws2.cell(x, 9).value or "")
            if HIZUKE not in mae:
                ws2.cell(x, 9).value = (
                    mae + f"　{HIZUKE}⚠ **県内他団体の前期計画は確かめずに"
                    f"進める（方針Ｃ・ご指示）。** 想定問答集 v6 の QB-64 で、"
                    f"比較を行っていないことをそのまま答える形にした。"
                    f"県調査の結果を素案 Ver.2.23・資料 v10・問答集 v6 に"
                    f"同じ版で反映した。")
                iro(ws2, x, 9)
    wb.save(WBS)

    futeki = []

    def judge(ok, mi):
        print(("   ○ " if ok else "   × ") + mi)
        if not ok:
            futeki.append(mi)

    print("業務工程管理表の更新（方針Ｃ・素案Ver.2.23）")
    judge(n in (0, 2), f"確認事項2件に書いた（{n}件）")
    judge(n_seika == 2, "成果品管理の成果品4・6 を更新した")
    wb2 = openpyxl.load_workbook(WBS, data_only=True)
    ws3 = wb2["03_確認事項一覧"]
    jotai = {str(ws3.cell(x, 8).value or "") for x in range(5, ws3.max_row + 1)
             if str(ws3.cell(x, 1).value or "").strip() in TSUIKI}
    judge(jotai <= {"確認待ち", "決着", "保留", "回答済"},
          "⚠ 列8（状態）を動かしていない")
    nos = sorted(int(s) for s in (str(ws3.cell(x, 1).value or "").strip()
                                  for x in range(5, ws3.max_row + 1))
                 if s.isdigit())
    judge(nos == list(range(1, max(nos) + 1)), f"欠番0（{len(nos)}件）")
    s = "\n".join(str(v) for sh in wb2 for row in sh.iter_rows(values_only=True)
                  for v in row if v is not None)
    for na in ("釜房みどりの園", "アルパイン", "蔵王グリーンホーム",
               "ながさか", "リリーハイム", "川崎市"):
        judge(na not in s, f"「{na}」が入っていない")
    judge(SOAN in s and SHIRYO in s and "想定問答集_R8.11_v6" in s,
          "成果品管理に素案 Ver.2.23・資料 v10・問答集 v6 が載っている")
    if futeki:
        sys.exit(1)
    print("   ○ 自己点検 不適合0件")


if __name__ == "__main__":
    main()
