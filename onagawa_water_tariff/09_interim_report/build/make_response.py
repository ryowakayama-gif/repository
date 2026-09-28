# -*- coding: utf-8 -*-
"""中間報告書に対する指摘事項の対応表（Word）の作成"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
NAVY, BLUE, RED, GREY = RGBColor(0x1F,0x38,0x64), RGBColor(0x2E,0x75,0xB6), RGBColor(0xC0,0,0), RGBColor(0x59,0x59,0x59)
GREEN = RGBColor(0x37,0x56,0x23)
JP, JPG = 'ＭＳ 明朝', 'ＭＳ ゴシック'

doc = Document()
st = doc.styles['Normal']; st.font.name = JP; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), JP)
for s in doc.sections:
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    s.top_margin = s.bottom_margin = Cm(1.5); s.left_margin = s.right_margin = Cm(1.5)

def setfont(run, name=JP, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = color

def P(text='', size=10.5, bold=False, color=None, align=None, space_after=6, font=JP):
    p = doc.add_paragraph()
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.line_spacing = 1.15
    if text: setfont(p.add_run(text), font, size, bold, color)
    return p

def H1(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(7)
    setfont(p.add_run(text), JPG, 13.5, True, NAVY)
    pbd = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'12'); b.set(qn('w:space'),'2'); b.set(qn('w:color'),'1F3864')
    pbd.append(b); p._p.get_or_add_pPr().append(pbd); return p

def H2(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    setfont(p.add_run(text), JPG, 11.5, True, BLUE); return p

def shade(cell, hexcolor):
    e = OxmlElement('w:shd'); e.set(qn('w:val'),'clear'); e.set(qn('w:fill'),hexcolor)
    cell._tc.get_or_add_tcPr().append(e)

def table(headers, rows, widths=None, size=8.5, note=None, align_right=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths:
        t.autofit = False
        for i, w in enumerate(widths):
            for row in t.rows: row.cells[i].width = Cm(w)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ''
        pr = c.paragraphs[0]; pr.alignment = WD_ALIGN_PARAGRAPH.CENTER; pr.paragraph_format.space_after = Pt(2)
        setfont(pr.add_run(str(h)), JPG, size, True, RGBColor(0xFF,0xFF,0xFF)); shade(c, '2E75B6')
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ''
            pr = cells[i].paragraphs[0]; pr.paragraph_format.space_after = Pt(2)
            txt = str(v)
            if align_right and i in align_right: pr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif i == 0: pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else: pr.alignment = WD_ALIGN_PARAGRAPH.LEFT
            col = GREEN if txt.startswith('◎') else (RED if txt.startswith('△') else None)
            setfont(pr.add_run(txt.lstrip('★')), JP, size, txt.startswith('★'), col)
            if widths: cells[i].width = Cm(widths[i])
        if widths:
            for i, w in enumerate(widths): cells[i].width = Cm(w)
    if note: P(note, size=8, color=GREY, space_after=8)
    else: P('', size=4, space_after=4)
    return t

# ============================== 表紙相当 ==============================
P('令和8年度　上下水道経営指標評価書作成等業務委託', size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2, font=JPG)
P('ビズアップ公共コンサルティング株式会社', size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10)
H1('女川町水道料金改定の検討　中間報告書　指摘事項対応表')
P('中間報告書（令和8年9月）に対する内部確認で指摘された14件について、'
  '対応内容と反映箇所を整理したものである。'
  '対応の結果、必要改定率は1.92倍から1.99倍に、給水原価は262.9円/㎥から272.4円/㎥に変更している。', space_after=10)

H2('1　主要数値の変更')
table(['指標','修正前','修正後','変更の理由'],
 [['減価償却費（5年計）','1,961,661千円','2,002,979千円','既存資産分を財政シミュレーションの計上額に直接接続（R02）'],
  ['総括原価（5年計）','1,134,492千円','1,175,810千円','上記の反映'],
  ['給水原価','262.9円/㎥','272.4円/㎥','上記の反映'],
  ['総括原価回収割合','52.1％','50.3％','上記の反映'],
  ['年間不足額','108,729千円/年','116,993千円/年','上記の反映'],
  ['★必要改定率','★1.92倍','★1.99倍','★上記の反映'],
  ['繰入金を控除しない場合','2.38倍','2.45倍','上記の反映。上限値としての表示は取りやめ（R13）'],
  ['標準家庭（20mm・20㎥）','4,795円','4,969円','必要改定率の変更'],
  ['口径別3-A　標準家庭','4,020円','4,166円','総括原価に合わせた水準係数の再調整'],
  ['産業特例による減収','約72百万円/年','算定していない','母集団が接続しないため算定を取り下げ（R04）']],
 widths=[5.0,4.0,4.0,13.5], size=9)

H2('2　指摘事項ごとの対応')
ROWS = [
 ('R01','他会計補助金272,997千円の全額控除を裏付けられない',
  '◎ 反映。控除する前提は維持したうえで、「繰出基準の整理が成立する場合の試算」であることを明示。'
  '算式（人件費＋支払利息）と説明欄の不一致、二重控除の照合、条項別区分の未確認を記載し、'
  '控除の有無で必要改定率が約0.46倍動くことを感応度として掲載',
  '5-3、4-3注記、5-9、12-1',
  '繰出基準に基づく区分の町確認'),
 ('R02','減価償却費の差替えで既存資産分が41,318千円不足',
  '◎ 反映。既存資産分を財政シミュレーションの計上額（令和6年度末取得分1,822,881千円＋'
  '令和7年度取得分32,456千円）に直接接続し直し、総括原価以下の数値を全面的に更新',
  '4-2、5-5、全章の数値',
  '取得価額の税処理、新規分に対応する長期前受金戻入益、供用開始年度の再構築'),
 ('R03','資産維持費の対象資産896,312千円の母数と期間が未確定',
  '◎ 反映。対象資産の出所が平成31年度末の残高であること、期末残高が直接入力の推計値であること、'
  '対象期間の期末日が1年ずれていることを明示。対象資産±100,000千円で必要改定率±0.025倍を掲載',
  '5-4、4-1注記、5-9',
  '固定資産台帳による再構築'),
 ('R04','19社モデルは実績に基づく財政効果として使えない',
  '◎ 反映。19社の数値を「仮定した条件の下での個別事例」に限定し、'
  '合計額・減収額・3-Bの収入と回収割合の掲載を取りやめ。'
  '100mmで口径全体実績の約5.71倍、75mmで全体超過という不整合を明示',
  '5-7、8-3、9-3、9-5、10-1',
  '使用者別調定明細による再計算'),
 ('R05','税込・税抜の混在が財政効果を変える',
  '◎ 反映。料金表・利用者影響は税込、収支・総括原価は税抜と整理し、橋渡し表を新設。'
  '令和7年度の税抜単価116.6円/㎥を明示し、128.2円/㎥が税込調定単価であることを注記',
  '5-2、4-3注記、各表の注記',
  '決算給水収益との期間帰属・調定取消の調整'),
 ('R06','口径別料金案の原価回収が実データで未検証',
  '◎ 反映。代表口径1つを割り当てた試算であること、件数がサンプル集計であり全口径データと'
  '一致しないこと、100％回収は試算の係数を目標に合わせて解いた結果であることを明示。'
  '実データによる参考値（約2.10倍）も併記',
  '5-6、8-1、8-2、9-5注記',
  '調定明細による検証'),
 ('R07','「赤字半減」の定義が資金収支と異なる',
  '◎ 反映。定義③を「料金不足額」と改称し、資産維持費が現金支出でないこと、'
  '元金償還・建設改良費・期首資金を含まないことを明示。'
  '1.124772倍の丸めである旨、営業損失・経常損失を財政シミュレーションの対応額で再計算（▲400,051千円、▲49,846千円）',
  '7-1、7-2、5-8',
  '仕様書がいう「現状赤字額」の定義の確認'),
 ('R08','産業特例に料金の逆転がある',
  '◎ 反映。25mm・50mmの499㎥／500㎥の比較表を追加し、約6割下がることを明示。'
  '超過部分への適用、料金上限方式、減免率方式、年間認定の4方式を比較対象として提示。'
  '塩竈市の制度が特定協同組合への供給区分であり要件の先例ではないこと、'
  '115円が独自に丸めた設定値であることを明記',
  '8-3（1）、10-3',
  '採用方式の決定と再試算'),
 ('R09','利用者影響の説明が料金表と矛盾する',
  '◎ 反映。「10㎥以下は下がる」の記載を削除し、13mmで月3.3㎥付近が分岐点であること、'
  '月6〜10㎥は1,265円→1,465円に上がることを明示。'
  '口径×使用水量の倍率表を新設。特殊用途の比較表にメーター使用料の扱いを注記し、'
  '船舶用（委託）が未試算であることを明示。住民説明章でも標準料金表と特例適用後を分離',
  '8-2（3）（4）、9-1、9-4、9-6、11章',
  '値上げ・値下げ件数、増額帯別件数の集計'),
 ('R10','需要予測と実施時期が接続していない',
  '◎ 反映。令和7年度実績1,024,576㎥から令和8年度推計901千㎥への約12％の段差が未説明であること、'
  '需要予測が大口需要を区分していないことを明示。'
  '初年度の適用月数・期首資金・元金償還を含む年次資金収支を最終報告の作業として掲載',
  '3-2（4）、10-4',
  '年次資金収支の作成と感応度分析'),
 ('R11','料金回収率という用語の分母が異なる',
  '◎ 反映。第4章以降の指標名を「総括原価回収割合」に変更し、'
  '第3章の決算指標と分母が異なることを用語定義表で明示。'
  '住民説明章の「町の一般会計で補っている」の表現も改めた',
  '4-3、5-8、3-1、11章',
  '―'),
 ('R12','制度・比較・地域背景の説明に修正を要する',
  '◎ 反映。①3％を「算定要領の標準値」と改め、②算定期間の根拠条項を明示、'
  '③「県内では女川町のみ」を削除し気仙沼市の逓減を併記、④仕様書適合の最終判定を保留、'
  '⑤出島大橋開通（令和6年12月19日）を反映し江島と分離、'
  '⑥原発需要・全国の主流・受容性の断定を判断・仮説の表現に修正',
  '1-1、2-2、2-3、4-1、11章',
  '仕様書原本の確認、送水切替の状況確認'),
 ('R13','「2.38倍が上限」は撤回する',
  '◎ 反映。1.99〜2.45倍を「繰入金の控除条件のみを変えた比較値」と明記し、'
  '上限・下限を示すものではないこと、他の条件を併せて変えれば2.5倍を超える場合もあることを注記',
  '4-3注記、12-4',
  '―'),
 ('R14','根拠資料の版と計算再現性を整える',
  '◎ 反映。旧版の試算資料を再現用として配布しないこと、'
  '固定値で保存された倍率と再計算値の区別、参照エラーの修復、'
  '建設改良費調査票の財源構成欄の数式を転記しないこと、'
  '3-Aの単価の丸めによる差を条例化前に統一することを掲載。'
  '各数値に採用版・対象年度・税区分・原票・確認者・確認日を付す方針を明記',
  '12-3、8-2注記',
  '採用資料の一本化'),
]
table(['No','指摘の要点','対応内容','反映箇所','残る課題'], ROWS,
 widths=[1.2,5.4,11.6,3.4,4.9], size=8,
 note='※ ◎は本改訂で反映済み。「残る課題」は最終報告に向けて確定させる事項であり、'
      '工程表および中間報告書12-1・12-2に対応する項目を掲げている。')

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

H2('3　構成の変更')
P('前提条件と限界を1か所で確認できるよう、第5章「本報告の前提と限界」を新設した。'
  'これに伴い旧第5章以降を1章ずつ繰り下げている。')
table(['','改訂前','改訂後'],
 [['第4章','総括原価の算定','総括原価の算定（変更なし）'],
  ['★第5章','料金改定シミュレーション①','★本報告の前提と限界（新設）'],
  ['第6章','料金改定シミュレーション②','料金改定シミュレーション①'],
  ['第7章','料金改定シミュレーション③','料金改定シミュレーション②'],
  ['第8章','料金体系の見直しに伴う影響','料金改定シミュレーション③'],
  ['第9章','改定案の比較と段階改定・経過措置','料金体系の見直しに伴う影響'],
  ['第10章','住民・議会説明用の要約','改定案の比較と段階改定・経過措置'],
  ['第11章','今後の作業と確認事項','住民・議会説明用の要約'],
  ['★第12章','―','★今後の作業と確認事項']],
 widths=[2.6,10.0,10.0], size=9)

H2('4　あわせて更新した資料')
table(['資料','更新内容'],
 [['女川町水道料金改定_中間報告書.docx','全面改訂（上記1〜3）'],
  ['女川町水道料金改定_最終報告に向けた工程表.docx',
   '必要改定率の表示を1.99〜2.45倍に修正。作業項目に減価償却費の再構築、資産維持費の再算定、'
   '水需要予測の組替え、産業用特例の方式比較、年次資金収支の作成を追加。'
   'データ取得依頼に使用者別調定明細、業務仕様書原本、出島大橋経由への送水切替を追加'],
  ['算定の基礎データ（作業用ファイル）','減価償却費を2,002,979千円に修正し、各パターンの水準を再計算']],
 widths=[8.0,14.6], size=9)

P('', space_after=14)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
setfont(p.add_run('以　上'), JP, 10.5)

cp = doc.core_properties
cp.title = '女川町水道料金改定の検討　中間報告書　指摘事項対応表'
cp.subject = '令和8年度上下水道経営指標評価書作成等業務委託'
cp.author = '若山諒太'
cp.last_modified_by = '若山諒太'
cp.comments = 'ビズアップ公共コンサルティング株式会社'

OUT = os.path.join(HERE, '..', '女川町水道料金改定_中間報告書_指摘事項対応表.docx')
doc.save(OUT); print('作成:', os.path.abspath(OUT))
