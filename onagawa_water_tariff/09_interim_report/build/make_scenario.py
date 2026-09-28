# -*- coding: utf-8 -*-
"""試算条件切替計算表（Excel）の作成。

「01_設定」の黄色いセルを変えると、総括原価・必要改定率・利用者への影響まで
すべて Excel の数式で再計算される。openpyxl は計算結果を保存しないため、
Excel で開いた時点で計算される（プレビューでは空欄に見えることがある）。
"""
import os
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
SRC = os.path.join(ROOT, '06_capital_plan_update', '09_assumption_switch_cost_calc.xlsx')
OUT = os.path.join(HERE, '..', 'onagawa_scenario_calc.xlsx')

NAVY, BLUE, RED, ORANGE, GREEN = '1F3864', '2E75B6', 'C00000', 'C55A11', '375623'
GREY, LIGHT, YELLOW = '595959', 'DDEBF7', 'FFF2CC'
F = 'Arial'
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
IN_FILL = PatternFill('solid', fgColor=YELLOW)
OUT_FILL = PatternFill('solid', fgColor='E2EFDA')


def title(ws, row, text, span, size=12):
    ws.cell(row=row, column=1, value=text).font = Font(name=F, size=size, bold=True, color='FFFFFF')
    for i in range(1, span + 1):
        ws.cell(row=row, column=i).fill = PatternFill('solid', fgColor=NAVY)
    ws.row_dimensions[row].height = 22


def note(ws, row, text, color=GREY, size=9):
    ws.cell(row=row, column=1, value=text).font = Font(name=F, size=size, color=color)


def band(ws, row, text, span, color=BLUE):
    ws.cell(row=row, column=1, value=text).font = Font(name=F, size=10, bold=True, color='FFFFFF')
    for i in range(1, span + 1):
        ws.cell(row=row, column=i).fill = PatternFill('solid', fgColor=color)


def header(ws, row, headers, size=9):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=F, size=size, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=BLUE)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 26


def widths(ws, vals):
    for i, w in enumerate(vals, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, row, col, value, fmt=None, bold=False, color=None, fill=None,
        align=None, size=9, wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=F, size=size, bold=bold, color=color)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    c.alignment = Alignment(horizontal=align or ('right' if isinstance(value, (int, float)) else 'left'),
                            vertical='center', wrap_text=wrap)
    return c


# ============================================================
# 総括原価の数式（設定セル参照でも直接値でも組み立てられる）
# ============================================================
KEY_D = "'07_計算用データ'!$G$6:$G$113"
VAL_D = "'07_計算用データ'!$E$6:$E$113"
VAL_L = "'07_計算用データ'!$F$6:$F$113"
KEY_I = "'07_計算用データ'!$G$118:$G$183"
VAL_I = "'07_計算用データ'!$D$118:$D$183"


def q(v):
    """設定値を数式に埋め込む。セル参照はそのまま、文字列は引用符付きにする。"""
    if isinstance(v, str) and (v.startswith("'") or v.startswith('$') or v.startswith('C')):
        return v
    if isinstance(v, str):
        return '"' + v + '"'
    return str(v)


def dep_key(plan, method, kanro, kiden):
    return f'{q(plan)}&"|"&{q(method)}&"|"&{q(kanro)}&"|"&{q(kiden)}'


def int_key(plan, rate, shokan):
    return f'{q(plan)}&"|"&TEXT({q(rate)},"0.0")&"|"&{q(shokan)}'


def f_maint(doryoku):
    return f'672327+IF({q(doryoku)}="実績単価に補正",47167,0)'


def f_dep(plan, method, kanro, kiden, kison):
    k = dep_key(plan, method, kanro, kiden)
    new = f'INDEX({VAL_D},MATCH({k},{KEY_D},0))'
    return (f'IF({q(method)}="固定資産台帳ベース",32188*5+{new},'
            f'IF({q(kison)}="差引方式",1814019,1855337)+{new})')


def f_int_new(plan, rate, shokan):
    return f'44630+INDEX({VAL_I},MATCH({int_key(plan, rate, shokan)},{KEY_I},0))'


def f_zasshi(z):
    return f'IF({q(z)}="算入する",36696,0)'


def f_am(torikata, ritsu, nensu):
    return (f'ROUND(IF({q(torikata)}="期首の帳簿価額",1042625,896312)'
            f'*{q(ritsu)}/100,0)*{q(nensu)}')


def f_choki(plan, method, kanro, kiden, choki):
    k = dep_key(plan, method, kanro, kiden)
    real = f'1542259+INDEX({VAL_L},MATCH({k},{KEY_D},0))'
    return (f'-IF({q(method)}="固定資産台帳ベース",0,'
            f'IF({q(choki)}="経営戦略の実額",{real},'
            f'IF({q(choki)}="成果品の推計値",450000,0)))')


def f_hojo(h):
    return (f'-IF({q(h)}="他会計補助金のみ",272997,'
            f'IF({q(h)}="他会計補助金＋補助金",691997,'
            f'IF({q(h)}="成果品の推計値",15000,0)))')


def f_sonota(s):
    return f'-IF({q(s)}="控除する",26494,0)'


def total_cost(plan, kanro, kiden, rate, shokan, kison, method, ritsu, torikata,
               nensu, choki, hojo, sonota, zasshi, doryoku):
    parts = [f_maint(doryoku), f_dep(plan, method, kanro, kiden, kison), '50585',
             f_int_new(plan, rate, shokan), f_zasshi(zasshi),
             f_am(torikata, ritsu, nensu), f_choki(plan, method, kanro, kiden, choki),
             f_hojo(hojo), f_sonota(sonota)]
    return '=' + '+'.join(f'({p})' for p in parts)


# 04_料金への影響 の行位置（先に決めておき、03シートから参照する）
Z_FIRST, Z_LAST = 7, 14      # 水量ゾーン8行
HELP_ROW = Z_LAST + 1        # 月額合計（件数加重）
REV_ROW = HELP_ROW + 1       # 年間収入
RATIO_ROW = REV_ROW + 1      # 総括原価回収割合

# 01_設定 のセル参照は、シートを書き出したあとに実際の行から組み立てる
S = {}
VOL = REV = RATE_IN = M_A = None

wb = Workbook()

# ============================== 01_設定 ==============================
ws = wb.active
ws.title = '01_設定'
ws.sheet_view.showGridLines = False
widths(ws, [4, 34, 30, 4, 64])
title(ws, 1, '女川町水道事業　試算条件切替計算表', 5, 13)
note(ws, 2, '黄色のセルを変えると、02以降のシートがすべて自動で計算し直されます。'
            'Excelで開いてご利用ください。')
note(ws, 3, '既定値は中間報告書（改訂版）の採用条件です。総括原価1,175,810千円・必要改定率1.99倍を再現します。', RED)

SETTINGS = [
 ('【投資と減価償却】', None, None, None),
 ('建設改良費の計画', '建設改良費調査票',
  ['建設改良費調査票', '経営戦略の計画', '調査票＋量水器費'],
  '調査票＝女川町回答（R8〜R17で35.15億円）。量水器費は経営戦略の1.06億円'),
 ('耐用年数（管路・構築物）', 38, [35, 38, 40], '経営戦略の設定は38年'),
 ('耐用年数（機械・電気）', 38, [15, 20, 25, 30, 38, 40],
  '調査票は4事業とも40年。鷲神高度処理設備は機械・電気設備を含むため要確認'),
 ('既存資産の償却費', '財政シミュに接続', ['財政シミュに接続', '差引方式'],
  '接続＝1,855,337千円（R6末1,822,881＋R7取得32,456）。差引方式＝1,814,019千円（41,318千円不足）'),
 ('減価償却費の算入方法', '損益計算書ベース', ['損益計算書ベース', '固定資産台帳ベース'],
  '台帳ベースを選ぶと長期前受金戻入益の控除は0になります'),
 ('【企業債】', None, None, None),
 ('借入利率（％）', 2.1, [round(2.0 + 0.1 * i, 1) for i in range(11)],
  '調査票は4事業とも未記入。2.1％は経営戦略の加重平均借入利率'),
 ('償還条件', '経営戦略方式（40年・据置なし）',
  ['経営戦略方式（40年・据置なし）', '調査票方式（25/40年・据置3年）'], ''),
 ('【資産維持費】', None, None, None),
 ('資産維持率（％）', 3.0, None, '日本水道協会「水道料金算定要領」の標準値は3％'),
 ('対象資産の取り方', '期首期末平均', ['期首期末平均', '期首の帳簿価額'],
  '期首期末平均＝896,312千円　期首＝1,042,625千円。いずれも推計値で未確定'),
 ('資産維持費の年数', 5, [5],
  '5年の料金算定なので5で固定。他の費目・収益・水量がすべて5年計のため、'
  'ここだけ短くすると期間がそろわない'),
 ('【控除項目】', None, None, None),
 ('長期前受金戻入益', '経営戦略の実額',
  ['経営戦略の実額', '成果品の推計値', '控除しない'],
  '実額＝1,542,259千円＋新規取得分。推計値＝450,000千円'),
 ('他会計補助金等', '他会計補助金のみ',
  ['他会計補助金のみ', '他会計補助金＋補助金', '成果品の推計値', '控除しない'],
  '他会計補助金272,997千円。繰出基準に基づく区分は町に未確認'),
 ('その他営業収益・受取利息等', '控除しない', ['控除する', '控除しない'],
  '控除する場合26,494千円。経営戦略の原価計算表は控除していない'),
 ('雑支出', '算入しない', ['算入する', '算入しない'], '算入する場合36,696千円'),
 ('動力費・薬品費', '実績単価に補正', ['実績単価に補正', '経営戦略のまま'],
  '経営戦略は1年分の費用を10年間の水量で割っており過小。補正額＋47,167千円'),
 ('【水量と収益】', None, None, None),
 ('有収水量　5年計（㎥）', 4316000, None, '経営戦略の水需要予測'),
 ('給水収益　5年計（千円・税抜）', 590845, None,
  '現行料金による見込額。令和8年4〜6月の基本料金等免除は未反映'),
 ('【料金と収入の換算】', None, None, None),
 ('端数処理', '四捨五入', ['四捨五入', '1円未満切捨て'],
  '★既定は四捨五入（中間報告書と同じ）。町の現行料金は1円未満切捨てとされているため、'
  '条例化の際は切捨てに切り替えて確認する。差は1件あたり1円以内・年間で約2万円'),
 ('消費税率', 0.10, None, '料金表は税込、収支・総括原価は税抜。収入の換算に使う'),

 ('【料金改定の試算】', None, None, None),
 ('パターン1の改定率', '=総括原価回収に必要な率', None,
  '空欄にすると03シートの必要改定率を使います。数値を入れるとその率で試算します'),
 ('パターン3-Aの水準係数の決め方', '自動（総括原価に合わせる）',
  ['自動（総括原価に合わせる）', '手入力'],
  '自動にすると、総括原価を回収する水準に毎回そろえます'),
 ('　手入力する場合の係数', 0.976194, None,
  '上を「手入力」にしたときだけ使います。03シートに実際に使われた係数が出ます'),
]

r = 5
input_rows = {}
for label, default, choices, memo in SETTINGS:
    if default is None:
        band(ws, r, label, 5)
        r += 1
        continue
    put(ws, r, 2, label, size=10)
    if label == 'パターン1の改定率':
        c = put(ws, r, 3, None, fmt='0.000"倍"', fill=IN_FILL, bold=True,
                color=RED, align='center', size=10)
    else:
        c = put(ws, r, 3, default, fill=IN_FILL, bold=True, color=RED, align='center', size=10)
    if isinstance(default, float) and label.endswith('（％）'):
        c.number_format = '0.0'
    if label == 'パターン3-Aの水準係数':
        c.number_format = '0.000000'
    if label.endswith('（㎥）') or label.endswith('（千円・税抜）'):
        c.number_format = '#,##0'
    if choices:
        dv = DataValidation(type='list', formula1='"' + ','.join(str(x) for x in choices) + '"',
                            allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(c)
    put(ws, r, 5, memo, size=8.5, color=GREY, wrap=True)
    input_rows[label] = r
    r += 1

# 実際に書き出した行から参照を組み立てる
def ref(label):
    return f"'01_設定'!$C${input_rows[label]}"

S.update(plan=ref('建設改良費の計画'), kanro=ref('耐用年数（管路・構築物）'),
         kiden=ref('耐用年数（機械・電気）'), rate=ref('借入利率（％）'),
         shokan=ref('償還条件'), kison=ref('既存資産の償却費'),
         method=ref('減価償却費の算入方法'), ritsu=ref('資産維持率（％）'),
         torikata=ref('対象資産の取り方'), nensu=ref('資産維持費の年数'),
         choki=ref('長期前受金戻入益'), hojo=ref('他会計補助金等'),
         sonota=ref('その他営業収益・受取利息等'), zasshi=ref('雑支出'),
         doryoku=ref('動力費・薬品費'))
HASU = ref('端数処理')
TAX = ref('消費税率')
VOL = ref('有収水量　5年計（㎥）')
M_MODE = ref('パターン3-Aの水準係数の決め方')
M_MANUAL = ref('　手入力する場合の係数')
REV = ref('給水収益　5年計（千円・税抜）')
RATE_IN = ref('パターン1の改定率')
M_A = "'03_指標と改定率'!$C$15"

r += 1
put(ws, r, 2, '設定の確認', size=10, bold=True)
chk = (f'=IF(ISNA(MATCH({dep_key(S["plan"], S["method"], S["kanro"], S["kiden"])},{KEY_D},0)),'
       f'"★耐用年数の組合せが参照表にありません",'
       f'IF(ISNA(MATCH({int_key(S["plan"], S["rate"], S["shokan"])},{KEY_I},0)),'
       f'"★借入利率が参照表にありません（2.0〜3.0を0.1刻みで選んでください）",'
       f'IF(ROUND({S["rate"]},1)<>{S["rate"]},'
       f'"★借入利率は0.1刻みでしか参照表にないため "&TEXT({S["rate"]},"0.00")&"％ は "'
       f'&TEXT(ROUND({S["rate"]},1),"0.0")&"％ として計算しています",'
       f'"設定は参照表にあります")))')
put(ws, r, 3, chk, fill=OUT_FILL, bold=True, color=NAVY, size=9)
ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
r += 2

band(ws, r, '算定結果（詳しくは02・03シート）', 5, NAVY)
r += 1
RESULTS = [
 ('総括原価　5年計（千円）', "='02_総括原価'!C14", '#,##0'),
 ('給水原価（円/㎥）', "='03_指標と改定率'!C6", '#,##0.0'),
 ('供給単価　R8〜R12・税抜（円/㎥）', "='03_指標と改定率'!C7", '#,##0.0'),
 ('総括原価回収割合', "='03_指標と改定率'!C8", '0.0%'),
 ('年間不足額（千円/年）', "='03_指標と改定率'!C9", '#,##0'),
 ('必要改定率（倍）', "='03_指標と改定率'!C10", '0.000'),
 ('赤字半減の改定率（倍）', "='03_指標と改定率'!C13", '0.000'),
 ('標準家庭　20mm・20㎥（円/月・改定後）', None, '#,##0'),
]
WS_SET = ws
RESULT_ROW = {}
for label, formula, fmt in RESULTS:
    put(ws, r, 2, label, size=10, bold=True)
    put(ws, r, 3, formula, fmt=fmt, fill=OUT_FILL, bold=True, color=NAVY, align='center', size=10)
    RESULT_ROW[label] = r
    r += 1
r += 1
note(ws, r, '※ 数式の計算結果はファイルに保存していません。Excelで開くと計算されます。')
note(ws, r + 1, '※ 既定値のまま開くと 総括原価1,175,810千円 / 給水原価272.4円/㎥ / '
                '回収割合50.3% / 必要改定率1.990倍 / 標準家庭（パターン1）4,969円 になります。')
note(ws, r + 2, '※ 標準家庭を赤字半減案（1.125倍）で見ると2,809円です。上の欄はパターン1の額です。')
note(ws, r + 3, '※ 他会計補助金等を「控除しない」にすると 1,448,807千円 / 2.452倍 になります。')

# ============================== 02_総括原価 ==============================
ws = wb.create_sheet('02_総括原価')
ws.sheet_view.showGridLines = False
widths(ws, [4, 34, 18, 4, 70])
title(ws, 1, '総括原価の算定（令和8〜12年度 5年計・税抜・千円）', 5)
note(ws, 2, '「01_設定」の選択に基づいて計算しています。')
header(ws, 4, ['', '費目', '金額', '', '内容・変わる条件'])
COST = [
 ('Ⅰ', '維持管理費', '=' + f_maint(S['doryoku']), '動力費・薬品費の補正で＋47,167千円'),
 ('Ⅱ', '減価償却費', '=' + f_dep(S['plan'], S['method'], S['kanro'], S['kiden'], S['kison']),
  '既存資産分＋新規分。計画・耐用年数・算入方法・既存分の取り方で変わる'),
 ('Ⅲ', '支払利息（既往債）', '=50585', '固定'),
 ('Ⅳ', '支払利息（新規債）', '=' + f_int_new(S['plan'], S['rate'], S['shokan']),
  '既往計画分44,630千円＋新規分。借入利率と償還条件で変わる'),
 ('Ⅴ', '雑支出', '=' + f_zasshi(S['zasshi']), '算入する場合36,696千円'),
 ('Ⅵ', '資産維持費', '=' + f_am(S['torikata'], S['ritsu'], S['nensu']),
  '対象資産×資産維持率×年数。対象資産は未確定'),
 ('Ⅶ', '控除：長期前受金戻入益', '=' + f_choki(S['plan'], S['method'], S['kanro'], S['kiden'], S['choki']),
  '補助財源で取得した資産の償却見合い'),
 ('Ⅷ', '控除：他会計補助金等', '=' + f_hojo(S['hojo']), '繰出基準に基づく区分は町に未確認'),
 ('Ⅸ', '控除：その他営業収益・受取利息等', '=' + f_sonota(S['sonota']), '控除する場合26,494千円'),
]
r = 5
COST_ROW = {}
for no, name, formula, memo in COST:
    COST_ROW[name] = r
    put(ws, r, 1, no, align='center')
    put(ws, r, 2, name)
    put(ws, r, 3, formula, fmt='#,##0')
    put(ws, r, 5, memo, size=8.5, color=GREY, wrap=True)
    r += 1
put(ws, r, 2, '総括原価　合計', bold=True, size=10)
put(ws, r, 3, f'=SUM(C5:C{r-1})', fmt='#,##0', bold=True, color=NAVY, fill=OUT_FILL, size=10)
TOTAL_ROW = r
r += 2
band(ws, r, '適用中の条件', 5)
r += 1
for label, ref in [('建設改良費の計画', S['plan']), ('耐用年数（管路・構築物）', S['kanro']),
                   ('耐用年数（機械・電気）', S['kiden']), ('既存資産の償却費', S['kison']),
                   ('減価償却費の算入方法', S['method']), ('借入利率（％）', S['rate']),
                   ('償還条件', S['shokan']), ('資産維持率（％）', S['ritsu']),
                   ('対象資産の取り方', S['torikata']), ('長期前受金戻入益', S['choki']),
                   ('他会計補助金等', S['hojo']), ('動力費・薬品費', S['doryoku'])]:
    put(ws, r, 2, label, size=8.5)
    put(ws, r, 3, '=' + ref, align='center', size=8.5)
    r += 1

# ============================== 03_指標と改定率 ==============================
ws = wb.create_sheet('03_指標と改定率')
ws.sheet_view.showGridLines = False
widths(ws, [4, 34, 18, 4, 70])
title(ws, 1, '給水原価・総括原価回収割合・必要改定率', 5)
note(ws, 2, '供給単価は税区分と対象年度で3種類ある。同じ指標として比較しないこと。')
header(ws, 4, ['', '指標', '算定値', '', '算定式・注意'])
TC = f"'02_総括原価'!C{TOTAL_ROW}"
IND = [
 ('総括原価　5年計（千円）', f'={TC}', '#,##0', '02シートの合計'),
 ('給水原価（円/㎥）', f'={TC}*1000/{VOL}', '#,##0.0', '総括原価×1,000÷有収水量'),
 ('供給単価　R8〜R12・税抜（円/㎥）', f'={REV}*1000/{VOL}', '#,##0.0', '給水収益×1,000÷有収水量'),
 ('総括原価回収割合', f'=C7/C6', '0.0%', '供給単価÷給水原価。決算の料金回収率とは分母が異なる'),
 ('年間不足額（千円/年）', f'=({TC}-{REV})/5', '#,##0', '（総括原価−給水収益）÷5年'),
 ('必要改定率（倍）', f'={TC}/{REV}', '0.000',
  '総括原価÷給水収益。★控除条件の前提つきの試算値であり確定水準ではない'),
 ('料金不足額　定義③（千円/年）', None, '#,##0',
  '給水収益＋基準内繰入金−（維持管理費＋支払利息＋資産維持費）÷5年。元金償還・建設改良費を含まない'),
 ('　半減に必要な増収額（千円/年）', '=-C11/2', '#,##0', ''),
 ('赤字半減の改定率（倍）', f'=1+C12/({REV}/5)', '0.000', '定義③の半減'),
 ('パターン1に適用する改定率（倍）', f'=IF(ISNUMBER({RATE_IN}),{RATE_IN},C10)', '0.000',
  '01_設定に数値を入れるとその率。空欄なら必要改定率'),
 ('3-Aに適用する水準係数', None, '0.000000',
  '自動の場合は「必要改定率×現行の月額合計÷水準係数1のときの月額合計」で求める。'
  'この式なら税区分と母集団の補正が相殺されるため、収入の推計方法に左右されない'),
]
r = 5
for name, formula, fmt, memo in IND:
    put(ws, r, 2, name, size=10, bold=name.startswith('必要') or name.startswith('総括原価回収'))
    r += 1
r = 5
for name, formula, fmt, memo in IND:
    if name.startswith('3-Aに適用する水準係数'):
        formula = (f'=IF({M_MODE}="手入力",{M_MANUAL},'
                   f"C10*'04_料金への影響'!$F${HELP_ROW}/'04_料金への影響'!$N${HELP_ROW})")
    elif name.startswith('料金不足額'):
        formula = (f'={REV}/5'
                   f'+IF({S["hojo"]}="控除しない",0,-(\'02_総括原価\'!C{COST_ROW["控除：他会計補助金等"]})/5)'
                   f'-(\'02_総括原価\'!C{COST_ROW["維持管理費"]}/5'
                   f'+\'02_総括原価\'!C{COST_ROW["支払利息（既往債）"]}/5'
                   f'+\'02_総括原価\'!C{COST_ROW["支払利息（新規債）"]}/5'
                   f'+\'02_総括原価\'!C{COST_ROW["資産維持費"]}/5)')
    put(ws, r, 3, formula, fmt=fmt, fill=OUT_FILL, bold=True, color=NAVY, align='center', size=10)
    put(ws, r, 5, memo, size=8.5, color=GREY, wrap=True)
    r += 1
r += 1
band(ws, r, '参考：段階改定（適用開始は令和9年4月以降）', 5)
r += 1
header(ws, r, ['', '段階', '改定率', '', '内容'])
r += 1
for name, formula, memo in [
        ('第1次（議決 令和9年3月）', '=C13', '定義③の料金不足額の半減'),
        ('第2次（令和13〜17年度）', '=C10', '総括原価の回収。その時点の投資計画で再算定')]:
    put(ws, r, 2, name)
    put(ws, r, 3, formula, fmt='0.000', align='center')
    put(ws, r, 5, memo, size=8.5, color=GREY)
    r += 1
r += 1
note(ws, r, '※ 改定率は通年ベース。令和8年度の料金は現行のままで、増収は令和9年度以降に生じる。', RED)

# ============================== 04_料金への影響 ==============================
ws = wb.create_sheet('04_料金への影響')
ws.sheet_view.showGridLines = False
widths(ws, [4, 18, 10, 8, 10, 13, 10, 13, 10, 13, 10, 4, 44, 12])
title(ws, 1, '利用者への影響（月額・税込）と年間収入（税抜）', 13)
note(ws, 2, 'パターン1は現行料金に改定率を一律に乗じたもの。3-Aは口径別基本料金＋段階逓増。')
note(ws, 3, '★件数・月平均水量は代表値による推計であり、原票との照合は未了。'
            '全使用者の調定明細による検証を経ていない。', RED)
note(ws, 4, '月額は税込。年間収入は税抜で、現行は給水収益（01_設定）の年平均をそのまま起点にしている。')

RATE = "'03_指標と改定率'!$C$14"
ZONES = [('① 0㎥', 0.0, 297, 13), ('② 1〜5㎥', 2.9, 513, 13), ('③ 6〜10㎥', 8.0, 590, 13),
         ('④ 11〜20㎥', 14.9, 856, 20), ('⑤ 21〜50㎥', 28.7, 672, 20),
         ('⑥ 51〜100㎥', 71.6, 55, 25), ('⑦ 101〜500㎥', 197.5, 55, 25),
         ('⑧ 501㎥超', 1030.5, 18, 50)]
METER = {13: 55, 20: 77, 25: 110, 40: 220, 50: 880, 75: 1100, 100: 1320}
BASE_A = {13: 770, 20: 1430, 25: 2310, 40: 5500, 50: 11000, 75: 22000, 100: 41800}
TIER = [(10, 91.3), (20, 192.5), (50, 258.5), (100, 280.5), (None, 324.5)]


def cur_formula(vcell, dia):
    return (f'IF({vcell}<=5,990,1210)+IF({vcell}>10,(MIN({vcell},100)-10)*121'
            f'+MAX(0,{vcell}-100)*110,0)+{METER[dia]}')


def p3a_unit_formula(vcell, dia):
    """水準係数を1としたときの3-A（丸めなし）。水準係数の算定に使う。"""
    tiers = (f'MIN({vcell},10)*91.3'
             f'+MAX(0,MIN({vcell},20)-10)*192.5'
             f'+MAX(0,MIN({vcell},50)-20)*258.5'
             f'+MAX(0,MIN({vcell},100)-50)*280.5'
             f'+MAX(0,{vcell}-100)*324.5')
    return f'{BASE_A[dia]}+({tiers})'


def p3a_formula(vcell, dia):
    tiers = (f'MIN({vcell},10)*91.3'
             f'+MAX(0,MIN({vcell},20)-10)*192.5'
             f'+MAX(0,MIN({vcell},50)-20)*258.5'
             f'+MAX(0,MIN({vcell},100)-50)*280.5'
             f'+MAX(0,{vcell}-100)*324.5')
    x = f'({BASE_A[dia]}+({tiers}))*{M_A}'
    return f'IF({HASU}="四捨五入",ROUND({x},0),ROUNDDOWN({x},0))'


header(ws, 6, ['', '水量ゾーン', '月平均\n水量', '件数', '適用\n口径', '現行', '', 'パターン1\n改定後',
               '現行比', 'パターン3-A', '現行比', '', '備考', '計算用\n3-A係数1'])
r = Z_FIRST
for name, vol, cnt, dia in ZONES:
    put(ws, r, 2, name)
    put(ws, r, 3, vol, fmt='#,##0.0', align='center')
    put(ws, r, 4, cnt, fmt='#,##0', align='center')
    put(ws, r, 5, f'{dia}mm', align='center')
    put(ws, r, 6, '=' + cur_formula(f'C{r}', dia), fmt='#,##0')
    put(ws, r, 8, f'=IF({HASU}="四捨五入",ROUND(F{r}*{RATE},0),ROUNDDOWN(F{r}*{RATE},0))', fmt='#,##0')
    put(ws, r, 9, f'=H{r}/F{r}', fmt='0.00"倍"', align='center')
    put(ws, r, 10, '=' + p3a_formula(f'C{r}', dia), fmt='#,##0')
    put(ws, r, 11, f'=J{r}/F{r}', fmt='0.00"倍"', align='center')
    put(ws, r, 14, '=' + p3a_unit_formula(f'C{r}', dia), fmt='#,##0.0', size=8, color=GREY)
    r += 1

# --- 月額合計（件数加重）。3-Aの水準係数を決めるために使う ---
put(ws, r, 2, '月額合計（件数加重・円）', bold=True, size=9)
put(ws, r, 4, f'=SUM(D{Z_FIRST}:D{Z_LAST})', fmt='#,##0', bold=True, align='center')
for col in (6, 8, 10, 14):
    L = get_column_letter(col)
    put(ws, r, col, f'=SUMPRODUCT($D${Z_FIRST}:$D${Z_LAST},{L}{Z_FIRST}:{L}{Z_LAST})',
        fmt='#,##0', size=9)
put(ws, r, 13, 'N列は水準係数を1としたときの3-A。03シートの水準係数の算定に使う',
    size=8.5, color=GREY, wrap=True)
r += 1

# --- 年間収入。起点を給水収益に接続する ---
put(ws, r, 2, '年間収入（百万円・税抜）', bold=True, size=9.5)
put(ws, r, 6, f'={REV}/5/1000', fmt='#,##0.0', bold=True, color=NAVY, fill=OUT_FILL)
put(ws, r, 8, f'=F{r}*{RATE}', fmt='#,##0.0', bold=True, color=NAVY, fill=OUT_FILL)
put(ws, r, 10, f'=F{r}*J{HELP_ROW}/F{HELP_ROW}', fmt='#,##0.0', bold=True, color=NAVY, fill=OUT_FILL)
put(ws, r, 13, '現行は給水収益（税抜）の年平均そのもの。パターン1は改定率を乗じ、'
               '3-Aはゾーン別試算の月額合計の比を乗じている', size=8.5, color=GREY, wrap=True)
r += 1

put(ws, r, 2, '総括原価回収割合', bold=True, size=9.5)
for col in (6, 8, 10):
    L = get_column_letter(col)
    put(ws, r, col, f"={L}{REV_ROW}/('03_指標と改定率'!$C$5/5/1000)",
        fmt='0.0%', bold=True, color=NAVY, fill=OUT_FILL)
put(ws, r, 13, '年間収入÷総括原価の年平均。水準係数が自動なら3-Aは100％になる',
    size=8.5, color=GREY, wrap=True)
r += 2

# --- ゾーン別試算と給水収益の関係（診断） ---
band(ws, r, 'ゾーン別試算と給水収益の関係（診断）', 14)
r += 1
DIAG = [
 ('ゾーン別試算の年額（税込）', f'=F{HELP_ROW}*12', '#,##0',
  '代表値×件数×12か月。全使用者の一部しか拾えていない'),
 ('　税抜に換算', f'=F{HELP_ROW}*12/(1+{TAX})', '#,##0', '（1＋消費税率）で割った額'),
 ('給水収益の年平均（税抜）', f'={REV}/5*1000', '#,##0', '01_設定の入力値'),
 ('含意される母集団補正率', None, '0.00000',
  '★給水収益÷税抜換算額。1.15前後なら、ゾーン別試算が拾えていない分が約15％あるということ。'
  '別データによる検証を経ていない'),
]
DIAG_FIRST = r
for name, formula, fmt, memo in DIAG:
    put(ws, r, 2, name, size=9)
    if formula is None:
        formula = f'=F{DIAG_FIRST+2}/F{DIAG_FIRST+1}'
    put(ws, r, 6, formula, fmt=fmt, size=9, fill=OUT_FILL)
    put(ws, r, 13, memo, size=8.5, color=(RED if memo.startswith('★') else GREY), wrap=True)
    r += 1
r += 1

band(ws, r, '標準家庭と特殊用途', 13)
r += 1
header(ws, r, ['', '区分', '月使用量', '', '口径', '現行', '', 'パターン1', '現行比',
               'パターン3-A', '現行比', '', '備考'])
r += 1
STD_ROW = r
BIZ_ROW = r + 3          # 中規模事業者（50mm・1,000㎥）の行
SPECIAL = [('標準家庭', 20, 20, '中間報告書の標準家庭'),
           ('少量使用（13mm・8㎥）', 8, 13, '3-Aでも現行より上がる'),
           ('13mm・3㎥', 3, 13, '3-Aで現行を下回るのはこの付近まで'),
           ('中規模事業者', 1000, 50, '仮定した事例。実在の事業者ではない'),
           ('大規模事業者', 3000, 75, '仮定した事例')]
for name, vol, dia, memo in SPECIAL:
    d = dia if dia in BASE_A else 100
    put(ws, r, 2, name)
    put(ws, r, 3, vol, fmt='#,##0', align='center')
    put(ws, r, 5, f'{dia}mm', align='center')
    put(ws, r, 6, '=' + cur_formula(f'C{r}', dia), fmt='#,##0')
    put(ws, r, 8, f'=IF({HASU}="四捨五入",ROUND(F{r}*{RATE},0),ROUNDDOWN(F{r}*{RATE},0))', fmt='#,##0')
    put(ws, r, 9, f'=H{r}/F{r}', fmt='0.00"倍"', align='center')
    put(ws, r, 10, '=' + p3a_formula(f'C{r}', d), fmt='#,##0')
    put(ws, r, 11, f'=J{r}/F{r}', fmt='0.00"倍"', align='center')
    put(ws, r, 13, memo, size=8.5, color=GREY, wrap=True)
    r += 1
# 01_設定 の結果欄に標準家庭（改定後）を接続する
WS_SET.cell(row=RESULT_ROW['標準家庭　20mm・20㎥（円/月・改定後）'], column=3,
            value=f"='04_料金への影響'!H{STD_ROW}")

r += 1
note(ws, r, '※ 湯屋用・船舶用は用途別の定額設計のため本表に含めていない（中間報告書9-4を参照）。')
note(ws, r + 1, '※ 現行料金の端数処理は1円未満の切捨て。条例化の際に試算とそろえる。')

# ============================== 05_口径別料金表 ==============================
ws = wb.create_sheet('05_口径別料金表')
ws.sheet_view.showGridLines = False
widths(ws, [4, 20, 16, 16, 4, 60])
title(ws, 1, 'パターン3-A　口径別料金表（税込・円）', 6)
note(ws, 2, '塩竈市の体系比率に「01_設定」の水準係数を乗じたもの。')
note(ws, 3, '★この係数は総括原価と一致するよう逆算した設定値であり、実際の調定データで検証した値ではない。', RED)
note(ws, 4, '※ 掲載額は設計値として四捨五入している。請求額の端数処理は「01_設定」の設定に従う（04シート）。')
header(ws, 6, ['', '口径', '現行メーター使用料', '案3-A 基本料金', '', '備考'])
r = 7
BASE_ROW = {}
for k in [13, 20, 25, 40, 50, 75, 100]:
    BASE_ROW[k] = r
    put(ws, r, 2, f'{k}mm', align='center')
    put(ws, r, 3, METER[k], fmt='#,##0')
    put(ws, r, 4, f'=ROUND({BASE_A[k]}*{M_A},0)', fmt='#,##0')
    r += 1
r += 1
header(ws, r, ['', '水量区分', '現行（一般用）', '案3-A 従量料金', '', '備考'])
r += 1
for (upper, rate), cur, memo in zip(TIER,
                                    ['基本水量', '121円/㎥', '121円/㎥', '121円/㎥', '110円/㎥（逓減）'],
                                    ['', '', '', '', '現行は100㎥超で単価が下がる']):
    lbl = f'〜{upper}㎥' if upper else '101㎥超'
    put(ws, r, 2, lbl, align='center')
    put(ws, r, 3, cur, align='center')
    put(ws, r, 4, f'={rate}*{M_A}', fmt='#,##0.0')
    put(ws, r, 6, memo, size=8.5, color=GREY)
    r += 1
r += 2
band(ws, r, '水準係数の調整', 6)
r += 1
put(ws, r, 2, '決め方', size=10)
put(ws, r, 3, f'={M_MODE}', align='center', fill=OUT_FILL, bold=True)
r += 1
put(ws, r, 2, '実際に使っている係数', size=10, bold=True)
put(ws, r, 3, f'={M_A}', fmt='0.000000', align='center', fill=OUT_FILL, bold=True, color=NAVY)
put(ws, r, 6, '「自動」なら総括原価を回収する水準に毎回そろえます。'
              '「手入力」にすると01_設定に入れた値をそのまま使うため、'
              '条件を変えても追随しません（04シートの回収割合で確認してください）。',
    size=8.5, color=GREY, wrap=True)
r += 1
put(ws, r, 2, '3-Aの年間収入（百万円）', size=10)
put(ws, r, 3, "='04_料金への影響'!J" + str(REV_ROW), fmt='#,##0.0', align='center', fill=OUT_FILL)
r += 1
put(ws, r, 2, '総括原価（年平均・百万円）', size=10)
put(ws, r, 3, "='03_指標と改定率'!C5/5/1000", fmt='#,##0.0', align='center', fill=OUT_FILL)
r += 1
put(ws, r, 2, '3-Aの回収割合', size=10, bold=True)
put(ws, r, 3, f'=C{r-2}/C{r-1}', fmt='0.0%', align='center', fill=OUT_FILL, bold=True, color=NAVY)
put(ws, r, 6, '100％から離れている場合は、上の決め方が「手入力」になっています。', size=8.5, color=GREY)
r += 2

band(ws, r, '掲載額と請求額の端数の差', 6)
r += 1
note(ws, r, '掲載額は基本料金を円単位・従量単価を小数第1位で丸めている。'
            '請求額は丸める前の単価で計算するため、両者は完全には一致しない。')
r += 1
header(ws, r, ['', '例', '掲載表から積み上げた額', '請求額（04シート）', '', '差'])
r += 1
put(ws, r, 2, '50mm・1,000㎥', align='center')
put(ws, r, 3, f'=D{BASE_ROW[50]}+10*89.1+10*187.9+30*252.3+50*273.8+900*316.8', fmt='#,##0')
put(ws, r, 4, f"='04_料金への影響'!J{BIZ_ROW}", fmt='#,##0')
put(ws, r, 6, f'=C{r}-D{r}', fmt='+#,##0"円";-#,##0"円"', align='center', bold=True, color=RED)
r += 2
note(ws, r, '※ 条例化の際は、料金表の額と請求額の端数処理を同一の規定にそろえる。')
r += 1
note(ws, r, '※ 塩竈市の基本料金比率：13mm 770／20mm 1,430／25mm 2,310／40mm 5,500／'
            '50mm 11,000／75mm 22,000／100mm 41,800（税込×1.1）')
note(ws, r + 1, '※ 従量料金の基礎単価：〜10㎥ 91.3／〜20㎥ 192.5／〜50㎥ 258.5／'
                '〜100㎥ 280.5／101㎥超 324.5円/㎥')

# ============================== 06_ケース比較 ==============================
ws = wb.create_sheet('06_ケース比較')
ws.sheet_view.showGridLines = False
widths(ws, [4, 40, 16, 14, 12, 14, 4, 46])
title(ws, 1, '主要ケースの比較', 8)
note(ws, 2, '費目の条件（C列の総括原価）は各行に書いた条件だけで計算しており、'
            '「01_設定」の費目設定を変えても動きません。')
note(ws, 3, 'ただし給水原価・回収割合・必要改定率（D〜F列）は「01_設定」の有収水量と給水収益を'
            '参照するため、そちらを変えると全行が動きます。', RED)

D = dict(plan='建設改良費調査票', kanro=38, kiden=38, rate=2.1,
         shokan='経営戦略方式（40年・据置なし）', kison='財政シミュに接続',
         method='損益計算書ベース', ritsu=3.0, torikata='期首期末平均', nensu=5,
         choki='経営戦略の実額', hojo='他会計補助金のみ', sonota='控除しない',
         zasshi='算入しない', doryoku='実績単価に補正')


def case(**over):
    p = dict(D)
    p.update(over)
    return total_cost(**p)


CASES = [
 ('★採用条件（中間報告書の改訂版）', {}, '基準内繰入金を控除。減価償却費は財政シミュに接続'),
 ('繰入金を控除しない', dict(hojo='控除しない'), '控除の根拠が確認できない場合'),
 ('他会計補助金＋補助金を控除', dict(hojo='他会計補助金＋補助金'), '補助金419,000千円も控除した場合'),
 ('既存資産の償却費を差引方式に戻す', dict(kison='差引方式'), '改訂前の算定。41,318千円不足する'),
 ('機械・電気を20年償却', dict(kiden=20), '鷲神高度処理設備の区分が確定した場合'),
 ('機械・電気を20年＋繰入金を控除しない', dict(kiden=20, hojo='控除しない'),
  'この2つを同時に変えた場合。表中で最も大きくなるのは長期前受金を推計値に戻した行'),
 ('借入利率2.0％', dict(rate=2.0), ''),
 ('借入利率3.0％', dict(rate=3.0), ''),
 ('資産維持率1.0％', dict(ritsu=1.0), ''),
 ('資産維持率5.0％', dict(ritsu=5.0), ''),
 ('対象資産を期首の帳簿価額に', dict(torikata='期首の帳簿価額'), '1,042,625千円を用いた場合'),
 ('量水器費を計上', dict(plan='調査票＋量水器費'), '経営戦略の1.06億円を戻す'),
 ('減価償却を固定資産台帳ベースに', dict(method='固定資産台帳ベース'), '長期前受金の控除は0になる'),
 ('長期前受金を成果品の推計値に', dict(choki='成果品の推計値'), '450,000千円。改訂前の前提'),
 ('動力費・薬品費を補正しない', dict(doryoku='経営戦略のまま'), '経営戦略の単価をそのまま使う'),
 ('その他営業収益等を控除し雑支出を算入', dict(sonota='控除する', zasshi='算入する'), ''),
]
header(ws, 5, ['', 'ケース', '総括原価\n5年計（千円）', '給水原価\n（円/㎥）', '回収割合',
               '必要改定率\n（倍）', '', '内容'])
r = 6
for name, over, memo in CASES:
    put(ws, r, 2, name, bold=name.startswith('★'), color=(NAVY if name.startswith('★') else None))
    put(ws, r, 3, case(**over), fmt='#,##0')
    put(ws, r, 4, f'=C{r}*1000/{VOL}', fmt='#,##0.0')
    put(ws, r, 5, f'={REV}*1000/{VOL}/D{r}', fmt='0.0%', align='center')
    put(ws, r, 6, f'=C{r}/{REV}', fmt='0.000', align='center',
        bold=True, color=NAVY)
    put(ws, r, 8, memo, size=8.5, color=GREY, wrap=True)
    if name.startswith('★'):
        for col in range(2, 7):
            ws.cell(row=r, column=col).fill = OUT_FILL
    r += 1
r += 1
note(ws, r, '※ 各行は他の費目条件を採用条件に固定したうえで、記載の条件だけを変えた値である。'
            '複数の条件が同時に動く場合は単純な加算にならない。')
note(ws, r + 1, '※ 必要改定率は給水収益に対する倍率であり、確定した改定水準ではない。')
note(ws, r + 2, '※ この表は総括原価の比較であり、3-Aの水準係数や利用者への影響は含まない。')

# ============================== 07_計算用データ ==============================
src = openpyxl.load_workbook(SRC)
s = src['05_計算用データ']
ws = wb.create_sheet('07_計算用データ')
ws.sheet_view.showGridLines = False
widths(ws, [22, 24, 12, 12, 14, 16, 46])
for row in range(1, s.max_row + 1):
    for col in range(1, 8):
        v = s.cell(row=row, column=col).value
        if v is None:
            continue
        c = ws.cell(row=row, column=col, value=v)
        c.font = Font(name=F, size=8.5)
        if row <= 2 or (isinstance(v, str) and v.startswith('【')):
            c.font = Font(name=F, size=9.5, bold=True, color=NAVY)
ws.cell(row=2, column=1,
        value='「02_総括原価」「06_ケース比較」が参照します。直接編集しないでください。').font = \
    Font(name=F, size=9, color=RED)
ws.sheet_state = 'visible'

# ============================== 08_参照元 ==============================
ws = wb.create_sheet('08_参照元')
ws.sheet_view.showGridLines = False
widths(ws, [4, 30, 18, 4, 70])
title(ws, 1, '固定値の出所', 5)
header(ws, 4, ['', '項目', '値（千円）', '', '出所・注意'])
REFS = [
 ('維持管理費（基礎）', 672327, '経営戦略の財政シミュレーション'),
 ('　動力費・薬品費の補正', 47167, '1年分の費用を10年間の水量で割る誤りの補正。R8〜R12の5年計'),
 ('既存資産の償却費（接続）', 1855337, 'R6末取得資産1,822,881＋R7取得資産32,456'),
 ('既存資産の償却費（差引方式）', 1814019, '2,032,440−218,421。改訂前の算定で41,318千円不足'),
 ('支払利息（既往債）', 50585, '経営戦略'),
 ('支払利息（新規債の既往計画分）', 44630, '経営戦略'),
 ('長期前受金戻入益（既往分）', 1542259, '経営戦略の実額'),
 ('他会計補助金', 272997, '★繰出基準に基づく区分は町に未確認'),
 ('　補助金', 419000, '算定根拠が確認できないため既定では控除しない'),
 ('その他営業収益・受取利息等', 26494, '経営戦略の原価計算表は控除していない'),
 ('雑支出', 36696, '経営戦略の原価計算表は算入していない'),
 ('対象資産（期首期末平均）', 896312, '★平成31年度末の残高に由来する推計値'),
 ('対象資産（期首の帳簿価額）', 1042625, '★同上'),
 ('固定資産台帳ベースの既存償却費（年）', 32188, '経営戦略'),
 ('有収水量　5年計（㎥）', 4316000, '経営戦略の水需要予測'),
 ('給水収益　5年計', 590845, '★令和8年4〜6月の基本料金等免除は未反映'),
 ('母集団補正率', 1.15145, '★ゾーン別試算の税抜収入を給水収益に合わせた値。別データによる検証が必要'),
 ('消費税率', 0.10, '料金表は税込、収支・総括原価は税抜'),
]
r = 5
for name, val, memo in REFS:
    put(ws, r, 2, name)
    put(ws, r, 3, val, fmt='#,##0')
    put(ws, r, 5, memo, size=8.5, color=(RED if memo.startswith('★') else GREY), wrap=True)
    r += 1
r += 1
note(ws, r, '★印は未確定。町の資料を受領して確定させる（中間報告書 第5章・第12章）。', RED)
r += 2
band(ws, r, 'この計算表でできないこと', 5)
r += 1
for t in ['使用者別の調定明細による検証。件数・口径分布は代表値による推計のまま',
          '母集団補正率1.15145の妥当性。ゾーン別試算が全使用者の一部しか拾えていないための補正で、'
          '別データによる検証を経ていない',
          '産業用特例（3-B）の減収額。対象者の母集団が接続しないため算定していない',
          '年次の資金収支。企業債元金償還・建設改良費・期首資金を含む資金残高',
          '令和8年4〜6月の基本料金等免除を織り込んだ令和8・9年度の実収入',
          '湯屋用・船舶用の影響（中間報告書9-4を参照）']:
    put(ws, r, 2, '・' + t, size=9)
    ws.cell(row=r, column=2).border = Border()
    r += 1

wb.properties.title = '女川町水道事業　試算条件切替計算表'
wb.properties.subject = '令和8年度上下水道経営指標評価書作成等業務委託'
wb.properties.creator = '若山諒太'
wb.properties.lastModifiedBy = '若山諒太'
wb.properties.category = 'ビズアップ公共コンサルティング株式会社'
wb.save(OUT)
print('作成:', os.path.abspath(OUT))
print('シート:', wb.sheetnames)
