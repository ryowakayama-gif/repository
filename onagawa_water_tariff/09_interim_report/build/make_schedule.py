# -*- coding: utf-8 -*-
"""最終報告書に向けた工程表の作成"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
NAVY, BLUE, RED, GREEN, ORANGE, GREY = (RGBColor(0x1F,0x38,0x64), RGBColor(0x2E,0x75,0xB6),
    RGBColor(0xC0,0,0), RGBColor(0x37,0x56,0x23), RGBColor(0xC5,0x5A,0x11), RGBColor(0x59,0x59,0x59))
JP, JPG = 'ＭＳ 明朝', 'ＭＳ ゴシック'

doc = Document()
st = doc.styles['Normal']; st.font.name = JP; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), JP)
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
sec.top_margin = sec.bottom_margin = Cm(1.6); sec.left_margin = sec.right_margin = Cm(1.6)

def setfont(run, name=JP, size=None, bold=None, color=None):
    run.font.name = name; run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = color

def P(text='', size=10.5, bold=False, color=None, align=None, space_after=6, font=JP):
    p = doc.add_paragraph()
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.line_spacing = 1.15
    if text: setfont(p.add_run(text), font, size, bold, color)
    return p

def H1(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(7)
    setfont(p.add_run(text), JPG, 13, True, NAVY)
    pbd = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'12'); b.set(qn('w:space'),'2'); b.set(qn('w:color'),'1F3864')
    pbd.append(b); p._p.get_or_add_pPr().append(pbd); return p

def H2(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    setfont(p.add_run(text), JPG, 11, True, BLUE); return p

def bullet(text, size=10):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3)
    setfont(p.add_run(text), JP, size); return p

def shade(cell, hexcolor):
    e = OxmlElement('w:shd'); e.set(qn('w:val'),'clear'); e.set(qn('w:fill'),hexcolor)
    cell._tc.get_or_add_tcPr().append(e)

def table(headers, rows, widths, size=9, note=None, hdr_color='2E75B6'):
    t = doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i,h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text=''; c.width = Cm(widths[i])
        pr = c.paragraphs[0]; pr.alignment = WD_ALIGN_PARAGRAPH.CENTER; pr.paragraph_format.space_after = Pt(1)
        setfont(pr.add_run(str(h)), JPG, size, True, RGBColor(0xFF,0xFF,0xFF)); shade(c, hdr_color)
    for r in rows:
        cells = t.add_row().cells
        for i,v in enumerate(r):
            cells[i].text=''; cells[i].width = Cm(widths[i])
            pr = cells[i].paragraphs[0]; pr.paragraph_format.space_after = Pt(1)
            txt=str(v)
            pr.alignment = WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
            bold = txt.startswith('★')
            setfont(pr.add_run(txt.lstrip('★')), JP, size, bold)
    if note: P(note, size=8.5, color=GREY, space_after=8)
    else: P('', size=4, space_after=4)
    return t

# ============================== 表題 ==============================
P('令和8年度　上下水道経営指標評価書作成等業務委託', size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
P('女川町上水道事業　水道料金改定の検討', size=13, bold=True, color=NAVY,
  align=WD_ALIGN_PARAGRAPH.CENTER, font=JPG, space_after=2)
P('最終報告書に向けた工程表', size=17, bold=True, color=NAVY,
  align=WD_ALIGN_PARAGRAPH.CENTER, font=JPG, space_after=6)
P('令和8年9月　ビズアップ公共コンサルティング株式会社', size=10,
  align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=10)

# ============================== 1 全体像 ==============================
H1('1　全体像とマイルストーン')
table(['時期','マイルストーン','内容'],
 [['令和8年9月末','★中間報告','業務仕様書の3パターンのシミュレーションと影響のとりまとめを提出（済）'],
  ['令和8年12月','総括原価の確定','決算確定値・固定資産台帳・実使用水量を反映し、必要改定率を確定'],
  ['令和9年1月','改定案の決定','A〜E案から採用案を絞り込み、町内調整を完了'],
  ['令和9年1〜2月','住民・事業者説明','説明会の開催、水産加工業者への個別説明'],
  ['令和9年2月末','★最終報告','最終報告書の提出'],
  ['令和9年3月','条例改正','議会への条例改正案の提出・議決'],
  ['令和9年4月以降','改定の実施','周知期間を経て新料金の適用開始']],
 [3.0,3.6,19.4], size=9.5,
 note='※ 業務の履行期限は令和9年3月末である。'
      '「改定案の決定・条例の議決」と「新料金の適用開始」は別の時点であり、'
      '適用開始は令和9年4月以降となる。したがって令和8年度の料金は現行のままであり、'
      '改定による増収は令和9年度以降に生じる。中間報告書の増収額は通年ベースの試算値である。')

# ============================== 2 工程表 ==============================
H1('2　工程表')
M = ['R8.10','R8.11','R8.12','R9.1','R9.2','R9.3']
FILL_MAIN, FILL_SUB, FILL_MS = '2E75B6', 'BDD7EE', 'C00000'
TASKS = [
 ('【データの確定】', None, None, None),
 ('　決算確定値（R6・R7）の受領と反映', '町→当方', [1,1,0,0,0,0], None),
 ('　固定資産台帳による対象資産の確定', '町→当方', [1,1,0,0,0,0], None),
 ('　建設改良費調査票の未記入事項の確定', '町', [1,0,0,0,0,0], None),
 ('　繰出基準の整理（基準内繰入金の確定）', '町', [1,1,0,0,0,0], None),
 ('　使用者別の調定明細（口径・月別水量）の取得', '町→当方', [1,1,0,0,0,0], None),
 ('　口径別の調定件数・水量の実績集計', '当方', [0,1,1,0,0,0], None),
 ('　業務仕様書原本による要件の照合', '町→当方', [1,0,0,0,0,0], None),
 ('【総括原価と料金表の確定】', None, None, None),
 ('　採用計算表の整備（本文の全指標を再現できる形に一本化）', '当方', [1,1,0,0,0,0], None),
 ('　減価償却費の再構築（税処理・長期前受金・供用開始年度）', '当方', [0,1,1,0,0,0], None),
 ('　資産維持費の対象資産の再算定（台帳との突合）', '当方', [0,1,1,0,0,0], None),
 ('　水需要予測の組替え（R7実績からの接続）', '当方', [0,1,1,0,0,0], None),
 ('　総括原価の再計算・必要改定率の確定', '当方', [0,1,1,0,0,0], None),
 ('　現行体系維持案の改定率の確定', '当方', [0,0,1,0,0,0], None),
 ('　赤字半減案の確定', '当方', [0,0,1,0,0,0], None),
 ('　口径別料金表案の精緻化（3-A・3-B）', '当方', [0,0,1,1,0,0], None),
 ('　産業用特例の方式の比較と採用方式による再試算', '当方', [0,0,1,1,0,0], None),
 ('　利用者影響の再計算（調定明細に各案を適用）', '当方', [0,0,1,1,0,0], None),
 ('　改定実施時期を反映した年次資金収支の作成', '当方', [0,0,0,1,0,0], None),
 ('　特殊用途区分（湯屋用・船舶用）の取扱い整理', '協議', [0,0,1,1,0,0], None),
 ('【改定案の決定と説明】', None, None, None),
 ('　改定案の絞り込み・町内調整', '協議', [0,0,1,1,0,0], None),
 ('　段階改定・経過措置の設計', '当方', [0,0,0,1,0,0], None),
 ('　住民・議会説明用資料の作成', '当方', [0,0,0,1,0,0], None),
 ('　水産加工業者への個別説明', '町', [0,0,0,1,1,0], None),
 ('　住民説明会の開催', '町', [0,0,0,1,1,0], None),
 ('【報告と条例改正】', None, None, None),
 ('　最終報告書の作成', '当方', [0,0,0,1,1,0], None),
 ('　★最終報告書の提出', '当方', [0,0,0,0,2,0], None),
 ('　条例改正案の策定', '町', [0,0,0,0,1,1], None),
 ('　★議会提出・議決', '町', [0,0,0,0,0,2], None),
 ('　改定の周知・システム改修', '町', [0,0,0,0,0,1], None),
]
t = doc.add_table(rows=1, cols=2+len(M)); t.style='Table Grid'
t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
W=[10.4,2.0]+[2.2]*len(M)
for i,h in enumerate(['作業項目','担当']+M):
    c=t.rows[0].cells[i]; c.text=''; c.width=Cm(W[i])
    pr=c.paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.CENTER; pr.paragraph_format.space_after=Pt(1)
    setfont(pr.add_run(h), JPG, 9, True, RGBColor(0xFF,0xFF,0xFF)); shade(c,'1F3864')
for name, owner, bars, _ in TASKS:
    cells=t.add_row().cells
    for i in range(2+len(M)): cells[i].width=Cm(W[i]); cells[i].text=''
    pr=cells[0].paragraphs[0]; pr.paragraph_format.space_after=Pt(1)
    if bars is None:
        setfont(pr.add_run(name), JPG, 9, True, NAVY)
        for i in range(2+len(M)): shade(cells[i],'DDEBF7')
        continue
    setfont(pr.add_run(name), JP, 9)
    pr=cells[1].paragraphs[0]; pr.alignment=WD_ALIGN_PARAGRAPH.CENTER; pr.paragraph_format.space_after=Pt(1)
    setfont(pr.add_run(owner), JP, 8.5, color=(RED if owner=='町' else (GREEN if owner=='協議' else None)))
    for j,v in enumerate(bars):
        c=cells[2+j]
        if v==1: shade(c, FILL_SUB)
        elif v==2:
            shade(c, FILL_MS)
            prc=c.paragraphs[0]; prc.alignment=WD_ALIGN_PARAGRAPH.CENTER; prc.paragraph_format.space_after=Pt(1)
            setfont(prc.add_run('▲'), JPG, 9, True, RGBColor(0xFF,0xFF,0xFF))
P('', size=4, space_after=4)
P('　　凡例：薄青＝作業期間　／　▲（赤）＝提出・議決　／　担当「町」＝女川町、「当方」＝ビズアップ公共コンサルティング、「協議」＝両者',
  size=8.5, color=GREY, space_after=10)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ============================== 3 月別の作業内容 ==============================
H1('3　月別の作業内容')
table(['時期','女川町にお願いする事項','当方で実施する作業','この時点での到達点'],
 [['令和8年10月',
   '・令和6年度・令和7年度決算の確定値の提供\n'
   '・令和8年4月1日時点の固定資産台帳残高の提供\n'
   '・建設改良費調査票の未記入事項の確定\n'
   '　（企業債の借入予定利率、償却年数の区分、\n'
   '　　江島海底送水管の資産種別等）\n'
   '・繰出基準に基づく基準内繰入金の整理',
   '・提供データの確認と整理\n'
   '・調査票の投資計画の確定\n'
   '・水産加工業の使用水量データの分析準備',
   '総括原価の算定に必要なデータが揃う'],
  ['令和8年11月',
   '・水産加工業者の事業者別・月別使用水量の提供\n'
   '・構造物及び設備の内訳の確認\n'
   '・量水器更新費の計上方針の決定',
   '・総括原価の再計算\n'
   '・必要改定率の確定\n'
   '・口径別の調定件数・水量の実績集計',
   '必要改定率が確定する'],
  ['令和8年12月',
   '・改定案の方向性についての内部協議',
   '・3パターンの数値の確定\n'
   '・口径別料金表案の精緻化\n'
   '・水産加工業影響の再試算\n'
   '・特殊用途区分の取扱いの整理',
   '3パターンの最終的な数値が固まる'],
  ['令和9年1月',
   '・改定案の決定\n'
   '・住民説明会の日程調整・開催通知\n'
   '・水産加工業者への個別説明の調整',
   '・段階改定・経過措置の設計\n'
   '・住民・議会説明用資料の作成\n'
   '・最終報告書の執筆開始',
   '採用する改定案が決まる'],
  ['令和9年2月',
   '・住民説明会の開催\n'
   '・水産加工業者への個別説明\n'
   '・条例改正案の策定着手',
   '・説明会での意見を踏まえた修正\n'
   '・最終報告書の完成・提出',
   '★最終報告書の提出'],
  ['令和9年3月',
   '・議会への条例改正案の提出\n'
   '・議決後の周知・システム改修の手配',
   '・議会説明資料の補助\n'
   '・条例改正案に関する技術的助言',
   '★条例改正の議決']],
 [2.4,7.6,7.4,5.0], size=8.5)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ============================== 4 クリティカルパス ==============================
H1('4　工程上の要点')
H2('4-1　クリティカルパス')
P('次の3点が遅れると、以降の工程が連鎖的に遅れる。', space_after=4)
table(['順','項目','期限','遅延した場合の影響'],
 [['1','決算確定値・固定資産台帳残高の提供','令和8年10月末',
   '総括原価が確定せず、必要改定率も確定しない。11月以降の全工程が止まる'],
  ['2','繰出基準に基づく基準内繰入金の整理','令和8年11月末',
   '必要改定率が1.99倍から2.45倍まで動くため、改定案の比較ができない'],
  ['3','改定案の決定','令和9年1月末',
   '住民説明・条例改正案の策定に入れず、令和9年3月議会に間に合わなくなる']],
 [1.2,7.0,3.0,15.0], size=9)

H2('4-2　並行して進められる作業')
for t in ['水産加工業の実使用水量の分析は、決算値の確定を待たずに着手できる',
          '口径別料金表案の骨格（体系の形）は既に設計済みであり、水準の調整のみが残る',
          '住民・議会説明用資料は中間報告書第11章を土台にできる。'
          'ただし数値の差替えだけでは足りず、繰入金控除の前提、口径別案の回収割合が'
          '試算上の設定値であること、産業用特例が未設計であることを、'
          '採用案の確定に合わせて書き改める必要がある',
          '条例改正案の条文構成は、改定案の決定前でも他団体の例を参照して準備できる']: bullet(t)

H2('4-3　工程が厳しくなった場合の対応')
table(['状況','対応の考え方'],
 [['調定明細の取得が遅れる',
   '現行体系を維持するA案・B案の条件付き比較までを最終報告の範囲とする。'
   '口径別へ移行するC案・D案の採用決定と条例用の料金表は、'
   '調定明細による検証の完了を条件とし、最終報告では参考案として位置づける。'
   '代替データを用いる場合は、対象率・欠損・想定される誤差・決裁の条件を明示する'],
  ['繰出基準の整理が固まらない',
   '基準内繰入金を控除する場合（1.99倍）と控除しない場合（2.45倍）の両論併記で最終報告を行う。'
   '中間報告と同じ扱いとなる'],
  ['令和9年3月議会に間に合わない',
   '条例改正を令和9年6月議会に送り、改定の実施時期を後ろ倒しする。'
   'その場合、令和9年度の収支見通しを再算定する必要がある'],
  ['改定幅について合意が得られない',
   '第1次改定の幅を小さくし、段階改定の回数を増やす。'
   'ただし総括原価の回収時期が後ろ倒しになるため、資金残高への影響を別途示す']],
 [6.0,20.0], size=9)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ============================== 5 データ取得依頼一覧 ==============================
H1('5　女川町にお願いするデータ・確認事項の一覧')
table(['No','事項','内容','期限','用途'],
 [['1','令和6年度・令和7年度決算の確定値',
   '損益計算書・貸借対照表の確定値。特に維持管理費、減価償却費、長期前受金戻入益、'
   '他会計補助金、支払利息','令和8年10月末','総括原価の確定'],
  ['2','固定資産台帳残高',
   '令和8年4月1日時点の帳簿価額（取水・浄水・配水の区分別）','令和8年10月末','資産維持費の対象資産'],
  ['3','企業債の借入予定利率','建設改良費調査票の4事業分','令和8年10月末','支払利息の算定'],
  ['4','償却年数の区分',
   '鷲神高度処理設備（機械・電気設備）、江島海底送水管の償却年数','令和8年10月末','減価償却費の算定'],
  ['5','江島海底送水管の属性','資産種別・更新／新規の別・補助／単独の別','令和8年10月末','施設区分別の集計'],
  ['6','基準内繰入金の整理',
   '繰出基準に基づき一般会計が負担すべき額。特に他会計補助金272,997千円および'
   '補助金419,000千円の性格','令和8年11月末','総括原価からの控除額'],
  ['7','構造物及び設備の内訳',
   '経営戦略26.4億円と調査票21.65億円の差4.75億円の内容','令和8年11月末','投資規模の確定'],
  ['8','量水器更新費の計上方針',
   '経営戦略の1.06億円を計上するか否か','令和8年11月末','投資規模の確定'],
  ['9','使用者別の調定明細',
   '使用者番号（匿名可）・業種・口径・用途・月別水量・調定額・減免の別。'
   '直近1〜3年度分。水産加工業を含む全使用者','令和8年11月末',
   '口径別料金案の検証、特例対象の特定、影響件数の集計'],
  ['10','湯屋用・船舶用の実績',
   '対象件数・使用水量・調定額。メーター使用料の取扱い（免除の有無）、'
   '船舶用（委託）の請求単位と対象者','令和8年11月末','特殊用途区分の取扱いの検討'],
  ['11','業務仕様書の原本',
   '資産維持率3％の要求水準、「現状赤字額」の定義、報告様式の要件','令和8年10月末',
   '仕様書適合の最終確認'],
  ['12','出島大橋経由への送水切替',
   '切替の完了時期、切替後の維持管理費、既存の海底送水管等の取扱い','令和8年11月末',
   '維持管理費と資産の前提'],
  ['13','令和8年度の基本料金等の免除',
   '令和8年4〜6月請求分の免除の内容（対象・免除額・件数）、'
   '交付金等による補填収入の額と会計処理','令和8年11月末',
   '令和8年度の収入見込みと資金残高'],
  ['14','端数処理の取扱い',
   '現行の端数処理（1円未満の切捨て）の適用範囲。'
   '基本料金・従量料金・メーター使用料それぞれの扱い','令和8年11月末',
   '条例案と試算の一致'],
  ['15','改定案の決定','A〜E案から採用する案の決定','令和9年1月末','最終報告書・条例改正案']],
 [1.0,4.6,9.6,2.8,4.4], size=8.5,
 note='※ 1〜6・11は総括原価と仕様書適合の確定に不可欠であり、最優先でお願いしたい。'
      '9は口径別料金案の検証に必須であり、これがない場合は代表値による推計のまま'
      '最終報告を行うこととなる。')
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ============================== 6 成果品 ==============================
H1('6　最終報告書の想定構成と中間報告からの主な変更点')
table(['章','内容','中間報告からの変更点'],
 [['第1章','業務の目的と検討範囲','★仕様書原本との照合結果を反映'],
  ['第2章','女川町水道事業の概要と現行料金体系','★出島大橋経由への送水切替の状況を反映'],
  ['第3章','財務・経営指標からみた課題','★令和6・7年度決算の確定値、組み替えた水需要予測を反映'],
  ['第4章','総括原価の算定','★決算確定値・固定資産台帳・基準内繰入金を反映し数値を確定'],
  ['第5章','前提と限界','★確定した事項を本文へ移し、残る限界のみを記載'],
  ['第6章','料金改定シミュレーション①現行体系維持案','★必要改定率を確定値で再算定'],
  ['第7章','料金改定シミュレーション②赤字半減案','★「現状赤字額」の定義を確定して再算定'],
  ['第8章','料金改定シミュレーション③口径別料金表案',
   '★調定明細により料金表を検証。産業用特例は境界の逆転が生じない方式へ'],
  ['第9章','料金体系の見直しに伴う影響',
   '★調定明細に基づく再計算。影響件数・増額帯別件数、特例の減収額を算定'],
  ['第10章','改定案の比較と段階改定・経過措置','★採用案を明示。年次の資金収支と感応度を追加'],
  ['第11章','住民・議会説明用の要約','★説明会での意見を踏まえて修正'],
  ['第12章','推奨案と実施に向けた留意事項','★新設。条例改正の論点、システム改修、周知方法を記載'],
  ['参考','県内料金体系統一化の動向','★新設。別途整理した検討資料を参考として添付']],
 [2.0,8.6,15.8], size=9,
 note='※ ★は中間報告から実質的な変更が生じる章。'
      '中間報告の第5章（前提と限界）に掲げた事項の確定状況を、最終報告で一覧として示す。')

H2('あわせて作成する資料')
table(['資料','内容','提出時期'],
 [['住民説明会用資料','中間報告書第11章を基に、図表中心の説明資料として作成','令和9年1月'],
  ['議会説明資料','改定の必要性・改定案・影響・経過措置を簡潔にまとめたもの','令和9年2月'],
  ['水産加工業者向け説明資料','規模別の影響額と経過措置の適用イメージ','令和9年1月'],
  ['条例改正案の参考資料','料金表の新旧対照、他団体の条文例','令和9年2月'],
  ['料金シミュレーション表（最終版）','前提条件を切り替えて再計算できる形の計算表','令和9年2月']],
 [5.4,14.6,6.4], size=9)

P('', space_after=14)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
setfont(p.add_run('以　上'), JP, 10.5)

OUT = os.path.join(HERE, '..', '女川町水道料金改定_最終報告に向けた工程表.docx')
cp = doc.core_properties
cp.title = '女川町水道料金改定の検討　最終報告に向けた工程表'
cp.subject = '令和8年度上下水道経営指標評価書作成等業務委託'
cp.author = '若山諒太'
cp.last_modified_by = '若山諒太'
cp.comments = 'ビズアップ公共コンサルティング株式会社'
doc.save(OUT); print('作成:', os.path.abspath(OUT))
