# -*- coding: utf-8 -*-
"""WBS（Excel）の作成。"""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from wbs_data import WBS, STALLED, RECOVERY, SPEC, TODAY, DEADLINE

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'onagawa_wbs_R8.xlsx')

NAVY, BLUE, RED, ORANGE, GREEN, YELLOW = '1F3864', '2E75B6', 'C00000', 'C55A11', '375623', 'FFF2CC'
GREY, LIGHT = '595959', 'DDEBF7'
F = 'Arial'

thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

STATE_FILL = {
    '完了':      PatternFill('solid', fgColor='E2EFDA'),
    '進行中':    PatternFill('solid', fgColor='FFF2CC'),
    '未着手':    PatternFill('solid', fgColor='F2F2F2'),
    '町回答待ち': PatternFill('solid', fgColor='FCE4E4'),
}
STATE_FONT = {
    '完了':      Font(name=F, size=9, color=GREEN),
    '進行中':    Font(name=F, size=9, color=ORANGE),
    '未着手':    Font(name=F, size=9, color=GREY),
    '町回答待ち': Font(name=F, size=9, bold=True, color=RED),
}


def title_row(ws, row, text, span, size=13):
    ws.cell(row=row, column=1, value=text)
    c = ws.cell(row=row, column=1)
    c.font = Font(name=F, size=size, bold=True, color='FFFFFF')
    c.alignment = Alignment(vertical='center')
    for i in range(1, span + 1):
        ws.cell(row=row, column=i).fill = PatternFill('solid', fgColor=NAVY)
    ws.row_dimensions[row].height = 24


def header_row(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=F, size=9, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=BLUE)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ============================== 01_WBS ==============================
wb = Workbook()
ws = wb.active
ws.title = '01_WBS'
ws.sheet_view.showGridLines = False

title_row(ws, 1, '令和8年度 上下水道経営指標評価書作成等業務委託（女川町）　WBS', 10)
ws['A2'] = f'作成日：{TODAY}　／　履行期限：{DEADLINE}　／　中間報告：令和8年9月末（提出済）'
ws['A2'].font = Font(name=F, size=9, color=GREY)
ws['A3'] = ('凡例：状態＝完了／進行中／未着手／町回答待ち　　'
            '★印の備考は工程上の要注意事項　　進捗率は当方の作業量に対する割合')
ws['A3'].font = Font(name=F, size=9, color=GREY)

HDR = ['WBS', '作業項目', '成果物', '担当', '開始', '終了', '状態', '進捗率', '備考・前提条件']
header_row(ws, 5, HDR)
widths(ws, [7, 34, 18, 7, 7, 7, 10, 8, 62])

r = 6
for num, lvl, name, deliv, owner, start, end, state, pct, note in WBS:
    if lvl == 0:
        ws.cell(row=r, column=1, value=num).font = Font(name=F, size=10, bold=True, color='FFFFFF')
        ws.cell(row=r, column=2, value=name).font = Font(name=F, size=10, bold=True, color='FFFFFF')
        for i in range(1, len(HDR) + 1):
            ws.cell(row=r, column=i).fill = PatternFill('solid', fgColor=BLUE)
            ws.cell(row=r, column=i).border = BOX
        ws.row_dimensions[r].height = 19
        r += 1
        continue
    vals = [num, '　' + name, deliv, owner, start, end, state, None, note]
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(name=F, size=9)
        c.border = BOX
        c.alignment = Alignment(vertical='top', wrap_text=(i in (2, 3, 9)),
                                horizontal='center' if i in (1, 4, 5, 6, 7, 8) else 'left')
    c = ws.cell(row=r, column=7)
    c.fill = STATE_FILL.get(state, PatternFill())
    c.font = STATE_FONT.get(state, Font(name=F, size=9))
    p = ws.cell(row=r, column=8, value=(pct or 0) / 100)
    p.number_format = '0%'
    p.font = Font(name=F, size=9, bold=(pct == 100))
    if note.startswith('★'):
        ws.cell(row=r, column=9).font = Font(name=F, size=9, bold=True, color=RED)
    r += 1

last = r - 1
ws.freeze_panes = 'A6'
ws.auto_filter.ref = f'A5:I{last}'

# 集計
r += 1
ws.cell(row=r, column=2, value='状態別の件数').font = Font(name=F, size=10, bold=True, color=NAVY)
r += 1
for label in ['完了', '進行中', '未着手', '町回答待ち']:
    n = sum(1 for x in WBS if x[1] == 1 and x[7] == label)
    ws.cell(row=r, column=2, value=label).font = STATE_FONT[label]
    ws.cell(row=r, column=2).fill = STATE_FILL[label]
    ws.cell(row=r, column=2).border = BOX
    c = ws.cell(row=r, column=3, value=n)
    c.font = Font(name=F, size=9)
    c.alignment = Alignment(horizontal='center')
    c.border = BOX
    r += 1
total = sum(1 for x in WBS if x[1] == 1)
done = sum(x[8] for x in WBS if x[1] == 1) / total / 100
ws.cell(row=r, column=2, value='合計').font = Font(name=F, size=9, bold=True)
ws.cell(row=r, column=2).border = BOX
c = ws.cell(row=r, column=3, value=total)
c.font = Font(name=F, size=9, bold=True)
c.alignment = Alignment(horizontal='center')
c.border = BOX
ws.cell(row=r, column=4, value='全体進捗').font = Font(name=F, size=9, bold=True)
c = ws.cell(row=r, column=5, value=done)
c.number_format = '0%'
c.font = Font(name=F, size=9, bold=True, color=NAVY)
c.alignment = Alignment(horizontal='center')

# ============================== 02_停滞事項 ==============================
ws2 = wb.create_sheet('02_停滞事項')
ws2.sheet_view.showGridLines = False
title_row(ws2, 1, '進捗が滞っている事項と解消の手立て', 8)
ws2['A2'] = f'作成日：{TODAY}'
ws2['A2'].font = Font(name=F, size=9, color=GREY)
header_row(ws2, 4, ['No', '区分', '滞っている事項', '起算日', '放置期間', '影響', '解消の手立て', '期限'])
widths(ws2, [5, 11, 40, 10, 10, 58, 58, 12])
r = 5
for no, kind, what, since, dur, impact, how, due in STALLED:
    vals = [no, kind, what, since, dur, impact, how, due]
    for i, v in enumerate(vals, start=1):
        c = ws2.cell(row=r, column=i, value=v)
        c.font = Font(name=F, size=9)
        c.border = BOX
        c.alignment = Alignment(vertical='top', wrap_text=(i in (3, 6, 7)),
                                horizontal='center' if i in (1, 2, 4, 5, 8) else 'left')
    ws2.cell(row=r, column=1).font = Font(name=F, size=9, bold=True, color=RED)
    ws2.cell(row=r, column=2).fill = PatternFill('solid', fgColor='FCE4E4')
    ws2.cell(row=r, column=2).font = Font(name=F, size=9, bold=True, color=RED)
    ws2.cell(row=r, column=3).font = Font(name=F, size=9, bold=True)
    ws2.cell(row=r, column=8).font = Font(name=F, size=9, bold=True, color=RED)
    r += 1
ws2.freeze_panes = 'A5'

r += 2
ws2.cell(row=r, column=1, value='リカバリの優先順位').font = Font(name=F, size=12, bold=True, color=NAVY)
r += 1
header_row(ws2, r, ['順', '作業', '所要', '期限', '効果'])
r += 1
for pri, task, cost, due, effect in RECOVERY:
    for i, v in enumerate([pri, task, cost, due, effect], start=1):
        c = ws2.cell(row=r, column=i, value=v)
        c.font = Font(name=F, size=9)
        c.border = BOX
        c.alignment = Alignment(vertical='top', wrap_text=(i in (2, 5)),
                                horizontal='center' if i in (1, 3, 4) else 'left')
    ws2.cell(row=r, column=1).font = Font(name=F, size=9, bold=True, color=NAVY)
    ws2.cell(row=r, column=4).font = Font(name=F, size=9, bold=True, color=RED)
    r += 1

# ============================== 03_仕様書対応 ==============================
ws3 = wb.create_sheet('03_仕様書対応')
ws3.sheet_view.showGridLines = False
title_row(ws3, 1, '業務仕様書の要求事項と成果物の対応', 5)
ws3['A2'] = ('※ 仕様書原本は未入手のため、令和8年6月2日作成の仕様書整合確認の記載に基づく。'
             '原本の受領後に本表を更新する。')
ws3['A2'].font = Font(name=F, size=9, color=RED)
header_row(ws3, 4, ['要求事項', '求められる内容', '現在の成果物', '達成度', '残作業'])
widths(ws3, [18, 42, 32, 9, 52])
r = 5
for req, what, now, rate, rest in SPEC:
    for i, v in enumerate([req, what, now, rate, rest], start=1):
        c = ws3.cell(row=r, column=i, value=v)
        c.font = Font(name=F, size=9)
        c.border = BOX
        c.alignment = Alignment(vertical='top', wrap_text=(i in (2, 3, 5)),
                                horizontal='center' if i == 4 else 'left')
    ws3.cell(row=r, column=1).font = Font(name=F, size=9, bold=True)
    val = int(rate.rstrip('％'))
    ws3.cell(row=r, column=4).font = Font(
        name=F, size=9, bold=True,
        color=(GREEN if val >= 100 else (ORANGE if val >= 60 else RED)))
    if rest.startswith('★'):
        ws3.cell(row=r, column=5).font = Font(name=F, size=9, bold=True, color=RED)
    r += 1
ws3.freeze_panes = 'A5'

wb.properties.title = '女川町 上下水道経営指標評価書作成等業務委託　WBS'
wb.properties.subject = '令和8年度上下水道経営指標評価書作成等業務委託'
wb.properties.creator = '若山諒太'
wb.properties.lastModifiedBy = '若山諒太'
wb.properties.category = 'ビズアップ公共コンサルティング株式会社'

wb.save(OUT)
print('作成:', os.path.abspath(OUT))
