"""障がい計画：令和8年度アンケート調査の図（第2部の各章に置く）。

原典：23_障がいアンケート分析/原本_受領_20261006/障がい福祉に関するアンケート調査_集計表_20261006.xlsx
  「集計_障がい者」「集計_障がい児」シートの件数と回答者数（n）。割合は 件数÷n（無回答を含む）で、
  分析報告書と同じ基数。図の割合が報告書の本文と一致することを ono_shogai_chosa.SHA で確かめる（zu_check）。
体裁は ono_figs の規則（配色・書体・値ラベル・dpi200）による。

前回計画（令和6年3月）は第2部の各章の「現状と課題」に調査のグラフを置いていた。今期も同じ置き方にする。
"""

import pathlib
import re

import openpyxl

import ono_figs as FG

plt = FG.plt
ROOT = pathlib.Path(__file__).parent / "小野町_引継ぎ_整理済"
XLSX = (ROOT / "23_障がいアンケート分析" / "原本_受領_20261006"
        / "障がい福祉に関するアンケート調査_集計表_20261006.xlsx")
SHIRYO = "資料：令和8年度 障がい福祉に関するアンケート調査"

# 選択肢の表記を図の中で短くする（意味を変えない範囲）
TANSHUKU = {
    "公共交通機関が少ない（公共交通機関が通っていない）": "公共交通機関が少ない",
    "外出先の建物の設備が不便（通路、トイレ、エレベーターなど）": "外出先の建物の設備が不便",
    "切符の買い方や乗り換えの方法がわかりにくい": "切符の買い方や乗り換えがわかりにくい",
    "本や新聞、雑誌の記事、テレビやラジオのニュース": "本・新聞・雑誌、テレビ・ラジオ",
    "家族や親せき、友人・知人": "家族や親せき、友人・知人",
    "障がい者団体や家族会（団体の機関誌など）": "障がい者団体や家族会",
    "病院のケースワーカーや介護保険のケアマネージャー": "病院のケースワーカー、ケアマネジャー",
    "相談支援事業所などの民間の相談窓口": "相談支援事業所などの相談窓口",
    "どこで情報を入手したらよいかわからない": "どこで入手したらよいかわからない",
    "安全なところまで、迅速に避難することができない": "安全なところまで迅速に避難できない",
    "被害状況、避難場所などの情報が入手できない": "被害状況・避難場所の情報が入手できない",
    "避難場所の設備（トイレ等）や生活環境が不安": "避難場所の設備や生活環境が不安",
    "補装具や日常生活用具の入手ができなくなる": "補装具・日常生活用具が入手できない",
    "ストマ（人工肛門・人工膀胱）": "ストマ",
    "人工呼吸器（レスピレーター）": "人工呼吸器",
    "中心静脈栄養（ＩＶＨ）": "中心静脈栄養",
    "ホームヘルパーなどサービス事業所の人": "ホームヘルパーなど事業所の人",
    "必要な在宅サービスが適切に利用できること": "必要な在宅サービスが適切に利用できる",
    "在宅で医療ケアなどが適切に得られること": "在宅で医療ケアなどが適切に得られる",
    "地域移行に向けた具体的な相談や体験利用ができること": "地域移行に向けた相談や体験利用",
    "会社勤めや、自営業、家業などで収入を得て仕事している": "収入を得る仕事をしている",
    "ボランティアなど、収入を得ない仕事をしている": "収入を得ない仕事をしている",
    "福祉施設、作業所等に通っている（就労継続支援Ａ型も含む）": "福祉施設・作業所等に通っている",
    "入所している施設や病院等で過ごしている": "施設や病院等で過ごしている",
    "大学、専門学校、職業訓練校などに通っている": "大学・専門学校・職業訓練校",
    "職場の上司や同僚に障がいの理解があること": "上司や同僚に障がいの理解がある",
    "勤務場所におけるバリアフリー等の配慮": "勤務場所のバリアフリー等の配慮",
    "就労後のフォローなど職場と支援機関の連携": "就労後のフォロー（職場と支援機関の連携）",
    "仕事についての職場外での相談対応、支援": "職場外での相談対応・支援",
    "幼稚園、保育所、障がい児通園施設などに通っている": "幼稚園・保育所・通園施設",
    "一般の小中学校、高校に通っている": "一般の小中学校・高校",
    "一般の小中学校の支援学級に通っている": "小中学校の支援学級",
    "特別支援学校（小中高等部）に通っている": "特別支援学校",
    "通院施設や保育所、幼稚園、学校の先生": "通院施設・保育所・学校の先生",
}
# 横棒の下に回す（灰色で示す）選択肢
SUE = ("その他", "特にない", "無回答", "情報は必要ない")

_CACHE = {}


def load(sheet):
    """{設問: {"n": 回答者数, "rows": [(項目, 選択肢, 件数)]}}"""
    if sheet in _CACHE:
        return _CACHE[sheet]
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    out = {}
    for row in wb[sheet].iter_rows(min_row=5, values_only=True):
        if not row[0] or row[3] is None:
            continue
        q = str(row[0])
        d = out.setdefault(q, {"n": int(row[5]), "rows": []})
        d["rows"].append((str(row[1]), str(row[2]), int(row[3])))
    _CACHE[sheet] = out
    return out


BETSU = ("身体障がい", "知的障がい", "精神障がい")


def load_betsu():
    """「集計_障がい別」：{設問: {"n": {障がい: n}, "rows": [(項目, 選択肢, {障がい: 件数})]}}。
    障がい別は所持している手帳による（重複所持は各区分に数える）。"""
    if "betsu" in _CACHE:
        return _CACHE["betsu"]
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    out, n, kou = {}, None, None
    for row in wb["集計_障がい別"].iter_rows(min_row=5, values_only=True):
        if row[2] == "回答者数（n）":
            n = dict(zip(BETSU, (int(row[3]), int(row[4]), int(row[5]))))
            continue
        if not row[0] or row[3] is None:
            continue
        if row[1]:
            kou = str(row[1])
        d = out.setdefault(str(row[0]), {"n": n, "rows": []})
        d["rows"].append((kou, str(row[2]), dict(zip(BETSU, (int(row[3]), int(row[4]), int(row[5]))))))
    _CACHE["betsu"] = out
    return out


def pct(sheet, q, sentaku, koumoku=None):
    d = load(sheet)[q]
    for k, s, c in d["rows"]:
        if s == sentaku and (koumoku is None or k == koumoku):
            return c / d["n"] * 100
    raise KeyError((sheet, q, sentaku, koumoku))


def _t(s):
    return TANSHUKU.get(s, s)


def yokobou(sheet, q, name, xlab="回答の割合（％）", nozoku=(), kensu=False, kyocho=1, kiso=None):
    """複数回答・単一回答を横棒で示す。多い順。上位 kyocho 件を赤で強調する。
    kiso＝(基数, 説明)。条件付き設問は、該当する方を基数にする（指定がなければ回答者全体）。"""
    d = load(sheet)[q]
    n = kiso[0] if kiso else d["n"]
    rows = [(s, c) for _k, s, c in d["rows"] if s not in nozoku and s != "無回答"]
    main = sorted([r for r in rows if r[0] not in SUE], key=lambda r: -r[1])
    sue = [r for r in rows if r[0] in SUE]
    rows = main + sue
    vals = [c / n * 100 for _s, c in rows]
    cols = [FG.RED if i < kyocho else FG.BLUE for i in range(len(main))] + [FG.GREY] * len(sue)
    fig, ax = plt.subplots(figsize=(7.4, 0.36 * len(rows) + 0.9))
    y = list(range(len(rows)))
    ax.barh(y, vals, color=cols, height=0.62)
    top = max(vals) * 1.22 if vals else 1
    for i, ((_s, c), v) in enumerate(zip(rows, vals)):
        lab = "%.1f%%" % v + ("（%d件）" % c if kensu else "")
        FG._lab(ax, v + top * 0.01, i, lab, FG.NAVY, sz=8)
    ax.set_yticks(y)
    ax.set_yticklabels([_t(s) for s, _c in rows], fontsize=8.5)
    ax.invert_yaxis()
    ax.set_xlim(0, top)
    ax.set_xlabel(xlab + "　n=%d" % n + ("（%s）" % kiso[1] if kiso else ""), fontsize=8.5)
    FG._xgrid(ax)
    fig.tight_layout()
    return FG._save(fig, name)


def gaito(sheet, q, sentaku):
    """条件付き設問の基数：前の設問で sentaku を選んだ方の数。"""
    return sum(c for _k, s, c in load(sheet)[q]["rows"] if s in sentaku)


BAND = [FG.NAVY, FG.BLUE, FG.PALE, "#E6E6E6"]


def obi(items, name, labels=None):
    """帯グラフ。items [(行の名前, sheet, 設問, 項目 or None)]。選択肢の順は最初の行に合わせる。"""
    rows, cats = [], None
    for nm, sheet, q, k in items:
        if sheet == "集計_障がい別":   # k は障がいの区分
            db = load_betsu()[q]
            d = {"n": db["n"][k]}
            rr = [(s, c[k]) for _kk, s, c in db["rows"]]
        else:
            d = load(sheet)[q]
            rr = [(s, c) for kk, s, c in d["rows"] if k is None or kk == k]
        if cats is None:
            cats = [s for s, _c in rr]
        if labels:   # 選択肢の語が行ごとに違う（できる／いる など）ときは、順番で対応させる
            if len(rr) != len(cats):
                raise ValueError("%s %s の選択肢の数が合わない" % (sheet, q))
            vs = [c / d["n"] * 100 for _s, c in rr]
        else:
            m = dict(rr)
            if set(m) != set(cats):
                raise ValueError("%s %s の選択肢が1行目と違う：%s" % (sheet, q, sorted(set(m) ^ set(cats))))
            vs = [m[s] / d["n"] * 100 for s in cats]
        rows.append((nm + "（n=%d）" % d["n"], vs))
    cols = (BAND[:len(cats) - 1] if len(cats) <= 4 else
            [FG.NAVY, FG.BLUE, "#7FAEDB", FG.PALE][:len(cats) - 1]) + ["#E6E6E6"]
    if cats[-1] != "無回答":
        cols[-1] = FG.PALE
    fig, ax = plt.subplots(figsize=(7.4, 0.62 * len(rows) + 1.3))
    for i, (_nm, vs) in enumerate(rows):
        left = 0
        for j, v in enumerate(vs):
            ax.barh(i, v, left=left, color=cols[j], height=0.58, edgecolor="white",
                    label=(labels or cats)[j] if i == 0 else None)
            if v >= 6:
                ax.text(left + v / 2, i, "%.1f" % v, ha="center", va="center", fontsize=8,
                        fontweight="bold", color="white" if j < 2 else FG.NAVY)
            left += v
    ax.set_yticks(range(len(rows)))
    ax.set_yticklabels([nm for nm, _v in rows], fontsize=8.5)
    ax.invert_yaxis()
    ax.set_xlim(0, 100)
    ax.set_xlabel("回答の割合（％）", fontsize=8.5)
    labs = labels or cats
    ncol = 2 if max(len(x) for x in labs) > 12 else min(len(cats), 5)
    ax.legend(fontsize=8, loc="lower center", ncol=ncol, bbox_to_anchor=(0.5, 1.0), frameon=False)
    fig.tight_layout()
    return FG._save(fig, name)


def riyou_iko(name="f80_shogai_riyou_iko"):
    """障がい福祉サービスの現在の利用と今後の利用意向（問31）。今後の多い順。"""
    d = load("集計_障がい者")
    n = d["問31-1"]["n"]
    ima = {k: c for k, s, c in d["問31-1"]["rows"] if s == "利用している"}
    kongo = {k: c for k, s, c in d["問31-2"]["rows"] if s == "利用したい"}
    ks = sorted(kongo, key=lambda k: -kongo[k])
    y = list(range(len(ks)))
    h = 0.36
    fig, ax = plt.subplots(figsize=(7.4, 0.36 * len(ks) + 1.2))
    ax.barh([i - h / 2 for i in y], [ima[k] / n * 100 for k in ks], height=h, color=FG.PALE, label="現在利用している")
    ax.barh([i + h / 2 for i in y], [kongo[k] / n * 100 for k in ks], height=h, color=FG.BLUE, label="今後利用したい")
    top = max(kongo.values()) / n * 100 * 1.2
    for i, k in enumerate(ks):
        FG._lab(ax, ima[k] / n * 100 + top * 0.01, i - h / 2, "%.1f" % (ima[k] / n * 100), FG.NAVY, sz=7)
        FG._lab(ax, kongo[k] / n * 100 + top * 0.01, i + h / 2, "%.1f" % (kongo[k] / n * 100), FG.BLUE, sz=7)
    ax.set_yticks(y)
    ax.set_yticklabels([SVC.get(k, k) for k in ks], fontsize=8.5)
    ax.invert_yaxis()
    ax.set_xlim(0, top)
    ax.set_xlabel("回答の割合（％）　n=%d" % n, fontsize=8.5)
    ax.legend(fontsize=8, loc="lower right", frameon=False)
    FG._xgrid(ax)
    fig.tight_layout()
    return FG._save(fig, name)


SVC = {"重度包括": "重度障害者等包括支援", "就労移行": "就労移行支援", "就労継続": "就労継続支援",
       "共同生活援助": "共同生活援助（グループホーム）", "施設入所": "施設入所支援", "相談支援": "相談支援",
       "就労選択": "就労選択支援", "宿泊型自立訓練": "宿泊型自立訓練"}

# 図の一覧：(番号の仮名, 置く章, 題名, 作る関数)
def _gobetsu(q_sha, q_ji):
    """障がい者全体・障がい別（身体・知的・精神）・障がい児の5行。"""
    return ([("障がいのある方 全体", "集計_障がい者", q_sha, None)]
            + [("　" + b, "集計_障がい別", q_sha, b) for b in BETSU]
            + [("障がいのある子ども", "集計_障がい児", q_ji, None)])


# 図の一覧：(仮名, 置く章, 題名, 作る関数)
ZU = [
    ("差別", "第2章", "障がいがあることで差別や嫌な思いをした経験",
     lambda: obi(_gobetsu("問35", "問32"), "f70_shogai_sabetsu", labels=["ある", "少しある", "ない", "無回答"])),
    ("差別場所", "第2章", "差別や嫌な思いをした場所（障がいのある方・複数回答）",
     lambda: yokobou("集計_障がい者", "問36-1", "f87_shogai_sabetsu_basho",
                     kiso=(gaito("集計_障がい者", "問35", ("ある", "少しある")), "「ある」「少しある」と答えた方"))),
    ("後見", "第2章", "成年後見制度の認知",
     lambda: obi([("障がいのある方", "集計_障がい者", "問37", None),
                  ("障がいのある子ども", "集計_障がい児", "問34", None)], "f71_shogai_kouken",
                 labels=["名前も内容も知っている", "名前は聞いたことがあるが内容は知らない", "名前も内容も知らない", "無回答"])),
    ("法認知", "第2章", "障害者虐待防止法・障害者差別解消法の認知（障がいのある方）",
     lambda: obi([("障害者虐待防止法", "集計_障がい者", "問39", None),
                  ("障害者差別解消法", "集計_障がい者", "問40", None)], "f72_shogai_hou")),
    ("外出頻度", "第3章", "1週間の外出の頻度",
     lambda: obi([("障がいのある方", "集計_障がい者", "問19", None),
                  ("障がいのある子ども", "集計_障がい児", "問25", None)], "f73_shogai_gaishutsu",
                 labels=["毎日外出する", "1週間に数回外出する", "めったに外出しない", "まったく外出しない", "無回答"])),
    ("外出困難", "第3章", "外出するときに困ること（障がいのある方・複数回答）",
     lambda: yokobou("集計_障がい者", "問22", "f74_shogai_komaru", kyocho=2)),
    ("情報", "第4章", "障がいや福祉サービスの情報の入手先（障がいのある方・複数回答）",
     lambda: yokobou("集計_障がい者", "問34", "f75_shogai_joho")),
    ("避難", "第5章", "火事や地震等の災害時に一人で避難できるか",
     lambda: obi(_gobetsu("問41", "問38"), "f76_shogai_hinan", labels=["できる", "できない", "わからない", "無回答"])),
    ("近所", "第5章", "家族が不在の場合や一人暮らしの場合に、近所に助けてくれる人がいるか",
     lambda: obi(_gobetsu("問42", "問39"), "f88_shogai_kinjo", labels=["いる", "いない", "わからない", "無回答"])),
    ("災害", "第5章", "災害時に困ること（障がいのある方・複数回答）",
     lambda: yokobou("集計_障がい者", "問43", "f77_shogai_saigai", kyocho=2)),
    ("医療", "第6章", "現在受けている医療ケア（障がいのある方・複数回答）",
     lambda: yokobou("集計_障がい者", "問16", "f78_shogai_iryo", nozoku=("中心静脈栄養（ＩＶＨ）",))),
    ("相談", "第7章", "悩みや困ったことの相談相手（障がいのある方・複数回答）",
     lambda: yokobou("集計_障がい者", "問33", "f79_shogai_soudan")),
    ("利用意向", "第7章", "障がい福祉サービスの現在の利用と今後の利用意向（障がいのある方）", riyou_iko),
    ("地域生活", "第7章", "地域で生活するために必要な支援（障がいのある方・複数回答）",
     lambda: yokobou("集計_障がい者", "問18", "f81_shogai_chiiki", kyocho=3)),
    ("児日中", "第8章", "障がいのある子どもの平日の日中の過ごし方",
     lambda: yokobou("集計_障がい児", "問17", "f82_shogai_ji_nicchu", kensu=True,
                     nozoku=("病院などのデイケアに通っている", "リハビリテーションを受けている"))),
    ("児支援", "第8章", "障がいのある子どもが地域で生活するために必要な支援（複数回答）",
     lambda: yokobou("集計_障がい児", "問24", "f83_shogai_ji_shien", kensu=True, kyocho=3)),
    ("日中", "第9章", "平日の日中の過ごし方（障がいのある方）",
     lambda: yokobou("集計_障がい者", "問23", "f84_shogai_nicchu")),
    ("就労", "第9章", "障がい者の就労支援として必要なこと（障がいのある方・複数回答）",
     lambda: yokobou("集計_障がい者", "問27", "f85_shogai_shuro", kyocho=3)),
    ("外出目的", "第10章", "外出の目的（障がいのある方・複数回答）",
     lambda: yokobou("集計_障がい者", "問21", "f86_shogai_mokuteki")),
]


def build_all():
    """{仮名: (章, 題名, png)}"""
    return {k: (sho, dai, f()) for k, sho, dai, f in ZU}


def zu_check():
    """図の割合が、分析報告書の本文の割合（ono_shogai_chosa.SHA）と一致するかを確かめる。"""
    import ono_shogai_chosa as C
    cases = [("集計_障がい者", "問35", ["ある", "少しある"], "25.9％"),
             ("集計_障がい者", "問37", ["名前も内容も知らない"], "36.7％"),
             ("集計_障がい者", "問40", ["知らない"], "78.8％"),
             ("集計_障がい者", "問22", ["外出にお金がかかる"], "25.1％"),
             ("集計_障がい者", "問43", ["投薬や治療が受けられない"], "51.7％"),
             ("集計_障がい者", "問41", ["できない", "わからない"], "55.6％"),
             ("集計_障がい者", "問33", ["家族や親せき"], "73.7％"),
             ("集計_障がい者", "問23", ["自宅で過ごしている"], "44.0％"),
             ("集計_障がい者", "問27", ["職場の上司や同僚に障がいの理解があること"], "35.9％"),
             ("集計_障がい者", "問18", ["経済的な負担の軽減"], "48.3％"),
             ("集計_障がい児", "問17", ["一般の小中学校、高校に通っている"], "40.0％")]
    sha = " ".join(t for _b, t, _c in C.SHA + C.JI)
    ng = []
    for sheet, q, ss, want in cases:
        d = load(sheet)[q]
        got = sum(c for _k, s, c in d["rows"] if s in ss) / d["n"] * 100
        if "%.1f％" % got != want:
            ng.append("%s %s %s：集計表 %.1f％／報告書 %s" % (sheet, q, ss, got, want))
        if want not in sha:
            ng.append("%s が報告書の主な結果（SHA・JI）にない" % want)
    return ng


if __name__ == "__main__":
    for k, (sho, dai, p) in build_all().items():
        print(sho, dai, p.name)
    print(zu_check() or "照合 すべて一致")
