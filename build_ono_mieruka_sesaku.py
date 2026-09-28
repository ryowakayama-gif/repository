"""小野町 見える化システム 施策反映に0が入っている件と、在宅サービスの利用回（日）数。

令和8年9月28日のご指示
  「添付ファイル確認下さい。在宅サービスの利用日数についても確認願います。」

**画面キャプチャ（推計名「将来推計0925-2」）で分かったこと**
  1. **自然体推計は6,511円で出ています。**
     画面上部の「第10期 4,829円（-1,682円）」の括弧は施策反映による増減です。
     6,511円 － 1,682円 ＝ 4,829円。
     根拠：将来推計1では「6,466円（0円）」で、
     「自然体推計に全て戻す」がグレー（＝施策反映なし）でした。
  2. **施策反映の在宅サービス利用回（日）数に0.0が入っています。**
     令和9年度以降がすべて0.0で、ピンク（＝施策反映値）で表示されています。
     **「自然体推計に全て戻す」を押してください。**
  3. **認定率の伸びの選択肢に「０（ゼロ）」はありません。**（画面で確定）
     ①当該市町村 令和3→令和8／②同 令和7→令和8／③同 令和6→令和8／
     ④同 令和6→令和7／⑤全国 令和6→令和7 の5つ。町は⑤を選択済み。

**在宅サービスの1人1月あたり利用回（日）数**
  令和6・7年度は当方の算定とほぼ一致する。令和8年度だけが跳ねている。
    訪問介護       R6 16.19／R7 15.64／**R8 19.34**回　当方15.5回
    短期入所生活介護 R6 9.28／R7 9.94／**R8 14.89**日　当方9.25日
    通所介護       R6 9.15／R7 9.26／R8 8.73回　当方9.14回

出力
  11_見える化出力依頼/小野町_見える化_施策反映に0が入っている件_YYYYMMDD.docx
  11_見える化出力依頼/小野町_見える化_在宅サービスの利用回数_YYYYMMDD.xlsx
"""

import pathlib
import warnings

warnings.filterwarnings("ignore")

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import build_ono_tanka as T

ROOT = pathlib.Path(__file__).parent
OUT = ROOT / "小野町_引継ぎ_整理済" / "11_見える化出力依頼"
ASOF = "20260928"
ASOF_JP = "令和8年9月28日"
JP_MIN = "游明朝"
JP_GO = "游ゴシック"

HEAD = PatternFill("solid", fgColor="1F3864")
NG = PatternFill("solid", fgColor="FCE4E4")
WARN = PatternFill("solid", fgColor="FFE0B2")
OK = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SHIZEN = 6511      # 自然体推計（4,829＋1,682）
GENZAI = 4829      # 画面の表示値（施策反映後）
SAISA = -1682      # 施策反映による増減
OUR_GETSU = 6048
          # 令和8年9月28日訂正。調整交付金見込額の基数に総合事業費を
          # 含めていなかった（6,095円→6,048円）。切上げ6,100円は不変。

# 認定率の伸びの選択肢（画面で確認）
SENTAKUSHI = [
    ("①", "当該市町村における令和3年度→令和8年度の伸び", "5年",
     "**5年の平均。令和7年8〜9月の断層も令和8年度の1か月分も5分の1に薄まる。**"
     "画面に出ている値は男の第1号被保険者で合計▲0.2％と小さい。**次善**"),
    ("②", "当該市町村における令和7年度→令和8年度の伸び", "1年",
     "**避ける。**令和8年度は国保連月報の1か月分である"),
    ("③", "当該市町村における令和6年度→令和8年度の伸び", "2年",
     "**避ける。**令和8年度の1か月分を2年の窓で拾う"),
    ("④", "当該市町村における令和6年度→令和7年度の伸び", "1年",
     "**避ける。**令和7年8〜9月の認定者数の断層（944人→792人）が入る。"
     "将来推計0925で短期入所生活介護が3.29倍に膨らんだ"),
    ("⑤", "全国における令和6年度→令和7年度の伸び", "1年",
     "**第一候補。**小野町の断層を含まない。**町はこれを選択済み**"),
]

# 在宅サービスの1人1月あたり利用回（日）数
RIYO = [
    ("訪問介護", "回", 993.8, 61.4, 800.8, 51.2, 1295.6, 67.0, 743, 48.0),
    ("訪問看護", "回", 265.6, 36.7, 281.1, 39.6, 286.3, 46.0, None, None),
    ("通所介護", "回", 1623.8, 177.5, 1509.1, 162.9, 1475.3, 169.0, 1436, 157.1),
    ("短期入所生活介護", "日", 248.8, 26.8, 270.3, 27.2, 893.2, 60.0, 222, 24.0),
]

# 訪問介護の要介護度別（画面から読み取り）
HOMON = [
    ("令和6年度", 8.2, 15.0, 19.7, 34.4, 22.1, "実績"),
    ("令和7年度", 9.2, 12.1, 19.4, 23.7, 39.1, "実績"),
    ("令和8年度", 9.0, 7.0, 12.4, 18.7, 61.3, "**実績（1か月分）**"),
    ("令和9〜11年度", 0.0, 0.0, 0.0, 0.0, 0.0, "**施策反映値（ピンク）**"),
]

TEJUN = [
    ("1", "**「施策反映　在宅サービス利用者数等」の画面で"
          "「自然体推計に全て戻す」を押す**",
     "**これが最優先です。**令和9年度以降の0.0が消え、"
     "自然体推計の値に戻ります"),
    ("2", "「利用回（日）数の施策反映へ」でも同じく"
          "「自然体推計に全て戻す」を押す",
     "利用率と利用回（日）数は別の画面です。**両方で押してください**"),
    ("3", "「施策反映　施設・居住系サービス利用者数」でも押す",
     "ボタンがグレーなら、その画面は編集されていません"),
    ("4", "「施策反映　認定者数」でも押す", "同上"),
    ("5", "**「保険料額の更新」を押して再計算する**",
     "**画面上部の括弧が「(0円)」に戻れば成功です**"),
    ("6", "第10期の保険料額を確認する",
     "**6,500円前後になるはずです**（自然体推計6,511円）"),
    ("7", "ワーニングチェックと総括表を出力する", "当方で突き合わせます"),
]

KAKUNIN = [
    ("1", "**「(-1,682円)」の意味をご確認ください。**"
          "当方は「施策反映による増減」と読んでいます。"
          "将来推計1では「(0円)」で、"
          "「自然体推計に全て戻す」がグレーでした"),
    ("2", "**利用率の伸びの選択画面のキャプチャをお送りください。**"
          "認定率の画面（5つの選択肢）は拝見しました。"
          "利用率の画面も同じ5つなら、⑤全国を選んでください"),
    ("3", "**令和8年度のチェックが本当に外れているかをご確認ください。**"
          "訪問介護の令和8年度が要介護5で61.3回と、"
          "令和7年度の39.1回から跳ねています。"
          "**1か月分のデータの特徴です**"),
    ("4", "**「自然体推計に全て戻す」を押したあと、"
          "令和9年度の値が令和7年度と令和8年度のどちらに近いかを"
          "お知らせください。**"
          "令和7年度（39.1回）に近ければ正しく、"
          "令和8年度（61.3回）に近ければ令和8年度が基準になっています"),
]


# ---------------------------------------------------------------- 計算

def riyo_rows():
    out = []
    for nm, tani, k6, n6, k7, n7, k8, n8, ok, on in RIYO:
        out.append((nm, tani, k6 / n6, k7 / n7, k8 / n8,
                    ok / on if ok else None, k8 / n8 / (k7 / n7)))
    return out


def selfcheck(R):
    bad = []
    if GENZAI - SAISA != SHIZEN:
        bad.append(f"自然体推計 {GENZAI - SAISA}≠{SHIZEN}")
    for nm, _t, r6, r7, r8, our, hi in R:
        if our and abs(our - r7) / r7 > 0.12:
            bad.append(f"{nm} 当方{our:.2f}と令和7年度{r7:.2f}の差が大きい")
    if len(SENTAKUSHI) != 5:
        bad.append(f"選択肢 {len(SENTAKUSHI)}件")
    return bad


# ---------------------------------------------------------------- 体裁

def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = JP_MIN
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), JP_MIN)
    return doc


def head(doc, text, level):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = JP_GO
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_GO)


def _emph(p, text):
    for i, part in enumerate(str(text).split("**")):
        if not part:
            continue
        r = p.add_run(part)
        r.font.name = JP_MIN
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_MIN)
        r.bold = bool(i % 2)


def body(doc, *paras):
    for t in paras:
        _emph(doc.add_paragraph(), t)


def note(doc, text):
    p = doc.add_paragraph()
    _emph(p, text)
    for r in p.runs:
        r.font.size = Pt(9)


def table(doc, header_row, rows, right_from=99):
    t = doc.add_table(rows=1, cols=len(header_row))
    t.style = "Table Grid"
    for i, h in enumerate(header_row):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = JP_GO
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_GO)
        r.bold = True
        r.font.size = Pt(9)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            _emph(p, v)
            for r in p.runs:
                r.font.size = Pt(9)
            if i >= right_from:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()


def style_head(ws, row=1):
    for c in ws[row]:
        if c.value is None:
            continue
        c.fill = HEAD
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        c.border = BORDER


def body_style(ws, first=2, wrap=()):
    for r in ws.iter_rows(min_row=first):
        for c in r:
            c.font = Font(size=9)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=c.column in wrap)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


# ---------------------------------------------------------------- Word

def build_docx(R):
    doc = new_doc()
    head(doc, "施策反映に0が入っている件と、在宅サービスの利用回（日）数", 0)
    body(doc, f"小野町　第10期介護保険事業計画　／　{ASOF_JP}")

    head(doc, "0　結論", 1)
    table(doc, ["#", "内容"], [
        ("1", f"**自然体推計は{SHIZEN:,}円で出ています。うまくいっています。**"
              f"画面上部の「第10期 {GENZAI:,}円（{SAISA:+,}円）」の括弧は"
              "施策反映による増減です。"
              f"**{SHIZEN:,}円 － {abs(SAISA):,}円 ＝ {GENZAI:,}円**（第1章）"),
        ("2", "**直すのは1か所です。"
              "「施策反映　在宅サービス利用者数等」の"
              "「自然体推計に全て戻す」を押してください。**"
              "令和9年度以降の利用回（日）数がすべて0.0で、"
              "ピンク（＝施策反映値）になっています（第2章）"),
        ("3", "**「０（ゼロ）」は認定率の選択肢にありませんでした。**"
              "5つの選択肢を拝見しました。"
              "**町が選ばれた⑤全国でよいと考えます**（第3章）"),
        ("4", "**在宅サービスの1人1月あたり利用回（日）数は、"
              "令和6・7年度は当方の算定とほぼ一致しています。**"
              "令和8年度だけが跳ねています（第4章）"),
    ])

    # ---- 1 自然体推計
    head(doc, "1　自然体推計は6,511円で出ています", 1)
    body(doc,
         "**画面上部の括弧の中は、施策反映による保険料の増減です。**"
         "将来推計1では「6,466円（0円）」で、"
         "そのとき「自然体推計に全て戻す」はグレーでした。"
         "**施策反映を入れていなければ括弧は（0円）になります。**")
    table(doc, ["区分", "金額", "内容"], [
        ("自然体推計", f"**{SHIZEN:,}円**",
         f"{GENZAI:,}円 ＋ {abs(SAISA):,}円。**伸びの設定は効いています**"),
        ("施策反映による増減", f"**{SAISA:+,}円**",
         "在宅サービスの利用回（日）数に0.0が入っているため"),
        ("画面の表示値", f"{GENZAI:,}円", "施策反映後"),
        ("（参考）当方の算定", f"{OUR_GETSU:,}円",
         "令和8年度を国保連月報の4〜6月提供分で置いたもの"),
    ], right_from=1)
    note(doc, f"※ 令和12年度も同じで、5,560円＋1,906円＝7,466円が"
              "自然体推計です。"
              "**※ この読み方が正しいかを、画面でご確認いただけると確実です。**")

    # ---- 2 施策反映
    head(doc, "2　施策反映に0.0が入っています", 1)
    body(doc,
         "**「施策反映　在宅サービス利用者数等」の利用回（日）数の画面で、"
         "令和9年度以降がすべて0.0になっています。**"
         "ピンクで表示されているのは、"
         "自然体推計値ではなく施策反映値として扱われているためです。")
    table(doc, ["年度", "要介護1", "要介護2", "要介護3", "要介護4", "要介護5",
                "区分"],
          [(y, f"{a:.1f}", f"{b:.1f}", f"{c:.1f}", f"{d:.1f}", f"{e:.1f}", k)
           for y, a, b, c, d, e, k in HOMON], right_from=1)
    note(doc, "※ 訪問介護の1人1月あたり利用回数（単位：回）。"
              "**令和9〜11年度が0.0のままでは、訪問介護の給付費が0になります。**")

    head(doc, "(1)　直し方", 2)
    table(doc, ["順", "すること", "留意点"],
          [(a, b, c) for a, b, c in TEJUN])
    body(doc,
         "**「自然体推計に全て戻す」は、施策反映の画面ごとにあります。**"
         "在宅サービスは「利用者数」と「利用回（日）数」で画面が分かれているため、"
         "両方で押してください。"
         "**ボタンがグレーであれば、その画面は編集されていません。**")

    # ---- 3 選択肢
    head(doc, "3　伸びの選択肢について", 1)
    body(doc,
         "**認定率の画面を拝見しました。「０（ゼロ）」はありません。**"
         "次の5つです。")
    table(doc, ["#", "選択肢", "窓", "評価"],
          [(a, b, c, d) for a, b, c, d in SENTAKUSHI])
    body(doc,
         "**町が選ばれた⑤全国でよいと考えます。**"
         "小野町の令和6年度から令和7年度への動きには"
         "令和7年8〜9月の認定者数の断層が入っていますが、"
         "全国の伸びには入りません。",
         "**①当該市町村における令和3年度→令和8年度の伸びも有力です。**"
         "5年の平均なので、断層も令和8年度の1か月分も5分の1に薄まります。"
         "画面に出ていた値は男の第1号被保険者で合計▲0.2％と小さく、"
         "穏当な水準です。"
         "**⑤で進めて、結果を見てから①と比べるのでも構いません。**")
    note(doc, "※ 利用率の伸びの選択画面は拝見していません。"
              "**同じ5つであれば⑤全国を選んでください。**"
              "違う選択肢が並んでいる場合は、キャプチャをお送りください。")

    # ---- 4 利用回（日）数
    head(doc, "4　在宅サービスの1人1月あたり利用回（日）数", 1)
    body(doc,
         "**令和6・7年度は当方の算定とほぼ一致しています。"
         "令和8年度だけが跳ねています。**")
    table(doc, ["サービス", "単位", "令和6年度", "令和7年度", "令和8年度",
                "当方（令和11年度）", "令和8年度÷令和7年度"],
          [(nm, tani, f"{r6:.2f}", f"{r7:.2f}", f"**{r8:.2f}**",
            f"{our:.2f}" if our else "―",
            f"**{hi:.2f}**" if hi > 1.2 else f"{hi:.2f}")
           for nm, tani, r6, r7, r8, our, hi in R], right_from=2)
    note(doc, "※ ワークシート「2_サービス別給付費」の回（日）数÷人数。"
              "いずれも1月あたり。"
              "**※ 当方の値は令和11年度の見込みです。**")
    body(doc,
         "**短期入所生活介護は令和7年度の9.94日から令和8年度14.89日へ"
         "1.50倍になっています。**"
         "利用者数も27.2人から60.0人へ2.21倍です。"
         "**1か月分のデータによるもので、実態ではありません。**",
         "**訪問介護の要介護5は、令和6年度22.1回・令和7年度39.1回・"
         "令和8年度61.3回と増え続けています。**"
         "一方で要介護2は15.0回・12.1回・7.0回と減っています。"
         "**利用者が数人しかいない区分では、"
         "1人の増減が1人あたりの回数を大きく動かします。**"
         "要介護度別の値だけを見て判断しないほうがよいところです。",
         "**当方の算定は、要介護度別の1人1月あたり回数を使っていません。**"
         "サービスごとの総回数を年報 様式2 から取り、"
         "施設・居住系・在宅の区分別の係数で延ばしています。"
         "**そのため、この種の振れの影響を受けません。**"
         "結果として、令和6・7年度の実績とほぼ同じ水準になっています。")

    # ---- 5 お願い
    head(doc, "5　お願いする事項", 1)
    table(doc, ["#", "内容"], [(a, b) for a, b in KAKUNIN])

    p = OUT / f"小野町_見える化_施策反映に0が入っている件_{ASOF}.docx"
    doc.save(p)
    return p


# ---------------------------------------------------------------- Excel

def build_xlsx(R):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("00_要点")
    ws.append(["#", "内容", "判定"])
    style_head(ws)
    for i, (a, ok) in enumerate([
        (f"自然体推計は{SHIZEN:,}円で出ている（{GENZAI:,}円＋{abs(SAISA):,}円）。"
         "伸びの設定は効いている", True),
        ("施策反映の在宅サービス利用回（日）数に0.0が入っている（ピンク表示）",
         False),
        ("「自然体推計に全て戻す」を押す。在宅は利用者数と利用回数の両画面で",
         False),
        ("「０（ゼロ）」は認定率の選択肢にない（5つの選択肢を画面で確認）", None),
        ("町が選んだ⑤全国でよい。①令和3→令和8（5年窓）も有力", True),
        ("在宅サービスの1人1月あたり回（日）数は令和6・7年度が当方とほぼ一致",
         True),
        ("令和8年度だけが跳ねている（短期入所生活介護9.94→14.89日）", False),
    ], start=1):
        ws.append([i, a, "○" if ok else ("×" if ok is False else "―")])
        ws.cell(ws.max_row, 3).fill = \
            OK if ok else (NG if ok is False else WARN)
    body_style(ws, wrap=(2,))
    widths(ws, [4, 76, 6])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("01_利用回数の対比")
    ws.append(["サービス", "単位", "令和6年度", "令和7年度", "令和8年度",
               "当方（令和11年度）", "R8÷R7"])
    style_head(ws)
    for nm, tani, r6, r7, r8, our, hi in R:
        ws.append([nm, tani, r6, r7, r8, our, hi])
        ws.cell(ws.max_row, 5).fill = NG if hi > 1.2 else WARN
        if our:
            ws.cell(ws.max_row, 6).fill = OK
    for r in range(2, ws.max_row + 1):
        for c in (3, 4, 5, 6, 7):
            ws.cell(r, c).number_format = "0.00"
    body_style(ws)
    widths(ws, [22, 6, 14, 14, 14, 18, 10])
    ws.freeze_panes = "A2"
    ws.append([])
    ws.append(["※ ワークシート「2_サービス別給付費」の回（日）数÷人数。"
               "いずれも1月あたり。"])
    ws.append(["※ 当方の値は令和11年度の見込み（build_ono_kaisu）。"])
    ws.append(["※ 令和8年度は国保連月報が1か月分のため跳ねている。"])

    ws = wb.create_sheet("02_訪問介護の要介護度別")
    ws.append(["年度", "要介護1", "要介護2", "要介護3", "要介護4", "要介護5",
               "区分"])
    style_head(ws)
    for y, a, b, c, d, e, k in HOMON:
        ws.append([y, a, b, c, d, e, k.replace("**", "")])
        if "施策反映" in k:
            for col in range(2, 7):
                ws.cell(ws.max_row, col).fill = NG
    for r in range(2, ws.max_row + 1):
        for c in range(2, 7):
            ws.cell(r, c).number_format = "0.0"
    body_style(ws)
    widths(ws, [16, 10, 10, 10, 10, 10, 28])
    ws.freeze_panes = "A2"
    ws.append([])
    ws.append(["※ 訪問介護の1人1月あたり利用回数（回）。画面から読み取り。"])
    ws.append(["※ 令和9〜11年度が0.0のままでは訪問介護の給付費が0になる。"])

    ws = wb.create_sheet("03_伸びの選択肢")
    ws.append(["#", "選択肢", "窓", "評価"])
    style_head(ws)
    for a, b, c, d in SENTAKUSHI:
        ws.append([a, b, c, d.replace("**", "")])
        ws.cell(ws.max_row, 2).fill = \
            OK if a == "⑤" else (WARN if a == "①" else NG)
    body_style(ws, wrap=(2, 4))
    widths(ws, [4, 42, 6, 80])
    ws.freeze_panes = "A2"
    ws.append([])
    ws.append(["※ 認定率の変化の画面（S211900）で確認。"
               "「０（ゼロ）」は選択肢にない。"])
    ws.append(["※ 利用率の伸びの画面は未確認。同じ5つなら⑤全国を選ぶ。"])

    ws = wb.create_sheet("04_手順")
    ws.append(["順", "すること", "留意点"])
    style_head(ws)
    for a, b, c in TEJUN:
        ws.append([a, b.replace("**", ""), c.replace("**", "")])
    body_style(ws, wrap=(2, 3))
    widths(ws, [4, 52, 62])
    ws.freeze_panes = "A2"

    p = OUT / f"小野町_見える化_在宅サービスの利用回数_{ASOF}.xlsx"
    wb.save(p)
    return p


# ---------------------------------------------------------------- main

def main():
    OUT.mkdir(parents=True, exist_ok=True)
    R = riyo_rows()
    bad = selfcheck(R)
    p1 = build_docx(R)
    p2 = build_xlsx(R)
    print("出力:", p1)
    print("     ", p2)
    print(f"  自然体推計 {SHIZEN:,}円（画面{GENZAI:,}円＋施策反映{abs(SAISA):,}円）")
    print("  1人1月あたり回（日）数")
    for nm, tani, r6, r7, r8, our, hi in R:
        o = f"／当方{our:.2f}" if our else ""
        print(f"    {nm:<16}R6 {r6:.2f}／R7 {r7:.2f}／R8 {r8:.2f}{tani}"
              f"（R8÷R7＝{hi:.2f}）{o}")
    print(f"  自己点検 {'OK' if not bad else 'NG: ' + ' / '.join(bad)}")


if __name__ == "__main__":
    main()
