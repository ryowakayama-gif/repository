"""障がい計画：事業所ヒアリングの依頼文一式（町長名版・受託者名版）と送付・日程管理表。

ヒアリングの中身（対象・質問・シート）は build_ono_shogai_hearing_sheet.py（令和8年10月22日）による。
依頼文の記載事項は、関係者団体等ヒアリング実施計画書（令和8年8月）の「参考　依頼文の記載事項」に従う
（件名・趣旨・依頼事項・事前準備のお願い・個人情報の取扱い・結果の利用・問合せ先・添付）。

名義（町長名か受託者名か）は町と協議中のため、両方を作る。町が決めた方を送る。

一式の中身
  1 依頼文
  2 日程調整・ご回答票（返信用）
  3 計画の概要とヒアリングのお願い（1枚）
  4 個人情報の取扱い
  ＋ ヒアリングシート（A 相談支援事業所／B サービス事業所。既存のExcel）を添付する

出力
  15_ヒアリング・調査設計/小野町_障がい計画_事業所ヒアリング_依頼文一式_町長名版_YYYYMMDD.docx
  15_ヒアリング・調査設計/小野町_障がい計画_事業所ヒアリング_依頼文一式_受託者名版_YYYYMMDD.docx
  15_ヒアリング・調査設計/小野町_障がい計画_事業所ヒアリング_送付・日程管理表_YYYYMMDD.xlsx
"""

import datetime
import pathlib
import re
import sys

import docx
import openpyxl
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl.styles import Alignment, Font, PatternFill

import build_ono_shogai_hearing_sheet as HS
import ono_shogai_zu as ZU

HERE = pathlib.Path(__file__).parent
OUT = HERE / "小野町_引継ぎ_整理済" / "15_ヒアリング・調査設計"
SHEET_XLSX = OUT / "小野町_障がい計画_事業所ヒアリングシート_20261022.xlsx"
ASOF = "20261023"
FONT = "游明朝"
SZ = [10.5]

# 日程（第2回自立支援協議会は令和8年11月下旬。結果の概要をそこで報告する）
SOFU = datetime.date(2026, 11, 4)      # 送付（目安）
KAITO = datetime.date(2026, 11, 10)    # 日程調整票の返信期限
KAISHI = datetime.date(2026, 11, 12)   # 実施期間
SHURYO = datetime.date(2026, 11, 20)
YOBI = "月火水木金土日"

# 問合せ先。町の連絡先は計画書の奥付（他メンバー素案）の記載による。受託者の担当は差し込む
TOWN = {"課": "小野町健康福祉課", "住所": "〒963-3492　福島県田村郡小野町大字小野新町字舘廻92",
        "TEL": "0247-72-6934", "FAX": "0247-72-3121"}
JUTAKU = "ビズアップ公共コンサルティング株式会社"
SASHIKOMI = ["【文書番号】", "【町長名】", "【町の担当者名】", "【受託者の担当者名】", "【受託者の電話】", "【受託者のメール】",
             "【事業所名】", "【送付日】"]

KENMEI = "おのまち障がい者計画・第4期障がい児福祉計画・第8期障がい福祉計画の策定に係るヒアリング調査へのご協力について（依頼）"


def wa(d):
    """令和の日付（曜日つき）。"""
    return "令和%d年%d月%d日（%s）" % (d.year - 2018, d.month, d.day, YOBI[d.weekday()])


def _font(run, sz=10.5, b=False):
    run.font.name = FONT
    run.font.size = Pt(sz)
    run.font.bold = b
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), FONT)


LS = [1.25]   # 行間（依頼文は1枚に収めるため狭くする）


def P(doc, text="", align=None, sz=None, b=False, ind=0, before=0, after=2):
    sz = sz or SZ[0]
    p = doc.add_paragraph()
    if align == "r":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Cm(ind)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = LS[0]
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            _font(p.add_run(part), sz, b or i % 2 == 1)
    return p


def T(doc, rows, widths, sz=9.5, head=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, widths):
        col.width = Cm(w)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.width = Cm(widths[j])
            c.paragraphs[0].text = ""
            for k, part in enumerate(re.split(r"\*\*", str(v))):
                if part:
                    _font(c.paragraphs[0].add_run(part), sz, (head and i == 0) or k % 2 == 1)
    return t


def page(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def sheet_items():
    """シートA・Bの、事前記入（確認）の質問番号と、当日伺う質問の数。"""
    out = {}
    for sh in ("A", "B"):
        ids = [i for i in HS.CLS if i.startswith(sh)]
        pre = [i for i in ids if HS.CLS[i][0] != HS.K_HEAR]
        out[sh] = (pre, len(ids) - len(pre), len(ids))
    return out


def letter(doc, ban):
    it = sheet_items()
    cho = ban == "町長名版"
    P(doc, "【文書番号】" if cho else "", align="r")
    P(doc, "令和8年11月　日（【送付日】）", align="r", after=8)
    P(doc, "【事業所名】")
    P(doc, "管理者　様", after=8)
    if cho:
        P(doc, "小野町長　【町長名】", align="r")
        P(doc, "（公印省略）", align="r", after=10)
    else:
        P(doc, JUTAKU, align="r")
        P(doc, "（小野町　障がい者計画等策定業務受託者）", align="r", after=10)
    P(doc, KENMEI, align="c", b=True, after=10)
    P(doc, "　時下ますますご清栄のこととお喜び申し上げます。" + ("日頃より本町の障がい福祉行政の推進に格別のご理解とご協力を賜り、厚く御礼申し上げます。"
                                                  if cho else "日頃より小野町の障がい福祉の推進にご尽力を賜り、厚く御礼申し上げます。"))
    P(doc, "　小野町では、現行の「おのまち障がい者計画・第3期障がい児福祉計画・第7期障がい福祉計画」の計画期間が"
           "令和8年度で満了することから、令和9年度から令和11年度までを計画期間とする次期計画を策定しております。")
    P(doc, "　次期計画では、障がい福祉サービス等の見込量と、その確保の方策を定めます。見込量を実際に利用できる量とするためには、"
           "事業所の皆様の受入れの状況や送迎、人材の確保などの実情を踏まえることが欠かせません。"
           "つきましては、ご多忙のところ誠に恐縮ですが、下記のとおりヒアリング調査にご協力くださいますようお願い申し上げます。")
    if not cho:
        P(doc, "　なお、本調査は、小野町からの委託を受けて当社が実施するものです。")
    P(doc, "記", align="c", before=6, after=6)
    P(doc, "１　内容", b=True)
    P(doc, "　同封のヒアリングシートに沿って、小野町にお住まいの方のご利用の状況、新たな受入れの余力、送迎、"
           "人材の確保、重度の方・医療的ケアが必要な方への対応、次期計画へのご意見を伺います。"
           "相談支援事業所の皆様にはシートA、サービス事業所・障がい児通所支援事業所の皆様にはシートBをお送りしています。", ind=0.4)
    P(doc, "２　方法と時間", b=True)
    P(doc, "　30分程度。訪問、電話、オンライン（Zoom等）のいずれかで、ご都合のよい方法をお選びください。", ind=0.4)
    P(doc, "３　実施期間", b=True)
    P(doc, "　%s から %s まで" % (wa(KAISHI), wa(SHURYO)), ind=0.4)
    P(doc, "４　事前のお願い", b=True)
    P(doc, "　シートの緑色の欄（シートAは%d問、シートBは%d問）は、当方が公表データ等をもとに記入しています。"
           "内容に誤りがないかご確認のうえ、同じ行の「ご回答」欄に人数・稼働率などをご記入いただけますと、当日の時間を短くできます。"
           "そのほかの質問（シートAは%d問、シートBは%d問）は当日伺いますので、ご記入いただく必要はありません。"
           % (len(it["A"][0]), len(it["B"][0]), it["A"][1], it["B"][1]), ind=0.4)
    P(doc, "５　ご回答", b=True)
    P(doc, "　同封の「日程調整・ご回答票」に、ご都合のよい日時と方法をご記入のうえ、%s までに、"
           "ファクス又は電子メールでご返信ください。日程が合わない場合も、その旨をご返信いただけますと幸いです。" % wa(KAITO), ind=0.4)
    P(doc, "６　結果の取扱い", b=True)
    P(doc, "　いただいたお話は、計画の策定と「アンケート調査結果及びヒアリング調査結果報告書」の作成にのみ用います。"
           "事業所名や、事業所が特定できる数値は公表しません。詳しくは同封の「個人情報の取扱い」をご覧ください。", ind=0.4)
    P(doc, "７　同封書類", b=True)
    P(doc, "　①ヒアリングシート（A又はB）　②日程調整・ご回答票　③計画の概要とヒアリングのお願い　④個人情報の取扱い", ind=0.4)
    P(doc, "８　お問合せ先", b=True)
    P(doc, "　（ヒアリングの日程・内容）%s　担当：【受託者の担当者名】　電話：【受託者の電話】　電子メール：【受託者のメール】" % JUTAKU, ind=0.4)
    P(doc, "　（計画の策定）%s　担当：【町の担当者名】　電話：%s　ファクス：%s" % (TOWN["課"], TOWN["TEL"], TOWN["FAX"]), ind=0.4)
    P(doc, "以上", align="r", before=6)


def kaito(doc):
    P(doc, "日程調整・ご回答票", align="c", sz=13, b=True, after=4)
    P(doc, "返信先：%s（ファクス又は電子メール）　返信期限：%s" % (JUTAKU, wa(KAITO)), align="c", sz=9.5, after=8)
    P(doc, "ファクス：【受託者のファクス】　電子メール：【受託者のメール】", align="c", sz=9.5, after=8)
    rows = [["項目", "ご記入欄"],
            ["事業所名", ""],
            ["ご回答者（職・氏名）", ""],
            ["ご連絡先（電話・電子メール）", ""],
            ["ご希望の方法", "□ 訪問　　□ 電話　　□ オンライン（Zoom等）"],
            ["第1希望の日時", "　　月　　日（　）　□午前　□午後　　時頃"],
            ["第2希望の日時", "　　月　　日（　）　□午前　□午後　　時頃"],
            ["第3希望の日時", "　　月　　日（　）　□午前　□午後　　時頃"],
            ["町職員の同席", "□ 希望する　　□ 希望しない　　□ どちらでもよい"],
            ["記録のための録音", "□ 差し支えない　　□ 録音しないでほしい"],
            ["ご参加が難しい場合", "□ 書面（シートへの記入）での回答を希望する　　□ 今回は見送る"],
            ["ご質問・ご連絡事項", "\n\n"]]
    T(doc, rows, [5.0, 11.0], sz=10)
    P(doc, "※ 実施期間は %s から %s までです。土日・祝日を除きます。" % (wa(KAISHI), wa(SHURYO)), sz=9, before=6)
    P(doc, "※ 録音は記録の正確を期すためのもので、要点の記録を作成した後に消去します。", sz=9)


def gaiyo(doc):
    sv = lambda k: ZU.pct("集計_障がい者", "問31-2", "利用したい", koumoku=k)
    P(doc, "計画の概要とヒアリングのお願い", align="c", sz=13, b=True, after=8)
    P(doc, "■ 策定する計画", b=True)
    T(doc, [["計画", "根拠", "定めること"],
            ["おのまち障がい者計画", "障害者基本法第11条第3項", "障がいのある方のための施策の基本的な方向"],
            ["第8期障がい福祉計画", "障害者総合支援法第88条第1項", "成果目標、障がい福祉サービス・地域生活支援事業の見込量と確保の方策"],
            ["第4期障がい児福祉計画", "児童福祉法第33条の20第1項", "障がい児通所支援・障がい児相談支援の見込量と確保の方策"]],
      [4.2, 4.6, 7.2], sz=9)
    P(doc, "計画期間：令和9年度から令和11年度までの3年間", sz=9.5, before=4, after=6)
    P(doc, "■ 策定の流れ", b=True)
    T(doc, [["時期", "内容"],
            ["令和8年8〜9月", "障がい福祉に関するアンケート調査（障がいのある方・障がいのある子どもの保護者）"],
            ["令和8年11月", "**事業所の皆様へのヒアリング（本調査）**"],
            ["令和8年11月下旬", "小野町自立支援協議会（第2回）でヒアリングの概要を報告"],
            ["令和8年12月〜令和9年1月", "計画素案の作成、自立支援協議会での協議、パブリックコメント"],
            ["令和9年3月", "計画の策定"]],
      [4.6, 11.4], sz=9)
    P(doc, "■ 事業所の皆様に伺いたい理由", b=True, before=6)
    P(doc, "アンケート調査では、今後利用したいサービスとして「相談支援」%.1f％、「居宅介護」%.1f％、「生活介護」%.1f％、"
           "「短期入所」%.1f％などが挙がりました（障がいのある方、回答259件）。"
           "一方で、こうした利用の希望を実際に受け止められるかは、事業所の皆様の受入れの余力、送迎の範囲、人材の確保によって決まります。"
           "**これらはアンケートや公表データでは分からないため、直接お話を伺いたいと考えています。**"
           % (sv("相談支援"), sv("居宅介護"), sv("生活介護"), sv("短期入所")), sz=9.5)
    P(doc, "伺ったお話は、計画の「見込量の確保の方策」と、相談支援の体制、人材の確保、重度の方・医療的ケアが必要な方への支援の施策に反映します。", sz=9.5)
    P(doc, "資料：令和8年度 障がい福祉に関するアンケート調査（問31：今後の障がい福祉サービスの利用意向。回答者全体に対する割合）", sz=8)


def kojin(doc):
    P(doc, "個人情報の取扱い", align="c", sz=13, b=True, after=8)
    for t in ["１　利用の目的　ヒアリングでいただいたお話は、おのまち障がい者計画・第4期障がい児福祉計画・第8期障がい福祉計画の策定と、"
              "「アンケート調査結果及びヒアリング調査結果報告書」の作成にのみ用います。",
              "２　公表の範囲　計画書・報告書には、複数の事業所のお話をまとめた形で載せます。事業所名、ご回答者のお名前、"
              "事業所が特定できる数値（定員・利用者数・職員数など）は公表しません。事業所名を出す必要が生じた場合は、事前にご承諾をいただきます。",
              "３　利用者の方の情報　個々の利用者の方が特定できる情報（お名前・住所・障がいの状況など）は伺いません。"
              "事例をお話しいただく場合も、個人が特定されない形でお願いします。",
              "４　記録　承諾をいただいた場合に限り録音し、要点の記録を作成した後に消去します。"
              "要点の記録は、ご確認いただいたうえで確定します。",
              "５　管理　記録は小野町及び受託者が適切に管理し、計画の策定後、定められた期間の経過後に廃棄します。"
              "受託者は、小野町との委託契約に定める秘密の保持の義務を負います。",
              "６　お問合せ　個人情報の取扱いに関するお問合せは、%s（電話 %s）までお願いします。" % (TOWN["課"], TOWN["TEL"])]:
        P(doc, t, sz=10, after=6)


def build(ban):
    doc = docx.Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(1.7)
    LS[0], SZ[0] = 1.05, 10.0
    letter(doc, ban)
    LS[0], SZ[0] = 1.25, 10.5
    page(doc)
    kaito(doc)
    page(doc)
    gaiyo(doc)
    page(doc)
    kojin(doc)
    p = OUT / f"小野町_障がい計画_事業所ヒアリング_依頼文一式_{ban}_{ASOF}.docx"
    doc.save(p)
    return p


def kanri():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "送付・日程管理"
    cols = ["No", "区分", "事業所名", "法人名", "サービス種別", "所在", "シート", "送付日", "回答日", "方法", "実施日時",
            "町職員の同席", "録音", "記録の確認日", "備考"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate([5, 14, 32, 28, 22, 10, 7, 10, 10, 10, 16, 10, 8, 12, 24]):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    n = 0
    for kind, r in HS.candidates():
        if not kind.startswith("優先"):
            continue
        n += 1
        sv = r["サービス種別"]
        ws.append([n, "町内", r["事業所の名称"], r["法人の名称"], sv, r["事業所住所（市区町村）"].replace("福島県", ""),
                   "A" if "相談支援" in sv else "B"] + [""] * 8)
    for k in range(5):
        ws.append([n + k + 1, "町外（選定待ち）", "", "", "", "", "", "", "", "", "", "", "", "", "町のデータで、小野町の方の利用が多い先を選ぶ"])
    ws.append([])
    ws.append(["", "※ 町内は障害福祉サービス等情報公表システム オープンデータ（令和8年3月版）の全事業所。町外は町のデータで5か所程度を選ぶ（合わせて10〜12か所）"])
    w2 = wb.create_sheet("日程")
    w2.append(["日", "曜", "午前", "午後", "備考"])
    d = KAISHI
    while d <= SHURYO:
        if d.weekday() < 5:
            w2.append([wa(d).split("（")[0], YOBI[d.weekday()], "", "", ""])
        d += datetime.timedelta(days=1)
    w3 = wb.create_sheet("差し込む事項")
    w3.append(["差し込む事項", "決める人", "期限", "状態"])
    for x, who, due in [("名義（町長名版か受託者名版か）", "町", "10月30日"), ("【文書番号】【町長名】（町長名版のみ）", "町", "送付前"),
                        ("【町の担当者名】", "町", "送付前"), ("【受託者の担当者名】【受託者の電話】【受託者のメール】【受託者のファクス】", "受託者", "送付前"),
                        ("【事業所名】（宛名。この表の事業所名を差し込む）", "受託者", "送付前"), ("【送付日】", "町・受託者", wa(SOFU) + "目安"),
                        ("町外の対象先（5か所程度）", "町", "11月上旬"), ("町職員の同席の要否", "町", "10月30日")]:
        w3.append([x, who, due, "未"])
    p = OUT / f"小野町_障がい計画_事業所ヒアリング_送付・日程管理表_{ASOF}.xlsx"
    wb.save(p)
    return p, n


def main():
    ng = []
    ps = [build(b) for b in ("町長名版", "受託者名版")]
    px, n_town = kanri()
    # 自己点検
    for d, nm in ((SOFU, "送付"), (KAITO, "返信期限"), (KAISHI, "実施の初日"), (SHURYO, "実施の最終日")):
        if d.weekday() >= 5:
            ng.append("%sが土日：%s" % (nm, d))
    if not (SOFU < KAITO < KAISHI <= SHURYO < datetime.date(2026, 11, 24)):
        ng.append("日程の順序が合わない（第2回協議会の前に終わること）")
    if datetime.date(2026, 11, 3) in (SOFU, KAITO, KAISHI, SHURYO):
        ng.append("文化の日（11月3日）に日程を置いている")
    it = sheet_items()
    if it["A"][2] + it["B"][2] != len(HS.CLS):
        ng.append("シートA・Bの質問数が仕分けと合わない")
    for p in ps:
        t = "\n".join(x.text for x in docx.Document(str(p)).paragraphs)
        t += "\n".join(c.text for tb in docx.Document(str(p)).tables for r in tb.rows for c in r.cells)
        for w in ("北塩原", "大雪", "川崎", "金ヶ崎", "福島町", "訂正", "お詫び", "おわび", "**"):
            if w in t:
                ng.append("%s に「%s」が出ている" % (p.name, w))
        rest = sorted(set(re.findall(r"【[^】]+】", t)) - set(SASHIKOMI) - {"【受託者のファクス】"})
        if rest:
            ng.append("%s に差し込みの一覧にない【】がある：%s" % (p.name, rest))
        for w in (KENMEI, "個人情報の取扱い", "日程調整・ご回答票", wa(KAITO), wa(KAISHI), wa(SHURYO)):
            if w not in t:
                ng.append("%s に「%s」がない" % (p.name, w))
    if not SHEET_XLSX.exists():
        ng.append("同封するヒアリングシート（Excel）がない")
    if n_town != 7:
        ng.append("町内の対象先が%d件（7件のはず）" % n_town)
    print("■ 事業所ヒアリングの依頼文一式")
    for p in ps + [px]:
        print("   ", p.relative_to(HERE))
    print("   送付 %s／返信 %s／実施 %s〜%s／町内の対象先 %d" % (wa(SOFU), wa(KAITO), wa(KAISHI), wa(SHURYO), n_town))
    print("   事前記入 シートA %d問・シートB %d問／当日伺う A %d問・B %d問"
          % (len(it["A"][0]), len(it["B"][0]), it["A"][1], it["B"][1]))
    print("\n■ 自己点検")
    for x in ng:
        print("   NG", x)
    print("   すべて適合" if not ng else "   **%d件**" % len(ng))
    return 1 if ng else 0


if __name__ == "__main__":
    sys.exit(main())
