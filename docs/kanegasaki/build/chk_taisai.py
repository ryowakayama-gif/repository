# -*- coding: utf-8 -*-
"""docx の体裁の実測値を一覧にする。前回計画書との対照表の「本計画」欄になる。"""
import sys, collections
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

for path in sys.argv[1:]:
    d = Document(path)
    print('■', path)
    s = d.sections[0]
    cm = lambda v: None if v is None else round(v.cm, 2)
    print('  用紙 %s×%s cm／余白 上%s 下%s 左%s 右%s／ヘッダ%s フッタ%s'
          % (cm(s.page_width), cm(s.page_height), cm(s.top_margin), cm(s.bottom_margin),
             cm(s.left_margin), cm(s.right_margin), cm(s.header_distance), cm(s.footer_distance)))
    n = d.styles['Normal']
    rf = n.element.rPr.rFonts
    print('  Normal　和文 %s／欧文 %s／%s pt／行間 %s／段落後 %s pt'
          % (rf.get(qn('w:eastAsia')), rf.get(qn('w:ascii')),
             n.font.size.pt if n.font.size else None,
             n.paragraph_format.line_spacing,
             n.paragraph_format.space_after.pt if n.paragraph_format.space_after else None))
    # 本文の run のフォント・サイズ・色の分布
    cnt = collections.Counter()
    for p in d.paragraphs:
        for r in p.runs:
            if not r.text.strip():
                continue
            f = r.font
            col = f.color.rgb if (f.color and f.color.type is not None and f.color.rgb) else None
            cnt[(str(f.name), f.size.pt if f.size else None, bool(r.bold), str(col))] += len(r.text)
    print('  本文の run（文字数の多い順・上位8）')
    for (fn, sz, b, col), c in cnt.most_common(8):
        print('    %-8s %-6s %-5s %-8s  %6d字' % (fn, sz, '太字' if b else '', col, c))
    # 表の中
    cnt2 = collections.Counter()
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        if not r.text.strip():
                            continue
                        f = r.font
                        col = f.color.rgb if (f.color and f.color.type is not None and f.color.rgb) else None
                        cnt2[(str(f.name), f.size.pt if f.size else None, bool(r.bold), str(col))] += len(r.text)
    print('  表の中の run（上位5）')
    for (fn, sz, b, col), c in cnt2.most_common(5):
        print('    %-8s %-6s %-5s %-8s  %6d字' % (fn, sz, '太字' if b else '', col, c))
    # ヘッダ・フッタ
    for lab, part in (('ヘッダ', s.header), ('フッタ', s.footer)):
        txt = ' / '.join(p.text.strip() for p in part.paragraphs if p.text.strip())
        fld = 'PAGE フィールドあり' if 'PAGE' in part._element.xml else ''
        print('  %s　%s %s' % (lab, txt or '（文字なし）', fld))
    print('  段落 %d／表 %d／図（インライン画像）%d'
          % (len(d.paragraphs), len(d.tables), part._element.xml.count('<pic:pic') + sum(
              p._element.xml.count('<pic:pic') for p in d.paragraphs)))
    print()
