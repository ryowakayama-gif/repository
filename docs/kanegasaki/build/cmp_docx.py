# -*- coding: utf-8 -*-
"""2つの docx を本文・表・図・書式で突き合わせる。復元の同一性の検証に用いる。"""
import sys, hashlib, zipfile, difflib
import docx
from docx.oxml.ns import qn


def sig(path):
    d = docx.Document(path)
    paras = {id(p._p): p for p in d.paragraphs}
    tbls = {id(t._tbl): t for t in d.tables}
    items = []
    for el in d.element.body:
        if el.tag == qn('w:tbl'):
            t = tbls[id(el)]
            rows = ['\x1f'.join(c.text for c in r.cells) for r in t.rows]
            items.append(('TBL', '\x1e'.join(rows)))
        elif el.tag == qn('w:p'):
            p = paras.get(id(el))
            if p is None:
                continue
            if any('PAGEREF' in (it.text or '') for it in p._p.findall('.//' + qn('w:instrText'))):
                items.append(('TOC', p.text.split('\t')[0]))
                continue
            if p._p.findall('.//' + qn('pic:pic')):
                items.append(('FIG', ''))
                continue
            t = p.text
            if not t.strip():
                continue
            runs = [(r.text, bool(r.bold),
                     float(r.font.size.pt) if r.font.size else None) for r in p.runs if r.text]
            items.append(('P', t + '\x1f' + repr(runs)))
    return items


def media(path):
    z = zipfile.ZipFile(path)
    return sorted(hashlib.md5(z.read(n)).hexdigest()
                  for n in z.namelist() if n.startswith('word/media/'))


a, b = sys.argv[1], sys.argv[2]
A, B = sig(a), sig(b)
print('%-28s 要素 %d（本文 %d・表 %d・図 %d・目次 %d）'
      % ('元', len(A), sum(1 for k, _ in A if k == 'P'), sum(1 for k, _ in A if k == 'TBL'),
         sum(1 for k, _ in A if k == 'FIG'), sum(1 for k, _ in A if k == 'TOC')))
print('%-28s 要素 %d（本文 %d・表 %d・図 %d・目次 %d）'
      % ('復元', len(B), sum(1 for k, _ in B if k == 'P'), sum(1 for k, _ in B if k == 'TBL'),
         sum(1 for k, _ in B if k == 'FIG'), sum(1 for k, _ in B if k == 'TOC')))
print('画像', '一致' if media(a) == media(b) else '不一致')

ka = ['%s|%s' % x for x in A]
kb = ['%s|%s' % x for x in B]
sm = difflib.SequenceMatcher(None, ka, kb, autojunk=False)
ndiff = 0
for op, i1, i2, j1, j2 in sm.get_opcodes():
    if op == 'equal':
        continue
    ndiff += max(i2 - i1, j2 - j1)
    print('\n--- %s 元[%d:%d] 復元[%d:%d]' % (op, i1, i2, j1, j2))
    for x in ka[i1:i2][:6]:
        print('  元  ', x[:160])
    for x in kb[j1:j2][:6]:
        print('  復元', x[:160])
print('\n差異のある要素', ndiff, '／', max(len(A), len(B)))
