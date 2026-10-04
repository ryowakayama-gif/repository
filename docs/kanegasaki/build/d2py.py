# -*- coding: utf-8 -*-
"""docx から fmt.py の呼出し（ch_*.py）を復元する。

　fmt.py の P／H1／H2／H3／BUL／TBL／SRC／FIG／TOC_HERE／BODY_HERE／BUILD_TOC が
　生成する書式を手がかりに、本文を組むPythonへ戻す。
　図は out/fig/ の再出力とバイト比較して元のファイル名を割り当てる。
"""
import sys, os, re, hashlib, zipfile
import docx
from docx.oxml.ns import qn

W = qn  # 短縮


def sz_of(r):
    return float(r.font.size.pt) if r.font.size else None


def col_of(r):
    try:
        if r.font.color is not None and r.font.color.type is not None:
            return str(r.font.color.rgb)
    except Exception:
        pass
    return None


def pt(v):
    """EMU/Length → pt（None は None）"""
    return None if v is None else round(v.pt, 1)


def cm(v):
    return None if v is None else round(v.cm, 2)


def enc(p):
    """段落の文字列を ** による太字表記に戻す。"""
    out = []
    for r in p.runs:
        t = r.text
        if not t:
            continue
        if r.bold:
            if out and out[-1][0]:
                out[-1] = (True, out[-1][1] + t)
            else:
                out.append((True, t))
        else:
            if out and not out[-1][0]:
                out[-1] = (False, out[-1][1] + t)
            else:
                out.append((False, t))
    s = ''
    for b, t in out:
        s += ('**%s**' % t) if b else t
    return s


def q(s):
    """Python のリテラルにする。"""
    return "'" + s.replace('\\', '\\\\').replace("'", "\\'").replace('\n', '\\n') + "'"


def has_pic(p):
    return bool(p._p.findall('.//' + qn('pic:pic')))


def pic_name(p, part, imgmap):
    for blip in p._p.findall('.//' + qn('a:blip')):
        rid = blip.get(qn('r:embed'))
        ip = part.related_parts[rid]
        h = hashlib.md5(ip.blob).hexdigest()
        return imgmap.get(h)
    return None


def is_pagebreak(p):
    return bool(p._p.findall('.//' + qn('w:br'))) and not p.text.strip()


def has_sectpr(p):
    pPr = p._p.find(qn('w:pPr'))
    return pPr is not None and pPr.find(qn('w:sectPr')) is not None


def is_tocline(p):
    for it in p._p.findall('.//' + qn('w:instrText')):
        if 'PAGEREF' in (it.text or ''):
            return True
    return False


def tbl_rows(t):
    rows = []
    for r in t.rows:
        row = []
        for c in r.cells:
            # セル内は段落1つを前提（fmt.TBL がそう作る）。複数なら改行で連結
            row.append('\n'.join(enc(p) for p in c.paragraphs).strip('\n'))
        rows.append(row)
    return rows


def tbl_widths(t):
    grid = t._tbl.find(qn('w:tblGrid'))
    if grid is None:
        return None
    ws = []
    for gc in grid.findall(qn('w:gridCol')):
        v = gc.get(qn('w:w'))
        if v is None:
            return None
        ws.append(round(int(v) / 567.0, 2))      # twips → cm
    return ws


def tbl_fs(t):
    if len(t.rows) < 2:
        return None
    for c in t.rows[1].cells:
        for p in c.paragraphs:
            for r in p.runs:
                if r.font.size:
                    return float(r.font.size.pt)
    return None


ALIGN = {1: 'WD_ALIGN_PARAGRAPH.CENTER', 2: 'WD_ALIGN_PARAGRAPH.RIGHT',
         3: 'WD_ALIGN_PARAGRAPH.JUSTIFY'}


def convert(path, out, body_pt=10.5):
    d = docx.Document(path)
    part = d.part
    z = zipfile.ZipFile(path)
    imgmap = {}
    for fn in sorted(os.listdir('out/fig')):
        imgmap[hashlib.md5(open('out/fig/' + fn, 'rb').read()).hexdigest()] = fn

    body = d.element.body
    paras = {id(p._p): p for p in d.paragraphs}
    tbls = {id(t._tbl): t for t in d.tables}

    L = []                      # 出力行
    figno = 0
    i = 0
    kids = list(body)
    n = len(kids)
    skip_toc = False
    seen_toc = False
    pend_after_tbl = False

    while i < n:
        el = kids[i]
        tag = el.tag
        if tag == qn('w:tbl'):
            t = tbls[id(el)]
            rows = tbl_rows(t)
            ws = tbl_widths(t)
            fs = tbl_fs(t)
            a = ['[\n' + ''.join('    [%s],\n' % ', '.join(q(c) for c in r) for r in rows) + ']']
            if ws:
                a.append('widths=[%s]' % ', '.join(str(x) for x in ws))
            if fs and abs(fs - 8.5) > 0.01:
                a.append('fs=%s' % fs)
            L.append('TBL(%s)' % ', '.join(a))
            pend_after_tbl = True
            i += 1
            continue
        if tag != qn('w:p'):
            i += 1
            continue
        p = paras.get(id(el))
        if p is None:
            i += 1
            continue
        txt = p.text.strip()

        # 表の直後に fmt.TBL が置く空段落は飛ばす
        if pend_after_tbl:
            pend_after_tbl = False
            if not txt and not has_pic(p):
                i += 1
                continue

        # 目次の範囲
        if skip_toc:
            if has_sectpr(p):
                L.append('BODY_HERE()')
                skip_toc = False
                i += 1
                continue
            i += 1
            continue
        if txt in ('目　次', '目次') and not seen_toc:
            seen_toc = True
            skip_toc = True
            L.append("TOC_HERE(%s)" % (q(txt) if txt != '目　次' else ''))
            i += 1
            continue
        if has_sectpr(p) and not txt:
            L.append('BODY_HERE()')
            i += 1
            continue
        if is_pagebreak(p):
            # TOC_HERE / BODY_HERE が自ら入れる改ページは落とす
            nxt = kids[i + 1] if i + 1 < n else None
            nt = paras.get(id(nxt)).text.strip() if (nxt is not None and nxt.tag == qn('w:p') and id(nxt) in paras) else ''
            if nt in ('目　次', '目次'):
                i += 1
                continue
            L.append('doc.add_page_break()')
            i += 1
            continue
        if has_pic(p):
            nm = pic_name(p, part, imgmap)
            L.append('# 画像の段落（キャプション無し）: %s' % nm)
            i += 1
            continue
        if not txt:
            L.append('P()')
            i += 1
            continue

        runs = [r for r in p.runs if r.text]
        s0 = sz_of(runs[0]) if runs else None
        c0 = col_of(runs[0]) if runs else None
        b0 = bool(runs[0].bold) if runs else False
        pf = p.paragraph_format
        bef, aft = pt(pf.space_before), pt(pf.space_after)
        ind = cm(pf.left_indent)
        al = p.alignment

        # 図（キャプション → 画像 → 出典）
        if s0 == 11.0 and b0 and c0 == '1F4D78' and txt.startswith('【図'):
            # FIG() がキャプション全体を太字の1ランにするため、** は付けない
            m = re.match(r'^【図(\d+)】(.*)$', txt)
            cap = m.group(2) if m else txt
            figno += 1
            j = i + 1
            nm, wid, src = None, None, None
            if j < n and kids[j].tag == qn('w:p'):
                pj = paras.get(id(kids[j]))
                if pj is not None and has_pic(pj):
                    nm = pic_name(pj, part, imgmap)
                    ext = pj._p.findall('.//' + qn('wp:extent'))
                    if ext:
                        wid = round(int(ext[0].get('cx')) / 360000.0, 2)
                    j += 1
            if j < n and kids[j].tag == qn('w:p'):
                pj = paras.get(id(kids[j]))
                if pj is not None and pj.text.strip():
                    rj = [r for r in pj.runs if r.text]
                    if rj and sz_of(rj[0]) == 8.5 and cm(pj.paragraph_format.left_indent) in (0.2, 0.19, 0.21):
                        src = enc(pj)
                        j += 1
            a = ["'out/fig/%s'" % nm, q(cap)]
            if wid and abs(wid - 15.5) > 0.05:
                a.append('width=%s' % wid)
            if src:
                a.append('src=%s' % q(src))
            L.append('FIG(%s)' % ', '.join(a))
            i = j
            continue

        e = enc(p)
        # 見出し
        if s0 == 15.0 and b0 and c0 == '2E74B5':
            L.append('H1(%s)' % q(txt)); i += 1; continue
        if s0 == 12.5 and b0 and c0 == '2E74B5':
            L.append('H2(%s)' % q(txt)); i += 1; continue
        if s0 == 11.0 and b0 and c0 == '1F4D78':
            L.append('H3(%s)' % q(txt)); i += 1; continue
        # 出典
        if s0 == 8.5 and ind in (0.2, 0.19, 0.21) and bef == 2.0 and aft == 8.0:
            L.append('SRC(%s)' % q(e)); i += 1; continue
        # 箇条書き
        if s0 == 10.0 and e.startswith('・') and ind in (0.4, 0.39, 0.41) and aft == 2.0:
            a = [q(e[1:])]
            if ind not in (0.4, 0.39, 0.41):
                a.append('ind=%s' % ind)
            L.append('BUL(%s)' % ', '.join(a)); i += 1; continue
        # 一般の段落
        a = [q(e)]
        if s0 is not None and abs(s0 - body_pt) > 0.01:
            a.append('sz=%s' % s0)
        if bef not in (0.0, None):
            a.append('before=%s' % bef)
        if aft not in (4.0, None):
            a.append('after=%s' % aft)
        if ind:
            a.append('ind=%s' % ind)
        if al is not None and int(al) in ALIGN:
            a.append('align=%s' % ALIGN[int(al)])
        if c0 and c0 != 'None' and not (s0 == 11.0 and b0):
            a.append('color=RGBColor(0x%s, 0x%s, 0x%s)' % (c0[0:2], c0[2:4], c0[4:6]))
        L.append('P(%s)' % ', '.join(a))
        i += 1

    with open(out, 'w', encoding='utf-8') as f:
        f.write('# -*- coding: utf-8 -*-\n')
        f.write('# %s の本文（docx から復元）\n\n' % os.path.basename(path))
        for line in L:
            f.write(line + '\n')
    print('%s → %s　%d 行　図%d' % (path, out, len(L), figno))


if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
