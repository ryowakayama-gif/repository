# -*- coding: utf-8 -*-
"""提供依頼資料（対象期間・定義を補った版）を組む。

　元の「提供依頼資料」（資料名9件）の No. と資料名はそのまま残し、
　対象期間・ほしい区分と定義・素案での用途・当方で確認済みの範囲・
　対応する照会番号・提出希望時期と、町の記入欄を加える。10件目に資-01を加える。
　体裁は元のファイルに合わせる（游ゴシック・細罫線）。印刷はA4横・幅合わせ。

　python3 -I gen_irai.py [出力先]
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else 'out/提供依頼資料_対象期間と定義を補った版.xlsx'
FN = '游ゴシック'

HDR = ['No.', '資料名', '対象期間・時点', 'ほしい区分・定義', '素案での用途',
       '当方で確認済みの範囲', '対応する照会番号', '提出希望時期',
       '【町記入】\n提出状況', '【町記入】\n提出（予定）日', '【町記入】\n回答・備考']
WID = [5, 30, 18, 46, 24, 40, 16, 12, 12, 13, 30]

ROWS = [
    (1, '介護給付費準備基金管理台帳・残高見込表・取崩予定額',
     '令和6〜8年度\n（令和8年度は決算見込）',
     '年度ごとの取崩額（繰入金）・積立額（利子を含む）・年度末残高。\n'
     '令和8年度は積立を含む決算見込。\n'
     '令和7年度は繰入金13,092,000円と基金運用状況調書の減額13,511,850円の差419,850円の内訳。\n'
     '第10期の取崩方針（案）があれば併せて',
     '第3章4\n第8章5⑩（取崩額の3案）',
     '令和7年度末残高271,741,782円（令和7年度決算）。\n'
     '令和8年度当初予算の繰入金21,581千円・積立金460千円（令和8年度末残高見込250,620,782円）',
     '至急《要回答4・39》', 'Ａ（至急）'),
    (2, '第1号被保険者・所得段階別人数実績表（介護保険料賦課資料）',
     '令和6〜8年度の\n当初賦課時点',
     '13段階別の人数。\n'
     '本町の条例上の割合（第2段階0.65ほか）と公費軽減後の割合を併記。\n'
     '住所地特例・適用除外（障害者支援施設等の入所者）の扱いを明記',
     '第8章5⑦⑧・第8章6\n別表5',
     '第9期の段階と割合（第9期計画書・概要版）。\n'
     '別表5は見える化システムの登録用データによる推計値（3か年計14,252人）であり、賦課の実績ではない',
     '《要回答40》《要回答70》', 'Ｂ'),
    (3, '介護保険料の調定・収納実績表、予定収納率の算定根拠',
     '令和6〜8年度\n（令和8年度は直近月まで）',
     '現年度分（特別徴収・普通徴収の別）と滞納繰越分の、調定額・収納額・収納率',
     '第8章5⑨（予定収納率）',
     '令和7年度決算書により現年度分の収納率99.80％',
     '《要回答40》', 'Ｂ'),
    (4, '要介護認定者の更新前後状態変化集計表（改善・維持・悪化）',
     '令和6・7年度\n（可能なら令和5年度から）',
     '更新前の要介護度別に、更新後の改善・維持・悪化の人数。\n'
     '更新と区分変更の別、死亡・未更新・転出の扱いを明記。\n'
     '要介護1から要介護2以上へ移った方の、移る前の利用サービス（わかる範囲で）',
     '第2章4（要介護1の重度化）\n第3章（維持・改善の割合）\n第9章3（重度・軽度別の認定率）',
     '見える化システムによる要介護1の更新時の移行率24.1％→57.0％。\n'
     '介護予防ケアマネジメントの維持・改善の割合76.5％→78.7％',
     '《要回答14》\n（本件に統合）', 'Ａ'),
    (5, '国保連の介護給付実績データ（要介護度・サービス種類・月別）',
     '令和6年4月サービス分〜\n令和8年8月サービス分\n（確定分）',
     'サービス種類別・要介護度別・月別の受給者数（実人数）と給付費。\n'
     '件数と実人数の別を明記（複数事業所を利用する方の重複の扱い）。\n'
     '介護予防サービス・総合事業を含む',
     '第3章2（計画値と実績の対比・別表1）\n第8章1〜4（令和8年度を実績に置き換える）',
     '令和7年度決算の保険給付費1,318,533,717円。\n'
     '月報は令和8年4月分まで受領していたが、作業環境の作り直しにより控えを失っている',
     '至急《要回答35》\n《要回答17》\n照会文書 第5の6', 'Ａ（至急）'),
    (6, '47行政区と6生活圏域の対応表、地区別高齢者人口集計',
     '令和8年3月末または9月末\n（基準日を明記）',
     '47行政区それぞれが属する生活圏（街・三ケ尻・南方・西部・永岡・北部）。\n'
     '行政区別の総人口・65歳以上・75歳以上・85歳以上',
     '第1章5・第2章1\n第4章9（地区別の状況）\n第7章1・第7章6',
     '6生活圏の定義と境界（金ケ崎町都市計画マスタープラン）。\n'
     '行政区別の人口・高齢化率（アンケート調査結果報告書 資料編）',
     '《要回答24》', 'Ａ'),
    (7, '介護サービス事業所・施設別の定員、稼働率、人員配置実績',
     '令和8年度（直近月）',
     '施設・居住系は定員と月末の入所者数、在宅系は利用者数と受入可能量。\n'
     '職種別の職員数（常勤換算）',
     '第7章2（サービス提供体制の確保方策）\n第8章3（見込量）\n資料編（事業所一覧）',
     '町内33事業所・15サービス種類（介護サービス情報公表システム 令和8年6月末時点）。\n'
     '同システムの定員の欄は任意記載で空欄が多い',
     '《要回答3》\n《要回答36・37》\n照会文書 第5の2', 'Ｂ'),
    (8, '通いの場・オレンジカフェの活動一覧と年間実績',
     '令和6・7年度\n（令和8年度は直近）',
     '場ごとの開催頻度（週1回以上の別）・参加者の実人数と延べ人数・運営主体。\n'
     '国の調査でいう通いの場とオレンジカフェの重なりの有無',
     '第3章1（成果指標②）\n第6章 基本方針1(2)（参加率の目標）\n第9章3',
     '国の調査による通いの場13か所・152人（令和6年度）。\nオレンジカフェ53か所（令和7年度）',
     '《要回答12》', 'Ａ'),
    (9, '一般介護予防事業評価事業・地域リハ活動支援事業の未実施理由／実施方針（確認回答）',
     '令和6〜8年度',
     '2事業それぞれについて、未実施の理由と第10期の実施方針。\n'
     '地域リハビリテーション活動支援事業は、胆江地域リハビリテーション広域支援センター（美希病院）の派遣調整の枠組みを使えるか',
     '第3章3\n第6章 基本方針1(2)(4)\n第9章2',
     '国の総合事業等の実施状況調査で2事業とも「実施無」。\n'
     '胆江圏域に地域リハビリテーション連絡協議会と派遣調整ガイドラインがあること（岩手県地域リハビリテーション連携指針）',
     '《要回答11》\n《要回答58》', 'Ａ'),
    (10, '【追加】認知症高齢者の日常生活自立度別人数（岩手県の調査への町の回答）',
     '令和3〜8年度',
     'Ⅰ・Ⅱa・Ⅱb・Ⅲa・Ⅲb・Ⅳ・Ｍの別。\n調査の基準日。第2号被保険者を含むか否か',
     '第2章3（認知症高齢者等の推移）\n第8章4（認知症高齢者数の推計）',
     '見える化システムＢ7（令和6年10月末。自立度Ⅱ以上69.3％）。\n'
     '県の調査による令和3〜5年度の値（480〜535人）。両者の系列が連続するか未確認',
     '《要回答51》', 'Ａ'),
]

thin = Side(style='thin', color='808080')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F = lambda **k: Font(name=FN, size=k.pop('size', 10), **k)
TOP = Alignment(vertical='top', wrap_text=True)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)
G_HDR = PatternFill('solid', fgColor='D9D9D9')      # 見出し
G_TOWN = PatternFill('solid', fgColor='FFF2CC')     # 町の記入欄
G_EX = PatternFill('solid', fgColor='F2F2F2')       # 記入例
G_ADD = PatternFill('solid', fgColor='E2EFDA')      # 追加した行

wb = Workbook()
ws = wb.active
ws.title = '提供依頼資料'
wb.properties.creator = '策定支援業務受託者'
wb.properties.title = '提供依頼資料（対象期間・定義を補った版）'

ws['A1'] = '提供依頼資料（対象期間・定義を補った版）'
ws['A1'].font = F(size=14, b=True)
ws['A2'] = ('金ケ崎町　高齢者福祉計画・第10期介護保険事業計画・認知症施策推進計画　　令和8年10月20日　'
            '計画素案の完成のために町にお願いする資料です。No.1〜9は元の一覧のとおり、No.10を追加しました。')
ws['A2'].font = F()
ws['A3'] = ('【記入方法】黄色の3列（提出状況・提出（予定）日・回答・備考）にご記入ください。提出状況はプルダウンから選びます。'
            '資料の写しでご提供いただく場合は、備考に「資料で提供」とご記入ください。5行目は記入例です。')
ws['A3'].font = F(b=True)
ws['A4'] = ('【提出希望時期】Ａ＝10月30日まで（第2回策定委員会 11月11日の資料に反映）　'
            'Ｂ＝11月20日まで（第3回策定委員会 12月16日の介護保険料の試算に反映）　至急＝第8章の算定の前提になるもの')
ws['A4'].font = F()
for r in (2, 3, 4):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HDR))
    ws.cell(r, 1).alignment = Alignment(vertical='center', wrap_text=True)
ws.row_dimensions[2].height = 30
ws.row_dimensions[3].height = 30
ws.row_dimensions[4].height = 30

HR = 6   # 見出し行
EX = 7   # 記入例
for c, h in enumerate(HDR, 1):
    x = ws.cell(HR, c, h)
    x.font = F(b=True)
    x.fill = G_TOWN if h.startswith('【町記入】') else G_HDR
    x.alignment = CEN
    x.border = BOX
ws.row_dimensions[HR].height = 36
# 記入例は見出しの直後に置く
ex = ['例', '（記入例）', '', '', '', '', '', '',
      '一部提出', '2026/10/28', '令和8年度は9月末時点の見込みで提出。積立額は確定後に追って提出']
for c, v in enumerate(ex, 1):
    x = ws.cell(EX, c, v)
    x.font = F(i=True, color='595959')
    x.fill = G_EX
    x.alignment = CEN if c in (1, 8, 9, 10) else TOP
    x.border = BOX
# 「5行目は記入例」の文言を実際の行番号に合わせる
ws['A3'] = ws['A3'].value.replace('5行目は記入例です。', '%d行目は記入例です。' % EX)

r = EX + 1
for row in ROWS:
    for c, v in enumerate(list(row) + ['', '', ''], 1):
        x = ws.cell(r, c, v)
        x.font = F(b=(c == 2))
        x.alignment = CEN if c in (1, 7, 8, 9, 10) else TOP
        x.border = BOX
        if c >= 9:
            x.fill = G_TOWN
        elif row[0] == 10:
            x.fill = G_ADD
    # 行の高さは中身から見積もる（全角1字≒列幅1.9、1行≒14pt）
    lines = 1
    for c, v in enumerate(row, 1):
        if isinstance(v, str):
            per = max(1, int(WID[c - 1] / 1.9))
            lines = max(lines, sum(-(-len(seg) // per) or 1 for seg in v.split('\n')))
    ws.row_dimensions[r].height = lines * 14 + 6
    r += 1
last = r - 1

dv = DataValidation(type='list', formula1='"未提出,提出済,一部提出,該当資料なし"', allow_blank=True)
dv.error = '一覧から選んでください'
dv.prompt = '未提出／提出済／一部提出／該当資料なし'
ws.add_data_validation(dv)
dv.add('I%d:I%d' % (EX + 1, last))
ws.cell(EX + 1, 10).number_format = 'yyyy/mm/dd'
for rr in range(EX + 1, last + 1):
    ws.cell(rr, 10).number_format = 'yyyy/mm/dd'

r = last + 2
NOTES = [
    '【補足】',
    '・No.1〜9の資料名は元の一覧のとおりです。対象期間・区分・定義は当方で補いました。資料の形式は問いません（既存の帳票・写しで構いません）。',
    '・No.4は、当方の照会番号《要回答14》（要介護1から要介護2へ移った方の移る前の利用サービス）を含むため、本件に統合しました。',
    '・No.10は、44件の「資料提供依頼一覧」の資-01に当たります。認知症高齢者等の人数の系列（県の調査と見える化システムＢ7）を確定するために加えました。',
    '・「当方で確認済みの範囲」は、町の公表資料・国の公表資料・既にご提供いただいた資料で当方が確認したものです。この範囲の再提供は不要です。',
    '・計画素案の《要確定》（町のご判断をいただく事項）17件は、資料のご依頼とは別にお願いします（確認事項照会 第2）。',
]
for t in NOTES:
    ws.cell(r, 1, t).font = F(b=(t == '【補足】'))
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HDR))
    ws.cell(r, 1).alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = 18 if t == '【補足】' else 30
    r += 1

for c, w in enumerate(WID, 1):
    ws.column_dimensions[chr(64 + c)].width = w
ws.freeze_panes = 'C%d' % (EX + 1)
ws.page_setup.orientation = 'landscape'
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = '%d:%d' % (HR, HR)
ws.print_options.horizontalCentered = True
ws.page_margins.left = ws.page_margins.right = 0.4
ws.page_margins.top = ws.page_margins.bottom = 0.5

wb.save(OUT)
print('保存', OUT, '　依頼', len(ROWS), '件　表', '%d〜%d行' % (HR, last))
