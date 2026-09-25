# -*- coding: utf-8 -*-
"""体裁を他メンバー作成レポート（report_honpen_20260828.docx）に合わせたヘルパー。
　書式のみを変更するものであり、API（P／H1／H2／H3／BUL／TBL／SRC）は従前と同一。
　本文の文字・数値・出典には一切手を加えない。

　合わせた項目
　　用紙・余白　　A4縦　上下左右2.0cm　ヘッダ／フッタ1.25cm
　　本文フォント　BIZ UDPゴシック 12pt　段落前後3pt
　　見出し　　　　H1 16pt 太字 #2E74B5／H2 14pt 太字 #2E74B5／H3 11pt 太字 #1F4D78
　　出典行　　　　8pt　段落前2pt・後8pt
　　表　　　　　　見出し行 塗り#1F3864・白文字11pt太字・中央／本文11pt
　　　　　　　　　罫線 上下#2E75B6(sz6)、左右・内側#BFBFBF(sz4)、セル余白左右10dxa
　　ヘッダ　　　　文書名を中央に配置
　　フッタ　　　　ページ番号を「- n -」の形で中央に配置
"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'BIZ UDPゴシック'
GFONT = 'BIZ UDPゴシック'
BODY_PT = 12          # 本文
TBL_PT = 11           # 表内
C_H1 = RGBColor(0x2E, 0x74, 0xB5)
C_H3 = RGBColor(0x1F, 0x4D, 0x78)
DOC_TITLE = '金ケ崎町 高齢者福祉計画・第10期介護保険事業計画・認知症施策推進計画'

doc = Document()
for s in doc.sections:
    s.page_width = Cm(21.0); s.page_height = Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = s.bottom_margin = Cm(2.0)
    s.header_distance = s.footer_distance = Cm(1.25)

st = doc.styles['Normal']
st.font.name = FONT; st.font.size = Pt(BODY_PT)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
st.element.rPr.rFonts.set(qn('w:ascii'), FONT)
st.element.rPr.rFonts.set(qn('w:hAnsi'), FONT)
st.paragraph_format.space_before = Pt(3)
st.paragraph_format.space_after = Pt(3)


def _f(run, f=FONT, sz=BODY_PT, b=False, color=None):
    run.font.name = f; run.font.size = Pt(sz); run.bold = b
    rf = run._element.rPr.rFonts
    for a in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        rf.set(qn(a), f)
    if color:
        run.font.color.rgb = color


def P(text='', sz=BODY_PT, b=False, f=FONT, before=3, after=3, ind=0, align=None, color=None):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if ind:
        pf.left_indent = Cm(ind)
    if align:
        p.alignment = align
    for seg in re.split(r'(\*\*.*?\*\*)', text):
        if not seg:
            continue
        if seg.startswith('**') and seg.endswith('**'):
            _f(p.add_run(seg[2:-2]), f, sz, True, color)
        else:
            _f(p.add_run(seg), f, sz, b, color)
    return p


def H1(t):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18); p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.keep_with_next = True
    _f(p.add_run(t), GFONT, 16, True, C_H1)
    return _register(p, t, 1)


def H2(t):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    _f(p.add_run(t), GFONT, 14, True, C_H1)
    return _register(p, t, 2)


def H3(t):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    _f(p.add_run(t), GFONT, 11, True, C_H3)
    return p


def BUL(t, ind=0.4):
    P('・' + t, ind=ind, before=2, after=2)


def _border(tbl):
    pr = tbl.tblPr
    for old in pr.findall(qn('w:tblBorders')):
        pr.remove(old)
    b = OxmlElement('w:tblBorders')
    for tag, sz, col in (('top', '6', '2E75B6'), ('left', '4', 'BFBFBF'),
                         ('bottom', '6', '2E75B6'), ('right', '4', 'BFBFBF'),
                         ('insideH', '4', 'BFBFBF'), ('insideV', '4', 'BFBFBF')):
        e = OxmlElement('w:' + tag)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), sz)
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), col)
        b.append(e)
    pr.append(b)
    m = OxmlElement('w:tblCellMar')
    for tag in ('left', 'right'):
        e = OxmlElement('w:' + tag)
        e.set(qn('w:w'), '10'); e.set(qn('w:type'), 'dxa')
        m.append(e)
    pr.append(m)


def TBL(rows, widths=None, fs=TBL_PT):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _border(t._tbl)
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row):
            c = t.cell(ri, ci)
            p = c.paragraphs[0]
            for r0 in list(p.runs):          # 空ランを残さない
                r0._element.getparent().remove(r0._element)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for seg in re.split(r'(\*\*.*?\*\*)', str(v)):
                if not seg:
                    continue
                bold = seg.startswith('**') and seg.endswith('**')
                txt = seg[2:-2] if bold else seg
                _f(p.add_run(txt), FONT, fs, bold or ri == 0,
                   RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else None)
            tcPr = c._tc.get_or_add_tcPr()
            if ri == 0:
                sh = OxmlElement('w:shd')
                sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), '1F3864')
                tcPr.append(sh)
            va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center')
            tcPr.append(va)
    if widths:
        for ci, w in enumerate(widths):
            for r in t.rows:
                r.cells[ci].width = Cm(w)
        grid = t._tbl.find(qn('w:tblGrid'))   # 列幅はグリッドにも書く
        if grid is not None:
            for ci, gc in enumerate(grid.findall(qn('w:gridCol'))):
                if ci < len(widths):
                    gc.set(qn('w:w'), str(int(Cm(widths[ci]).twips)))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def SRC(t):
    P(t, sz=8, ind=0.2, before=2, after=8)


FIGNO = [0]                      # 図番号の連番


def FIG(path, caption, width=15.5, src=None):
    """図を1つ置く。キャプションは図の上に【図N】として置き、出典は図の下に置く。

    　第1回策定委員会資料の体裁に合わせる（キャプションを上、出典を8ptで下）。
    """
    FIGNO[0] += 1
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(10); pf.space_after = Pt(2)
    pf.keep_with_next = True          # キャプションと図を離さない
    _f(p.add_run('【図%d】%s' % (FIGNO[0], caption)), sz=11, b=True, color=C_H3)
    q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_before = Pt(0)
    q.paragraph_format.space_after = Pt(2 if src else 8)
    q.paragraph_format.keep_together = True
    q.paragraph_format.keep_with_next = bool(src)   # 図と出典を離さない
    q.add_run().add_picture(path, width=Cm(width))
    if src:
        SRC(src)
    return FIGNO[0]


def _fld(p, instr):
    for tag, extra in (('begin', None), (None, instr), ('end', None)):
        r = OxmlElement('w:r')
        if tag:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), tag)
        else:
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = extra
        r.append(e); p._p.append(r)


def _headfoot(title=DOC_TITLE):
    s = doc.sections[0]
    hp = s.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _f(hp.add_run(title), FONT, 9)
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _f(fp.add_run('- '), FONT, 10)
    _fld(fp, ' PAGE ')
    _f(fp.add_run(' -'), FONT, 10)
    for r in fp.runs:
        _f(r, FONT, 10)


_headfoot()


# ══════════════════════════════════════════════════════════════════
# 　目次（クリックで本文へ飛べるリンクと、PAGEREF によるページ番号）
# ══════════════════════════════════════════════════════════════════
# 　目次の各行は、本文の見出しに置いたブックマークへのリンク（w:hyperlink）と、
# 　ページ番号のフィールド（PAGEREF）で構成する。
# 　　・リンクはフィールドではないため、更新しなくてもクリックで移動できる。
# 　　・ページ番号はフィールドであり、Wordで開いた時点で更新される。
# 　　　更新前に表示する値は out/_toc_pages.json（mk_toc.py が組版結果から作る）
# 　　　から読む。ファイルがない場合は「－」を表示する。
import json as _json
import os as _os
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_SECTION

TEXTW = 17.0                     # 本文の幅（cm）＝21.0−2.0−2.0
HEADINGS = []                    # 目次に載せる見出し （level, text, bookmark名）
_TOC_ANCHOR = [None]


def _fldchar(kind):
    e = OxmlElement('w:fldChar')
    e.set(qn('w:fldCharType'), kind)
    return e


def _bookmark(p, name):
    """段落の先頭と末尾にブックマークを置く。目次のリンク先となる。"""
    bid = str(len(HEADINGS) + 100)
    st_ = OxmlElement('w:bookmarkStart')
    st_.set(qn('w:id'), bid)
    st_.set(qn('w:name'), name)
    en_ = OxmlElement('w:bookmarkEnd')
    en_.set(qn('w:id'), bid)
    pPr = p._p.find(qn('w:pPr'))       # pPr は段落の先頭になければならない
    if pPr is None:
        p._p.insert(0, st_)
    else:
        pPr.addnext(st_)
    p._p.append(en_)


def _register(p, text, level):
    """見出しを目次に登録し、ブックマークを置く。"""
    name = '_Toc10_%03d' % (len(HEADINGS) + 1)
    HEADINGS.append((level, text, name))
    _bookmark(p, name)
    return p


def PGNUM_START(section, start):
    """セクションのページ番号を start から数え直す。"""
    sp = section._sectPr
    e = sp.find(qn('w:pgNumType'))
    if e is None:
        e = OxmlElement('w:pgNumType')
        # CT_SectPr は要素の順序が定められており、pgNumType は cols より前に置く
        after = None
        for tag in ('w:cols', 'w:formProt', 'w:vAlign', 'w:noEndnote',
                    'w:titlePg', 'w:textDirection', 'w:docGrid'):
            after = sp.find(qn(tag))
            if after is not None:
                break
        if after is None:
            sp.append(e)
        else:
            after.addprevious(e)
    e.set(qn('w:start'), str(start))


def _setup_section(s, title=DOC_TITLE, pageno=True):
    s.page_width = Cm(21.0); s.page_height = Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = s.bottom_margin = Cm(2.0)
    s.header_distance = s.footer_distance = Cm(1.25)
    s.header.is_linked_to_previous = False
    s.footer.is_linked_to_previous = False
    hp = s.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(hp.runs):
        r.text = ''
    _f(hp.add_run(title), FONT, 9)
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(fp.runs):
        r.text = ''
    if pageno:
        _f(fp.add_run('- '), FONT, 10)
        _fld(fp, ' PAGE ')
        _f(fp.add_run(' -'), FONT, 10)
        for r in fp.runs:
            _f(r, FONT, 10)
    return s


def TOC_HERE(title='目　次'):
    """目次を置く位置に目印を置く。BUILD_TOC でここに差し込まれる。"""
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    _f(p.add_run(title), GFONT, 16, True, C_H1)
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_after = Pt(10)
    _f(q.add_run('各行をクリックすると本文へ移動します。'
                 'ページ番号はWordで開いた時点で自動的に更新されます'), FONT, 9)
    _TOC_ANCHOR[0] = doc.add_paragraph()   # 差込みの目印
    return _TOC_ANCHOR[0]


def BODY_HERE():
    """本文の開始位置でセクションを切り、ページ番号を1から数え直す。"""
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    _setup_section(s)
    PGNUM_START(s, 1)
    # 表紙・目次のセクションにはページ番号を出さない
    _setup_section(doc.sections[0], pageno=False)
    return s


def _toc_run(parent, txt=None, child=None, bold=False, size=10.5):
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    f = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        f.set(qn(a), FONT)
    rPr.append(f)
    if bold:
        rPr.append(OxmlElement('w:b'))
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(size * 2)))
    rPr.append(sz)
    r.append(rPr)
    if txt is not None:
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = txt
        r.append(t)
    if child is not None:
        r.append(child)
    parent.append(r)
    return r


def _toc_line(level, text, name, pages):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(1 if level == 1 else 0)
    pf.space_after = Pt(1 if level == 1 else 0)
    pf.line_spacing = 1.05
    pf.left_indent = Cm(0 if level == 1 else 0.9)
    pf.tab_stops.add_tab_stop(Cm(TEXTW), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    hl = OxmlElement('w:hyperlink')
    hl.set(qn('w:anchor'), name)
    p._p.append(hl)
    _toc_run(hl, txt=text, bold=(level == 1), size=11 if level == 1 else 10)
    _toc_run(hl, child=OxmlElement('w:tab'), size=10)
    _toc_run(hl, child=_fldchar('begin'), size=10)
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = ' PAGEREF %s \\h ' % name
    _toc_run(hl, child=it, size=10)
    _toc_run(hl, child=_fldchar('separate'), size=10)
    _toc_run(hl, txt=str(pages.get(name, '－')), bold=(level == 1), size=10)
    _toc_run(hl, child=_fldchar('end'), size=10)
    return p


def _update_fields():
    """Wordで開いた時点でページ番号のフィールドを更新させる。"""
    stg = doc.settings.element
    uf = stg.find(qn('w:updateFields'))
    if uf is None:
        uf = OxmlElement('w:updateFields')
        after = None
        for tag in ('w:compat', 'w:docVars', 'w:rsids', 'w:themeFontLang'):
            after = stg.find(qn(tag))
            if after is not None:
                break
        if after is None:
            stg.append(uf)
        else:
            after.addprevious(uf)
    uf.set(qn('w:val'), 'true')
    z = stg.find(qn('w:zoom'))          # 既定テンプレートの w:zoom は必須属性を欠く
    if z is not None and z.get(qn('w:percent')) is None:
        z.set(qn('w:percent'), '100')


def BUILD_TOC(pagemap='out/_toc_pages.json'):
    """目印の位置に目次を差し込む。"""
    if _TOC_ANCHOR[0] is None:
        return
    pages = {}
    if _os.path.exists(pagemap):
        with open(pagemap, encoding='utf-8') as f:
            pages = _json.load(f)
    paras = [_toc_line(lv, tx, nm, pages) for lv, tx, nm in HEADINGS]
    anchor = _TOC_ANCHOR[0]
    for q in paras:
        anchor._p.addprevious(q._p)
    anchor._p.getparent().remove(anchor._p)
    _update_fields()
