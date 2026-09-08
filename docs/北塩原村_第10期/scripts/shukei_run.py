# -*- coding: utf-8 -*-
"""集計プログラム本体。
   コードブックと回答データを読み、派生変数を作り、単純集計とクロス集計を
   xlsx に出力する。本番データが来る前にダミーデータで通しの検証を行う。

   使い方:
     python3 shukei_run.py                       # ダミーデータで実行
     python3 shukei_run.py --needs <csv> --zaitaku <csv> --out <xlsx>
"""
import os, sys, json, csv, argparse, collections

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import shukei_data as SD
from shukei_derive import derive_needs, derive_zaitaku, JUDGE_PROVISIONAL

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

NAVY, LGREY, YELLOW = "1F3864", "F2F2F2", "FFF2CC"
F_H = Font(name="Meiryo UI", size=9, bold=True, color="FFFFFF")
F_B = Font(name="Meiryo UI", size=9)
F_T = Font(name="Meiryo UI", size=10, bold=True, color=NAVY)
FILL_H = PatternFill("solid", fgColor=NAVY)
FILL_G = PatternFill("solid", fgColor=LGREY)
FILL_Y = PatternFill("solid", fgColor=YELLOW)
BD = Border(*[Side(style="thin", color="BFBFBF")] * 4)

# 集計軸の定義（コード → (表示名, 値を取り出す関数, 表示順)）
AXES = {
    "A": ("全体", lambda r: "全体", ["全体"]),
    "B": ("地区", lambda r: r.get("地区"), ["北山", "大塩", "桧原", "裏磐梯"]),
    "C": ("年齢階級", lambda r: r.get("年齢階級"), ["65〜74歳", "75〜84歳", "85歳以上"]),
    "D": ("性別", lambda r: r.get("性別"), ["男性", "女性"]),
    "E": ("認定状況", lambda r: r.get("認定状況"),
          ["認定なし", "要支援1・2", "要介護1・2", "要介護3〜5"]),
    "F": ("世帯類型", lambda r: SETAI.get(str(r.get("N_問1_1", "")), None),
          ["単身", "夫婦のみ", "その他"]),
    "G": ("前回比較", lambda r: None, []),          # 第9期の集計表の受領後に実装
    "H": ("生活機能4層", lambda r: r.get("L01") or None,
          ["第1層　元気・活躍層", "第2層　元気・未参加層",
           "第3層　リスク層", "第4層　要支援・要介護層"]),
    "I": ("要介護度", lambda r: r.get("要介護度"), ["要支援1・2", "要介護1・2", "要介護3〜5"]),
    "J": ("介護者の勤務形態", lambda r: r.get("介護者の勤務形態"),
          ["フルタイム", "パートタイム", "働いていない"]),
}
SETAI = {"1": "単身", "2": "夫婦のみ", "3": "夫婦のみ", "4": "その他", "5": "その他"}


def read_csv(path):
    return list(csv.DictReader(open(path, encoding="utf-8-sig")))


# 集計仕様書 基本方針7：無回答は集計対象に含め、分母から除外しない。
#   したがって％の分母は当該区分の回答者数（n）である。


def tally_single(rows, col, opts):
    """単一回答の度数。戻り値は (ラベル, 件数) のリストと分母（n）"""
    cnt = collections.Counter()
    for r in rows:
        cnt[str(r.get(col, "")).strip()] += 1
    out = [(f"{n}．{lab}", cnt.get(str(n), 0)) for n, lab in opts]
    out.append(("無回答", len(rows) - sum(c for _l, c in out)))
    return out, len(rows)


def tally_multi(rows, var, opts):
    """複数回答の度数。分母は回答者数（基本方針6）。合計は100％を超える。"""
    out = []
    for n, lab in opts:
        c = sum(1 for r in rows if str(r.get(f"{var}_c{n}", "")).strip() == "1")
        out.append((f"{n}．{lab}", c))
    sel = sum(1 for r in rows
              if any(str(r.get(f"{var}_c{n}", "")).strip() == "1" for n, _ in opts))
    out.append(("無回答", len(rows) - sel))
    return out, len(rows)


def tally_numeric(rows, col):
    vals = []
    for r in rows:
        v = str(r.get(col, "")).strip()
        if v:
            try:
                vals.append(float(v))
            except ValueError:
                pass
    if not vals:
        return [], 0
    vals.sort()
    n = len(vals)
    out = [("有効回答数", n), ("平均", round(sum(vals) / n, 1)),
           ("最小", round(vals[0], 1)),
           ("中央値", round(vals[n // 2], 1)),
           ("最大", round(vals[-1], 1)),
           ("無回答", len(rows) - n)]
    return out, n


def sheet_header(ws, title, note=""):
    ws.append([title])
    ws.cell(ws.max_row, 1).font = F_T
    if note:
        ws.append([note])
        ws.cell(ws.max_row, 1).font = Font(name="Meiryo UI", size=8, color="808080")
    ws.append([])


def write_table(ws, head, rows, widths=None):
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(ws.max_row, c)
        cell.font, cell.fill, cell.border = F_H, FILL_H, BD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
        for c in range(1, len(head) + 1):
            cell = ws.cell(ws.max_row, c)
            cell.font, cell.border = F_B, BD
            if c > 1:
                cell.alignment = Alignment(horizontal="right")
    if widths:
        from openpyxl.utils import get_column_letter
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.append([])


def pct(n, d):
    return round(n / d * 100, 1) if d else 0.0


def build_tanjun(wb, book, data, hyos, sheet_name):
    ws = wb.create_sheet(sheet_name)
    sheet_header(ws, f"単純集計　{sheet_name}",
                 f"n＝{len(data)}　％は回答者数に対する割合（無回答を分母に含む・基本方針7）。"
                 "複数回答は合計が100％を超える（基本方針6）")
    for q in book:
        if q["票"] not in hyos or not q["列"]:
            continue
        fmt = q["形式"]
        if fmt == "自由":
            continue
        ws.append([f"{q['番号']}　{q['設問文'][:60]}"])
        ws.cell(ws.max_row, 1).font = F_T
        ws.append([f"区分：{q['区分']}／形式：{fmt}／クロス軸：{q['クロス軸']}"
                   + ("／見える化：登録対象" if q["見える化"] == "○" else "")])
        ws.cell(ws.max_row, 1).font = Font(name="Meiryo UI", size=8, color="808080")
        if fmt == "数値":
            for col in q["列"]:
                out, _v = tally_numeric(data, col)
                write_table(ws, [col, "値"], [[a, b] for a, b in out], [34, 12])
        elif fmt.startswith("複数"):
            out, valid = tally_multi(data, q["変数"], q["選択肢"])
            write_table(ws, ["選択肢", "件数", "％"],
                        [[a, b, pct(b, valid)] for a, b in out], [46, 10, 10])
        else:
            out, valid = tally_single(data, q["列"][0], q["選択肢"])
            write_table(ws, ["選択肢", "件数", "％"],
                        [[a, b, pct(b, valid)] for a, b in out], [46, 10, 10])
    return ws


def build_axis(wb, book, data, hyos, sheet_name, axis_codes):
    """設問×集計軸のクロス（％表）"""
    ws = wb.create_sheet(sheet_name)
    sheet_header(ws, f"軸別集計　{sheet_name}",
                 "各セルは当該区分の回答者数を分母とした％。複数回答は合計が100％を超える")
    for code in axis_codes:
        name, getter, order = AXES[code]
        if not order:
            continue
        groups = {k: [r for r in data if getter(r) == k] for k in order}
        ws.append([f"■ 軸{code}　{name}　" + "／".join(f"{k} n={len(v)}" for k, v in groups.items())])
        ws.cell(ws.max_row, 1).font = F_T
        for q in book:
            if q["票"] not in hyos or not q["列"] or q["形式"] in ("自由", "数値", "マトリクス"):
                continue
            if code not in q["クロス軸"].split(","):
                continue
            head = ["選択肢", "全体"] + order
            rows = []
            base = (tally_multi(data, q["変数"], q["選択肢"]) if q["形式"].startswith("複数")
                    else tally_single(data, q["列"][0], q["選択肢"]))
            per = {}
            for k, g in groups.items():
                per[k] = (tally_multi(g, q["変数"], q["選択肢"]) if q["形式"].startswith("複数")
                          else tally_single(g, q["列"][0], q["選択肢"]))
            for i, (lab, c) in enumerate(base[0]):
                row = [lab, pct(c, base[1])]
                for k in order:
                    o, v = per[k]
                    row.append(pct(o[i][1], v))
                rows.append(row)
            ws.append([f"{q['番号']}　{q['設問文'][:50]}"])
            ws.cell(ws.max_row, 1).font = Font(name="Meiryo UI", size=9, bold=True)
            write_table(ws, head, rows, [42, 9] + [11] * len(order))
    return ws


def crosstab(rows, get_r, ord_r, get_c, ord_c):
    tbl = {a: {b: 0 for b in ord_c} for a in ord_r}
    for r in rows:
        a, b = get_r(r), get_c(r)
        if a in tbl and b in tbl[a]:
            tbl[a][b] += 1
    return tbl


def build_cross(wb, needs, zai):
    """分析編のクロス表（doc25 CROSS）のうち、軸が確定しているものを出力する"""
    ws = wb.create_sheet("分析クロス")
    sheet_header(ws, "分析編クロス表", "件数と行％。網掛けは分析編の中核となる表")

    D01 = (lambda r: r.get("D01_3") or None, ["0疾病", "1〜2疾病", "3疾病以上"])
    R99 = (lambda r: r.get("R99_3") or None, ["0領域", "1〜2領域", "3領域以上"])
    CERT = (lambda r: r.get("認定状況"), AXES["E"][2])
    L01 = (lambda r: r.get("L01") or None, AXES["H"][2])
    CHIKU = (lambda r: r.get("地区"), AXES["B"][2])
    AGE = (lambda r: r.get("年齢階級"), AXES["C"][2])
    P01 = (lambda r: {"1": "参加あり", "0": "参加なし"}.get(str(r.get("P01", ""))), ["参加あり", "参加なし"])
    S01 = (lambda r: {"1": "受領あり", "0": "受領なし"}.get(str(r.get("S01", ""))), ["受領あり", "受領なし"])
    S02 = (lambda r: {"1": "提供あり", "0": "提供なし"}.get(str(r.get("S02", ""))), ["提供あり", "提供なし"])
    YOKAI = (lambda r: r.get("要介護度"), AXES["I"][2])
    KINMU = (lambda r: r.get("介護者の勤務形態"), AXES["J"][2])
    Q6 = (lambda r: single_lab(r, "ZA_問6"), None)

    PLAN = [
        ("X-01", "疾病数 × リスク該当領域数", needs, D01, R99, True),
        ("X-02", "認定状況 × リスク該当領域数", needs, CERT, R99, True),
        ("X-04", "支え合いの受領 × 提供", needs, S01, S02, True),
        ("X-05", "生活機能4層 × 地区", needs, L01, CHIKU, False),
        ("X-06", "生活機能4層 × 年齢階級", needs, L01, AGE, False),
        ("X-12", "疾病数 × 社会参加", needs, D01, P01, False),
        ("X-16", "要介護度 × 就労継続の見通し", zai, YOKAI,
         (lambda r: single_lab(r, "ZB_問10"), None), True),
        ("X-17", "介護者の勤務形態 × 不安に感じる介護", zai, KINMU, None, False),
        ("X-15", "施設入所の検討状況 × 村内施設があれば", zai, Q6,
         (lambda r: single_lab(r, "ZA_問6_1"), None), True),
    ]
    for no, title, rows, rspec, cspec, key in PLAN:
        if cspec is None:
            continue
        get_r, ord_r = rspec
        get_c, ord_c = cspec
        if ord_r is None:
            ord_r = sorted({get_r(r) for r in rows if get_r(r)})
        if ord_c is None:
            ord_c = sorted({get_c(r) for r in rows if get_c(r)})
        tbl = crosstab(rows, get_r, ord_r, get_c, ord_c)
        ws.append([f"{no}　{title}" + ("　★分析編の中核" if key else "")])
        ws.cell(ws.max_row, 1).font = F_T
        if key:
            ws.cell(ws.max_row, 1).fill = FILL_Y
        head = [""] + list(ord_c) + ["計"]
        body = []
        for a in ord_r:
            tot = sum(tbl[a].values())
            body.append([a] + [tbl[a][b] for b in ord_c] + [tot])
            body.append([f"　{a}（％）"] + [pct(tbl[a][b], tot) for b in ord_c] + [100.0 if tot else 0.0])
        write_table(ws, head, body, [30] + [13] * (len(ord_c) + 1))
    return ws


LAB_CACHE = {}


def single_lab(r, col):
    v = str(r.get(col, "")).strip()
    opts = LAB_CACHE.get(col)
    if not opts or not v:
        return None
    for n, lab in opts:
        if str(n) == v:
            return f"{n}．{lab[:22]}"
    return None


def build_index(wb, book, n_needs, n_zai):
    ws = wb.create_sheet("目次", 0)
    sheet_header(ws, "第10期北塩原村 アンケート調査　集計結果",
                 "集計仕様書（doc25）に基づく。設問の定義はコードブック（data/集計_コードブック.json）による。")
    write_table(ws, ["シート", "内容"],
                [["ニーズ単純集計", "介護予防・日常生活圏域ニーズ調査の単純集計"],
                 ["ニーズ軸別集計", "同上を地区・年齢階級・性別・認定状況・生活機能4層で分けたもの"],
                 ["在宅単純集計", "在宅介護実態調査（A票・B票）の単純集計"],
                 ["在宅軸別集計", "同上を地区・要介護度・介護者の勤務形態で分けたもの"],
                 ["分析クロス", "分析編に用いるクロス表"],
                 ["派生変数", "リスク判定・生活機能4層などの算出結果と判定基準"],
                 ["設問一覧", "コードブック（設問・形式・選択肢数・データ列）"]],
                [22, 78])
    write_table(ws, ["調査", "回答数"],
                [["介護予防・日常生活圏域ニーズ調査", n_needs],
                 ["在宅介護実態調査", n_zai]], [40, 14])
    return ws


def build_derived_sheet(wb, needs):
    ws = wb.create_sheet("派生変数")
    sheet_header(ws, "派生変数", "【暫定】はリスク判定の基準が国の手引きに明記されておらず、村との確認を要するもの")
    BINARY = {f"R{i:02d}" for i in range(1, 11)} | {"P01", "S01", "S02", "S99"}
    rows = []
    for code, name, src, rule, kind in SD.DERIVED:
        if code in BINARY:
            n = sum(1 for r in needs if str(r.get(code, "")) == "1")
        else:
            n = ""            # 該当／非該当ではない変数（件数の欄は使わない）
        mark = "【暫定】" if code in JUDGE_PROVISIONAL else ""
        rows.append([code, name, src, mark + rule.replace("【要確認】", ""), kind,
                     n if n != "" else "―",
                     pct(n, len(needs)) if isinstance(n, int) else "―"])
    write_table(ws, ["コード", "名称", "元の設問", "判定基準", "区分", "該当数", "％"],
                rows, [8, 22, 18, 54, 10, 10, 8])

    # リスク該当領域数の分布
    cnt = collections.Counter(str(r.get("R99", "")) for r in needs)
    write_table(ws, ["リスク該当領域数", "件数", "％"],
                [[k if k else "無回答", cnt[k], pct(cnt[k], len(needs))]
                 for k in sorted(cnt, key=lambda x: (x == "", x))], [22, 12, 10])
    # 生活機能4層
    cnt = collections.Counter(r.get("L01") or "判定不能" for r in needs)
    write_table(ws, ["生活機能4層", "件数", "％"],
                [[k, cnt[k], pct(cnt[k], len(needs))] for k in AXES["H"][2] + ["判定不能"]
                 if k in cnt], [30, 12, 10])
    return ws


def build_book_sheet(wb, book):
    ws = wb.create_sheet("設問一覧")
    sheet_header(ws, "設問一覧（コードブック）", "調査票のdocxから抽出し、集計仕様書と突き合わせたもの")
    write_table(ws, ["票", "番号", "設問文", "形式", "区分", "選択肢数", "データ列", "クロス軸"],
                [[q["票"], q["番号"], q["設問文"][:70], q["形式"], q["区分"],
                  len(q["選択肢"]), len(q["列"]), q["クロス軸"]] for q in book],
                [8, 12, 62, 10, 10, 9, 9, 14])
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--needs", default=os.path.join(BASE, "data", "dummy_ニーズ.csv"))
    ap.add_argument("--zaitaku", default=os.path.join(BASE, "data", "dummy_在宅.csv"))
    ap.add_argument("--out", default="/home/user/repository/output/10_集計結果.xlsx")
    a = ap.parse_args()

    book = json.load(open(os.path.join(BASE, "data", "集計_コードブック.json"), encoding="utf-8"))
    for q in book:
        if q["列"]:
            LAB_CACHE[q["列"][0]] = q["選択肢"]

    needs = [derive_needs(r) for r in read_csv(a.needs)]
    zai = [derive_zaitaku(r) for r in read_csv(a.zaitaku)]

    wb = Workbook()
    wb.remove(wb.active)
    build_index(wb, book, len(needs), len(zai))
    build_tanjun(wb, book, needs, ("ニーズ",), "ニーズ単純集計")
    build_axis(wb, book, needs, ("ニーズ",), "ニーズ軸別集計", ["B", "C", "D", "E", "H"])
    build_tanjun(wb, book, zai, ("在宅A", "在宅B"), "在宅単純集計")
    build_axis(wb, book, zai, ("在宅A", "在宅B"), "在宅軸別集計", ["B", "I", "J"])
    build_cross(wb, needs, zai)
    build_derived_sheet(wb, needs)
    build_book_sheet(wb, book)
    for ws in wb:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A4"
    wb.save(a.out)
    print(f"保存: {a.out}")
    print(f"  ニーズ {len(needs)}件 / 在宅 {len(zai)}件 / シート {len(wb.sheetnames)}")
    print("  " + " / ".join(wb.sheetnames))


if __name__ == "__main__":
    main()
