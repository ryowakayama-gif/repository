# -*- coding: utf-8 -*-
"""令和8年度 交付金 評価指標（都道府県分）集計表から福島県の得点を取り出し、
   市町村分による本村の目標別得点と政策領域ごとに対照する。

   受領（2026/09/25）
     001732598.xlsx  令和8年度 該当状況調査票集計表（都道府県分）47都道府県
     001732599.pdf   同 評価指標（都道府県分）全38頁

   受領物は都道府県分であり、本村および福島県内59市町村の得点は載っていない。
   本村の評価は既収録の市町村分（2026/09/02受領）による。
"""
import csv, os, sys

SRC = "/root/.claude/uploads/134138ca-61f7-57d3-9e9b-5f081a1a345d/06a90cb2-001732598.xlsx"
OUT = "/home/user/repository/docs/北塩原村_第10期/data/交付金_都道府県分_令和8年度.csv"

# 都道府県分 集計表の合計列（0起点）。ヘッダーを縦に連結して同定した
COL = {"推進Ⅰ": 39, "推進Ⅱ": 59, "推進Ⅲ": 124, "推進Ⅳ": 157, "推進合計": 158,
       "支援Ⅰ": 247, "支援Ⅱ": 282, "支援Ⅲ": 306, "支援Ⅳ": 339, "支援合計": 340,
       "総合計": 341, "今年度順位": 342, "前年度得点": 5, "前年度順位": 6}
# 集計表のヘッダーに埋め込まれた統計（配点／全国合計／平均／平均率／中位／標準偏差）
STAT_ROWS = 18

# 市町村分（2026/09/02受領）による本村の令和8年度 目標別得点（doc20 §4-1・§4-2）
MURA = {"推進Ⅰ": (76, 62.3), "推進Ⅱ": (62, 69.2), "推進Ⅲ": (46, 50.8), "推進Ⅳ": (45, 47.8),
        "支援Ⅰ": (71, 57.8), "支援Ⅱ": (32, 51.1), "支援Ⅲ": (70, 68.3), "支援Ⅳ": (45, 47.8)}
NAME = {"Ⅰ": ("持続可能な地域のあるべき姿をかたちにする", "介護予防／日常生活支援を推進する"),
        "Ⅱ": ("公正・公平な給付を行う体制を構築する", "認知症総合支援を推進する"),
        "Ⅲ": ("介護人材の確保その他のサービス提供基盤の整備を推進する", "在宅医療・在宅介護連携の体制を構築する"),
        "Ⅳ": ("高齢者がその状況に応じて可能な限り自立した日常生活を営む",) * 2}

def load():
    from openpyxl import load_workbook
    wb = load_workbook(SRC, read_only=True, data_only=True)
    rows = list(wb["全国集計（都道府県）"].iter_rows(values_only=True))
    H = rows[:STAT_ROWS]
    FUK = next(r for r in rows[STAT_ROWS:] if r[2] == "福島県")
    def stat(c):
        # 見出し列を縦に読み、数値だけを順に拾う（配点・全国合計・平均・平均率・中位・標準偏差）
        v = [float(r[c]) for r in H
             if c < len(r) and isinstance(r[c], (int, float))]
        return v
    return FUK, stat

if __name__ == "__main__":
    FUK, stat = load()
    recs = []
    print("■ 福島県の交付金評価（都道府県分・令和8年度）\n")
    print(f"{'区分':8s} {'配点':>5s} {'福島県':>7s} {'全国平均':>8s} {'差':>7s} {'全国中位':>8s}")
    print("─" * 52)
    for k in ("推進Ⅰ", "推進Ⅱ", "推進Ⅲ", "推進Ⅳ", "推進合計",
              "支援Ⅰ", "支援Ⅱ", "支援Ⅲ", "支援Ⅳ", "支援合計", "総合計"):
        c = COL[k]
        s = stat(c)
        hai, _tot, avg, _rate, med = s[0], s[1], s[2], s[3], s[4]
        f = FUK[c]
        mark = "" if k.endswith("合計") else ""
        print(f"{k:8s} {hai:5.0f} {f:7.0f} {avg:8.1f} {f-avg:+7.1f} {med:8.0f}{mark}")
        recs.append({"区分": k, "配点": hai, "福島県": f, "全国平均": round(avg, 1),
                     "差": round(f - avg, 1), "全国中央値": med})
    print("─" * 52)
    print(f"今年度順位 {FUK[COL['今年度順位']]}位／47　"
          f"前年度（令和7年度）{FUK[COL['前年度得点']]}点・{FUK[COL['前年度順位']]}位")

    print("\n\n■ 政策領域ごとの対照（本村＝市町村分／福島県＝都道府県分）\n")
    print("  同じ目標番号でも、市町村分は市町村自身の取組、都道府県分は県による市町村支援を評価する。")
    print("  厳密な比較ではなく、同じ政策領域が県と市町村のどちらで弱いかを見るための対照である。\n")
    print(f"{'交付金':5s} {'目標':4s} {'政策領域':28s} "
          f"{'本村':>5s} {'全国':>6s} {'差':>6s} │ {'福島県':>6s} {'全国':>6s} {'差':>6s}  判定")
    print("─" * 112)
    hantei = []
    for kf, idx in (("推進", 0), ("支援", 1)):
        for mk in "ⅠⅡⅢⅣ":
            k = f"{kf}{mk}"
            mv, ma = MURA[k]
            c = COL[k]
            s = stat(c)
            fv, fa = FUK[c], s[2]
            md, fd = mv - ma, fv - fa
            # 本村の不足幅と県の不足幅を比べ、県全体の傾向か本村固有かを切り分ける
            if md >= 0:
                j = "本村は全国平均以上"
            elif fd <= md:
                j = "県全体の傾向（県の不足が本村と同等以上）"
            elif fd <= md / 2:
                j = "県も弱いが本村がより弱い"
            elif abs(md) < 5:
                j = "本村固有（ただし差は小さい）"
            else:
                j = "★本村固有の弱点"
            hantei.append((k, NAME[mk][idx], md, fd, j))
            print(f"{kf:5s} {mk:4s} {NAME[mk][idx][:26]:28s} "
                  f"{mv:5.0f} {ma:6.1f} {md:+6.1f} │ {fv:6.0f} {fa:6.1f} {fd:+6.1f}  {j}")
    print("─" * 112)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8-sig", newline="") as fp:
        w = csv.DictWriter(fp, fieldnames=["区分", "配点", "福島県", "全国平均", "差", "全国中央値"])
        w.writeheader(); w.writerows(recs)
    print(f"\n保存: {OUT}")
