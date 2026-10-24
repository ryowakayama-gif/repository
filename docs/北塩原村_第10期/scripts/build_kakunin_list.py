# -*- coding: utf-8 -*-
"""確認事項一覧（単独のExcel）を書き出す。

【なぜ別に書き出すか】
  確認事項は WBS進捗管理表（成果品05）の1シートとして持っているが、
  村や他の点検者（他AIを含む）に渡すときは、確認事項だけを1冊で見たい。
  wbs_kakunin（題・内容・優先度・状態）、wbs_pending（影響度）、
  wbs_kubun（尋ね方の区分）を突き合わせて1表にする。数は書き写さない。

  python3 scripts/build_kakunin_list.py
"""
import os
import sys

sys.dont_write_bytecode = True
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import paths as P
import wbs_kakunin as K
import wbs_kubun as B
import wbs_pending as PD

OUT = P.out("20_北塩原村第10期_確認事項一覧.xlsx")
F = "BIZ UDPゴシック"
THIN = Side(style="thin", color="BFBFBF")
BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="595959")


def _eikyo(dai):
    v = PD.IMPACT.get(dai)
    if not v:
        return "", ""
    lv = v[0]
    return lv, PD.LEVEL.get(lv, ("", ""))[0]


def _kubun(dai):
    v = B.KUBUN.get(dai)
    return (v[0], v[1]) if v else ("Ｄ", B.KIJUN["Ｄ"])


def sheet(wb, title, rows, kaiketsu=False):
    ws = wb.create_sheet(title)
    heads = ["No.", "分類", "確認事項", "内容・経緯", "関連WBS", "出典", "優先度", "状態"]
    widths = [6, 10, 40, 90, 14, 16, 7, 9]
    if not kaiketsu:
        heads += ["影響度", "影響度の意味", "尋ね方の区分", "当方が先行して行うこと"]
        widths += [7, 10, 9, 40]
    for i, (h, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=F, size=9, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BD
        ws.column_dimensions[c.column_letter].width = w
    for n, r in enumerate(rows, 1):
        bunrui, dai, naiyo, wbs, src, yu, jo = r[:7]
        vals = [n, bunrui, dai, naiyo, wbs, src, yu, jo]
        if not kaiketsu:
            lv, lvn = _eikyo(dai)
            kb, saki = _kubun(dai)
            vals += [lv, lvn, kb, saki]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=n + 1, column=i, value=v)
            c.font = Font(name=F, size=9)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BD
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "見方"
    mitai = [
        ("確認事項一覧", ""),
        ("基準日", PD.BASE_DATE if hasattr(PD, "BASE_DATE") else ""),
        ("未解決", "%d件" % len(K.K)),
        ("解決済", "%d件" % len(K.SOLVED)),
        ("優先度", "S＝9月3日の調査票第1稿提出・9月9日の校了・9月18日の発送を守るために決着が必要（キックオフ議事録で未決着のもの）／"
                 "A＝令和8年9月中に確定が必要／B＝第1回策定委員会（11月）まで／C＝第2回以降（成果品05と同じ定義。起票時の区分のまま）"),
        ("影響度", "／".join("%d＝%s（%s）" % (k, v[0], v[1]) for k, v in PD.LEVEL.items())),
        ("尋ね方の区分", "Ｂ＝当方で進める（村には尋ねない）／Ｃ＝他の計画・公表資料で先行し不足分のみ照会／Ｄ＝村のご回答・内部データが必要"),
        ("出どころ", "scripts/wbs_kakunin.py・wbs_pending.py・wbs_kubun.py（成果品05 WBS進捗管理表の確認事項と同じもの）"),
    ]
    for i, (a, b) in enumerate(mitai, 1):
        ws0.cell(row=i, column=1, value=a).font = Font(name=F, size=10, bold=True)
        ws0.cell(row=i, column=2, value=b).font = Font(name=F, size=10)
    ws0.column_dimensions["A"].width = 16
    ws0.column_dimensions["B"].width = 110
    sheet(wb, "未解決", K.K)
    sheet(wb, "解決済", K.SOLVED, kaiketsu=True)
    wb.save(OUT)
    print("保存:", OUT, "未解決%d件・解決済%d件" % (len(K.K), len(K.SOLVED)))


if __name__ == "__main__":
    main()
