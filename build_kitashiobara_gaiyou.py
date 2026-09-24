# -*- coding: utf-8 -*-
"""
北塩原村 第8期障がい福祉計画・第4期障がい児福祉計画
素案 概要版ジェネレータ

出力: output/北塩原村_計画素案_概要版.docx（A4縦・4ページ）

現行計画の概要版（source/現行計画/第4次障がい者計画等_概要版.pdf、A4・4ページ）
の構成を踏襲する。ただし次の点が異なる。

  - 今回の改定対象は障がい福祉計画（第8期）と障がい児福祉計画（第4期）の
    2計画である。障がい者計画（第4次）は令和11年度までを計画期間としており
    改定対象外のため、その基本理念・基本目標・施策体系は「継承する内容」
    として整理する。
  - 現行概要版の「4 SDGｓとの関連について」は、今回は扱わない方針のため
    「国の基本指針の見直しのポイント」に差し替える。
  - 成果目標は第8期で8項目に再編されたため、項目ごとの囲みではなく
    一覧表にまとめる。

数値は output/北塩原村_計画素案.docx 及び kitashiobara_common.py と同じ
出典から取り、村に確認中のものには【要更新】を付す。
"""

import os

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from kitashiobara_common import SERVICES_ADULT, SERVICES_CHILD

OUT_DIR = "/home/user/repository/output"
OUT_FILE = f"{OUT_DIR}/北塩原村_計画素案_概要版.docx"

FONT = "BIZ UDPゴシック"
DARK = "1F3864"        # 章見出しの帯（計画素案の表頭と同色）
ACCENT = "FF7F0E"      # 障がい分野の識別色（build_excel.py の規約）
BAND = "DDEBF7"        # 小見出しの帯
NOTE_FILL = "FFF2CC"   # 【要更新】ボックス
BOX_FILL = "F2F2F2"
BORDER_COLOR = "AAAAAA"

PAGE_W = 10206         # A4（210mm）－左右余白15mm ＝ 180mm（twips）

# 1ページに収まる行数の目安（A4縦・上下余白15mm・9pt・行送り固定12pt）
LINES_PER_PAGE = 57
CHARS_PER_LINE = 54    # 9pt 全角で1行に入るおよその文字数


# ============================================================
# 書式ヘルパー
# ============================================================
def _el(tag, **attrs):
    e = OxmlElement(tag if ":" in tag else f"w:{tag}")
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def _cell_shade(cell, fill):
    cell._tc.get_or_add_tcPr().append(_el("shd", val="clear", color="auto", fill=fill))


def _cell_margins(cell, top=30, left=80, bottom=30, right=80):
    mar = _el("tcMar")
    for tag, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        mar.append(_el(tag, w=v, type="dxa"))
    cell._tc.get_or_add_tcPr().append(mar)


def _run(para, text, size=9, bold=False, color=None):
    run = para.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = _el("rFonts")
        rpr.insert(0, rf)
    rf.set(qn("w:eastAsia"), FONT)
    return run


def para(doc, text="", size=9, bold=False, align=None, space_after=2,
         indent=0, color=None, line=12):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = Pt(line)
    if indent:
        pf.left_indent = Pt(indent)
        pf.first_line_indent = Pt(-indent)
    if align is not None:
        p.alignment = align
    if text:
        _run(p, text, size=size, bold=bold, color=color)
    return p


def band(doc, number, title):
    """章見出しの帯（濃紺地に白抜き）。"""
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    cell = t.cell(0, 0)
    cell.width = Mm(180)
    _cell_shade(cell, DARK)
    _cell_margins(cell, top=50, bottom=50, left=110)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(14)
    _run(p, f"{number}　{title}" if number else title,
         size=11.5, bold=True, color="FFFFFF")
    _tbl_width(t)
    _no_borders(t)
    para(doc, "", size=4, space_after=0, line=5)
    return t


def sub(doc, text, fill=BAND):
    """小見出しの帯（淡色地）。"""
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    cell = t.cell(0, 0)
    cell.width = Mm(180)
    _cell_shade(cell, fill)
    _cell_margins(cell, top=20, bottom=20, left=100)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(12)
    _run(p, text, size=9.5, bold=True, color="1F3864")
    _tbl_width(t)
    _no_borders(t)
    para(doc, "", size=3, space_after=0, line=4)
    return t


def bullets(doc, lines, size=9):
    for line in lines:
        para(doc, f"・{line}", size=size, indent=9)


def _tbl_width(t):
    t._tbl.tblPr.append(_el("tblW", w=PAGE_W, type="dxa"))
    t._tbl.tblPr.append(_el("tblLayout", type="fixed"))


def _no_borders(t):
    borders = _el("tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_el(edge, val="nil"))
    t._tbl.tblPr.append(borders)


def _borders(t, color=BORDER_COLOR):
    borders = _el("tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_el(edge, val="single", sz=4, space=0, color=color))
    t._tbl.tblPr.append(borders)


def table(doc, rows, widths, header=True, size=8.5, aligns=None, fill_first=None):
    """本文の表。widths は twips で合計 PAGE_W とする。"""
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.autofit = False
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            cell = t.cell(r, c)
            cell.width = Mm(widths[c] * 180 / PAGE_W)
            _cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(11)
            if aligns and aligns[c] == "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            is_head = header and r == 0
            if is_head:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(p, str(value), size=size, bold=is_head,
                 color="FFFFFF" if is_head else None)
            if is_head:
                _cell_shade(cell, DARK)
            elif fill_first and c == 0:
                _cell_shade(cell, fill_first)
    _tbl_width(t)
    _borders(t)
    # 列幅を明示する
    grid = t._tbl.find(qn("w:tblGrid"))
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(widths[i]))
    para(doc, "", size=4, space_after=0, line=5)
    return t


def box(doc, lines, fill=BOX_FILL, border=BORDER_COLOR, size=9, bold_first=False):
    """囲み（強調・注記）。"""
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    cell = t.cell(0, 0)
    cell.width = Mm(180)
    _cell_shade(cell, fill)
    _cell_margins(cell, top=60, bottom=60, left=140, right=140)
    first = cell.paragraphs[0]
    for i, line in enumerate(lines):
        p = first if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(12)
        _run(p, line, size=size, bold=(bold_first and i == 0))
    _tbl_width(t)
    _borders(t, border)
    para(doc, "", size=4, space_after=0, line=5)
    return t


def note(doc, text):
    return box(doc, [text], fill=NOTE_FILL, border="C9A03B", size=8)


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    _run(p, "", size=1).add_break(WD_BREAK.PAGE)


# ============================================================
# 本文データ
# ============================================================
HYOKI_TEXT = (
    "本計画では「障害」などの「害」の字の表記について、字に対する印象に配慮するとともに、"
    "障がい者の人権をより尊重する観点から、国の法令等に基づく法律用語や施設名等の固有名称を除き、"
    "可能な限り「害」の字をひらがなで表記しています。このため、本計画では「がい」と「害」の字が"
    "混在する表現となっています。"
)

ICHIZUKE = [
    ["区分", "障害者計画", "障害福祉計画", "障害児福祉計画"],
    ["根拠法", "障害者基本法\n第11条第３項", "障害者総合支援法\n第88条第１項",
     "児童福祉法\n第33条の20第１項"],
    ["内容", "障がい者施策の基本的方向性について定める計画",
     "障害福祉サービス等の見込量とその確保策を定める計画",
     "障害児通所支援等の提供体制と確保策を定める計画"],
    ["今回の改定", "対象外\n（第４次計画を継続）", "第８期として策定", "第４期として策定"],
]

KIKAN = [
    ["計画", "計画期間", "今回の取扱い"],
    ["第４次北塩原村障がい者計画", "令和６年度〜令和11年度（６年間）",
     "改定対象外。基本理念・基本目標・施策体系を本計画に継承"],
    ["第８期北塩原村障がい福祉計画", "令和９年度〜令和11年度（３年間）", "今回策定"],
    ["第４期北塩原村障がい児福祉計画", "令和９年度〜令和11年度（３年間）", "今回策定"],
]

SHISHIN_POINTS = [
    "成果目標が７項目から８項目になりました。「障がい福祉人材の確保・定着、"
    "当事者視点に立ったケアの充実のための生産性向上」が新たに加わっています。",
    "就労選択支援（令和７年10月開始）の積極的な利用促進が求められ、"
    "就労選択支援に係る成果目標が新設されました。",
    "障がい児支援では、児童発達支援センター等の「４つの中核機能」の確保と、"
    "インクルージョン（地域社会への参加・包容）推進のための協議の場の設置が求められています。",
    "のぞまないセルフプラン（自己作成のサービス等利用計画）の解消が、"
    "相談支援体制の成果目標に位置づけられました。",
    "障害福祉サービス等情報公表制度の公表率に関する成果目標が新設されました。",
    "人口減少地域におけるサービスの維持・確保、高次脳機能障害のある方への支援、"
    "災害時におけるサービス提供の確保、スポーツ・健康増進活動による社会参加の促進などが、"
    "新たな視点として示されています。",
]

TECHOU = [
    ["区分", "平成31年", "令和２年", "令和３年", "令和４年", "令和５年", "令和11年\n（中位推計）"],
    ["身体障害者手帳", "126人", "121人", "128人", "122人", "114人", "103人"],
    ["療育手帳", "13人", "13人", "14人", "12人", "12人", "11人"],
    ["精神障害者保健福祉手帳", "24人", "28人", "27人", "26人", "30人", "31人"],
    ["合計", "163人", "162人", "169人", "160人", "156人", "145人"],
]

RINEN = "「障がいのあるなしに関わらず、お互いの人格や個性を尊重し、多様な価値観を認め合い、誰もが自分らしく輝くむら」"

MOKUHYOU = [
    ("①　障がいへの理解を深め、交流を育むむら",
     "障がいの有無に関わらず、個人の人格や個性を尊重し合うことを目指し、"
     "障がいを理由とする差別の解消や障がいのある人に関する正しい理解が必要であるため、"
     "広報等を活用した啓発活動を推進します。"),
    ("②　共に支え合い、誰もが安心して暮らせるむら",
     "すべての人が共に協力し合い、誰もが住み慣れた地域で安心して暮らすことを目指し、"
     "障がいにつながる疾病を予防・軽減するため、保健・医療との連携や"
     "災害などの緊急時における安全・安心の確保に取り組みます。"),
    ("③　みんなが輝き、自立した生活を送れるむら",
     "生きがいを持って活動できる社会を目指し、障がいのある人も様々な可能性の中から"
     "自分らしい生き方を選択できるように、ライフステージに沿った切れ目のない支援を進めます。"),
]

SHISAKU = [
    ["基本施策", "第８期計画で加わる視点", "目標値（現状→令和11年度）"],
    ["①　啓発・広報",
     "合理的配慮の提供義務化の周知、意思疎通支援従事者の養成・派遣体制の整備、"
     "障がい者等に対する虐待の防止",
     "差別や偏見を感じている人の割合\n25.3％→15％"],
    ["②　保健・医療",
     "精神障害にも対応した地域包括ケアシステムの構築、医療的ケア児等への支援、"
     "高次脳機能障害のある方への支援",
     "特定健康診査の受診率\n51.3％→60％"],
    ["③　福祉",
     "地域生活支援拠点等の年１回以上の検証、強度行動障害を有する方の支援ニーズの把握、"
     "のぞまないセルフプランの解消、人口減少地域におけるサービスの維持・確保",
     "福祉サービスに満足している人の割合\n52％→60％"],
    ["④　教育・育成",
     "児童発達支援センターの４つの中核機能の確保、インクルージョン推進のための"
     "協議の場の設置、こども家庭センターとの連携",
     "医療的ケア児の受入れ体制の整備\n無→有"],
    ["⑤　雇用・就業",
     "就労選択支援の利用促進、法定雇用率2.7％への引上げを踏まえた事業主啓発",
     "障がい者就労施設等からの物品調達件数\n０件→３件"],
    ["⑥　生活環境",
     "災害時における障害福祉サービス提供の確保、個別避難計画の作成促進、"
     "冬季・積雪時の避難支援体制の整備",
     "公共施設トイレの設置率\n56.8％→100％\n個別避難計画\n１件→10件"],
    ["⑦　スポーツ・文化",
     "スポーツ・健康増進活動による社会参加等の促進、"
     "地域共生社会の実現に向けた取組の一層の推進",
     "外出の目的が趣味・スポーツ等である人の割合\n28.1％→40％"],
]

SEIKA = [
    ["成果目標", "現状（基準）", "令和11年度の目標値"],
    ["（１）福祉施設の入所者の\n　　　地域生活への移行",
     "施設入所者数　４人\n（令和４年度末。令和７年度末の実数は確認中）",
     "地域生活移行者数　０人\n施設入所者削減見込み　０人\n（入所者は重度の方が中心のため、個別の移行"
     "可能性を自立支援協議会で検討します）"],
    ["（２）精神障害にも対応した\n　　　地域包括ケアシステムの構築",
     "精神病床に１年以上入院している方　５人\n（第７期計画時点）",
     "退院者数　１人\n心のサポーター養成講座の実施【新規】\nＫ６等によるこころの状態を把握する機会の提供【新規】"],
    ["（３）福祉施設から\n　　　一般就労への移行等",
     "一般就労移行者数・就労定着支援利用者数\n（令和６年度実績を確認中）",
     "一般就労移行者数　１人\n就労定着支援事業利用者数　１人\n就労選択支援利用者数　年１〜２人【新規】"],
    ["（４）障がい児支援の\n　　　提供体制の整備等",
     "児童発達支援センター　０か所\n保育所等訪問支援事業所　０か所\n医療的ケア児　１人",
     "近隣市町村との連携により４つの中核機能へのアクセスを確保（０か所）\n"
     "医療的ケア児等コーディネーター　１人\nインクルージョン推進の協議の場を会津北部圏域で設置【新規】\n"
     "医療的ケア児等に関する協議の場を圏域で設置【新規】"],
    ["（５）地域生活支援の充実",
     "会津北部地域生活支援拠点（４町村共同）　１か所\nコーディネーター　１人\n"
     "強度行動障害のある方　０人（令和４年度末）",
     "拠点の運用状況の検証　年１回以上\n強度行動障害を有する方の支援ニーズを圏域で把握【新規】"],
    ["（６）相談支援体制の\n　　　充実・強化等",
     "基幹相談支援センター　未設置\n協議会専門部会　０か所\nセルフプラン率　０％",
     "基幹相談支援センター　１か所（４町村広域）\n協議会専門部会　２か所\nのぞまないセルフプラン　０件【新規】"],
    ["（７）障がい福祉人材の確保・定着\n　　　及び生産性向上【新規】",
     "県が実施する研修への村職員の参加　１人\n（令和４年度）",
     "県研修への参加　２人以上\n県のワンストップ窓口等の支援策を村内事業所へ周知"],
    ["（８）障がい福祉サービス等の\n　　　質の向上に係る体制の構築",
     "県が実施する研修への村職員の参加　１人",
     "県研修への参加　２人\n村内対象事業所の情報公表制度　公表率100％【新規】"],
]

PDCA = [
    ["段階", "内容"],
    ["Plan（計画）", "成果目標及びサービスの見込量（活動指標）を定めます。"],
    ["Do（実行）", "計画に基づきサービスの提供体制を確保し、事業を実施します。"],
    ["Check（評価）",
     "少なくとも年１回、実績を把握し、成果目標及び活動指標の達成状況を分析・評価します。"],
    ["Action（改善）",
     "評価の結果を踏まえ、必要に応じて計画の変更やサービス提供体制の確保に係る対策を講じます。"],
]

SCHEDULE = [
    ["時期", "内容"],
    ["令和８年度中", "アンケート調査の実施・集計"],
    ["令和８年度中", "北塩原村障がい者自立支援協議会（計画素案の審議）"],
    ["令和８年度中", "庁議（計画案の決定）"],
    ["令和８年度中", "パブリックコメントの実施"],
    ["令和９年３月", "計画の確定・公表（予定）"],
]

# 見込量の基準値は kitashiobara_common.py の個別積上げ用マスタから取る
MIKOMI_PICK = [
    "居宅介護", "生活介護", "自立訓練（生活訓練）", "就労移行支援",
    "就労継続支援A型", "就労継続支援B型", "就労定着支援", "共同生活援助",
    "施設入所支援", "計画相談支援",
    "児童発達支援", "放課後等デイサービス", "障がい児相談支援",
]


def mikomi_rows():
    """主なサービスの基準値（令和8年度）を2列組の表にする。"""
    base = {name: users for name, _u, _q, users, *_rest in SERVICES_ADULT}
    base.update({name: users for name, _u, _q, users, *_rest in SERVICES_CHILD})
    missing = [n for n in MIKOMI_PICK if n not in base]
    if missing:
        raise LookupError(f"見込量マスタに無いサービスです: {missing}")

    pairs = [(n, f"{base[n]}人／月") for n in MIKOMI_PICK]
    half = (len(pairs) + 1) // 2
    left, right = pairs[:half], pairs[half:]
    rows = [["サービス", "利用者数", "サービス", "利用者数"]]
    for i in range(half):
        a = left[i]
        b = right[i] if i < len(right) else ("", "")
        rows.append([a[0], a[1], b[0], b[1]])
    return rows


# ============================================================
# ページ数の見積り（4ページに収まっているかの確認用）
# ============================================================
def estimate_lines(doc):
    """改ページで区切って、各ページのおよその行数を数える。"""
    pages = [0]
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = "".join(t.text or "" for t in child.iter(qn("w:t")))
            if child.findall(".//" + qn("w:br") + "[@" + qn("w:type") + "='page']"):
                pages.append(0)
                continue
            pages[-1] += max(1, -(-len(text) // CHARS_PER_LINE)) if text else 0.4
        elif child.tag == qn("w:tbl"):
            for tr in child.findall(qn("w:tr")):
                longest = 1
                for tc in tr.findall(qn("w:tc")):
                    text = "".join(t.text or "" for t in tc.iter(qn("w:t")))
                    hard = text.count("\n") + 1
                    width = int(tc.find(qn("w:tcPr")).find(qn("w:tcW")).get(qn("w:w"))) \
                        if tc.find(qn("w:tcPr")) is not None and \
                        tc.find(qn("w:tcPr")).find(qn("w:tcW")) is not None else PAGE_W
                    chars = max(4, int(CHARS_PER_LINE * width / PAGE_W))
                    longest = max(longest, max(hard, -(-len(text) // chars)))
                pages[-1] += longest * 0.92
    return pages


# ============================================================
# 各ページ
# ============================================================
def page1(doc):
    # 表紙帯
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    cell = t.cell(0, 0)
    cell.width = Mm(180)
    _cell_shade(cell, DARK)
    _cell_margins(cell, top=120, bottom=120, left=160)
    lines = [
        ("◆　第８期北塩原村障がい福祉計画", 13),
        ("◆　第４期北塩原村障がい児福祉計画", 13),
        ("（令和９年度〜令和11年度）　素案　概要版", 11),
    ]
    for i, (text, size) in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(size + 6)
        _run(p, text, size=size, bold=True, color="FFFFFF")
    _tbl_width(t)
    _no_borders(t)
    para(doc, "", size=4, space_after=0, line=5)

    p = para(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=4)
    _run(p, "作成年月　令和８年９月　／　発行　北塩原村　／　"
            "編集　北塩原村保健福祉課　福祉係　（TEL）0241-23-3113", size=8.5)

    band(doc, "１", "計画策定の背景・趣旨")
    para(doc,
         "本村では、令和６年３月に「第４次北塩原村障がい者計画・第７期北塩原村障がい福祉計画・"
         "第３期北塩原村障がい児福祉計画」を策定し、障がいのある人もない人も互いに人格と個性を"
         "尊重し合う共生社会の実現を目指してきました。")
    para(doc,
         "このうち「第７期北塩原村障がい福祉計画」及び「第３期北塩原村障がい児福祉計画」は"
         "令和８年度で計画期間が終了することから、令和８年３月31日に告示された国の基本指針、"
         "これまでの本村の取組み及び障がいのある方のニーズを踏まえ、令和９年度から令和11年度までの"
         "３年間について、障がい福祉サービス等の見込量とその確保のための方策を定めるものです。")
    para(doc,
         "なお、「第４次北塩原村障がい者計画」は令和11年度までを計画期間としており、"
         "今回の改定の対象ではありません。その基本理念・基本目標及び施策体系は、"
         "本計画においても共有する理念・目標・体系として継承します。", space_after=4)

    sub(doc, "「障害」と「障がい」の表記について")
    para(doc, HYOKI_TEXT, space_after=4)

    band(doc, "２", "計画の位置づけ")
    para(doc, "本計画は、以下の法律に基づいて策定する法定計画です。")
    table(doc, ICHIZUKE, [1250, 2985, 2985, 2986], aligns=["center", None, None, None],
          fill_first=BAND)

    band(doc, "３", "計画の期間")
    table(doc, KIKAN, [3000, 2700, 4506], fill_first=BAND)

    band(doc, "４", "国の基本指針の見直しのポイント")
    para(doc,
         "第８期障害福祉計画及び第４期障害児福祉計画に係る国の基本指針は令和８年３月31日に"
         "告示され、主な見直し事項として次の点が示されています。")
    bullets(doc, SHISHIN_POINTS, size=8.5)


def page2(doc):
    band(doc, "５", "本村の障がい者の現状")
    para(doc,
         "本村の人口は令和６年４月１日現在2,394人で、平成31年（2,743人）から349人減少し、"
         "高齢化率は41.4％となっています。令和11年の人口は2,185人と推計されています。")

    sub(doc, "障がい者手帳の所持者数の推移と将来推計")
    table(doc, TECHOU, [2206, 1200, 1200, 1200, 1200, 1200, 2000],
          aligns=[None, "center", "center", "center", "center", "center", "center"])
    box(doc, [
        "障がい者手帳をお持ちの方は156人（令和５年４月１日現在）で、人口の約6.4％、"
        "約15人に１人の割合です。",
        "　身体障がいのある人　114人（4.7％）／知的障がいのある人　12人（0.5％）／"
        "精神障がいのある人　30人（1.2％）",
        "令和11年には合計145人と見込まれます。身体障がいのある方は減少、"
        "精神障がいのある方は増加する傾向にあります。",
    ])
    note(doc, "【要更新】手帳所持者数は現行計画（令和５年４月１日現在）の数値です。"
              "令和８年４月１日現在の最新値を村に確認しており、確定後に更新します。"
              "令和11年の推計値は、トレンド延長法と所持率法の中位値です。")

    band(doc, "６", "計画の基本理念")
    box(doc, [RINEN], fill=NOTE_FILL, border=ACCENT, size=10.5, bold_first=True)
    para(doc,
         "障がいの有無に関係なく、地域に暮らす人々がお互いに支え合いながら共に生きる"
         "「地域共生社会」の理念を基本としています。本村は、特性の異なる多様な地域が"
         "手を取り合い、各々の個性を磨きながら生活しています。これを踏まえ、人格や個性、"
         "価値観を認め合いながら、誰もが自分らしく生活できる地域社会の実現を基本理念とします。",
         space_after=4)

    band(doc, "７", "基本目標")
    para(doc, "基本理念を実現するため、次の３つの目標を設定し、各種施策を推進します。")
    for title, text in MOKUHYOU:
        sub(doc, title)
        para(doc, text, space_after=4)


def page3(doc):
    band(doc, "８", "継承する施策体系と第８期計画で加わる視点")
    para(doc,
         "基本理念及び基本目標のもとに、第４次北塩原村障がい者計画が定める７つの基本施策を"
         "継承します。第８期計画で新たに加わる視点は、次のとおり各施策に接続します。")
    table(doc, SHISAKU, [1500, 5706, 3000], size=8)
    note(doc, "【要更新】目標値は第４次北塩原村障がい者計画が令和11年度を目標年度として定めた"
              "ものです。このうちアンケートで測る指標は、本計画の策定に係るアンケート調査の"
              "集計結果により現状値を更新します。")

    band(doc, "９", "令和11年度に向けた成果目標")
    para(doc,
         "障がいのある人の自立支援の観点から、地域生活への移行や就労支援等の課題に対応するため、"
         "令和11年度を目標年度として、国の基本指針に即し、次の８項目について成果目標を設定します。"
         "第７期計画までは７項目でしたが、第８期計画では（７）が新設され８項目になりました。")
    table(doc, SEIKA[:7], [2400, 3400, 4406], size=8)


def page4(doc):
    sub(doc, "９　令和11年度に向けた成果目標（続き）", fill=BAND)
    table(doc, [SEIKA[0]] + SEIKA[7:], [2400, 3400, 4406], size=8)
    note(doc, "【要更新】現状（基準）の数値は、村に令和７年度末時点の実績を確認しているところです。"
              "目標値は北塩原村障がい者自立支援協議会での審議を経て確定します。")

    band(doc, "10", "障がい福祉サービス等の見込量")
    bullets(doc, [
        "本村は利用者数が少なく、１人の増減が見込量を大きく動かすため、サービスごとに"
        "「継続・流入・流出」を１人単位で積み上げる方法により見込量を算定します。"
        "国が示す推計方法（過去の変化率の平均、人口当たりの利用率）は、検証に用います。",
        "本村は全域が過疎地域であるため、国の基本指針が定める地域差の是正に関する"
        "算定方法（別表第五）の適用対象外です。",
        "介護保険に相当するサービスがない共同生活援助等は、65歳到達後も引き続き"
        "障がい福祉サービスを利用できます。年齢到達を理由に一律に打ち切ることはしません。",
    ], size=8.5)
    para(doc,
         "・令和７年度（給付実績が完結した直近の年度）を基準年度として第1次概算を行いました。"
         "第８期の３か年の給付費は約1億7,286万円（各年度約5,762万円）と見込んでいます。",
         size=8.5, indent=9)
    para(doc, "＜主なサービスの利用者数（令和８年度・基準値）＞", size=8.5, bold=True)
    table(doc, mikomi_rows(), [3053, 2050, 3053, 2050], size=8,
          aligns=[None, "center", None, "center"])
    note(doc, "【要更新】基準値は現行計画（第７期）の令和８年度見込量です。"
              "令和７年度の給付実績を村と精査のうえ、令和９〜11年度の見込量を確定します。")

    band(doc, "11", "計画の推進体制")
    para(doc,
         "計画に定める事項については、PDCAサイクルにより進行管理を行います。"
         "評価は保健福祉課が行い、北塩原村障がい者自立支援協議会に報告して意見を聴いたうえで、"
         "結果を村ホームページ等で公表します。")
    table(doc, PDCA, [1800, 8406], size=8.5, fill_first=BAND)

    band(doc, "12", "成年後見制度の利用促進")
    para(doc,
         "本村では、認知症、知的障がい、その他の精神上の障がいがあることにより、"
         "日常生活を支える様々な行為や買い物、財産管理が難しい事例が見られます。"
         "高齢化の進展を踏まえると、成年後見制度の必要性は今後も高まっていくことが見込まれます。")
    para(doc,
         "令和４年７月より中核機関「会津権利擁護・成年後見センター」を会津圏域11市町村で設置し、"
         "制度の周知啓発、広報活動、相談対応を行っています。令和６年度からは市民後見人養成事業が"
         "追加されました。本計画期間においても同センターとの連携を継続します。", space_after=4)

    band(doc, "", "今後の予定")
    table(doc, SCHEDULE, [2400, 7806], size=8.5, fill_first=BAND)


# ============================================================
def setup_section(doc):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)
    # 既定スタイルを本文書体に合わせる
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(9)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = _el("rFonts")
        rpr.insert(0, rf)
    rf.set(qn("w:eastAsia"), FONT)
    rf.set(qn("w:ascii"), FONT)
    rf.set(qn("w:hAnsi"), FONT)


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    doc = docx.Document()
    setup_section(doc)

    page1(doc)
    page_break(doc)
    page2(doc)
    page_break(doc)
    page3(doc)
    page_break(doc)
    page4(doc)

    doc.save(OUT_FILE)

    pages = estimate_lines(doc)
    print(f"作成: {OUT_FILE}")
    print(f"  ページ数: {len(pages)}")
    for i, lines in enumerate(pages, 1):
        mark = "" if lines <= LINES_PER_PAGE else "　← 収まりません"
        print(f"  p.{i}  推定 {lines:.0f} 行 / 目安 {LINES_PER_PAGE} 行{mark}")


if __name__ == "__main__":
    main()
