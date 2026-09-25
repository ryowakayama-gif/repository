"""小野町 見える化システムの設定の確認結果と、小野町自身の交付金の得点分析。

令和8年9月25日に受領
  画面キャプチャ8枚（将来推計の設定・施策反映の3画面・第9期の画面）
  見える化システムの出力 58件（W126〜W155、J1〜J15、F15〜F27）

**この受領により、推定にとどまっていた3点が確定した。**
  ① 令和8年度が何か月分か  → 令和8年5月の1か月（設定画面で確定）
  ② 現在の推計方法の設定    → 基本推計・認定率の伸びは令和6→令和7・利用率の伸び0
  ③ 施策反映が入っているか  → 入っていない（3画面とも「自然体推計に全て戻す」が押せない）

**あわせて解消しなかったことも確定した。**
  第9期の推計パターンはシステムに残っていない（画面が空で保険料も「―」）。

出力
  11_見える化出力依頼/小野町_見える化システムの設定の確認結果_YYYYMMDD.docx
  08_協議会資料/小野町_交付金_小野町の得点分析_YYYYMMDD.xlsx
"""

import pathlib
import warnings

warnings.filterwarnings("ignore")

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = pathlib.Path(__file__).parent
BASE = ROOT / "小野町_引継ぎ_整理済"
SRC = BASE / "11_見える化出力依頼" / "原本_見える化出力_20260925"
OUT1 = BASE / "11_見える化出力依頼"
OUT2 = BASE / "08_協議会資料"
ASOF = "20260925"
ASOF_JP = "令和8年9月25日"
JP_MIN = "游明朝"
JP_GO = "游ゴシック"

HEAD = PatternFill("solid", fgColor="1F3864")
NG = PatternFill("solid", fgColor="FCE4E4")
WARN = PatternFill("solid", fgColor="FFE0B2")
OK = PatternFill("solid", fgColor="E2EFDA")
CALC = PatternFill("solid", fgColor="EAF1FB")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# システムが表示している保険料（暫定値）と、当方の算定。
SYS_D10 = 6466          # 第10期（暫定値・月額）
SYS_R12 = 7298          # 令和12年度（暫定値・月額）
OUR_D10 = 6210          # 当方の標準（見える化の仕様にそろえた自然体推計）
PCT1 = 57.9             # 総給付費1％＝保険料の月額（円）

# 画面から読み取った設定。キャプチャは原本として保存してある。
SETTEI = [
    ("推計名", "将来推計1", "―"),
    ("公表時点", "**暫定版データ**", "確定版の公表後に再出力が要る"),
    ("（1）1 認定者数の自然体推計の粒度",
     "**性別／年齢5歳階級別／要介護度別に推計する（初期値）**",
     "**いわゆる基本推計。包括推計ではない**"),
    ("（1）2 介護サービス利用者数の粒度",
     "サービス別／要介護度別に推計する（初期値）", "―"),
    ("（2）1 令和8年度の認定者数",
     "実績値の設定なし（＝月報の値をそのまま使用）",
     "**月報は令和8年5月の1か月分である（下表）**"),
    ("（2）2 認定率の伸び",
     "**当該市町村における令和6年度→令和7年度の伸び**",
     "**完結年度どうしの伸びであり、当方の判断と一致する**"),
    ("（3）1 令和8年度の施設・居住系利用者数",
     "編集なし（＝月報から計算した実績見込み値）", "同上"),
    ("（3）2 施設・居住系の利用率等の伸び", "**0（ゼロ）**",
     "**利用率の変化を0とする。当方の算定と一致**"),
    ("（4）1・2 令和8年度の在宅利用者数・回（日）数",
     "編集なし（＝月報から計算した実績見込み値）", "同上"),
    ("（4）3 在宅サービス利用率の伸び", "**0（ゼロ）**", "同上"),
    ("（4）4 在宅1人1月あたり利用回（日）数の伸び", "**0（ゼロ）**", "同上"),
    ("（5）1・2 1人1月あたり給付費",
     "独自設定なし（＝最新年度の実績をベース）", "―"),
]

HOKOKU = [
    ("令和6年度", "**年報を使用します**", "12か月・確報"),
    ("令和7年度", "**月報を使用する。5月〜4月の12か月すべてにチェック**",
     "**12か月＝完結している**"),
    ("令和8年度", "**月報を使用します。5月のみチェック**"
                "（6月〜4月は未チェック）",
     "**令和8年5月の1か月分である。**"
     "当方が利用者数の整数率（非ゼロ26件がすべて整数）から"
     "判定していたことが、設定画面で裏づけられた"),
]

SISAKU = [
    ("認定者数", "**グレーアウト（押せない）**", "施策反映は入っていない"),
    ("施設・居住系サービス利用者数", "**グレーアウト（押せない）**", "同上"),
    ("在宅サービス利用者数等", "**グレーアウト（押せない）**", "同上"),
]

# 画面で読み取れた、令和8年5月の1か月が第10期に固定されている実例（訪問介護）。
HOMON_RIYORITSU = [
    ("令和6年度", 10.1, 16.3, 18.0, 10.0, 5.2),
    ("令和7年度", 12.5, 12.3, 10.2, 13.2, 3.7),
    ("**令和8年度**", 14.4, 9.2, 17.3, 11.5, 11.9),
    ("令和9年度", 14.4, 9.2, 17.3, 11.5, 11.9),
    ("令和10年度", 14.4, 9.2, 17.3, 11.5, 11.9),
    ("令和11年度", 14.4, 9.2, 17.3, 11.5, 11.9),
]
HOMON_KAISU = [
    ("令和6年度", 9.1, 12.0, 19.2, 34.4, 18.8),
    ("令和7年度", 8.9, 12.0, 20.4, 24.8, 33.8),
    ("**令和8年度**", 7.4, 7.0, 15.5, 19.5, 44.8),
    ("令和9年度", 7.4, 7.0, 15.5, 19.5, 44.8),
    ("令和10年度", 7.4, 7.0, 15.5, 19.5, 44.8),
    ("令和11年度", 7.4, 7.0, 15.5, 19.5, 44.8),
]
HOMON_NINZU = [("令和6年度", 68), ("令和7年度", 51), ("**令和8年度**", 62),
               ("令和9年度", 60), ("令和10年度", 61)]


# ---------------------------------------------------------------- 交付金

def read_w(k):
    """W指標のファイルから (項目, {全国/福島県/小野町}) を返す。"""
    f = sorted(SRC.glob(f"{k}_*.xlsx"))
    if not f:
        return []
    wb = openpyxl.load_workbook(f[0], data_only=True)
    nm = [s for s in wb.sheetnames if s.startswith("表形式")]
    if not nm:
        return []
    ws = wb[nm[0]]
    hr = next((r for r in range(1, 8)
               if "全国" in [str(ws.cell(r, c).value) for c in range(1, 14)]),
              None)
    if hr is None:
        return []
    cols = {ws.cell(hr, c).value: c for c in range(1, 14)
            if ws.cell(hr, c).value in ("全国", "福島県", "小野町")}
    out = []
    for r in range(hr + 1, ws.max_row + 1):
        lab = ws.cell(r, 2).value
        if not lab or str(lab).startswith(("行", "（時点）", "（出典）")):
            continue
        d = {k2: ws.cell(r, cc).value for k2, cc in cols.items()}
        if any(isinstance(v, (int, float)) for v in d.values()):
            out.append((str(lab), d))
    return out


def kofukin():
    """小野町の交付金の得点（令和5年＝令和6年度交付金）。"""
    sogo = read_w("W126")
    mok = {"推進": read_w("W127"), "支援": read_w("W129")}
    gun = {"推進": read_w("W128"), "支援": read_w("W130")}
    meisai, zero, ue = [], [], []
    for k in [f"W{i}" for i in range(131, 138)]:
        for lab, d in read_w(k):
            if "合計" in lab:
                continue
            t, z, ken = d.get("小野町"), d.get("全国"), d.get("福島県")
            if not (isinstance(t, (int, float)) and isinstance(z, (int, float))):
                continue
            meisai.append((k, lab, t, z, ken))
            if t == 0:
                zero.append((k, lab, t, z, ken))
            elif t > z:
                ue.append((k, lab, t, z, ken))
    zero.sort(key=lambda x: -x[3])
    ue.sort(key=lambda x: -(x[2] - x[3]))
    return {"総合": sogo, "目標": mok, "指標群": gun,
            "明細": meisai, "0点": zero, "上回る": ue}


# ---------------------------------------------------------------- 体裁

def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = JP_MIN
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), JP_MIN)
    return doc


def head(doc, text, level):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = JP_GO
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_GO)


def _emph(p, text):
    for i, part in enumerate(str(text).split("**")):
        if not part:
            continue
        r = p.add_run(part)
        r.font.name = JP_MIN
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_MIN)
        r.bold = bool(i % 2)


def body(doc, *paras):
    for t in paras:
        _emph(doc.add_paragraph(), t)


def note(doc, text):
    p = doc.add_paragraph()
    _emph(p, text)
    for r in p.runs:
        r.font.size = Pt(9)


def table(doc, header_row, rows, right_from=99):
    t = doc.add_table(rows=1, cols=len(header_row))
    t.style = "Table Grid"
    for i, h in enumerate(header_row):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = JP_GO
        r._element.rPr.rFonts.set(qn("w:eastAsia"), JP_GO)
        r.bold = True
        r.font.size = Pt(9)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            _emph(p, v)
            for r in p.runs:
                r.font.size = Pt(9)
            if i >= right_from:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()


# ---------------------------------------------------------------- Word

def build_docx(K):
    doc = new_doc()
    head(doc, "見える化システムの設定の確認結果", 0)
    body(doc, f"小野町　第10期介護保険事業計画　／　{ASOF_JP}")

    head(doc, "0　結論", 1)
    table(doc, ["#", "確かめたこと", "結果"], [
        ("1", "令和8年度は何か月分か",
         "**令和8年5月の1か月である。**"
         "介護保険事業状況報告の設定で、令和8年度は5月のみにチェックが"
         "入っている（令和7年度は5月〜4月の12か月すべて）。"
         "**当方が利用者数の整数率から判定していたことが裏づけられた**"),
        ("2", "現在の推計方法の設定",
         "**基本推計（性別／年齢5歳階級別／要介護度別）。**"
         "認定率の伸びは令和6年度→令和7年度、"
         "施設・居住系と在宅の利用率の伸びおよび"
         "1人1月あたり利用回（日）数の伸びはいずれも0（ゼロ）"),
        ("3", "施策反映が入っているか",
         "**入っていない。**3画面とも「自然体推計に全て戻す」が"
         "グレーアウトしている。"
         "**他町村の案件で起きた「伸びの窓を変えても見込量が動かない」"
         "事態は、小野町では起きていない**"),
        ("4", "第9期の推計パターン",
         "**残っていない。**第9期の画面は推計名が空欄で保険料も「―」である。"
         "**第9期の乖離の原因は、当方の推定にとどまる**"),
        ("5", "システムが示す保険料との差",
         f"**システム{SYS_D10:,}円（暫定値）に対し当方{OUR_D10:,}円で"
         f"{SYS_D10 - OUR_D10:+,}円。**"
         "原因は令和8年度の置き方にある（第3章）"),
    ])

    head(doc, "1　令和8年度は令和8年5月の1か月である", 1)
    body(doc,
         "**「介護保険事業状況報告の設定」画面で確定しました。**"
         "「公表時点データを使用する」が選ばれています。")
    table(doc, ["年度", "設定", "意味"],
          [(y, s, m) for y, s, m in HOKOKU])
    body(doc,
         "**これは推定ではなく設定そのものです。**"
         "当方はこれまで、見える化ワークシートの利用者数の非ゼロ26件が"
         "すべて整数であることから令和8年度を1か月分と判定していましたが、"
         "**判定の方法と結論の両方が正しかったことが確かめられました。**")

    head(doc, "2　推計方法の設定", 1)
    table(doc, ["設定項目", "現在の設定", "当方の算定との関係"],
          [(k, v, m) for k, v, m in SETTEI])
    body(doc,
         "**認定率の伸びが「令和6年度→令和7年度」である点は重要です。**"
         "選べる伸びのうち完結年度どうしで組めるのはこれだけであり、"
         "当方が伸びの窓の検討で到達した結論と一致します。",
         "**利用率と1人1月あたり利用回（日）数の伸びがいずれも0である点も、"
         "当方の自然体推計と同じです。**"
         "令和8年9月25日の再検証で延ばし方を見える化の仕様にそろえた際の"
         "前提が、画面で裏づけられました。")

    head(doc, "3　システムの6,466円と当方の6,210円の差", 1)
    body(doc,
         f"**差は{SYS_D10 - OUR_D10:+,}円です。**"
         "設定は同じなのに差が出るのは、**令和8年度の水準の置き方**が"
         "違うためです。",
         "**システムは令和8年5月の1か月をそのまま基点にしています。**"
         "当方は令和7年度の実績に国保連月報3か月（令和8年4〜6月提供分）の"
         "前年同期比0.9598を乗じて令和8年度を置いています。")
    table(doc, ["区分", "システム", "当方"], [
        ("令和8年度の置き方",
         "**令和8年5月の1か月**（月報1か月）",
         "**令和7年度実績×0.9598**（月報3か月の前年同期比）"),
        ("第10期への延ばし方", "利用率・回数の伸び0", "同じ"),
        ("保険料基準額（月額）", f"**{SYS_D10:,}円**（暫定値）",
         f"**{OUR_D10:,}円**"),
    ], right_from=1)
    body(doc,
         "**1か月がどれだけ振れるかは、訪問介護の画面で見えます。**"
         "利用率（％）と1人1月あたり利用回数は、"
         "令和8年度の値が令和9年度以降そのまま固定されます。")
    table(doc, ["年度", "要介護1", "要介護2", "要介護3", "要介護4", "要介護5"],
          [(y, f"{a}", f"{b}", f"{c}", f"{d}", f"{e}")
           for y, a, b, c, d, e in HOMON_RIYORITSU], right_from=1)
    note(doc, "※ 訪問介護の在宅サービス利用率（％）。見える化システムの画面による。")
    table(doc, ["年度", "要介護1", "要介護2", "要介護3", "要介護4", "要介護5"],
          [(y, f"{a}", f"{b}", f"{c}", f"{d}", f"{e}")
           for y, a, b, c, d, e in HOMON_KAISU], right_from=1)
    note(doc, "※ 訪問介護の1人1月あたり利用回数（回）。同上。")
    body(doc,
         "**要介護5の1人1月あたり利用回数は、令和7年度33.8回から"
         "令和8年度44.8回へ32.5％上がり、その44.8回が令和11年度まで続きます。**"
         "要介護4は逆に24.8回から19.5回へ21.4％下がります。"
         "**1か月の振れが3年間に効いているということです。**",
         "利用者数も同じです。訪問介護の在宅サービス利用者数は"
         "令和6年度68人・令和7年度51人に対し、令和8年度62人（令和7年度の1.22倍）で、"
         "令和9年度60人・令和10年度61人と続きます。")
    table(doc, ["年度", "訪問介護の利用者数（人／月）"],
          [(y, f"{v}") for y, v in HOMON_NINZU], right_from=1)
    body(doc,
         "**当方が月報3か月を使っているのは、1か月の振れを避けるためです。**"
         f"総給付費1％が保険料の月額{PCT1}円にあたるため、"
         f"{SYS_D10 - OUR_D10}円の差は総給付費で"
         f"約{(SYS_D10 - OUR_D10) / PCT1:.1f}％にあたります。",
         "**どちらを計画値とするかは協議会にお諮りする事項です。**"
         "令和8年度の残りの月の月報が届けば、3か月より長い期間で"
         "確かめられます。")

    head(doc, "4　施策反映は入っていない", 1)
    table(doc, ["画面", "「自然体推計に全て戻す」の状態", "意味"],
          [(k, v, m) for k, v, m in SISAKU])
    body(doc,
         "**3画面ともグレーアウトしており、自然体推計のままです。**"
         "他町村の案件では、このボタンが押せる状態のまま総括表を出したために"
         "伸びの窓を変えても見込量が1値も動かないという事態が起きていましたが、"
         "**小野町では起きていません。**")

    head(doc, "5　第9期の推計パターンは残っていない", 1)
    body(doc,
         "**第9期の画面は、推計名が空欄で保険料額も「―」でした。**"
         "第9期策定時の基準年度・伸びの窓・推計手法は、"
         "システムからは読み取れません。",
         "**したがって、第9期の乖離（実績／見込みが令和6年度98.6％・"
         "令和7年度91.1％）の原因は、当方の推定にとどまります。**"
         "第9期の見込みが3か年ほぼ同値（振れ幅0.42％）で自然体推計そのものであり、"
         "その水準が令和5年度の実績より7.3％高いことから、"
         "基準年度の置き方に原因があったとみるのが自然ですが、"
         "**設定そのものを確かめる手段はありません。**"
         "町の起案文書が残っていればそれによります。")

    head(doc, "6　「推計パターン毎の乖離状況」ボタンについて", 1)
    body(doc,
         "**推計方法の設定画面に「推計パターン毎の乖離状況」ボタンがあります"
         "（（1）1・2 の右）。**"
         "当方はこれまで、推計手法（基本推計／包括推計）のどちらが小野町で"
         "当たるかを確かめるには厚生労働省の保険者別配布シートが要ると"
         "申し上げてきましたが、**このボタンで確認できる可能性があります。**",
         "**恐れ入りますが、2つのボタンを押していただき、"
         "表示された内容をお知らせください。**"
         "国が標準の推計方法を当てはめた誤差（推計値÷実績値）が"
         "示されていれば、推計手法の選択の妥当性がその場で確定します。")

    head(doc, "7　町にお願いしたいこと", 1)
    table(doc, ["#", "内容", "優先度"], [
        ("1", "**「推計パターン毎の乖離状況」の2つのボタンの表示内容**"
              "（第6章）", "A"),
        ("2", "**確定版データの公表後の再出力。**"
              "現在の表示は「暫定版データ」である", "A"),
        ("3", "令和8年度の月報が6月分以降も公表され次第、"
              "設定画面でチェックを追加したうえでの再出力", "A"),
        ("4", "第9期策定時の起案文書（推計の設定が分かるもの）", "B"),
    ])

    OUT1.mkdir(parents=True, exist_ok=True)
    p = OUT1 / f"小野町_見える化システムの設定の確認結果_{ASOF}.docx"
    doc.save(p)
    print("出力:", p)


# ---------------------------------------------------------------- Excel

def style_head(ws, row):
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row, c)
        cell.fill = HEAD
        cell.font = Font(bold=True, color="FFFFFF", size=9, name=JP_GO)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = BORDER


def xbody(ws, first, wrap=()):
    for r in range(first + 1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.font = Font(size=9, name=JP_MIN)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=c in wrap)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def notes(ws, lines):
    ws.append([])
    for t in lines:
        ws.append([t])
        ws.cell(ws.max_row, 1).font = Font(size=9, italic=True, name=JP_MIN)


def build_xlsx(K):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("00_小野町の得点")
    ws.append(["小野町の交付金の得点　（令和6年度交付金＝令和5年の取組の評価）"])
    ws.cell(1, 1).font = Font(bold=True, size=13, name=JP_GO)
    ws.append([f"出所：地域包括ケア「見える化」システム W126〜W137"
               f"（令和8年9月25日出力）"])
    ws.append([])
    ws.append(["区分", "項目", "小野町", "全国", "福島県", "全国比"])
    hr = ws.max_row
    for lab, d in K["総合"]:
        t, z, ken = d.get("小野町"), d.get("全国"), d.get("福島県")
        ws.append(["総合", lab, t, z, ken, (t / z) if z else None])
    for kbn in ("推進", "支援"):
        for lab, d in K["目標"][kbn]:
            t, z, ken = d.get("小野町"), d.get("全国"), d.get("福島県")
            ws.append([f"{kbn}・目標別", lab, t, z, ken, (t / z) if z else None])
        for lab, d in K["指標群"][kbn]:
            t, z, ken = d.get("小野町"), d.get("全国"), d.get("福島県")
            ws.append([f"{kbn}・指標群別", lab, t, z, ken,
                       (t / z) if z else None])
    style_head(ws, hr)
    for r in range(hr + 1, ws.max_row + 1):
        for c in (3, 4, 5):
            ws.cell(r, c).number_format = "#,##0.0"
        ws.cell(r, 6).number_format = "0.0%"
        v = ws.cell(r, 6).value
        if isinstance(v, (int, float)):
            ws.cell(r, 6).fill = NG if v < 0.7 else WARN if v < 1.0 else OK
    xbody(ws, hr, wrap=(2,))
    widths(ws, [16, 46, 10, 10, 10, 10])
    notes(ws, [
        "**※ 見える化の「令和5年(2023年)」は令和6年度交付金にあたる**"
        "（令和5年の取組を評価したもの）。全国平均422.4点は"
        "令和6年度交付金の全国平均と一致する。",
        "**※ 令和7・8年度の市町村分は見える化にまだ入っていない。**",
        "※ 赤＝全国の7割未満／橙＝全国未満／緑＝全国以上。",
    ])

    ws = wb.create_sheet("01_0点の項目")
    ws.append(["小野町が0点で、全国平均が高い項目"])
    ws.cell(1, 1).font = Font(bold=True, size=12, name=JP_GO)
    ws.append([])
    ws.append(["指標ID", "評価項目", "小野町", "全国平均", "福島県平均"])
    hr = ws.max_row
    for k, lab, t, z, ken in K["0点"]:
        ws.append([k, lab, t, z, ken])
    style_head(ws, hr)
    for r in range(hr + 1, ws.max_row + 1):
        for c in (3, 4, 5):
            ws.cell(r, c).number_format = "#,##0.0"
        ws.cell(r, 3).fill = NG
    xbody(ws, hr, wrap=(2,))
    widths(ws, [10, 54, 10, 12, 12])
    notes(ws, [
        "**※ 全国平均が高い順。上にあるものほど取りこぼしが大きい。**",
        "※ 交付金の評価指標は、計画に書けば翌年度の取組として得点になるもの"
        "（体制・取組）と、実績の積上げを要するもの（活動）がある。",
    ])

    ws = wb.create_sheet("02_全国を上回る項目")
    ws.append(["小野町が全国平均を上回る項目"])
    ws.cell(1, 1).font = Font(bold=True, size=12, name=JP_GO)
    ws.append([])
    ws.append(["指標ID", "評価項目", "小野町", "全国平均", "福島県平均", "差"])
    hr = ws.max_row
    for k, lab, t, z, ken in K["上回る"]:
        ws.append([k, lab, t, z, ken, t - z])
    style_head(ws, hr)
    for r in range(hr + 1, ws.max_row + 1):
        for c in (3, 4, 5, 6):
            ws.cell(r, c).number_format = "#,##0.0"
        ws.cell(r, 3).fill = OK
    xbody(ws, hr, wrap=(2,))
    widths(ws, [10, 54, 10, 12, 12, 10])

    ws = wb.create_sheet("03_評価項目の全明細")
    ws.append(["指標ID", "評価項目", "小野町", "全国平均", "福島県平均",
               "全国比"])
    hr = ws.max_row
    for k, lab, t, z, ken in K["明細"]:
        ws.append([k, lab, t, z, ken, (t / z) if z else None])
    style_head(ws, hr)
    for r in range(hr + 1, ws.max_row + 1):
        for c in (3, 4, 5):
            ws.cell(r, c).number_format = "#,##0.0"
        ws.cell(r, 6).number_format = "0.0%"
    xbody(ws, hr, wrap=(2,))
    widths(ws, [10, 54, 10, 12, 12, 10])
    ws.freeze_panes = "A2"

    OUT2.mkdir(parents=True, exist_ok=True)
    p = OUT2 / f"小野町_交付金_小野町の得点分析_{ASOF}.xlsx"
    wb.save(p)
    print("出力:", p)


def main():
    K = kofukin()
    build_docx(K)
    build_xlsx(K)
    so = {lab: d for lab, d in K["総合"]}
    g = so.get("推進・支援合計", {})
    print(f"  小野町の交付金 総合{g.get('小野町')}点／"
          f"全国{g.get('全国')}／福島県{g.get('福島県')}"
          f"（全国比{g.get('小野町') / g.get('全国') * 100:.1f}％）")
    print(f"  0点の項目 {len(K['0点'])}件／全国を上回る項目 {len(K['上回る'])}件")


if __name__ == "__main__":
    main()
